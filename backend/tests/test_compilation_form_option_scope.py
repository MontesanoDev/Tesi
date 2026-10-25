"""A valid SOURCE value cannot silently choose a printed alternative (offline API)."""
import copy
import json
from io import BytesIO

import pytest
import test_compilation_resilient_segments as rev2
from docx import Document
from test_compilation_global_sessions import api, provider

from app import compilation_global_dialogue as dialogue
from app import compilation_profile

__all__ = ["api", "provider"]
OPTION = "[] utilizzare i dati del soggetto delegato:"


def form():
    doc = Document()
    doc.add_paragraph("[ ] utilizzare i dati del richiedente;")
    doc.add_paragraph("OPPURE")
    doc.add_paragraph(OPTION)
    doc.add_paragraph("")
    doc.add_paragraph("Dato del delegato: ____________")
    doc.add_paragraph("")
    doc.add_paragraph("Dato indipendente: ____________")
    stream = BytesIO()
    doc.save(stream)
    return stream.getvalue()


def activate(monkeypatch, enabled):
    monkeypatch.setenv("MAPI_GENERATION_REVISION", "3")
    monkeypatch.setenv("MAPI_DIALOGUE_REVISION", "1")
    original = compilation_profile.new_session_flags
    monkeypatch.setattr(compilation_profile, "new_session_flags", lambda: {
        **original(), **({"form_option_scope_revision": 1} if enabled else {})})


def writer(inp, ids):
    return json.dumps(rev2.envelope(inp, {c: rev2.slot(inp, c)
                                        for c in rev2.dato_ids(inp, ids)}))


@pytest.mark.anyio
@pytest.mark.parametrize("enabled", [False, True])
async def test_valid_source_and_json_do_not_choose_an_omitted_option(
    api, provider, monkeypatch, enabled,
):
    activate(monkeypatch, enabled)
    s = await rev2.started(api, monkeypatch, form(), rev2.Model({0: [writer], 1: [writer]}))
    response = await rev2.analyze(api, s)
    assert response.status_code == 200, response.text
    s = response.json()
    selected, _, _ = dialogue.export_selection(s)
    assert "p6.s0" in selected  # Independent, documented data.
    export = await rev2.draft(api, s, ["p4.s0", "p6.s0"])
    if enabled:
        assert export.status_code == 409, export.text
        assert "p4.s0" not in selected
        field = next(f for f in s["global_compilation"]["fields"]
                     if f["candidate_id"] == "p4.s0")
        assert any("opzione FORM non accertata" in issue for issue in field["issues"])
        assert (await rev2.draft(api, s, selected)).status_code == 200
    else:
        assert export.status_code == 200, export.text


def state(*, condition=True):
    catalog = rev2.model_form(form())
    return {"generation_revision": 3, "form_option_scope_revision": 1, "catalog": catalog,
            "dialogue": {"answers": {}}, "audit": {"fields": [
                {"candidate_id": cid, "conditions": [], "validation": "GROUNDED_PROPOSAL"}
                for cid in catalog["candidates"]], "conditions": {"c": {
                    "form_evidence": {"paragraph_id": "paragraph:2", "quote": OPTION},
                    "operator": "semantic", "result": "unknown", "applies_to": []}}
                if condition else {}}}


def confirm(data, *, key="condition:c", action="CONDITION_TRUE", ids=None):
    from app import compilation_scope_review as scope

    ids = ["p4.s0"] if ids is None else ids
    record = {"action": action, "applies": action == "CONDITION_TRUE",
              "user_quote": "Scelgo il soggetto delegato", "affected_fields": ids,
              "provenance": "USER"}
    record["scope_validation"] = scope.proof(scope.target(data, key, ids), record)
    data["dialogue"]["answers"][key] = record


def projected(data):
    from app import compilation_form_option_scope as options

    data["audit"] = options.apply(data["audit"], data["catalog"])
    return data


def test_the_condition_is_propagated_without_becoming_a_source_fact_or_user_decision():
    from app import compilation_scope_review as scope
    from app.compilation_segments import applicability_states

    data = projected(state())
    assert data["audit"]["conditions"]["c"]["result"] == "unknown"
    assert data["audit"]["conditions"]["c"]["applies_to"] == ["p4.s0"]
    assert data["dialogue"]["answers"] == {}
    assert scope.blocked_source_fields(data) == {"p4.s0"}
    assert applicability_states(data)["p4.s0"]["state"] == "CONDIZIONATA"
    confirm(data)
    assert not scope.blocked_source_fields(data)
    assert applicability_states(data)["p4.s0"]["state"] == "VERIFICATA"
    data["dialogue"]["answers"].clear()  # Revocation.
    assert scope.blocked_source_fields(data) == {"p4.s0"}


@pytest.mark.anyio
@pytest.mark.parametrize("created_with_opt_in", [False, True])
async def test_real_creation_profile_is_persisted_across_an_environment_change(
    api, provider, monkeypatch, created_with_opt_in,
):
    monkeypatch.delenv("MAPI_COMPILATION_PROFILE", raising=False)
    if created_with_opt_in:
        monkeypatch.setenv("MAPI_FORM_OPTION_SCOPE_REVISION", "1")
    else:
        monkeypatch.delenv("MAPI_FORM_OPTION_SCOPE_REVISION", raising=False)
    s = await rev2.started(api, monkeypatch, form(), rev2.Model({0: [writer], 1: [writer]}))
    initial_flags = {k: v for k, v in s["global_compilation"].items() if k.endswith("revision")}
    if created_with_opt_in:
        monkeypatch.delenv("MAPI_FORM_OPTION_SCOPE_REVISION", raising=False)
    else:
        monkeypatch.setenv("MAPI_FORM_OPTION_SCOPE_REVISION", "1")
    response = await rev2.analyze(api, s)
    assert response.status_code == 200, response.text
    data = rev2.state_of(response.json())["global_compilation"]
    assert {k: v for k, v in data.items() if k.endswith("revision")} == initial_flags
    selected, _, _ = dialogue.export_selection(response.json())
    assert ("p4.s0" not in selected) == created_with_opt_in


@pytest.mark.parametrize("defect", ["missing", "foreign", "partial", "invalid", "detached"])
def test_valid_json_or_other_conditions_do_not_cover_the_required_option(defect):
    from app import compilation_form_option_scope as options
    from app import compilation_scope_review as scope
    from app.compilation_segments import applicability_states

    data = state(condition=defect != "missing")
    if defect not in {"missing", "detached"}:
        condition = data["audit"]["conditions"]["c"]
        if defect == "foreign":
            condition["form_evidence"]["paragraph_id"] = "paragraph:0"
        elif defect == "partial":
            condition["form_evidence"]["quote"] = "soggetto delegato"
        else:
            condition["validation_errors"] = ["Invalid FORM citation"]
    if defect != "detached":
        projected(data)
    confirm(data)  # Even a simulated SUPPORTED receipt does not fix an omitted condition.
    assert options.uncovered_fields(data) == {"p4.s0"}
    assert scope.blocked_source_fields(data) == {"p4.s0"}
    assert "p4.s0" in scope.blocked_user_fields(data)
    assert applicability_states(data)["p4.s0"]["state"] == "NON_ACCERTATA"
    if defect == "missing":
        assert data["audit"]["conditions"] == {}


@pytest.mark.parametrize("change", ["negative", "skip", "foreign", "stale", "other_fields",
                                   "source_only", "source_false", "source_reuse"])
def test_a_decision_only_authorizes_its_verified_current_branch(change):
    from app import compilation_scope_review as scope

    data = projected(state())
    if change in {"negative", "skip"}:
        confirm(data, action="CONDITION_FALSE" if change == "negative" else "SKIP")
    elif change == "foreign":
        confirm(data, key="condition:other")
    elif change == "other_fields":
        confirm(data, ids=["p6.s0"])
    elif change == "source_only":
        data["audit"]["conditions"]["c"]["result"] = "true"
    elif change == "source_reuse":
        data["audit"]["fields"][0]["source_reuse"] = {"origin": "SOURCE"}
    else:
        confirm(data)
        if change == "source_false":
            data["audit"]["conditions"]["c"]["result"] = "false"
        else:
            data["audit"]["conditions"]["c"]["reason"] = "Changed meaning"
    assert scope.blocked_source_fields(data) == {"p4.s0"}


def test_a_user_value_receipt_is_not_a_choice_of_the_omitted_option():
    from app import compilation_scope_review as scope

    data = projected(state(condition=False))
    data["dialogue"]["user_values"] = {"p4.s0": {"value": "ABC123", "user_quote": "ABC123"}}
    receipt = {"action": "VALUE", "value": "ABC123", "user_quote": "ABC123",
               "provenance": "USER", "affected_fields": ["p4.s0"]}
    receipt["scope_validation"] = scope.proof(scope.target(data, "field:p4.s0", ["p4.s0"]),
                                             receipt)
    data["dialogue"]["answers"]["field:p4.s0"] = receipt
    assert "p4.s0" in scope.blocked_user_fields(data)


def test_projection_is_repeatable_and_cannot_grant_permission_when_removed():
    from app import compilation_form_option_scope as options

    data = projected(state())
    original = copy.deepcopy(data)
    projected(data)
    assert data == original
    data["audit"].pop("form_option_scope")
    assert options.blocked_fields(data) == {"p4.s0"}
    data["audit"]["conditions"].clear()
    assert options.uncovered_fields(data) == {"p4.s0"}


@pytest.mark.parametrize("change", ["table", "text_between", "inline", "example_prefix",
                                   "checked", "no_colon", "no_pair", "target_heading",
                                   "empty_control", "editable_option"])
def test_ambiguous_or_interrupted_structure_does_not_create_a_dependency(change):
    from app import compilation_form_option_scope as options

    catalog = state()["catalog"]
    blocks = catalog["blocks"]
    option = next(p for b in blocks for p in b["paragraphs"] if p["id"] == "paragraph:2")
    target = next(p for b in blocks for p in b["paragraphs"] if p["id"] == "paragraph:4")
    blank = next(b for b in blocks if b["id"] == "paragraph:3")
    if change == "table":
        blank["tables"] = [{"id": "table:99", "rows": []}]
    elif change == "text_between":
        blank["paragraphs"][0]["text"] = "Dati generali indipendenti"
    elif change == "inline":
        option["text"] += " [ ] altro:"
    elif change == "example_prefix":
        option["text"] = "Esempio: " + option["text"]
    elif change == "checked":
        option["text"] = option["text"].replace("[]", "[X]")
    elif change == "no_colon":
        option["text"] = option["text"].rstrip(":")
    elif change == "no_pair":
        blocks[0]["paragraphs"][0]["text"] = "Testo descrittivo"
    elif change == "target_heading":
        target["all_text_direct_bold"] = True
    elif change == "empty_control":
        blank["paragraphs"][0]["controls"] = [{"id": "c", "offset": 0}]
    else:
        catalog["candidates"]["p2.s0"] = {"context_id": "paragraph:2", "type": "text"}
    assert options.regions(catalog) == []


def test_only_the_first_data_paragraph_is_owned_and_quote_can_omit_the_box_marker():
    from app import compilation_form_option_scope as options

    data = state()
    data["audit"]["conditions"]["c"]["form_evidence"]["quote"] = OPTION.removeprefix("[]")
    projected(data)
    assert options.regions(data["catalog"])[0]["candidate_ids"] == ["p4.s0"]
    confirm(data)
    assert options.blocked_fields(data) == set()


def test_creation_requires_explicit_opt_in_and_does_not_upgrade_existing_sessions(monkeypatch):
    from app import compilation_form_option_scope as options

    monkeypatch.delenv("MAPI_COMPILATION_PROFILE", raising=False)
    monkeypatch.delenv("MAPI_FORM_OPTION_SCOPE_REVISION", raising=False)
    old = compilation_profile.new_session_flags()
    monkeypatch.setenv("MAPI_FORM_OPTION_SCOPE_REVISION", "1")
    new = compilation_profile.new_session_flags()
    assert "form_option_scope_revision" not in old
    assert not options.enabled(old)
    assert options.enabled(new)
    assert new["compilation_profile"] != old["compilation_profile"]
    assert {k: new[k] for k in old if k != "compilation_profile"} == {
        k: v for k, v in old.items() if k != "compilation_profile"}
    monkeypatch.setenv("MAPI_COMPILATION_PROFILE", "storico")
    assert not options.enabled(compilation_profile.new_session_flags())
    assert not options.enabled({**new, "generation_revision": 2})


@pytest.mark.anyio
async def test_export_after_the_own_condition_is_verified_and_revocation_blocks_again(
    api, provider, monkeypatch,
):
    from app import compilation_scope_review as scope
    from app.compilation_chat import ClarificationReply

    activate(monkeypatch, True)

    def with_condition(inp, ids):
        raw = json.loads(writer(inp, ids))
        raw["conditions"]["c"] = rev2.condition(inp, "paragraph:2", OPTION)
        return json.dumps(raw)

    s = await rev2.started(api, monkeypatch, form(),
                          rev2.Model({0: [with_condition], 1: [with_condition]}))
    response = await rev2.analyze(api, s)
    assert response.status_code == 200, response.text
    s = response.json()
    data = rev2.state_of(s)["global_compilation"]
    target = next(q for q in dialogue.question_pool(data) if q["key"].startswith("condition:"))
    assert target["field_ids"] == ["p4.s0"]
    question = next(q for q in s["chat"]["question"]["slots"] if q["key"] == target["key"])
    reply = ClarificationReply(slot=question["slot"], action="CONDITION_TRUE",
                               user_quote="Scelgo il soggetto delegato")
    proof = scope.proof(scope.target(data, target["key"], target["field_ids"]), reply)
    s, errors = dialogue.apply_replies("alpha", s["id"], s["version"], [reply], reply.user_quote,
                                       scope_verifications={reply.slot: proof})
    assert not errors
    s = (await rev2.analyze(api, s)).json()
    ids, conditions, _ = dialogue.export_selection(s)
    assert set(ids) == {"p4.s0", "p6.s0"}
    exported = await rev2.draft(api, s, ids, conditions=conditions)
    assert exported.status_code == 200, exported.text
    data = rev2.state_of(exported.json())["global_compilation"]
    data["dialogue"]["answers"].clear()
    assert scope.blocked_source_fields(data) == {"p4.s0"}
