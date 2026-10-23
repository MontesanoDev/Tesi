"""FORM applicability is independent of a true SOURCE value and MODEL review."""

import hashlib
import json
from copy import deepcopy
from io import BytesIO
from pathlib import Path

import pytest
from docx import Document

from app.compilation_global import build_global_input, validate_global_proposal
from app.compilation_global_dialogue import export_selection
from app.compilation_global_sessions import initial_state, review_fields
from app.compilation_source_reuse import reuse_selected_person
from app.docx_templates import fill_docx, inspect_docx

FIXTURE = Path(__file__).parent / "fixtures/form-applicability-v1"


def load(name):
    return json.loads((FIXTURE / name).read_text())


def frozen(variant):
    inp = load("input.json")
    layout = inspect_docx((FIXTURE / "original.docx").read_bytes())
    data = {**initial_state(layout), **load(f"{variant}-3-state.json"), "catalog": inp["form"]}
    data["audit"] = validate_global_proposal((FIXTURE / f"{variant}-3-response.txt").read_text(),
                                            layout, inp["project_id"], inp["sources"])
    data["audit"] = reuse_selected_person(data, layout, inp["project_id"], inp["sources"])
    data["fields"] = review_fields(data)
    return layout, data


@pytest.mark.parametrize("variant", ["B0", "E1"])
def test_exact_original_counterexamples_keep_true_fact_but_do_not_write_inactive_branch(variant):
    layout, data = frozen(variant)
    field = next(f for f in data["audit"]["fields"] if f["candidate_id"] == "t0.r8.c1")
    assert field["validation"] == "GROUNDED_PROPOSAL"
    assert field["conditions"] == []  # The original LLM omission is replayed unchanged.
    ids, _, users = export_selection({"global_compilation": data})
    assert "t0.r8.c1" not in ids + users
    expected = set(load(f"{variant}-3-writes.json")["text_values"]) - {"t0.r8.c1"}
    # Independent completeness now quarantines the ambiguous combined contact
    # label; these are still the exact original MODEL responses and SOURCE facts.
    expected -= {"t0.r11.c1", "t25.r3.c1"}
    assert set(ids) == expected
    values = {f["candidate_id"]: f.get("proposed_value", f.get("value"))
              for f in data["audit"]["fields"] if f["candidate_id"] in ids}
    written = Document(BytesIO(fill_docx(layout, values)))
    assert written.tables[0].cell(8, 1).text == ""
    assert "Via Giovanni Amendola 172/C" in written.tables[0].cell(7, 1).text


def test_original_counterexample_fixture_has_frozen_content_hashes():
    for name, digest in load("manifest.json")["files"].items():
        assert hashlib.sha256((FIXTURE / name).read_bytes()).hexdigest() == digest


def test_source_and_original_proposal_are_not_rewritten_by_export():
    _, data = frozen("B0")
    audit = deepcopy(data["audit"])
    export_selection({"global_compilation": data})
    assert audit == data["audit"]


def prepared(text="Sede operativa (se diversa dalla sede legale) ________", *,
             reference="Via Uno 1", extra="", subject="Aurora", layout_kind="paragraph",
             current_label="Sede operativa", reference_label="Sede legale"):
    doc = Document()
    if layout_kind == "paragraph":
        doc.add_paragraph(text)
    else:
        table = doc.add_table(rows=1, cols=2)
        parts = text.replace("________", "").strip().split("\n")
        table.cell(0, 0).text = parts[0]
        for part in parts[1:]:
            table.cell(0, 0).add_paragraph(part)
        if layout_kind == "reordered":
            doc.add_paragraph("Altra sezione")
            doc.add_table(rows=1, cols=3)
        if layout_kind == "merged":
            from docx.oxml import OxmlElement

            table.cell(0, 0)._tc.get_or_add_tcPr().append(OxmlElement("w:vMerge"))
    stream = BytesIO()
    doc.save(stream)
    layout = inspect_docx(stream.getvalue())
    quote = f"{current_label} di {subject}: Via Uno 1."
    content = quote
    if reference is not None:
        content += f"\n{reference_label} di {subject}: {reference}."
    content += "\n" + extra
    sources = [{"chunk_id": -1, "file_id": -1, "chunk_index": 0, "role": "source",
                "scope": "global", "project_id": None, "category": "company",
                "source_name": "dati.txt", "content": content}]
    catalog = build_global_input(layout, "alpha", sources)["form"]
    paragraphs = {p["id"]: p for b in catalog["blocks"] for p in b["paragraphs"]}
    raw = {"version": 12, "uncertain_paragraph_ids": [], "coverage": [], "choice_groups": {},
           "clarifications": {}, "conditions": {}, "bindings": {}, "fields": {},
           "facts": {"f": {"entity": "company", "subject": subject, "aspect": "location",
                    "property": "sede_operativa", "value": "Via Uno 1",
                    "evidence": [{"chunk_id": -1, "quote": quote}]}}}
    for cid, candidate in catalog["candidates"].items():
        if candidate["type"] != "text":
            continue
        anchors = [p for p in candidate["anchor_paragraph_ids"] if paragraphs[p]["text"].strip()]
        if not anchors:
            continue
        pid = anchors[0]
        marked = paragraphs[pid].get("text_with_fields", paragraphs[pid]["text"])
        if ";" in marked:
            marked = next(c.strip() for c in marked.split(";") if f"[[{cid}]]" in c)
        raw["fields"][cid] = {"label": "Dato", "kind": "data", "disposition": "value",
                              "conditions": [], "property": "sede_operativa", "fact_id": "f",
                              "form": {"paragraph_id": pid,
                                       "quote": marked}}
    return layout, sources, raw, catalog


def state(case):
    layout, sources, raw, catalog = case
    data = {**initial_state(layout), "catalog": catalog, "artifact_path": "alpha/session/attempt",
            "audit": validate_global_proposal(json.dumps(raw), layout, "alpha", sources),
            "dialogue": {"phase": "READY", "answers": {}, "user_values": {},
                         "reviewed_candidate_ids": list(catalog["candidates"])}}
    data["fields"] = review_fields(data)
    return data


def check(data, cid="p0.s0"):
    return data["audit"]["form_applicability"][cid]


@pytest.mark.parametrize("layout_kind", ["paragraph", "table", "reordered"])
@pytest.mark.parametrize("condition,expected", [
    ("se diversa dalla sede legale", "false"),
    ("se non diversa dalla sede legale", "true"),
    ("se coincidente con la sede legale", "true"),
    ("se non coincidente con la sede legale", "false"),
    ("solo se differente dalla sede legale", "false"),
    ("qualora identica alla sede legale", "true"),
    ("quando uguale alla sede legale", "true"),
    ("se non è distinta dalla sede legale", "true"),
    ("se la Sede operativa è diversa dalla sede legale", "false"),
])
def test_equivalent_local_contexts_and_negations_keep_three_valued_meaning(
    layout_kind, condition, expected,
):
    text = f"Sede operativa\n({condition})" if layout_kind != "paragraph" else (
        f"Sede operativa ({condition}) ________")
    data = state(prepared(text, layout_kind=layout_kind))
    cid = "p0.s0" if layout_kind == "paragraph" else "t0.r0.c1"
    result = check(data, cid)
    assert result["result"] == expected
    ids, _, _ = export_selection({"global_compilation": data})
    assert (cid in ids) == (expected == "true")
    assert result["subject"] == "Aurora"
    assert all(p["role"] == "source" for p in result["source_evidence"])
    assert all(p["role"] == "form" for c in result["contexts"] for p in c["form_evidence"])


@pytest.mark.parametrize("text", [
    "Sede operativa ________ (se diversa dalla sede legale)",
    "Sede operativa (compilare solo se diversa dalla sede legale) ________",
    "Sede operativa: ________ qualora diversa dalla sede legale",
])
def test_condition_before_or_after_the_exact_marker_cannot_be_omitted(text):
    data = state(prepared(text))
    assert check(data)["result"] == "false"
    assert export_selection({"global_compilation": data})[0] == []


@pytest.mark.parametrize("relation,expected", [
    ("diversa dalla sede legale", "true"), ("non coincidente con la sede legale", "true"),
    ("non uguale alla sede legale", "true"), ("coincidente con la sede legale", "false"),
    ("non diversa dalla sede legale", "false"),
])
def test_explicit_subject_bound_source_relations_prove_true_or_false(relation, expected):
    data = state(prepared(reference=None, extra=f"Sede operativa di Aurora: {relation}."))
    assert check(data)["result"] == expected
    assert check(data)["basis"] == "explicit_source_relation"
    assert bool(export_selection({"global_compilation": data})[0]) == (expected == "true")


@pytest.mark.parametrize("reference,extra", [
    (None, ""), ("Via Due 2", ""),  # Different strings are not semantic inequality.
    (None, "Sede legale di Altra Società: Via Uno 1."),
    (None, "Sede operativa di Altra Società: diversa dalla sede legale."),
    (None, "Se la sede operativa di Aurora è diversa dalla sede legale."),
    (None, "Non è certo che la sede operativa di Aurora sia diversa dalla sede legale."),
    (None, "Sede operativa di Aurora: forse diversa dalla sede legale."),
    (None, "Sede operativa di Aurora: non è noto se coincidente con la sede legale."),
    ("non disponibile", ""), ("forse Via Uno 1", ""),
    ("Via Uno 1", "Sede operativa di Aurora: diversa dalla sede legale."),
    ("Via Uno 1", "Sede legale di Aurora: Via Due 2."),
    (None, "Sede operativa di Aurora: diversa dalla sede legale.\n"
            "Sede operativa di Aurora: coincidente con la sede legale."),
])
def test_missing_unproved_negated_hypothetical_or_conflicting_evidence_stays_unknown(
    reference, extra,
):
    data = state(prepared(reference=reference, extra=extra))
    assert check(data)["result"] == "unknown"
    assert export_selection({"global_compilation": data})[0] == []


@pytest.mark.parametrize("left,right", [
    ("Recapito secondario", "Recapito principale"),
    ("Indirizzo di fatturazione", "Indirizzo di consegna"),
    ("Domicilio eletto", "Domicilio professionale"),
])
def test_comparison_grammar_does_not_depend_on_office_or_benchmark_names(left, right):
    data = state(prepared(f"{left} (se diverso dal {right}) ________",
                          current_label=left, reference_label=right))
    assert check(data)["result"] == "false"
    assert check(data)["contexts"][0]["predicate"]["left"] == left.casefold()


@pytest.mark.parametrize("condition", [
    "se diversa e comunque autorizzata dalla sede legale",
    "se diversa dalla sede legale o se richiesta dal committente",
    "se non non diversa dalla sede legale",
])
def test_unrecognised_comparisons_are_unknown_not_unconditional(condition):
    data = state(prepared(f"Sede operativa ({condition}) ________"))
    assert check(data)["result"] == "unknown"
    assert export_selection({"global_compilation": data})[0] == []


def test_true_and_false_alternative_branches_are_selected_independently():
    data = state(prepared("Sede operativa (se diversa dalla sede legale) ________; "
                          "Sede operativa (se coincidente con la sede legale) ________"))
    assert check(data, "p0.s0")["result"] == "false"
    assert check(data, "p0.s1")["result"] == "true"
    assert export_selection({"global_compilation": data})[0] == ["p0.s1"]


@pytest.mark.parametrize("relation,reference,expected", [
    ("", "Via Uno 1", ["p0.s1"]),
    ("Sede operativa di Aurora: diversa dalla sede legale.", None, ["p0.s0"]),
    ("", None, []),
])
def test_explicit_else_is_the_complement_of_the_same_local_predicate(relation, reference, expected):
    data = state(prepared("Sede operativa (se diversa dalla sede legale) ________; "
                          "altrimenti Sede operativa ________",
                          reference=reference, extra=relation))
    assert export_selection({"global_compilation": data})[0] == expected
    assert check(data, "p0.s1")["contexts"][0]["alternative_of"]


def test_else_without_a_structurally_demonstrated_parent_stays_unknown():
    data = state(prepared("altrimenti Sede operativa ________"))
    assert check(data)["result"] == "unknown"
    assert export_selection({"global_compilation": data})[0] == []


def test_predicate_does_not_spread_to_adjacent_row_or_paragraph_or_other_table():
    doc = Document()
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Sede operativa (se diversa dalla sede legale)"
    table.cell(1, 0).text = "Sede operativa"
    doc.add_paragraph("Sede operativa ________")
    doc.add_table(rows=1, cols=2).cell(0, 0).text = "Sede operativa"
    from app.compilation_catalog import build_compilation_catalog
    from app.compilation_form_applicability import requirements

    stream = BytesIO()
    doc.save(stream)
    found = requirements(build_compilation_catalog(inspect_docx(stream.getvalue())))
    assert set(found) == {"t0.r0.c1"}


@pytest.mark.parametrize("result", ["false", "unknown"])
def test_source_user_value_export_assent_or_model_condition_cannot_override_form_veto(result):
    case = prepared(reference="Via Uno 1" if result == "false" else None)
    data = state(case)
    audit = deepcopy(data["audit"])
    data["dialogue"]["answers"] = {"condition:invented": {"applies": True,
        "action": "CONDITION_TRUE", "provenance": "USER", "user_quote": "Genera la bozza"}}
    data["dialogue"]["user_values"] = {"p0.s0": {"value": "Via Uno 1", "provenance": "USER",
                                                "user_quote": "Via Uno 1"}}
    # Even an old or stale public selectable flag is insufficient.
    data["fields"][0]["selectable"] = True
    assert export_selection({"global_compilation": data})[::2] == ([], [])
    public = {k: v for k, v in data.items() if k not in {"audit", "catalog", "writable_ids"}}
    assert export_selection({"global_compilation": public})[::2] == ([], [])
    assert data["audit"] == audit


def test_true_form_predicate_does_not_supply_procedural_user_authorization():
    from app.compilation_procedural_consent import requirements

    data = state(prepared("Sede operativa (se coincidente con la sede legale) ________, "
                          "che autorizzo per le comunicazioni"))
    assert check(data)["result"] == "true"
    assert export_selection({"global_compilation": data})[0] == []
    requirement = requirements(data)[0]
    data["dialogue"]["answers"][requirement["key"]] = {
        "provenance": "USER", "action": "CONDITION_TRUE", "applies": True,
        "user_quote": "Sì, autorizzo questa dichiarazione", **requirement}
    assert export_selection({"global_compilation": data})[0] == ["p0.s0"]
    assert check(data)["USER_authorization_used"] is False


def test_existing_v12_without_new_check_requires_reanalysis_without_mutating_session():
    data = state(prepared())
    del data["audit"]["form_applicability"]
    before = deepcopy(data)
    fields = review_fields(data)
    assert fields[0]["form_applicability"]["result"] == "unknown"
    assert fields[0]["selectable"] is False
    assert data == before


def test_form_catalog_and_model_request_are_unchanged_by_the_validator():
    case = prepared()
    layout, sources, raw, _ = case
    before = build_global_input(layout, "alpha", sources)
    copied = deepcopy(before)
    validate_global_proposal(json.dumps(raw), layout, "alpha", sources)
    assert before == copied == build_global_input(layout, "alpha", sources)


@pytest.mark.parametrize("prefix,suffix", [
    ("## Esempio ipotetico\n\n", ""),
    ("Se fosse verificata questa situazione:\n", ""),
    ("Le seguenti affermazioni non sono confermate:\n", ""),
    ("```\n", "\n```"),
])
def test_literal_assertions_cannot_drop_a_source_hypothesis_or_example_context(prefix, suffix):
    case = prepared()
    case[1][0]["content"] = prefix + case[1][0]["content"] + suffix
    data = state(case)
    assert check(data)["result"] == "unknown"
    assert export_selection({"global_compilation": data})[0] == []


def test_a_true_predicate_does_not_relax_the_value_property_grounding():
    case = prepared("Sede operativa (se coincidente con la sede legale) ________")
    case[2]["fields"]["p0.s0"]["property"] = "incarico"
    data = state(case)
    assert check(data)["result"] == "true"
    assert data["audit"]["fields"][0]["validation"] == "REJECTED"
    assert export_selection({"global_compilation": data})[0] == []


def test_recorded_veto_survives_an_incomplete_cached_catalog():
    from app.compilation_form_applicability import blocked_fields

    data = state(prepared())
    del data["catalog"]["candidates"]["p0.s0"]["type"]
    assert blocked_fields(data) == {"p0.s0"}


def test_ambiguous_merge_keeps_the_local_comparison_unknown():
    data = state(prepared(layout_kind="merged"))
    assert check(data, "t0.r0.c1")["result"] == "unknown"
    assert export_selection({"global_compilation": data})[0] == []


def test_multiple_destinations_in_one_undivided_comparative_clause_are_unknown():
    data = state(prepared("Sede operativa (se diversa dalla sede legale) "
                          "________ recapito ________"))
    assert all(c["result"] == "unknown" for c in data["audit"]["form_applicability"].values())
    assert export_selection({"global_compilation": data})[0] == []
