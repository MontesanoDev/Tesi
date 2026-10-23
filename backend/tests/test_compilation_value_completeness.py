"""Backend compound requirements; frozen original and independent negative gates."""

import hashlib
import json
from copy import deepcopy
from io import BytesIO
from pathlib import Path

import pytest
from docx import Document
from test_docx_formtext import form_text

from app.compilation_branch_evidence import reconcile_branches
from app.compilation_global import build_global_input, validate_global_proposal
from app.compilation_global_dialogue import export_selection
from app.compilation_global_sessions import initial_state, review_fields
from app.compilation_source_reuse import reuse_selected_person
from app.compilation_value_completeness import blocked_fields, value_check
from app.docx_templates import fill_docx, inspect_docx

FIXTURE = Path(__file__).parent / "fixtures"


def prepared(label="Numero e data d'iscrizione", value="123", *, style="table", source=None,
             quote=None, entity="company", subject="Aurora", role=None):
    doc = Document()
    if role:
        doc.add_paragraph(role)
    if style == "inline":
        doc.add_paragraph(label + " ________")
    elif style == "formtext":
        form_text(doc.add_paragraph(label + " "))
    elif style == "table":
        doc.add_table(rows=1, cols=2).cell(0, 0).text = label
    elif style == "previous":
        doc.add_table(rows=2, cols=1).cell(0, 0).text = label
    elif style == "header":
        table = doc.add_table(rows=3, cols=2)
        table.cell(0, 0).text, table.cell(0, 1).text = label, "Altro campo"
    stream = BytesIO()
    doc.save(stream)
    layout = inspect_docx(stream.getvalue())
    text = source if source is not None else f"Iscrizione di {subject}: {value}."
    sources = [{"chunk_id": -7, "file_id": -3, "chunk_index": 0, "source_name": "registri.txt",
                "role": "source", "scope": "global", "project_id": None,
                "category": "company", "content": text}]
    inp = build_global_input(layout, "alpha", sources)
    catalog = inp["form"]
    cid = ("t0.r2.c0" if style == "header" else "t0.r1.c0" if style == "previous"
           else next(iter(catalog["candidates"])))
    paragraph = next(p for b in catalog["blocks"] for p in b["paragraphs"]
                     if label in p["text"])
    marked = paragraph.get("text_with_fields", paragraph["text"])
    raw = {"version": 12, "uncertain_paragraph_ids": [], "coverage": [], "choice_groups": {},
        "clarifications": {}, "bindings": {}, "conditions": {},
        "facts": {"f": {"entity": entity, "subject": subject, "aspect": "identity",
            "property": "documented_property", "value": value,
            "evidence": [{"chunk_id": -7, "quote": quote if quote is not None else text}]}},
        "fields": {cid: {"label": "Etichetta MODEL deliberatamente generica", "kind": "data",
            "disposition": "value", "conditions": [], "property": "documented_property",
            "fact_id": "f", "form": {"paragraph_id": paragraph["id"], "quote": marked}}}}
    if role:
        role_paragraph = catalog["blocks"][0]["paragraphs"][0]
        raw["bindings"]["b"] = {"form": {"paragraph_id": role_paragraph["id"], "quote": role},
            "required_role": role, "source_role": role,
            "evidence": [{"chunk_id": -7, "quote": text}]}
        raw["fields"][cid]["binding_id"] = "b"
    return layout, sources, raw, catalog, cid


def state(case):
    layout, sources, raw, catalog, cid = case
    audit = validate_global_proposal(json.dumps(raw), layout, "alpha", sources, expected_version=12)
    data = {**initial_state(layout), "catalog": catalog, "audit": audit,
            "dialogue": {"phase": "READY", "answers": {}, "user_values": {},
                         "reviewed_candidate_ids": [cid]}}
    data["fields"] = review_fields(data)
    return data


def check(data, cid):
    return data["audit"]["value_completeness"][cid]


@pytest.mark.parametrize("style", ["table", "previous", "header", "inline", "formtext"])
@pytest.mark.parametrize("value,expected,missing", [
    ("01234560981", "INCOMPLETE", "date"),
    ("14 marzo 2019", "INCOMPLETE", "number"),
    ("n. 123 del 14 marzo 2019", "COMPLETE", None),
])
def test_required_number_and_date_in_all_supported_physical_contexts(
    style, value, expected, missing,
):
    case = prepared(value=value, style=style)
    data, cid = state(case), case[4]
    assert data["audit"]["facts"]["f"]["validation"] == "GROUNDED_PROPOSAL"
    assert check(data, cid)["result"] == expected
    if missing:
        assert missing in check(data, cid)["missing_components"]
    ids, _, users = export_selection({"global_compilation": data})
    assert (cid in ids) is (expected == "COMPLETE") and not users
    if expected == "COMPLETE":
        assert all(c["evidence"] for c in check(data, cid)["components"].values())
        written = Document(BytesIO(fill_docx(case[0], {cid: value})))
        assert value in written.element.xml


@pytest.mark.parametrize("label", ["numero e data d'iscrizione", "Data e numero d'iscrizione",
                                   "Numero ed anno d'iscrizione"])
def test_true_fact_is_not_complete_just_because_model_calls_it_value(label):
    case = prepared(label=label)
    case[2]["fields"][case[4]]["form"]["quote"] = label.split()[0]
    data = state(case)
    assert data["audit"]["fields"][0]["validation"] == "GROUNDED_PROPOSAL"
    assert check(data, case[4])["result"] == "INCOMPLETE"
    assert not data["fields"][0]["selectable"]


@pytest.mark.parametrize("date", ["14/03/2019", "14/3/2019", "2019-03-14", "14-03-2019",
                                  "14.03.2019", "14 marzo 2019"])
def test_equivalent_complete_date_formats_do_not_change_grounding(date):
    case = prepared(value="n. 123 del " + date)
    data = state(case)
    result = check(data, case[4])
    assert result["result"] == "COMPLETE"
    assert result["components"]["date"]["canonical"] == "2019-03-14"


@pytest.mark.parametrize("value", ["n.123 del 31/02/2019", "n.123 del 14/03/19", "123 del ignoto"])
def test_invalid_or_unproven_date_never_fills_a_required_component(value):
    case = prepared(value=value)
    assert check(state(case), case[4])["result"] != "COMPLETE"


@pytest.mark.parametrize("label,value,expected", [
    ("Numero o data d'iscrizione", "123", "COMPLETE"),
    ("Numero oppure data d'iscrizione", "14 marzo 2019", "COMPLETE"),
    ("Numero/data d'iscrizione", "123", "UNKNOWN"),
    ("Numero + data d'iscrizione", "123", "UNKNOWN"),
    ("Numero e data d'iscrizione (data facoltativa)", "123", "COMPLETE"),
    ("Numero e data d'iscrizione", "n.123 del 14/03/2019 (sez. A)", "COMPLETE"),
])
def test_cumulative_alternative_ambiguous_and_optional_requirements(label, value, expected):
    case = prepared(label=label, value=value)
    assert check(state(case), case[4])["result"] == expected


@pytest.mark.parametrize("text", [
    "Soggetto: Aurora; iscrizione R-1: numero 123; iscrizione R-2: data 14/03/2019",
    "Iscrizione; Soggetto: Aurora; numero 123; Soggetto: Beta; data 14/03/2019",
    "Registro delle imprese: 123; Albo: data 14/03/2019",
    "Iscrizione numero 123\nAltra iscrizione data 14/03/2019",
])
def test_even_one_literal_span_cannot_join_different_records_or_subjects(text):
    case = prepared(value=text, source=text)
    data = state(case)
    assert data["audit"]["facts"]["f"]["validation"] == "GROUNDED_PROPOSAL"
    assert check(data, case[4])["result"] == "UNKNOWN"
    assert not export_selection({"global_compilation": data})[0]


@pytest.mark.parametrize("second_record", ["R-1 di Aurora", "R-2 di Aurora", "R-1 di Beta"])
def test_multiple_source_spans_do_not_bypass_the_unchanged_flat_value_contract(second_record):
    a = "Iscrizione R-1 di Aurora: numero 123."
    b = f"Iscrizione {second_record}: data 14/03/2019."
    case = prepared(value="123 del 14/03/2019", source=a + "\n" + b)
    case[2]["facts"]["f"]["evidence"] = [{"chunk_id": -7, "quote": a},
                                          {"chunk_id": -7, "quote": b}]
    data = state(case)
    assert data["audit"]["facts"]["f"]["validation"] == "REJECTED"
    assert check(data, case[4])["result"] == "UNKNOWN"
    assert not data["fields"][0]["selectable"]


@pytest.mark.parametrize("value,expected", [("Lecce (LE), 14 novembre 1982", "COMPLETE"),
                                          ("14/11/1982, Lecce (LE)", "COMPLETE"),
                                          ("14 novembre 1982", "INCOMPLETE"),
                                          ("Lecce (LE)", "INCOMPLETE")])
def test_place_and_date_require_both_components(value, expected):
    case = prepared(label="Data e luogo di nascita", value=value)
    assert check(state(case), case[4])["result"] == expected


@pytest.mark.parametrize("value,expected", [
    ("Via Roma 12, 70126 Bari (BA), Italia", "COMPLETE"),
    ("Via Roma 12, Bari (BA)", "COMPLETE"),
    ("Via Roma 12", "INCOMPLETE"),
])
def test_address_components_only_when_explicitly_required(value, expected):
    case = prepared(label="Sede legale (via, città, prov.)", value=value)
    assert check(state(case), case[4])["result"] == expected


@pytest.mark.parametrize("value,expected", [("Ada Neri", "COMPLETE"),
                                          ("Ing. Ada Neri", "COMPLETE"),
                                          ("Ada", "INCOMPLETE")])
def test_name_and_surname_use_full_documented_person_identity(value, expected):
    case = prepared(label="Nome e cognome", value=value, entity="person", subject="Ada Neri",
        role="Direttore tecnico", source=f"Direttore tecnico: Ada Neri; Nominativo: {value}.")
    data = state(case)
    assert check(data, case[4])["result"] == expected
    if expected == "COMPLETE":
        assert data["fields"][0]["selectable"]


def test_complete_shape_never_repairs_preexisting_personal_role_rejection():
    case = prepared(label="Data e luogo di nascita", value="Lecce, 14/11/1982", entity="person",
                    subject="Ada Neri", source="Socia Ada Neri: Lecce, 14/11/1982.")
    data = state(case)
    assert check(data, case[4])["result"] == "COMPLETE"
    assert not data["fields"][0]["selectable"]
    assert data["audit"]["fields"][0]["subject_binding"]["result"] != "literal_match"


@pytest.mark.parametrize("label", ["Ordine/Collegio", "Ragione sociale",
                                   "Servizi di analisi e verifica", "Data",
                                   "Numero iscrizione Albo", "Poteri conferiti / qualifica"])
def test_conjunction_or_slash_alone_is_not_a_compound_requirement(label):
    case = prepared(label=label)
    assert not state(case)["audit"]["value_completeness"]


def test_separate_inline_destinations_are_not_treated_as_one_compound_value():
    case = prepared(label="Numero ________ e data", style="inline")
    data = state(case)
    assert not data["audit"]["value_completeness"]


def test_open_description_stays_unknown_without_guessing_its_required_parts():
    case = prepared(label="Qualifica e generalità del Tecnico", value="Ingegnere Ada Neri")
    assert check(state(case), case[4])["result"] == "UNKNOWN"


def test_generic_user_assent_or_narrator_review_cannot_supply_missing_date():
    case = prepared()
    data, cid = state(case), case[4]
    data["dialogue"].update(reviewed_candidate_ids=[cid], held_candidate_ids=[], phase="READY")
    data["dialogue"]["answers"]["generic"] = {"action": "CONDITION_TRUE", "applies": True,
                                              "user_quote": "Confermo tutto", "provenance": "USER"}
    assert not export_selection({"global_compilation": data})[0]
    data["dialogue"]["user_values"][cid] = {"value": "123 del 14/03/2019",
        "user_quote": "123 del 14/03/2019", "provenance": "USER"}
    assert cid in blocked_fields(data)
    assert not export_selection({"global_compilation": data})[2]


def test_user_procedural_place_date_needs_actual_components_in_the_bound_reply():
    case = prepared(label="Luogo e data", value="Bari")
    data, cid = state(case), case[4]
    assert value_check(data, cid, "Bari", user_quote="Bari")["result"] == "INCOMPLETE"
    assert value_check(data, cid, "Bari, 21 ottobre 2026",
                       user_quote="Bari, 21 ottobre 2026")["result"] == "COMPLETE"


def test_old_v12_without_completeness_check_is_not_silently_upgraded():
    case = prepared()
    data, cid = state(case), case[4]
    del data["audit"]["value_completeness"]
    assert cid in blocked_fields(data)
    assert "value_completeness" not in data["audit"]


def original():
    folder = FIXTURE / "value-completeness-v1"
    inp = json.loads((FIXTURE / "form-applicability-v1/input.json").read_text())
    layout = inspect_docx((FIXTURE / "form-applicability-v1/original.docx").read_bytes())
    data = json.loads((folder / "state.json").read_text())["global_compilation"]
    data["audit"] = validate_global_proposal((folder / "response.txt").read_text(), layout,
                                            inp["project_id"], inp["sources"])
    data["audit"] = reuse_selected_person(data, layout, inp["project_id"], inp["sources"])
    data["audit"] = reconcile_branches(data, layout, inp["project_id"], inp["sources"])
    data["fields"] = review_fields(data)
    return layout, data


def test_exact_original_source_fact_stays_true_but_partial_ui_value_cannot_be_written():
    layout, data = original()
    cid = "t25.r3.c1"
    item = next(f for f in data["audit"]["fields"] if f["candidate_id"] == cid)
    assert item["proposed_value"] == "01234560981"
    assert item["proposal_grounding"] == "GROUNDED_PROPOSAL"
    assert check(data, cid)["result"] == "INCOMPLETE"
    assert check(data, cid)["missing_components"] == ["date"]
    ids, _, users = export_selection({"global_compilation": data})
    assert cid not in ids + users and "t25.r2.c0" in ids
    assert "t0.r8.c1" not in ids + users
    values = {f["candidate_id"]: f.get("proposed_value", f.get("value"))
              for f in data["audit"]["fields"] if f["candidate_id"] in ids}
    docx = Document(BytesIO(fill_docx(layout, values)))
    assert docx.tables[25].cell(3, 2).text == ""  # Logical grid follows label gridSpan=2.
    assert docx.tables[25].cell(2, 0).text == "Bari"


def test_frozen_original_fixture_hashes_and_model_input_remain_unchanged():
    folder = FIXTURE / "value-completeness-v1"
    for name, digest in json.loads((folder / "manifest.json").read_text())["files"].items():
        assert hashlib.sha256((folder / name).read_bytes()).hexdigest() == digest
    case = prepared(value="n.123 del 14/03/2019")
    before, raw = build_global_input(case[0], "alpha", case[1]), deepcopy(case[2])
    state(case)
    assert raw == case[2] and before == build_global_input(case[0], "alpha", case[1])


@pytest.mark.parametrize("result", ["true", "false", "unknown"])
def test_complete_value_still_needs_independent_form_applicability(result):
    value = "Via Roma 12, 70126 Bari (BA)"
    content = f"Indirizzo di Aurora: {value}."
    if result == "false":
        content += f"\nIndirizzo principale di Aurora: {value}."
    if result == "true":
        content += "\nIndirizzo di Aurora: diverso dall'indirizzo principale."
    case = prepared(label="Indirizzo (via, città, prov.) (se diverso dall'indirizzo principale)",
                    value=value, source=content)
    data, cid = state(case), case[4]
    assert check(data, cid)["result"] == "COMPLETE"
    assert data["audit"]["form_applicability"][cid]["result"] == result
    assert (cid in export_selection({"global_compilation": data})[0]) is (result == "true")


@pytest.mark.parametrize("value", ["n.123 dal 2019", "n.123 dal 14/03/2019"])
def test_full_date_can_satisfy_required_registration_year(value):
    case = prepared(label="Numero ed anno d'iscrizione", value=value)
    assert check(state(case), case[4])["result"] == "COMPLETE"


@pytest.mark.parametrize("label,value,expected", [
    ("Cellulare + Telefono", "+39 080 000 2040", "UNKNOWN"),
    ("Cellulare e telefono", "Cellulare: 3331234567; Telefono: 0801234567", "COMPLETE"),
    ("Cellulare + Telefono", "Cellulare: 3331234567; Telefono: 0801234567", "COMPLETE"),
    ("Cellulare e telefono", "Cellulare: non fornito; Telefono: 0801234567", "INCOMPLETE"),
])
def test_distinct_contacts_are_not_declared_complete_from_a_single_label(label, value, expected):
    case = prepared(label=label, value=value, source="Recapiti di Aurora: " + value)
    assert check(state(case), case[4])["result"] == expected


@pytest.mark.parametrize("value,expected", [
    ("Progettazione del plesso scolastico, 14 marzo 2019", "COMPLETE"),
    ("Progettazione del plesso scolastico", "INCOMPLETE"),
    ("Contratto C-1 oggetto Progettazione plesso; Contratto C-2 data 14/03/2019", "UNKNOWN"),
])
def test_contract_description_and_date_are_bound_to_one_contract(value, expected):
    case = prepared(label="Oggetto del contratto di appalto e data di stipula", value=value,
                    source="Contratto di Aurora: " + value)
    assert check(state(case), case[4])["result"] == expected


@pytest.mark.parametrize("value,expected", [
    ("Notaio: Ada Neri; repertorio: 123; raccolta: 456", "COMPLETE"),
    ("Notaio: Ada Neri; repertorio: 123", "INCOMPLETE"),
])
def test_explicit_required_qualifiers_are_not_optional(value, expected):
    case = prepared(label="Estremi procura (notaio, repertorio, raccolta)", value=value,
                    source="Procura di Aurora: " + value)
    assert check(state(case), case[4])["result"] == expected


def test_separate_number_and_date_slots_both_continue_to_accept_atomic_source_values():
    text = "Iscrizione R-1 di Aurora: numero 123; data 14/03/2019."
    case = prepared(label="Numero ________ e data", style="inline", source=text)
    layout, sources, raw, catalog, first = case
    second = next(cid for cid in catalog["candidates"] if cid != first)
    paragraph = catalog["blocks"][0]["paragraphs"][0]
    raw["facts"]["g"] = {**raw["facts"]["f"], "property": "date", "value": "14/03/2019"}
    raw["fields"][first]["form"]["quote"] = f"Numero [[{first}]]"
    raw["fields"][second] = {**raw["fields"][first], "property": "date", "fact_id": "g",
        "form": {"paragraph_id": paragraph["id"], "quote": f"e data [[{second}]]"}}
    data = state(case)
    assert not data["audit"]["value_completeness"]
    data["dialogue"]["reviewed_candidate_ids"] = [first, second]
    ids, _, _ = export_selection({"global_compilation": data})
    assert set(ids) == {first, second}
    written = Document(BytesIO(fill_docx(layout, {first: "123", second: "14/03/2019"})))
    assert written.paragraphs[1].text == "Numero 123 e data 14/03/2019"


def test_two_different_dates_do_not_become_an_arbitrary_complete_date_choice():
    case = prepared(value="n.123 del 14/03/2019 o 15/03/2019")
    result = check(state(case), case[4])
    assert result["result"] == "UNKNOWN"
    assert result["basis"] == "ambiguous_dates_for_one_required_component"


@pytest.mark.parametrize("value,expected", [
    ("Ordine Ingegneri di Bari n.123 del 14/03/2019", "COMPLETE"),
    ("n.123 del 14/03/2019", "INCOMPLETE"),
])
def test_explicit_required_registration_authority_is_preserved(value, expected):
    case = prepared(label="Iscrizione Ordine-data e numero", value=value)
    assert check(state(case), case[4])["result"] == expected


@pytest.mark.parametrize("qualifier", ["raccolta: non fornito", "repertorio: da fornire",
                                        "notaio: ignoto"])
def test_a_required_qualifier_label_is_not_itself_a_documented_value(qualifier):
    parts = {"notaio": "Ada Neri", "repertorio": "123", "raccolta": "456"}
    key, value = qualifier.split(": ")
    parts[key] = value
    case = prepared(label="Estremi procura (notaio, repertorio, raccolta)",
                    value="; ".join(k + ": " + v for k, v in parts.items()))
    assert check(state(case), case[4])["result"] == "INCOMPLETE"


@pytest.mark.parametrize("record", ["Documento", "Certificato", "Protocollo"])
def test_number_and_date_are_not_limited_to_registration_forms(record):
    value = "n.123 del 14/03/2019"
    case = prepared(label=f"Numero e data del {record}", value=value,
                    source=f"{record} di Aurora: {value}")
    assert check(state(case), case[4])["result"] == "COMPLETE"


@pytest.mark.parametrize("value", [
    "numero 123; altra pratica data 14/03/2019",
    "numero 123 di Aurora; data 14/03/2019 di Boreale",
    "numero 123 della prima iscrizione; data 14/03/2019 della successiva",
])
def test_unrelated_clauses_do_not_form_a_record_even_inside_one_literal_value(value):
    case = prepared(value=value)
    assert check(state(case), case[4])["result"] == "UNKNOWN"


@pytest.mark.parametrize("label,value,expected", [
    ("Numero e data iscrizione e sezione", "n.123 del 14/03/2019", "INCOMPLETE"),
    ("Numero e data iscrizione e sezione", "n.123 del 14/03/2019 (sez. A)", "COMPLETE"),
    ("Numero e data iscrizione, sezione facoltativa", "n.123 del 14/03/2019", "COMPLETE"),
])
def test_optional_and_required_registration_qualifiers_are_different(label, value, expected):
    case = prepared(label=label, value=value)
    assert check(state(case), case[4])["result"] == expected


@pytest.mark.parametrize("qualifier", ["SEZ. a", "Sezione A"])
@pytest.mark.parametrize("required", [True, False])
def test_qualifier_letter_case_does_not_reject_an_otherwise_complete_value(qualifier, required):
    label = "Numero e data iscrizione" + (" e sezione" if required else "")
    case = prepared(label=label, value="n.123 del 14/03/2019 (" + qualifier + ")")
    assert check(state(case), case[4])["result"] == "COMPLETE"
