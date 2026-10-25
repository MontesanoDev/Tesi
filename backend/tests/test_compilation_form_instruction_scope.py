"""Revision 2: an explicit conditional FORM instruction governs its member blocks (offline API)."""
import json
from io import BytesIO

import pytest
import test_compilation_resilient_segments as rev2
from docx import Document
from test_compilation_global_sessions import api, provider

from app import compilation_global_dialogue as dialogue
from app import compilation_profile

__all__ = ["api", "provider"]
SINGLE = "(da compilare nel caso di partecipazione in forma singola)"
GROUP = "(da compilare nel caso di partecipazione in forma di RT non ancora costituito)"
MEMBERS = ["p4.s0", "p7.s0"]


def bold(doc, text):
    doc.add_paragraph().add_run(text).bold = True


def form():
    doc = Document()
    bold(doc, SINGLE)                                            # paragraph:0
    doc.add_paragraph("Dato del concorrente: ____________")      # paragraph:1
    doc.add_paragraph("*** ***")                                 # paragraph:2
    bold(doc, GROUP)                                             # paragraph:3
    doc.add_paragraph("Dato del componente: ____________")       # paragraph:4
    bold(doc, "(al fine di fare emergere la natura del singolo componente)")  # paragraph:5
    doc.add_paragraph("*** *** ****")                            # paragraph:6
    doc.add_paragraph("Dato del secondo componente: ____________")  # paragraph:7
    bold(doc, "CHIEDE")                                          # paragraph:8
    doc.add_paragraph("Dato indipendente: ____________")         # paragraph:9
    stream = BytesIO()
    doc.save(stream)
    return stream.getvalue()


def activate(monkeypatch, revision):
    monkeypatch.setenv("MAPI_GENERATION_REVISION", "3")
    monkeypatch.setenv("MAPI_DIALOGUE_REVISION", "1")
    original = compilation_profile.new_session_flags
    monkeypatch.setattr(compilation_profile, "new_session_flags", lambda: {
        **original(), **({"form_option_scope_revision": revision} if revision else {})})


def writer(inp, ids):
    return json.dumps(rev2.envelope(inp, {c: rev2.slot(inp, c)
                                        for c in rev2.dato_ids(inp, ids)}))


def instructions(catalog, revision=2):
    from app import compilation_form_option_scope as options

    return {r["paragraph_id"]: r["candidate_ids"] for r in options.regions(catalog, revision)
            if r.get("kind") == "instruction"}


def test_an_instruction_governs_its_member_blocks_up_to_the_next_heading():
    catalog = rev2.model_form(form())
    # The bold parenthetical note inside the member block does not end the scope; the next
    # instruction and the heading do. The data after the heading is independent.
    assert instructions(catalog) == {"paragraph:0": ["p1.s0"], "paragraph:3": MEMBERS}
    assert instructions(catalog, revision=1) == {}


@pytest.mark.parametrize("change, expected", [
    ("not_bold", None), ("own_field", None), ("no_marker", None), ("example_prefix", None),
    ("table", ["p4.s0"]), ("heading_note", ["p4.s0"]), ("no_tail", None),
])
def test_ambiguous_or_interrupted_structure_bounds_or_removes_the_scope(change, expected):
    catalog = rev2.model_form(form())
    blocks = catalog["blocks"]
    paragraph = {p["id"]: p for b in blocks for p in b["paragraphs"]}
    instruction = paragraph["paragraph:3"]
    if change == "not_bold":
        instruction["all_text_direct_bold"] = False
    elif change == "own_field":
        catalog["candidates"]["p3.s0"] = {"context_id": "paragraph:3", "type": "text"}
    elif change == "no_marker":
        instruction["text"] = "(compilare se si partecipa in RT)"
    elif change == "example_prefix":
        instruction["text"] = "Esempio: " + GROUP
    elif change == "no_tail":
        instruction["text"] = "(da compilare nel caso di)"
    elif change == "table":
        next(b for b in blocks if b["id"] == "paragraph:6")["tables"] = [
            {"id": "table:99", "rows": []}]
    else:
        paragraph["paragraph:5"]["text"] = "DICHIARA"
    assert instructions(catalog).get("paragraph:3") == expected


def state(*, condition=True, revision=2, quote=GROUP):
    catalog = rev2.model_form(form())
    return {"generation_revision": 3, "form_option_scope_revision": revision,
            "catalog": catalog, "writable_ids": sorted(catalog["candidates"]),
            "dialogue": {"answers": {}}, "audit": {"fields": [
                {"candidate_id": cid, "conditions": [], "validation": "GROUNDED_PROPOSAL"}
                for cid in catalog["candidates"]], "conditions": {"c": {
                    "form_evidence": {"paragraph_id": "paragraph:3", "quote": quote},
                    "operator": "semantic", "result": "unknown", "applies_to": []}}
                if condition else {}}}


def confirm(data, *, key="condition:c", action="CONDITION_TRUE", ids=None):
    from app import compilation_scope_review as scope

    ids = list(MEMBERS) if ids is None else ids
    record = {"action": action, "applies": action == "CONDITION_TRUE",
              "user_quote": "Partecipiamo come RT", "affected_fields": ids,
              "provenance": "USER"}
    record["scope_validation"] = scope.proof(scope.target(data, key, ids), record)
    data["dialogue"]["answers"][key] = record


def projected(data):
    from app import compilation_form_option_scope as options

    data["audit"] = options.apply(data["audit"], data["catalog"],
                                  data["form_option_scope_revision"])
    return data


def test_member_data_needs_the_instruction_condition_and_a_verified_user_choice():
    from app import compilation_scope_review as scope
    from app.compilation_segments import applicability_states

    data = projected(state())
    assert data["audit"]["conditions"]["c"]["applies_to"] == MEMBERS
    assert data["audit"]["form_option_scope"]["revision"] == 2
    # The single-participation block has no condition of its own: it stays blocked too.
    assert scope.blocked_source_fields(data) == {"p1.s0", *MEMBERS}
    assert applicability_states(data)["p4.s0"]["state"] == "CONDIZIONATA"
    assert applicability_states(data)["p1.s0"]["state"] == "NON_ACCERTATA"
    confirm(data)
    assert scope.blocked_source_fields(data) == {"p1.s0"}
    assert applicability_states(data)["p7.s0"]["state"] == "VERIFICATA"
    assert "p9.s0" not in scope.blocked_source_fields(data)
    data["dialogue"]["answers"].clear()  # Revocation.
    assert scope.blocked_source_fields(data) == {"p1.s0", *MEMBERS}


@pytest.mark.parametrize("defect", ["missing", "foreign", "partial", "invalid", "detached"])
def test_other_or_incomplete_conditions_do_not_cover_the_instruction(defect):
    from app import compilation_form_option_scope as options
    from app import compilation_scope_review as scope
    from app.compilation_global_sessions import review_fields
    from app.compilation_segments import applicability_states

    data = state(condition=defect != "missing")
    if defect not in {"missing", "detached"}:
        condition = data["audit"]["conditions"]["c"]
        if defect == "foreign":
            condition["form_evidence"]["paragraph_id"] = "paragraph:0"
            condition["form_evidence"]["quote"] = SINGLE
        elif defect == "partial":
            condition["form_evidence"]["quote"] = "da compilare nel caso di partecipazione"
        else:
            condition["validation_errors"] = ["Invalid FORM citation"]
    if defect != "detached":
        projected(data)
    confirm(data)  # Even a simulated SUPPORTED receipt does not fix an omitted condition.
    assert set(MEMBERS) <= options.instruction_uncovered_fields(data)
    assert set(MEMBERS) <= scope.blocked_source_fields(data)
    assert set(MEMBERS) <= scope.blocked_user_fields(data)
    assert applicability_states(data)["p4.s0"]["state"] == "NON_ACCERTATA"
    proposed = next(f for f in data["audit"]["fields"] if f["candidate_id"] == "p7.s0")
    proposed.update(value="ABC123", proposal={"value": "ABC123"}, form_evidence={
        "paragraph_id": "paragraph:7", "quote": "Dato del secondo componente"})
    field = next(f for f in review_fields(data) if f["candidate_id"] == "p7.s0")
    assert any("istruzione FORM non accertata" in issue for issue in field["issues"])


@pytest.mark.parametrize("quote", [GROUP, GROUP.strip("()"), GROUP.strip("()") + ".",
                                   GROUP.upper()])
def test_the_whole_instruction_may_be_cited_without_parentheses_or_final_punctuation(quote):
    from app import compilation_form_option_scope as options

    data = projected(state(quote=quote))
    confirm(data)
    assert not set(MEMBERS) & options.blocked_fields(data)


@pytest.mark.parametrize("change", ["negative", "skip", "foreign", "other_fields",
                                   "source_only", "source_false"])
def test_a_decision_only_authorizes_its_verified_current_branch(change):
    from app import compilation_scope_review as scope

    data = projected(state())
    if change in {"negative", "skip"}:
        confirm(data, action="CONDITION_FALSE" if change == "negative" else "SKIP")
    elif change == "foreign":
        confirm(data, key="condition:other")
    elif change == "other_fields":
        confirm(data, ids=["p9.s0"])
    elif change == "source_only":
        data["audit"]["conditions"]["c"]["result"] = "true"
    else:
        confirm(data)
        data["audit"]["conditions"]["c"]["result"] = "false"
    assert set(MEMBERS) <= scope.blocked_source_fields(data)


def test_revision_one_ignores_instructions_and_stays_unchanged():
    from app import compilation_form_option_scope as options
    from app import compilation_scope_review as scope

    data = projected(state(revision=1))
    assert options.uncovered_fields(data) == set()
    assert data["audit"]["form_option_scope"] == {"revision": 1, "origin": "BACKEND",
                                                  "regions": []}
    assert not scope.blocked_source_fields(data)


def test_creation_requires_the_explicit_revision_two_opt_in(monkeypatch):
    from app import compilation_form_option_scope as options

    monkeypatch.delenv("MAPI_COMPILATION_PROFILE", raising=False)
    monkeypatch.delenv("MAPI_FORM_OPTION_SCOPE_REVISION", raising=False)
    default = compilation_profile.new_session_flags()
    monkeypatch.setenv("MAPI_FORM_OPTION_SCOPE_REVISION", "1")
    first = compilation_profile.new_session_flags()
    monkeypatch.setenv("MAPI_FORM_OPTION_SCOPE_REVISION", "2")
    second = compilation_profile.new_session_flags()
    assert "form_option_scope_revision" not in default
    assert first["form_option_scope_revision"] == 1
    assert first["compilation_profile"] == "sezioni-opzioni-2026-10-25"
    assert second["form_option_scope_revision"] == 2
    assert second["compilation_profile"] == "sezioni-istruzioni-2026-10-25"
    assert options.enabled(second)
    assert {k: v for k, v in second.items() if k not in {"compilation_profile",
                                                       "form_option_scope_revision"}} == {
        k: v for k, v in default.items() if k != "compilation_profile"}
    monkeypatch.setenv("MAPI_COMPILATION_PROFILE", "storico")
    assert not options.enabled(compilation_profile.new_session_flags())


@pytest.mark.anyio
@pytest.mark.parametrize("revision", [None, 2])
async def test_member_data_without_a_condition_is_not_exported(api, provider, monkeypatch,
                                                               revision):
    activate(monkeypatch, revision)
    s = await rev2.started(api, monkeypatch, form(), rev2.Model({0: [writer], 1: [writer]}))
    response = await rev2.analyze(api, s)
    assert response.status_code == 200, response.text
    s = response.json()
    selected, _, _ = dialogue.export_selection(s)
    assert "p9.s0" in selected  # Independent, documented data after the heading.
    export = await rev2.draft(api, s, ["p4.s0", "p9.s0"])
    if revision:
        assert export.status_code == 409, export.text
        assert not {"p1.s0", *MEMBERS} & set(selected)
        assert (await rev2.draft(api, s, selected)).status_code == 200
    else:
        assert export.status_code == 200, export.text


@pytest.mark.anyio
async def test_member_data_is_exported_after_the_instruction_is_confirmed(
    api, provider, monkeypatch,
):
    from app import compilation_scope_review as scope
    from app.compilation_chat import ClarificationReply

    activate(monkeypatch, 2)

    def with_condition(inp, ids):
        raw = json.loads(writer(inp, ids))
        raw["conditions"]["c"] = rev2.condition(inp, "paragraph:3", GROUP)
        return json.dumps(raw)

    s = await rev2.started(api, monkeypatch, form(),
                          rev2.Model({0: [with_condition], 1: [with_condition]}))
    response = await rev2.analyze(api, s)
    assert response.status_code == 200, response.text
    s = response.json()
    data = rev2.state_of(s)["global_compilation"]
    target = next(q for q in dialogue.question_pool(data) if q["key"].startswith("condition:"))
    assert target["field_ids"] == MEMBERS
    question = next(q for q in s["chat"]["question"]["slots"] if q["key"] == target["key"])
    reply = ClarificationReply(slot=question["slot"], action="CONDITION_TRUE",
                               user_quote="Partecipiamo come RT non ancora costituito")
    proof = scope.proof(scope.target(data, target["key"], target["field_ids"]), reply)
    s, errors = dialogue.apply_replies("alpha", s["id"], s["version"], [reply], reply.user_quote,
                                       scope_verifications={reply.slot: proof})
    assert not errors
    s = (await rev2.analyze(api, s)).json()
    ids, conditions, _ = dialogue.export_selection(s)
    assert set(ids) == {*MEMBERS, "p9.s0"}  # The single block has no condition: still empty.
    exported = await rev2.draft(api, s, ids, conditions=conditions)
    assert exported.status_code == 200, exported.text
    data = rev2.state_of(exported.json())["global_compilation"]
    data["dialogue"]["answers"].clear()
    assert set(MEMBERS) <= scope.blocked_source_fields(data)


LONG = ("(da compilare nel caso di partecipazione in forma di RT non ancora costituito, "
        "ovvero di Consorzio Ordinario di cui all’art. 68 del D. Lgs n. 36/2023)")


@pytest.mark.parametrize("quote, covered", [
    (LONG, True),
    ("da compilare nel caso di partecipazione in forma di RT non ancora costituito", True),
    ("(da compilare nel caso di partecipazione in forma di RT non ancora costituito, ovvero "
     "di Consorzio Ordinario", True),  # Trailing legal reference left out.
    ("da compilare nel caso di partecipazione in forma di RT", False),  # Before 1st alternative.
    ("da compilare nel caso di partecipazione in forma di RT non ancora costitu", False),
    ("partecipazione in forma di RT non ancora costituito", False),  # Inner fragment.
    ("RT non ancora costituito, ovvero di Consorzio Ordinario", False),
])
def test_an_anchored_prefix_must_reach_the_first_alternative_of_the_instruction(quote, covered):
    from app import compilation_form_option_scope as options

    data = state(quote=quote)
    instruction = next(p for b in data["catalog"]["blocks"] for p in b["paragraphs"]
                       if p["id"] == "paragraph:3")
    instruction["text"] = LONG
    projected(data)
    confirm(data)
    assert (not set(MEMBERS) & options.blocked_fields(data)) is covered
    assert bool(set(MEMBERS) & options.instruction_uncovered_fields(data)) is not covered
