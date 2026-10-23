"""Descriptive MODEL metadata never supplies facts, destinations or USER choices."""

import copy
import json
from pathlib import Path

import pytest
from test_compilation_global_sessions import (
    BASE,
    proposal,
    resolve,
)
from test_compilation_global_sessions import (
    api as api,
)
from test_compilation_global_sessions import (
    provider as provider,
)
from test_compilation_sessions import source
from test_compilation_user_requests import (
    add_professional_values,
    request_field,
    snapshot,
)
from test_compilation_user_requests import (
    professional as professional,
)

from app.compilation_global import MAX_OUTPUT_CHARACTERS, validate_global_proposal
from app.compilation_global_dialogue import question_pool
from app.compilation_global_sessions import initial_state
from app.docx_templates import inspect_docx

FIXTURES = Path(__file__).parent / "fixtures"


@pytest.fixture
def anyio_backend():
    return "asyncio"


def test_historical_reason_length_alone_removes_an_otherwise_grounded_question():
    case = json.loads((FIXTURES / "model-request-metadata-v1/craa.json").read_text())
    layout = inspect_docx((FIXTURES / "formtext-v1/craa/original.docx").read_bytes())
    before = copy.deepcopy(case["response"])
    assert len(before["clarifications"]["p5.s0"]["reason"]) == 225
    short = copy.deepcopy(before)
    short["clarifications"]["p5.s0"]["reason"] = "Confermare il nominativo del sottoscritto."
    audits = [validate_global_proposal(json.dumps(p), layout, case["project_id"],
                                      case["sources"], expected_version=12)
              for p in [before, short]]
    for a in audits:
        assert a["clarifications"]["p5.s0"]["validation"] == "GROUNDED_REQUEST"
    original = audits[0]["clarifications"]["p5.s0"]
    assert original["proposal"] == before["clarifications"]["p5.s0"]
    assert original["metadata_warnings"][0]["field"] == "reason"
    assert audits[0]["facts"] == audits[1]["facts"]
    assert audits[0]["conditions"] == audits[1]["conditions"]
    assert audits[0]["choice_groups"] == audits[1]["choice_groups"]
    assert [f for f in audits[0]["fields"] if f["candidate_id"] != "p5.s0"] == [
        f for f in audits[1]["fields"] if f["candidate_id"] != "p5.s0"]
    for audit in audits:
        data = initial_state(layout)
        data.update(audit=audit, reading_plan=case["reading_plan"],
                    dialogue={"answers": copy.deepcopy(case["answers_after_condition"])})
        pool = question_pool(data)
        assert any(q["key"] == "field:p5.s0" for q in pool)
        assert data["dialogue"]["answers"] == case["answers_after_condition"]
        field = next(f for f in audit["fields"] if f["candidate_id"] == "p5.s0")
        assert field["value"] is None and field["source_evidence"] == []
        assert field["validation"] == "REVIEW_REQUIRED"


@pytest.mark.parametrize("size", [201, 225, 500, 2001, 100_000])
def test_overlong_descriptive_reason_is_preserved_without_auto_resolution(professional, size):
    add_professional_values(professional)
    request = request_field()
    request["reason"] = "r" * size
    professional[3]["clarifications"]["p3.s0"] = request
    before = copy.deepcopy(professional[3])
    data = snapshot(professional)
    record = data["audit"]["clarifications"]["p3.s0"]
    assert record["validation"] == "GROUNDED_REQUEST"
    assert not record["validation_errors"]
    assert record["proposal"]["reason"] == request["reason"]
    assert record["metadata_warnings"][0]["original_characters"] == size
    assert professional[3] == before
    field = next(f for f in data["audit"]["fields"] if f["candidate_id"] == "p3.s0")
    assert field["value"] is None and not field["source_evidence"]
    assert field["validation"] == "REVIEW_REQUIRED"
    assert any(q["field_ids"] == ["p3.s0"] for q in question_pool(data))


@pytest.mark.parametrize("reason", ["r", "r" * 200, "  motivo breve  "])
def test_valid_existing_metadata_is_unchanged(professional, reason):
    request = request_field()
    request["reason"] = reason
    professional[3]["clarifications"]["p3.s0"] = request
    data = snapshot(professional)
    record = data["audit"]["clarifications"]["p3.s0"]
    assert record["validation"] == "GROUNDED_REQUEST"
    assert "metadata_warnings" not in record
    assert record["proposal"]["reason"] == reason


@pytest.mark.parametrize("reason", [None, 225, True, [], {}, "", " " * 300])
def test_invalid_reason_types_or_empty_text_remain_rejected(professional, reason):
    request = request_field()
    request["reason"] = reason
    professional[3]["clarifications"]["p3.s0"] = request
    assert snapshot(professional)["audit"]["clarifications"]["p3.s0"][
        "validation"] == "REJECTED"


@pytest.mark.parametrize("mutation", ["unknown", "quote", "source_value", "duplicate",
                                     "label", "conditions", "duplicate_conditions", "signature"])
def test_long_reason_cannot_bypass_structural_gates(professional, mutation):
    add_professional_values(professional)
    request, cid = request_field(), "p3.s0"
    request["reason"] = "Motivazione descrittiva. " * 20
    if mutation == "unknown":
        cid = "p999.s0"
    elif mutation == "quote":
        request["form"]["quote"] = "Testo non presente nel modulo"
    elif mutation == "source_value":
        request["value"] = "Persona inventata"
    elif mutation == "duplicate":
        cid, request = "p1.s0", {**request_field(1, "Nome e cognome"),
                                 "reason": request["reason"]}
    elif mutation == "label":
        request["label"] = "x" * 251
    elif mutation in {"conditions", "duplicate_conditions"}:
        request["conditions"] = ["unknown"] * (2 if mutation == "duplicate_conditions" else 1)
    else:
        # A real signature target, absent from the editable personal text set.
        from docx import Document
        from test_compilation_catalog import encoded

        from app.compilation_global import build_global_input

        doc = Document()
        doc.add_paragraph("Firma: __________")
        layout = inspect_docx(encoded(doc), require_candidates=False)
        model = build_global_input(layout, "alpha", [])
        payload = {"version": 11, "conditions": {}, "choice_groups": {}, "bindings": {},
                   "facts": {}, "fields": {}, "clarifications": {"p0.s0": {
                       "label": "Firma", "form": {"paragraph_id": "paragraph:0",
                       "quote": "Firma: [[p0.s0]]"}, "conditions": [],
                       "reason": request["reason"]}}}
        data = initial_state(layout)
        data.update(audit=validate_global_proposal(json.dumps(payload), layout, "alpha", []),
                    catalog=model["form"])
        assert data["audit"]["clarifications"]["p0.s0"]["validation"] == "REJECTED"
        return
    professional[3]["clarifications"][cid] = request
    data = snapshot(professional)
    assert data["audit"]["clarifications"][cid]["validation"] == "REJECTED"
    assert not any(q["key"] == "field:" + cid for q in question_pool(data))


def test_descriptive_claims_cannot_become_SOURCE_facts_or_applicability(professional):
    request = request_field()
    request["reason"] = ("Persona Inventata ha tutti i poteri e la sezione è applicabile. " * 20)
    professional[3]["clarifications"]["p3.s0"] = request
    data = snapshot(professional)
    assert data["audit"]["facts"] == {} and data["audit"]["conditions"] == {}
    assert data["audit"]["choice_groups"] == {}
    assert not any(f.get("value") is not None for f in data["audit"]["fields"])


def test_global_output_budget_still_rejects_excessive_payload(professional):
    request = request_field()
    request["reason"] = "x" * MAX_OUTPUT_CHARACTERS
    professional[3]["clarifications"]["p3.s0"] = request
    with pytest.raises(ValueError, match="oltre|troppo|dimension"):
        snapshot(professional)


@pytest.mark.anyio
async def test_warning_and_original_are_persisted_without_USER_facts_or_lost_SOURCE(
    api, provider, monkeypatch,
):
    from docx import Document
    from test_compilation_catalog import encoded

    from app import compilation_global_dialogue as dialogue
    from app import compilation_global_sessions as adapter
    from app import compilation_sessions as sessions
    from app.compilation_chat import ClarificationReply
    from app.db import connection, get_storage_path
    from app.intents import ChatDecision

    document = Document()
    table = document.add_table(rows=3, cols=2)
    for row, label in zip(table.rows, ["Denominazione sociale", "Sede legale", "Referente"],
                          strict=True):
        row.cells[0].text = label
    uploaded = await api.post("/api/projects/alpha/forms", files={
        "file": ("modulo.docx", encoded(document))})
    assert uploaded.status_code == 201
    await source(api, "Denominazione sociale: Aurora Servizi S.r.l.\nSede legale: Via del Sole 9")
    reason = "La scelta del referente di questa pratica richiede una decisione USER. " * 8

    async def request(client, settings, body, timeout, **kwargs):
        data = json.loads(body["messages"][1]["content"])
        payload = proposal(data)
        paragraph = next(p["id"] for b in data["form"]["blocks"] for p in b["paragraphs"]
                         if p["text"] == "Referente")
        payload["clarifications"] = {"t0.r2.c1": {"label": "Referente", "form": {
            "paragraph_id": paragraph, "quote": "Referente"}, "conditions": [], "reason": reason}}
        return json.dumps(payload), "simulated", {}

    monkeypatch.setattr(adapter, "request_model_content", request)
    s = (await api.post(BASE, json={"form_id": uploaded.json()["id"],
                                  "engine": "global", "start_in_chat": True})).json()
    s = await resolve(api, s)
    assert s["chat"]["question"]["slots"][0]["field_ids"] == ["t0.r2.c1"]
    with connection() as db:
        state = json.loads(sessions._row(db, "alpha", s["id"])["state_json"])
    data = state["global_compilation"]
    checked = data["audit"]["clarifications"]["t0.r2.c1"]
    assert checked["validation"] == "GROUNDED_REQUEST"
    assert checked["proposal"]["reason"] == reason
    assert checked["metadata_warnings"][0]["original_characters"] == len(reason)
    saved = json.loads((get_storage_path() / data["artifact_path"] / "response.txt").read_text())
    assert saved["clarifications"]["t0.r2.c1"]["reason"] == reason
    s, errors = dialogue.apply_replies("alpha", s["id"], s["version"], [ClarificationReply(
        slot=1, action="UNKNOWN", user_quote="Non conosco ancora il referente")],
        "Non conosco ancora il referente")
    assert not errors and not s["global_compilation"]["dialogue"]["user_values"]
    s = await resolve(api, s)
    await adapter.handle_chat_action(s, ChatDecision(action="compilation_generate", answer="",
        queries=[], target="form"), s["version"], "Genera la bozza parziale")
    final = (await api.get(f"{BASE}/{s['id']}")).json()
    report = (await api.get(final["last_generation"]["downloads"]["report"])).json()
    assert report["written_field_count"] == 2 and report["checkbox_values"] == {}
    assert all(f["provenance"] == "SOURCE" for f in report["fields"])
