"""Conversational global review, real persistence/validators and simulated transports."""

import json
from io import BytesIO

import pytest
from docx import Document
from test_compilation_global_sessions import BASE, proposal, resolve, start
from test_compilation_global_sessions import api as global_api
from test_compilation_global_sessions import provider as global_provider

from app import compilation_chat
from app import compilation_global_dialogue as dialogue
from app import compilation_global_sessions as adapter
from app import compilation_sessions as sessions
from app.compilation_chat import ClarificationDecision, ClarificationReply
from app.db import connection, get_storage_path
from app.generation import GenerationError
from app.intents import ChatDecision, bind_turn_plan, parse_turn_plan, routing_context

api, provider = global_api, global_provider


@pytest.fixture
def anyio_backend():
    return "asyncio"


async def conditional(api, monkeypatch):
    async def request(client, settings, body, timeout, **kwargs):
        data = json.loads(body["messages"][1]["content"])
        return json.dumps(proposal(data, conditional=True)), "simulated", {}

    monkeypatch.setattr(adapter, "request_model_content", request)
    return await resolve(api, await start(api))


def reply(session, action, message, slot=1, value=None):
    return dialogue.apply_replies("alpha", session["id"], session["version"], [
        ClarificationReply(slot=slot, action=action, value=value, user_quote=message)], message)


@pytest.mark.anyio
async def test_natural_summary_persisted_bounded_and_no_technical_identifiers(api, provider):
    s = await resolve(api, await start(api))
    message = s["global_compilation"]["dialogue"]["message"]
    assert "Ho trovato" in message and "bozza parziale" in message
    assert dialogue.no_technical_ids(message)
    assert not s["last_generation"] and s["summary"]["resolved"] == 0
    requests = list(get_storage_path().rglob("dialogue/*/request.json"))
    assert len(requests) == 1
    request = json.loads(requests[0].read_text())
    body = request["body"]
    context = json.loads(body["messages"][1]["content"])
    assert body["max_tokens"] == 2048 and len(provider) == 1
    assert "SCHEMA OUTPUT:" in body["messages"][0]["content"]
    assert '"question"' in body["messages"][0]["content"]
    assert "t0.r0.c1" not in json.dumps(context)
    assert len(context["question_pool"]) <= 24
    assert "SOURCE" not in message
    assert (await api.get(f"{BASE}/{s['id']}")).json()["global_compilation"]["dialogue"] == (
        s["global_compilation"]["dialogue"])


@pytest.mark.anyio
@pytest.mark.parametrize("invalid_group", [None, {}, {
    "label": "Scelta condizionale", "subject": "Richiedente", "property": "partecipazione",
    "selection": "single", "min_selected": 0, "conditions": [],
    "form": {"paragraph_id": "paragraph:0", "quote": "Dati aziendali"},
    "options": {"p999.cb0": {"label": "x" * 251, "form": {
        "paragraph_id": "paragraph:0", "quote": "Dati aziendali"}}},
}])
async def test_invalid_choice_group_is_saved_without_500_or_losing_source_proposals(
    api, provider, monkeypatch, invalid_group,
):
    async def request(client, settings, body, timeout, **kwargs):
        payload = proposal(json.loads(body["messages"][1]["content"]))
        payload["choice_groups"] = {"invalid": invalid_group}
        return json.dumps(payload), "simulated", {}

    monkeypatch.setattr(adapter, "request_model_content", request)
    s = await resolve(api, await start(api))
    assert s["status"] == "WAITING_FOR_USER" and s["last_error"] is None
    assert s["global_compilation"]["phase"] == "REVIEW"
    assert s["global_compilation"]["proposal_count"] == 2
    assert not s["global_compilation"]["dialogue"].get("error")
    assert s["chat"]["question"]["kind"] == "generate"
    assert s == (await api.get(f"{BASE}/{s['id']}")).json()
    audit = (await api.get(f"{BASE}/{s['id']}/global-report")).json()
    assert audit["choice_groups"]["invalid"]["validation"] == "REJECTED"
    assert audit["choice_groups"]["invalid"]["validation_errors"]
    assert "min_selected" not in audit["choice_groups"]["invalid"]
    ids, conditions, users = dialogue.export_selection(s)
    generated = await api.post(f"{BASE}/{s['id']}/global-draft", json={
        "version": s["version"], "candidate_ids": ids, "allow_unresolved": True})
    assert generated.status_code == 200, generated.text
    report = (await api.get(generated.json()["last_generation"]["downloads"]["report"])).json()
    assert report["written_field_count"] == 2 and report["checkbox_values"] == {}
    assert all(f["provenance"] == "SOURCE" for f in report["fields"])
    assert not conditions and not users and s["summary"]["resolved"] == 0


@pytest.mark.anyio
@pytest.mark.parametrize("action", ["UNKNOWN", "SKIP", "REFUSE", "CONDITION_FALSE"])
async def test_control_defers_or_excludes_without_value_and_without_repeating(
    api, provider, monkeypatch, action,
):
    s = await conditional(api, monkeypatch)
    assert s["chat"]["question"]["slots"][0]["kind"] == "applicability"
    s, errors = reply(s, action, "Una risposta semantica pertinente")
    assert not errors and s["chat"]["auto_continue"] and not s["chat"]["question"]
    record = s["global_compilation"]["dialogue"]["answers"]["condition:c"]
    assert record["provenance"] == "USER" and "value" not in record
    s = await resolve(api, s)
    assert s["chat"]["question"]["kind"] == "generate"
    assert s["global_compilation"]["attempts"] == 1
    assert not s["last_generation"] and not dialogue.export_selection(s)[0]


@pytest.mark.anyio
async def test_conversational_confirmation_then_explicit_generation_preserves_source(api, provider,
                                                                                   monkeypatch):
    s = await conditional(api, monkeypatch)
    context = compilation_chat.planner_context(s)
    route = routing_context(context)
    assert route["slots"][0]["slot"] == 1
    assert "field_ids" not in json.dumps(route) and "condition_id" not in json.dumps(route)
    plan = parse_turn_plan({"intent": "answer_fields", "answer": "", "queries": [],
        "target": "form", "replies": [{"slot": 1, "action": "CONDITION_TRUE",
                                      "user_quote": "Sì, si applica"}]}, {1})
    decision = bind_turn_plan(plan, context, "Sì, si applica")
    assert isinstance(decision, ClarificationDecision)
    await compilation_chat.handle_group_decision("alpha", s["conversation_id"], decision, s,
                                                "Sì, si applica", s["version"])
    s = sessions.get_session("alpha", s["id"])
    assert not s["last_generation"]
    s = await resolve(api, s)
    decision = ChatDecision(action="compilation_generate", answer="", queries=[], target="form")
    _, command = await adapter.handle_chat_action(s, decision, s["version"],
                                                 "Sì, genera la bozza parziale")
    assert command["action"] == "generated"
    s = sessions.get_session("alpha", s["id"])
    report = (await api.get(s["last_generation"]["downloads"]["report"])).json()
    assert report["written_field_count"] == 2 and report["resolved"] == 0
    assert all(f["provenance"] == "SOURCE" for f in report["fields"])
    assert report["user_review"]["user_quote"] == "Sì, genera la bozza parziale"
    assert report["user_review"]["confirmed_condition_ids"] == ["c"]


@pytest.mark.anyio
@pytest.mark.parametrize("consent", ["si", "siii"])
async def test_skipped_conditions_then_chat_consent_exports_blank_without_confirming_facts(
    api, provider, monkeypatch, consent,
):
    from app import intents

    s = await conditional(api, monkeypatch)
    s, errors = reply(s, "SKIP", "lascia in bianco")
    assert not errors and s["last_generation"] is None
    s = await resolve(api, s)
    assert s["chat"]["question"]["kind"] == "generate"
    assert "senza aggiungere valori" in s["chat"]["question"]["message"]
    before = s["global_compilation"]["dialogue"]["answers"]

    async def route(client, settings, body, timeout, **kwargs):
        context = body["messages"][1]["content"]
        assert '"question_kind": "generate"' in context
        return json.dumps({"intent": "generate", "answer": "", "queries": [],
                           "target": "form"}), "simulated", {}

    monkeypatch.setattr(intents, "request_model_content", route)
    response = await api.post("/api/projects/alpha/answer", json={
        "question": consent, "document_id": s["form_id"],
        "conversation_id": s["conversation_id"], "compilation_session_id": s["id"],
        "compilation_version": s["version"],
    })
    assert response.status_code == 200, response.text
    assert response.json()["compilation"]["action"] == "generated"
    assert "senza inserire alcun valore" in response.json()["answer"]
    s = (await api.get(f"{BASE}/{s['id']}")).json()
    assert s["status"] == "GENERATED" and s["chat"]["question"] is None
    assert s["global_compilation"]["dialogue"]["answers"] == before
    assert s["summary"]["resolved"] == s["summary"]["user_provided"] == 0
    run = s["last_generation"]
    report = (await api.get(run["downloads"]["report"])).json()
    assert report["written_field_count"] == 0 and report["fields"] == []
    assert report["user_review"]["user_quote"] == consent
    assert report["user_review"]["allow_empty"] is True
    assert report["user_review"]["confirmed_condition_ids"] == []
    original = Document(BytesIO((await api.get(run["downloads"]["template"])).content))
    draft = Document(BytesIO((await api.get(run["downloads"]["docx"])).content))
    assert [t._element.xml for t in draft.tables] == [t._element.xml for t in original.tables]
    assert [p.text for p in draft.paragraphs if not p.text.startswith("BOZZA NON VERIFICATA")] == [
        p.text for p in original.paragraphs]
    assert len(provider) == 0  # The conditional transport is separate; no new analysis.


@pytest.mark.anyio
@pytest.mark.parametrize("allow_empty,quote", [(False, "Sì"), (True, ""), (True, "  ")])
async def test_empty_export_requires_explicit_opt_in_and_user_consent(
    api, provider, allow_empty, quote,
):
    s = await resolve(api, await start(api))
    response = await api.post(f"{BASE}/{s['id']}/global-draft", json={
        "version": s["version"], "allow_unresolved": True, "allow_empty": allow_empty,
        "user_quote": quote,
    })
    assert response.status_code == 422
    assert (await api.get(f"{BASE}/{s['id']}")).json()["last_generation"] is None


@pytest.mark.anyio
async def test_partial_export_keeps_eligible_source_and_excludes_skipped_conditional_value(
    api, provider, monkeypatch,
):
    async def mixed(client, settings, body, timeout, **kwargs):
        data = proposal(json.loads(body["messages"][1]["content"]), conditional=True)
        data["fields"]["t0.r1.c1"].update(disposition="value", conditions=[])
        data["conditions"]["c"]["scope"] = [{"kind": "candidate", "id": "t0.r0.c1"}]
        return json.dumps(data), "simulated", {}

    monkeypatch.setattr(adapter, "request_model_content", mixed)
    s = await resolve(api, await start(api))
    s, errors = reply(s, "SKIP", "lascia vuoto ciò che non sai")
    assert not errors
    s = await resolve(api, s)
    assert "con i dati disponibili" in s["chat"]["question"]["message"]
    decision = ChatDecision(action="compilation_generate", answer="", queries=[], target="form")
    _, command = await adapter.handle_chat_action(s, decision, s["version"], "Sì")
    assert command["action"] == "generated"
    s = sessions.get_session("alpha", s["id"])
    report = (await api.get(s["last_generation"]["downloads"]["report"])).json()
    assert report["written_field_count"] == 1
    assert report["fields"][0]["cell_id"] == "t0.r1.c1"
    assert report["fields"][0]["provenance"] == "SOURCE"
    doc = Document(BytesIO((await api.get(s["last_generation"]["downloads"]["docx"])).content))
    assert doc.tables[0].cell(0, 1).text == ""
    assert doc.tables[0].cell(1, 1).text == "Via del Sole 9"


@pytest.mark.anyio
async def test_pause_resume_preserves_group_without_reanalysis(api, provider, monkeypatch):
    s = await conditional(api, monkeypatch)
    question = s["chat"]["question"]
    s = adapter.control_session("alpha", s["id"], s["version"], "pause")
    assert s["chat"]["paused"] and s["chat"]["question"] is None
    assert not s["chat"]["auto_continue"]
    s = adapter.control_session("alpha", s["id"], s["version"], "resume")
    assert s["chat"]["question"] == question and not s["chat"]["auto_continue"]
    assert s["global_compilation"]["attempts"] == 1


@pytest.mark.anyio
async def test_bad_summary_safe_explicit_retry_without_reanalysis(api, provider, monkeypatch):
    async def invalid(*args, **kwargs):
        return '{"message":"ID t0.r0.c1 da compilare", "questions":[]}', "simulated", {}

    monkeypatch.setattr(dialogue, "request_model_content", invalid)
    s = await resolve(api, await start(api))
    assert s["global_compilation"]["dialogue"]["error"] == "ValueError"
    assert not s["last_generation"] and not s["chat"]["auto_continue"]
    s = adapter.control_session("alpha", s["id"], s["version"], "resume")
    assert s["chat"]["auto_continue"]
    s = await resolve(api, s)
    assert s["global_compilation"]["attempts"] == 1 and len(provider) == 1


@pytest.mark.anyio
async def test_historical_review_gets_summary_only(api, provider):
    s = await resolve(api, await start(api))
    with connection() as db:
        row = sessions._row(db, "alpha", s["id"])
        state = json.loads(row["state_json"])
        state["global_compilation"].pop("dialogue")
        s = sessions.save_state(db, row, state, "fixture_historical_review")
    assert s["chat"]["auto_continue"] and not s["chat"]["question"]
    s = await resolve(api, s)
    assert s["global_compilation"]["dialogue"]["message"] and len(provider) == 1


@pytest.mark.anyio
async def test_bad_reply_does_not_mutate_condition(api, provider, monkeypatch):
    s = await conditional(api, monkeypatch)
    for quote, message in [("inventato", "No"), ("Sì o forse no", "Sì o forse no")]:
        s, errors = dialogue.apply_replies("alpha", s["id"], s["version"], [
            ClarificationReply(slot=1, action="CONDITION_TRUE", user_quote=quote)], message)
        assert errors == [1]
        assert not s["global_compilation"]["dialogue"]["answers"]
        assert not s["chat"]["auto_continue"] and not s["last_generation"]


@pytest.mark.anyio
async def test_provider_failure_in_summary_keeps_audit_without_retry(api, provider, monkeypatch):
    async def fail(*args, **kwargs):
        raise GenerationError("simulated summary failure")

    monkeypatch.setattr(dialogue, "request_model_content", fail)
    s = await resolve(api, await start(api))
    assert s["status"] == "WAITING_FOR_USER" and s["global_compilation"]["proposal_count"] == 2
    assert "non sono riuscito" in s["global_compilation"]["dialogue"]["message"]
    assert not s["last_generation"] and len(provider) == 1


def test_question_pool_omits_unassessed_ids_and_prioritizes_shared_conditions():
    fields = [{"candidate_id": f"t0.r{i}.c1", "label": f"Informazione {i}", "kind": "data",
               "disposition": "review", "form_evidence": {"quote": f"Informazione {i}"},
               "conditions": ["wide"] if i < 5 else ["local"]} for i in range(6)]
    data = {"audit": {"fields": fields + [{"candidate_id": "p123.s0", "proposal": None}],
                      "conditions": {k: {"result": "unknown", "form_evidence": {"quote": k}}
                                     for k in ["local", "wide"]}},
            "writable_ids": [f["candidate_id"] for f in fields]}
    pool = dialogue.question_pool(data)
    assert pool[0]["condition_id"] == "wide" and pool[0]["impact"] == 5
    assert len(pool) == 2
    plan = dialogue.DialoguePlan(message="Ho trovato informazioni nelle fonti.", questions=[
        dialogue.DialogueQuestion(slot=1, question="La prima sezione si applica?")])
    result = dialogue.render_dialogue(plan, pool)
    assert len(result["question"]["slots"]) == 1
    assert "t0.r" not in result["message"] and "p123" not in result["message"]


@pytest.mark.anyio
async def test_user_value_only_for_active_slot_and_kept_distinct_from_source(api, provider):
    s = await resolve(api, await start(api))
    # A grounded review question, never a blank/unassessed XML candidate.
    with connection() as db:
        row = sessions._row(db, "alpha", s["id"])
        state = json.loads(row["state_json"])
        data = state["global_compilation"]
        slot = {"slot": 1, "kind": "value", "key": "field:t0.r1.c1", "label": "Sede legale",
                "field_ids": ["t0.r1.c1"], "condition": "", "context": "Sede legale"}
        data["dialogue"]["question"] = {"kind": "clarifications", "slots": [slot],
                                        "message": "Quale sede legale devo usare?", "field_ids": []}
        s = sessions.save_state(db, row, state, "fixture_active_value_review")
    before = s["global_compilation"]["fields"][0]
    s, errors = reply(s, "VALUE", "Usa Via della Luna 8", value="Via della Luna 8")
    assert not errors and s["global_compilation"]["fields"][0] == before
    user = s["global_compilation"]["dialogue"]["user_values"]["t0.r1.c1"]
    assert user["provenance"] == "USER" and user["value"] == "Via della Luna 8"
    s = await resolve(api, s)
    ids, conditions, user_ids = dialogue.export_selection(s)
    assert ids == ["t0.r0.c1"] and user_ids == ["t0.r1.c1"]
    response = await api.post(f"{BASE}/{s['id']}/global-draft", json={"version": s["version"],
        "candidate_ids": ids, "confirmed_condition_ids": conditions, "user_candidate_ids": user_ids,
        "allow_unresolved": True})
    assert response.status_code == 200, response.text
    s = response.json()
    report = (await api.get(s["last_generation"]["downloads"]["report"])).json()
    assert [f["provenance"] for f in report["fields"]] == ["SOURCE", "USER"]
    assert report["fields"][1]["source_evidence"] == []
    doc = Document(BytesIO((await api.get(s["last_generation"]["downloads"]["docx"])).content))
    assert doc.tables[0].cell(1, 1).text == "Via della Luna 8"


@pytest.mark.anyio
async def test_partial_group_reply_updates_valid_slot_and_keeps_ambiguous_one_open(api, provider):
    s = await resolve(api, await start(api))
    with connection() as db:
        row = sessions._row(db, "alpha", s["id"])
        state = json.loads(row["state_json"])
        data = state["global_compilation"]
        for f in data["audit"]["fields"]:
            f.update(disposition="review")
        slots = [{"slot": i + 1, "kind": "value", "key": "field:" + cid,
                  "label": label, "field_ids": [cid], "condition": "", "context": label}
                 for i, (cid, label) in enumerate([
                     ("t0.r0.c1", "Denominazione sociale"), ("t0.r1.c1", "Sede legale")])]
        data["dialogue"]["question"] = {"kind": "clarifications", "slots": slots,
            "message": "Quali valori devo usare per denominazione e sede?", "field_ids": []}
        s = sessions.save_state(db, row, state, "fixture_group_review")
    message = "Denominazione: Orione S.r.l.; per la sede non sono sicuro."
    s, errors = dialogue.apply_replies("alpha", s["id"], s["version"], [
        ClarificationReply(slot=1, action="VALUE", value="Orione S.r.l.",
                           user_quote="Denominazione: Orione S.r.l."),
        ClarificationReply(slot=2, action="CLARIFY", user_quote="per la sede non sono sicuro.")
    ], message)
    assert errors == [2] and s["chat"]["auto_continue"]
    assert list(s["global_compilation"]["dialogue"]["user_values"]) == ["t0.r0.c1"]
    with connection() as db:
        data = json.loads(sessions._row(db, "alpha", s["id"])["state_json"])["global_compilation"]
    pool = dialogue.question_pool(data)
    assert [q["field_ids"] for q in pool] == [["t0.r1.c1"]]
    assert s["global_compilation"]["dialogue"]["user_turns"] == 1


@pytest.mark.anyio
async def test_narrator_can_only_remove_proposals_and_hold_persists(api, provider, monkeypatch):
    async def hold(*args, **kwargs):
        return json.dumps({"message": "Ho trovato i dati, ma la destinazione della sede è dubbia.",
                           "questions": [], "held_proposal_slots": [2]}), "simulated", {}

    monkeypatch.setattr(dialogue, "request_model_content", hold)
    s = await resolve(api, await start(api))
    assert dialogue.export_selection(s)[0] == ["t0.r0.c1"]
    assert s["global_compilation"]["dialogue"]["held_candidate_ids"] == ["t0.r1.c1"]
    assert s["global_compilation"]["proposal_count"] == 2  # Original audit unchanged.
    assert not s["last_generation"]


@pytest.mark.anyio
async def test_cannot_forge_user_export_outside_an_active_reply(api, provider):
    s = await resolve(api, await start(api))
    r = await api.post(f"{BASE}/{s['id']}/global-draft", json={"version": s["version"],
        "candidate_ids": [], "user_candidate_ids": ["t0.r0.c1"], "allow_unresolved": True})
    assert r.status_code == 422
    assert (await api.get(f"{BASE}/{s['id']}")).json()["last_generation"] is None


@pytest.mark.anyio
async def test_uncertainty_can_defer_without_becoming_a_factual_value(api, provider, monkeypatch):
    s = await conditional(api, monkeypatch)
    s, errors = reply(s, "UNKNOWN", "Forse, non ho abbastanza informazioni per confermarlo.")
    assert not errors and s["chat"]["auto_continue"]
    assert s["global_compilation"]["dialogue"]["answers"]["condition:c"]["action"] == "UNKNOWN"
    assert not s["global_compilation"]["dialogue"]["user_values"]


@pytest.mark.anyio
async def test_narration_keeps_validated_person_role_context_without_ids(api, provider):
    s = await resolve(api, await start(api))
    with connection() as db:
        data = json.loads(sessions._row(db, "alpha", s["id"])["state_json"])["global_compilation"]
    # Exercise serialization of a binding already checked by V10, not new role inference.
    data["audit"]["fields"][0]["subject_binding"] = {
        "form_evidence": {"paragraph_id": "paragraph:99", "quote": "Responsabile qualità"}}
    context = dialogue.model_context(data, [])
    assert context["document_values"][0]["required_role_context"] == "Responsabile qualità"
    assert "paragraph:99" not in json.dumps(context)
