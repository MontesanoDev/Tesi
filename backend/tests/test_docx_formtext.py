"""Native FORMTEXT destinations; deterministic structure/writes, no LLM calls."""

import json
from io import BytesIO
from zipfile import ZipFile

import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from lxml import etree
from test_compilation_catalog import native_checkbox

from app.compilation_catalog import build_compilation_catalog
from app.compilation_global import build_global_input, validate_global_proposal
from app.compilation_sessions import candidate_snapshots, validate_value
from app.docx_templates import (
    DocumentInputError,
    fill_docx,
    inspect_docx,
    is_signature_target,
    text_of,
)


def encoded(doc):
    stream = BytesIO()
    doc.save(stream)
    return stream.getvalue()


def form_text(paragraph, *, result="\u2002" * 5, instruction=" FORMTEXT ",
              name="RepeatedName", kind=None, maximum=None, enabled=True):
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    data = etree.SubElement(begin, qn("w:ffData"))
    etree.SubElement(data, qn("w:name")).set(qn("w:val"), name)
    etree.SubElement(data, qn("w:enabled")).set(qn("w:val"), str(int(enabled)))
    properties = etree.SubElement(data, qn("w:textInput"))
    if kind:
        etree.SubElement(properties, qn("w:type")).set(qn("w:val"), kind)
    if maximum is not None:
        etree.SubElement(properties, qn("w:maxLength")).set(qn("w:val"), str(maximum))
    paragraph.add_run()._r.append(begin)
    # The code may be split across runs, independently of the result.
    for part in [instruction[:4], instruction[4:]]:
        node = OxmlElement("w:instrText")
        node.text = part
        paragraph.add_run()._r.append(node)
    separator = OxmlElement("w:fldChar")
    separator.set(qn("w:fldCharType"), "separate")
    paragraph.add_run()._r.append(separator)
    for character in result:
        paragraph.add_run(character).bold = True
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    paragraph.add_run()._r.append(end)
    return begin, separator, end


def layout(doc):
    return inspect_docx(encoded(doc), require_candidates=False)


def test_duplicate_field_names_get_distinct_stable_positions_and_visible_context():
    doc = Document()
    p = doc.add_paragraph("Nome: ")
    form_text(p)
    p.add_run("; Comune: ")
    form_text(p)
    doc.add_paragraph("Ragione sociale: ")
    form_text(doc.paragraphs[-1])
    original = encoded(doc)
    parsed = inspect_docx(original)
    assert parsed.candidate_ids == {"p0.ft0", "p0.ft1", "p1.ft0"}
    catalog = build_compilation_catalog(parsed)
    assert catalog == build_compilation_catalog(inspect_docx(original))
    marked = catalog["blocks"][0]["paragraphs"][0]["text_with_fields"]
    assert marked == "Nome: [[p0.ft0]]; Comune: [[p0.ft1]]"
    assert catalog["candidates"]["p0.ft1"]["slot_context"]["before"] == "; Comune: "
    assert len(candidate_snapshots(parsed, "example", 1, "form.docx")) == 3


@pytest.mark.parametrize("empty", [False, True])
def test_writer_updates_only_result_and_preserves_field_properties_and_other_parts(empty):
    doc = Document()
    doc.sections[0].header.paragraphs[0].text = "Intestazione intatta"
    p = doc.add_paragraph("Società: ")
    form_text(p, result="" if empty else "___")
    p.add_run(" / Campo non richiesto: ")
    untouched = form_text(p)[0]
    before_untouched = etree.tostring(untouched)
    original = encoded(doc)
    parsed = inspect_docx(original)
    before_markers = [etree.tostring(n) for n in parsed.document.element.iter(qn("w:fldChar"))]
    before_styles = [etree.tostring(n) for n in parsed.document.element.iter(qn("w:rPr"))]
    value = "D'Amico & Soci <S.r.l.>\nBari\tBA"
    first = fill_docx(parsed, {"p0.ft0": value})
    assert first == fill_docx(parsed, {"p0.ft0": value})
    output = Document(BytesIO(first))
    assert output.paragraphs[1].text == (
        "Società: " + value + " / Campo non richiesto: " + "\u2002" * 5
    )
    assert [etree.tostring(n) for n in output.element.iter(qn("w:fldChar"))] == before_markers
    assert [etree.tostring(n) for n in output.paragraphs[1]._p.iter(qn("w:rPr"))] == before_styles
    remaining = list(output.element.iter(qn("w:ffData")))[1].getparent()
    assert etree.tostring(remaining) == before_untouched
    with ZipFile(BytesIO(original)) as before, ZipFile(BytesIO(first)) as after:
        assert before.testzip() is after.testzip() is None
        assert before.namelist() == after.namelist()
        assert all(before.read(n) == after.read(n) for n in before.namelist()
                   if n != "word/document.xml")
    assert parsed.original == original


def test_table_field_is_not_duplicated_as_empty_cell_and_checkbox_stays_independent():
    doc = Document()
    cells = doc.add_table(rows=1, cols=2).rows[0].cells
    cells[0].text = "Denominazione"
    p = cells[1].paragraphs[0]
    native_checkbox(p)
    p.add_run(" Società: ")
    form_text(p)
    parsed = inspect_docx(encoded(doc))
    assert parsed.candidate_ids == {"p1.ft0"}
    catalog = build_compilation_catalog(parsed)
    assert set(catalog["candidates"]) == {"p1.cb0", "p1.ft0"}
    draft = fill_docx(parsed, {"p1.ft0": "Azienda documentata"}, checkbox_values={})
    out = Document(BytesIO(draft))
    assert text_of(next(out.element.iter(qn("w:ffData")))) == ""
    assert "Azienda documentata" in out.tables[0].cell(0, 1).text
    assert next(out.element.iter(qn("w:checkBox"))) is not None


@pytest.mark.parametrize("kind", [
    "instruction", "metadata", "separator", "end", "duplicate_separator",
    "nested", "cross_paragraph", "disabled", "locked", "number", "format",
    "prefilled", "drawing", "content_control", "merge_continuation",
])
def test_unsupported_or_unsafe_controls_never_become_writable(kind):
    doc = Document()
    doc.add_paragraph("Dato ordinario: ____")
    p = doc.add_paragraph("Dato: ")
    begin, separator, end = form_text(p)
    if kind == "instruction":
        nodes = list(p._p.iter(qn("w:instrText")))
        nodes[0].text, nodes[1].text = " DATE ", ""
    elif kind == "metadata":
        begin.remove(begin.find(qn("w:ffData")))
    elif kind in {"separator", "end"}:
        node = separator if kind == "separator" else end
        node.getparent().remove(node)
    elif kind == "duplicate_separator":
        node = OxmlElement("w:fldChar")
        node.set(qn("w:fldCharType"), "separate")
        separator.addnext(node)
    elif kind == "nested":
        inner = OxmlElement("w:fldChar")
        inner.set(qn("w:fldCharType"), "begin")
        separator.addnext(inner)
        closing = OxmlElement("w:fldChar")
        closing.set(qn("w:fldCharType"), "end")
        end.addprevious(closing)
    elif kind == "cross_paragraph":
        doc.add_paragraph()._p.append(end.getparent())
    elif kind == "disabled":
        begin.find(".//" + qn("w:enabled")).set(qn("w:val"), "0")
    elif kind == "locked":
        begin.set(qn("w:fldLock"), "true")
    elif kind == "number":
        etree.SubElement(begin.find(".//" + qn("w:textInput")),
                         qn("w:type")).set(qn("w:val"), "number")
    elif kind == "format":
        etree.SubElement(begin.find(".//" + qn("w:textInput")),
                         qn("w:format")).set(qn("w:val"), "UPPERCASE")
    elif kind == "prefilled":
        results = list(p._p.iter(qn("w:t")))
        results[1].text = "Dato preesistente"
    elif kind == "drawing":
        separator.addnext(OxmlElement("w:drawing"))
    elif kind == "content_control":
        wrapper = OxmlElement("w:sdt")
        content = etree.SubElement(wrapper, qn("w:sdtContent"))
        p._p.addprevious(wrapper)
        content.append(p._p)
    else:
        table = doc.add_table(rows=2, cols=1)
        table.cell(0, 0).text = "Intestazione"
        cell = table.cell(1, 0)
        cell._tc.append(p._p)
        merge = etree.SubElement(cell._tc.get_or_add_tcPr(), qn("w:vMerge"))
        merge.set(qn("w:val"), "continue")
    parsed = layout(doc)
    assert parsed.candidate_ids == {"p0.s0"}
    with pytest.raises(DocumentInputError):
        fill_docx(parsed, {"p1.ft0": "Non autorizzato"})
    # An unrelated valid write must preserve the entire excluded structure,
    # including malformed fields, prefilled results and their original XML.
    before = [etree.tostring(n) for n in parsed.document.element.body][1:]
    output = Document(BytesIO(fill_docx(parsed, {"p0.s0": "Valore verificato"})))
    assert [etree.tostring(n) for n in output.element.body][2:] == before


def test_existing_prefilled_field_does_not_renumber_later_text_control():
    doc = Document()
    p = doc.add_paragraph("Precompilato: ")
    form_text(p, result="Valore originale")
    p.add_run("; Nuovo: ")
    form_text(p)
    assert layout(doc).candidate_ids == {"p0.ft1"}


@pytest.mark.parametrize("label", ["Firma: ", "Sottoscrizione: "])
def test_signatures_remain_write_protected(label):
    doc = Document()
    form_text(doc.add_paragraph(label))
    parsed = inspect_docx(encoded(doc))
    assert is_signature_target(parsed, "p0.ft0")
    with pytest.raises(DocumentInputError, match="firme"):
        fill_docx(parsed, {"p0.ft0": "Nome Cognome"})


def test_native_constraints_and_existing_value_checks_still_apply():
    doc = Document()
    p = doc.add_paragraph("PEC: ")
    form_text(p)
    p.add_run("; Codice: ")
    form_text(p, maximum=4)
    parsed = inspect_docx(encoded(doc))
    assert parsed.field_types["p0.ft0"] == "email"
    for target, value in [("p0.ft0", "pec-incompleta"), ("p0.ft1", "12345")]:
        with pytest.raises(DocumentInputError):
            fill_docx(parsed, {target: value})
    fill_docx(parsed, {"p0.ft0": "ufficio@azienda.demo", "p0.ft1": "1234"})


def test_global_destination_grounding_and_SOURCE_checks_reuse_existing_contract():
    doc = Document()
    form_text(doc.add_paragraph("Denominazione: "))
    parsed = inspect_docx(encoded(doc))
    quote = "Denominazione: Società Esempio"
    source = {"chunk_id": -7, "file_id": -3, "chunk_index": 0,
              "source_name": "azienda.txt", "role": "source", "scope": "global",
              "category": "company", "project_id": None, "content": quote}
    payload = {"version": 12, "uncertain_paragraph_ids": [], "facts": {
        "company": {"subject": "Società Esempio", "entity": "company", "aspect": "identity",
                    "property": "name", "value": "Società Esempio",
                    "evidence": [{"chunk_id": -7, "quote": quote}]}},
        "bindings": {}, "conditions": {}, "choice_groups": {},
        "fields": {"p0.ft0": {"label": "Denominazione", "kind": "data",
                   "disposition": "value", "property": "name", "fact_id": "company",
                   "conditions": [],
                   "form": {"paragraph_id": "paragraph:0",
                            "quote": "Denominazione: [[p0.ft0]]"}}},
        "coverage": [], "clarifications": {}}
    inputs = build_global_input(parsed, "example", [source])
    assert "p0.ft0" in inputs["form"]["candidates"]
    audit = validate_global_proposal(json.dumps(payload), parsed, "example", [source],
                                     expected_version=12)
    assert audit["fields"][0]["validation"] == "GROUNDED_PROPOSAL"
    assert audit["fields"][0]["written_value"] is None
    checked = validate_value(parsed, {"id": "p0.ft0", "label": "Denominazione"},
                             "Società Esempio", source, quote)
    assert not checked["validation_codes"]
    assert checked["written_value"] == "Società Esempio"
    rejected = validate_value(parsed, {"id": "p0.ft0", "label": "Denominazione"},
                              "Altra società", source, quote)
    assert rejected["validation_codes"]


@pytest.mark.parametrize("blanks", [("", "___"), ("", "", "___")])
def test_adjacent_empty_results_do_not_shift_the_following_field(blanks):
    doc = Document()
    p = doc.add_paragraph("Dati: ")
    for blank in blanks:
        form_text(p, result=blank)
    parsed = inspect_docx(encoded(doc))
    requested = {f"p0.ft{i}": f"Valore{i}" for i in range(len(blanks))}
    output = Document(BytesIO(fill_docx(parsed, requested)))
    assert output.paragraphs[1].text == "Dati: " + "".join(requested.values())


def test_partitioned_email_controls_keep_existing_rejection():
    doc = Document()
    p = doc.add_paragraph("PEC: ")
    form_text(p)
    p.add_run("@")
    form_text(p)
    p.add_run(".")
    form_text(p)
    parsed = inspect_docx(encoded(doc))
    assert set(parsed.field_types.values()) == {"email_parts"}
    for cid in parsed.candidate_ids:
        with pytest.raises(DocumentInputError, match="diviso in parti"):
            fill_docx(parsed, {cid: "ufficio@azienda.demo"})
