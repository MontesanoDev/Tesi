"""Structural writes only; semantic authorization is checked by the global adapter."""

from io import BytesIO
from zipfile import ZipFile

import pytest
from docx import Document
from docx.oxml.ns import qn
from lxml import etree
from test_compilation_catalog import W14, encoded, native_checkbox

from app.compilation_catalog import _checkboxes
from app.docx_templates import DocumentInputError, fill_docx, inspect_docx


def states(data):
    body = Document(BytesIO(data)).element.body
    positions = {p: i for i, p in enumerate(body.iter(qn("w:p")))}
    return {c["id"]: c["initial_checked"] for controls in _checkboxes(body, positions).values()
            for _, c in controls}


@pytest.mark.parametrize("legacy", [True, False])
@pytest.mark.parametrize("initial,target", [(False, True), (True, False)])
def test_native_state_cache_original_and_regeneration(legacy, initial, target):
    doc = Document()
    p = doc.add_paragraph("Scelta: ")
    native_checkbox(p, legacy=legacy, checked=initial)
    p.add_run(" Opzione descritta")
    original = encoded(doc)
    layout = inspect_docx(original, require_candidates=False)
    generated = fill_docx(layout, {}, checkbox_values={"p0.cb0": target})
    assert states(generated)["p1.cb0"] is target  # Draft notice is one paragraph.
    assert layout.original == original
    back = fill_docx(layout, {}, checkbox_values={"p0.cb0": initial})
    assert states(back)["p1.cb0"] is initial
    if not legacy:
        result = Document(BytesIO(generated))
        content = next(result.element.iter(qn("w:sdtContent")))
        assert "".join(n.text or "" for n in content.iter(qn("w:t"))) == (
            "☒" if target else "☐")
    with ZipFile(BytesIO(original)) as a, ZipFile(BytesIO(generated)) as b:
        assert a.namelist() == b.namelist()
        assert all(a.read(n) == b.read(n) for n in a.namelist() if n != "word/document.xml")


def test_text_and_multiple_controls_keep_ids_and_unselected_state():
    doc = Document()
    p = doc.add_paragraph()
    native_checkbox(p, checked=True)
    p.add_run(" Prima ")
    native_checkbox(p, legacy=False)
    p.add_run(" Seconda")
    doc.add_paragraph("Dato: ________")
    original = encoded(doc)
    generated = fill_docx(inspect_docx(original), {"p1.s0": "Valore"},
                          checkbox_values={"p0.cb1": True})
    assert states(generated) == {"p1.cb0": True, "p1.cb1": True}
    assert Document(BytesIO(generated)).paragraphs[-1].text == "Dato: Valore"


def test_custom_content_control_glyph_and_font_are_preserved():
    doc = Document()
    control = native_checkbox(doc.add_paragraph(), legacy=False)
    for name, value in [("checked", "2611"), ("unchecked", "2610")]:
        state = etree.SubElement(control, f"{{{W14}}}{name}State")
        state.set(f"{{{W14}}}val", value)
        state.set(f"{{{W14}}}font", "Segoe UI Symbol")
    data = fill_docx(inspect_docx(encoded(doc), require_candidates=False), {},
                     checkbox_values={"p0.cb0": True})
    result = Document(BytesIO(data))
    assert next(result.element.iter(qn("w:sdtContent"))).find(
        ".//" + qn("w:t")).text == "☑"
    assert any(n.get(qn("w:ascii")) == "Segoe UI Symbol"
               for n in result.element.iter(qn("w:rFonts")))


@pytest.mark.parametrize("values", [{"p9.cb0": True}, {"p0.cb0": 1}, {"p0.cb0": "true"}])
def test_unknown_ids_and_non_boolean_values_are_rejected(values):
    doc = Document()
    native_checkbox(doc.add_paragraph())
    with pytest.raises(DocumentInputError):
        fill_docx(inspect_docx(encoded(doc), require_candidates=False), {}, checkbox_values=values)


def test_unrelated_content_inside_control_is_never_overwritten():
    doc = Document()
    control = native_checkbox(doc.add_paragraph(), legacy=False)
    wrapper = control.getparent().getparent()
    wrapper.find(".//" + qn("w:t")).text = "Testo da conservare"
    with pytest.raises(DocumentInputError, match="visiva"):
        fill_docx(inspect_docx(encoded(doc), require_candidates=False), {},
                  checkbox_values={"p0.cb0": True})
