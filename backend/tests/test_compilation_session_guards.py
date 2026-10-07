"""Additional transport, stale-evidence and structural validation regressions."""

import json
from io import BytesIO

import httpx
import pytest
from docx import Document
from test_compilation_sessions import BASE, create, docx, resolve, simulate, source
from test_compilation_sessions import anyio_backend as anyio_backend
from test_compilation_sessions import api as api

from app import compilation_session_resolution as resolution
from app import compilation_sessions as sessions
from app.compilation_session_models import CandidateMeanings
from app.config import AISettings, use_ai_settings
from app.db import connection
from app.docx_templates import DocumentInputError


@pytest.mark.anyio
@pytest.mark.parametrize("provider", ["deepseek", "ollama"])
async def test_bounded_pipeline_with_real_json_transport_and_schema(api, monkeypatch, provider):
    state = await create(api, docx(("Denominazione sociale", "Sede legale", "Forma giuridica")))
    await source(api, "Denominazione sociale: Mapi Ingegneria S.r.l.")
    calls = []

    def handler(request):
        body = json.loads(request.content)
        user = json.loads(body["messages"][1]["content"])
        schema = user["schema_output"]
        assert "JSON" in body["messages"][0]["content"]
        if provider == "ollama":
            assert body["format"] == schema
        else:
            assert body["response_format"] == {"type": "json_object"}
        calls.append(schema["title"])
        if schema["title"] == "CandidateMeanings":
            fields = [
                {
                    "candidate_id": f["candidate_id"],
                    "classification": "data",
                    "form_quote": f["label_hint"],
                    "requirement": {
                        "name": f["label_hint"],
                        "form_quote": f["label_hint"],
                        "form_citation_id": 1,
                    },
                    "entity": "company",
                    "reason": "Requisito della cella",
                }
                for f in user["candidates"]
            ]
            reply = {"fields": fields}
        elif schema["title"] == "CompilationSourcePlan":
            reply = {"clusters": [{"requirement_ids": [1, 2, 3]}], "fields": []}
        else:
            assert schema["title"] == "CandidateMatches"
            s = user["sources"][0]
            reply = {
                "fields": [
                    {
                        "candidate_id": "t0.r0.c1",
                        "supports": [
                            {
                                "source_id": s["source_id"],
                                "quote": s["content"],
                                "value": "Mapi Ingegneria S.r.l.",
                            }
                        ],
                        "reason": "Dato esplicito nella SOURCE",
                    }
                ]
            }
        if provider == "ollama":
            data = {
                "model": "mock",
                "done": True,
                "done_reason": "stop",
                "message": {"content": json.dumps(reply)},
            }
        else:
            data = {
                "model": "mock",
                "choices": [
                    {
                        "finish_reason": "stop",
                        "message": {
                            "content": json.dumps(reply),
                        },
                    }
                ],
            }
        return httpx.Response(200, json=data)

    original_client = httpx.AsyncClient
    monkeypatch.setattr(
        httpx,
        "AsyncClient",
        lambda **kwargs: original_client(
            transport=httpx.MockTransport(handler),
            **kwargs,
        ),
    )
    with use_ai_settings(
        AISettings(api_key="mock", model="mock", provider=provider, base_url="http://mock.test")
    ):
        result = await resolution.resolve_session("alpha", state["id"], 1)
    assert calls == ["CandidateMeanings", "CompilationSourcePlan", "CandidateMatches"]
    assert result["summary"]["resolved"] == 1
    assert result["summary"]["missing"] == 2


@pytest.mark.anyio
async def test_source_changed_during_model_call_cannot_become_resolved(api, monkeypatch):
    state = await create(api, docx(("Denominazione sociale",)))
    await source(api, "Denominazione sociale: Mapi Ingegneria S.r.l.")
    simulate(monkeypatch, {"t0.r0.c1": ["Mapi Ingegneria S.r.l."]})
    match = resolution.match_candidates

    async def changing(*args):
        result = await match(*args)
        with connection() as db:
            db.execute("DELETE FROM global_document_chunks")
        return result

    monkeypatch.setattr(resolution, "match_candidates", changing)
    response = await api.post(f"{BASE}/{state['id']}/resolve", json={"version": 1})
    assert response.status_code == 409
    result = sessions.get_session("alpha", state["id"])
    assert result["status"] == "FAILED"
    assert result["fields"][0]["status"] == "PENDING"
    assert result["fields"][0]["value"] is None


@pytest.mark.anyio
async def test_form_injected_as_source_fails_the_factual_gate(api):
    state = await create(api, docx(("Denominazione sociale",)))
    field = state["fields"][0]
    field["requirement"] = {
        "name": "Denominazione sociale",
        "form_quote": "Denominazione sociale",
        "form_citation_id": 1,
    }
    with pytest.raises(ValueError, match="SOURCE"):
        resolution.validate_support(
            field, {"role": "form", "content": "Azienda S.r.l."}, "Azienda S.r.l.", "Azienda S.r.l."
        )


@pytest.mark.anyio
async def test_parser_failure_creates_no_partial_session(api, monkeypatch):
    original = await create(api)

    def failure(_):
        raise DocumentInputError("Contenuto DOCX non supportato")

    monkeypatch.setattr(sessions, "inspect_docx", failure)
    result = await api.post(BASE, json={"form_id": original["form_id"]})
    assert result.status_code == 422
    with connection() as db:
        assert db.execute("SELECT COUNT(*) FROM compilation_sessions").fetchone()[0] == 1
        assert db.execute("SELECT COUNT(*) FROM compilation_session_revisions").fetchone()[0] == 1


@pytest.mark.anyio
async def test_omitted_candidates_do_not_starve_following_batches(api, monkeypatch):
    state = await create(api, docx(tuple(f"Campo {i}" for i in range(20))))
    simulate(monkeypatch, classify_hook=lambda fields: CandidateMeanings(fields=[]))
    first = await resolve(api, state)
    second = await resolve(api, first)
    assert first["summary"]["pending"] == second["summary"]["pending"] == 20
    ids1 = first["last_resolution"]["candidate_ids"]
    ids2 = second["last_resolution"]["candidate_ids"]
    assert ids2[:8] == [f"t0.r{i}.c1" for i in range(12, 20)]
    assert set(ids1) | set(ids2) == {f"t0.r{i}.c1" for i in range(20)}


@pytest.mark.anyio
async def test_paragraph_placeholders_and_stable_ids_survive_export(api):
    document = Document()
    document.add_paragraph("Denominazione sociale: {{azienda}}; Sede legale: {{sede}}")
    stream = BytesIO()
    document.save(stream)
    original = stream.getvalue()
    state = await create(api, original)
    assert [f["candidate_id"] for f in state["fields"]] == ["p0.s0", "p0.s1"]
    assert state["fields"][1]["location"]["slot"] == 2
    url = f"{BASE}/{state['id']}"
    state = (
        await api.patch(
            url + "/fields",
            json={
                "version": 1,
                "fields": [
                    {"field_id": "p0.s0", "value": "Mapi"},
                    {"field_id": "p0.s1", "value": "Bari"},
                ],
            },
        )
    ).json()
    result = await api.post(url + "/finalize", json={"version": state["version"]})
    assert result.status_code == 200, result.text
    draft = await api.get(result.json()["last_generation"]["downloads"]["docx"])
    assert Document(BytesIO(draft.content)).paragraphs[1].text == (
        "Denominazione sociale: Mapi; Sede legale: Bari"
    )


@pytest.mark.anyio
async def test_concurrent_user_update_rolls_back_outdated_generation(api, monkeypatch):
    state = await create(api)
    real_fill = sessions.fill_docx
    from app.compilation_session_models import UserFieldInput

    def fill_with_concurrent_update(layout, values):
        sessions.update_fields(
            "alpha",
            state["id"],
            1,
            [
                UserFieldInput(field_id="t0.r0.c1", value="Correzione utente"),
            ],
        )
        return real_fill(layout, values)

    monkeypatch.setattr(sessions, "fill_docx", fill_with_concurrent_update)
    response = await api.post(
        f"{BASE}/{state['id']}/finalize", json={"version": 1, "allow_unresolved": True}
    )
    assert response.status_code == 409
    result = sessions.get_session("alpha", state["id"])
    assert result["fields"][0]["value"] == "Correzione utente"
    assert result["last_generation"] is None
    with connection() as db:
        assert db.execute("SELECT COUNT(*) FROM document_compilations").fetchone()[0] == 0


@pytest.mark.anyio
async def test_decorative_candidate_is_not_a_missing_obligation(api, monkeypatch):
    state = await create(api, docx(("",), context="Spazio separatore"))

    def classify(fields):
        return CandidateMeanings(
            fields=[
                {
                    "candidate_id": f["id"],
                    "classification": "decorative",
                    "requirement": None,
                    "form_quote": "Spazio separatore",
                    "reason": "Cella priva di etichetta o dato richiesto",
                }
                for f in fields
            ]
        )

    simulate(monkeypatch, classify_hook=classify)
    result = await resolve(api, state)
    assert result["summary"]["not_applicable"] == result["summary"]["total"]
    assert result["summary"]["missing"] == result["summary"]["resolved"] == 0
    assert result["status"] == "READY"
    assert all(f["provenance"] == "FORM" and f["value"] is None for f in result["fields"])
