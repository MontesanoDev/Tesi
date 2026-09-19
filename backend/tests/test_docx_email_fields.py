import json
from io import BytesIO

import pytest
from docx import Document

from app import document_compilation as compilation
from app.docx_templates import DocumentInputError, fill_docx, inspect_docx


@pytest.fixture
def anyio_backend():
    return "asyncio"


def email_layout(text="E-mail: ....\u2026..", *, table=False):
    doc = Document()
    if table:
        cells = doc.add_table(rows=1, cols=2).rows[0].cells
        cells[0].text = text
    else:
        doc.add_paragraph(text)
    buffer = BytesIO()
    doc.save(buffer)
    return inspect_docx(buffer.getvalue())


def sources(content):
    return compilation.CompilationSources([{
        "id": "company:1", "document_id": 1, "source_name": "azienda.md",
        "chunk_index": 0, "content": content, "scope": "company", "source_kind": "source",
    }], 1, len(content))


def proposal(layout, value, quote, *, target=None, label="Recapito"):
    return {
        "cell_id": target or next(iter(layout.candidate_ids)), "label": label,
        "entity": "company", "kind": "data", "status": "proposed", "value": value,
        "evidence": [{"source_id": "company:1", "quote": quote}],
        "reason": "Dato presente nella fonte",
    }


def response(*fields):
    return json.dumps({"fields": fields, "warnings": []})


@pytest.mark.parametrize("label", [
    "E-mail:", "email", "PEC", "P.E.C.", "Indirizzo di posta elettronica certificata:",
    "Email aziendale", "PEC (obbligatoria)",
])
@pytest.mark.parametrize("table", [False, True])
def test_type_comes_from_template_not_model_label(label, table):
    layout = email_layout(label if table else f"{label} ....\u2026..", table=table)
    target = next(iter(layout.candidate_ids))
    assert layout.field_types[target] == "email"
    quote = "PEC: impresa@pec.demo"
    field = compilation.validate_proposals(
        response(proposal(layout, "pec", quote, label="Altro dato")), layout, sources(quote),
    )["fields"][0]
    assert field["written_value"] is None
    assert field["validation_codes"] == ["invalid_email"]
    assert field["repairable"]
    with pytest.raises(DocumentInputError, match="Email/PEC"):
        fill_docx(layout, {target: "pec"})


@pytest.mark.parametrize("value", [
    "segreteria", "impresa.pec.demo", "impresa@pec", "a@@pec.demo", "a@pec..demo",
    "a..b@pec.demo", "a@-pec.demo", "a@pec_demo.it", "a@pec.demo; b@pec.demo",
    "Persona <a@pec.demo>", "a@pec.demo\nNota", "a@pec.demo...", "a @pec.demo",
])
def test_invalid_or_multiple_recapiti_stay_unwritten_even_when_literal(value):
    layout = email_layout()
    quote = f"Recapito fornito: {value}"
    field = compilation.validate_proposals(
        response(proposal(layout, value, quote)), layout, sources(quote),
    )["fields"][0]
    assert "invalid_email" in field["validation_codes"]
    assert field["written_value"] is None


@pytest.mark.parametrize("value", [
    "impresa@pec.demo", "segreteria+gare@azienda.example", "Nome.Cognome@Azienda.IT",
    "citt\u00e0@azienda.example",
])
def test_complete_address_is_preserved_not_normalized_and_dns_is_disabled(monkeypatch, value):
    import email_validator.deliverability

    def no_dns(*args, **kwargs):
        pytest.fail("La compilazione non deve interrogare il DNS")

    monkeypatch.setattr(email_validator.deliverability, "validate_email_deliverability", no_dns)
    layout = email_layout("PEC: ....\u2026..; Fine.")
    quote = f"PEC: {value}."
    field = compilation.validate_proposals(
        response(proposal(layout, value, quote)), layout, sources(quote),
    )["fields"][0]
    assert field["written_value"] == value
    doc = Document(BytesIO(fill_docx(layout, {field["cell_id"]: value})))
    assert doc.paragraphs[1].text == f"PEC: {value}; Fine."


@pytest.mark.parametrize("value,full", [
    ("ferri@pec.demo", "luca.ferri@pec.demo"),
    ("luca@pec.it", "luca@pec.it.example"),
    ("luca@pec.it", "luca@pec.it-extra.example"),
])
def test_valid_mailbox_substring_is_not_a_full_address_in_evidence(value, full):
    layout = email_layout()
    quote = f"PEC: {full}"
    field = compilation.validate_proposals(
        response(proposal(layout, value, quote)), layout, sources(quote),
    )["fields"][0]
    assert field["written_value"] is None
    assert field["validation_codes"] == ["partial_email_evidence"]
    assert field["repairable"]


@pytest.mark.parametrize("text", [
    "E-mail: ___@___.___", "PEC: {{utente}}@{{dominio}}.{{suffisso}}",
    "PEC: nome@____", "E-mail: ____@dominio.example", "PEC: ____ . ____",
    "PEC: nome@dominio.____",
])
def test_explicit_address_parts_are_all_blocked_including_direct_writer(text):
    layout = email_layout(text)
    assert set(layout.field_types.values()) == {"email_parts"}
    assert set(layout.field_types) == layout.candidate_ids
    quote = "PEC: impresa@pec.demo"
    for target in layout.candidate_ids:
        field = compilation.validate_proposals(
            response(proposal(layout, "impresa@pec.demo", quote, target=target)),
            layout, sources(quote),
        )["fields"][0]
        assert field["validation_codes"] == ["unsupported_email_layout"]
        assert not field["repairable"]
        with pytest.raises(DocumentInputError, match="separatore|separatori"):
            fill_docx(layout, {target: "impresa@pec.demo"})


def test_type_does_not_leak_to_following_tax_id_or_phone_fields():
    layout = email_layout("Email: ___ CF: ___ Telefono: ___")
    assert layout.field_types == {"p0.s0": "email"}
    quote = "CF: DEMO123. Telefono: 080123456"
    raw = proposal(layout, "DEMO123", quote, target="p0.s1")
    field = compilation.validate_proposals(response(raw), layout, sources(quote))["fields"][0]
    assert field["written_value"] == "DEMO123"


def test_named_marker_and_column_header_are_typed():
    assert email_layout("{{email}}; altro: {{nome}}").field_types == {"p0.s0": "email"}
    doc = Document()
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Persona"
    table.cell(0, 1).text = "Email"
    table.cell(1, 0).text = "Anna"
    buffer = BytesIO()
    doc.save(buffer)
    assert inspect_docx(buffer.getvalue()).field_types == {"t0.r1.c1": "email"}


@pytest.mark.anyio
@pytest.mark.parametrize("fixed", [True, False])
async def test_email_uses_only_one_targeted_repair_and_keeps_other_fields(monkeypatch, fixed):
    layout = email_layout("Email: ....\u2026...; Nome: ____")
    quote = "Email: ufficio@impresa.demo. Nome: Impresa"
    bad = proposal(layout, "ufficio", quote, target="p0.s0")
    good = proposal(layout, "Impresa", quote, target="p0.s1", label="Nome")
    calls = []

    async def model(prompt):
        payload = json.loads(prompt)
        calls.append(payload)
        if len(calls) == 1:
            return response(bad, good), "test-model", 10
        assert payload["target_ids"] == ["p0.s0"]
        assert payload["correction_request"][0]["validation_codes"] == ["invalid_email"]
        value = "ufficio@impresa.demo" if fixed else "ufficio"
        return response({**bad, "value": value}), "test-model", 12

    monkeypatch.setattr(compilation, "request_field_proposals", model)
    report, _, tokens, execution = await compilation.compile_field_batches(
        layout, sources(quote), "Progetto diverso", "Solo dati disponibili",
    )
    assert len(calls) == 2
    assert tokens == 22
    assert execution["repaired_fields"] == int(fixed)
    assert report["fields"][1]["written_value"] == "Impresa"
    first = report["fields"][0]
    assert first["written_value"] == ("ufficio@impresa.demo" if fixed else None)
    assert first["repair"]["status"] == ("corrected" if fixed else "unresolved")
    values = {f["cell_id"]: f["written_value"] for f in report["fields"] if f["written_value"]}
    output = Document(BytesIO(fill_docx(layout, values)))
    assert output.paragraphs[1].text == (
        "Email: ufficio@impresa.demo; Nome: Impresa" if fixed
        else "Email: ....\u2026...; Nome: Impresa"
    )
