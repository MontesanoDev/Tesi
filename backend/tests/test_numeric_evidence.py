import json
from io import BytesIO

import pytest
from docx import Document

from app import document_compilation as compilation
from app.docx_templates import fill_docx, inspect_docx


@pytest.fixture
def anyio_backend():
    return "asyncio"


def inputs(value, source, quote=None):
    document = Document()
    document.add_paragraph("Identificativo: ____; Denominazione: ____")
    buffer = BytesIO()
    document.save(buffer)
    layout = inspect_docx(buffer.getvalue())
    sources = compilation.CompilationSources([{
        "id": "company:1", "document_id": 1, "source_name": "azienda.txt",
        "chunk_index": 0, "content": source, "scope": "company", "source_kind": "source",
    }], 1, len(source))
    proposal = {
        "cell_id": "p0.s0", "label": "Identificativo", "entity": "company",
        "kind": "data", "status": "proposed", "value": value,
        "evidence": [{"source_id": "company:1", "quote": quote or source}],
        "reason": "Dato presente nella fonte",
    }
    return layout, sources, proposal


def response(*fields):
    return json.dumps({"fields": fields, "warnings": []})


@pytest.mark.parametrize("value,source,quote", [
    ("1234567890", "Partita IVA: 01234567890.", None),
    ("1234567890", "Partita IVA: 01234567890.", "1234567890"),
    ("0123456789", "Partita IVA: 01234567890.", None),
    ("23456789", "Partita IVA: 01234567890.", "23456789"),
    ("01234567890", "Identificativo: IT01234567890.", None),
    ("ABC123", "Codice: ABC123Z.", "ABC123"),
    ("123", "Codice: ABC123Z. Altro numero: 123.", "ABC123Z"),
    ("IVA 1234567890", "Partita IVA 12345678901.", None),
    ("1234567890.", "Partita IVA: 01234567890.", None),
])
def test_partial_numeric_tokens_are_blocked_even_when_the_quote_is_trimmed(value, source, quote):
    layout, sources, proposal = inputs(value, source, quote)
    field = compilation.validate_proposals(response(proposal), layout, sources)["fields"][0]
    assert field["written_value"] is None
    assert field["validation_codes"] == ["partial_numeric_evidence"]


@pytest.mark.parametrize("value,source,quote", [
    ("01234567890", "Partita IVA: 01234567890.", None),
    ("01234567890", "Partita IVA: (01234567890).", "01234567890"),
    ("IT01234567890", "Identificativo: IT01234567890.", None),
    ("abc123", "Codice: ABC123.", "ABC123"),
    ("12", "Data: 25/12/2026.", "25/12/2026"),
    ("20", "Civico 120; numero autonomo 20.", "20"),
    ("Societa 2", "Denominazione:\nSocieta\u00a02.", "Denominazione: Societa 2."),
    ("Esempio", "Denominazione: Esempio S.r.l.", None),
])
def test_complete_tokens_and_explicit_date_components_keep_their_literal_values(
    value, source, quote,
):
    layout, sources, proposal = inputs(value, source, quote)
    field = compilation.validate_proposals(response(proposal), layout, sources)["fields"][0]
    assert field["written_value"] == value
    assert field["validation_codes"] == []
    document = Document(BytesIO(fill_docx(layout, {field["cell_id"]: value})))
    assert document.paragraphs[1].text == f"Identificativo: {value}; Denominazione: ____"


@pytest.mark.anyio
@pytest.mark.parametrize("fixed", [True, False])
async def test_numeric_validation_in_one_call_preserves_unrelated_fields(monkeypatch, fixed):
    layout, sources, bad = inputs("1234567890", "Partita IVA: 01234567890. Societa: Esempio.")
    if fixed:
        bad["value"] = "01234567890"
    good = {**bad, "cell_id": "p0.s1", "label": "Denominazione", "value": "Esempio"}
    calls = []

    async def model(prompt, **_options):
        payload = json.loads(prompt)
        calls.append(payload)
        assert payload["target_ids"] == ["p0.s0", "p0.s1"]
        return response(bad, good), "test", 10

    monkeypatch.setattr(compilation, "request_field_proposals", model)
    report, _, tokens, execution = await compilation.compile_fields_once(
        layout, sources, "Progetto", "",
    )
    assert len(calls) == execution["requests"] == 1
    assert tokens == 10
    assert report["fields"][0]["validation_codes"] == (
        [] if fixed else ["partial_numeric_evidence"]
    )
    values = {f["cell_id"]: f["written_value"] for f in report["fields"] if f["written_value"]}
    document = Document(BytesIO(fill_docx(layout, values)))
    expected = "01234567890" if fixed else "____"
    assert document.paragraphs[1].text == f"Identificativo: {expected}; Denominazione: Esempio"
