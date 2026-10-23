"""Rev 3 candidate, offline: an option's condition also covers the option's continuation (R1).

Counterexample recorded in rev 2 (minervino-1, `p37.*`): a SOURCE value written in the
paragraphs that follow «☐ Professionista appartenente allo studio associato» although the model's
condition on that option was open and its scope listed only the first paragraph.
"""

import json
from copy import deepcopy
from io import BytesIO

import pytest
import test_compilation_resilient_segments as rev2
from docx import Document
from test_compilation_catalog import native_checkbox
from test_compilation_global_sessions import api, provider

from app import compilation_global_dialogue as dialogue
from app import compilation_global_sessions as adapter
from app import compilation_scope_review as scope
from app.compilation_branch_scope import continuations, enabled, propagate
from app.compilation_segments import applicability_states, unverified_fields
from app.docx_templates import inspect_docx

__all__ = ["api", "provider"]

STUDIO = " Professionista appartenente allo studio associato avente la denominazione"
COMPANY = " Legale rappresentante della società di ingegneria"


@pytest.fixture(autouse=True)
def candidate_runtime(monkeypatch):
    monkeypatch.setenv("MAPI_GENERATION_REVISION", "3")
    monkeypatch.setenv("MAPI_DIALOGUE_REVISION", "1")
    monkeypatch.setenv("MAPI_DERIVATION_REVISION", "1")
    monkeypatch.setenv("MAPI_BRANCH_SCOPE_REVISION", "1")


def option(doc, text, *, glyph=None):
    paragraph = doc.add_paragraph()
    if glyph:
        paragraph.add_run(glyph + text)
    else:
        native_checkbox(paragraph)
        paragraph.add_run(text)
    return paragraph


def branches_docx(*, glyph=None, after=()):
    """Two exclusive options, each followed by its own data paragraphs (Minervino layout)."""
    doc = Document()
    doc.add_paragraph().add_run("IN QUALITÀ DI").bold = True       # paragraph:0
    option(doc, STUDIO, glyph=glyph)                                # paragraph:1
    doc.add_paragraph("Dato 2: ____________")                       # paragraph:2
    doc.add_paragraph("Dato 3: ____________")                       # paragraph:3
    option(doc, COMPANY, glyph=glyph)                               # paragraph:4
    doc.add_paragraph("Dato 5: ____________")                       # paragraph:5
    for kind in after:
        if kind == "empty":
            doc.add_paragraph("")
        elif kind == "separator":
            doc.add_paragraph("*** *** ***")
        elif kind == "dots":
            doc.add_paragraph("..............................")
        elif kind == "bold":
            doc.add_paragraph().add_run("DICHIARA ALTRESÌ").bold = True
        elif kind == "table":
            doc.add_table(rows=1, cols=2).rows[0].cells[0].text = "Dato tabella:"
        elif kind == "bracket":
            doc.add_paragraph("[ ] di non avvalersi dei requisiti di altro operatore")
        doc.add_paragraph("Dato finale: ____________")
    stream = BytesIO()
    doc.save(stream)
    return stream.getvalue()


def catalog_of(data):
    return rev2.model_form(data)


def ids_in(catalog, paragraph_id):
    return [c for c, x in catalog["candidates"].items() if x["context_id"] == paragraph_id]


def text_of(catalog, paragraph_id):
    return next(p["text"] for b in catalog["blocks"] for p in b["paragraphs"]
                if p["id"] == paragraph_id)


# Region structure (catalog only).


@pytest.mark.parametrize("glyph", [None, "☐"])
def test_a_run_goes_from_the_option_to_the_next_option(glyph):
    catalog = catalog_of(branches_docx(glyph=glyph))
    runs = continuations(catalog)
    assert runs["paragraph:1"] == ids_in(catalog, "paragraph:2") + ids_in(catalog, "paragraph:3")
    assert runs["paragraph:4"] == ids_in(catalog, "paragraph:5")
    assert "paragraph:0" not in runs  # A heading is not an option.


@pytest.mark.parametrize("boundary", ["empty", "separator", "bold", "table"])
def test_structural_boundaries_end_the_last_run(boundary):
    catalog = catalog_of(branches_docx(after=[boundary]))
    final = next(c for c, x in catalog["candidates"].items()
                 if text_of(catalog, x["context_id"]).startswith("Dato finale"))
    assert final not in continuations(catalog)["paragraph:4"]


def test_a_dotted_fill_in_line_continues_the_run_and_brackets_are_not_options():
    catalog = catalog_of(branches_docx(after=["dots"]))
    dotted = [c for c, x in catalog["candidates"].items()
              if text_of(catalog, x["context_id"]).startswith("....")]
    assert dotted and set(dotted) <= set(continuations(catalog)["paragraph:4"])
    # «[ ]» stays outside V14 and R1 alike: neither an option nor a boundary (a longer run
    # can only withhold more, never admit a write).
    catalog = catalog_of(branches_docx(after=["bracket"]))
    final = next(c for c, x in catalog["candidates"].items()
                 if text_of(catalog, x["context_id"]).startswith("Dato finale"))
    assert len(continuations(catalog)) == 2 and final in continuations(catalog)["paragraph:4"]


# Gate behaviour on a Minervino-like state.


def state(conditions, *, revision=3, branch=1):
    catalog = catalog_of(branches_docx())
    fields = []
    for paragraph in ("paragraph:2", "paragraph:3", "paragraph:5"):
        for cid in ids_in(catalog, paragraph):
            fields.append({"candidate_id": cid, "label": "Codice", "conditions": [],
                           "validation": "GROUNDED_PROPOSAL", "kind": "data",
                           "form_evidence": {"paragraph_id": paragraph,
                                             "quote": text_of(catalog, paragraph)}})
    data = {"generation_revision": revision, "catalog": catalog, "dialogue": {"answers": {}},
            "audit": {"version": 12, "conditions": {}, "fields": fields}}
    if branch:
        data["branch_scope_revision"] = branch
    for key, (anchor, result, declared) in conditions.items():
        data["audit"]["conditions"][key] = {
            "result": result, "reason": "Ramo dell'opzione",
            "form_evidence": {"paragraph_id": anchor, "quote": text_of(catalog, anchor)},
            "applies_to": sorted({*ids_in(catalog, anchor), *declared})}
        for field in fields:
            if field["candidate_id"] in declared:
                field["conditions"].append(key)
    return data


def minervino_like(**kwargs):
    """Open condition on the studio option; the model's scope lists only paragraph 2."""
    catalog = catalog_of(branches_docx())
    data = state({"c35": ("paragraph:1", "unknown", ids_in(catalog, "paragraph:2"))}, **kwargs)
    return data, ids_in(catalog, "paragraph:2"), ids_in(catalog, "paragraph:3"), ids_in(
        catalog, "paragraph:5")


def extended(data):
    if enabled(data):
        data["audit"] = propagate(data["audit"], data["catalog"])
    return data


def receipt(data, key, ids, action="CONDITION_TRUE"):
    record = {"action": action, "applies": action == "CONDITION_TRUE", "user_quote": "Sì",
              "affected_fields": sorted(ids), "provenance": "USER"}
    record["scope_validation"] = scope.proof(scope.target(data, key, sorted(ids)), record)
    data["dialogue"]["answers"][key] = record
    return record


def test_the_recorded_counterexample_without_r1_and_its_closure_with_r1():
    data, declared, forgotten, other = minervino_like(branch=None)
    assert not enabled(data)
    blocked = scope.blocked_source_fields(extended(data))
    assert set(declared) <= blocked and not set(forgotten) & blocked  # Rev 2/3 counterexample.
    data, declared, forgotten, other = minervino_like()
    data = extended(data)
    assert set(declared + forgotten) <= scope.blocked_source_fields(data)
    assert set(forgotten) <= unverified_fields(data)
    assert not set(other) & unverified_fields(data)
    assert data["audit"]["branch_scope"] == {"revision": 1, "origin": "BACKEND",
        "option_continuations": {"c35": {"paragraph_id": "paragraph:1",
                                         "candidate_ids": forgotten}}}
    assert set(forgotten) <= set(data["audit"]["conditions"]["c35"]["applies_to"])
    states = applicability_states(data)
    assert all(states[c]["state"] == "CONDIZIONATA"
               and "condizione_aperta:c35" in states[c]["reasons"] for c in forgotten)
    assert states[other[0]]["state"] == "VERIFICATA"


@pytest.mark.parametrize("answer", ["CONDITION_FALSE", "SKIP", "UNKNOWN"])
def test_a_negative_or_ambiguous_user_answer_keeps_the_continuation_out(answer):
    data, declared, forgotten, _ = minervino_like()
    data = extended(data)
    record = {"action": answer, "applies": False if answer == "CONDITION_FALSE" else None,
              "user_quote": "No" if answer == "CONDITION_FALSE" else "Non so",
              "affected_fields": declared + forgotten, "provenance": "USER"}
    data["dialogue"]["answers"]["condition:c35"] = record
    assert set(declared + forgotten) <= scope.blocked_source_fields(data)


def test_a_verified_true_answer_on_the_extended_scope_admits_the_continuation():
    data, declared, forgotten, _ = minervino_like()
    data = extended(data)
    receipt(data, "condition:c35", declared + forgotten)
    assert not set(declared + forgotten) & scope.blocked_source_fields(data)


def test_a_receipt_bound_to_the_model_scope_only_does_not_spread_to_the_continuation():
    data, declared, forgotten, _ = minervino_like()
    receipt(data, "condition:c35", declared)  # Given before the extension (stale scope).
    data = extended(data)
    blocked = scope.blocked_source_fields(data)
    assert set(forgotten) <= blocked
    assert not set(declared) & blocked  # Still valid for the destinations it was checked on.


def test_a_revoked_answer_or_a_changed_condition_blocks_again():
    data, declared, forgotten, _ = minervino_like()
    data = extended(data)
    receipt(data, "condition:c35", declared + forgotten)
    changed = deepcopy(data)
    changed["audit"]["conditions"]["c35"]["reason"] = "Altro significato"
    assert set(forgotten) <= scope.blocked_source_fields(changed)
    receipt(data, "condition:c35", declared + forgotten, action="CONDITION_FALSE")
    assert set(forgotten) <= scope.blocked_source_fields(data)


def test_a_true_answer_on_the_other_branch_never_admits_this_continuation():
    catalog = catalog_of(branches_docx())
    studio, company = ids_in(catalog, "paragraph:2"), ids_in(catalog, "paragraph:5")
    data = extended(state({"c35": ("paragraph:1", "unknown", studio),
                           "c39": ("paragraph:4", "unknown", [])}))
    receipt(data, "condition:c39", company)
    blocked = scope.blocked_source_fields(data)
    assert not set(company) & blocked
    assert set(studio + ids_in(catalog, "paragraph:3")) <= blocked


def test_the_applicable_branch_keeps_its_source_writes():
    catalog = catalog_of(branches_docx())
    company = ids_in(catalog, "paragraph:5")
    data = extended(state({"c39": ("paragraph:4", "true", [])}))
    assert data["audit"]["branch_scope"]["option_continuations"]["c39"]["candidate_ids"] == company
    assert not set(company) & unverified_fields(data)
    assert applicability_states(data)[company[0]]["state"] == "VERIFICATA"


def test_r1_never_creates_conditions_or_reads_headings():
    """Documented limit: without an option-anchored condition nothing changes (rev 2, 18/25)."""
    catalog = catalog_of(branches_docx())
    for conditions in ({}, {"h": ("paragraph:0", "unknown", [])}):
        data = state(conditions)
        before = deepcopy(data["audit"])
        data = extended(data)
        assert data["audit"]["branch_scope"]["option_continuations"] == {}
        assert data["audit"]["fields"] == before["fields"]
        assert not set(ids_in(catalog, "paragraph:3")) & unverified_fields(data)


def test_propagation_is_idempotent_and_skips_rejected_conditions():
    data, _, forgotten, _ = minervino_like()
    once = propagate(deepcopy(data["audit"]), data["catalog"])
    twice = propagate(deepcopy(once), data["catalog"])
    assert [f["conditions"] for f in twice["fields"]] == [f["conditions"] for f in once["fields"]]
    data["audit"]["conditions"]["c35"]["validation_errors"] = ["Condizione non valida"]
    assert propagate(data["audit"], data["catalog"])["branch_scope"][
        "option_continuations"] == {}


@pytest.mark.parametrize("revision", [None, 1, 2])
def test_only_revision3_with_the_flag_is_affected(revision):
    data, _, forgotten, _ = minervino_like(revision=revision)
    assert not enabled(data)
    assert not set(forgotten) & scope.blocked_source_fields(extended(data))


def test_new_session_profile_requires_both_opt_ins(monkeypatch):
    layout = inspect_docx(branches_docx())
    assert adapter.initial_state(layout)["branch_scope_revision"] == 1
    monkeypatch.setenv("MAPI_GENERATION_REVISION", "2")
    assert "branch_scope_revision" not in adapter.initial_state(layout)
    monkeypatch.delenv("MAPI_GENERATION_REVISION")
    assert "branch_scope_revision" not in adapter.initial_state(layout)
    monkeypatch.setenv("MAPI_GENERATION_REVISION", "3")
    monkeypatch.delenv("MAPI_BRANCH_SCOPE_REVISION")
    assert "branch_scope_revision" not in adapter.initial_state(layout)


# Whole session with a simulated provider: analysis, question, export.


def wrong_branch_writer(inp, ids):
    """Valid V12 JSON: open condition on the studio option, scope limited to paragraph 2,
    values also in paragraph 3 (forgotten) and in the company branch."""
    targets = [c for c in ids if rev2.paragraph_of(inp, c)["text"].startswith("Dato")]
    first, forgotten, company = targets[0], targets[1], targets[2]
    quote = rev2.paragraph_of(inp, first)
    studio = next(p for b in inp["form"]["blocks"] for p in b["paragraphs"]
                  if p["id"] == "paragraph:1")
    condition = rev2.condition(inp, "paragraph:1", studio["text"].strip(),
                               scope=[{"kind": "candidate", "id": first}])
    assert quote["id"] == "paragraph:2"
    return json.dumps(rev2.envelope(inp, {first: rev2.slot(inp, first, conditions=["c"]),
                                          forgotten: rev2.slot(inp, forgotten),
                                          company: rev2.slot(inp, company)},
                                    conditions={"c": condition}))


@pytest.mark.anyio
@pytest.mark.parametrize("flag", ["1", None])
async def test_session_with_r1_asks_before_writing_the_forgotten_continuation(
    api, provider, monkeypatch, flag,
):
    if flag is None:
        monkeypatch.delenv("MAPI_BRANCH_SCOPE_REVISION")
    model = rev2.Model({0: [wrong_branch_writer], 1: [wrong_branch_writer]})
    session = await rev2.started(api, monkeypatch, branches_docx(), model)
    response = await rev2.analyze(api, session)
    assert response.status_code == 200, response.text
    session = response.json()
    data = rev2.state_of(session)["global_compilation"]
    first, forgotten, company = "p2.s0", "p3.s0", "p5.s0"
    ids, _, _ = dialogue.export_selection(session)
    assert company in ids and first not in ids
    if flag is None:
        assert "branch_scope" not in data["audit"]
        assert forgotten in ids  # Rev 3 alone: the recorded counterexample stays exportable.
        return
    assert data["branch_scope_revision"] == 1
    assert data["audit"]["branch_scope"]["option_continuations"]["c"]["candidate_ids"] == [
        forgotten]
    assert forgotten not in ids
    question = next(q for q in dialogue.question_pool(data) if q["key"] == "condition:c")
    assert {first, forgotten} <= set(question["field_ids"])
    refused = await rev2.draft(api, session, [*ids, forgotten])
    assert refused.status_code == 409
    exported = await rev2.draft(api, session, ids)
    assert exported.status_code == 200, exported.text
