"""Organization check: a named role attestation at another organization never grounds a person."""

import json
from io import BytesIO

import pytest
from docx import Document
from test_compilation_global_sessions import BASE, api, provider, resolve
from test_compilation_sessions import source

from app import compilation_global_dialogue as dialogue
from app import compilation_global_sessions as adapter
from app import compilation_sessions as sessions
from app.compilation_attribution_safety import (
    apply_attribution_safety,
    foreign_attestations,
    own_organization,
)
from app.compilation_global import validate_global_proposal
from app.compilation_source_reuse import reuse_selected_person
from app.docx_templates import inspect_docx

__all__ = ["api", "provider"]

CF = "RSSMRA80A01H501U"
OPERATOR = "Alfa Ingegneria S.r.l."
HEADER = "la quale ha come Direttore Tecnico:"
COMPANY_LINE = f"- Denominazione: {OPERATOR}"


def src(chunk, content, category="company"):
    return {"chunk_id": chunk, "file_id": chunk - 20, "chunk_index": 0,
            "source_name": f"s{-chunk}.md", "role": "source", "scope": "global",
            "category": category, "project_id": None, "content": content}


def named(owner):
    return f"- **Mario Bianchi:** direttore tecnico {owner}, codice fiscale {CF}."


def form(with_operator=True):
    doc = Document()
    doc.add_paragraph(HEADER)
    doc.add_paragraph("C.F. ____________")
    if with_operator:
        doc.add_paragraph("Denominazione ____________")
    stream = BytesIO()
    doc.save(stream)
    return inspect_docx(stream.getvalue())


def proposal(line, chunk=-1, company_chunk=-9, with_operator=True):
    facts = {"cf": {"subject": "Mario Bianchi", "entity": "person", "aspect": "identity",
                    "property": "codice_fiscale", "value": CF,
                    "evidence": [{"chunk_id": chunk, "quote": f"codice fiscale {CF}"}]},
             "co": {"subject": OPERATOR, "entity": "company", "aspect": "identity",
                    "property": "denominazione", "value": OPERATOR,
                    "evidence": [{"chunk_id": company_chunk, "quote": COMPANY_LINE}]}}
    fields = {"p1.s0": {"label": "C.F.", "kind": "data", "disposition": "value", "conditions": [],
                        "form": {"paragraph_id": "paragraph:1", "quote": "C.F. [[p1.s0]]"},
                        "fact_id": "cf", "property": "codice_fiscale", "binding_id": "b",
                        "reason": "dato SOURCE"}}
    if with_operator:
        fields["p2.s0"] = {"label": "Denominazione", "kind": "data", "disposition": "value",
                           "conditions": [], "fact_id": "co", "property": "denominazione",
                           "form": {"paragraph_id": "paragraph:2",
                                    "quote": "Denominazione [[p2.s0]]"},
                           "binding_id": None, "reason": "dato SOURCE"}
    return {"version": 12, "uncertain_paragraph_ids": [], "choice_groups": {},
            "clarifications": {}, "coverage": [], "conditions": {}, "facts": facts,
            "fields": fields,
            "bindings": {"b": {"form": {"paragraph_id": "paragraph:0", "quote": HEADER},
                               "required_role": "direttore tecnico",
                               "source_role": "direttore tecnico",
                               "evidence": [{"chunk_id": chunk, "quote": line}]}}}


def audited(owner, *, with_operator=True, category="company"):
    line = named(owner)
    sources = [src(-1, line, category), src(-9, COMPANY_LINE)]
    audit = validate_global_proposal(json.dumps(proposal(line, with_operator=with_operator)),
                                     form(with_operator), "alpha", sources)
    return audit, {s["chunk_id"]: s for s in sources}


def cf_field(audit):
    return next(f for f in audit["fields"] if f["candidate_id"] == "p1.s0")


@pytest.mark.parametrize("owner, with_operator", [
    ("di Beta Costruzioni S.r.l.", True),
    ("di Beta Costruzioni S.r.l.", False),  # No grounded company: an explicit name cannot match.
    ("di cantiere", True),
])
def test_explicit_foreign_organization_removes_the_grounding(owner, with_operator):
    audit, source_map = audited(owner, with_operator=with_operator)
    assert cf_field(audit)["validation"] == "GROUNDED_PROPOSAL"  # The literal gate accepts it.
    result, blocked = apply_attribution_safety(audit, source_map)
    field = cf_field(result)
    assert list(blocked) == ["p1.s0"]
    assert field["validation"] == "REJECTED" and field["value"] is None
    assert "proposal_grounding" not in field and "proposed_value" not in field
    assert field["coverage"]["state"] == "REJECTED_PROPOSAL"
    assert field["attribution_safety"]["foreign_attestations"][0]["chunk_id"] == -1
    assert result["attribution_safety"] == {"revision": 1, "blocked": ["p1.s0"]}
    others = [f for f in result["fields"] if f["candidate_id"] != "p1.s0"]
    assert others == [f for f in audit["fields"] if f["candidate_id"] != "p1.s0"]


@pytest.mark.parametrize("owner", [f"di {OPERATOR}", "della società"])
def test_own_or_generic_organization_is_unchanged(owner):
    audit, source_map = audited(owner)
    assert cf_field(audit)["validation"] == "GROUNDED_PROPOSAL"
    assert apply_attribution_safety(audit, source_map) == (audit, {})


def test_general_kb_is_never_an_attestation():
    audit, source_map = audited("di Beta Costruzioni S.r.l.")
    source_map[-1] = {**source_map[-1], "category": "general"}
    assert apply_attribution_safety(audit, source_map)[1] == {}


def test_owner_grammar_is_closed_and_name_based():
    operators = {"alfaingegneriasrl"}
    assert own_organization("Alfa Ingegneria S.r.l.", operators)
    assert own_organization("Società Alfa Ingegneria S.r.l.", operators)
    assert own_organization("società nello scenario dal 2 ottobre 2026", operators)
    assert not own_organization("società Beta S.r.l.", operators)
    assert not own_organization("società mandante", operators)
    assert not own_organization("Alfa Ingegneria S.r.l.", set())
    # Elided forms the literal gate does not accept yet: checked on the attestation itself.
    sources = {-1: src(-1, named("dell'impresa ausiliaria")),
               -2: src(-2, named("dell'operatore economico"))}
    found = foreign_attestations(sources, "Mario Bianchi", "direttore tecnico", operators)
    assert [f["organization"] for f in found] == ["impresa ausiliaria"]


def test_bold_signatory_role_reuse_checks_its_organization():
    """V13 reuses the signatory's documented office from the bold named form."""
    for owner, expected in [("della società dal 2 ottobre 2026", set()),
                            ("di Beta Costruzioni S.r.l.", {"p9.s0"})]:
        line = f"- **Mario Bianchi**: **amministratore unico {owner}**. Nomina documentata."
        role = {"candidate_id": "p9.s0", "validation": "REVIEW_REQUIRED",
                "proposal_grounding": "GROUNDED_PROPOSAL", "proposed_value": "amministratore unico",
                "value": None, "validation_errors": [],
                "source_reuse": {"property": "role", "subject": "Mario Bianchi",
                                 "fact": {"subject": "Mario Bianchi", "entity": "person",
                                          "property": "role", "value": "amministratore unico"}}}
        operator = {"candidate_id": "p8.s0", "validation": "GROUNDED_PROPOSAL", "fact_id": "co",
                    "value": OPERATOR}
        audit = {"fields": [operator, role], "facts": {"co": {"subject": OPERATOR,
                                                              "entity": "company"}}}
        blocked = apply_attribution_safety(audit, {-1: src(-1, line)})[1]
        assert set(blocked) == expected


def test_only_new_sessions_run_the_check():
    layout = form()
    audit, source_map = audited("di Beta Costruzioni S.r.l.")
    catalog = adapter.initial_state(layout)
    assert catalog["attribution_safety_revision"] == 1
    sources = list(source_map.values())
    data = {**catalog, "audit": audit}
    new = reuse_selected_person(data, layout, "alpha", sources)
    assert cf_field(new)["validation"] == "REJECTED"
    assert new["counts"].get("REJECTED", 0) >= 1
    previous = {k: v for k, v in data.items() if k != "attribution_safety_revision"}
    old = reuse_selected_person(previous, layout, "alpha", sources)
    assert cf_field(old)["validation"] == "GROUNDED_PROPOSAL"
    assert "attribution_safety" not in old


@pytest.mark.anyio
@pytest.mark.parametrize("revision", ["NEW_SESSION", "OLD_SESSION"])
async def test_real_session_never_exports_a_foreign_organization_role(
    api, provider, monkeypatch, revision,
):
    stream = BytesIO()
    doc = Document()
    for text in (HEADER, "C.F. ____________", "Denominazione ____________"):
        doc.add_paragraph(text)
    doc.save(stream)
    uploaded = await api.post("/api/projects/alpha/forms", files={
        "file": ("modulo.docx", stream.getvalue())})
    await source(api, f"{named('di Beta Costruzioni S.r.l.')}\n{COMPANY_LINE}", name="persone.md")

    async def model(client, settings, body, timeout, **kwargs):
        inp = json.loads(body["messages"][1]["content"])
        chunk = next(s["chunk_id"] for s in inp["sources"] if "Beta" in s["content"])
        return json.dumps(proposal(named("di Beta Costruzioni S.r.l."), chunk, chunk)), \
            "simulated", {}

    monkeypatch.setattr(adapter, "request_model_content", model)
    session = (await api.post(BASE, json={"form_id": uploaded.json()["id"], "engine": "global",
                                         "start_in_chat": True})).json()
    if revision == "OLD_SESSION":
        from app.db import connection

        with connection() as db:
            db.execute("BEGIN IMMEDIATE")
            row, state = adapter._row_state(db, "alpha", session["id"], session["version"])
            state["global_compilation"].pop("attribution_safety_revision")
            session = sessions.save_state(db, row, state, "test_previous_revision")
    session = await resolve(api, session)
    ids = dialogue.export_selection(session)[0]
    assert "p2.s0" in ids  # The operator itself is unaffected.
    if revision == "OLD_SESSION":
        assert "p1.s0" in ids  # Previous sessions keep their behaviour.
        return
    assert "p1.s0" not in ids
    bypass = await api.post(f"{BASE}/{session['id']}/global-draft", json={
        "version": session["version"], "candidate_ids": ["p1.s0"], "allow_unresolved": True,
        "confirmed_condition_ids": [], "user_quote": "Genera la bozza"})
    assert bypass.status_code == 422
