"""Ordinary label/value tables, local conditions and unchanged factual gates."""

from io import BytesIO

import pytest
from docx import Document
from test_compilation_sessions import api as api
from test_compilation_sessions import create, resolve, simulate
from test_compilation_sessions import source as upload_source

from app import compilation_session_resolution as resolution
from app.compilation_semantics import (
    FormReview,
    SourceReview,
    review_meanings,
    review_sources,
    validate_persisted_semantics,
)
from app.compilation_session_models import CandidateMatches, CandidateMeanings
from app.compilation_sessions import candidate_snapshots
from app.docx_templates import inspect_docx

LABELS = (
    "Impresa candidata", "Forma giuridica", "Indirizzo sede legale",
    "Sede operativa (se diversa dalla sede legale)", "Codice fiscale azienda",
    "Partita IVA azienda", "Telefono aziendale",
)
CONDITION = "se diversa dalla sede legale"


@pytest.fixture
def anyio_backend():
    return "asyncio"


def table_document(heading="Richiesta di ammissione"):
    document = Document()
    document.add_paragraph(heading)
    table = document.add_table(rows=len(LABELS), cols=2)
    for row, label in zip(table.rows, LABELS, strict=True):
        row.cells[0].text = label
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def fields_for_table(heading="Richiesta di ammissione"):
    layout = inspect_docx(table_document(heading))
    return layout, candidate_snapshots(layout, "alpha", 1, "anagrafica.docx")


def table_meanings(fields, *, explicit_condition=True):
    return CandidateMeanings(fields=[{
        "candidate_id": f["id"], "classification": "data", "entity": "company",
        "requirement": {"name": f["label_hint"].split(" (")[0],
                        "form_quote": f["label_hint"], "form_citation_id": 1},
        "form_quote": f["label_hint"], "reason": "Proprietà dell'impresa candidata",
        "condition": CONDITION if explicit_condition and CONDITION in f["label_hint"] else "",
        "semantic": {"subject_anchor": LABELS[0], "subject_relation": "represented_organization",
                     "context_role": "ORGANIZATION_PROFILE", "context_quote": f["label_hint"],
                     "section_id": "paragraph:0"},
    } for f in fields])


async def approve_form(task, data, schema):
    assert schema is FormReview
    return FormReview(fields=[{
        "candidate_id": p["candidate"]["id"], "accepted": True,
        "condition_complete": True, "reason": "Label, slot, soggetto e condizione verificati",
    } for p in data["proposals"]])


async def approve_source(task, data, schema):
    assert schema is SourceReview
    return SourceReview(supports=[{
        "candidate_id": p["candidate_id"], "kind": p["kind"], "index": p["index"],
        "accepted": True, "reason": "Relazione verificata nella SOURCE",
    } for p in data["proposals"]])


def evidence(content, *, role="source"):
    return {"chunk_id": -1, "file_id": -1, "source_name": "profilo.txt", "chunk_index": 0,
            "content": content, "role": role, "scope": "global", "category": "company",
            "project_id": None, "document_metadata": "test"}


@pytest.mark.anyio
async def test_ordinary_table_labels_do_not_become_company_type_conditions():
    _, fields = fields_for_table()
    ordinary = [f for f in fields if CONDITION not in f["label_hint"]]
    meanings = table_meanings(ordinary)

    async def assess(task, data, schema):
        assert len(data["proposals"]) == len(ordinary)
        assert all(not p["interpretation"]["condition"] for p in data["proposals"])
        return await approve_form(task, data, schema)

    await review_meanings(ordinary, meanings, assess)
    resolution.apply_meanings(ordinary, meanings)
    assert all(f["requirement"] and not f["condition"] for f in ordinary)
    assert all(f["entity"] == "company" and f["provenance"] is None for f in ordinary)
    assert all(f["search"]["status"] == "not_searched" for f in ordinary)
    for field in ordinary:
        validate_persisted_semantics(field)


@pytest.mark.anyio
async def test_local_row_condition_is_anchored_to_its_own_form_row():
    _, fields = fields_for_table()
    conditional = next(f for f in fields if CONDITION in f["label_hint"])
    meanings = table_meanings([conditional])
    await review_meanings([conditional], meanings, approve_form)
    resolution.apply_meanings([conditional], meanings)
    assert conditional["requirement"] and conditional["condition"] == CONDITION
    section_id = conditional["semantic"]["section_id"]
    assert section_id == conditional["structural"]["row_section_id"]
    sections = {s["id"]: s["text"] for s in conditional["structural"]["form_sections"]}
    assert CONDITION in sections[section_id]
    assert not conditional.get("applicability") and conditional["provenance"] is None
    validate_persisted_semantics(conditional)


@pytest.mark.anyio
async def test_review_can_restore_explicit_row_condition_omitted_by_classifier():
    _, fields = fields_for_table()
    field = next(f for f in fields if CONDITION in f["label_hint"])
    meanings = table_meanings([field], explicit_condition=False)

    async def assess(task, data, schema):
        return FormReview(fields=[{
            "candidate_id": field["id"], "accepted": True, "condition_complete": True,
            "reason": "La parentesi condiziona solo questa riga",
            "section_condition": {"condition": CONDITION,
                "section_id": field["structural"]["row_section_id"], "condition_kind": "other"},
        }])

    await review_meanings([field], meanings, assess)
    resolution.apply_meanings([field], meanings)
    assert field["condition"] == CONDITION and field["requirement"]


@pytest.mark.anyio
async def test_incomplete_condition_review_still_rejects_table_field():
    _, fields = fields_for_table()
    field = next(f for f in fields if CONDITION in f["label_hint"])
    meanings = table_meanings([field], explicit_condition=False)

    async def assess(task, data, schema):
        return FormReview(fields=[{
            "candidate_id": field["id"], "accepted": True, "condition_complete": False,
            "reason": "Omette una condizione esplicita nella label",
        }])

    await review_meanings([field], meanings, assess)
    resolution.apply_meanings([field], meanings)
    assert field["status"] == "PENDING" and field["requirement"] is None
    assert field["search"]["status"] == "not_searched"


@pytest.mark.anyio
async def test_table_section_type_remains_a_reviewed_condition():
    _, fields = fields_for_table("Se l'impresa candidata è una cooperativa")
    field = fields[0]
    meanings = table_meanings([field])
    meanings.fields[0].semantic.subject_anchor = "cooperativa"

    async def assess(task, data, schema):
        return FormReview(fields=[{
            "candidate_id": field["id"], "accepted": True, "condition_complete": True,
            "reason": "La sezione è esplicitamente condizionata",
            "section_condition": {"condition": "cooperativa", "section_id": "paragraph:0",
                                  "condition_kind": "subject_type"},
        }])

    await review_meanings([field], meanings, assess)
    resolution.apply_meanings([field], meanings)
    assert field["condition"] == "cooperativa" and field["requirement"]
    assert not field.get("applicability")


@pytest.mark.anyio
@pytest.mark.parametrize("invalid", ["other_slot", "person", "invented_quote", "invented_property"])
async def test_table_interpretation_still_needs_its_own_anchor_subject_and_review(invalid):
    _, fields = fields_for_table()
    field = fields[0]
    meanings = table_meanings([field])
    item = meanings.fields[0]
    if invalid == "other_slot":
        item.semantic.form_anchor = fields[1]["structural"]["slot_anchor"]
    elif invalid == "person":
        item.entity = "person"
        item.requirement.person_role = "firmatario"
    elif invalid == "invented_quote":
        item.semantic.context_quote = "Impresa candidata | Partita IVA azienda"
    else:
        item.requirement.name = "Fatturato annuale"

    async def assess(task, data, schema):
        assert invalid == "invented_property"
        return FormReview(fields=[{
            "candidate_id": field["id"], "accepted": False, "condition_complete": True,
            "reason": "Proprietà non richiesta da questa riga",
        }])

    await review_meanings([field], meanings, assess)
    resolution.apply_meanings([field], meanings)
    assert field["status"] == "PENDING" and field["requirement"] is None
    assert field["search"]["status"] == "not_searched"


@pytest.mark.anyio
@pytest.mark.parametrize("value,role,resolved", [
    ("Aurora Progetti S.r.l.", "source", True),
    ("Aurora Progetti S.r.l.", "form", False),
    ("Impresa candidata: ________", "source", False),
    ("Impresa inventata", "source", False),
])
async def test_accepted_table_label_does_not_relax_source_gates(value, role, resolved):
    layout, fields = fields_for_table()
    field = fields[0]
    meanings = table_meanings([field])
    await review_meanings([field], meanings, approve_form)
    resolution.apply_meanings([field], meanings)
    proof = evidence(
        "Impresa candidata: Aurora Progetti S.r.l.\nImpresa candidata: ________", role=role,
    )
    matches = CandidateMatches(fields=[{
        "candidate_id": field["id"], "reason": "Proposta da verificare", "supports": [
            {"source_id": 1, "quote": proof["content"], "value": value},
        ],
    }])
    coverage = {field["id"]: [1]}
    await review_sources([field], matches, [proof], coverage, approve_source)
    resolution.apply_matches(layout, [field], matches, [proof], coverage)
    assert field["status"] == ("RESOLVED" if resolved else "MISSING")
    assert field["provenance"] == ("SOURCE" if resolved else None)


@pytest.mark.anyio
async def test_conditional_table_value_without_applicability_is_not_resolved():
    layout, fields = fields_for_table()
    field = next(f for f in fields if CONDITION in f["label_hint"])
    meanings = table_meanings([field])
    await review_meanings([field], meanings, approve_form)
    resolution.apply_meanings([field], meanings)
    proof = evidence("Sede operativa: Via del Molo 8, Genova.")
    matches = CandidateMatches(fields=[{
        "candidate_id": field["id"], "reason": "Solo il valore, non la condizione", "supports": [
            {"source_id": 1, "quote": proof["content"], "value": "Via del Molo 8, Genova"},
        ],
    }])
    coverage = {field["id"]: [1]}
    await review_sources([field], matches, [proof], coverage, approve_source)
    resolution.apply_matches(layout, [field], matches, [proof], coverage)
    assert field["status"] == "AMBIGUOUS" and field["value"] is None
    assert field["provenance"] is None and field["condition"] == CONDITION


@pytest.mark.anyio
async def test_table_form_acceptance_reaches_real_retrieval_and_persistence(api, monkeypatch):
    state = await create(api, table_document())
    values = ["Aurora Progetti S.r.l.", "Società a responsabilità limitata",
              "Via del Porto 19, Genova", "Via del Molo 8, Genova", "01234567890",
              "IT01234567890", "+39 010 123456"]
    await upload_source(api, "\n".join(label + ': ' + value for label, value
                                      in zip(LABELS, values, strict=True)))
    mapping = {f["id"]: [value] for f, value in zip(state["fields"], values, strict=True)}
    simulate(monkeypatch, mapping)

    async def classify(fields):
        return table_meanings(fields)

    async def request(task, data, schema):
        if schema is FormReview:
            return await approve_form(task, data, schema)
        return await approve_source(task, data, schema)

    monkeypatch.setattr(resolution, "classify_candidates", classify)
    monkeypatch.setattr(resolution, "request_structured", request)
    state = await resolve(api, state)
    for field, value in zip(state["fields"], values, strict=True):
        assert field["requirement"] and field["search"]["status"] == "searched"
        if CONDITION in field["label_hint"]:
            assert field["condition"] == CONDITION and field["status"] == "AMBIGUOUS"
        else:
            assert field["status"] == "RESOLVED" and field["value"] == value
            assert field["provenance"] == "SOURCE"
        assert field["form_evidence"]["role"] == "form"
    again = await api.get(f"/api/projects/alpha/compilation-sessions/{state['id']}")
    assert again.json() == state
