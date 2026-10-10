"""Compact coverage and conservative guards; not model-quality assertions."""

import copy
import json
from io import BytesIO

import pytest
from docx import Document

from app.compilation_global import build_global_input, validate_global_proposal
from app.docx_templates import inspect_docx


@pytest.fixture
def case():
    doc = Document()
    table = doc.add_table(rows=3, cols=2)
    table.cell(0, 0).text = "Recapito principale"
    table.cell(1, 0).text = "Recapito secondario (solo se diverso dal principale)"
    table.cell(2, 0).text = "Responsabile - Nome"
    stream = BytesIO()
    doc.save(stream)
    layout = inspect_docx(stream.getvalue())
    source = {"chunk_id": -7, "file_id": -3, "chunk_index": 0,
              "source_name": "profilo.txt", "role": "source", "scope": "global",
              "category": "company", "project_id": None,
              "content": "Principale: Via Luna 4; secondario: Via Luna 4; "
                         "Responsabile: Dott. Ada Neri, Registro n. 7526"}
    fact = {"subject": "Organizzazione", "entity": "company", "aspect": "location",
            "property": "recapito principale", "value": "Via Luna 4",
            "evidence": [{"chunk_id": -7, "quote": source["content"]}]}
    payload = {"version": 4, "facts": {"main": fact, "secondary": {
        **copy.deepcopy(fact), "property": "recapito secondario"}},
        "reviewed_paragraph_ids": ["paragraph:0", "paragraph:2", "paragraph:4"],
        "uncertain_paragraph_ids": [],
        "conditions": {"different": {
            "form": {"paragraph_id": "paragraph:2", "quote": "solo se diverso dal principale"},
            "aspect": "location", "operator": "different", "fact_ids": ["main", "secondary"],
            "reason": "Confronto recapiti", "applies_to": ["t0.r1.c1"]}},
        "fields": {"t0.r0.c1": {
            "label": "Recapito principale", "kind": "data", "disposition": "value",
            "property": "recapito principale", "fact_id": "main", "conditions": [],
            "form": {"paragraph_id": "paragraph:0", "quote": "Recapito principale"}}},
        "deferred": ["t0.r1.c1", "t0.r2.c1"]}
    return layout, [source], payload


def audit(case):
    return validate_global_proposal(json.dumps(case[2]), case[0], "alpha", case[1])


def test_compact_input_requests_exact_scan_without_interpreting_docx(case):
    data = build_global_input(case[0], "alpha", case[1])
    assert data["form_review_paragraph_ids"] == ["paragraph:0", "paragraph:2", "paragraph:4"]
    assert set(data["schema_output"]["properties"]) == {
        "version", "facts", "conditions", "reviewed_paragraph_ids",
        "uncertain_paragraph_ids", "fields", "deferred"}


def test_deferred_positions_are_covered_but_never_invented_missing_fields(case):
    result = audit(case)
    assert result["candidate_count"] == 3
    assert result["deferred_count"] == 2
    assert result["counts"] == {"GROUNDED_PROPOSAL": 1, "REVIEW_REQUIRED": 2}
    assert result["resolved"] == result["writes"] == 0
    for field in result["fields"][1:]:
        assert field["validation"] == "REVIEW_REQUIRED"
        assert "value" not in field and "label" not in field
        assert field["semantic_review_required"] and field["written_value"] is None


@pytest.mark.parametrize("mutation", ["omitted", "invented", "duplicate", "overlap"])
def test_compact_candidate_partition_must_be_exact(case, mutation):
    payload = case[2]
    if mutation == "omitted":
        payload["deferred"].pop()
    elif mutation == "invented":
        payload["deferred"][0] = "invented"
    elif mutation == "duplicate":
        payload["deferred"].append(payload["deferred"][0])
    else:
        payload["deferred"].append("t0.r0.c1")
    with pytest.raises(ValueError, match="Coverage"):
        audit(case)


@pytest.mark.parametrize("mutation", ["omitted", "invented", "duplicate",
                                          "uncertain_unknown", "uncertain_duplicate"])
def test_inventory_receipt_is_exact_and_uncertainty_cannot_name_unknown_paragraphs(case, mutation):
    payload = case[2]
    if mutation == "omitted":
        payload["reviewed_paragraph_ids"].remove("paragraph:2")
    elif mutation == "invented":
        payload["reviewed_paragraph_ids"].append("paragraph:99")
    elif mutation == "duplicate":
        payload["reviewed_paragraph_ids"].append("paragraph:2")
    elif mutation == "uncertain_unknown":
        payload["uncertain_paragraph_ids"] = ["paragraph:99"]
    else:
        payload["uncertain_paragraph_ids"] = ["paragraph:2", "paragraph:2"]
    with pytest.raises(ValueError, match="Inventario"):
        audit(case)


def test_condition_inverse_inventory_is_backend_derived_from_canonical_origin(case):
    result = audit(case)
    assert result["condition_scan"] == {
        "paragraph:0": [], "paragraph:2": ["different"], "paragraph:4": []}


def test_model_cannot_return_a_second_conflicting_condition_map(case):
    case[2]["condition_scan"] = {"paragraph:0": ["different"]}
    with pytest.raises(ValueError, match="Extra inputs"):
        audit(case)


def test_condition_origin_requires_an_actual_form_paragraph(case):
    case[2]["conditions"]["different"]["form"]["paragraph_id"] = "paragraph:99"
    with pytest.raises(ValueError, match="Inventario"):
        audit(case)


def test_clause_in_a_wrong_real_paragraph_is_still_rejected_by_literal_gate(case):
    case[2]["conditions"]["different"]["form"]["paragraph_id"] = "paragraph:0"
    result = audit(case)
    assert result["conditions"]["different"]["validation"] == "REJECTED"


def test_local_condition_is_attached_even_if_assignment_omits_it(case):
    field = copy.deepcopy(case[2]["fields"]["t0.r0.c1"])
    field.update(label="Recapito secondario", property="recapito secondario", fact_id="secondary",
                 form={"paragraph_id": "paragraph:2", "quote": "Recapito secondario"})
    case[2]["fields"]["t0.r1.c1"] = field
    case[2]["deferred"].remove("t0.r1.c1")
    # Even omitting the target in the semantic scope cannot hide a LOCAL clause.
    case[2]["conditions"]["different"]["applies_to"] = []
    result = audit(case)
    assert result["fields"][1]["validation"] == "REJECTED"
    assert "falsa" in result["fields"][1]["validation_errors"][0]


def test_uncertain_scan_blocks_value_without_inventing_conditions(case):
    case[2]["uncertain_paragraph_ids"] = ["paragraph:0"]
    result = audit(case)
    assert result["fields"][0]["validation"] == "REJECTED"
    assert "incerto" in result["fields"][0]["validation_errors"][0]


def test_unconditional_scan_is_not_independent_semantic_acceptance(case):
    result = audit(case)
    assert result["condition_inventory"]["semantic_completeness_verified"] is False
    assert result["fields"][0]["semantic_review_required"]
    assert result["fields"][0]["written_value"] is None


def test_fact_cannot_be_reused_as_a_different_property(case):
    case[2]["fields"]["t0.r0.c1"]["property"] = "numero registro"
    field = audit(case)["fields"][0]
    assert field["validation"] == "REJECTED"
    assert "proprietà" in field["validation_errors"][0]


def test_composite_profile_is_not_a_fact_for_each_atomic_property(case):
    payload = case[2]
    payload["facts"]["profile"] = {
        "subject": "Ada Neri", "entity": "person", "aspect": "professional",
        "property": "profilo responsabile", "value": "Dott. Ada Neri, Registro n. 7526",
        "evidence": [{"chunk_id": -7, "quote": case[1][0]["content"]}]}
    payload["fields"]["t0.r2.c1"] = {
        "label": "Nome responsabile", "kind": "data", "disposition": "value",
        "property": "nome responsabile", "fact_id": "profile", "conditions": [],
        "form": {"paragraph_id": "paragraph:4", "quote": "Nome"}}
    payload["deferred"].remove("t0.r2.c1")
    assert audit(case)["fields"][2]["validation"] == "REJECTED"
    # A distinct atom from the same quote is allowed as a PROPOSAL.
    payload["facts"]["name"] = {**payload["facts"]["profile"],
                                "property": "nome responsabile", "value": "Ada Neri"}
    payload["fields"]["t0.r2.c1"]["fact_id"] = "name"
    del payload["facts"]["profile"]
    field = audit(case)["fields"][2]
    assert field["validation"] == "GROUNDED_PROPOSAL"
    assert field["value"] == "Ada Neri"
    assert field["source_evidence"][0]["chunk_id"] == -7


@pytest.mark.parametrize("value", ["non disponibile nel dataset", "non noto",
                                       "Not provided in the sources", "____", "N/A"])
def test_unavailability_is_not_a_usable_source_value(case, value):
    case[1][0]["content"] += "; Informazione: " + value
    case[2]["facts"]["main"].update(value=value, evidence=[{
        "chunk_id": -7, "quote": "Informazione: " + value}])
    result = audit(case)
    assert result["facts"]["main"]["validation"] == "REJECTED"
    assert result["fields"][0]["validation"] == "REJECTED"


def test_legitimate_negative_fact_is_not_an_availability_placeholder(case):
    text = "Non possiede la qualifica aggiuntiva"
    case[1][0]["content"] += "; " + text
    case[2]["facts"]["negative"] = {**copy.deepcopy(case[2]["facts"]["main"]),
        "property": "qualifica aggiuntiva", "value": text,
        "evidence": [{"chunk_id": -7, "quote": text}]}
    case[2]["conditions"]["qualification"] = {
        "form": {"paragraph_id": "paragraph:4", "quote": "Responsabile"},
        "aspect": "location", "operator": "semantic", "fact_ids": ["negative"],
        "reason": "Proposta da valutare", "applies_to": ["t0.r2.c1"]}
    assert audit(case)["facts"]["negative"]["validation"] == "GROUNDED_PROPOSAL"


def test_scope_cannot_name_an_invented_candidate(case):
    case[2]["conditions"]["different"]["applies_to"] = ["invented"]
    with pytest.raises(ValueError, match="Inventario"):
        audit(case)


def test_semantic_section_scope_does_not_authorize_nonlocal_exclusion(case):
    payload = case[2]
    payload["conditions"]["different"]["applies_to"] = ["t0.r2.c1"]
    payload["fields"]["t0.r2.c1"] = {
        "label": "Nome responsabile", "kind": "data", "disposition": "not_applicable",
        "conditions": [], "form": {"paragraph_id": "paragraph:4", "quote": "Nome"}}
    payload["deferred"].remove("t0.r2.c1")
    field = audit(case)["fields"][2]
    assert field["validation"] == "REJECTED"
    assert "Non applicabilità" in field["validation_errors"][0]


def test_invalid_assignment_remains_local_without_weakening_partition(case):
    case[2]["fields"]["t0.r2.c1"] = {"made_up": "unusable"}
    case[2]["deferred"].remove("t0.r2.c1")
    result = audit(case)
    assert result["fields"][0]["validation"] == "GROUNDED_PROPOSAL"
    assert result["fields"][2]["validation"] == "REJECTED"


def test_unreferenced_source_facts_are_rejected_instead_of_cataloguing_kb(case):
    case[2]["facts"]["unrelated"] = copy.deepcopy(case[2]["facts"]["main"])
    with pytest.raises(ValueError, match="non referenziati"):
        audit(case)


@pytest.mark.parametrize("disposition", ["review", "missing", "ignore"])
def test_unresolved_or_decorative_details_must_use_compact_deferred_group(case, disposition):
    case[2]["fields"]["t0.r0.c1"].update(disposition=disposition, fact_id=None)
    # The main fact remains a dependency of the comparison, not an orphan.
    assert audit(case)["fields"][0]["validation"] == "REJECTED"


@pytest.mark.parametrize("origin", ["form", "user", "other_project"])
def test_compact_provenance_does_not_expand_allowed_corpus(case, origin):
    source = case[1][0]
    if origin == "other_project":
        source.update(chunk_id=7, file_id=3, project_id="beta", scope="project:beta", category=None)
    else:
        source["role"] = origin
    with pytest.raises(ValueError):
        audit(case)
