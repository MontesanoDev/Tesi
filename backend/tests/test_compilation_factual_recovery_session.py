"""V13 persistence and ordinary export; provider transport alone is simulated."""

import json
from copy import deepcopy
from io import BytesIO

import pytest
from docx import Document
from test_compilation_factual_recovery import inline_case, role_case
from test_compilation_global_sessions import BASE, api, provider, resolve
from test_compilation_selected_person import NAME, PROFILE
from test_compilation_sessions import source

from app import compilation_global_sessions as adapter
from app.compilation_global_dialogue import export_selection
from app.db import connection
from app.repository import save_conversation_turn

__all__ = ["api", "provider"]


@pytest.fixture
def anyio_backend():
    return "asyncio"


async def prepare(api, monkeypatch, layout, sources, payload, *, revision=13, choice=None):
    upload = (await api.post("/api/projects/alpha/forms", files={
        "file": ("modulo.docx", layout.original)})).json()
    for index, item in enumerate(sources):
        await source(api, item["content"], name=f"profilo-{index}.md")
    calls = []

    async def propose(client, settings, body, timeout, **kwargs):
        calls.append(body)
        actual = json.loads(body["messages"][1]["content"])
        result = deepcopy(payload)
        for part in [*result["facts"].values(), *result["bindings"].values()]:
            for proof in part["evidence"]:
                original = next(s for s in sources if s["chunk_id"] == proof["chunk_id"])
                match = next(s for s in actual["sources"]
                             if original["content"] in s["content"])
                proof["chunk_id"] = match["chunk_id"]
        return json.dumps(result), "simulated", {}

    monkeypatch.setattr(adapter, "request_model_content", propose)
    session = (await api.post(BASE, json={"form_id": upload["id"], "engine": "global",
                                        "start_in_chat": True})).json()
    with connection() as db:
        row, state = adapter._row_state(db, "alpha", session["id"], session["version"])
        assert state["global_compilation"]["backend_revision"] == 13
        if revision == 12:
            state["global_compilation"].pop("backend_revision")
            db.execute("UPDATE compilation_sessions SET state_json=? WHERE id=?",
                       (json.dumps(state), row["id"]))
    if choice:
        save_conversation_turn(session["conversation_id"], {"question": choice,
            "answer": "Avvio analisi", "generation_status": "direct",
            "compilation": {"session_id": session["id"], "action": "start"}})
    return await resolve(api, session), calls


@pytest.mark.anyio
@pytest.mark.parametrize("table", [False, True])
@pytest.mark.parametrize("selection", ["scoped", "role_only", "generic", "two_names"])
async def test_missing_model_name_only_closed_user_choice_reaches_export_and_reopen(
    api, provider, monkeypatch, table, selection,
):
    _, layout, sources = inline_case(table=table)
    sources[0]["content"] = PROFILE
    texts = {"scoped": f"Il firmatario scelto per questa domanda è {NAME}.",
             "role_only": f"{NAME} è rappresentante legale.",
             "generic": "Confermo tutto, genera la bozza.",
             "two_names": f"Il firmatario scelto per questa domanda è {NAME} oppure Bea Mori."}
    payload = {"version": 12, "facts": {}, "fields": {}, "bindings": {}, "conditions": {},
               "clarifications": {}, "choice_groups": {}, "coverage": []}
    session, calls = await prepare(api, monkeypatch, layout, sources, payload,
                                   choice=texts[selection])
    reopened = (await api.get(f"{BASE}/{session['id']}")).json()
    assert reopened == session
    ids, conditions, users = export_selection(reopened)
    assert len(ids) == (1 if selection == "scoped" else 0)
    assert not conditions and not users and len(calls) == 1
    if ids:
        response = await api.post(f"{BASE}/{session['id']}/global-draft", json={
            "version": session["version"], "candidate_ids": ids, "allow_unresolved": True,
            "user_quote": "Genera la bozza parziale"})
        assert response.status_code == 200, response.text
        downloads = response.json()["last_generation"]["downloads"]
        report = (await api.get(downloads["report"])).json()
        assert report["written_field_count"] == 1
        assert report["fields"][0]["provenance"] == "SOURCE"
        assert report["fields"][0]["written_value"] == NAME
        output = Document(BytesIO((await api.get(downloads["docx"])).content))
        assert NAME in output.element.xml
        if table:
            assert output.tables[0].cell(0, 1).text == ""
    assert len(calls) == 1


@pytest.mark.anyio
@pytest.mark.parametrize("revision", [12, 13])
async def test_old_persisted_session_does_not_acquire_v13_on_analysis_or_read(
    api, provider, monkeypatch, revision,
):
    layout, sources, payload = role_case()
    session, calls = await prepare(api, monkeypatch, layout, sources, payload, revision=revision)
    assert bool(export_selection(session)[0]) is (revision == 13)
    for _ in range(2):
        assert (await api.get(f"{BASE}/{session['id']}")).json() == session
    with connection() as db:
        _, state = adapter._row_state(db, "alpha", session["id"], session["version"])
        assert state["global_compilation"].get("backend_revision", 12) == revision
    assert len(calls) == 1


@pytest.mark.anyio
@pytest.mark.parametrize("mutation", ["none", "value", "role", "registration", "removed"])
async def test_cross_record_review_and_saved_selection_revalidate_current_sources_before_export(
    api, provider, monkeypatch, mutation,
):
    layout, sources, payload = role_case()
    session, calls = await prepare(api, monkeypatch, layout, sources, payload)
    ids, conditions, users = export_selection(session)
    assert ids == ["p1.s0"] and not conditions and not users
    request = {"version": session["version"], "candidate_ids": ids, "allow_unresolved": True,
               "user_quote": "Esporta la bozza parziale"}
    with connection() as db:
        if mutation == "value":
            db.execute("UPDATE global_document_chunks SET content=replace(content,"
                       " 'Ingegnere', 'Architetto')")
        elif mutation == "role":
            db.execute("UPDATE global_document_chunks SET content=replace(content,"
                       " 'Direttore tecnico', 'Socio')")
        elif mutation == "registration":
            db.execute("UPDATE global_document_chunks SET content=replace(content, '6234', '9123')"
                       " WHERE content LIKE '%Direttore tecnico%'")
        elif mutation == "removed":
            db.execute("DELETE FROM global_document_chunks "
                       "WHERE content LIKE '%Direttore tecnico%'")
    response = await api.post(f"{BASE}/{session['id']}/global-draft", json=request)
    if mutation == "none":
        assert response.status_code == 200, response.text
        downloads = response.json()["last_generation"]["downloads"]
        report = (await api.get(downloads["report"])).json()
        assert report["written_field_count"] == 1
        assert report["fields"][0]["provenance"] == "SOURCE"
        output = Document(BytesIO((await api.get(downloads["docx"])).content))
        assert "Ingegnere" in "\n".join(p.text for p in output.paragraphs)
    else:
        assert response.status_code == 409, response.text
        assert (await api.get(f"{BASE}/{session['id']}")).json()["last_generation"] is None
    assert len(calls) == 1
