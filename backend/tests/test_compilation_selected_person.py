"""Registered signatory choices reuse literal SOURCE, never USER biography."""

import copy
import json
from io import BytesIO

import pytest
from docx import Document

from app.compilation_global import TASK, build_global_input, validate_global_proposal
from app.docx_templates import inspect_docx

NAME = "Ada Neri"
PROFILE = """### Ada Neri — anagrafica
- Nome: Ada.
- Cognome: Neri.
- Luogo di nascita: Torino (TO).
- Data di nascita: 12 gennaio 1985.
- Comune di residenza: Milano (MI).
- Indirizzo di residenza: Via Roma 24.
- Codice fiscale personale: NREADA85A52L219R.
"""


def case():
    doc = Document()
    doc.add_paragraph("Il sottoscritto ________, nato a ________ (________) il ________, residente")
    doc.add_paragraph("in ________ (________), via ________, C.F. ________, in qualità di")
    doc.add_paragraph("________ dell’operatore economico ________ con sede in")
    stream = BytesIO()
    doc.save(stream)
    layout = inspect_docx(stream.getvalue(), require_candidates=False)
    src = {"chunk_id": -7, "file_id": -3, "chunk_index": 0, "source_name": "profili.md",
           "role": "source", "scope": "global", "category": "company", "project_id": None,
           "content": '- **Società:** Aurora Servizi S.r.l.\n' + PROFILE
           + '\n- **Ada Neri**: **amministratrice unica della società**.'}
    payload = {"version": 12, "fields": {}, "facts": {}, "bindings": {}, "conditions": {},
               "choice_groups": {}, "clarifications": {}, "coverage": []}
    payload["facts"] = {"company": {"subject": "Aurora Servizi S.r.l.", "entity": "company",
        "aspect": "identity", "property": "denominazione", "value": "Aurora Servizi S.r.l.",
        "evidence": [{"chunk_id": -7, "quote": "- **Società:** Aurora Servizi S.r.l."}]}}
    payload["fields"] = {"p2.s1": {"label": "Operatore economico", "kind": "data",
        "disposition": "value", "conditions": [], "property": "denominazione", "fact_id": "company",
        "form": {"paragraph_id": "paragraph:2", "quote": "operatore economico [[p2.s1]]"}}}
    audit = validate_global_proposal(json.dumps(payload), layout, "alpha", [src])
    catalog = build_global_input(layout, "alpha", [src])["form"]
    receipt = {"action": "VALUE", "value": NAME, "user_quote": NAME, "provenance": "USER",
               "label": "Il sottoscritto"}
    data = {"audit": audit, "catalog": catalog, "writable_ids": list(catalog["candidates"]),
            "dialogue": {"answers": {"field:p0.s0": receipt}, "user_values": {
                "p0.s0": {"value": NAME, "user_quote": NAME, "provenance": "USER"}}}}
    return data, layout, [src]


def reused(data, layout, sources):
    from app.compilation_source_reuse import reuse_selected_person

    return reuse_selected_person(data, layout, "alpha", sources)


def test_binding_prompt_describes_supported_composed_proofs_without_same_chunk_ban():
    assert "Ruolo e valore devono essere nello stesso chunk SOURCE." not in TASK
    assert "titolo" in TASK and "intestazioni" in TASK and "omonim" in TASK


def test_selected_person_yields_eight_source_proposals_without_touching_original_audit():
    data, layout, sources = case()
    before = copy.deepcopy(data)
    audit = reused(data, layout, sources)
    fields = [f for f in audit["fields"] if f.get("source_reuse")]
    assert len(fields) == 8
    assert {f["proposed_value"] for f in fields} == {
        "Torino", "TO", "12 gennaio 1985", "Milano", "MI", "Roma 24",
        "NREADA85A52L219R", "amministratrice unica"}
    assert all(f["proposal_grounding"] == "GROUNDED_PROPOSAL" for f in fields)
    assert all(f["source_evidence"] and f["written_value"] is None for f in fields)
    assert all(f["source_reuse"]["identity_evidence"] for f in fields)
    assert audit["writes"] == audit["resolved"] == 0 and data == before


def test_street_component_does_not_repeat_printed_type_or_convert_other_street_types():
    data, layout, sources = case()
    field = next(f for f in reused(data, layout, sources)["fields"] if f["candidate_id"] == "p1.s2")
    assert field["proposed_value"] == "Roma 24"
    sources[0]["content"] = sources[0]["content"].replace("Via Roma 24", "Corso Roma 24")
    assert not next(f for f in reused(data, layout, sources)["fields"]
                    if f["candidate_id"] == "p1.s2").get("source_reuse")


def test_personal_role_is_not_transferred_to_a_different_operator():
    data, layout, sources = case()
    sources[0]["content"] = sources[0]["content"].replace(
        "- **Società:** Aurora Servizi S.r.l.", "- **Società:** Altra Società S.r.l.")
    # Keep the original company SOURCE independently available and its proposal valid.
    sources.append({**sources[0], "chunk_id": -8, "file_id": -4,
                    "content": "- **Società:** Aurora Servizi S.r.l."})
    audit = reused(data, layout, sources)
    assert len([f for f in audit["fields"] if f.get("source_reuse")]) == 7
    assert not next(f for f in audit["fields"] if f["candidate_id"] == "p2.s0").get("source_reuse")


@pytest.mark.parametrize("change", ["no_receipt", "wrong_receipt", "partial_name", "form_role",
                                     "homonym", "conflicting_birth", "different_person"])
def test_unproven_identity_or_destination_never_reuses_personal_properties(change):
    data, layout, sources = case()
    if change == "no_receipt":
        data["dialogue"]["answers"] = {}
    elif change == "wrong_receipt":
        data["dialogue"]["answers"]["field:p0.s0"]["value"] = "Bea Mori"
    elif change == "partial_name":
        for value in [data["dialogue"]["answers"]["field:p0.s0"],
                      data["dialogue"]["user_values"]["p0.s0"]]:
            value.update(value="Ada", user_quote="Ada")
    elif change == "form_role":
        data["catalog"]["blocks"][0]["paragraphs"][0]["text_with_fields"] = (
            data["catalog"]["blocks"][0]["paragraphs"][0]["text_with_fields"].replace(
                "sottoscritto", "direttore tecnico"))
    elif change == "homonym":
        sources.append({**sources[0], "chunk_id": -8, "file_id": -4,
                        "content": PROFILE.replace("85A52", "86A52")})
    elif change == "conflicting_birth":
        sources[0]["content"] += "\n" + PROFILE.replace("12 gennaio", "13 gennaio")
    else:
        sources[0]["content"] = sources[0]["content"].replace("Ada Neri", "Bea Mori")
    assert not any(f.get("source_reuse") for f in reused(data, layout, sources)["fields"])


def test_missing_or_revoked_role_does_not_block_biography_or_create_signatory_powers():
    data, layout, sources = case()
    sources[0]["content"] = sources[0]["content"].replace(
        "**amministratrice", "**ex amministratrice")
    audit = reused(data, layout, sources)
    assert len([f for f in audit["fields"] if f.get("source_reuse")]) == 7
    assert not next(f for f in audit["fields"] if f["candidate_id"] == "p2.s0").get("source_reuse")


def test_foreign_project_and_general_kb_cannot_supply_personal_data():
    data, layout, sources = case()
    sources[0]["category"] = "general"
    assert not any(f.get("source_reuse") for f in reused(data, layout, sources)["fields"])
    sources[0].update(chunk_id=7, file_id=3, category=None, project_id="beta", scope="project:beta")
    with pytest.raises(ValueError, match="progetto"):
        reused(data, layout, sources)


def test_homonym_in_another_source_table_is_not_hidden_by_unique_profile_heading():
    from test_compilation_binding_identity import HEADER, ROW, TABLE

    data, layout, sources = case()
    sources.append({**sources[0], "chunk_id": -8, "file_id": -4,
                    "content": TABLE.replace(ROW, ROW.replace(
                        "NREADA85A41L219R", "NREADA86A41L219R"))})
    assert HEADER in sources[1]["content"]
    assert not any(f.get("source_reuse") for f in reused(data, layout, sources)["fields"])


def test_same_identifier_asserted_for_different_people_blocks_reuse():
    data, layout, sources = case()
    sources.append({**sources[0], "chunk_id": -8, "file_id": -4,
                    "content": PROFILE.replace("Ada Neri", "Bea Mori").replace(
                        "Nome: Ada", "Nome: Bea").replace("Cognome: Neri", "Cognome: Mori")})
    assert not any(f.get("source_reuse") for f in reused(data, layout, sources)["fields"])


def test_no_automatic_applicability_or_override_of_user_values():
    data, layout, sources = case()
    data["audit"]["fields"][1]["conditions"] = ["unknown"]
    data["audit"]["conditions"]["unknown"] = {"result": "unknown"}
    data["dialogue"]["user_values"]["p1.s3"] = {"value": "user supplied"}
    audit = reused(data, layout, sources)
    assert not next(f for f in audit["fields"] if f["candidate_id"] == "p1.s3").get("source_reuse")
    conditional = next(f for f in audit["fields"] if f["candidate_id"] == "p0.s1")
    assert conditional["condition_checks"]["unknown"]["result"] == "unknown"


def test_source_presence_and_omission_are_distinct_from_absence_and_user_conditions():
    from test_compilation_binding_identity import TABLE

    from app.compilation_coverage import source_availability_context

    data, _, sources = case()
    sources[0]["content"] += "\n" + TABLE
    field = data["audit"]["fields"][0]
    field["condition_checks"] = {"c": {"result": "unknown"}}
    ctx = source_availability_context(data["audit"], sources)
    assert ctx["absent_from_sources"]["verified"] is False
    assert ctx["not_retrieved"]["verified"] is False
    assert ctx["present_not_proposed"]["records"]
    assert ctx["conditional_or_waiting"] == 1
    ctx = source_availability_context(data["audit"], sources, {"condition:c": {"applies": True}})
    assert ctx["conditional_or_waiting"] == 0


def test_previous_contracts_do_not_acquire_new_questions_or_implicit_reuse():
    from app.compilation_source_reuse import signatory_fields, source_reuse_requested

    data, layout, sources = case()
    data["audit"]["version"] = 11
    assert signatory_fields(data) == {}
    assert source_reuse_requested(data) is False
    assert reused(data, layout, sources) == data["audit"]
