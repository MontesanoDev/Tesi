"""V12 evidence representation regressions; no real provider or export."""

import json
from io import BytesIO

import pytest
from docx import Document

from app.compilation_global import validate_global_proposal
from app.docx_templates import inspect_docx

TITLE = "## 7. Soci della società — scenario sintetico"
HEADER = "| Persona | Numero iscrizione |"
ROW = "| Ada Neri | 6234 |"
TABLE = f"{TITLE}\n\n{HEADER}\n| --- | --- |\n{ROW}\n| Bea Mori | 9999 |"
ROLE = "Responsabile della verifica"
NARRATIVE = f"- **Ada Neri:** {ROLE.lower()} di Alba; numero iscrizione 6234."


def source(content, chunk=-7):
    return {"chunk_id": chunk, "file_id": chunk - 20, "chunk_index": 0,
            "source_name": "fonte.txt", "role": "source", "scope": "global",
            "category": "company", "project_id": None, "content": content}


def case(role="Socio", content=TABLE, quote=ROW):
    doc = Document()
    doc.add_paragraph(role)
    doc.add_paragraph("Numero iscrizione: ________")
    stream = BytesIO()
    doc.save(stream)
    layout = inspect_docx(stream.getvalue(), require_candidates=False)
    proposal = {
        "version": 12, "uncertain_paragraph_ids": [], "conditions": {},
        "choice_groups": {}, "clarifications": {}, "coverage": [],
        "bindings": {"b": {
            "form": {"paragraph_id": "paragraph:0", "quote": role},
            "required_role": role, "source_role": role,
            "evidence": [{"chunk_id": -7, "quote": TITLE},
                         {"chunk_id": -7, "quote": ROW}]}},
        "facts": {"f": {
            "subject": "Ada Neri", "entity": "person", "aspect": "professional",
            "property": "numero_iscrizione", "value": "6234",
            "evidence": [{"chunk_id": -7, "quote": quote}]}},
        "fields": {"p1.s0": {
            "label": "Numero iscrizione", "kind": "data", "disposition": "value",
            "form": {"paragraph_id": "paragraph:1", "quote": "Numero iscrizione: [[p1.s0]]"},
            "conditions": [], "property": "numero_iscrizione", "fact_id": "f",
            "binding_id": "b"}},
    }
    return layout, [source(content)], proposal


def audit(data):
    return validate_global_proposal(json.dumps(data[2]), data[0], "alpha", data[1])


def narrative_case(content=NARRATIVE, quote=None):
    data = case(ROLE)
    data[1][0]["content"] = f"### Dr. Ada Neri - {ROLE}\nNome: Dr. Ada Neri."
    data[1].append(source(content, -9))
    data[2]["bindings"]["b"]["evidence"] = [
        {"chunk_id": -7, "quote": data[1][0]["content"].splitlines()[0]}]
    data[2]["facts"]["f"].update(subject="Dr. Ada Neri", evidence=[
        {"chunk_id": -9, "quote": content if quote is None else quote}])
    return data


def assert_blocked(result):
    field = result["fields"][0]
    assert field["subject_binding"]["result"] == "review_required"
    assert field.get("value") is field["written_value"] is None
    assert "proposal_grounding" not in field
    assert result["writes"] == result["resolved"] == 0


def test_table_title_headers_and_same_subject_row_support_role_without_repetition():
    result = audit(case())
    field = result["fields"][0]
    assert field["subject_binding"]["result"] == "literal_match"
    assert field["value"] == "6234"
    attribution = field["subject_binding"]["record_attributions"][0]
    assert attribution["format"] == "markdown_role_table"
    assert attribution["title"] == TITLE
    assert attribution["header"] == HEADER
    assert attribution["row"] == ROW
    assert field["subject_binding"]["source_evidence"][0]["content_sha256"]
    assert result["bindings"]["b"]["validation"] == "GROUNDED_PROPOSAL"
    assert result["writes"] == result["resolved"] == 0


def test_fact_evidence_can_explicitly_prove_person_role_and_value_in_its_own_chunk():
    result = audit(narrative_case())
    field = result["fields"][0]
    assert field["value"] == "6234"
    check = field["subject_binding"]
    assert check["result"] == "literal_match"
    assert check["record_attributions"][0]["format"] == "named_role_line"
    assert check["source_evidence"][0]["chunk_id"] == -9
    assert check["source_evidence"][0]["category"] == "company"
    assert check["recovery"]["original_validation_errors"]


def test_invalid_value_quote_is_still_rejected_even_with_explicit_person_role():
    content = NARRATIVE + "\n- **Bea Mori:** altra persona, provincia Bari (BA)."
    data = narrative_case(content, NARRATIVE)
    data[2]["facts"]["f"].update(value="Bari (BA)", property="provincia")
    data[2]["fields"]["p1.s0"]["property"] = "provincia"
    result = audit(data)
    assert result["facts"]["f"]["validation"] == "REJECTED"
    assert result["facts"]["f"]["validation_errors"] == [
        "Estratto/valore non presente nella SOURCE"]
    assert result["fields"][0]["validation"] == "REJECTED"
    assert_blocked(result)


@pytest.mark.parametrize("make_case", [case, narrative_case])
def test_recovered_binding_keeps_unknown_form_applicability_nonwritable(make_case):
    data = make_case()
    data[2]["conditions"]["c"] = {
        "form": {"paragraph_id": "paragraph:0", "quote": data[2]["bindings"]["b"]["required_role"]},
        "aspect": "participation", "operator": "semantic", "fact_ids": [],
        "scope": [{"kind": "candidate", "id": "p1.s0"}], "reason": "Da confermare"}
    data[2]["fields"]["p1.s0"].update(
        disposition="conditional", conditions=["c"], reason="Ramo da confermare")
    result = audit(data)
    field = result["fields"][0]
    assert field["subject_binding"]["result"] == "literal_match"
    assert field["proposal_grounding"] == "GROUNDED_PROPOSAL"
    assert field["proposed_value"] == "6234"
    assert field["condition_checks"]["c"]["result"] == "unknown"
    assert field["validation"] == "REVIEW_REQUIRED"
    assert field["value"] is field["written_value"] is None
    assert result["writes"] == result["resolved"] == 0


@pytest.mark.parametrize("content,quote", [
    (TABLE.replace(TITLE, "## Non soci della società"), ROW),
    (TABLE.replace(TITLE, "## Ex soci della società"), ROW),
    (TABLE.replace(TITLE, "## Soci non confermati"), ROW),
    (TABLE.replace(TITLE, "## Soci potenziali"), ROW),
    (TABLE.replace(TITLE, "## Altre persone"), ROW),
    (TABLE.replace(ROW, "| Carla Riva | 6234 |"), "| Carla Riva | 6234 |"),
    (TABLE.replace("| Bea Mori | 9999 |", ROW), ROW),
    (TABLE.replace("| Bea Mori | 9999 |", "| Ada Neri | 9999 |"), ROW),
    (TABLE.replace(ROW, "| Ada Neri | non iscritta, 6234 |"),
     "| Ada Neri | non iscritta, 6234 |"),
    (f"{TITLE}\n\n## Altre persone\n{HEADER}\n| --- | --- |\n{ROW}", ROW),
    (f"{TITLE}\n{HEADER}\n| --- | --- |\n| Bea Mori | 9999 |\n\n"
     f"{HEADER}\n| --- | --- |\n{ROW}", ROW),
    (TABLE.replace("| --- | --- |", "| --- |"), ROW),
    (TABLE.replace(ROW, "| Ada Neri | 6234 | valore extra |"),
     "| Ada Neri | 6234 | valore extra |"),
    (TABLE, "| Bea Mori | 9999 |"),
])
def test_table_proof_cannot_borrow_title_person_or_value_from_another_record(content, quote):
    data = case(content=content, quote=quote)
    data[2]["bindings"]["b"]["evidence"] = [{"chunk_id": -7, "quote": quote}]
    if "9999" in quote:
        data[2]["facts"]["f"]["value"] = "9999"
    assert_blocked(audit(data))


@pytest.mark.parametrize("content,quote", [
    (NARRATIVE.replace("Ada Neri", "Bea Mori"), None),
    (NARRATIVE.replace(ROLE.lower(), "altro ruolo"), None),
    (NARRATIVE.replace(ROLE.lower(), "non " + ROLE.lower()), None),
    (NARRATIVE.replace(ROLE.lower(), "ex " + ROLE.lower()), None),
    (NARRATIVE.replace("numero iscrizione", "non iscritta, numero"), None),
    ("Non nominata: " + NARRATIVE, NARRATIVE),
    (f"- **Ada Neri:** {ROLE.lower()} di Alba;\n- **Bea Mori:** numero 6234.", None),
    (f"- **Ada Neri:** {ROLE.lower()} di Alba; **Bea Mori:** numero 6234.", None),
])
def test_fact_fallback_does_not_join_other_people_negations_or_separate_records(content, quote):
    assert_blocked(audit(narrative_case(content, quote)))


@pytest.mark.parametrize("change", ["general", "foreign", "form", "wrong_role", "wrong_form",
                                     "missing_binding", "invalid_binding_quote"])
def test_fact_recovery_does_not_bypass_scope_form_or_role_checks(change):
    data = narrative_case()
    binding = data[2]["bindings"]["b"]
    if change == "general":
        data[1][-1]["category"] = "general"
    elif change == "foreign":
        data[1][-1].update(scope="project:beta", project_id="beta")
    elif change == "form":
        data[1][-1]["role"] = "form"
    elif change == "wrong_role":
        binding["source_role"] = "Altro ruolo"
    elif change == "wrong_form":
        binding["form"]["quote"] = "Ruolo inesistente"
    elif change == "missing_binding":
        data[2]["bindings"] = {}
    else:
        binding["evidence"][0]["quote"] = "Citazione inventata"
    if change in {"foreign", "form"}:
        with pytest.raises(ValueError):
            audit(data)
    else:
        assert_blocked(audit(data))


def test_every_fact_citation_needs_its_own_person_role_support():
    data = narrative_case()
    other = source("- **Bea Mori:** altro ruolo; numero iscrizione 6234.", -11)
    data[1].append(other)
    data[2]["facts"]["f"]["evidence"].append({"chunk_id": -11, "quote": other["content"]})
    assert_blocked(audit(data))


@pytest.mark.parametrize("role,title", [
    ("Tecnico abilitato", "## 7. Tecnici abilitati della società"),
    ("Responsabile", "## Responsabili della società"),
])
def test_table_singular_plural_is_structural_and_not_specific_to_one_role(role, title):
    data = case(role, TABLE.replace(TITLE, title))
    data[2]["bindings"]["b"]["evidence"][0]["quote"] = title
    assert audit(data)["fields"][0]["value"] == "6234"


def test_valid_fact_record_can_recover_unsupported_binding_role_quote():
    data = narrative_case()
    data[1][0]["content"] = "Nome: Dr. Ada Neri."
    data[2]["bindings"]["b"]["evidence"][0]["quote"] = data[1][0]["content"]
    result = audit(data)
    assert result["bindings"]["b"]["validation"] == "REJECTED"
    assert result["fields"][0]["subject_binding"]["result"] == "literal_match"
    assert result["fields"][0]["value"] == "6234"


def test_fact_recovery_cannot_use_a_role_from_a_different_form_context():
    data = narrative_case()
    doc = data[0].document
    doc.add_paragraph(ROLE)
    stream = BytesIO()
    doc.save(stream)
    layout = inspect_docx(stream.getvalue(), require_candidates=False)
    data = layout, data[1], data[2]
    data[2]["bindings"]["b"]["form"]["paragraph_id"] = "paragraph:2"
    result = audit(data)
    assert result["bindings"]["b"]["validation"] == "GROUNDED_PROPOSAL"
    assert_blocked(result)


def test_v10_historical_table_rule_is_not_implicitly_upgraded():
    data = case()
    data[2]["version"] = 10
    data[2].pop("coverage")
    data[2].pop("choice_groups")
    data[2].pop("clarifications")
    assert_blocked(audit(data))


def later_table_person_case():
    """The second person's row is separated from the table heading by Ada."""
    data = case()
    row = "| Bea Mori | 9999 |"
    data[2]["facts"]["f"].update(subject="Bea Mori", value="9999", evidence=[
        {"chunk_id": -7, "quote": row}])
    data[2]["bindings"]["b"]["evidence"] = [
        {"chunk_id": -7, "quote": TITLE}, {"chunk_id": -7, "quote": row}]
    return data


def test_separate_literal_title_and_later_person_row_are_grounded():
    result = audit(later_table_person_case())
    field = result["fields"][0]
    assert field["value"] == "9999"
    assert field["subject_binding"]["record_attributions"][0]["subject"] == "Bea Mori"
    assert result["writes"] == result["resolved"] == 0


def test_concatenating_non_adjacent_table_parts_remains_rejected():
    data = later_table_person_case()
    data[2]["bindings"]["b"]["evidence"] = [{"chunk_id": -7,
        "quote": f"{TITLE}\n\n{HEADER}\n| --- | --- |\n| Bea Mori | 9999 |"}]
    assert_blocked(audit(data))


def test_separate_spans_do_not_attribute_another_persons_value():
    data = later_table_person_case()
    data[2]["facts"]["f"].update(value="6234", evidence=[{"chunk_id": -7, "quote": ROW}])
    assert_blocked(audit(data))


def test_four_literal_spans_cannot_exceed_binding_contract():
    data = later_table_person_case()
    data[2]["bindings"]["b"]["evidence"] = [{"chunk_id": -7, "quote": q}
        for q in [TITLE, HEADER, "| --- | --- |", "| Bea Mori | 9999 |"]]
    assert_blocked(audit(data))
