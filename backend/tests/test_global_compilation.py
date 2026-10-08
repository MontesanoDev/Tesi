"""Independent one-call experiment: hostile proposals cannot authorize writes."""

import asyncio
import json
import sqlite3
from copy import deepcopy
from io import BytesIO
from zipfile import ZipFile

import httpx
import pytest
from docx import Document
from lxml import etree

from app.config import AISettings
from app.docx_templates import DRAFT_NOTICE, inspect_docx, text_of
from app.global_compilation import GlobalPlan, represent_form, request_body, validate_plan
from scripts import global_compilation_experiment as experiment


@pytest.fixture
def case():
    document = Document()
    document.add_paragraph("Domanda di ammissione")
    table = document.add_table(rows=3, cols=2)
    for row, label in zip(table.rows, ["Operatore economico", "Codice fiscale operatore economico",
                                       "Codice fiscale"], strict=True):
        row.cells[0].text = label
    document.add_paragraph("7.d)da compilare in caso di SOCIETÀ DI INGEGNERIA")
    document.add_paragraph("Requisiti del direttore tecnico:")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Nome e cognome:"
    table.cell(1, 0).text = "Ordine professionale di appartenenza:"
    document.add_paragraph("7.e)da compilare in caso di STUDIO ASSOCIATO")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Nome e cognome:"
    document.add_paragraph("7.f)da compilare in caso di CONSORZIO STABILE")
    document.add_paragraph("Il consorzio è costituito dalle seguenti consorziate:")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Ragione sociale:"
    document.sections[0].footer.paragraphs[0].text = "Nota integrale nel piè di pagina"
    output = BytesIO()
    document.save(output)
    layout = inspect_docx(output.getvalue())
    form = represent_form(layout) | {"project_id": "alpha"}
    text = ("## Società\n- Ragione sociale: Nuova S.r.l.\n"
            "- Tipologia di attività: società di ingegneria.\n- Codice fiscale: ZZ123456.\n"
            "### Ing. Ada Verdi - direttore tecnico\n"
            "- Nominativo completo: Ing. Ada Verdi.\n"
            "- Ordine professionale: Ordine degli Ingegneri di Roma.\n")
    source = {"id": "source-a", "role": "source", "scope": "global", "content": text,
              "lines": [{"line": i, "text": t} for i, t in enumerate(text.splitlines(), 1)]}

    def proof(line):
        return {"origin": "SOURCE", "reference_id": "source-a", "line": line,
                "quote": source["lines"][line - 1]["text"]}

    entities = [
        {"id": "company", "kind": "organization", "name": "Nuova S.r.l.", "role": "operator",
         "legal_type": "società di ingegneria", "evidence": [proof(2)]},
        {"id": "director", "kind": "person", "name": "Ada Verdi", "role": "technical_director",
         "legal_type": "", "evidence": [proof(6)]},
    ]
    sections = [{"section_id": s["id"], "condition_quote": s["condition_quote"],
                 "condition_subject_id": "company", "applicability": "UNKNOWN", "evidence": [],
                 "reason": "Da verificare"} for s in form["sections"]]
    engineering = next(s for s in sections if "INGEGNERIA" in s["condition_quote"])
    engineering.update(applicability="APPLICABLE", evidence=[proof(3)])
    fields = [{"candidate_id": c["id"], "section_id": c["section_id"], "subject_id": "company",
               "label": c["label_hint"], "form_quote": c["label_hint"], "status": "UNKNOWN",
               "value": None, "evidence": [], "reason": "Dato non accertato"}
              for c in form["candidates"]]
    for index, subject, value, line in [
        (0, "company", "Nuova S.r.l.", 2), (1, "company", "ZZ123456", 4),
        (3, "director", "Ing. Ada Verdi", 6),
        (4, "director", "Ordine degli Ingegneri di Roma", 7),
    ]:
        fields[index].update(subject_id=subject, value=value, evidence=[proof(line)],
                             status="PROPOSED")
    plan = {"entities": entities, "sections": sections, "fields": fields, "questions": [],
            "warnings": [], "coverage_count": len(fields), "complete": True}
    return layout, form, [source], [], plan


def validate(case, reason="stop"):
    layout, form, sources, users, plan = case
    return validate_plan(GlobalPlan.model_validate(plan), form, sources, users, layout,
                         finish_reason=reason)


def test_global_plan_accepts_explicit_company_and_director_properties(case):
    result = validate(case)
    assert result["coverage_complete"]
    assert len(result["writes"]) == 4
    assert result["source_proposals"] == result["source_validated"] == 4
    assert not result["document_completed"] and not result["ready_for_submission"]
    assert "t0.r2.c1" not in result["writes"]


@pytest.mark.parametrize("change", ["foreign_id", "duplicate", "omitted",
                                  "not_complete", "foreign_subject", "foreign_section"])
def test_global_coverage_and_scope_never_claim_completion_or_export(case, change):
    plan = case[-1]
    if change == "foreign_id":
        plan["fields"][0]["candidate_id"] = "invented"
    elif change == "duplicate":
        plan["fields"].append(deepcopy(plan["fields"][0]))
    elif change == "omitted":
        plan["fields"].pop()
    elif change == "not_complete":
        plan["complete"] = False
    elif change == "foreign_subject":
        plan["fields"][0]["subject_id"] = "invented"
    else:
        plan["sections"][0]["section_id"] = "invented"
    result = validate(case)
    assert not result["coverage_complete"] and not result["writes"]


@pytest.mark.parametrize("reason", ["length", "content_filter", None])
def test_non_stop_response_cannot_write_even_if_json_is_complete(case, reason):
    result = validate(case, reason)
    assert not result["coverage_complete"] and result["writes"] == {}


@pytest.mark.parametrize("change", ["form_source", "other_project", "foreign_source",
                                  "wrong_line", "wrong_property", "other_person", "wrong_role",
                                  "invented_condition", "wrong_section", "fake_entity_proof"])
def test_literal_quote_cannot_bypass_subject_property_scope_or_condition(case, change):
    plan, source = case[-1], case[2][0]
    if change == "form_source":
        source["role"] = "form"
    elif change == "other_project":
        source["scope"] = "project:beta"
    elif change == "foreign_source":
        plan["fields"][0]["evidence"][0]["reference_id"] = "invented"
    elif change == "wrong_line":
        plan["fields"][0]["evidence"][0]["line"] = 999
    elif change == "wrong_property":
        plan["fields"][1].update(value="Nuova S.r.l.", evidence=plan["fields"][0]["evidence"])
    elif change == "other_person":
        plan["entities"][1]["name"] = "Bea Neri"
    elif change == "wrong_role":
        plan["entities"][1]["role"] = "representative"
    elif change == "invented_condition":
        plan["sections"][1]["condition_quote"] = "Condizione inesistente"
    elif change == "wrong_section":
        plan["fields"][3]["section_id"] = "document"
    else:
        plan["entities"][1]["evidence"] = []
    result = validate(case)
    assert len(result["writes"]) < 4
    assert not result["ready_for_submission"]


def test_legal_type_does_not_exclude_other_branches_or_define_participation(case):
    for section in case[-1]["sections"]:
        if "ASSOCIATO" in section["condition_quote"] or "CONSORZIO" in section["condition_quote"]:
            section.update(applicability="NOT_APPLICABLE",
                           evidence=[case[-1]["sections"][1]["evidence"][0]])
    result = validate(case)
    assert all(s["applicability"] == "UNKNOWN" for s in result["sections"][2:])
    assert not any(f["status"] == "NOT_APPLICABLE" for f in result["fields"])


def test_explicit_user_exclusion_is_scoped_to_its_recorded_condition(case):
    section = case[-1]["sections"][2]
    user = {"id": "user-a", "text": "Non siamo uno studio associato.",
            "section_id": section["section_id"], "condition": "STUDIO ASSOCIATO", "applies": False}
    case[3].append(user)
    section.update(applicability="NOT_APPLICABLE", evidence=[
        {"origin": "USER", "reference_id": "user-a", "line": 0, "quote": user["text"]},
    ])
    result = validate(case)
    assert result["sections"][2]["applicability"] == "NOT_APPLICABLE"
    assert result["fields"][5]["status"] == "NOT_APPLICABLE"
    # A different, local USER condition cannot exclude this whole section.
    user["condition"] = "se diverso"
    assert validate(case)["sections"][2]["applicability"] == "UNKNOWN"


def test_company_cf_cannot_fill_unqualified_personal_cf(case):
    case[-1]["fields"][2].update(status="PROPOSED", value="ZZ123456",
                                 evidence=deepcopy(case[-1]["fields"][1]["evidence"]))
    result = validate(case)
    assert result["fields"][2]["status"] == "UNKNOWN"


def test_known_other_person_cannot_fill_an_ambiguous_personal_subject(case):
    source, plan = case[2][0], case[-1]
    additions = ["### Arch. Bea Neri - referente", "- Nominativo completo: Arch. Bea Neri.",
                 "- Codice fiscale: PERSONAL123."]
    start = len(source["lines"]) + 1
    source["lines"].extend({"line": i, "text": t}
                           for i, t in enumerate(additions, start))
    source["content"] += "\n".join(additions)

    def evidence(line):
        return {"origin": "SOURCE", "reference_id": source["id"], "line": line,
                "quote": source["lines"][line - 1]["text"]}

    plan["entities"].append({"id": "other-person", "kind": "person", "name": "Bea Neri",
                             "role": "other", "legal_type": "", "evidence": [evidence(start+1)]})
    item = plan["fields"][2]
    item.update(status="PROPOSED", subject_id="other-person", value="PERSONAL123",
                evidence=[evidence(start+2)])
    result = validate(case)
    assert item["candidate_id"] not in result["writes"]
    assert "non determinabile" in result["fields"][2]["errors"][0]


def test_valid_independent_source_support_survives_an_unsupported_pdf_property(case):
    item = case[-1]["fields"][0]
    other = deepcopy(case[2][0])
    other.update(id="pdf-source", content="ANAGRAFICA Nuova S.r.l.",
                 lines=[{"line": 1, "text": "ANAGRAFICA Nuova S.r.l."}])
    case[2].append(other)
    item["evidence"].insert(0, {"origin": "SOURCE", "reference_id": "pdf-source", "line": 1,
                               "quote": "ANAGRAFICA Nuova S.r.l."})
    result = validate(case)
    assert result["writes"][item["candidate_id"]] == "Nuova S.r.l."
    field = result["fields"][0]
    assert len(field["validated_evidence"]) == len(field["rejected_evidence"]) == 1
    assert "fuori dal catalogo" in field["rejected_evidence"][0]["error"]


def test_composed_non_literal_values_are_not_repaired_into_other_values(case):
    item = case[-1]["fields"][1]
    item["value"] = "ZZ123456 (valore simulato)"
    result = validate(case)
    assert item["candidate_id"] not in result["writes"]
    assert result["fields"][1]["rejected_evidence"]


def test_wrong_declared_count_is_a_warning_and_actual_ids_authorize_partial_export(case):
    case[-1]["coverage_count"] = 400
    result = validate(case)
    assert result["candidate_ids_complete"]
    assert result["coverage_complete"] and len(result["writes"]) == 4
    assert result["actual_coverage_count"] == len(case[0].candidate_ids)
    assert result["returned_candidate_items"] == len(case[-1]["fields"])
    assert result["returned_unique_candidate_ids"] == len(case[0].candidate_ids)
    assert result["declared_coverage_count"] == 400
    assert any("coverage_count incoerente" in e for e in result["global_warnings"])
    assert result["global_errors"] == []


def add_source(case, text, source_id="pdf-source"):
    source = {"id": source_id, "role": "source", "scope": "global", "content": text,
              "lines": [{"line": i, "text": t} for i, t in enumerate(text.splitlines(), 1)]}
    case[2].append(source)

    def evidence(line):
        return {"origin": "SOURCE", "reference_id": source_id, "line": line,
                "quote": source["lines"][line - 1]["text"]}
    return source, evidence


def test_labeled_pdf_table_supports_company_properties_without_colons(case):
    _, evidence = add_source(case, "DENOMINAZIONE Nuova S.r.l.\n"
                                  "CODICE FISCALE E PARTITA IVA ZZ123456 (dato simulato)")
    first, fiscal = case[-1]["fields"][:2]
    first["evidence"] = [evidence(1)]
    fiscal.update(value="ZZ123456 (dato simulato)", evidence=[evidence(2)])
    case[-1]["entities"][0]["evidence"] = [evidence(1)]
    result = validate(case)
    assert result["writes"][first["candidate_id"]] == first["value"]
    assert result["writes"][fiscal["candidate_id"]] == fiscal["value"]
    assert result["fields"][1]["validated_evidence"] == fiscal["evidence"]
    assert not result["ready_for_submission"]


@pytest.mark.parametrize("change", ["other_owner", "no_owner", "wrong_property", "wrong_scope",
                                  "narrative", "partial_quote", "wrong_role"])
def test_pdf_parser_extension_cannot_bypass_grounding(case, change):
    text = "DENOMINAZIONE Nuova S.r.l.\nCODICE FISCALE ZZ123456"
    if change == "other_owner":
        text = text.replace("Nuova", "Diversa")
    elif change == "no_owner":
        text = "CODICE FISCALE ZZ123456"
    elif change == "wrong_property":
        text = text.replace("CODICE FISCALE", "TELEFONO")
    elif change == "narrative":
        text = text.replace("CODICE FISCALE", "Il documento descrive il CODICE FISCALE")
    source, evidence = add_source(case, text)
    item = case[-1]["fields"][1]
    item["evidence"] = [evidence(len(source["lines"]))]
    if change == "wrong_scope":
        source["scope"] = "project:another"
    elif change == "partial_quote":
        item["evidence"][0]["quote"] = "CODICE FISCALE"
    elif change == "wrong_role":
        item["subject_id"] = "director"
    assert item["candidate_id"] not in validate(case)["writes"]


def test_pdf_multiple_named_companies_keep_their_own_properties(case):
    _, evidence = add_source(case, "DENOMINAZIONE Nuova S.r.l.\nCODICE FISCALE FIRST123\n"
                                  "DENOMINAZIONE Diversa S.r.l.\nCODICE FISCALE ZZ123456")
    item = case[-1]["fields"][1]
    item["evidence"] = [evidence(4)]
    assert item["candidate_id"] not in validate(case)["writes"]


def test_pdf_director_and_adjacent_order_keep_a_personal_owner(case):
    _, evidence = add_source(case, "DENOMINAZIONE Nuova S.r.l.\n"
                                  "DIRETTORE TECNICO Ing. Ada Verdi - iscrizione n. 9876,\n"
                                  "Ordine degli Ingegneri di Roma\n"
                                  "CODICE FISCALE ZZ123456")
    plan = case[-1]
    plan["entities"][1]["evidence"] = [evidence(2)]
    plan["fields"][3]["evidence"] = [evidence(2)]
    plan["fields"][4]["evidence"] = [evidence(3)]
    plan["fields"][1]["evidence"] = [evidence(4)]
    result = validate(case)
    assert plan["fields"][3]["candidate_id"] in result["writes"]
    assert plan["fields"][4]["candidate_id"] in result["writes"]
    assert plan["fields"][1]["candidate_id"] not in result["writes"]


@pytest.mark.parametrize("condition,accepted", [
    ("[da compilare solo a cura di società di professionisti, società di ingegneria, "
     "consorzi e loro consorziate esecutrici]", True),
    ("[da compilare solo a cura di società di professionisti, società di ingegneria, "
     "consorzi art. 77, co. 3 lettera f) e loro consorziate esecutrici e società tra "
     "professionisti (v. art. 12 L. 100/2020 e D.M. n. 21/2021)]", True),
    ("da compilare a cura di società di ingegneria e consorzio stabile", False),
    ("da compilare a cura di società di ingegneria, solo se consorziata esecutrice", False),
    ("da compilare a cura di società di ingegneria, con poteri di rappresentanza", False),
    ("da compilare a cura di società di ingegneria, iscritte all'albo", False),
    ("da compilare in caso di SOCIETÀ DI INGEGNERIA nel caso di rappresentante", False),
    ("da compilare in caso di SOCIETÀ DI INGEGNERIA di cui all'art. 1, solo se consorziata", False),
    ("da compilare in caso di partecipazione singola", False),
])
def test_explicit_type_enumeration_preserves_other_conditions(case, condition, accepted):
    case[1]["sections"][1]["condition_quote"] = condition
    case[-1]["sections"][1]["condition_quote"] = condition
    result = validate(case)
    assert (result["sections"][1]["applicability"] == "APPLICABLE") == accepted
    assert (case[-1]["fields"][3]["candidate_id"] in result["writes"]) == accepted


def test_a_bad_section_proof_does_not_cancel_an_independent_valid_one(case):
    section = case[-1]["sections"][1]
    section["evidence"].insert(0, {"origin": "SOURCE", "reference_id": "invented",
                                   "line": 1, "quote": "Prova non valida"})
    result = validate(case)
    assert result["sections"][1]["applicability"] == "APPLICABLE"
    assert len(result["sections"][1]["rejected_evidence"]) == 1


@pytest.mark.parametrize("change", [None, "other_owner", "negated", "different_condition",
                                  "unproven_subject", "missing_fact", "truncated_quote"])
def test_local_address_equality_exclusion_requires_exact_source_relation(case, change):
    item = case[-1]["fields"][0]
    candidate = case[1]["candidates"][0]
    candidate.update(label_hint="Sede operativa (se diversa dalla sede legale)",
                     local_form="Sede operativa (se diversa dalla sede legale) | ")
    text = ("DENOMINAZIONE Nuova S.r.l.\n"
            "Rapporto tra le sedi: sede operativa coincidente con la sede legale.")
    if change == "other_owner":
        text = text.replace("Nuova", "Diversa")
    elif change == "negated":
        text = text.replace("coincidente", "non coincidente")
    elif change == "missing_fact":
        text = text.replace("coincidente con", "non documentata rispetto a")
    _, evidence = add_source(case, text)
    item.update(form_quote=candidate["label_hint"], status="NOT_APPLICABLE", value=None,
                evidence=[evidence(2)])
    if change == "different_condition":
        candidate["label_hint"] = candidate["local_form"] = "Sede operativa (se in consorzio)"
        item["form_quote"] = candidate["label_hint"]
    elif change == "unproven_subject":
        item["subject_id"] = "director"
    elif change == "truncated_quote":
        item["evidence"][0]["quote"] = "Rapporto tra le sedi"
    field = validate(case)["fields"][0]
    independently_proven = change in {None, "unproven_subject", "truncated_quote"}
    assert (field["status"] == "NOT_APPLICABLE") == independently_proven
    if independently_proven:
        assert field["condition_subject"]["entity_id"] == "company"
        assert field["validated_condition_evidence"][0]["quote"] == text.splitlines()[1]
    if change == "unproven_subject":
        assert field["proposal_errors"]
    if change == "truncated_quote":
        # Bad model quote still rejected; independent full proof used.
        assert field["rejected_evidence"]


def test_mixed_question_preserves_missing_data_without_reasking_resolved_data(case):
    resolved, missing = case[-1]["fields"][0], case[-1]["fields"][2]
    case[-1]["questions"] = [{"question": "Confermare la società e fornire il dato mancante?",
                              "candidate_ids": [resolved["candidate_id"], missing["candidate_id"]],
                              "section_ids": ["document"]}]
    result = validate(case)
    assert not result["questions"]
    review = result["questions_requiring_review"][0]
    assert review["resolved_candidate_ids"] == [resolved["candidate_id"]]
    assert review["unresolved_candidate_ids"] == [missing["candidate_id"]]


def test_full_form_contains_footer_and_non_candidate_text(case):
    form = case[1]
    encoded = json.dumps(form, ensure_ascii=False)
    assert "Nota integrale nel piè di pagina" in encoded
    assert "consorzio è costituito dalle seguenti consorziate" in encoded
    assert len(form["candidates"]) == len(case[0].candidate_ids)


def test_section_headings_inside_tables_govern_later_tables_without_losing_text():
    document = Document()
    document.add_table(rows=1, cols=1).cell(0, 0).text = (
        "3.b)da compilare in caso di SOCIETÀ DI INGEGNERIA"
    )
    document.add_table(rows=1, cols=2).cell(0, 0).text = "Denominazione sociale:"
    document.add_paragraph("Requisiti del direttore tecnico:")
    document.add_table(rows=1, cols=2).cell(0, 0).text = "Nome e cognome:"
    document.add_table(rows=1, cols=1).cell(0, 0).text = (
        "3.c)da compilare in caso di CONSORZIO STABILE"
    )
    document.add_paragraph("Il consorzio è costituito dalle seguenti consorziate:")
    document.add_table(rows=1, cols=2).cell(0, 0).text = "Ragione sociale:"
    document.add_paragraph("Testo finale senza candidate, integralmente conservato.")
    output = BytesIO()
    document.save(output)
    form = represent_form(inspect_docx(output.getvalue()))
    assert len(form["sections"]) == 3
    assert form["candidates"][0]["section_id"] == form["sections"][1]["id"]
    assert form["candidates"][1]["subject_role_hint"] == "technical_director"
    assert form["candidates"][2]["subject_role_hint"] == "consortium_member"
    assert form["candidates"][2]["section_id"] == form["sections"][2]["id"]
    assert "Testo finale senza candidate" in json.dumps(form)


def test_request_does_not_apply_ollama_context_or_small_session_output_limit(case):
    settings = AISettings("secret", "generic-deepseek", "https://api.deepseek.com",
                          context_window=32768)
    body = request_body(settings, case[1], case[2], case[3])
    assert body["max_tokens"] == 98304
    assert "secret" not in json.dumps(body)
    assert body["thinking"]["type"] == "disabled"
    payload = json.loads(body["messages"][1]["content"])
    assert set(payload) == {"FORM", "SOURCE", "USER", "schema_output"}


def test_oversized_input_is_rejected_without_truncation(case):
    settings = AISettings("secret", "generic-deepseek", "https://api.deepseek.com")
    with pytest.raises(ValueError, match="nessun taglio"):
        request_body(settings, {"content": "X" * 1_000_000}, case[2], case[3])


def test_real_entry_point_sends_exactly_one_request_and_protects_originals(case, tmp_path,
                                                                        monkeypatch):
    layout, form, sources, users, plan = case
    output = tmp_path / "isolated"
    output.mkdir()
    settings = AISettings("secret", "generic-deepseek", "https://api.deepseek.com")
    body = request_body(settings, form, sources, users)
    experiment.dump(output / "manifest.json", {
        "status": "prepared", "project_id": "alpha", "model": settings.model,
        "base_url": settings.base_url,
        "request_sha256": experiment.digest(json.dumps(body, ensure_ascii=False).encode()),
    })
    experiment.dump(output / "request.json", body)
    experiment.dump(output / "inputs.json", {"FORM": form, "SOURCE": sources, "USER": users})
    (output / "original.docx").write_bytes(layout.original)
    experiment.dump(output / "preservation-before.json", {"tables": {}, "files": {}})
    monkeypatch.setattr(experiment, "fingerprint", lambda: {"tables": {}, "files": {}})
    monkeypatch.setattr(experiment, "resolve_project_settings", lambda _: settings)
    requests = []

    def respond(request):
        requests.append(request)
        return httpx.Response(200, json={"choices": [{"finish_reason": "stop",
                              "message": {"content": json.dumps(plan)}}],
                              "usage": {"prompt_tokens": 123, "completion_tokens": 456}})

    client_class = httpx.AsyncClient
    monkeypatch.setattr(experiment.httpx, "AsyncClient", lambda **kwargs: client_class(
        transport=httpx.MockTransport(respond), **kwargs,
    ))
    asyncio.run(experiment.live(output))
    assert len(requests) == 1
    with pytest.raises(ValueError, match="vietati retry"):
        asyncio.run(experiment.live(output))
    assert len(requests) == 1
    report = json.loads((output / "validation.json").read_text())
    assert report["usage"]["prompt_tokens"] == 123 and len(report["writes"]) == 4
    assert (output / "bozza-globale-parziale.docx").exists()
    assert (output / "original.docx").read_bytes() == layout.original


def frozen_run(case, path, *, reason="stop"):
    layout, form, sources, users, plan = case
    path.mkdir()
    body = request_body(AISettings("secret", "generic-deepseek", "https://api.deepseek.com"),
                        form, sources, users)
    experiment.dump(path / "request.json", body)
    experiment.dump(path / "manifest.json", {
        "status": "responded", "model_calls": 1,
        "request_sha256": experiment.digest(json.dumps(body, ensure_ascii=False).encode()),
    })
    experiment.dump(path / "inputs.json", {"FORM": form, "SOURCE": sources, "USER": users})
    experiment.dump(path / "response.json", {
        "choices": [{"finish_reason": reason, "message": {"content": json.dumps(plan)}}],
    })
    (path / "original.docx").write_bytes(layout.original)


def test_offline_replay_uses_zero_provider_calls_and_exports_only_verified_cells(case, tmp_path,
                                                                               monkeypatch):
    historical, output = tmp_path / "historical", tmp_path / "replay"
    case[-1]["coverage_count"] = 400
    frozen_run(case, historical)
    before = {p.name: p.read_bytes() for p in historical.iterdir()}

    def forbidden(*args, **kwargs):
        pytest.fail("Offline replay must not resolve profiles or call a provider")

    monkeypatch.setattr(experiment, "resolve_project_settings", forbidden)
    monkeypatch.setattr(experiment.httpx, "AsyncClient", forbidden)
    report = experiment.offline_replay(historical, output)
    assert len(report["writes"]) == 4
    assert not report["document_completed"] and not report["ready_for_submission"]
    assert before == {p.name: p.read_bytes() for p in historical.iterdir()}
    audit = json.loads((output / "offline-replay.json").read_text())
    assert audit["additional_model_calls"] == 0 and audit["historical_inputs_unchanged"]
    assert audit["copied_inputs_identical"]
    with (ZipFile(BytesIO(case[0].original)) as original,
          ZipFile(output / report["export"]["path"]) as exported):
        assert original.namelist() == exported.namelist()
        for name in original.namelist():
            if name != "word/document.xml":
                assert original.read(name) == exported.read(name)
    exported_doc = Document(output / report["export"]["path"])
    assert exported_doc.paragraphs[0].text == DRAFT_NOTICE
    for ti, table in enumerate(case[0].document.tables):
        for ri, row in enumerate(table.rows):
            for ci, cell in enumerate(row.cells):
                cid = f"t{ti}.r{ri}.c{ci}"
                new_cell = exported_doc.tables[ti].rows[ri].cells[ci]
                if cid in report["writes"]:
                    assert text_of(new_cell._tc) == report["writes"][cid]
                else:
                    assert etree.tostring(new_cell._tc) == etree.tostring(cell._tc)
    with pytest.raises(FileExistsError):
        experiment.offline_replay(historical, output)


@pytest.mark.parametrize("change", ["missing", "duplicate", "foreign", "truncated", "schema"])
def test_offline_replay_refuses_export_for_structural_or_truncation_errors(case, tmp_path, change):
    plan = case[-1]
    plan["coverage_count"] = 400  # This warning must never mask a real failure.
    if change == "missing":
        plan["fields"].pop()
    elif change == "duplicate":
        plan["fields"].append(deepcopy(plan["fields"][0]))
    elif change == "foreign":
        plan["fields"][0]["candidate_id"] = "foreign"
    elif change == "schema":
        plan["fields"][0]["status"] = "UNSAFE"
    historical, output = tmp_path / "historical", tmp_path / "replay"
    frozen_run(case, historical, reason="length" if change == "truncated" else "stop")
    report = experiment.offline_replay(historical, output)
    assert not report["coverage_complete"] and report["writes"] == {}
    assert not report["export"]["generated"]
    assert not (output / "bozza-globale-parziale.docx").exists()


@pytest.mark.parametrize("change", ["request", "inputs", "original"])
def test_offline_replay_rejects_altered_frozen_input(case, tmp_path, change):
    historical, output = tmp_path / "historical", tmp_path / "replay"
    frozen_run(case, historical)
    if change == "request":
        path = historical / "request.json"
        data = json.loads(path.read_text())
        data["temperature"] = 0.9
        experiment.dump(path, data)
    elif change == "inputs":
        path = historical / "inputs.json"
        data = json.loads(path.read_text())
        data["SOURCE"][0]["scope"] = "project:another"
        experiment.dump(path, data)
    else:
        document = Document()
        document.add_paragraph("Originale diverso")
        document.save(historical / "original.docx")
    with pytest.raises(ValueError):
        experiment.offline_replay(historical, output)
    assert not (output / "bozza-globale-parziale.docx").exists()


def test_snapshot_and_source_selection_preserve_form_and_project_isolation(tmp_path, monkeypatch):
    # Read-only loader: another project's source, generated artifacts and templates
    # cannot enter SOURCE, while USER project facts keep their distinct provenance.
    dbpath, storage, knowledge = tmp_path / "app.db", tmp_path / "uploads", tmp_path / "knowledge"
    storage.mkdir()
    knowledge.mkdir()
    document = Document()
    document.add_table(rows=1, cols=2).cell(0, 0).text = "Operatore economico"
    original = BytesIO()
    document.save(original)
    (storage / "original.docx").write_bytes(original.getvalue())
    (storage / "source.txt").write_text("Tutto il contenuto della fonte.\nUltima linea.")
    (knowledge / "project-facts.md").write_text("Decisione USER di progetto")
    state = {"template_sha256": experiment.digest(original.getvalue()), "fields": []}
    with sqlite3.connect(dbpath) as db:
        db.executescript("""
            CREATE TABLE compilation_sessions(id, project_id, form_id, version, state_json);
            CREATE TABLE project_files(id, project_id, name, kind, storage_path);
            CREATE TABLE global_documents(id, name, category, storage_path);
        """)
        db.execute("INSERT INTO compilation_sessions VALUES (?,?,?,?,?)",
                   ("session", "alpha", 1, 7, json.dumps(state)))
        db.executemany("INSERT INTO project_files VALUES (?,?,?,?,?)", [
            (1, "alpha", "original.docx", "form", "original.docx"),
            (2, "alpha", "source.txt", "source", "source.txt"),
            (3, "beta", "foreign.txt", "source", "foreign.txt"),
            (4, "alpha", "draft.md", "artifact", "draft.md"),
            (5, "alpha", "project-facts.md", "artifact", "project-facts.md"),
        ])
    for name, value in [("MAPI_DB_PATH", dbpath), ("MAPI_STORAGE_PATH", storage),
                        ("MAPI_KNOWLEDGE_PATH", knowledge)]:
        monkeypatch.setenv(name, str(value))
    monkeypatch.setattr(experiment, "resolve_project_settings", lambda _: AISettings(
        "secret", "generic-deepseek", "https://api.deepseek.com",
    ))
    before = experiment.fingerprint()
    experiment.prepare("session", tmp_path / "run", 98304)
    data = json.loads((tmp_path / "run/inputs.json").read_text())
    assert [s["id"] for s in data["SOURCE"]] == ["project:2"]
    assert data["SOURCE"][0]["content"].endswith("Ultima linea.")
    assert data["USER"][0]["origin"] == "USER"
    assert before == experiment.fingerprint()


@pytest.mark.parametrize("model_state", ["UNKNOWN", "APPLICABLE", "NOT_APPLICABLE"])
@pytest.mark.parametrize("wrong_subject", [False, True])
def test_verified_user_section_precedes_llm_state_and_subject(case, model_state, wrong_subject):
    section = case[-1]["sections"][2]
    case[3].append({"id": "recorded-choice", "section_id": section["section_id"],
                    "condition": "STUDIO ASSOCIATO", "applies": False,
                    "text": "Decisione verificata: non siamo uno studio associato."})
    section.update(applicability=model_state, evidence=[],
                   condition_subject_id="director" if wrong_subject else "company")
    field = case[-1]["fields"][5]
    field.update(status="PROPOSED", value="Ing. Ada Verdi", subject_id="director")
    case[-1]["questions"] = [{"question": "Si applica lo studio associato?",
                              "candidate_ids": [field["candidate_id"]],
                              "section_ids": [section["section_id"]]}]
    result = validate(case)
    assert result["sections"][2]["applicability"] == "NOT_APPLICABLE"
    assert result["sections"][2]["decision_origin"] == "USER"
    assert result["sections"][2]["condition_subject_id"] == "company"
    assert result["fields"][5]["status"] == "NOT_APPLICABLE"
    assert field["candidate_id"] not in result["writes"]
    assert not result["questions"]
    assert section["applicability"] == model_state  # Raw proposal not rewritten.


@pytest.mark.parametrize("model_state", ["UNKNOWN", "NOT_APPLICABLE", "PROPOSED"])
def test_verified_user_field_value_is_preserved_without_llm_role_or_citation(case, model_state):
    field = case[-1]["fields"][0]
    case[3].append({"id": "recorded-value", "candidate_id": field["candidate_id"],
                    "value": "Nuova S.r.l.", "text": "Operatore confermato: Nuova S.r.l."})
    field.update(status=model_state, value="Diversa S.r.l." if model_state == "PROPOSED" else None,
                 subject_id="director", evidence=[])
    result = validate(case)
    assert result["writes"][field["candidate_id"]] == "Nuova S.r.l."
    assert result["fields"][0]["decision_origin"] == "USER"
    assert result["fields"][0]["preserved_user_decisions"][0]["id"] == "recorded-value"


@pytest.mark.parametrize("invalid", ["other_project", "source_origin", "unverified", "duplicate"])
def test_only_verified_in_scope_user_records_are_authoritative(case, invalid):
    section = case[-1]["sections"][2]
    user = {"id": "recorded-choice", "section_id": section["section_id"],
            "condition": "STUDIO ASSOCIATO", "applies": False, "text": "Non studio associato"}
    if invalid == "other_project":
        user["project_id"] = "foreign"
    elif invalid == "source_origin":
        user["origin"] = "SOURCE"
    elif invalid == "unverified":
        user["status"] = "UNKNOWN"
    else:
        case[3].append(deepcopy(user))
    case[3].append(user)
    assert validate(case)["sections"][2]["applicability"] == "UNKNOWN"


def test_conflicting_verified_user_decisions_are_preserved_without_choosing(case):
    section = case[-1]["sections"][2]
    for i, answer in enumerate([True, False]):
        case[3].append({"id": f"choice-{i}", "section_id": section["section_id"],
                        "condition": "STUDIO ASSOCIATO", "applies": answer,
                        "text": f"Decisione registrata: {answer}"})
    result = validate(case)["sections"][2]
    assert result["applicability"] == "UNKNOWN" and result["errors"]
    assert len(result["validated_evidence"]) == 2


@pytest.mark.parametrize("model_state", ["UNKNOWN", "NOT_APPLICABLE"])
def test_positive_source_condition_is_checked_before_clarifications(case, model_state):
    section = case[-1]["sections"][1]
    section.update(applicability=model_state, condition_subject_id="director", evidence=[])
    case[-1]["questions"] = [{"question": "È applicabile questa sezione?",
                              "candidate_ids": [], "section_ids": [section["section_id"]]}]
    result = validate(case)
    checked = result["sections"][1]
    assert checked["applicability"] == "APPLICABLE" and checked["decision_origin"] == "SOURCE"
    assert checked["condition_subject_id"] == "company"
    assert checked["proposal_errors"]
    assert len(result["writes"]) == 4 and not result["questions"]


@pytest.mark.parametrize("invalid", [None, "other_owner", "wrong_scope", "form", "qualified"])
def test_negative_section_needs_explicit_source_for_correct_operator(case, invalid):
    source, _ = add_source(case, "- Ragione sociale: Nuova S.r.l.\n- Studio associato: No.")
    if invalid == "other_owner":
        source["lines"][0]["text"] = "- Ragione sociale: Esterna S.r.l."
    elif invalid == "wrong_scope":
        source["scope"] = "project:foreign"
    elif invalid == "form":
        source["role"] = "form"
    elif invalid == "qualified":
        source["lines"][1]["text"] = "- Studio associato: No, salvo altri incarichi."
    section = case[-1]["sections"][2]
    section.update(applicability="NOT_APPLICABLE", condition_subject_id="director", evidence=[])
    checked = validate(case)["sections"][2]
    assert (checked["applicability"] == "NOT_APPLICABLE") == (invalid is None)
    if invalid is None:
        assert checked["condition_subject_id"] == "company"
        assert checked["validated_evidence"][0]["quote"] == "- Studio associato: No."


def test_source_conflict_does_not_cancel_authoritative_user_but_blocks_source_guess(case):
    add_source(case, "- Ragione sociale: Nuova S.r.l.\n- Studio associato: No.", "source-no")
    add_source(case, "- Ragione sociale: Nuova S.r.l.\n- Studio associato: Si.", "source-yes")
    assert validate(case)["sections"][2]["applicability"] == "UNKNOWN"
    sid = case[-1]["sections"][2]["section_id"]
    case[3].append({"id": "final-choice", "section_id": sid, "condition": "STUDIO ASSOCIATO",
                    "applies": False, "text": "Decisione verificata: non studio associato."})
    assert validate(case)["sections"][2]["applicability"] == "NOT_APPLICABLE"


@pytest.mark.parametrize("model_state", ["UNKNOWN", "PROPOSED"])
def test_local_source_condition_precedes_value_or_missing_model_decision(case, model_state):
    field, candidate = case[-1]["fields"][0], case[1]["candidates"][0]
    candidate.update(label_hint="Sede operativa (se diversa dalla sede legale)",
                     local_form="Sede operativa (se diversa dalla sede legale)")
    add_source(case, "- Ragione sociale: Nuova S.r.l.\n"
                    "- Rapporto tra le sedi: sede operativa coincidente con la sede legale.")
    field.update(status=model_state, value="Via Roma 8" if model_state == "PROPOSED" else None,
                 form_quote=candidate["local_form"], evidence=[])
    case[-1]["questions"] = [{"question": "Qual è la sede operativa?",
                              "candidate_ids": [field["candidate_id"]],
                              "section_ids": ["document"]}]
    result = validate(case)
    assert result["fields"][0]["status"] == "NOT_APPLICABLE"
    assert field["candidate_id"] not in result["writes"] and not result["questions"]


def role_case():
    document = Document()
    for caption, label in [
        ("I soggetti titolari di poteri di amministrazione e rappresentanza sono:",
         "Nome e cognome"),
        ("Rivestono la qualifica di Direttore Tecnico i seguenti soggetti:", "Nome e cognome"),
        ("Il consorzio è costituito dalle seguenti consorziate:", "Ragione sociale"),
    ]:
        document.add_paragraph(caption)
        document.add_table(rows=1, cols=2).cell(0, 0).text = label
    data = BytesIO()
    document.save(data)
    layout = inspect_docx(data.getvalue())
    form = represent_form(layout) | {"project_id": "generic-project"}
    # Mimic a historical frozen FORM with imprecise hints; original captions remain complete.
    for c in form["candidates"]:
        c["subject_context"], c["subject_role_hint"] = "", "operator"
    text = ("- Ragione sociale: Consorzio Alfa.\n"
            "### Ing. Lina Blu - legale rappresentante\n"
            "- Nominativo completo: Ing. Lina Blu.\n"
            "### Ing. Ugo Gialli - direttore tecnico\n"
            "- Nominativo completo: Ing. Ugo Gialli.\n"
            "### Impresa Beta S.r.l. - consorziata esecutrice\n"
            "- Ragione sociale: Impresa Beta S.r.l.\n")
    source = {"id": "roles", "role": "source", "scope": "project:generic-project",
              "content": text, "lines": [{"line": i, "text": t}
                                         for i, t in enumerate(text.splitlines(), 1)]}
    def evidence(line):
        return {"origin": "SOURCE", "reference_id": "roles", "line": line,
                "quote": source["lines"][line-1]["text"]}
    entities = [{"id": eid, "name": name, "kind": kind, "role": role,
                 "legal_type": "", "evidence": [evidence(line)]}
                for eid, name, kind, role, line in [
                    ("operator", "Consorzio Alfa", "organization", "operator", 1),
                    ("representative", "Ing. Lina Blu", "person", "representative", 3),
                    ("director", "Ing. Ugo Gialli", "person", "technical_director", 5),
                    ("member", "Impresa Beta S.r.l.", "organization", "consortium_member", 7),
                ]]
    fields = [{"candidate_id": c["id"], "section_id": c["section_id"], "subject_id": eid,
               "label": c["label_hint"], "form_quote": c["label_hint"], "status": "PROPOSED",
               "value": name, "evidence": [evidence(line)], "reason": "Prova esplicita"}
              for c, (eid, name, line) in zip(form["candidates"], [
                  ("representative", "Ing. Lina Blu", 3), ("director", "Ing. Ugo Gialli", 5),
                  ("member", "Impresa Beta S.r.l.", 7),
              ], strict=True)]
    plan = {"entities": entities, "sections": [{"section_id": "document", "condition_quote": "",
             "condition_subject_id": "", "applicability": "APPLICABLE", "evidence": [],
             "reason": "Incondizionata"}], "fields": fields, "questions": [], "warnings": [],
            "coverage_count": len(fields), "complete": True}
    return layout, form, [source], [], plan


def test_original_form_captions_preserve_representative_director_and_member_roles():
    result = validate(role_case())
    assert len(result["writes"]) == 3
    assert [f["required_role"] for f in result["fields"]] == [
        "representative", "technical_director", "consortium_member",
    ]


@pytest.mark.parametrize("target,donor", [(0, 1), (1, 0), (2, 0), (2, 1), (2, "operator")])
def test_literal_fact_cannot_be_transferred_to_another_form_role(target, donor):
    data = role_case()
    field = data[-1]["fields"][target]
    if donor == "operator":
        entity = data[-1]["entities"][0]
        field.update(subject_id=entity["id"], value=entity["name"], evidence=entity["evidence"])
    else:
        other = data[-1]["fields"][donor]
        field.update(subject_id=other["subject_id"], value=other["value"],
                     evidence=other["evidence"])
    result = validate(data)
    assert result["fields"][target]["status"] == "UNKNOWN"
    assert field["candidate_id"] not in result["writes"]


def participation_case(case):
    candidate, field = case[1]["candidates"][0], case[-1]["fields"][0]
    candidate.update(label_hint="Forma di partecipazione:", local_form="Forma di partecipazione:")
    field.update(status="UNKNOWN", value=None, evidence=[], form_quote=candidate["local_form"])
    return candidate["id"]


def test_unknown_participation_has_one_stable_question_even_if_llm_duplicates_or_omits(case):
    cid = participation_case(case)
    case[-1]["questions"] = [{"question": text, "candidate_ids": ids, "section_ids": ["document"]}
                              for text, ids in [
                                  ("Qual è la modalità di partecipazione?", [cid]),
                                  ("Come intendete partecipare?", []),
                                  ("Si partecipa in forma singola o RTI?", []),
                                  ("Forma di partecipazione e dati del firmatario?",
                                   [cid, case[1]["candidates"][2]["id"]]),
                              ]]
    first = validate(case)
    assert first["participation"]["status"] == "UNKNOWN"
    assert first["participation"]["value"] is None and len(first["questions"]) == 1
    assert first["questions_requiring_review"][0]["unresolved_candidate_ids"] == [
        case[1]["candidates"][2]["id"],
    ]
    case[-1]["questions"] = []
    second = validate(case)
    assert second["questions"] == first["questions"]
    assert first["participation"]["decision_id"] == second["participation"]["decision_id"]


@pytest.mark.parametrize("origin", ["USER", "SOURCE"])
def test_explicit_current_practice_participation_is_resolved_before_questions(case, origin):
    cid = participation_case(case)
    if origin == "USER":
        case[3].append({"id": "participation", "condition": "Modalità di partecipazione",
                        "value": "in RTI", "text": "Decisione verificata: in RTI"})
    else:
        source, _ = add_source(case, "- Ragione sociale: Nuova S.r.l.\n"
                                    "- Modalità di partecipazione alla gara: in RTI.")
        source["scope"] = "project:alpha"
    case[-1]["questions"] = [{"question": "Forma di partecipazione?", "candidate_ids": [cid],
                              "section_ids": ["document"]}]
    result = validate(case)
    assert result["participation"]["status"] == "RESOLVED"
    assert result["participation"]["origin"] == origin
    assert not result["questions"]


@pytest.mark.parametrize("scope", ["global", "project:foreign"])
def test_participation_is_not_inferred_from_company_or_other_practice(case, scope):
    participation_case(case)
    source, _ = add_source(case, "- Ragione sociale: Nuova S.r.l.\n"
                                "- Forma giuridica: S.r.l.\n"
                                "- Modalità di partecipazione: singola.")
    source["scope"] = scope
    result = validate(case)
    assert result["participation"]["status"] == "UNKNOWN" and len(result["questions"]) == 1


def test_participation_user_precedes_conflicting_source_without_guessing(case):
    participation_case(case)
    source, _ = add_source(case, "- Ragione sociale: Nuova S.r.l.\n"
                                "- Modalità di partecipazione: singola.")
    source["scope"] = "project:alpha"
    case[3].append({"id": "mode", "condition": "Forma di partecipazione", "value": "RTI",
                    "text": "Decisione verificata: RTI"})
    result = validate(case)
    assert result["participation"]["value"] == "RTI" and not result["questions"]


def test_conflicting_source_address_relations_remain_unknown(case):
    candidate, field = case[1]["candidates"][0], case[-1]["fields"][0]
    candidate.update(label_hint="Sede operativa (se diversa dalla sede legale)",
                     local_form="Sede operativa (se diversa dalla sede legale)")
    field.update(status="UNKNOWN", value=None, evidence=[], form_quote=candidate["local_form"])
    for source_id, relation in [("same", "coincidente"), ("different", "non coincidente")]:
        add_source(case, "- Ragione sociale: Nuova S.r.l.\n"
                        f"- Rapporto tra le sedi: sede operativa {relation} con la sede legale.",
                   source_id)
    result = validate(case)["fields"][0]
    assert result["status"] == "UNKNOWN" and any("conflitto" in e for e in result["errors"])


def test_verified_positive_local_user_condition_is_not_overruled_by_source_or_llm(case):
    candidate, field = case[1]["candidates"][0], case[-1]["fields"][0]
    candidate.update(label_hint="Sede operativa (se diversa dalla sede legale)",
                     local_form="Sede operativa (se diversa dalla sede legale)")
    _, evidence = add_source(case, "- Ragione sociale: Nuova S.r.l.\n"
                            "- Rapporto tra le sedi: sede operativa coincidente "
                            "con la sede legale.\n"
                            "- Sede operativa: Via Roma 8.")
    field.update(form_quote=candidate["local_form"], status="NOT_APPLICABLE", value=None,
                 evidence=[])
    case[3].append({"id": "local-condition", "candidate_id": candidate["id"],
                    "condition": "se diversa dalla sede legale", "applies": True,
                    "text": "Decisione verificata: la sede operativa è diversa."})
    first = validate(case)["fields"][0]
    assert first["status"] == "UNKNOWN"  # No value, but true USER condition is preserved.
    assert first["condition_rule"] == "recorded_user_local_condition"
    assert first["preserved_user_decisions"][0]["applies"] is True
    field.update(status="PROPOSED", value="Via Roma 8", evidence=[evidence(3)])
    assert validate(case)["writes"][candidate["id"]] == "Via Roma 8"


def test_user_fragment_cannot_exclude_a_whole_conditional_section(case):
    section = case[-1]["sections"][1]
    section["condition_quote"] = "da compilare in caso di SOCIETÀ DI INGEGNERIA con rappresentante"
    case[1]["sections"][1]["condition_quote"] = section["condition_quote"]
    case[3].append({"id": "partial", "section_id": section["section_id"],
                    "condition": "rappresentante", "applies": False,
                    "text": "Non abbiamo indicato il rappresentante."})
    assert validate(case)["sections"][1]["applicability"] == "UNKNOWN"


def test_user_local_fragment_is_not_a_verified_whole_condition(case):
    field, candidate = case[-1]["fields"][0], case[1]["candidates"][0]
    candidate.update(label_hint="Sede operativa (se diversa dalla sede legale)",
                     local_form="Sede operativa (se diversa dalla sede legale)")
    field.update(status="NOT_APPLICABLE", value=None, form_quote=candidate["local_form"])
    case[3].append({"id": "partial", "candidate_id": candidate["id"], "condition": "se",
                    "applies": False, "text": "Condizione non verificata completamente"})
    assert validate(case)["fields"][0]["status"] == "UNKNOWN"


@pytest.mark.parametrize("scope", ["global", "project:foreign", "project:alpha"])
def test_execution_as_consortium_member_requires_current_project_proof(case, scope):
    condition = "da compilare in caso di CONSORZIATA ESECUTRICE di Consorzio Stabile"
    case[1]["sections"][2]["condition_quote"] = condition
    case[-1]["sections"][2].update(condition_quote=condition, applicability="NOT_APPLICABLE")
    source, _ = add_source(case, "- Ragione sociale: Nuova S.r.l.\n"
                                "- Consorziata esecutrice di consorzio stabile: No.")
    source["scope"] = scope
    checked = validate(case)["sections"][2]
    assert (checked["applicability"] == "NOT_APPLICABLE") == (scope == "project:alpha")


def test_user_state_survives_incomplete_model_but_does_not_enable_export(case):
    field = case[-1]["fields"][0]
    case[3].append({"id": "saved", "candidate_id": field["candidate_id"],
                    "value": "Nuova S.r.l.", "text": "Nuova S.r.l."})
    case[-1]["fields"].pop(0)
    result = validate(case)
    assert result["fields"][0]["value"] == "Nuova S.r.l."
    assert result["fields"][0]["preserved_user_decisions"]
    assert not result["coverage_complete"] and not result["writes"]


def test_different_person_names_and_ids_keep_same_role_safety():
    data = role_case()
    replacement = {"Lina Blu": "Nora Viola", "Ugo Gialli": "Leo Bruni",
                   "Impresa Beta": "Ditta Delta"}
    for entity in data[-1]["entities"]:
        old_id = entity["id"]
        entity["id"] = f"unique-{old_id}"
        for field in data[-1]["fields"]:
            if field["subject_id"] == old_id:
                field["subject_id"] = entity["id"]
    for old, new in replacement.items():
        data[-1].update(json.loads(json.dumps(data[-1]).replace(old, new)))
        data[2][0].update(json.loads(json.dumps(data[2][0]).replace(old, new)))
    assert len(validate(data)["writes"]) == 3
    field = data[-1]["fields"][0]
    other = data[-1]["fields"][1]
    field.update(subject_id=other["subject_id"], value=other["value"], evidence=other["evidence"])
    assert validate(data)["fields"][0]["status"] == "UNKNOWN"


def test_recorded_user_value_may_have_no_conditional_predicate(case):
    field = case[-1]["fields"][0]
    case[3].append({"id": "saved-name", "candidate_id": field["candidate_id"],
                    "value": "Nuova S.r.l.", "condition": None, "applies": None,
                    "text": "Nuova S.r.l."})
    field.update(status="UNKNOWN", value=None, evidence=[])
    assert validate(case)["writes"][field["candidate_id"]] == "Nuova S.r.l."


@pytest.mark.parametrize("value", ["non so", "Società a responsabilità limitata", "da decidere"])
def test_unknown_answer_or_legal_form_does_not_resolve_participation(case, value):
    participation_case(case)
    case[3].append({"id": "mode", "condition": "Modalità di partecipazione", "value": value,
                    "text": f"Risposta registrata: {value}"})
    result = validate(case)
    assert result["participation"]["status"] == "UNKNOWN"
    assert result["participation"]["preserved_user_decisions"][0]["value"] == value
    assert "nessuna scelta dedotta" in result["participation"]["errors"][0]
    assert len(result["questions"]) == 1


def test_question_for_unresolved_field_is_not_lost_because_another_section_is_excluded(case):
    sid = case[-1]["sections"][2]["section_id"]
    case[3].append({"id": "excluded", "section_id": sid, "condition": "STUDIO ASSOCIATO",
                    "applies": False, "text": "Non studio associato"})
    cid = case[1]["candidates"][2]["id"]
    case[-1]["questions"] = [{"question": "Fornire il codice fiscale personale mancante.",
                              "candidate_ids": [cid], "section_ids": [sid]}]
    assert validate(case)["questions"][0]["candidate_ids"] == [cid]


def test_signatory_label_is_not_overridden_by_an_inherited_person_role():
    data = role_case()
    candidate, field = data[1]["candidates"][1], data[-1]["fields"][1]
    candidate.update(label_hint="Il/la sottoscritto/a", local_form="Il/la sottoscritto/a")
    field.update(form_quote=candidate["local_form"])
    result = validate(data)
    assert result["fields"][1]["required_role"] == "signatory"
    assert result["fields"][1]["status"] == "UNKNOWN"
