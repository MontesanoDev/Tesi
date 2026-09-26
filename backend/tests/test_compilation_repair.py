import asyncio
import json
from io import BytesIO

import pytest
from docx import Document

from app import document_compilation as compilation
from app.docx_templates import fill_docx, inspect_docx, text_of
from app.generation import GenerationError

COMPANY = "Impresa Esempio S.r.l."
LEGAL_FORM = "Societa a responsabilita limitata"


@pytest.fixture
def anyio_backend():
    return "asyncio"


def layout_for(kind="table"):
    doc = Document()
    doc.add_paragraph("Modulo dimostrativo indipendente dal bando Catanzaro")
    if kind == "paragraph":
        for label, marker in (("Denominazione", "nome"), ("Forma", "forma"), ("Firma", "firma")):
            doc.add_paragraph(f"{label}: {{{{{marker}}}}}")
    else:
        if kind == "shifted":
            extra = doc.add_table(rows=1, cols=1)
            extra.cell(0, 0).text = "Tabella descrittiva non compilabile"
        table = doc.add_table(rows=3, cols=2)
        for row, label in zip(table.rows, ("Denominazione", "Forma", "Firma"), strict=True):
            row.cells[0].text = label
    output = BytesIO()
    doc.save(output)
    return inspect_docx(output.getvalue())


def context():
    content = f"Denominazione: {COMPANY}. Forma giuridica: {LEGAL_FORM}."
    return compilation.CompilationSources([{
        "id": "company:77", "document_id": 77, "source_name": "impresa.txt",
        "chunk_index": 0, "content": content, "scope": "company", "source_kind": "source",
    }], 1, len(content))


def proposal(target, *, label="Denominazione", value=COMPANY, **changes):
    return {
        "cell_id": target, "label": label, "entity": "company", "kind": "data",
        "status": "proposed", "value": value,
        "evidence": [{"source_id": "company:77", "quote": context().selected[0]["content"]}],
        "reason": "Dato esplicito nella fonte", **changes,
    }


def response(*fields):
    return json.dumps({"fields": fields, "warnings": []})


def targets(layout):
    return (*layout.cells, *layout.slots)


@pytest.mark.parametrize("status", ["missing", "needs_review", "not_applicable"])
def test_omitted_empty_evidence_is_allowed_only_for_unwritten_fields(status):
    layout = layout_for()
    raw = proposal(targets(layout)[0], status=status, value=None)
    raw.pop("evidence")
    field = compilation.validate_proposals(response(raw), layout, context())["fields"][0]
    assert field["status"] == status
    assert field["written_value"] is None
    assert field["evidence"] == []
    assert field["repairable"] is False


@pytest.mark.parametrize("status", ["missing", "needs_review", "not_applicable"])
@pytest.mark.parametrize("label", ["", "   "])
def test_unwritten_unlabelled_cell_uses_its_coordinate_without_inventing_meaning(status, label):
    layout = layout_for()
    target = targets(layout)[0]
    raw = proposal(target, label=label, status=status, value=None, evidence=[])
    field = compilation.validate_proposals(response(raw), layout, context())["fields"][0]
    assert field["label"] == target
    assert field["written_value"] is None
    assert raw["label"] == label


@pytest.mark.parametrize("case", ["proposed", "value", "null", "missing", "unknown_id"])
def test_unwritten_label_fallback_does_not_relax_write_or_coordinate_validation(case):
    layout = layout_for()
    raw = proposal(targets(layout)[0], label="", status="needs_review", value=None)
    if case == "proposed":
        raw.update(status="proposed", value=COMPANY)
    elif case == "value":
        raw["value"] = COMPANY
    elif case == "null":
        raw["label"] = None
    elif case == "missing":
        raw.pop("label")
    else:
        raw["cell_id"] = "t999.r0.c0"
    with pytest.raises(GenerationError):
        compilation.validate_proposals(response(raw), layout, context())


@pytest.mark.parametrize("case", [
    "proposed", "value", "null", "string", "object", "missing_status", "missing_value",
])
def test_empty_evidence_default_does_not_relax_malformed_or_writable_proposals(case):
    layout = layout_for()
    raw = proposal(targets(layout)[0], status="needs_review", value=None)
    raw.pop("evidence")
    if case == "proposed":
        raw.update(status="proposed", value=COMPANY)
    elif case == "value":
        raw["value"] = COMPANY
    elif case in ("null", "string", "object"):
        raw["evidence"] = {"null": None, "string": "", "object": {}}[case]
    else:
        raw.pop(case.removeprefix("missing_"))
    with pytest.raises(GenerationError):
        compilation.validate_proposals(response(raw), layout, context())


def test_user_input_is_citable_without_becoming_a_document_or_affecting_coverage():
    layout, sources = layout_for(), context()
    before = sources.coverage()
    entered = "Sottoscrittore: Giulia Bianchi. Partecipazione singola."
    payload = json.loads(compilation.build_prompt(layout, sources, "Altro progetto", entered))
    user = next(s for s in payload["sources"] if s["id"] == "user:instructions")
    assert user["origin"] == "user"
    assert user["scope"] == "user"
    assert user["source_kind"] == "user_instructions"
    assert user["document_id"] is None
    assert user["chunk_index"] is None
    assert payload["source_coverage"] == sources.coverage() == before
    assert "origin" not in sources.selected[0]
    raw = proposal(targets(layout)[0], entity="person", value="Giulia Bianchi", evidence=[{
        "source_id": "user:instructions", "quote": "Sottoscrittore: Giulia Bianchi.",
    }])
    field = compilation.validate_proposals(
        response(raw), layout, sources, instructions=entered,
    )["fields"][0]
    assert field["written_value"] == "Giulia Bianchi"
    assert field["evidence"][0]["origin"] == "user"
    assert field["evidence"][0]["fragment"] is None
    assert field["evidence"][0]["document_id"] is None
    absent = compilation.validate_proposals(response(raw), layout, sources)["fields"][0]
    assert absent["written_value"] is None
    assert "unknown_source" in absent["validation_codes"]
    assert json.loads(compilation.build_prompt(layout, sources, "Altro", "  "))[
        "user_instructions_source_id"
    ] is None


@pytest.mark.parametrize("kind,origin", [
    ("source", "document"), ("project_facts", "user"), ("call_facts", "extracted"),
])
def test_persisted_source_origin_is_explicit(kind, origin):
    sources = context()
    sources.selected[0]["source_kind"] = kind
    assert compilation.source_catalog(sources)[0]["origin"] == origin


@pytest.mark.parametrize("bad_evidence", [
    [{"source_id": "project:999", "quote": COMPANY}],
    [{"source_id": "company:77", "quote": "Sottoscrittore: Giulia Bianchi"}],
])
def test_false_citation_blocks_only_its_field_even_with_another_valid_citation(bad_evidence):
    layout = layout_for()
    first, second, _ = targets(layout)
    good = proposal(first)
    bad = proposal(second, evidence=[*good["evidence"], *bad_evidence])
    report = compilation.validate_proposals(response(good, bad), layout, context())
    assert report["fields"][0]["written_value"] == COMPANY
    blocked = report["fields"][1]
    assert blocked["status"] == "needs_review"
    assert blocked["written_value"] is None
    assert blocked["repairable"] is True
    assert blocked["rejected_evidence"][0]["quote"] == bad_evidence[0]["quote"]
    assert len(blocked["evidence"]) == 1


@pytest.mark.parametrize("kind", ["table", "shifted", "paragraph"])
@pytest.mark.parametrize("status", ["proposed", "missing", "needs_review", "not_applicable"])
def test_out_of_batch_proposals_are_audited_but_never_classified_or_written(kind, status):
    layout = layout_for(kind)
    first, second, _ = targets(layout)
    extra = proposal(
        second, status=status, value=COMPANY if status == "proposed" else None,
    )
    report = compilation.validate_proposals(
        response(extra, proposal(first)), layout, context(), allowed_ids={first},
    )
    assert [field["cell_id"] for field in report["fields"]] == [first]
    assert report["fields"][0]["written_value"] == COMPANY
    assert second in report["unclassified_fields"]
    if kind != "paragraph":
        assert second in report["unclassified_cells"]
    assert report["rejected_proposals"] == [{
        "proposal": extra,
        "validation_code": "outside_batch",
        "message": "Proposta scartata: campo esterno al gruppo richiesto",
        "target_ids": [first],
    }]
    all_rejected = compilation.validate_proposals(
        response(extra), layout, context(), allowed_ids=set(),
    )
    assert all_rejected["fields"] == []
    assert set(all_rejected["unclassified_fields"]) == set(targets(layout))


@pytest.mark.parametrize("failure", ["unknown", "unwritable", "duplicate", "bad_schema"])
def test_scope_filter_does_not_relax_structural_validation(failure):
    layout = layout_for()
    first, second, _ = targets(layout)
    extras = [proposal(second)]
    if failure == "unknown":
        extras = [proposal("t999.r0.c1")]
    elif failure == "unwritable":
        extras = [proposal("t0.r0.c0")]
    elif failure == "duplicate":
        extras.append(proposal(second))
    else:
        extras = [proposal(second, value={"unexpected": "object"})]
    with pytest.raises(GenerationError):
        compilation.validate_proposals(
            response(proposal(first), *extras), layout, context(), allowed_ids={first},
        )


def test_scope_filter_keeps_evidence_and_signature_checks_for_assigned_fields():
    layout = layout_for()
    first, second, signature = targets(layout)
    report = compilation.validate_proposals(
        response(proposal(first, evidence=[]), proposal(second), proposal(signature)),
        layout, context(), allowed_ids={first, signature},
    )
    assert len(report["fields"]) == 2
    assert all(field["written_value"] is None for field in report["fields"])
    assert "value_not_in_quote" in report["fields"][0]["validation_codes"]
    assert "protected_field" in report["fields"][1]["validation_codes"]
    assert len(report["rejected_proposals"]) == 1


@pytest.mark.anyio
@pytest.mark.parametrize("kind", ["table", "shifted", "paragraph"])
@pytest.mark.parametrize("error", ["accents", "wrong_quote", "unknown_source"])
async def test_one_targeted_repair_is_independent_of_project_and_field_position(
    monkeypatch, kind, error,
):
    layout = layout_for(kind)
    first, second, _ = targets(layout)
    good = proposal(first)
    bad = proposal(second, label="Forma", value=LEGAL_FORM)
    if error == "accents":
        bad["value"] = "Societ\u00e0 a responsabilit\u00e0 limitata"
    elif error == "wrong_quote":
        bad["evidence"][0]["quote"] = "Forma giuridica non presente"
    else:
        bad["evidence"][0]["source_id"] = "company:999"
    calls = []

    async def model(prompt):
        body = json.loads(prompt)
        calls.append(body)
        if len(calls) == 1:
            return response(good, bad), "first-model", 10
        assert body["target_ids"] == [second]
        assert body["correction_request"][0]["cell_id"] == second
        assert body["correction_request"][0]["validation_codes"]
        return response(proposal(second, label="Forma", value=LEGAL_FORM)), "repair-model", 12

    monkeypatch.setattr(compilation, "request_field_proposals", model)
    report, name, tokens, execution = await compilation.compile_field_batches(
        layout, context(), "Progetto indipendente", "Solo i dati disponibili",
    )
    assert len(calls) == 2
    assert calls[0]["sources"] == calls[1]["sources"]
    assert execution["repaired_fields"] == execution["repair_attempted_fields"] == 1
    assert execution["repair_requests"] == 1
    assert tokens == 22
    assert name == "first-model, repair-model"
    assert "repair" not in report["fields"][0]
    repaired = report["fields"][1]
    assert repaired["written_value"] == LEGAL_FORM
    assert repaired["repair"]["initial_proposal"]["value"] == bad["value"]
    assert repaired["repair"]["status"] == "corrected"
    assert repaired["validation_notes"] == []
    data = fill_docx(layout, {f["cell_id"]: f["written_value"] for f in report["fields"]})
    output = Document(BytesIO(data))
    assert COMPANY in text_of(output.element.body)
    assert LEGAL_FORM in text_of(output.element.body)


@pytest.mark.anyio
@pytest.mark.parametrize("failure", [
    "same_error", "omitted", "missing", "identity", "timeout", "length",
])
async def test_failed_repair_keeps_good_fields_and_never_runs_a_second_repair(monkeypatch, failure):
    layout = layout_for()
    first, second, _ = targets(layout)
    good = proposal(first)
    bad = proposal(second, evidence=[{"source_id": "missing", "quote": COMPANY}])
    calls = 0

    async def model(_prompt):
        nonlocal calls
        calls += 1
        if calls == 1:
            return response(good, bad), "test", 10
        if failure == "timeout":
            raise GenerationError("Timeout del provider")
        if failure == "length":
            raise compilation.TruncatedCompilationError(8)
        if failure == "omitted":
            return response(), "test", 8
        if failure == "missing":
            return response(proposal(second, status="missing", value=None, evidence=[])), "test", 8
        if failure == "identity":
            return response(proposal(second, entity="person")), "test", 8
        return response(bad), "test", 8

    monkeypatch.setattr(compilation, "request_field_proposals", model)
    report, _, tokens, execution = await compilation.compile_field_batches(
        layout, context(), "Altro progetto", "",
    )
    assert calls == 2
    assert execution["repair_requests"] == 1
    assert execution["repaired_fields"] == 0
    assert tokens == (None if failure == "timeout" else 18)
    assert report["fields"][0]["written_value"] == COMPANY
    blocked = report["fields"][1]
    assert blocked["written_value"] is None
    assert blocked["status"] == "needs_review"  # Not disguised as a missing source datum.
    assert blocked["repair"]["attempted"] is True
    assert blocked["repair"]["status"] == "unresolved"
    assert any("proposte bloccate" in warning for warning in report["warnings"])


@pytest.mark.anyio
@pytest.mark.parametrize("kind", ["table", "paragraph"])
@pytest.mark.parametrize("include_correction", [False, True])
async def test_out_of_batch_repair_cannot_modify_previously_accepted_fields(
    monkeypatch, kind, include_correction,
):
    layout = layout_for(kind)
    first, second, _ = targets(layout)
    calls = []

    async def model(prompt):
        body = json.loads(prompt)
        calls.append(body)
        if len(calls) == 1:
            return response(proposal(first), proposal(second, evidence=[])), "test", 10
        assert body["target_ids"] == [second]
        extra = proposal(first, value="Impresa")
        own = [proposal(second)] if include_correction else []
        return response(extra, *own), "test", 12

    monkeypatch.setattr(compilation, "request_field_proposals", model)
    report, _, tokens, execution = await compilation.compile_field_batches(
        layout, context(), "Altro progetto", "",
    )
    assert len(calls) == 2
    assert tokens == 22
    assert execution["out_of_batch_proposals"] == 1
    assert execution["repair_requests"] == 1
    assert execution["repaired_fields"] == int(include_correction)
    assert report["fields"][0]["written_value"] == COMPANY
    assert "repair" not in report["fields"][0]
    repaired = report["fields"][1]
    assert repaired["written_value"] == (COMPANY if include_correction else None)
    assert repaired["repair"]["status"] == ("corrected" if include_correction else "unresolved")
    assert report["rejected_proposals"][0]["proposal"]["cell_id"] == first
    assert report["rejected_proposals"][0]["phase"] == "repair"
    assert report["rejected_proposals"][0]["request"] == 2
    assert report["rejected_proposals"][0]["target_ids"] == [second]


@pytest.mark.anyio
@pytest.mark.parametrize("failure", ["json", "unknown", "duplicate"])
async def test_structural_repair_errors_still_abort_the_document(monkeypatch, failure):
    layout = layout_for()
    first, second, _ = targets(layout)
    calls = 0

    async def model(_prompt):
        nonlocal calls
        calls += 1
        if calls == 1:
            return response(proposal(first), proposal(second, evidence=[])), "test", 10
        if failure == "json":
            return '{"fields":', "test", 2
        if failure == "unknown":
            return response(proposal("t999.r0.c1")), "test", 2
        return response(proposal(second), proposal(second)), "test", 2

    monkeypatch.setattr(compilation, "request_field_proposals", model)
    with pytest.raises(GenerationError):
        await compilation.compile_field_batches(layout, context(), "Altro", "")
    assert calls == 2


@pytest.mark.anyio
async def test_repair_budget_exhaustion_leaves_a_report_without_extra_calls(monkeypatch):
    layout = layout_for()
    first, second, _ = targets(layout)
    calls = []
    monkeypatch.setattr(compilation, "MAX_BATCH_REQUESTS", 2)
    monkeypatch.setattr(compilation, "FIELDS_PER_BATCH", 2)

    async def model(prompt):
        body = json.loads(prompt)
        calls.append(body)
        assert "correction_request" not in body
        if len(calls) == 1:
            return response(proposal(first), proposal(second, evidence=[])), "test", 9
        return response(), "test", 9

    monkeypatch.setattr(compilation, "request_field_proposals", model)
    report, _, tokens, execution = await compilation.compile_field_batches(
        layout, context(), "Altro", "",
    )
    assert len(calls) == 2
    assert tokens == 18
    assert execution["completed_batches"] == 2
    assert execution["repair_requests"] == 0
    assert execution["repair_skipped_fields"] == 1
    assert report["fields"][1]["repair"]["attempted"] is False
    assert "Limite" in report["fields"][1]["repair"]["message"]


@pytest.mark.anyio
async def test_global_timeout_also_applies_during_repair(monkeypatch):
    layout = layout_for()
    calls = 0
    monkeypatch.setattr(compilation, "COMPILATION_TIMEOUT_SECONDS", 0.02)

    async def model(_prompt):
        nonlocal calls
        calls += 1
        if calls == 1:
            return response(proposal(targets(layout)[0], evidence=[])), "test", 9
        await asyncio.sleep(1)
        return response(), "test", 9

    monkeypatch.setattr(compilation, "request_field_proposals", model)
    with pytest.raises(GenerationError, match="tempo massimo complessivo"):
        await compilation.compile_field_batches(layout, context(), "Altro", "")
    assert calls == 2


@pytest.mark.anyio
@pytest.mark.parametrize("kind", ["declaration", "choice", "signature", "disguised_signature"])
async def test_user_sources_and_repairs_cannot_enable_protected_fields(monkeypatch, kind):
    layout = layout_for()
    target = targets(layout)[2 if kind == "disguised_signature" else 0]
    calls = 0

    async def model(_prompt):
        nonlocal calls
        calls += 1
        return response(proposal(
            target, kind="data" if kind == "disguised_signature" else kind,
            value="Giulia Bianchi", entity="person", evidence=[{
                "source_id": "user:instructions", "quote": "Giulia Bianchi",
            }],
        )), "test", 4

    monkeypatch.setattr(compilation, "request_field_proposals", model)
    report, _, _, execution = await compilation.compile_field_batches(
        layout, context(), "Altro", "Giulia Bianchi",
    )
    assert calls == 1
    assert execution["repair_requests"] == 0
    assert report["fields"][0]["written_value"] is None
    assert report["fields"][0]["validation_codes"] == ["protected_field"]


@pytest.mark.parametrize("original,altered", [
    ("2026-09-17", "2026-09-18"), ("01234567890", "01234567891"), ("1250000", "1520000"),
])
def test_dates_identifiers_and_amounts_are_not_fuzzy_matched(original, altered):
    layout, sources = layout_for(), context()
    sources.selected[0]["content"] = original
    field = compilation.validate_proposals(response(proposal(
        targets(layout)[0], value=altered,
        evidence=[{"source_id": "company:77", "quote": original}],
    )), layout, sources)["fields"][0]
    assert field["written_value"] is None
    assert "value_not_in_quote" in field["validation_codes"]
