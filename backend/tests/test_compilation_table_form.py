"""Physical FORM edges and independent negative gates; no provider inference."""

import json
from copy import deepcopy
from io import BytesIO
from pathlib import Path

import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from app.compilation_catalog import build_compilation_catalog
from app.compilation_global import _form_bindings, build_global_input, validate_global_proposal
from app.compilation_global_dialogue import export_selection
from app.compilation_global_sessions import initial_state, review_fields
from app.compilation_table_form import relations
from app.docx_templates import fill_docx, inspect_docx

FIXTURE = Path(__file__).parent / "fixtures"


def layout_for(rows, *, edit=None, preceding="", columns=None):
    doc = Document()
    if preceding:
        doc.add_paragraph(preceding)
    table = doc.add_table(rows=len(rows), cols=columns or len(rows[0]))
    for ri, row in enumerate(rows):
        for ci, text in enumerate(row):
            table.cell(ri, ci).text = text
    if edit:
        edit(doc, table)
    stream = BytesIO()
    doc.save(stream)
    return inspect_docx(stream.getvalue(), require_candidates=False)


def catalog_for(layout):
    catalog = build_compilation_catalog(layout)
    catalog["_table_form_links"] = relations(layout, catalog)
    return catalog


def empty_proposal():
    return {"version": 12, "uncertain_paragraph_ids": [], "coverage": [], "choice_groups": {},
            "clarifications": {}, "facts": {}, "fields": {}, "bindings": {}, "conditions": {}}


def proposed(layout, cid, *, pid=None, entity="company", subject="Aurora",
             content="Identificativo di Aurora: VAL-73.", value="VAL-73", binding=None):
    catalog = catalog_for(layout)
    edge = catalog["_table_form_links"].get(cid, {})
    if pid is None:
        pid = edge["form_evidence"][0]["paragraph_id"]
    text = next(p["text"] for b in catalog["blocks"] for p in b["paragraphs"] if p["id"] == pid)
    raw = empty_proposal()
    raw["facts"]["f"] = {"entity": entity, "subject": subject, "aspect": "identity",
                          "property": "identificativo", "value": value,
                          "evidence": [{"chunk_id": -7, "quote": content}]}
    raw["fields"][cid] = {"label": "Etichetta MODEL", "kind": "data", "disposition": "value",
                          "conditions": [], "property": "identificativo", "fact_id": "f",
                          "form": {"paragraph_id": pid, "quote": text}}
    sources = [{"chunk_id": -7, "file_id": -3, "chunk_index": 0, "source_name": "profili.txt",
                "role": "source", "scope": "global", "project_id": None,
                "category": "company", "content": content}]
    if binding:
        role, source_role = binding
        role_p = next(p for b in catalog["blocks"] for p in b["paragraphs"] if p["text"] == role)
        raw["bindings"]["b"] = {"required_role": role, "source_role": source_role,
                                "form": {"paragraph_id": role_p["id"], "quote": role},
                                "evidence": [{"chunk_id": -7, "quote": content}]}
        raw["fields"][cid]["binding_id"] = "b"
    return raw, sources


def checked(layout, raw, sources):
    return validate_global_proposal(json.dumps(raw), layout, "alpha", sources, expected_version=12)


def field(audit, cid):
    return next(f for f in audit["fields"] if f["candidate_id"] == cid)


@pytest.mark.parametrize("rows,cid,label", [
    ([["Identificativo", ""]], "t0.r0.c1", "t0.r0.c0"),
    ([["", "Identificativo"]], "t0.r0.c0", "t0.r0.c1"),
    ([["Identificativo"], [""]], "t0.r1.c0", "t0.r0.c0"),
    ([["Identificativo", "Recapito"], ["", ""], ["", ""]], "t0.r2.c1", "t0.r0.c1"),
    ([["A", "", "B", ""]], "t0.r0.c3", "t0.r0.c2"),
])
def test_exact_label_edges_and_physical_writes(rows, cid, label):
    layout = layout_for(rows)
    catalog = catalog_for(layout)
    edge = catalog["_table_form_links"][cid]
    assert edge["result"] == "structural_match" and edge["label_cell_id"] == label
    assert edge["original_sha256"] == layout.sha256 and edge["role"] == "form"
    raw, sources = proposed(layout, cid)
    item = field(checked(layout, raw, sources), cid)
    assert item["validation"] == "GROUNDED_PROPOSAL", item
    assert item["semantic_review_required"] and item["written_value"] is None
    written = Document(BytesIO(fill_docx(layout, {cid: "VAL-73"})))
    _, ri, ci = cid.split(".")
    assert written.tables[0].cell(int(ri[1:]), int(ci[1:])).text == "VAL-73"


def test_horizontal_merges_use_grid_spans_not_physical_column_indices():
    def merge(doc, table):
        table.cell(0, 0).merge(table.cell(0, 2))
        table.cell(1, 0).merge(table.cell(1, 2))
    layout = layout_for([["A", "", "", "B"], ["", "", "", ""]], edit=merge)
    links = relations(layout, build_compilation_catalog(layout))
    assert links["t0.r1.c0"]["grid_interval"] == [0, 3]
    assert links["t0.r1.c0"]["label_cell_id"] == "t0.r0.c0"
    assert links["t0.r1.c1"]["grid_interval"] == [3, 4]
    assert links["t0.r1.c1"]["label_cell_id"] == "t0.r0.c1"


def test_vertical_merges_with_equal_extents_are_one_unambiguous_pair():
    def merge(doc, table):
        table.cell(0, 0).merge(table.cell(1, 0))
        table.cell(0, 1).merge(table.cell(1, 1))
    layout = layout_for([["A", ""], ["", ""]], edit=merge)
    edge = relations(layout, build_compilation_catalog(layout))["t0.r0.c1"]
    assert edge["label_cell_id"] == "t0.r0.c0" and edge["vertical_extent"] == [0, 1]


def test_label_is_all_fixed_paragraphs_in_one_cell():
    layout = layout_for([["Identificativo", ""]],
                        edit=lambda d, t: t.cell(0, 0).add_paragraph("dell'organizzazione"))
    edge = catalog_for(layout)["_table_form_links"]["t0.r0.c1"]
    assert [p["quote"] for p in edge["form_evidence"]] == ["Identificativo", "dell'organizzazione"]


@pytest.mark.parametrize("rows,cid", [
    ([["A", "", "B"]], "t0.r0.c1"),
    ([["A", "", ""]], "t0.r0.c1"),
    ([["", ""], ["", ""]], "t0.r1.c0"),
    ([["A", "B"], ["", "Interruzione"], ["", ""]], "t0.r2.c0"),
    ([["A", "B"], ["", ""], ["", "Soggetto fisso"]], "t0.r2.c0"),
])
def test_ambiguous_or_interrupted_layouts_have_no_proven_label(rows, cid):
    edge = catalog_for(layout_for(rows))["_table_form_links"][cid]
    assert edge["result"] == "not_demonstrated" and "form_evidence" not in edge


def test_partial_header_span_does_not_distribute_label_to_independent_values():
    layout = layout_for([["A", ""], ["", ""]],
                        edit=lambda d, t: t.cell(0, 0).merge(t.cell(0, 1)))
    assert all(e["result"] == "not_demonstrated" for e in
               catalog_for(layout)["_table_form_links"].values())


def test_one_vertical_label_for_several_independent_targets_is_ambiguous():
    layout = layout_for([["A", ""], ["", ""]],
                        edit=lambda d, t: t.cell(0, 0).merge(t.cell(1, 0)))
    links = catalog_for(layout)["_table_form_links"]
    assert links["t0.r0.c1"]["result"] == links["t0.r1.c1"]["result"] == "not_demonstrated"


@pytest.mark.parametrize("wrong", ["other_column", "other_table", "duplicate_label_other_table"])
def test_citations_cannot_cross_columns_tables_or_duplicate_sections(wrong):
    def extra(doc, table):
        doc.add_paragraph("Sezione diversa")
        other = doc.add_table(rows=2, cols=2)
        other.cell(0, 0).text = "A" if wrong == "duplicate_label_other_table" else "Altro"
        other.cell(0, 1).text = "B"
    layout = layout_for([["A", "B"], ["", ""]], edit=extra)
    cat = catalog_for(layout)
    target = "t0.r1.c0"
    if wrong == "other_column":
        pid = cat["_table_form_links"]["t0.r1.c1"]["form_evidence"][0]["paragraph_id"]
    else:
        pid = cat["_table_form_links"]["t1.r1.c0"]["form_evidence"][0]["paragraph_id"]
    raw, sources = proposed(layout, target, pid=pid)
    assert field(checked(layout, raw, sources), target)["validation"] == "REJECTED"


def test_nested_table_labels_never_escape_their_cell_but_own_relations_work():
    def nested(doc, table):
        inner = table.cell(0, 0).add_table(rows=2, cols=1)
        inner.cell(0, 0).text = "A"
    layout = layout_for([["", ""]], edit=nested)
    links = catalog_for(layout)["_table_form_links"]
    assert links["t0.r0.c1"]["result"] == "not_demonstrated"
    assert links["t1.r1.c0"]["label_cell_id"] == "t1.r0.c0"


@pytest.mark.parametrize("entity,binding,grounded", [
    ("company", None, False), ("person", None, False),
    ("person", ("Direttore tecnico", "Socio"), False),
    ("person", ("Direttore tecnico", "Direttore tecnico"), True),
])
def test_source_type_and_actual_role_remain_independent(entity, binding, grounded):
    layout = layout_for([["Nome e cognome", "Codice"], ["", ""]], preceding="Direttore tecnico")
    raw, sources = proposed(layout, "t0.r1.c0", entity=entity, subject="Ada Neri", value="Ada Neri",
                            content="Direttore tecnico: Ada Neri; Socio: Ada Neri.",
                            binding=binding)
    item = field(checked(layout, raw, sources), "t0.r1.c0")
    assert (item["validation"] == "GROUNDED_PROPOSAL") is grounded


def test_different_people_cannot_supply_columns_of_the_same_new_record():
    layout = layout_for([["Nome e cognome", "Codice"], ["", ""]], preceding="Direttore tecnico")
    raw, sources = proposed(layout, "t0.r1.c0", entity="person", subject="Ada Neri",
        value="Ada Neri",
        content="Direttore tecnico: Ada Neri; Codice: NRAX1.",
        binding=("Direttore tecnico", "Direttore tecnico"))
    raw["facts"]["g"] = {**raw["facts"]["f"], "subject": "Eva Blu", "value": "Eva Blu"}
    raw["facts"]["g"]["evidence"] = [{"chunk_id": -8, "quote": "Direttore tecnico: Eva Blu."}]
    sources.append({**sources[0], "chunk_id": -8, "content": "Direttore tecnico: Eva Blu."})
    edge = catalog_for(layout)["_table_form_links"]["t0.r1.c1"]
    raw["fields"]["t0.r1.c1"] = {**raw["fields"]["t0.r1.c0"], "fact_id": "g",
                                 "form": {k: edge["form_evidence"][0][k]
                                          for k in ["paragraph_id", "quote"]}}
    result = checked(layout, raw, sources)
    assert all(field(result, cid)["validation"] == "REJECTED" for cid in raw["fields"])


@pytest.mark.parametrize("kind", ["previous", "header"])
@pytest.mark.parametrize("result", ["true", "false", "unknown"])
def test_new_form_edges_preserve_independent_applicability_and_export_veto(kind, result):
    label = "Sede operativa (se diversa dalla sede legale)"
    rows = [[label], [""]] if kind == "previous" else [[label, "Recapito"], ["", ""], ["", ""]]
    cid = "t0.r1.c0" if kind == "previous" else "t0.r2.c0"
    layout = layout_for(rows)
    content = "Sede operativa di Aurora: Via Uno 1."
    if result == "false":
        content += "\nSede legale di Aurora: Via Uno 1."
    if result == "true":
        content += "\nSede operativa di Aurora: diversa dalla sede legale."
    raw, sources = proposed(layout, cid, content=content, value="Via Uno 1")
    audit = checked(layout, raw, sources)
    assert audit["form_applicability"][cid]["result"] == result
    data = {**initial_state(layout), "audit": audit, "catalog": build_global_input(
        layout, "alpha", sources)["form"],
        "dialogue": {"phase": "READY", "reviewed_candidate_ids": [cid]}}
    data["fields"] = review_fields(data)
    ids, _, users = export_selection({"global_compilation": data})
    assert (cid in ids) is (result == "true") and not users


def test_thirteen_original_barriers_now_have_exact_destination_specific_proof():
    baseline = json.loads((FIXTURE / "table-form-v1/catanzaro.json").read_text())
    layout = inspect_docx((FIXTURE / "form-applicability-v1/original.docx").read_bytes())
    catalog = build_compilation_catalog(layout)
    paragraphs, old = _form_bindings(catalog)
    catalog["_table_form_links"] = relations(layout, catalog)
    _, new = _form_bindings(catalog)
    assert len(baseline["blocked"]) == 13
    for barrier in baseline["blocked"]:
        cid = barrier["candidate_id"]
        assert not any(paragraphs[p].strip() for p in old[cid])
        edge = catalog["_table_form_links"][cid]
        assert edge["result"] == "structural_match"
        expected = {p["paragraph_id"] for p in edge["form_evidence"]}
        assert {p for p in new[cid] if paragraphs[p].strip()} == expected
        assert len(expected) == 1
    assert catalog["_table_form_links"]["t25.r2.c0"]["form_evidence"][0]["paragraph_id"] == \
        "paragraph:439"


def test_original_truthful_cciaa_counterexample_can_be_grounded_but_not_auto_authorized():
    inp = json.loads((FIXTURE / "form-applicability-v1/input.json").read_text())
    layout = inspect_docx((FIXTURE / "form-applicability-v1/original.docx").read_bytes())
    raw = empty_proposal()
    raw["facts"]["f"] = {"subject": "Mapi Ingegneria S.r.l.", "entity": "company",
        "aspect": "identity", "property": "cciaa", "value": "Bari", "evidence": [{
            "chunk_id": -52,
            "quote": "- **Camera di Commercio competente nello scenario:** Bari."}]}
    raw["fields"]["t25.r2.c0"] = {"label": "Camera di Commercio", "kind": "data",
        "disposition": "value", "form": {"paragraph_id": "paragraph:439",
            "quote": "iscrizione alla CCIAA di: (per gli stranieri registri equipollenti)"},
        "conditions": [], "property": "cciaa", "fact_id": "f"}
    audit = checked(layout, raw, inp["sources"])
    item = field(audit, "t25.r2.c0")
    assert item["validation"] == "GROUNDED_PROPOSAL" and not item["validation_errors"]
    data = {**initial_state(layout), "audit": audit, "catalog": inp["form"],
            "dialogue": {"reviewed_candidate_ids": []}}
    data["fields"] = review_fields(data)
    assert export_selection({"global_compilation": data})[0] == []
    assert "iscrizione alla CCIAA" in data["fields"][0]["form_context"]
    assert raw["fields"]["t25.r2.c0"]["conditions"] == []


def test_validation_does_not_mutate_model_input_or_facts():
    layout = layout_for([["A", "B"], ["", ""]])
    raw, sources = proposed(layout, "t0.r1.c0")
    before, frozen = build_global_input(layout, "alpha", sources), deepcopy(raw)
    checked(layout, raw, sources)
    assert raw == frozen and build_global_input(layout, "alpha", sources) == before
    assert "_table_form_links" not in before["form"]


@pytest.mark.parametrize("placeholder", ["________", "........", " "])
def test_parser_owned_whole_cell_placeholders_keep_their_label(placeholder):
    layout = layout_for([["Denominazione sociale", placeholder]])
    raw, sources = proposed(layout, "t0.r0.c1", value="Aurora",
                            content="Denominazione sociale: Aurora.")
    item = field(checked(layout, raw, sources), "t0.r0.c1")
    assert item["validation"] == "GROUNDED_PROPOSAL"


def test_unknown_grid_offsets_and_unpaired_vertical_merge_have_no_new_edge():
    def edit(doc, table):
        bad = OxmlElement("w:gridBefore")
        bad.set(qn("w:val"), "unknown")
        table.rows[1]._tr.get_or_add_trPr().append(bad)
        table.cell(0, 0)._tc.get_or_add_tcPr().append(OxmlElement("w:vMerge"))
    layout = layout_for([["A", "B"], ["", ""]], edit=edit)
    assert not any(e["result"] == "structural_match" for e in
                   catalog_for(layout)["_table_form_links"].values())


def test_narrower_property_edge_never_removes_a_preexisting_condition_scope():
    def header(doc, table):
        table.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
    layout = layout_for([["Condizione del ramo", "Identificativo"], ["", ""]], edit=header)
    raw, sources = proposed(layout, "t0.r1.c1")
    pid = catalog_for(layout)["_table_form_links"]["t0.r1.c0"]["form_evidence"][0]["paragraph_id"]
    raw["conditions"]["c"] = {"form": {"paragraph_id": pid, "quote": "Condizione del ramo"},
        "aspect": "other", "operator": "semantic", "fact_ids": [], "reason": "Da decidere",
        "scope": []}
    raw["fields"]["t0.r1.c1"].update(disposition="conditional", conditions=["c"],
                                     reason="Ramo ancora da confermare")
    audit = checked(layout, raw, sources)
    assert set(audit["conditions"]["c"]["applies_to"]) == {"t0.r1.c0", "t0.r1.c1"}
    assert field(audit, "t0.r1.c1")["conditions"] == ["c"]
    assert field(audit, "t0.r1.c1")["validation"] == "REVIEW_REQUIRED"
    assert field(audit, "t0.r1.c1")["condition_checks"]["c"]["result"] == "unknown"


@pytest.mark.parametrize("kind,label", [("declaration", "Dichiarazione"),
                                        ("data", "Autorizzo le comunicazioni alla PEC")])
def test_new_table_proof_does_not_authorize_a_declaration(kind, label):
    layout = layout_for([[label], [""]])
    raw, sources = proposed(layout, "t0.r1.c0")
    raw["fields"]["t0.r1.c0"]["kind"] = kind
    audit = checked(layout, raw, sources)
    item = field(audit, "t0.r1.c0")
    assert audit["facts"]["f"]["validation"] == "GROUNDED_PROPOSAL"
    assert item["form_table_link"]["result"] == "structural_match"
    assert item["form_authorization_check"]["result"] == "unknown"
    assert item["validation"] == "REJECTED" and item.get("value") is None
