"""Literal FORM branch, registered participation and named SOURCE role stay separate."""

import copy
import json

import pytest
from docx import Document
from test_compilation_catalog import encoded, native_checkbox

from app.compilation_global import build_global_input, validate_global_proposal


def case():
    doc = Document()
    doc.add_paragraph("IN QUALITÀ DI")
    p = doc.add_paragraph()
    native_checkbox(p)
    p.add_run(" Professionista singolo")
    doc.add_paragraph("Recapito: ________")
    p = doc.add_paragraph()
    native_checkbox(p)
    p.add_run(" Legale rappresentante di società di ingegneria")
    doc.add_paragraph("Denominazione: ________")
    doc.add_paragraph(
        "la quale ha come Responsabile Tecnico (per le società di ingegneria) "
        "o Responsabile Progettazione (per le società di professionisti):"
    )
    doc.add_paragraph("Nome e cognome: ________")
    doc.add_paragraph("C.F.: ________")
    p = doc.add_paragraph()
    native_checkbox(p)
    p.add_run(" Mandatario di raggruppamento")
    doc.add_paragraph("CHIEDE")
    from app.docx_templates import inspect_docx

    layout = inspect_docx(encoded(doc))
    sources = [
        dict(
            chunk_id=-1,
            file_id=-1,
            chunk_index=0,
            source_name="azienda.md",
            role="source",
            scope="global",
            project_id=None,
            category="company",
            content="Denominazione: Vega S.r.l.\nTipologia: società di ingegneria",
        ),
        dict(
            chunk_id=-2,
            file_id=-2,
            chunk_index=0,
            source_name="persone.md",
            role="source",
            scope="global",
            project_id=None,
            category="company",
            content="- **Nora Serra:** Responsabile Tecnico di Vega S.r.l.; C.F. SRRNRO85A41F205D.",
        ),
    ]
    catalog = build_global_input(layout, "alpha", sources)["form"]
    ps = {p["id"]: p for b in catalog["blocks"] for p in b["paragraphs"]}
    form = {"paragraph_id": "paragraph:5", "quote": ps["paragraph:5"]["text"]}
    payload = {
        "version": 12,
        "coverage": [],
        "uncertain_paragraph_ids": [],
        "choice_groups": {},
        "clarifications": {},
        "conditions": {
            "company": {
                "form": {"paragraph_id": "paragraph:3", "quote": ps["paragraph:3"]["text"]},
                "aspect": "participation",
                "operator": "semantic",
                "fact_ids": [],
                "reason": "Qualità della domanda",
                "scope": [
                    {"kind": "candidate", "id": c} for c in ["p3.cb0", "p4.s0", "p6.s0", "p7.s0"]
                ],
            },
            "technical": {
                "form": form,
                "aspect": "professional",
                "operator": "semantic",
                "fact_ids": ["type", "role"],
                "reason": "Responsabile aziendale",
                "scope": [{"kind": "candidate", "id": c} for c in ["p6.s0", "p7.s0"]],
            },
        },
        "bindings": {
            "person": {
                "form": form,
                "required_role": "Responsabile Tecnico",
                "source_role": "Responsabile Tecnico",
                "evidence": [{"chunk_id": -2, "quote": sources[1]["content"]}],
            }
        },
        "facts": {},
        "fields": {},
    }
    for fid, entity, aspect, subject, prop, value, chunk, quote in [
        (
            "company",
            "company",
            "legal_type",
            "Vega S.r.l.",
            "denominazione",
            "Vega S.r.l.",
            -1,
            "Denominazione: Vega S.r.l.",
        ),
        (
            "type",
            "company",
            "legal_type",
            "Vega S.r.l.",
            "tipologia",
            "società di ingegneria",
            -1,
            "Tipologia: società di ingegneria",
        ),
        (
            "role",
            "person",
            "professional",
            "Nora Serra",
            "ruolo",
            "Responsabile Tecnico",
            -2,
            sources[1]["content"],
        ),
        (
            "name",
            "person",
            "identity",
            "Nora Serra",
            "nome",
            "Nora Serra",
            -2,
            sources[1]["content"],
        ),
        (
            "cf",
            "person",
            "identity",
            "Nora Serra",
            "codice_fiscale",
            "SRRNRO85A41F205D",
            -2,
            sources[1]["content"],
        ),
    ]:
        payload["facts"][fid] = dict(
            entity=entity,
            aspect=aspect,
            subject=subject,
            property=prop,
            value=value,
            evidence=[dict(chunk_id=chunk, quote=quote)],
        )
    for cid, fid, conditions in [
        ("p4.s0", "company", ["company"]),
        ("p6.s0", "name", ["company", "technical"]),
        ("p7.s0", "cf", ["company", "technical"]),
    ]:
        payload["fields"][cid] = dict(
            label=payload["facts"][fid]["property"],
            kind="data",
            disposition="conditional",
            reason="Applicabilità del ramo da verificare",
            property=payload["facts"][fid]["property"],
            fact_id=fid,
            conditions=conditions,
            form={
                "paragraph_id": catalog["candidates"][cid]["context_id"],
                "quote": ps[catalog["candidates"][cid]["context_id"]]["text_with_fields"],
            },
            **({"binding_id": "person"} if fid != "company" else {}),
        )
    audit = validate_global_proposal(
        json.dumps(payload), layout, "alpha", sources, expected_version=12
    )
    assert all(
        f.get("proposal_grounding") == "GROUNDED_PROPOSAL"
        for f in audit["fields"]
        if f["candidate_id"] in payload["fields"]
    )
    receipt = {
        "provenance": "USER",
        "action": "CONDITION_TRUE",
        "applies": True,
        "question": "Confermi la qualità di legale rappresentante di società di ingegneria?",
        "user_quote": "Vega S.r.l. partecipa soltanto come società di ingegneria.",
    }
    data = {
        "catalog": catalog,
        "audit": audit,
        "writable_ids": list(catalog["candidates"]),
        "dialogue": {
            "answers": {
                "condition:company": receipt,
                "condition:technical": {
                    "provenance": "USER",
                    "action": "SKIP",
                    "user_quote": "Non fornisco una conferma aggiuntiva sul ruolo tecnico.",
                },
            },
            "user_values": {},
        },
    }
    return data, layout, sources


def reconciled(data, layout, sources):
    from app.compilation_branch_evidence import reconcile_branches

    return reconcile_branches(data, layout, "alpha", sources)


def test_composed_branch_requires_distinct_user_form_and_named_source_proofs():
    data, layout, sources = case()
    before = copy.deepcopy(data)
    result = reconciled(data, layout, sources)
    assert data == before
    condition = result["conditions"]["technical"]
    assert condition["result"] == "true"
    proof = condition["branch_evidence"]
    assert proof["receipt_key"] == "condition:company"
    assert proof["role"] == "Responsabile Tecnico" and proof["subject"] == "Nora Serra"
    assert proof["operator"] == "Vega S.r.l." and proof["source_evidence"]
    for field in result["fields"]:
        if field["candidate_id"] in {"p6.s0", "p7.s0"}:
            assert field["condition_checks"]["technical"]["result"] == "true"
            assert field["written_value"] is None
    assert data["dialogue"]["answers"]["condition:technical"]["action"] == "SKIP"


def test_explicit_confirmation_prefix_does_not_change_the_registered_subject():
    data, layout, sources = case()
    receipt = data['dialogue']['answers']['condition:company']
    receipt['user_quote'] = ('Confermo quanto già indicato nella richiesta iniziale: '
                             + receipt['user_quote'])
    assert reconciled(data, layout, sources)['conditions']['technical']['result'] == 'true'


def test_unknown_cardinality_does_not_trust_an_arbitrary_user_value_receipt_key():
    data, layout, sources = case()
    data['audit'] = reconciled(data, layout, sources)
    value = copy.deepcopy(data['audit']['receipt_choice_values']['p3.cb0'])
    data['audit'].pop('receipt_choice_values')
    data['dialogue']['user_values']['p3.cb0'] = value
    from app.compilation_choices import materialize_choices

    with pytest.raises(ValueError, match='ricevuta'):
        materialize_choices(data, [], ['p3.cb0'])


@pytest.mark.parametrize(
    "change",
    [
        "no_receipt",
        "false",
        "uncertain",
        "other_type",
        "generic_yes",
        "quoted_choice",
        "wrong_scope",
        "wrong_form",
        "other_company",
        "other_person",
        "other_role",
        "negative_role",
        "hypothetical",
        "role_missing",
        "binding_rejected",
        "fact_rejected",
        "general_source",
        "direct_denial",
        "role_conflict",
    ],
)
def test_missing_or_non_equivalent_proofs_cannot_prove_branch(change):
    data, layout, sources = case()
    receipt = data["dialogue"]["answers"]["condition:company"]
    cond = data["audit"]["conditions"]["technical"]
    if change == "no_receipt":
        data["dialogue"]["answers"].pop("condition:company")
    elif change == "false":
        receipt.update(applies=False, action="CONDITION_FALSE")
    elif change == "uncertain":
        receipt["user_quote"] = "Forse Vega S.r.l. partecipa come società di ingegneria."
    elif change == "other_type":
        receipt["user_quote"] = "Vega S.r.l. partecipa come società di professionisti."
    elif change == "generic_yes":
        receipt["user_quote"] = "Sì."
    elif change == "quoted_choice":
        receipt["user_quote"] = 'La fonte dice: "Vega S.r.l. partecipa come società di ingegneria".'
    elif change == "wrong_scope":
        data["audit"]["conditions"]["company"]["applies_to"] = ["p3.cb0"]
    elif change == "wrong_form":
        cond["form_evidence"]["quote"] = "Responsabile per qualsiasi società"
    elif change == "binding_rejected":
        data["audit"]["bindings"]["person"]["validation"] = "REJECTED"
    elif change == "fact_rejected":
        data["audit"]["facts"]["name"]["validation"] = "REJECTED"
    elif change == "direct_denial":
        data["dialogue"]["answers"]["condition:technical"].update(
            action="CONDITION_FALSE", applies=False
        )
    elif change == "role_conflict":
        sources.append(
            {
                **sources[1],
                "chunk_id": -3,
                "content": "Nora Serra: non è Responsabile Tecnico di Vega S.r.l.",
            }
        )
    else:
        old = sources[1]["content"]
        changed = {
            "other_company": old.replace("Vega S.r.l.", "Altra S.r.l."),
            "other_person": old.replace("Nora Serra", "Ada Mori"),
            "other_role": old.replace("Responsabile Tecnico", "Socio"),
            "negative_role": old.replace(":** ", ":** non è ", 1),
            "hypothetical": old.replace(":** ", ":** forse ", 1),
            "role_missing": "Nora Serra: C.F. SRRNRO85A41F205D.",
            "general_source": old,
        }[change]
        sources[1]["content"] = changed
        if change == "general_source":
            sources[1]["category"] = "general"
    assert reconciled(data, layout, sources)["conditions"]["technical"]["result"] == "unknown"


def test_checkbox_receipt_maps_one_positive_and_keeps_cardinality_unknown():
    data, layout, sources = case()
    result = reconciled(data, layout, sources)
    choice = result["receipt_choice_values"]["p3.cb0"]
    assert choice["value"] is True and choice["provenance"] == "USER"
    assert choice["receipt_key"] == "condition:company"
    group = result["choice_groups"][choice["choice_group_id"]]
    assert group["validation"] == "GROUNDED_GROUP" and group["selection"] == "unknown"
    assert set(group["options"]) == {"p1.cb0", "p3.cb0", "p8.cb0"}
    assert set(result["receipt_choice_values"]) == {"p3.cb0"}
    from app.compilation_branch_evidence import effective_user_values
    from app.compilation_choices import materialize_choices

    data["audit"] = result
    assert effective_user_values(data)["p3.cb0"] == choice
    assert materialize_choices(data, [], ["p3.cb0"]) == ({"p3.cb0": True}, [])


@pytest.mark.parametrize(
    "change",
    [
        "absent",
        "false",
        "skip",
        "partial_form",
        "scope",
        "ambiguous_control",
        "checked_peer",
        "cross_heading",
    ],
)
def test_no_checkbox_without_exact_receipt_and_bounded_group(change):
    data, layout, sources = case()
    parent = data["audit"]["conditions"]["company"]
    if change == "absent":
        data["dialogue"]["answers"] = {}
    elif change == "false":
        data["dialogue"]["answers"]["condition:company"].update(
            applies=False, action="CONDITION_FALSE"
        )
    elif change == "skip":
        data["dialogue"]["answers"]["condition:company"].update(action="SKIP", applies=None)
    elif change == "partial_form":
        parent["form_evidence"]["quote"] = "società di ingegneria"
    elif change == "scope":
        parent["applies_to"].remove("p3.cb0")
    elif change == "ambiguous_control":
        data["catalog"]["candidates"]["p3.cb1"] = copy.deepcopy(
            data["catalog"]["candidates"]["p3.cb0"]
        )
    elif change == "checked_peer":
        data["catalog"]["candidates"]["p1.cb0"]["initial_checked"] = True
    elif change == "cross_heading":
        for b in data["catalog"]["blocks"]:
            for p in b["paragraphs"]:
                if p["id"] == "paragraph:2":
                    p["text"] = "DICHIARAZIONI"
                    p.pop("text_with_fields", None)
        data["catalog"]["candidates"].pop("p2.s0")
    assert not reconciled(data, layout, sources).get("receipt_choice_values")


def test_stale_or_tampered_derived_checkbox_cannot_pass_export():
    data, layout, sources = case()
    data["audit"] = reconciled(data, layout, sources)
    from app.compilation_branch_evidence import effective_user_values
    from app.compilation_choices import materialize_choices

    data["dialogue"]["answers"]["condition:company"]["applies"] = False
    assert "p3.cb0" not in effective_user_values(data)
    with pytest.raises((ValueError, KeyError)):
        materialize_choices(data, [], ["p3.cb0"])


@pytest.mark.parametrize('message', [
    'Lascia il campo vuoto.',
    'Lascialo vuoto, non fornisco una conferma aggiuntiva sul ruolo tecnico.',
])
def test_explicit_request_for_blank_is_not_a_deferral_of_confirmation_only(message):
    data, layout, sources = case()
    data['dialogue']['answers']['condition:technical']['user_quote'] = message
    assert reconciled(data, layout, sources)['conditions']['technical']['result'] == 'unknown'
