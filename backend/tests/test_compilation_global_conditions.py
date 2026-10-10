"""Conservative condition gates, not assertions of real model quality."""

import copy
import json
from io import BytesIO

import pytest
from docx import Document

from app.compilation_global import build_global_input, global_request, validate_global_proposal
from app.config import AISettings
from app.docx_templates import inspect_docx


@pytest.fixture
def case():
    doc = Document()
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Recapito principale"
    table.cell(1, 0).text = "Recapito secondario\n(solo se diverso dal principale)"
    stream = BytesIO()
    doc.save(stream)
    layout = inspect_docx(stream.getvalue())
    source = {"chunk_id": -7, "file_id": -3, "chunk_index": 0,
              "source_name": "profilo.txt", "role": "source", "scope": "global",
              "category": "company", "project_id": None,
              "content": "Recapito principale: Via Luna 4; secondario: Via Luna 4"}
    fact = {"subject": "Organizzazione", "entity": "company", "aspect": "location",
            "property": "recapito principale", "value": "Via Luna 4",
            "evidence": [{"chunk_id": -7, "quote": source["content"]}]}
    payload = {"version": 2, "facts": {"main": fact,
               "secondary": {**copy.deepcopy(fact), "property": "recapito secondario"}},
               "conditions": {"different": {
                   "form": {"paragraph_id": "paragraph:2",
                            "quote": "solo se diverso dal principale"},
                   "aspect": "location", "operator": "different",
                   "fact_ids": ["main", "secondary"],
                   "reason": "Confrontare i due recapiti della stessa organizzazione"}},
               "fields": {
                   "t0.r0.c1": {"label": "Recapito principale", "kind": "data",
                                "disposition": "value", "fact_id": "main", "conditions": [],
                                "form": {"paragraph_id": "paragraph:0",
                                         "quote": "Recapito principale"}},
                   "t0.r1.c1": {"label": "Recapito secondario", "kind": "data",
                                "disposition": "value", "fact_id": "secondary",
                                "conditions": ["different"],
                                "form": {"paragraph_id": "paragraph:2",
                                         "quote": "Recapito secondario"}},
               }}
    return layout, [source], payload


def audit(case):
    layout, sources, payload = case
    return validate_global_proposal(json.dumps(payload), layout, "alpha", sources)


def test_shared_contract_is_requested_without_a_second_call(case):
    layout, sources, _ = case
    schema = build_global_input(layout, "alpha", sources)["schema_output"]
    assert set(schema["properties"]) == {
        "version", "facts", "conditions", "reviewed_paragraph_ids",
        "uncertain_paragraph_ids", "fields", "deferred"}


@pytest.mark.parametrize("thinking", [False, True])
def test_thinking_comparison_preserves_input_contract_and_budgets(case, thinking):
    layout, sources, _ = case
    data = build_global_input(layout, "alpha", sources)
    settings = AISettings("secret", "deepseek-flash", "https://example.invalid")
    body = global_request(settings, data, thinking=thinking)
    assert body["thinking"]["type"] == ("enabled" if thinking else "disabled")
    assert body.get("reasoning_effort") == ("low" if thinking else None)
    assert body["max_tokens"] == 49152
    assert len(body["messages"]) == 2
    assert json.loads(body["messages"][1]["content"]) == data


def test_same_addresses_block_value_under_different_condition(case):
    result = audit(case)
    assert result["conditions"]["different"]["result"] == "false"
    assert result["fields"][1]["validation"] == "REJECTED"
    assert "falsa" in " ".join(result["fields"][1]["validation_errors"])
    assert result["fields"][0]["validation"] == "GROUNDED_PROPOSAL"
    assert result["resolved"] == result["writes"] == 0


def test_false_condition_can_ground_non_applicability_without_a_value(case):
    case[2]["fields"]["t0.r1.c1"].update(disposition="not_applicable", fact_id=None)
    result = audit(case)
    assert result["fields"][1]["validation"] == "GROUNDED_PROPOSAL"
    assert result["fields"][1]["source_evidence"][0]["chunk_id"] == -7
    assert result["fields"][1]["written_value"] is None


@pytest.mark.parametrize("disposition", ["value", "missing", "not_applicable"])
def test_unknown_condition_requires_review_instead_of_asserting_applicability(case, disposition):
    case[2]["conditions"]["different"].update(operator="semantic", fact_ids=[])
    case[2]["fields"]["t0.r1.c1"].update(disposition=disposition,
                                           fact_id="secondary" if disposition == "value" else None)
    result = audit(case)
    assert result["conditions"]["different"]["result"] == "unknown"
    assert result["fields"][1]["validation"] == "REJECTED"


def test_unknown_condition_is_kept_as_review(case):
    case[2]["conditions"]["different"].update(operator="semantic", fact_ids=[])
    case[2]["fields"]["t0.r1.c1"].update(disposition="review", fact_id=None)
    assert audit(case)["fields"][1]["validation"] == "GROUNDED_PROPOSAL"


def test_type_cannot_decide_participation_even_when_values_differ(case):
    case[2]["facts"]["main"]["aspect"] = "legal_type"
    case[2]["conditions"]["different"]["aspect"] = "participation"
    case[2]["fields"]["t0.r1.c1"].update(disposition="not_applicable", fact_id=None)
    result = audit(case)
    assert result["conditions"]["different"]["result"] == "unknown"
    assert "dimension" in result["conditions"]["different"]["validation_errors"][0]
    assert result["fields"][1]["validation"] == "REJECTED"


def test_two_subjects_cannot_be_used_as_equivalent_addresses(case):
    case[2]["facts"]["secondary"]["subject"] = "Altro ente"
    result = audit(case)
    assert result["conditions"]["different"]["result"] == "unknown"
    assert result["fields"][1]["validation"] == "REJECTED"


def test_different_spelling_is_not_proof_of_different_real_values(case):
    case[1][0]["content"] += "; altra grafia: Via Luna n. 4"
    case[2]["facts"]["secondary"].update(value="Via Luna n. 4", evidence=[
        {"chunk_id": -7, "quote": case[1][0]["content"]},
    ])
    result = audit(case)
    assert result["conditions"]["different"]["result"] == "unknown"


def test_true_same_condition_allows_a_proposal_but_never_a_resolved_value(case):
    doc = Document(BytesIO(case[0].original))
    doc.tables[0].cell(1, 0).text = "Recapito secondario\n(solo se coincide col principale)"
    stream = BytesIO()
    doc.save(stream)
    case = (inspect_docx(stream.getvalue()), *case[1:])
    case[2]["conditions"]["different"]["form"]["quote"] = "solo se coincide col principale"
    case[2]["conditions"]["different"]["operator"] = "same"
    result = audit(case)
    assert result["conditions"]["different"]["result"] == "true"
    assert result["fields"][1]["validation"] == "GROUNDED_PROPOSAL"
    assert result["fields"][1]["semantic_review_required"]
    assert result["fields"][1]["written_value"] is None


@pytest.mark.parametrize("mutation", ["same_id", "same_property", "other_aspect"])
def test_comparison_cannot_manufacture_a_false_condition(case, mutation):
    if mutation == "same_id":
        case[2]["conditions"]["different"]["fact_ids"] = ["main", "main"]
    elif mutation == "same_property":
        case[2]["facts"]["secondary"]["property"] = "recapito principale"
    else:
        case[2]["conditions"]["different"]["aspect"] = "other"
        for fact in case[2]["facts"].values():
            fact["aspect"] = "other"
    result = audit(case)
    assert result["conditions"]["different"]["result"] == "unknown"
    assert result["fields"][1]["validation"] == "REJECTED"


def test_a_nonlocal_false_condition_cannot_exclude_a_destination(case):
    doc = Document(BytesIO(case[0].original))
    text = "Questa altra sezione si applica solo se i recapiti sono diversi"
    doc.add_paragraph(text)
    stream = BytesIO()
    doc.save(stream)
    layout = inspect_docx(stream.getvalue())
    data = build_global_input(layout, "alpha", case[1])
    paragraph = next(p for block in data["form"]["blocks"] for p in block["paragraphs"]
                     if p["text"] == text)
    case[2]["conditions"]["different"]["form"] = {
        "paragraph_id": paragraph["id"], "quote": text,
    }
    case[2]["fields"]["t0.r1.c1"].update(disposition="not_applicable", fact_id=None)
    result = audit((layout, *case[1:]))
    assert result["conditions"]["different"]["result"] == "false"
    assert result["fields"][1]["validation"] == "REJECTED"


@pytest.mark.parametrize("change", [
    {"value": "Via Luna 4"}, {"evidence": []}, {"conditions": None},
    {"conditions": ["different", "different"]}, {"choice": True},
])
def test_explicit_assignment_schema_cannot_bypass_shared_proofs(case, change):
    case[2]["fields"]["t0.r0.c1"].update(change)
    assert audit(case)["fields"][0]["validation"] == "REJECTED"


def test_unknown_field_is_not_inferred_from_a_fact(case):
    case[2]["fields"]["invented"] = case[2]["fields"].pop("t0.r0.c1")
    with pytest.raises(ValueError, match="Coverage"):
        audit(case)


@pytest.mark.parametrize("change", [
    {"evidence": []}, {"value": "Inventato"}, {"evidence": [{"chunk_id": -99, "quote": "x"}]},
])
def test_bad_fact_quarantines_only_dependent_assignments(case, change):
    case[2]["facts"]["secondary"].update(change)
    result = audit(case)
    assert result["facts"]["secondary"]["validation"] == "REJECTED"
    assert result["fields"][0]["validation"] == "GROUNDED_PROPOSAL"
    assert result["fields"][1]["validation"] == "REJECTED"


@pytest.mark.parametrize("mutation", ["fact", "condition", "form", "condition_form", "coverage"])
def test_unknown_references_and_unsupported_form_are_rejected(case, mutation):
    field = case[2]["fields"]["t0.r1.c1"]
    if mutation == "fact":
        field["fact_id"] = "unknown"
    elif mutation == "condition":
        field["conditions"] = ["unknown"]
    elif mutation == "form":
        field["form"]["quote"] = "Premessa inventata dall'utente"
    elif mutation == "condition_form":
        case[2]["conditions"]["different"]["form"]["quote"] = "condizione inventata"
    else:
        case[2]["fields"].pop("t0.r0.c1")
        with pytest.raises(ValueError, match="Coverage"):
            audit(case)
        return
    assert audit(case)["fields"][1]["validation"] == "REJECTED"


def test_shared_fact_reused_without_duplicating_provenance_or_claiming_semantics(case):
    case[2]["fields"]["t0.r1.c1"].update(fact_id="main", conditions=[])
    result = audit(case)
    assert result["counts"] == {"GROUNDED_PROPOSAL": 2}
    for field in result["fields"]:
        assert field["fact_id"] == "main"
        assert field["source_evidence"][0]["file_id"] == -3
        assert len(field["source_evidence"][0]["content_sha256"]) == 64
        assert field["semantic_review_required"] and field["written_value"] is None


@pytest.mark.parametrize("origin", ["form", "user", "general"])
def test_fact_cannot_disguise_form_user_or_general_as_company_proof(case, origin):
    if origin == "general":
        case[1][0]["category"] = "general"
        assert audit(case)["facts"]["main"]["validation"] == "REJECTED"
    else:
        case[1][0]["role"] = origin
        with pytest.raises(ValueError, match="Ruolo SOURCE"):
            audit(case)


def test_duplicate_or_extra_json_remains_rejected(case):
    layout, sources, payload = case
    for raw in [json.dumps(payload) + "}", json.dumps(payload).replace(
            '"version": 2', '"version": 2, "version": 2')]:
        with pytest.raises(ValueError):
            validate_global_proposal(raw, layout, "alpha", sources)
