"""Session/domain regressions: isolated DB/storage, real FTS/Qdrant, simulated AI."""

import asyncio
import importlib
import json
from io import BytesIO
from pathlib import Path

import httpx
import pytest
from docx import Document
from ollama import Client, ListResponse

from app import compilation_session_resolution as resolution
from app import compilation_sessions as sessions
from app.compilation_session_models import CandidateMatches, CandidateMeanings
from app.compilation_sources import CompilationSearch
from app.db import connection, get_storage_path, init_database
from app.generation import GenerationError
from app.main import app
from app.repository import create_project
from app.retrieval_settings import RetrievalInput, save_settings
from app.schemas import ProjectCreate
from app.source_planning import SourceClusters, validate_source_plan

BASE = "/api/projects/alpha/compilation-sessions"
DEMO = (
    Path(__file__).resolve().parents[2]
    / "demo-documents/bandi/catanzaro-dl-cse/modello/domanda-partecipazione.docx"
)


@pytest.fixture
def anyio_backend():
    return "asyncio"


def docx(labels=("Denominazione sociale", "Sede legale"), context="Dati aziendali"):
    document = Document()
    document.add_paragraph(context)
    table = document.add_table(rows=len(labels), cols=2)
    for index, label in enumerate(labels):
        table.cell(index, 0).text = label
    stream = BytesIO()
    document.save(stream)
    return stream.getvalue()


@pytest.fixture(params=["fts5", "qdrant"])
async def api(tmp_path, monkeypatch, request):
    monkeypatch.setenv("MAPI_DB_PATH", str(tmp_path / "sessions.db"))
    monkeypatch.setenv("MAPI_STORAGE_PATH", str(tmp_path / "uploads"))
    monkeypatch.setenv("MAPI_KNOWLEDGE_PATH", str(tmp_path / "knowledge"))
    init_database()
    for name in ("Alpha", "Beta"):
        create_project(ProjectCreate(title=name, description="Test sessioni"))
    monkeypatch.setattr(
        Client,
        "list",
        lambda *args: ListResponse(
            models=[
                {"model": "bge-m3", "digest": "test-session-v1"},
            ]
        ),
    )
    monkeypatch.setattr(
        Client,
        "embed",
        lambda *args, **kwargs: {
            "embeddings": [[1.0, 0.5, 0.1] for _ in kwargs["input"]],
        },
    )
    save_settings(RetrievalInput(backend=request.param))
    async with httpx.AsyncClient(
        transport=httpx.ASGITransport(app=app), base_url="http://test"
    ) as client:
        yield client


async def create(api, content=None, project="alpha"):
    data = content if content is not None else docx()
    uploaded = await api.post(
        f"/api/projects/{project}/forms", files={"file": ("domanda-partecipazione.docx", data)}
    )
    assert uploaded.status_code == 201, uploaded.text
    response = await api.post(
        f"/api/projects/{project}/compilation-sessions", json={"form_id": uploaded.json()["id"]}
    )
    assert response.status_code == 201, response.text
    return response.json()


async def source(api, text, *, project=None, category="company", name="fonte.txt"):
    path = f"/api/projects/{project}/files" if project else "/api/global-knowledge/files"
    response = await api.post(
        path,
        files={"file": (name, text.encode())},
        data={"category": category} if not project else {},
    )
    assert response.status_code == 201, response.text
    return response.json()


def simulate(monkeypatch, values=None, *, role="", relationship="single", classify_hook=None):
    """Only the AI is simulated; every support still traverses all backend gates."""
    calls = {"classify": 0, "planner": 0, "match": 0}
    values = values or {}

    # These tests concern field resolution; document planning has dedicated tests.
    from app import compilation_document_plan
    monkeypatch.setattr(compilation_document_plan, "needs_document_plan", lambda state: False)
    monkeypatch.setattr(resolution, "needs_document_plan", lambda state: False)

    async def classify(fields):
        calls["classify"] += 1
        if classify_hook:
            return classify_hook(fields)
        return CandidateMeanings(
            fields=[
                {
                    "candidate_id": f["id"],
                    "classification": "data",
                    "requirement": {
                        "name": f["label"].rstrip(":"),
                        "form_quote": f["context"],
                        "form_citation_id": 1,
                        "person_role": role,
                    },
                    "form_quote": f["context"],
                    "entity": "person" if role else "company",
                    "reason": "Campo strutturale",
                    "kind": "data",
                }
                for f in fields
                if f["label"].strip()
            ]
        )

    async def plan(fields, request):
        calls["planner"] += 1
        from app.requirement_checks import Requirement

        requirements = [Requirement.model_validate(f["requirement"]) for f in fields]
        return CompilationSearch(validate_source_plan(
            SourceClusters(
                clusters=[
                    {"requirement_ids": list(range(start, min(start + 4, len(requirements) + 1)))}
                    for start in range(1, len(requirements) + 1, 4)
                ]
            ),
            requirements,
        ), {i: [r.name] for i, r in enumerate(requirements, 1)})

    async def match(fields, sources, coverage):
        calls["match"] += 1
        proposals = []
        for f in fields:
            supports = []
            for value in values.get(f["id"], []):
                for index in coverage.get(f["id"], []):
                    s = sources[index - 1]
                    if value in s["content"]:
                        supports.append({"source_id": index, "quote": s["content"], "value": value})
                        break
            if f["id"] in coverage:
                proposals.append(
                    {
                        "candidate_id": f["id"],
                        "supports": supports,
                        "relationship": relationship,
                        "reason": "Riscontro nelle fonti",
                    }
                )
        return CandidateMatches(fields=proposals)

    monkeypatch.setattr(resolution, "classify_candidates", classify)
    monkeypatch.setattr(resolution, "plan_compilation_search", plan)
    monkeypatch.setattr(resolution, "match_candidates", match)
    return calls


async def resolve(api, session, ids=None):
    response = await api.post(
        f"{BASE}/{session['id']}/resolve",
        json={
            "version": session["version"],
            **({"field_ids": ids} if ids else {}),
        },
    )
    assert response.status_code == 200, response.text
    return response.json()


@pytest.mark.anyio
async def test_create_resume_snapshot_and_original_immutable(api):
    original = docx()
    state = await create(api, original)
    assert state["status"] == "CREATED" and state["version"] == 1
    assert state["summary"]["pending"] == 2
    assert [f["id"] for f in state["fields"]] == ["t0.r0.c1", "t0.r1.c1"]
    # New module instance/connection: state is in SQLite, never an object cache.
    importlib.reload(sessions)
    again = (await api.get(f"{BASE}/{state['id']}")).json()
    assert again == state
    downloaded = await api.get(f"/api/projects/alpha/forms/{state['form_id']}/download")
    assert downloaded.content == original
    with connection() as db:
        assert db.execute("SELECT original FROM compilation_sessions").fetchone()[0] == original
        assert db.execute("SELECT COUNT(*) FROM document_compilations").fetchone()[0] == 0
    assert not list(get_storage_path().rglob("bozza.docx"))


@pytest.mark.anyio
async def test_resolved_missing_and_form_cannot_verify(api, monkeypatch):
    state = await create(api, docx(context="Sede legale: Roma. È solo il modulo."))
    await source(api, "Denominazione sociale: Mapi Ingegneria S.r.l.")
    simulate(monkeypatch, {"t0.r0.c1": ["Mapi Ingegneria S.r.l."], "t0.r1.c1": ["Roma"]})
    result = await resolve(api, state)
    first, second = result["fields"]
    assert first["status"] == "RESOLVED" and first["provenance"] == "SOURCE"
    assert first["source_evidence"][0]["role"] == "source"
    assert first["source_evidence"][0]["category"] == "company"
    assert first["source_evidence"][0]["chunk_id"] < 0
    assert second["status"] == "MISSING" and second["value"] is None
    assert result["status"] == "WAITING_FOR_USER"
    assert result["summary"]["resolved"] == result["summary"]["missing"] == 1


@pytest.mark.anyio
@pytest.mark.parametrize(
    "relationship,expected", [("alternatives", "AMBIGUOUS"), ("conflict", "CONFLICTING")]
)
async def test_competing_supported_values_are_not_arbitrarily_chosen(
    api,
    monkeypatch,
    relationship,
    expected,
):
    state = await create(api, docx(("Sede legale",)))
    await source(api, "Sede legale: Via Uno, Bari.", name="uno.txt")
    await source(api, "Sede legale: Via Due, Bari.", name="due.txt")
    simulate(
        monkeypatch, {"t0.r0.c1": ["Via Uno, Bari", "Via Due, Bari"]}, relationship=relationship
    )
    result = await resolve(api, state)
    field = result["fields"][0]
    assert field["status"] == expected and field["value"] is None
    assert len(field["alternatives"]) == 2
    assert all(a["evidence"]["role"] == "source" for a in field["alternatives"])
    assert (
        await api.post(f"{BASE}/{state['id']}/finalize", json={"version": result["version"]})
    ).status_code == 409


@pytest.mark.anyio
async def test_user_update_local_history_concurrency_and_regeneration(api, monkeypatch):
    original = docx()
    state = await create(api, original)
    await source(api, "Denominazione sociale: Mapi Ingegneria S.r.l.")
    simulate(monkeypatch, {"t0.r0.c1": ["Mapi Ingegneria S.r.l."]})
    state = await resolve(api, state)
    first = state["fields"][0]
    stale_version = state["version"]
    url = f"{BASE}/{state['id']}"
    change = {"field_id": "t0.r1.c1", "value": "Via Uno, Bari"}
    response = await api.patch(
        url + "/fields",
        json={
            "version": state["version"],
            "fields": [change],
        },
    )
    state = response.json()
    assert state["status"] == "READY" and state["fields"][0] == first
    assert state["fields"][1]["status"] == "USER_PROVIDED"
    assert state["fields"][1]["provenance"] == "USER"
    assert state["fields"][1]["source_evidence"] == []
    stale = await api.patch(url + "/fields", json={"version": stale_version, "fields": [change]})
    assert stale.status_code == 409
    outputs = []
    for index, value in enumerate(("Via Uno, Bari", "Via Due, Bari")):
        if index:
            state = (
                await api.patch(
                    url + "/fields",
                    json={
                        "version": state["version"],
                        "fields": [
                            {"field_id": "t0.r1.c1", "value": value},
                        ],
                    },
                )
            ).json()
        response = await api.post(url + "/finalize", json={"version": state["version"]})
        assert response.status_code == 200, response.text
        state = response.json()
        assert state["status"] == "GENERATED"
        downloads = state["last_generation"]["downloads"]
        draft = await api.get(downloads["docx"])
        outputs.append(draft.content)
        document = Document(BytesIO(draft.content))
        assert document.tables[0].cell(1, 1).text == value
        assert document.tables[0].cell(0, 1).text == "Mapi Ingegneria S.r.l."
        assert (await api.get(downloads["template"])).content == original
        report = (await api.get(downloads["report"])).json()
        assert report["fields"][1]["evidence"][0]["origin"] == "user"
        assert report["ready_for_submission"] is False
    assert outputs[0] != outputs[1]
    history = (await api.get(url + "/revisions")).json()
    update = next(e for e in reversed(history) if e["action"] == "user_update")
    assert len(update["changes"]) == 1
    assert update["changes"][0]["before"]["value"] == "Via Uno, Bari"
    assert update["changes"][0]["after"]["value"] == "Via Due, Bari"
    with connection() as db:
        assert db.execute("SELECT COUNT(*) FROM document_compilations").fetchone()[0] == 2
        assert (
            db.execute(
                "SELECT COUNT(*) FROM document_chunks WHERE content LIKE '%Via Due%'",
            ).fetchone()[0]
            == 0
        )


@pytest.mark.anyio
async def test_user_not_applicable_and_explicit_draft(api):
    state = await create(api)
    url = f"{BASE}/{state['id']}"
    state = (
        await api.patch(
            url + "/fields",
            json={
                "version": 1,
                "fields": [
                    {
                        "field_id": "t0.r0.c1",
                        "action": "not_applicable",
                        "reason": "Sezione non pertinente",
                    }
                ],
            },
        )
    ).json()
    assert state["summary"]["not_applicable"] == 1 and state["summary"]["missing"] == 0
    assert state["summary"]["pending"] == 1
    blocked = await api.post(url + "/finalize", json={"version": state["version"]})
    assert blocked.status_code == 409
    result = await api.post(
        url + "/finalize", json={"version": state["version"], "allow_unresolved": True}
    )
    assert result.status_code == 200, result.text
    output = await api.get(result.json()["last_generation"]["downloads"]["docx"])
    assert Document(BytesIO(output.content)).tables[0].cell(1, 1).text == ""


@pytest.mark.anyio
async def test_existing_email_and_signature_validators_and_user_choice(api):
    state = await create(api, docx(("PEC", "Firma", "Scelta")))
    url = f"{BASE}/{state['id']}"
    for field, value in (("t0.r0.c1", "incompleta"), ("t0.r1.c1", "Mario Rossi")):
        response = await api.patch(
            url + "/fields",
            json={
                "version": 1,
                "fields": [
                    {"field_id": "t0.r2.c1", "value": "No"},
                    {"field_id": field, "value": value},
                ],
            },
        )
        assert response.status_code == 422
        assert (await api.get(url)).json()["version"] == 1
    response = await api.patch(
        url + "/fields",
        json={
            "version": 1,
            "fields": [
                {"field_id": "t0.r2.c1", "value": "No"},
            ],
        },
    )
    assert response.status_code == 200, response.text
    assert response.json()["fields"][2]["provenance"] == "USER"


@pytest.mark.anyio
async def test_project_and_conversation_isolation(api, monkeypatch):
    state = await create(api, docx(("Sede legale",)))
    await source(api, "Sede legale: Via Segreta, Altro Progetto", project="beta")
    simulate(monkeypatch, {"t0.r0.c1": ["Via Segreta, Altro Progetto"]})
    result = await resolve(api, state)
    assert result["fields"][0]["status"] == "MISSING"
    assert result["last_resolution"]["source_count"] == 0
    bad_base = BASE.replace("alpha", "beta")
    assert (await api.post(bad_base, json={"form_id": state["form_id"]})).status_code == 404
    assert (await api.get(f"{bad_base}/{state['id']}")).status_code == 404
    assert (
        await api.patch(
            f"{bad_base}/{state['id']}/fields",
            json={
                "version": 1,
                "fields": [
                    {"field_id": "t0.r0.c1", "value": "Rubato"},
                ],
            },
        )
    ).status_code == 404
    assert (
        await api.post(BASE, json={"form_id": state["form_id"], "conversation_id": "inesistente"})
    ).status_code == 404


@pytest.mark.anyio
async def test_provider_failure_and_retry_do_not_lose_fields(api, monkeypatch):
    state = await create(api)
    url = f"{BASE}/{state['id']}"
    state = (
        await api.patch(
            url + "/fields",
            json={
                "version": 1,
                "fields": [
                    {"field_id": "t0.r0.c1", "value": "Mapi"},
                ],
            },
        )
    ).json()

    async def failure(_):
        raise GenerationError("Provider interrotto")

    monkeypatch.setattr(resolution, "classify_candidates", failure)
    result = await api.post(url + "/resolve", json={"version": state["version"]})
    assert result.status_code == 502
    failed = (await api.get(url)).json()
    assert failed["status"] == "FAILED" and failed["lease_until"] is None
    assert failed["fields"] == state["fields"]
    simulate(monkeypatch)
    retried = await resolve(api, failed)
    assert retried["fields"][0] == state["fields"][0]
    assert retried["fields"][1]["status"] == "MISSING"


@pytest.mark.anyio
async def test_claim_prevents_overwrites_and_expired_claim_is_recoverable(api, monkeypatch):
    state = await create(api)
    claim, _, _ = sessions.claim_resolution("alpha", state["id"], 1, None)
    url = f"{BASE}/{state['id']}"
    assert (await api.post(url + "/resolve", json={"version": claim["version"]})).status_code == 409
    with connection() as db:
        row = sessions._row(db, "alpha", state["id"])
        stored = json.loads(row["state_json"])
        stored["lease_until"] = "2000-01-01T00:00:00+00:00"
        db.execute(
            "UPDATE compilation_sessions SET state_json=? WHERE id=?",
            (json.dumps(stored), state["id"]),
        )
    simulate(monkeypatch)
    recovered = await resolve(api, claim)
    assert recovered["status"] == "WAITING_FOR_USER"
    assert recovered["summary"]["missing"] == 2


@pytest.mark.anyio
async def test_changed_source_blocks_export_and_snapshot_survives_form_deletion(api, monkeypatch):
    state = await create(api, docx(("Denominazione sociale",)))
    await source(api, "Denominazione sociale: Mapi Ingegneria S.r.l.")
    simulate(monkeypatch, {"t0.r0.c1": ["Mapi Ingegneria S.r.l."]})
    state = await resolve(api, state)
    assert state["status"] == "READY"
    assert (await api.delete(f"/api/projects/alpha/forms/{state['form_id']}")).status_code == 204
    url = f"{BASE}/{state['id']}"
    state = (await api.get(url)).json()
    assert state["form_id"] is None
    generated = await api.post(url + "/finalize", json={"version": state["version"]})
    assert generated.status_code == 200, generated.text
    state = generated.json()
    with connection() as db:
        db.execute("UPDATE global_document_chunks SET content='Fonte cambiata'")
    blocked = await api.post(url + "/finalize", json={"version": state["version"]})
    assert blocked.status_code == 409


@pytest.mark.anyio
async def test_storage_failure_rolls_back_output_and_marks_failed(api, monkeypatch):
    state = await create(api)
    real_write = Path.write_bytes

    def fail(path, data):
        if path.name == "bozza.docx":
            raise OSError("Disco pieno")
        return real_write(path, data)

    monkeypatch.setattr(Path, "write_bytes", fail)
    response = await api.post(
        f"{BASE}/{state['id']}/finalize", json={"version": 1, "allow_unresolved": True}
    )
    assert response.status_code == 500
    assert (await api.get(f"{BASE}/{state['id']}")).json()["status"] == "FAILED"
    with connection() as db:
        assert db.execute("SELECT COUNT(*) FROM document_compilations").fetchone()[0] == 0
    assert not list(get_storage_path().rglob("template.docx"))


@pytest.mark.anyio
async def test_batch_is_bounded_and_not_searched_stays_pending(api, monkeypatch):
    state = await create(api, docx(tuple(f"Campo {i}" for i in range(40))))
    calls = simulate(monkeypatch)
    first = await resolve(api, state)
    assert len(first["last_resolution"]["candidate_ids"]) == 12
    assert first["summary"]["missing"] == 12 and first["summary"]["pending"] == 28
    assert calls == {"classify": 1, "planner": 1, "match": 0}
    assert first["last_resolution"]["query_count"] <= 12
    second = await resolve(api, first)
    assert second["summary"]["pending"] == 16
    assert second["fields"][:12] == first["fields"][:12]

    async def insufficient(fields, request):
        from app.requirement_checks import Requirement

        reqs = [Requirement.model_validate(f["requirement"]) for f in fields]
        return CompilationSearch(validate_source_plan(SourceClusters(clusters=[]), reqs),
                                 {i: [r.name] for i, r in enumerate(reqs, 1)})

    monkeypatch.setattr(resolution, "plan_compilation_search", insufficient)
    third = await resolve(api, second)
    assert third["summary"]["missing"] == 30
    assert third["summary"]["pending"] == 10
    assert sum(f["search"]["status"] == "coverage_limit" for f in third["fields"]) == 6


@pytest.mark.anyio
async def test_false_form_quote_and_wrong_field_source_never_resolve(api, monkeypatch):
    state = await create(api, docx(("Denominazione sociale", "Direttore tecnico")))
    await source(api, "Denominazione sociale: Mapi Ingegneria S.r.l.")
    simulate(monkeypatch, {"t0.r1.c1": ["Mapi Ingegneria S.r.l."]})
    result = await resolve(api, state)
    assert result["fields"][1]["status"] == "MISSING"
    assert result["fields"][1]["validation_errors"]

    def false_claim(fields):
        return CandidateMeanings(
            fields=[
                {
                    "candidate_id": fields[0]["id"],
                    "classification": "data",
                    "requirement": {
                        "name": "ISO 9001",
                        "form_quote": "ISO 9001",
                        "form_citation_id": 1,
                    },
                    "form_quote": "ISO 9001",
                    "reason": "Premessa utente non nel FORM",
                }
            ]
        )

    simulate(monkeypatch, classify_hook=false_claim)
    result = await resolve(api, result, ["t0.r0.c1"])
    assert result["fields"][0]["status"] == "PENDING"
    assert result["fields"][0]["requirement"] is None


@pytest.mark.anyio
async def test_real_demo_structural_candidates_verified_separately(api, monkeypatch):
    state = await create(api, DEMO.read_bytes())
    assert state["summary"]["total"] == 279
    # Names, legal form and director data alone do not prove branch applicability.
    # This focused placement test starts from an explicit, already verified USER
    # type decision; UNKNOWN and rejected type proofs have generic regressions.
    with connection() as db:
        saved = json.loads(db.execute(
            "SELECT state_json FROM compilation_sessions WHERE id=?", (state["id"],),
        ).fetchone()[0])
        engineering = next(f for f in saved["fields"] if f["id"] == "t25.r0.c1")
        engineering["form_dependency_user"] = {
            "condition": engineering["form_dependency"]["condition"], "applies": True,
            "provenance": "USER", "user_quote": "Siamo una società di ingegneria",
        }
        db.execute("UPDATE compilation_sessions SET state_json=? WHERE id=?",
                   (json.dumps(saved), state["id"]))
    text = (
        "Denominazione sociale: Mapi Ingegneria S.r.l.\n"
        "Forma giuridica: Società a responsabilità limitata\nSede legale: Via Test 12, Bari.\n"
        "Direttore tecnico: Nome e cognome Elisa Romano; qualifica professionale Ingegnere; "
        "Ordine professionale di appartenenza Ordine degli Ingegneri di Bari; "
        "iscrizione professionale simulata n. 8421."
    )
    await source(api, text)
    values = {
        "t25.r0.c1": ["Mapi Ingegneria S.r.l."],
        "t25.r4.c1": ["Società a responsabilità limitata"],
        "t25.r4.c3": ["Via Test 12, Bari"],
        "t26.r0.c1": ["Elisa Romano"],
        "t26.r1.c1": ["Ingegnere"],
        "t26.r3.c1": ["Ordine degli Ingegneri di Bari"],
        "t26.r4.c1": ["8421"],
    }
    simulate(monkeypatch, values)
    first = await resolve(api, state, ["t25.r0.c1", "t25.r4.c1", "t25.r4.c3"])
    simulate(monkeypatch, values, role="direttore tecnico")
    second = await resolve(api, first, [f"t26.r{i}.c1" for i in range(5)])
    fields = {f["id"]: f for f in second["fields"]}
    for candidate, value in values.items():
        assert fields[candidate]["value"] == value[0], fields[candidate]
        assert fields[candidate]["status"] == "RESOLVED"
        assert fields[candidate]["source_evidence"][0]["role"] == "source"
    assert fields["t26.r2.c1"]["status"] == "MISSING"  # qualification date
    assert second["summary"]["pending"] + second["summary"]["not_applicable"] == 271
    assert second["summary"]["not_applicable"] > 0
    assert fields["t40.r0.c1"]["status"] == "PENDING"  # Participation remains UNKNOWN.
    # Fixture data are simulated: this is not an evaluation of a live model.


@pytest.mark.anyio
async def test_automatic_exclusion_requires_explicit_source(api, monkeypatch):
    state = await create(api, docx(("Sede studio",), "Sezione studio associato"))
    await source(api, "Sezione studio associato: non applicabile. Sede studio non prevista.")

    def classify(fields):
        f = fields[0]
        return CandidateMeanings(
            fields=[
                {
                    "candidate_id": f["id"],
                    "classification": "data",
                    "condition": "Sezione studio associato",
                    "requirement": {
                        "name": "Sede studio",
                        "form_quote": f["context"],
                        "form_citation_id": 1,
                    },
                    "form_quote": f["context"],
                    "reason": "Campo condizionale",
                }
            ]
        )

    simulate(monkeypatch, classify_hook=classify)

    async def match(fields, sources, coverage):
        return CandidateMatches(
            fields=[
                {
                    "candidate_id": fields[0]["id"],
                    "supports": [],
                    "exclusion": {
                        "source_id": 1,
                        "quote": sources[0]["content"],
                        "value": "non applicabile",
                    },
                    "reason": "Esclusione esplicita",
                }
            ]
        )

    monkeypatch.setattr(resolution, "match_candidates", match)
    result = await resolve(api, state)
    assert result["fields"][0]["status"] == "NOT_APPLICABLE", result["fields"][0]
    assert result["fields"][0]["provenance"] == "SOURCE"
    assert result["summary"]["missing"] == 0


@pytest.mark.anyio
async def test_cancelled_resolution_persists_failure(api, monkeypatch):
    state = await create(api)

    async def cancelled(_):
        raise asyncio.CancelledError()

    monkeypatch.setattr(resolution, "classify_candidates", cancelled)
    with pytest.raises(asyncio.CancelledError):
        await resolution.resolve_session("alpha", state["id"], 1)
    assert sessions.get_session("alpha", state["id"])["status"] == "FAILED"
