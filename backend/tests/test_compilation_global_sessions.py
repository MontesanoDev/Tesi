"""Default chat adapter: isolated DB/storage, simulated provider, actual V10 gates/rendering."""

import asyncio
import hashlib
import json
from io import BytesIO

import pytest
from docx import Document
from test_compilation_sessions import api as legacy_api
from test_compilation_sessions import docx, source

from app import compilation_chat
from app import compilation_global_sessions as global_sessions
from app import compilation_sessions as sessions
from app.config import AISettings
from app.db import connection, get_storage_path
from app.generation import GenerationError
from app.intents import ChatDecision

api = legacy_api
BASE = "/api/projects/alpha/compilation-sessions"


@pytest.fixture
def anyio_backend():
    return "asyncio"


@pytest.fixture
def provider(monkeypatch):
    monkeypatch.setattr(compilation_chat, "CHAT_ENGINE", "global")
    monkeypatch.setattr(
        global_sessions,
        "get_ai_settings",
        lambda: AISettings(
            api_key="test-secret", model="simulated", base_url="https://provider.invalid"
        ),
    )
    # Routes pin project settings separately; the transport boundary alone is simulated.
    from app import ai_profiles

    monkeypatch.setattr(
        ai_profiles, "resolve_project_settings", lambda project: global_sessions.get_ai_settings()
    )
    calls = []

    async def request(client, settings, body, timeout, **kwargs):
        data = json.loads(body["messages"][1]["content"])
        calls.append(data)
        assert data["user_values"] == []
        assert all(s["role"] == "source" for s in data["sources"])
        return json.dumps(proposal(data)), "simulated", {"total_tokens": 10}

    monkeypatch.setattr(global_sessions, "request_model_content", request)
    from app import compilation_global_dialogue as dialogue

    monkeypatch.setattr(dialogue, "get_ai_settings", global_sessions.get_ai_settings)

    async def summarize(client, settings, body, timeout, **kwargs):
        data = json.loads(body["messages"][1]["content"])
        return json.dumps({"message": "Ho trovato dati nelle fonti. Restano parti da verificare.",
                           "questions": [{"slot": q["slot"],
                                          "question": "Questa situazione si applica al tuo caso?"}
                                         for q in data["question_pool"][:3]]}), "simulated", {}

    monkeypatch.setattr(dialogue, "request_model_content", summarize)
    return calls


def proposal(data, *, conditional=False, invalid_source=False):
    paragraphs = {p["id"]: p["text"] for b in data["form"]["blocks"] for p in b["paragraphs"]}
    facts, fields = {}, {}
    labels = ["Denominazione sociale", "Sede legale"]
    for index, label in enumerate(labels):
        key = f"t0.r{index}.c1"
        evidence = next(s for s in data["sources"] if "Aurora" in s["content"])
        quote = evidence["content"].splitlines()[index]
        value = quote.partition(":")[2].strip()
        fact = f"f{index}"
        facts[fact] = {
            "subject": "Aurora Servizi S.r.l.",
            "entity": "company",
            "aspect": "identity" if index == 0 else "location",
            "property": label,
            "value": value,
            "evidence": [
                {"chunk_id": 99999 if invalid_source else evidence["chunk_id"], "quote": quote}
            ],
        }
        paragraph = next(
            p
            for p in data["form"]["candidates"][key]["anchor_paragraph_ids"]
            if paragraphs[p] == label
        )
        fields[key] = {
            "label": label,
            "kind": "data",
            "disposition": "conditional" if conditional else "value",
            "form": {"paragraph_id": paragraph, "quote": label},
            "conditions": ["c"] if conditional else [],
            "fact_id": fact,
            "property": label,
        }
        if conditional:
            fields[key]["reason"] = "Applicabilità da confermare"
    conditions = (
        {
            "c": {
                "form": {"paragraph_id": "paragraph:0", "quote": "Dati aziendali"},
                "aspect": "participation",
                "operator": "semantic",
                "fact_ids": [],
                "scope": [{"kind": "table", "id": "table:0"}],
                "reason": "Applicabilità richiede conferma",
            }
        }
        if conditional
        else {}
    )
    return {
        "version": data["schema_output"]["properties"]["version"]["const"],
        "uncertain_paragraph_ids": [],
        "facts": facts,
        "fields": fields,
        "bindings": {},
        "conditions": conditions,
    }


async def start(api, *, conversation_id=None):
    upload = await api.post("/api/projects/alpha/forms", files={"file": ("modulo.docx", docx())})
    assert upload.status_code == 201, upload.text
    await source(api, "Denominazione sociale: Aurora Servizi S.r.l.\nSede legale: Via del Sole 9")
    response = await api.post(
        BASE,
        json={
            "form_id": upload.json()["id"],
            "engine": "global",
            "start_in_chat": True,
            "conversation_id": conversation_id,
        },
    )
    assert response.status_code == 201, response.text
    return response.json()


async def resolve(api, session):
    response = await api.post(
        f"{BASE}/{session['id']}/resolve", json={"version": session["version"], "automatic": True}
    )
    assert response.status_code == 200, response.text
    return response.json()


@pytest.mark.anyio
async def test_default_chat_start_and_one_call_persistent_review(api, provider):
    s = await start(api)
    decision = ChatDecision(action="compile", answer="", target="form", queries=[])
    text, command = await compilation_chat.handle_decision(
        "alpha", s["conversation_id"], decision, {"form_id": s["form_id"]}, None, "compilalo"
    )
    assert command["session_id"] == s["id"]
    assert "Analizzo" in text
    assert s["engine"] == "global" and s["summary"]["total"] == 2
    s = await resolve(api, s)
    assert len(provider) == 1
    assert s["status"] == "WAITING_FOR_USER" and not s["chat"]["auto_continue"]
    assert s["global_compilation"]["proposal_count"] == 2
    assert s["summary"]["resolved"] == 0 and s["last_generation"] is None
    assert s == (await api.get(f"{BASE}/{s['id']}")).json()
    duplicate = await api.post(
        BASE,
        json={
            "form_id": s["form_id"],
            "engine": "global",
            "conversation_id": s["conversation_id"],
            "start_in_chat": True,
        },
    )
    assert duplicate.json()["id"] == s["id"] and duplicate.json()["version"] == s["version"]
    assert (
        await api.post(f"{BASE}/{s['id']}/resolve", json={"version": s["version"]})
    ).status_code == 409
    assert len(provider) == 1
    report = await api.get(f"{BASE}/{s['id']}/global-report")
    assert report.status_code == 200 and report.json()["resolved"] == 0


@pytest.mark.anyio
async def test_review_export_and_regeneration_from_original(api, provider):
    s = await resolve(api, await start(api))
    with connection() as db:
        original = bytes(
            db.execute(
                "SELECT original FROM compilation_sessions WHERE id=?", (s["id"],)
            ).fetchone()[0]
        )
    original_sha = hashlib.sha256(original).hexdigest()
    for ids in [["t0.r0.c1"], ["t0.r1.c1"]]:
        request = {"version": s["version"], "candidate_ids": ids, "allow_unresolved": True}
        response = await api.post(f"{BASE}/{s['id']}/global-draft", json=request)
        assert response.status_code == 200, response.text
        s = response.json()
        assert not s["chat"]["auto_continue"] and s["summary"]["resolved"] == 0
        generated = s["last_generation"]
        doc = Document(BytesIO((await api.get(generated["downloads"]["docx"])).content))
        assert doc.tables[0].cell(0, 1).text == (
            "Aurora Servizi S.r.l." if ids[0] == "t0.r0.c1" else ""
        )
        assert doc.tables[0].cell(1, 1).text == ("Via del Sole 9" if ids[0] == "t0.r1.c1" else "")
        report = (await api.get(generated["downloads"]["report"])).json()
        assert report["template_sha256"] == original_sha
        assert report["fields"][0]["provenance"] == "SOURCE"
        assert report["fields"][0]["reviewed_by"] == "USER"
        assert report["user_review"]["provenance"] == "USER" and report["resolved"] == 0
        assert report["written_field_count"] == 1
        assert (
            hashlib.sha256((await api.get(generated["downloads"]["template"])).content).hexdigest()
            == original_sha
        )
    assert len(provider) == 1


@pytest.mark.anyio
async def test_unknown_applicability_requires_explicit_confirmation(api, provider, monkeypatch):
    async def conditional(client, settings, body, timeout, **kwargs):
        return (
            json.dumps(proposal(json.loads(body["messages"][1]["content"]), conditional=True)),
            "simulated",
            {},
        )

    monkeypatch.setattr(global_sessions, "request_model_content", conditional)
    s = await resolve(api, await start(api))
    assert s["global_compilation"]["proposal_count"] == 2
    request = {"version": s["version"], "candidate_ids": ["t0.r0.c1"], "allow_unresolved": True}
    assert (await api.post(f"{BASE}/{s['id']}/global-draft", json=request)).status_code == 409
    request["confirmed_condition_ids"] = ["c"]
    generated = await api.post(f"{BASE}/{s['id']}/global-draft", json=request)
    assert generated.status_code == 200, generated.text
    assert generated.json()["status"] == "GENERATED"


@pytest.mark.anyio
@pytest.mark.parametrize(
    "change", ["no_consent", "unknown_candidate", "duplicate", "ordinary_finalize"]
)
async def test_no_implicit_or_unvalidated_export(api, provider, change):
    s = await resolve(api, await start(api))
    request = {"version": s["version"], "candidate_ids": ["t0.r0.c1"], "allow_unresolved": True}
    if change == "no_consent":
        request["allow_unresolved"] = False
    if change == "unknown_candidate":
        request["candidate_ids"] = ["invented"]
    if change == "duplicate":
        request["candidate_ids"] = ["t0.r0.c1", "t0.r0.c1"]
    route = "finalize" if change == "ordinary_finalize" else "global-draft"
    body = {"version": s["version"], "allow_unresolved": True} if route == "finalize" else request
    assert (await api.post(f"{BASE}/{s['id']}/{route}", json=body)).status_code in {409, 422}
    assert (await api.get(f"{BASE}/{s['id']}")).json()["last_generation"] is None


@pytest.mark.anyio
async def test_source_mutation_blocks_export(api, provider):
    s = await resolve(api, await start(api))
    with connection() as db:
        db.execute("UPDATE global_document_chunks SET content=content || ' Changed'")
    response = await api.post(
        f"{BASE}/{s['id']}/global-draft",
        json={"version": s["version"], "candidate_ids": ["t0.r0.c1"], "allow_unresolved": True},
    )
    assert response.status_code == 409, response.text
    assert (await api.get(f"{BASE}/{s['id']}")).json()["last_generation"] is None


@pytest.mark.anyio
async def test_invalid_source_is_not_selectable(api, provider, monkeypatch):
    async def invalid(client, settings, body, timeout, **kwargs):
        return (
            json.dumps(proposal(json.loads(body["messages"][1]["content"]), invalid_source=True)),
            "simulated",
            {},
        )

    monkeypatch.setattr(global_sessions, "request_model_content", invalid)
    s = await resolve(api, await start(api))
    assert s["global_compilation"]["proposal_count"] == 0
    assert all(not f["selectable"] for f in s["global_compilation"]["fields"])
    assert s["summary"]["resolved"] == 0


@pytest.mark.anyio
async def test_isolation_and_legacy_preserved(api, provider):
    s = await start(api)
    legacy = sessions.create_session("alpha", s["form_id"], s["conversation_id"])
    assert legacy["engine"] == "legacy" and legacy["id"] != s["id"]
    assert len(legacy["fields"]) == 2
    for route in ["", "/global-report"]:
        assert (
            await api.get(f"/api/projects/beta/compilation-sessions/{s['id']}{route}")
        ).status_code == 404
    assert (
        await api.post(
            f"/api/projects/beta/compilation-sessions/{s['id']}/resolve",
            json={"version": s["version"]},
        )
    ).status_code == 404
    assert len(provider) == 0


@pytest.mark.anyio
async def test_failure_persists_without_automatic_retry(api, provider, monkeypatch):
    async def fail(*args, **kwargs):
        raise GenerationError("simulated failure")

    monkeypatch.setattr(global_sessions, "request_model_content", fail)
    s = await start(api)
    response = await api.post(f"{BASE}/{s['id']}/resolve", json={"version": s["version"]})
    assert response.status_code == 502
    s = (await api.get(f"{BASE}/{s['id']}")).json()
    assert s["status"] == "FAILED" and s["global_compilation"]["attempts"] == 1
    assert not s["chat"]["auto_continue"] and not s["last_generation"]
    assert list(get_storage_path().rglob("metrics.json"))
    resumed = global_sessions.control_session("alpha", s["id"], s["version"], "resume")
    assert resumed["id"] == s["id"] and resumed["chat"]["auto_continue"]


@pytest.mark.anyio
async def test_pause_wins_over_in_flight_proposal(api, provider, monkeypatch):
    entered, finish = asyncio.Event(), asyncio.Event()

    async def wait(client, settings, body, timeout, **kwargs):
        entered.set()
        await finish.wait()
        return json.dumps(proposal(json.loads(body["messages"][1]["content"]))), "simulated", {}

    monkeypatch.setattr(global_sessions, "request_model_content", wait)
    s = await start(api)
    task = asyncio.create_task(
        api.post(f"{BASE}/{s['id']}/resolve", json={"version": s["version"]})
    )
    await asyncio.wait_for(entered.wait(), 5)
    paused = global_sessions.control_session("alpha", s["id"], s["version"], "pause")
    finish.set()
    response = await task
    assert response.status_code == 409
    after = (await api.get(f"{BASE}/{s['id']}")).json()
    assert after["version"] == paused["version"] and after["chat"]["paused_by_user"]
    assert after["global_compilation"]["proposal_count"] == 0 and after["last_generation"] is None
