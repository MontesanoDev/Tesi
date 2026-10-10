"""Structural contract tests; no simulated or real AI provider is involved."""

import json
from hashlib import sha256
from io import BytesIO
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

import pytest
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from lxml import etree

from app.compilation_catalog import W14, build_compilation_catalog, document_blocks
from app.docx_templates import DocumentInputError, inspect_docx

DEMO = Path(__file__).resolve().parents[2] / "demo-documents"


def encoded(doc):
    stream = BytesIO()
    doc.save(stream)
    return stream.getvalue()


def catalog(doc):
    layout = inspect_docx(encoded(doc), require_candidates=False)
    return build_compilation_catalog(layout)


def native_checkbox(paragraph, *, legacy=True, checked=False, default=None):
    if legacy:
        begin = OxmlElement("w:fldChar")
        begin.set(qn("w:fldCharType"), "begin")
        data = etree.SubElement(begin, qn("w:ffData"))
        control = etree.SubElement(data, qn("w:checkBox"))
        if default is not None:
            etree.SubElement(control, qn("w:default")).set(qn("w:val"), str(int(default)))
        if checked is not None:
            etree.SubElement(control, qn("w:checked")).set(qn("w:val"), str(int(checked)))
        paragraph.add_run()._r.append(begin)
        end = OxmlElement("w:fldChar")
        end.set(qn("w:fldCharType"), "end")
        paragraph.add_run()._r.append(end)
    else:
        sdt = OxmlElement("w:sdt")
        properties = etree.SubElement(sdt, qn("w:sdtPr"))
        control = etree.SubElement(properties, f"{{{W14}}}checkbox")
        etree.SubElement(control, f"{{{W14}}}checked").set(
            f"{{{W14}}}val", str(int(checked)),
        )
        content = etree.SubElement(sdt, qn("w:sdtContent"))
        run = etree.SubElement(content, qn("w:r"))
        etree.SubElement(run, qn("w:t")).text = "☑" if checked else "☐"
        paragraph._p.append(sdt)
    return control


def paragraphs(result):
    return [p for block in result["blocks"] for p in block["paragraphs"]]


def test_complete_ordered_context_once_with_two_same_named_fields_in_different_sections():
    doc = Document()
    doc.add_heading("Categoria Alfa", level=1)
    doc.add_paragraph("Solo quando si applica la categoria Alfa.")
    doc.add_paragraph("Nome e cognome: __________")
    doc.add_heading("Categoria Beta", level=1)
    doc.add_heading("Referente", level=2)
    doc.add_paragraph("Nome e cognome: __________")
    result = catalog(doc)
    assert [p["text"] for p in paragraphs(result)] == [p.text for p in doc.paragraphs]
    assert result["candidates"]["p2.s0"]["section_ids"] == ["paragraph:0"]
    assert result["candidates"]["p5.s0"]["section_ids"] == ["paragraph:3", "paragraph:4"]
    assert json.dumps(result, ensure_ascii=False).count("Solo quando si applica") == 1
    assert all("status" not in field and "value" not in field
               for field in result["candidates"].values())
    assert result["role"] == "form"


def test_unstyled_conditional_instructions_remain_literal_not_guessed_sections():
    doc = Document()
    doc.add_paragraph("Compilare solo se il richiedente appartiene alla categoria Gamma")
    doc.add_paragraph("Identificativo: __________")
    result = catalog(doc)
    assert paragraphs(result)[0]["text"] == doc.paragraphs[0].text
    assert result["candidates"]["p1.s0"]["section_ids"] == []
    assert "applicability" not in json.dumps(result)


def test_hierarchy_uses_inherited_outline_levels_instead_of_style_names():
    doc = Document()
    custom = doc.styles.add_style("Custom unrelated name", WD_STYLE_TYPE.PARAGRAPH)
    custom.base_style = doc.styles["Heading 2"]
    doc.add_heading("Contesto principale", level=1)
    doc.add_paragraph("Contesto dipendente", style=custom)
    doc.add_paragraph("Identificativo: __________")
    result = catalog(doc)
    assert paragraphs(result)[1]["outline_level"] == 1
    assert result["candidates"]["p2.s0"]["section_ids"] == ["paragraph:0", "paragraph:1"]


def test_tab_stops_are_layout_only_but_run_tabs_and_line_breaks_are_text():
    doc = Document()
    p = doc.add_paragraph("Etichetta")
    properties = p._p.get_or_add_pPr()
    tabs = etree.SubElement(properties, qn("w:tabs"))
    tab_stop = etree.SubElement(tabs, qn("w:tab"))
    tab_stop.set(qn("w:val"), "left")
    tab_stop.set(qn("w:pos"), "720")
    p.add_run("\tValore\nAltro: __________")
    assert paragraphs(catalog(doc))[0]["text"] == "Etichetta\tValore\nAltro: __________"


@pytest.mark.parametrize("legacy", [True, False])
@pytest.mark.parametrize("checked", [True, False])
def test_checkbox_only_native_forms_are_catalogued_without_changing_v1_support(legacy, checked):
    doc = Document()
    p = doc.add_paragraph("Opzione: ")
    native_checkbox(p, legacy=legacy, checked=checked)
    data = encoded(doc)
    with pytest.raises(DocumentInputError, match="Nessun campo supportato"):
        inspect_docx(data)
    layout = inspect_docx(data, require_candidates=False)
    result = build_compilation_catalog(layout)
    assert not layout.candidate_ids
    assert list(result["candidates"]) == ["p0.cb0"]
    control = result["candidates"]["p0.cb0"]
    assert control["initial_checked"] is checked
    assert control["type"] == "checkbox"
    assert control["context_id"] == "paragraph:0"
    assert result["checkbox_write_supported"] is False
    assert sha256(data).hexdigest() == result["template_sha256"]


def test_checkbox_positions_distinguish_two_choices_in_one_paragraph():
    doc = Document()
    p = doc.add_paragraph()
    native_checkbox(p)
    p.add_run(" Sì ")
    native_checkbox(p)
    p.add_run(" No")
    result = catalog(doc)
    assert paragraphs(result)[0]["controls"] == [
        {"id": "p0.cb0", "offset": 0}, {"id": "p0.cb1", "offset": 4},
    ]
    assert len(result["candidates"]) == 2


def test_block_level_content_control_maps_to_its_content_paragraph():
    doc = Document()
    p = doc.add_paragraph()
    control = native_checkbox(p, legacy=False, checked=True)
    sdt = control.getparent().getparent()
    content = sdt.find(qn("w:sdtContent"))
    run = content[0]
    content.remove(run)
    inner = etree.SubElement(content, qn("w:p"))
    inner.append(run)
    p._p.addprevious(sdt)
    p._p.getparent().remove(p._p)
    result = catalog(doc)
    assert list(result["candidates"]) == ["p0.cb0"]
    assert result["candidates"]["p0.cb0"]["initial_checked"] is True
    assert paragraphs(result)[0]["controls"] == [{"id": "p0.cb0", "offset": 0}]


def test_legacy_default_checked_state_is_preserved_as_original_not_source():
    doc = Document()
    native_checkbox(doc.add_paragraph(), checked=None, default=True)
    field = catalog(doc)["candidates"]["p0.cb0"]
    assert field["initial_checked"] is True
    assert "provenance" not in field


def test_unknown_checkbox_state_fails_without_guessing_negative():
    doc = Document()
    control = native_checkbox(doc.add_paragraph())
    control.find(qn("w:checked")).set(qn("w:val"), "unknown")
    with pytest.raises(DocumentInputError, match="Stato checkbox"):
        catalog(doc)


def test_rows_link_checkbox_to_label_in_another_cell_and_preserve_text_candidates():
    doc = Document()
    table = doc.add_table(rows=2, cols=2)
    native_checkbox(table.cell(0, 0).paragraphs[0])
    table.cell(0, 1).text = "Categoria Delta"
    table.cell(1, 0).text = "Denominazione"
    layout = inspect_docx(encoded(doc))
    result = build_compilation_catalog(layout)
    assert {key for key, field in result["candidates"].items() if field["type"] == "text"} == (
        layout.candidate_ids
    )
    row = result["blocks"][0]["tables"][0]["rows"][0]
    assert row[0]["paragraph_ids"] == [result["candidates"]["p0.cb0"]["context_id"]]
    label_id = row[1]["paragraph_ids"][0]
    assert next(p["text"] for p in paragraphs(result) if p["id"] == label_id) == "Categoria Delta"


def test_existing_document_map_projection_preserves_ids_boundaries_and_candidates():
    doc = Document()
    doc.add_paragraph("Istruzione generale")
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Denominazione"
    doc.add_paragraph("Altro dato: __________")
    layout = inspect_docx(encoded(doc))
    assert document_blocks(layout) == [
        {"id": "paragraph:0", "text": "Istruzione generale", "candidate_ids": []},
        {"id": "paragraph:1", "text": "Denominazione", "candidate_ids": ["t0.r0.c1"]},
        {"id": "paragraph:3", "text": "Altro dato: __________", "candidate_ids": ["p3.s0"]},
    ]


def test_nested_tables_and_physical_merges_preserve_context_relationships():
    doc = Document()
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).merge(table.cell(0, 1)).text = "Intestazione comune"
    nested = table.cell(1, 1).add_table(rows=1, cols=2)
    nested.cell(0, 0).text = "Campo interno"
    result = catalog(doc)
    outer, inner = result["blocks"][0]["tables"]
    assert outer["rows"][0][0]["grid_span"] == "2"
    assert len(outer["rows"][0]) == 1
    assert outer["rows"][1][1]["table_ids"] == [inner["id"]]
    outer_ids = outer["rows"][1][1]["paragraph_ids"]
    inner_ids = inner["rows"][0][0]["paragraph_ids"]
    assert not set(outer_ids) & set(inner_ids)
    assert "t1.r0.c1" in result["candidates"]


def test_signature_candidates_remain_write_protected():
    doc = Document()
    doc.add_paragraph("Firma: __________")
    result = catalog(doc)
    assert result["candidates"]["p0.s0"]["write_protected"] is True


def test_headers_and_footers_are_form_context_without_writable_ids():
    doc = Document()
    doc.add_paragraph("Identificativo: __________")
    doc.sections[0].header.paragraphs[0].text = "Istruzione in intestazione"
    native_checkbox(doc.sections[0].footer.paragraphs[0])
    result = catalog(doc)
    assert len(result["candidates"]) == 1
    assert result["read_only_parts"] == [
        {"part": "word/footer1.xml", "paragraphs": [], "read_only_checkbox_count": 1},
        {"part": "word/header1.xml", "paragraphs": ["Istruzione in intestazione"],
         "read_only_checkbox_count": 0},
    ]


def test_footnote_reference_keeps_its_literal_instruction_and_id_once():
    doc = Document()
    p = doc.add_paragraph("Identificativo: __________")
    reference = etree.SubElement(p.add_run()._r, qn("w:footnoteReference"))
    reference.set(qn("w:id"), "7")
    notes = etree.Element(qn("w:footnotes"))
    note = etree.SubElement(notes, qn("w:footnote"))
    note.set(qn("w:id"), "7")
    note_p = etree.SubElement(note, qn("w:p"))
    etree.SubElement(etree.SubElement(note_p, qn("w:r")), qn("w:t")).text = (
        "Istruzione specifica della nota"
    )
    data = encoded(doc)
    stream = BytesIO()
    with ZipFile(BytesIO(data)) as original, ZipFile(stream, "w", ZIP_DEFLATED) as output:
        for entry in original.infolist():
            output.writestr(entry, original.read(entry))
        output.writestr("word/footnotes.xml", etree.tostring(notes))
    layout = inspect_docx(stream.getvalue(), require_candidates=False)
    result = build_compilation_catalog(layout)
    assert paragraphs(result)[0]["references"] == [
        {"part": "word/footnotes.xml", "id": "7", "offset": len(p.text)},
    ]
    assert result["read_only_parts"][0]["notes"] == [
        {"id": "7", "paragraphs": ["Istruzione specifica della nota"]},
    ]
    assert json.dumps(result, ensure_ascii=False).count("Istruzione specifica della nota") == 1


@pytest.mark.parametrize("relative,count", [
    ("bandi/catanzaro-dl-cse/modello/domanda-partecipazione.docx", 28),
    ("bandi/minervino-elenco-sia/originali/domanda-iscrizione.docx", 13),
])
def test_real_native_checkbox_formats_have_stable_ids_without_double_counting(relative, count):
    original = (DEMO / relative).read_bytes()
    layout = inspect_docx(original, require_candidates=False)
    xml_before = etree.tostring(layout.document.element)
    result = build_compilation_catalog(layout)
    assert sum(c["type"] == "checkbox" for c in result["candidates"].values()) == count
    assert sum(c["type"] == "text" for c in result["candidates"].values()) == len(
        layout.candidate_ids,
    )
    assert result == build_compilation_catalog(inspect_docx(original, require_candidates=False))
    assert etree.tostring(layout.document.element) == xml_before
    assert layout.original == original == (DEMO / relative).read_bytes()


def test_over_budget_never_silently_omits_context(monkeypatch):
    from app import compilation_catalog

    doc = Document()
    doc.add_paragraph("Identificativo: __________")
    monkeypatch.setattr(compilation_catalog, "MAX_CATALOG_CHARACTERS", 5)
    with pytest.raises(DocumentInputError, match="nessun testo omesso"):
        catalog(doc)


def test_combined_text_and_checkbox_budget_is_bounded():
    doc = Document()
    p = doc.add_paragraph("Identificativo: __________")
    for _ in range(400):
        native_checkbox(p)
    # Complex content protects the placeholder. Use a separate writable slot.
    doc.add_paragraph("Altro identificativo: __________")
    with pytest.raises(DocumentInputError, match="400 candidate"):
        catalog(doc)
