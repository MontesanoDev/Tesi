"""Conditional FORM caption omitted from otherwise valid JSON: no implied applicability."""
import json
from io import BytesIO

import pytest
import test_compilation_resilient_segments as rev2
from docx import Document
from test_compilation_global_sessions import api, provider

from app import compilation_global_dialogue as dialogue
from app import compilation_global_sessions as adapter
from app import compilation_scope_review as scope
from app import compilation_table_scope as table_scope
from app.compilation_chat import ClarificationReply
from app.compilation_segments import applicability_states
from app.docx_templates import inspect_docx

__all__ = ["api", "provider"]


def form(*, instruction="2.a) da compilare in caso di ADESIONE AL PROGETTO",
         between=None, bold=True):
    doc = Document()
    doc.add_table(rows=1, cols=1).cell(0, 0).paragraphs[0].add_run(instruction).bold = bold
    if between is not None:
        doc.add_paragraph(between)
    doc.add_table(rows=1, cols=1).cell(0, 0).text = "Dato riservato: ____________"
    doc.add_paragraph("Dato indipendente: ____________")
    stream = BytesIO()
    doc.save(stream)
    return stream.getvalue()


def writer(inp, ids):
    targets = rev2.dato_ids(inp, ids)
    return json.dumps(rev2.envelope(inp, {c: rev2.slot(inp, c) for c in targets}))


@pytest.mark.anyio
@pytest.mark.parametrize("flag", [None, "1"])
async def test_valid_json_cannot_silently_omit_the_conditional_caption(
    api, provider, monkeypatch, flag,
):
    monkeypatch.setenv("MAPI_GENERATION_REVISION", "3")
    monkeypatch.setenv("MAPI_DIALOGUE_REVISION", "1")
    if flag:
        monkeypatch.setenv("MAPI_TABLE_SCOPE_REVISION", flag)
    else:
        monkeypatch.delenv("MAPI_TABLE_SCOPE_REVISION", raising=False)
    s = await rev2.started(api, monkeypatch, form(), rev2.Model({0: [writer], 1: [writer]}))
    response = await rev2.analyze(api, s)
    assert response.status_code == 200, response.text
    s = response.json()
    data = rev2.state_of(s)["global_compilation"]
    conditional = next(cid for cid, c in data["catalog"]["candidates"].items()
                       if c.get("location", {}).get("table_id") == "table:1")
    selected, _, _ = dialogue.export_selection(s)
    if not flag:
        assert conditional in selected
        return
    assert conditional not in selected
    assert (await rev2.draft(api, s, [*selected, conditional])).status_code == 409
    assert selected
    assert (await rev2.draft(api, s, selected)).status_code == 200


def data_for(*, condition=True):
    catalog = rev2.model_form(form())
    cid = next(c for c, x in catalog["candidates"].items()
               if x.get("location", {}).get("table_id") == "table:1")
    data = {"generation_revision": 3, "table_scope_revision": 1, "catalog": catalog,
            "dialogue": {"answers": {}}, "audit": {"fields": [
                {"candidate_id": c, "conditions": [], "validation": "GROUNDED_PROPOSAL"}
                for c in catalog["candidates"]], "conditions": {}}}
    if condition:
        data["audit"]["conditions"]["c"] = {"result": "unknown", "applies_to": [],
            "form_evidence": {"paragraph_id": "paragraph:0", "quote":
                catalog["blocks"][0]["paragraphs"][0]["text"]}}
    return data, cid


def confirm(data, cid, *, key="condition:c", action="CONDITION_TRUE"):
    record = {"action": action, "applies": action == "CONDITION_TRUE", "user_quote": "Sì",
              "affected_fields": [cid], "provenance": "USER"}
    record["scope_validation"] = scope.proof(scope.target(data, key, [cid]), record)
    data["dialogue"]["answers"][key] = record


@pytest.mark.parametrize("kwargs", [
    {"instruction": "DATI GENERALI"}, {"bold": False},
    {"instruction": "Esempio: da compilare in caso di ADESIONE"},
    {"between": "Altra sezione autonoma"},
])
def test_no_inferred_scope_through_unrecognized_or_interrupted_structure(kwargs):
    assert table_scope.scopes(rev2.model_form(form(**kwargs))) == []


@pytest.mark.parametrize("instruction", [
    "da compilare in caso di ADESIONE", "7.b) da compilare nel caso di AFFIDAMENTO",
    "4) DA COMPILARE IN CASO DI PARTECIPAZIONE CONGIUNTA",
])
def test_explicit_instruction_is_not_tied_to_entity_names_or_a_tender(instruction):
    regions = table_scope.scopes(rev2.model_form(form(instruction=instruction, between="")))
    assert len(regions) == 1 and regions[0]["table_id"] == "table:1"


def test_known_condition_is_propagated_without_becoming_a_true_condition():
    data, cid = data_for()
    table_scope.apply(data["audit"], data["catalog"])
    assert data["audit"]["conditions"]["c"]["result"] == "unknown"
    assert data["audit"]["conditions"]["c"]["applies_to"] == [cid]
    assert cid in scope.blocked_source_fields(data)
    confirm(data, cid)
    assert cid not in scope.blocked_source_fields(data)
    assert applicability_states(data)[cid]["state"] == "VERIFICATA"
    data["audit"]["conditions"]["c"]["result"] = "false"
    assert cid in scope.blocked_source_fields(data)


@pytest.mark.parametrize("action", ["SKIP", "UNKNOWN", "CONDITION_FALSE"])
def test_non_authorizing_answers_do_not_unlock_a_branch(action):
    data, cid = data_for()
    table_scope.apply(data["audit"], data["catalog"])
    confirm(data, cid, action=action)
    assert cid in scope.blocked_source_fields(data)


def test_missing_condition_is_unknown_even_with_another_verified_user_answer():
    data, cid = data_for(condition=False)
    confirm(data, cid, key="condition:another")
    table_scope.apply(data["audit"], data["catalog"])
    assert data["audit"]["conditions"] == {}  # Never invent a condition.
    assert cid in scope.blocked_source_fields(data)
    assert cid in scope.blocked_user_fields(data)
    assert applicability_states(data)[cid]["state"] == "NON_ACCERTATA"


def test_a_literal_user_value_and_a_scope_receipt_are_not_an_omitted_branch_decision():
    data, cid = data_for(condition=False)
    data["dialogue"]["user_values"] = {cid: {"value": "ABC123", "user_quote": "ABC123"}}
    record = {"action": "VALUE", "value": "ABC123", "user_quote": "ABC123",
              "provenance": "USER", "affected_fields": [cid]}
    record["scope_validation"] = scope.proof(scope.target(data, "field:" + cid, [cid]), record)
    data["dialogue"]["answers"]["field:" + cid] = record
    assert cid in scope.blocked_user_fields(data)


def test_rejected_foreign_or_detached_conditions_do_not_count_as_coverage():
    for change in ("rejected", "foreign", "detached"):
        data, cid = data_for()
        if change == "rejected":
            data["audit"]["conditions"]["c"]["validation_errors"] = ["bad FORM quote"]
        if change == "foreign":
            data["audit"]["conditions"]["c"]["form_evidence"]["paragraph_id"] = "paragraph:2"
        if change != "detached":
            table_scope.apply(data["audit"], data["catalog"])
        assert cid in table_scope.blocked_fields(data)


def test_new_profile_requires_both_flags_and_existing_sessions_are_not_upgraded(monkeypatch):
    monkeypatch.setenv("MAPI_TABLE_SCOPE_REVISION", "1")
    for generation in ("0", "1", "2"):
        monkeypatch.setenv("MAPI_GENERATION_REVISION", generation)
        assert "table_scope_revision" not in adapter.initial_state(inspect_docx(form()))
    monkeypatch.setenv("MAPI_GENERATION_REVISION", "3")
    assert adapter.initial_state(inspect_docx(form()))["table_scope_revision"] == 1
    data, cid = data_for(condition=False)
    data.pop("table_scope_revision")
    assert not table_scope.blocked_fields(data)
    data["table_scope_revision"] = 1
    data["generation_revision"] = 2
    assert not table_scope.blocked_fields(data)


def test_projection_is_idempotent_and_missing_projection_does_not_authorize():
    data, cid = data_for()
    table_scope.apply(data["audit"], data["catalog"])
    before = json.dumps(data["audit"], sort_keys=True)
    table_scope.apply(data["audit"], data["catalog"])
    assert json.dumps(data["audit"], sort_keys=True) == before
    data, cid = data_for(condition=False)
    assert "table_scope" not in data["audit"]
    assert cid in table_scope.blocked_fields(data)


@pytest.mark.parametrize("ambiguous", ["editable", "nested", "outside"])
def test_caption_must_be_fixed_simple_and_contain_the_instruction(ambiguous):
    catalog = rev2.model_form(form())
    if ambiguous == "editable":
        catalog["candidates"]["caption"] = {"location": {"table_id": "table:0"}}
    elif ambiguous == "nested":
        catalog["blocks"][0]["tables"][0]["rows"][0][0]["table_ids"] = ["table:99"]
    else:
        catalog["blocks"][0]["tables"][0]["rows"][0][0]["paragraph_ids"] = []
    assert not table_scope.scopes(catalog)


@pytest.mark.anyio
async def test_actual_export_after_verified_condition_keeps_all_other_gates(
    api, provider, monkeypatch,
):
    monkeypatch.setenv("MAPI_GENERATION_REVISION", "3")
    monkeypatch.setenv("MAPI_DIALOGUE_REVISION", "1")
    monkeypatch.setenv("MAPI_TABLE_SCOPE_REVISION", "1")

    def with_condition(inp, ids):
        raw = json.loads(writer(inp, ids))
        caption = inp["form"]["blocks"][0]["paragraphs"][0]
        raw["conditions"]["c"] = rev2.condition(inp, caption["id"], caption["text"])
        return json.dumps(raw)

    s = await rev2.started(api, monkeypatch, form(),
                          rev2.Model({0: [with_condition], 1: [with_condition]}))
    response = await rev2.analyze(api, s)
    assert response.status_code == 200, response.text
    s = response.json()
    data = rev2.state_of(s)["global_compilation"]
    target = next(q for q in dialogue.question_pool(data) if q["key"].startswith("condition:"))
    assert target["field_ids"] == ["p1.s0"]
    question = next(q for q in s["chat"]["question"]["slots"] if q["key"] == target["key"])
    reply = ClarificationReply(slot=question["slot"], action="CONDITION_TRUE",
                               user_quote="Confermo l'adesione al progetto")
    proof = scope.proof(scope.target(data, target["key"], target["field_ids"]), reply)
    s, errors = dialogue.apply_replies("alpha", s["id"], s["version"], [reply], reply.user_quote,
                                       scope_verifications={reply.slot: proof})
    assert not errors
    s = (await rev2.analyze(api, s)).json()
    ids, conditions, _ = dialogue.export_selection(s)
    assert "p1.s0" in ids
    exported = await rev2.draft(api, s, ids, conditions=conditions)
    assert exported.status_code == 200, exported.text
