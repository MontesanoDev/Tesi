"""Software grounding checks, not a semantic verdict on real model responses."""

import copy
import json
from io import BytesIO

import pytest
from docx import Document
from pydantic import ValidationError

from app.compilation_global import build_global_input, validate_global_proposal
from app.compilation_global_contract import AlternativeAssignment
from app.docx_templates import inspect_docx


@pytest.fixture
def case():
    doc = Document()
    doc.add_heading("Dati dell'organizzazione", 1)
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Recapito corrispondenza"
    table.cell(1, 0).text = "Codice pratica"
    stream = BytesIO()
    doc.save(stream)
    layout = inspect_docx(stream.getvalue())
    source = {"chunk_id": -7, "file_id": -3, "chunk_index": 0,
              "source_name": "profilo.txt", "role": "source", "scope": "global",
              "category": "company", "project_id": None,
              "content": "Alba: recapito Via Luna 4; altro recapito Via Sole 8; "
                         "codice pratica AZ-98765."}
    facts = {}
    for key, value, prop, aspect in [
        ("a", "Via Luna 4", "recapito corrispondenza", "location"),
        ("b", "Via Sole 8", "recapito corrispondenza", "location"),
        ("c", "AZ-98765", "codice pratica", "identity"),
    ]:
        facts[key] = {"subject": "Alba", "entity": "company", "aspect": aspect,
                      "property": prop, "value": value,
                      "evidence": [{"chunk_id": -7, "quote": source["content"]}]}
    fields = {
        "t0.r0.c1": {"label": "Recapito corrispondenza", "kind": "data",
                     "disposition": "review", "property": "recapito corrispondenza",
                     "review_kind": "ambiguous", "alternative_fact_ids": ["a", "b"],
                     "conditions": [], "reason": "Nessuna priorità documentata",
                     "form": {"paragraph_id": "paragraph:1", "quote": "Recapito corrispondenza"}},
        "t0.r1.c1": {"label": "Codice pratica", "kind": "data", "disposition": "value",
                     "property": "codice pratica", "fact_id": "c", "conditions": [],
                     "form": {"paragraph_id": "paragraph:3", "quote": "Codice pratica"}},
    }
    return layout, [source], {"version": 6, "conditions": {}, "fields": fields, "facts": facts}


def audit(case):
    return validate_global_proposal(json.dumps(case[2]), case[0], "alpha", case[1])


def field(result):
    return result["fields"][0]


@pytest.mark.parametrize("kind", ["ambiguous", "conflicting"])
def test_alternatives_have_individual_source_provenance_and_no_chosen_value(case, kind):
    case[2]["fields"]["t0.r0.c1"]["review_kind"] = kind
    result = audit(case)
    reviewed = field(result)
    assert reviewed["validation"] == "REVIEW_REQUIRED"
    assert reviewed["review_kind"] == kind
    assert reviewed["value"] is reviewed["written_value"] is reviewed["fact_id"] is None
    assert [a["value"] for a in reviewed["alternatives"]] == ["Via Luna 4", "Via Sole 8"]
    for alternative in reviewed["alternatives"]:
        proof = alternative["source_evidence"][0]
        assert (proof["chunk_id"], proof["file_id"], proof["role"], proof["category"]) == (
            -7, -3, "source", "company")
        assert proof["scope"] == "global" and proof["project_id"] is None
        assert alternative["value"] in proof["quote"]
        assert proof["content_sha256"]
    assert result["fields"][1]["validation"] == "GROUNDED_PROPOSAL"
    assert result["resolved"] == result["writes"] == 0
    assert all(f["semantic_review_required"] for f in result["fields"])


@pytest.mark.parametrize("change", [{"subject": "Altro soggetto"}, {"entity": "person"},
    {"aspect": "identity"}, {"property": "sede legale"}, {"value": "Via Luna 4"},
    {"value": "Via inventata"}, {"evidence": []},
    {"evidence": [{"chunk_id": -99, "quote": "Via Sole 8"}]}])
def test_invalid_or_unrelated_alternative_is_local_and_cannot_be_chosen(case, change):
    case[2]["facts"]["b"].update(change)
    result = audit(case)
    assert field(result)["validation"] == "REJECTED"
    assert field(result)["value"] is field(result)["written_value"] is None
    assert result["fields"][1]["validation"] == "GROUNDED_PROPOSAL"


@pytest.mark.parametrize("change", [
    {"alternative_fact_ids": ["a"]}, {"alternative_fact_ids": ["a", "a"]},
    {"alternative_fact_ids": ["a", "missing"]}, {"alternative_fact_ids": ["a", "b", "a", "b", "a"]},
    {"property": "altro requisito"}, {"fact_id": "a"}, {"choice": True}, {"value": "Via Luna 4"},
    {"form": {"paragraph_id": "paragraph:3", "quote": "Codice pratica"}},
    {"form": {"paragraph_id": "paragraph:1", "quote": "Requisito inventato"}},
    {"conditions": ["inesistente"]}, {"kind": "signature"},
])
def test_review_cannot_bypass_existing_anchor_condition_or_no_value_gates(case, change):
    case[2]["fields"]["t0.r0.c1"].update(change)
    result = audit(case)
    assert field(result)["validation"] == "REJECTED"
    assert field(result)["written_value"] is None
    assert result["fields"][1]["validation"] == "GROUNDED_PROPOSAL"


def test_known_alternatives_block_silent_selection_even_if_model_emits_value(case):
    first = case[2]["fields"]["t0.r0.c1"]
    first.update(disposition="value", fact_id="a")
    first.pop("alternative_fact_ids")
    first.pop("review_kind")
    # The second fact is referenced by a separate grounded condition. It cannot
    # disappear just because the value assignment itself chose only the first.
    case[2]["conditions"]["context"] = {
        "form": {"paragraph_id": "paragraph:3", "quote": "Codice pratica"},
        "aspect": "location", "operator": "semantic", "fact_ids": ["b"],
        "reason": "Contesto da verificare", "scope": []}
    result = audit(case)
    reviewed = field(result)
    assert reviewed["proposal"]["fact_id"] == "a"
    assert reviewed["validation"] == "REVIEW_REQUIRED"
    assert reviewed["review_kind"] == "unselected_alternatives"
    assert reviewed["value"] is reviewed["written_value"] is None
    assert {a["fact_id"] for a in reviewed["alternatives"]} == {"a", "b"}


def test_distinct_properties_do_not_contaminate_single_supported_value(case):
    case[2]["fields"]["t0.r0.c1"] = {
        "label": "Recapito corrispondenza", "kind": "data", "disposition": "value",
        "fact_id": "a", "property": "recapito corrispondenza", "conditions": [],
        "form": {"paragraph_id": "paragraph:1", "quote": "Recapito corrispondenza"}}
    result = audit(case)
    assert field(result)["validation"] == "GROUNDED_PROPOSAL"
    assert "b" in result["unused_fact_ids"]
    assert result["fields"][1]["validation"] == "GROUNDED_PROPOSAL"


def test_unknown_applicability_remains_unknown_in_review(case):
    case[2]["conditions"]["unknown"] = {
        "form": {"paragraph_id": "paragraph:1", "quote": "Recapito corrispondenza"},
        "aspect": "location", "operator": "semantic", "fact_ids": [],
        "reason": "Applicabilità da verificare", "scope": []}
    result = audit(case)
    assert field(result)["validation"] == "REVIEW_REQUIRED"
    assert field(result)["condition_checks"]["unknown"]["result"] == "unknown"
    assert field(result)["written_value"] is None


@pytest.mark.parametrize("role", ["form", "user"])
def test_alternatives_require_real_authorized_source_role(case, role):
    case[1][0]["role"] = role
    with pytest.raises(ValueError, match="Ruolo SOURCE"):
        audit(case)


def test_source_in_another_project_cannot_enter_alternatives(case):
    case[1][0].update(chunk_id=7, file_id=3, scope="project:beta", project_id="beta", category=None)
    with pytest.raises(ValueError, match="fuori progetto"):
        audit(case)


def test_general_kb_cannot_prove_company_alternatives(case):
    case[1][0]["category"] = "general"
    assert all(f["validation"] == "REJECTED" for f in audit(case)["fields"])


def test_schema_is_bounded_and_review_has_no_writable_value(case):
    data = build_global_input(case[0], "alpha", case[1])
    schema = data["schema_output"]
    assert schema["properties"]["version"]["const"] == 6
    assert schema["properties"]["facts"]["maxProperties"] == 128
    assert schema["properties"]["fields"]["maxProperties"] == 400
    assert schema["$defs"]["AlternativeAssignment"]["properties"]["alternative_fact_ids"] == {
        "items": {"maxLength": 80, "minLength": 1, "type": "string"},
        "maxItems": 4, "minItems": 2, "title": "Alternative Fact Ids", "type": "array"}
    item = copy.deepcopy(case[2]["fields"]["t0.r0.c1"])
    item["value"] = "Via Luna 4"
    with pytest.raises(ValidationError):
        AlternativeAssignment.model_validate(item)


def test_v5_replay_does_not_silently_accept_new_review_records(case):
    case[2]["version"] = 5
    result = audit(case)
    assert result["version"] == 5
    assert field(result)["validation"] == "REJECTED"
    assert result["fields"][1]["validation"] == "GROUNDED_PROPOSAL"
