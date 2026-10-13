"""Bounded reading/attribution tests, not real model semantic evaluation."""

import copy
import json
from io import BytesIO

import pytest
from docx import Document

from app.compilation_catalog import MAX_CONTINUATION_PARAGRAPHS, build_compilation_catalog
from app.compilation_global import build_global_input, validate_global_proposal
from app.compilation_role_records import record_support
from app.docx_templates import inspect_docx

ROLE = "Responsabile della verifica"


def layout_for(doc):
    stream = BytesIO()
    doc.save(stream)
    return inspect_docx(stream.getvalue(), require_candidates=False)


@pytest.fixture
def case():
    doc = Document()
    doc.add_paragraph(ROLE)
    for label in ["Nome completo", "Qualifica", "Numero registrazione"]:
        doc.add_paragraph(label + ": ________")
    source = {"chunk_id": -7, "file_id": -3, "chunk_index": 2,
              "source_name": "profilo.txt", "role": "source", "scope": "global",
              "category": "company", "project_id": None,
              "content": "### Dr. Ada Neri - responsabile della verifica\n"
              "Nome completo: Dr. Ada Neri.\nQualifica: Perito.\nNumero registrazione: 6234.\n"
              "### Bea Mori - altro incarico\nNumero registrazione: 9999."}
    fields, facts = {}, {}
    for index, (label, value, quote) in enumerate([
        ("Nome completo", "Dr. Ada Neri", "Nome completo: Dr. Ada Neri."),
        ("Qualifica", "Perito", "Qualifica: Perito."),
        ("Numero registrazione", "6234", "Numero registrazione: 6234."),
    ]):
        cid, ref = f"p{index + 1}.s0", f"f{index}"
        facts[ref] = {"subject": "Ada Neri", "entity": "person", "aspect": "professional",
                      "property": label, "value": value,
                      "evidence": [{"chunk_id": -7, "quote": quote}]}
        fields[cid] = {"label": label, "kind": "data", "disposition": "value",
                       "form": {"paragraph_id": f"paragraph:{index + 1}",
                                "quote": f"{label}: [[{cid}]]"}, "conditions": [],
                       "fact_id": ref, "property": label, "binding_id": "b"}
    binding = {"form": {"paragraph_id": "paragraph:0", "quote": ROLE},
               "required_role": ROLE, "source_role": ROLE,
               "evidence": [{"chunk_id": -7, "quote": source["content"].splitlines()[0]}]}
    return layout_for(doc), [source], {"version": 10, "fields": fields, "facts": facts,
                                      "conditions": {}, "bindings": {"b": binding}}


def audit(case):
    return validate_global_proposal(json.dumps(case[2]), case[0], "alpha", case[1],
                                    expected_version=10)


def test_continuous_fields_share_literal_origin_and_keep_separate_destinations(case):
    data = build_global_input(case[0], "alpha", case[1])
    assert data["schema_output"]["properties"]["version"]["const"] == 12
    for cid, candidate in data["form"]["candidates"].items():
        assert candidate["continuation_origin_paragraph_id"] == "paragraph:0"
        assert "paragraph:0" not in candidate["anchor_paragraph_ids"]
        assert candidate["anchor_paragraph_ids"] == [candidate["context_id"]]
    result = audit(case)
    assert result["binding_checks"] == {"literal_match": 3}
    assert [f["value"] for f in result["fields"]] == ["Dr. Ada Neri", "Perito", "6234"]
    for field in result["fields"]:
        binding = field["subject_binding"]
        assert binding["record_attributions"][0]["format"] == "person_role_heading"
        proof = binding["source_evidence"][0]
        assert (proof["chunk_id"], proof["file_id"], proof["chunk_index"]) == (-7, -3, 2)
        assert (proof["role"], proof["category"]) == ("source", "company")
        assert proof["content_sha256"]
        assert field["written_value"] is None and field["semantic_review_required"]
    assert result["resolved"] == result["writes"] == 0


def test_v9_historical_role_and_source_rules_do_not_change(case):
    case[2]["version"] = 9
    result = validate_global_proposal(json.dumps(case[2]), case[0], "alpha", case[1])
    assert result["binding_checks"] == {"review_required": 3}
    assert all(f.get("value") is None for f in result["fields"])
    with pytest.raises(ValueError, match="V10"):
        audit(case)


@pytest.mark.parametrize("barrier", ["prose", "blank", "heading", "table", "checkbox"])
def test_continuation_breaks_at_structural_boundaries(barrier):
    from test_compilation_catalog import native_checkbox

    doc = Document()
    doc.add_paragraph("Introduzione al primo soggetto")
    doc.add_paragraph("Dato: ________")
    if barrier == "prose":
        doc.add_paragraph("Introduzione al secondo soggetto")
    elif barrier == "blank":
        doc.add_paragraph()
    elif barrier == "heading":
        doc.add_heading("Altro blocco", 1)
    elif barrier == "table":
        doc.add_table(rows=1, cols=1).cell(0, 0).text = "Altro blocco"
    else:
        native_checkbox(doc.add_paragraph("Altra opzione"))
    doc.add_paragraph("Altro dato: ________")
    catalog = build_compilation_catalog(layout_for(doc))
    last = [f for f in catalog["candidates"].values() if f["type"] == "text"][-1]
    assert last.get("continuation_origin_paragraph_id") != "paragraph:0"


def test_continuation_is_bounded_and_never_guesses_a_section():
    doc = Document()
    doc.add_paragraph("Introduzione")
    for _ in range(MAX_CONTINUATION_PARAGRAPHS + 1):
        doc.add_paragraph("Dato: ________")
    catalog = build_compilation_catalog(layout_for(doc))
    assert catalog["candidates"][f"p{MAX_CONTINUATION_PARAGRAPHS}.s0"][
        "continuation_origin_paragraph_id"] == "paragraph:0"
    assert "continuation_origin_paragraph_id" not in catalog["candidates"][
        f"p{MAX_CONTINUATION_PARAGRAPHS + 1}.s0"]
    assert all(f["section_ids"] == [] for f in catalog["candidates"].values())


@pytest.mark.parametrize("header,kind", [
    ("### Dr. Ada Neri - responsabile della verifica", "person_role_heading"),
    ("## Ada Neri — Responsabile della verifica", "person_role_heading"),
    ("RESPONSABILE DELLA VERIFICA Dr. Ada Neri - iscrizione 6234", "uppercase_role_person_row"),
])
def test_explicit_source_formats_without_domain_titles(header, kind):
    source = header + "\nNumero registrazione: 6234."
    support = record_support(source, header, ROLE, "Ada Neri", ["Numero registrazione: 6234."])
    assert support["format"] == kind


@pytest.mark.parametrize("source,quote", [
    ("### Bea Mori - Responsabile della verifica\nNumero registrazione: 6234.", None),
    ("### Ada Neri - altro incarico\nNumero registrazione: 6234.", None),
    ("### Ada Neri - non Responsabile della verifica\nNumero registrazione: 6234.", None),
    ("NON RESPONSABILE DELLA VERIFICA Ada Neri - iscrizione 6234",
     "RESPONSABILE DELLA VERIFICA Ada Neri"),
    ("La responsabile della verifica è Ada Neri. Numero registrazione: 6234.", None),
    ("### Dr. Ada Neri - responsabile della verifica\n### Bea Mori - altro incarico\n"
     "Numero registrazione: 6234.", "Dr. Ada Neri - responsabile della verifica"),
    ("RESPONSABILE DELLA VERIFICA Ada Neri\nALTRO INCARICO Bea Mori - numero registrazione 6234",
     "RESPONSABILE DELLA VERIFICA Ada Neri"),
])
def test_wrong_person_role_negation_narrative_or_record_is_not_supported(source, quote):
    assert record_support(source, quote or source, ROLE, "Ada Neri", ["6234"]) is None


def test_fact_from_other_person_in_same_chunk_is_rejected_without_contaminating_siblings(case):
    case[2]["facts"]["f2"].update(value="9999", evidence=[
        {"chunk_id": -7, "quote": "Numero registrazione: 9999."}])
    result = audit(case)
    assert result["fields"][0]["validation"] == "GROUNDED_PROPOSAL"
    assert result["fields"][1]["validation"] == "GROUNDED_PROPOSAL"
    assert result["fields"][2]["value"] is None
    assert result["fields"][2]["subject_binding"]["result"] == "review_required"


@pytest.mark.parametrize("change", ["form", "user", "general", "foreign", "other_role"])
def test_source_authority_and_exact_role_still_required(case, change):
    if change in {"form", "user"}:
        case[1][0]["role"] = change
        with pytest.raises(ValueError):
            audit(case)
    elif change == "general":
        case[1][0]["category"] = "general"
        assert audit(case)["binding_checks"] == {"review_required": 3}
    elif change == "foreign":
        case[1][0].update(scope="project:other", project_id="other")
        with pytest.raises(ValueError):
            audit(case)
    else:
        case[2]["bindings"]["b"]["required_role"] = "Altro incarico"
        assert audit(case)["binding_checks"] == {"review_required": 3}


def test_unknown_applicability_keeps_person_hypothesis_nonwritable(case):
    case[2]["conditions"]["c"] = {
        "form": {"paragraph_id": "paragraph:0", "quote": ROLE},
        "aspect": "professional", "operator": "semantic", "fact_ids": [],
        "scope": [{"kind": "candidate", "id": "p3.s0"}], "reason": "Da confermare"}
    case[2]["fields"]["p3.s0"].update(disposition="conditional", conditions=["c"],
                                    reason="Da confermare")
    field = audit(case)["fields"][-1]
    assert field["proposed_value"] == "6234"
    assert field["value"] is field["written_value"] is None
    assert field["condition_checks"]["c"]["result"] == "unknown"


def test_person_record_in_another_chunk_cannot_verify_value(case):
    other = copy.deepcopy(case[1][0])
    other.update(chunk_id=-8, content="Ada Neri. Numero registrazione: 6234.")
    case[1].append(other)
    case[2]["facts"]["f2"]["evidence"][0]["chunk_id"] = -8
    field = audit(case)["fields"][-1]
    assert field["value"] is None
    assert field["subject_binding"]["result"] == "review_required"
