"""Persistent API/export gates with simulated transport and temporary FTS5/Qdrant."""

import json
from copy import deepcopy
from io import BytesIO

import pytest
from docx import Document
from test_compilation_form_applicability import prepared
from test_compilation_global_sessions import BASE, api, provider, resolve
from test_compilation_sessions import source

from app import compilation_global_sessions as adapter
from app import compilation_sessions as sessions
from app.db import connection

__all__ = ["api", "provider"]


@pytest.fixture
def anyio_backend():
    return "asyncio"


async def start(api, monkeypatch, result):
    layout, sources, raw, _ = prepared(layout_kind="table",
        reference=None if result in {"true", "unknown"} else "Via Uno 1",
        extra="Sede operativa di Aurora: diversa dalla sede legale." if result == "true" else "")
    uploaded = await api.post("/api/projects/alpha/forms", files={
        "file": ("modulo.docx", layout.original)})
    await source(api, sources[0]["content"], name="sedi.txt")

    async def proposal(client, settings, body, timeout, **kwargs):
        inp = json.loads(body["messages"][1]["content"])
        evidence = next(s for s in inp["sources"] if "Sede operativa di Aurora:" in s["content"])
        value = deepcopy(raw)
        value["facts"]["f"]["evidence"][0]["chunk_id"] = evidence["chunk_id"]
        return json.dumps(value), "simulated", {}

    monkeypatch.setattr(adapter, "request_model_content", proposal)
    session = (await api.post(BASE, json={"form_id": uploaded.json()["id"], "engine": "global",
                                         "start_in_chat": True})).json()
    return await resolve(api, session)


@pytest.mark.anyio
@pytest.mark.parametrize("result", ["true", "false", "unknown"])
async def test_persisted_three_valued_gate_and_actual_docx_export(
    api, provider, monkeypatch, result,
):
    session = await start(api, monkeypatch, result)
    field = session["global_compilation"]["fields"][0]
    assert field["form_applicability"]["result"] == result
    assert field["selectable"] is (result == "true")
    request = {"version": session["version"], "candidate_ids": ["t0.r0.c1"],
               "allow_unresolved": True, "user_quote": "Genera la bozza parziale"}
    response = await api.post(f"{BASE}/{session['id']}/global-draft", json=request)
    if result != "true":
        assert response.status_code == 409, response.text
        assert "Applicabilità FORM" in response.json()["detail"]
        assert (await api.get(f"{BASE}/{session['id']}")).json()["last_generation"] is None
    else:
        assert response.status_code == 200, response.text
        downloads = response.json()["last_generation"]["downloads"]
        report = (await api.get(downloads["report"])).json()
        assert report["written_field_count"] == 1
        assert report["fields"][0]["provenance"] == "SOURCE"
        written = Document(BytesIO((await api.get(downloads["docx"])).content))
        assert written.tables[0].cell(0, 1).text == "Via Uno 1"


@pytest.mark.anyio
@pytest.mark.parametrize("result", ["false", "unknown"])
@pytest.mark.parametrize("as_user", [False, True])
async def test_fresh_gate_blocks_stale_audit_and_direct_source_or_user_selection(
    api, provider, monkeypatch, result, as_user,
):
    session = await start(api, monkeypatch, result)
    with connection() as db:
        row, state = adapter._row_state(db, "alpha", session["id"], session["version"])
        data = state["global_compilation"]
        # Simulate an old/stale in-memory verdict. Original SOURCE and MODEL response stay exact.
        data["audit"]["form_applicability"]["t0.r0.c1"]["result"] = "true"
        if as_user:
            data["dialogue"].setdefault("user_values", {})["t0.r0.c1"] = {
                "value": "Via Uno 1", "provenance": "USER", "user_quote": "Via Uno 1"}
        session = sessions.save_state(db, row, state, "test_stale_predicate")
    request = {"version": session["version"], "candidate_ids": [] if as_user else ["t0.r0.c1"],
               "user_candidate_ids": ["t0.r0.c1"] if as_user else [], "allow_unresolved": True,
               "user_quote": "Genera la bozza, confermo tutto", "confirmed_condition_ids": []}
    response = await api.post(f"{BASE}/{session['id']}/global-draft", json=request)
    assert response.status_code == 409, response.text
    assert response.json()["detail"] == "Applicabilità FORM cambiata o non dimostrata"
    assert (await api.get(f"{BASE}/{session['id']}")).json()["last_generation"] is None
