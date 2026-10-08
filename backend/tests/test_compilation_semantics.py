"""Structural anchors, independent semantic gates, relational SOURCE and empty values."""
import json
from io import BytesIO

import pytest
from docx import Document

from app import compilation_session_resolution as resolution
from app.compilation_clarifications import synchronize
from app.compilation_semantics import (
    FormReview,
    HistoricalServices,
    SourceReview,
    empty_value_reason,
    propagate_section_exclusions,
    review_meanings,
    review_sources,
)
from app.compilation_session_models import CandidateMatches, CandidateMeanings
from app.compilation_sessions import candidate_snapshots, validate_value
from app.docx_templates import inspect_docx


@pytest.fixture
def anyio_backend():
    return "asyncio"


def layout_fields(text="Delegato della cooperativa ______ con recapito in ______ Comune ______"):
    doc = Document()
    doc.add_paragraph(text)
    output = BytesIO()
    doc.save(output)
    layout = inspect_docx(output.getvalue())
    return layout, candidate_snapshots(layout, "alpha", 1, "modulo.docx")


def meaning(field, name, **semantic_overrides):
    return {
        "candidate_id": field["id"], "classification": "data", "entity": "company",
        "requirement": {"name": name, "form_quote": field["label_hint"],
                        "form_citation_id": 1, "person_role": ""},
        "form_quote": field["label_hint"], "reason": "Proprietà della organizzazione rappresentata",
        "semantic": {
            "form_anchor": field["structural"]["slot_anchor"],
            "subject_anchor": "cooperativa", "subject_relation": "represented_organization",
            "context_role": "ORGANIZATION_PROFILE", **semantic_overrides,
        },
    }


async def approve_form(task, data, schema):
    assert schema is FormReview
    return FormReview(fields=[{
        "candidate_id": p["candidate"]["id"], "accepted": True, "condition_complete": True,
        "reason": "Slot e relazione controllati nel FORM",
    } for p in data["proposals"]])


async def classified(fields, names, request=approve_form, **semantic):
    meanings = CandidateMeanings(fields=[meaning(f, name, **semantic)
                                        for f, name in zip(fields, names, strict=True)])
    await review_meanings(fields, meanings, request)
    resolution.apply_meanings(fields, meanings)
    return meanings


def source(content):
    return {"chunk_id": -1, "file_id": 2, "source_name": "profilo.txt", "chunk_index": 0,
            "content": content, "role": "source", "scope": "global", "category": "company",
            "project_id": None, "document_metadata": "test"}


async def source_review(task, data, schema):
    assert schema is SourceReview
    return SourceReview(supports=[{
        "candidate_id": p["candidate_id"], "kind": p["kind"], "index": p["index"],
        "accepted": True, "relation_quote": p["source"]["content"],
        "reason": "Proprietà, soggetto e relazione verificati",
    } for p in data["proposals"]])


@pytest.mark.anyio
async def test_composite_slots_nonliteral_names_and_company_relation():
    layout, fields = layout_fields()
    await classified(fields, ["Denominazione aziendale", "Indirizzo sede", "Comune sede"])
    assert [f["requirement"]["name"] for f in fields] == [
        "Denominazione aziendale", "Indirizzo sede", "Comune sede",
    ]
    assert all(f["entity"] == "company" and not f["requirement"]["person_role"] for f in fields)
    evidence = source("La cooperativa Aurora ha sede in Via del Porto 19, Genova.")
    values = ["Aurora", "Via del Porto 19", "Genova"]
    matches = CandidateMatches(fields=[{
        "candidate_id": f["id"], "supports": [{"source_id": 1, "quote": evidence["content"],
                                               "value": value}], "reason": "Dato documentato",
    } for f, value in zip(fields, values, strict=True)])
    coverage = {f["id"]: [1] for f in fields}
    await review_sources(fields, matches, [evidence], coverage, source_review)
    resolution.apply_matches(layout, fields, matches, [evidence], coverage)
    assert [f["value"] for f in fields] == values
    assert all(f["status"] == "RESOLVED" and f["provenance"] == "SOURCE" for f in fields)
    for f in fields:  # The same persisted review must remain usable at finalization.
        resolution.validate_support(f, evidence, evidence["content"], f["value"])


@pytest.mark.anyio
async def test_unsupported_semantic_property_is_rejected_before_source():
    _, fields = layout_fields()

    async def deny(task, data, schema):
        return FormReview(fields=[{"candidate_id": fields[0]["id"], "accepted": False,
                                   "condition_complete": True,
                                   "reason": "Il blank richiede il nome, non il fatturato"}])

    await classified(fields[:1], ["Ricavi annuali"], deny)
    assert fields[0]["status"] == "PENDING" and fields[0]["requirement"] is None
    assert fields[0]["search"]["status"] == "not_searched"
    assert "fatturato" in fields[0]["validation_errors"][0]


@pytest.mark.anyio
async def test_anchor_of_other_slot_cannot_authorize_meaning():
    _, fields = layout_fields()
    await classified(fields[:1], ["Denominazione"], form_anchor="Comune [[p0.s2]]")
    assert fields[0]["requirement"] is None
    assert "Anchor" in fields[0]["validation_errors"][0]


@pytest.mark.anyio
async def test_backend_binds_empty_anchor_to_actual_slot():
    _, fields = layout_fields()
    await classified(fields[:1], ["Denominazione"], form_anchor="")
    assert fields[0]["semantic"]["form_anchor"] == fields[0]["structural"]["slot_anchor"]
    assert fields[0]["requirement"]["name"] == "Denominazione"


@pytest.mark.anyio
@pytest.mark.parametrize("condition", [None, "cooperativa", "società straniera"])
async def test_section_condition_review_must_be_present_and_literal(condition):
    _, fields = layout_fields()

    async def assess(task, data, schema):
        return FormReview(fields=[{
            "candidate_id": fields[0]["id"], "accepted": True, "condition_complete": True,
            "reason": "Sezione condizionata",
            "section_condition": {"condition": condition, "section_id": "paragraph:0",
                                  "condition_kind": "subject_type"} if condition else None,
        }])

    await classified(fields[:1], ["Denominazione"], assess, section_id="paragraph:0")
    if condition in {None, "cooperativa"}:
        assert fields[0]["condition"] == "cooperativa"
        assert fields[0]["semantic"]["condition_kind"] == "subject_type"
        resolution.validate_persisted_semantics(fields[0])
    else:
        assert fields[0]["requirement"] is None
        assert fields[0]["status"] == "PENDING"


@pytest.mark.anyio
async def test_implicit_type_condition_cannot_drop_foreign_section_qualifier():
    _, fields = layout_fields("Delegato della cooperativa stabilita all'estero ______")

    async def incomplete(task, data, schema):
        assert data["proposals"][0]["interpretation"]["condition"] == "cooperativa"
        return FormReview(fields=[{
            "candidate_id": fields[0]["id"], "accepted": True, "condition_complete": False,
            "reason": "La tipologia non comprende lo stabilimento all'estero",
        }])

    await classified(fields, ["Denominazione"], incomplete, section_id="paragraph:0")
    assert fields[0]["requirement"] is None and fields[0]["status"] == "PENDING"
    assert "qualificatori" in fields[0]["validation_errors"][0]


@pytest.mark.anyio
async def test_company_type_condition_cannot_include_representative_person_role():
    _, fields = layout_fields()
    item = meaning(fields[0], "Denominazione", section_id="paragraph:0",
                   condition_kind="subject_type")
    item["condition"] = "Delegato della cooperativa"
    meanings = CandidateMeanings(fields=[item])
    await review_meanings(fields[:1], meanings, approve_form)
    resolution.apply_meanings(fields[:1], meanings)
    assert fields[0]["status"] == "PENDING" and fields[0]["requirement"] is None
    assert "senza ruolo personale" in fields[0]["validation_errors"][0]


@pytest.mark.anyio
async def test_represented_organization_cannot_have_person_role():
    _, fields = layout_fields()
    item = meaning(fields[0], "Denominazione")
    item["entity"] = "person"
    item["requirement"]["person_role"] = "Delegato"
    meanings = CandidateMeanings(fields=[item])
    await review_meanings(fields, meanings, approve_form)
    resolution.apply_meanings(fields, meanings)
    assert fields[0]["requirement"] is None
    assert "Organizzazione rappresentata" in fields[0]["validation_errors"][0]


@pytest.mark.anyio
@pytest.mark.parametrize("entity", ["authority", "other"])
async def test_organization_subject_does_not_require_company_entity(entity):
    _, fields = layout_fields("Servizi già eseguiti: ente committente ______")
    field = fields[0]
    item = meaning(field, "Committente pregresso", subject_anchor="ente committente",
                   subject_relation="organization", context_role="PAST_SERVICE",
                   context_quote="Servizi già eseguiti")
    item["entity"] = entity
    meanings = CandidateMeanings(fields=[item])
    await review_meanings([field], meanings, approve_form)
    resolution.apply_meanings([field], meanings)
    assert field["requirement"]["name"] == "Committente pregresso"
    assert field["entity"] == entity
    assert field["semantic"]["context_role"] == "PAST_SERVICE"


@pytest.mark.anyio
async def test_section_type_applicability_from_source_shared_by_slots():
    layout, fields = layout_fields()
    meanings = CandidateMeanings(fields=[{
        **meaning(f, name, section_id="paragraph:0", condition_kind="subject_type"),
        "condition": "cooperativa",
    } for f, name in zip(fields, ["Denominazione", "Indirizzo", "Comune"], strict=True)])
    await review_meanings(fields, meanings, approve_form)
    resolution.apply_meanings(fields, meanings)
    evidence = source("Aurora è una cooperativa. Sede: Via del Porto 19, Genova.")
    matches = CandidateMatches(fields=[{
        "candidate_id": fields[0]["id"], "supports": [],
        "applicability": [{"source_id": 1, "quote": evidence["content"],
                           "value": "Aurora è una cooperativa", "applies": True}],
        "reason": "Tipologia attestata",
    }])
    coverage = {fields[0]["id"]: [1]}
    await review_sources(fields, matches, [evidence], coverage, source_review)
    resolution.apply_matches(layout, fields, matches, [evidence], coverage)
    state = {"fields": fields, "status": "CREATED", "chat_workflow": {
        "clarification_mode": "grouped", "steps": 0,
        "started_at": "2099-01-01T00:00:00+00:00", "analysis_attempts": {},
        "source_attempts": {}, "active_clarifications": [],
    }}
    synchronize(state)
    assert all(f["applicability"]["applies"] for f in fields)
    assert all(f["applicability"]["provenance"] == "SOURCE" for f in fields)


@pytest.mark.anyio
@pytest.mark.parametrize("extra", ["unrequested", "duplicate"])
async def test_bad_source_review_item_preserves_valid_value_and_applicability(extra):
    layout, fields = layout_fields()
    field = fields[0]
    item = meaning(field, "Denominazione", section_id="paragraph:0", condition_kind="subject_type")
    item["condition"] = "cooperativa"
    meanings = CandidateMeanings(fields=[item])
    await review_meanings([field], meanings, approve_form)
    resolution.apply_meanings([field], meanings)
    evidence = source("Aurora è una cooperativa.")
    matches = CandidateMatches(fields=[{"candidate_id": field["id"], "reason": "Fonte",
        "supports": [{"source_id": 1, "quote": evidence["content"], "value": "Aurora"}] * 2,
        "applicability": [{"source_id": 1, "quote": evidence["content"],
                           "value": "cooperativa", "applies": True}]}])

    async def assess(task, data, schema):
        verdicts = [{"candidate_id": field["id"], "kind": p["kind"], "index": p["index"],
                     "accepted": True, "reason": "Proposta verificata"} for p in data["proposals"]]
        verdicts.append({"candidate_id": field["id"], "accepted": True, "reason": "Indice errato",
                         "kind": "applicability" if extra == "unrequested" else "value",
                         "index": 1})
        return SourceReview(supports=verdicts)

    coverage = {field["id"]: [1]}
    await review_sources([field], matches, [evidence], coverage, assess)
    resolution.apply_matches(layout, [field], matches, [evidence], coverage)
    assert field["status"] == "RESOLVED" and field["value"] == "Aurora"
    assert field["applicability"]["applies"] and field["applicability"]["provenance"] == "SOURCE"
    assert any("ignorata" in error for error in field["validation_errors"])
    if extra == "duplicate":
        assert not field["source_proposals"][1]["accepted"]


@pytest.mark.anyio
@pytest.mark.parametrize("historical", [False, True])
async def test_source_must_document_past_service_relationship(historical):
    layout, fields = layout_fields("cooperativa: servizi già eseguiti. Committente ______")
    await classified(fields, ["Committente del servizio"], context_role="PAST_SERVICE",
                     context_quote="servizi già eseguiti", subject_relation="contract_party")
    evidence = source("Aurora ha eseguito nel 2021 un servizio per Comune di Levante."
                      if historical else "Avviso corrente. Stazione appaltante: Comune di Levante.")
    matches = CandidateMatches(fields=[{
        "candidate_id": fields[0]["id"], "supports": [{"source_id": 1,
            "quote": evidence["content"], "value": "Comune di Levante"}], "reason": "Proposta",
    }])

    async def assess(task, data, schema):
        if schema is HistoricalServices:
            assert set(data) == {"sources"}  # No property/target to bias the classification.
            return HistoricalServices(facts=[{
                "source_id": 1, "performer_quote": "Aurora", "service_quote": "un servizio",
                "performance_quote": evidence["content"],
            }] if historical else [])
        return SourceReview(supports=[{
            "candidate_id": fields[0]["id"], "kind": "value", "index": 0,
            "accepted": True, "relation_quote": evidence["content"],
            "reason": "Proprietà apparentemente corretta: non basta per il contesto storico",
        }])

    coverage = {fields[0]["id"]: [1]}
    await review_sources(fields, matches, [evidence], coverage, assess)
    resolution.apply_matches(layout, fields, matches, [evidence], coverage)
    assert fields[0]["status"] == ("RESOLVED" if historical else "MISSING")
    assert fields[0]["source_proposals"][0]["accepted"] is historical
    if historical:
        resolution.validate_support(fields[0], evidence, evidence["content"], "Comune di Levante")


@pytest.mark.anyio
@pytest.mark.parametrize("performer,performance", [
    ("Inventata", "Aurora ha eseguito un servizio per Comune di Levante."),
    ("Aurora", "Aurora ha eseguito un servizio inesistente per Comune di Levante."),
])
async def test_historical_proof_must_bind_literal_performer_and_performance(performer, performance):
    layout, fields = layout_fields("cooperativa: servizi già eseguiti. Committente ______")
    await classified(fields, ["Committente"], context_role="PAST_SERVICE",
                     context_quote="servizi già eseguiti", subject_relation="contract_party")
    evidence = source("Aurora ha eseguito un servizio per Comune di Levante.")
    matches = CandidateMatches(fields=[{"candidate_id": fields[0]["id"], "reason": "Proposta",
        "supports": [{"source_id": 1, "quote": evidence["content"],
                      "value": "Comune di Levante"}]}])

    async def assess(task, data, schema):
        assert schema is HistoricalServices
        return HistoricalServices(facts=[{"source_id": 1, "performer_quote": performer,
            "service_quote": "un servizio", "performance_quote": performance}])

    coverage = {fields[0]["id"]: [1]}
    await review_sources(fields, matches, [evidence], coverage, assess)
    resolution.apply_matches(layout, fields, matches, [evidence], coverage)
    assert fields[0]["status"] == "MISSING"


@pytest.mark.anyio
async def test_semantic_review_cannot_replace_literal_source_provenance():
    layout, fields = layout_fields()
    await classified(fields[:1], ["Denominazione"])
    evidence = source("Organizzazione: Aurora")
    matches = CandidateMatches(fields=[{
        "candidate_id": fields[0]["id"], "supports": [{"source_id": 1,
            "quote": "Organizzazione: Inventata", "value": "Inventata"}], "reason": "Proposta",
    }])
    coverage = {fields[0]["id"]: [1]}
    await review_sources(fields[:1], matches, [evidence], coverage, source_review)
    resolution.apply_matches(layout, fields[:1], matches, [evidence], coverage)
    assert fields[0]["status"] == "MISSING"
    assert "SOURCE" in fields[0]["validation_errors"][0]


@pytest.mark.parametrize("value", ["________", "Identificativo: ________", "N/A", "....",
                                   "Campo non compilato", "Identificativo:", "Identificativo"])
def test_empty_information_never_passes_proposal_validation(value):
    layout, fields = layout_fields("Identificativo ______")
    fields[0].update(label="Identificativo", entity="other", kind="data")
    evidence = source("Identificativo: " + value)
    result = validate_value(layout, fields[0], value, evidence, evidence["content"])
    assert "empty_information" in result["validation_codes"]
    assert empty_value_reason(value, "Identificativo")


@pytest.mark.parametrize("value", ["Aurora", "Via del Porto 19", "Genova", "16100", "GE",
                                   "+39 010 123456", "posta@pec.demo", "NA"])
def test_information_values_are_not_empty(value):
    assert empty_value_reason(value) is None


def test_exclusivity_requires_reviewed_explicit_form_group():
    def section(ident, group):
        return {"id": ident, "condition": ident, "entity": "company", "status": "MISSING",
                "semantic": {"section_id": ident, "exclusive_group_id": group},
                "semantic_validation": {"accepted": True},
                "form_evidence": {"template_sha256": "same-original"}}
    a, b, unrelated = section("a", "choose-one"), section("b", "choose-one"), section("c", "")
    a["applicability"] = {"condition": "a", "applies": True, "provenance": "SOURCE",
                          "evidence": [{"quote": "Tipo A"}]}
    propagate_section_exclusions({"fields": [a, b, unrelated]})
    assert b["status"] == "NOT_APPLICABLE" and b["provenance"] == "SOURCE"
    assert unrelated["status"] == "MISSING" and "applicability" not in unrelated


@pytest.mark.parametrize("value", ["________", "Riferimento: ________", "Riferimento: N/A"])
def test_placeholder_source_cannot_become_resolved(value):
    layout, fields = layout_fields("Riferimento ______")
    field = fields[0]
    meanings = CandidateMeanings(fields=[{
        "candidate_id": field["id"], "classification": "data", "entity": "other",
        "requirement": {"name": "Riferimento", "form_quote": "Riferimento",
                        "form_citation_id": 1}, "form_quote": "Riferimento", "reason": "Campo",
    }])
    resolution.apply_meanings(fields, meanings)
    evidence = source("Riferimento: " + value)
    matches = CandidateMatches(fields=[{"candidate_id": field["id"], "reason": "Proposta",
        "supports": [{"source_id": 1, "value": value, "quote": evidence["content"]}]}])
    resolution.apply_matches(layout, fields, matches, [evidence], {field["id"]: [1]})
    assert field["status"] == "MISSING" and field["value"] is None
    assert not field["source_proposals"][0]["accepted"]


def test_table_keeps_section_context_beyond_immediate_instructions():
    doc = Document()
    doc.add_paragraph("Elencare contratti già eseguiti dalla candidata nel triennio.")
    for i in range(8):
        doc.add_paragraph(f"Istruzione di compilazione numero {i}.")
    table = doc.add_table(rows=7, cols=2)
    table.cell(0, 0).text = "Committente"
    table.cell(0, 1).text = "Importo"
    output = BytesIO()
    doc.save(output)
    layout = inspect_docx(output.getvalue())
    fields = candidate_snapshots(layout, "alpha", 1, "esperienze.docx")
    for field in fields:
        assert "contratti già eseguiti" in field["structural"]["form_sections"][0]["text"]


@pytest.mark.parametrize("schema,key,extra", [
    (FormReview, "fields", {"condition_complete": True}),
    (SourceReview, "supports", {"kind": "value", "index": 0}),
])
def test_partial_review_preserves_valid_item_and_never_approves_malformed(schema, key, extra):
    review = resolution.parse_structured(json.dumps({key: [
        {"candidate_id": "a", "accepted": True, "reason": "Verificato", **extra},
        {"candidate_id": "b", "accepted": "invalid", "reason": "Non verificato", **extra},
    ]}), schema)
    assert [v.candidate_id for v in getattr(review, key)] == ["a"]


def test_partial_historical_facts_keep_valid_proof():
    review = resolution.parse_structured(json.dumps({"facts": [
        {"source_id": 1, "performer_quote": "Aurora", "service_quote": "un servizio",
         "performance_quote": "Aurora ha eseguito un servizio."},
        {"source_id": 2, "performance_quote": "Incomplete"},
    ]}), HistoricalServices)
    assert [fact.source_id for fact in review.facts] == [1]


def test_duplicate_source_review_discards_only_its_own_proposal():
    item = {"candidate_id": "a", "kind": "value", "index": 0,
            "accepted": True, "reason": "Verificato"}
    review = resolution.parse_structured(json.dumps({"supports": [
        item, {**item, "candidate_id": "b"}, {**item, "accepted": False},
    ]}), SourceReview)
    assert [v.candidate_id for v in review.supports] == ["b"]


@pytest.mark.anyio
async def test_administrator_is_not_evidence_of_practice_signatory():
    layout, fields = layout_fields("Firmatario della candidatura ______ cooperativa")
    field = fields[0]
    item = meaning(field, "Nominativo sottoscrittore", subject_relation="signatory",
                   context_role="PERSON_PROFILE", subject_anchor="Firmatario della candidatura")
    item["entity"] = "person"
    meanings = CandidateMeanings(fields=[item])
    await review_meanings(fields, meanings, approve_form)
    resolution.apply_meanings(fields, meanings)
    evidence = source("Amministratrice della cooperativa Aurora: Ada Neri.")
    matches = CandidateMatches(fields=[{"candidate_id": field["id"], "reason": "Proposta",
        "supports": [{"source_id": 1, "quote": evidence["content"], "value": "Ada Neri"}]}])

    async def deny(task, data, schema):
        return SourceReview(supports=[{"candidate_id": field["id"], "kind": "value", "index": 0,
            "accepted": False, "reason": "La carica non prova la scelta del firmatario"}])

    coverage = {field["id"]: [1]}
    await review_sources(fields, matches, [evidence], coverage, deny)
    resolution.apply_matches(layout, fields, matches, [evidence], coverage)
    assert field["status"] == "MISSING" and field["provenance"] is None
