"""Document-level relationships, independent of tender names/layout conventions."""

from copy import deepcopy
from io import BytesIO

import pytest
from docx import Document

from app.compilation_chat import new_cycle
from app.compilation_clarifications import pending_slots, synchronize
from app.compilation_conditions import owner
from app.compilation_document_plan import (
    DocumentPlan,
    DocumentPlanReview,
    document_blocks,
    verified_plan,
)
from app.compilation_sessions import candidate_snapshots
from app.docx_templates import inspect_docx


def scenario(*, approved=True):
    doc = Document()
    doc.add_paragraph("Indicare una sola categoria del richiedente")
    doc.add_paragraph("Richiedente appartenente alla categoria Alfa")
    table = doc.add_table(rows=4, cols=2)
    table.cell(0, 0).text = "Nominativo"
    table.cell(0, 1).text = "Qualifica"
    doc.add_paragraph("Richiedente appartenente alla categoria Beta")
    doc.add_paragraph("Denominazione: __________________")
    doc.add_paragraph("Dati comuni a tutti i richiedenti")
    doc.add_paragraph("Referente: __________________")
    stream = BytesIO()
    doc.save(stream)
    layout = inspect_docx(stream.getvalue())
    blocks = document_blocks(layout)
    instruction, first, _, second, _, common, _ = blocks
    proposal = DocumentPlan(sections=[
        {"start_id": first["id"], "end_before_id": second["id"],
         "condition": "categoria Alfa", "kind": "subject_type", "reason": "Ramo Alfa"},
        {"start_id": second["id"], "end_before_id": common["id"],
         "condition": "categoria Beta", "kind": "subject_type", "reason": "Ramo Beta"},
    ], choices=[{"instruction_id": instruction["id"], "quote": instruction["text"],
                 "section_ids": [first["id"], second["id"]], "reason": "Scelta unica"}])
    review = DocumentPlanReview(accepted_sections=[first["id"], second["id"]],
                               accepted_choices=[instruction["id"]] if approved else [],
                               reason="Relazioni verificate nell'originale")
    fields = candidate_snapshots(layout, "project-one", 8, "arbitrary.docx")
    for index, field in enumerate(fields):
        field.update(status="MISSING", condition="", entity="person" if index % 2 else "company",
                     requirement={"name": field["label_hint"],
                                  "person_role": "membro" if index % 2 else ""},
                     search={"status": "searched"})
    state = {"template_sha256": layout.sha256, "fields": fields,
             "status": "WAITING_FOR_USER", "chat_workflow": new_cycle(),
             "document_plan": verified_plan(layout, proposal, review)}
    return layout, proposal, review, state


def test_one_condition_for_members_and_organization_with_different_field_roles():
    _, _, _, state = scenario()
    synchronize(state)
    slots = pending_slots(state)
    alpha = [s for s in slots if s["condition"] == "categoria Alfa"]
    assert len(alpha) == 1 and len(alpha[0]["field_ids"]) == 6
    assert len({f["requirement"]["person_role"] for f in state["fields"][:6]}) == 2


def test_source_selects_only_reviewed_alternative_and_preserves_common_fields():
    _, _, _, state = scenario()
    synchronize(state)
    beta = state["fields"][-2]
    beta.update(condition="categoria Beta", applicability={
        "condition": "categoria Beta", "applies": True, "provenance": "SOURCE",
        "evidence": [{"role": "source", "quote": "Il richiedente appartiene alla categoria Beta"}],
    })
    synchronize(state)
    assert all(f["status"] == "NOT_APPLICABLE" for f in state["fields"][:6])
    assert state["fields"][-1]["status"] == "MISSING"
    assert not any(s["condition"] == "categoria Alfa" for s in pending_slots(state))
    before = deepcopy(state)
    synchronize(state)
    assert before == state


def test_no_exclusion_from_unreviewed_choice_even_with_positive_type():
    _, _, _, state = scenario(approved=False)
    synchronize(state)
    beta = state["fields"][-2]
    beta.update(condition="categoria Beta", applicability={
        "condition": "categoria Beta", "applies": True, "provenance": "USER", "user_quote": "sì",
    })
    synchronize(state)
    assert all(f["status"] == "MISSING" for f in state["fields"][:6])


@pytest.mark.parametrize("mutation", ["quote", "unknown_id", "overlap", "duplicate", "review"])
def test_invalid_map_cannot_authorize_branch_propagation(mutation):
    layout, proposal, review, _ = scenario()
    if mutation == "quote":
        proposal.sections[0].condition = "categoria inesistente"
    elif mutation == "unknown_id":
        proposal.sections[0].end_before_id = "not-in-document"
    elif mutation == "overlap":
        proposal.sections[0].end_before_id = None
    elif mutation == "duplicate":
        proposal.sections.append(proposal.sections[0].model_copy())
    else:
        review.accepted_sections = []
    result = verified_plan(layout, proposal, review)
    assert proposal.sections[0].start_id not in {s["start_id"] for s in result["sections"]}
    assert not result["choices"]


def test_map_and_scope_cannot_cross_original_or_project():
    _, _, _, state = scenario()
    synchronize(state)
    field = state["fields"][0]
    field["condition"] = "categoria Alfa"
    other = deepcopy(field)
    other["form_evidence"]["project_id"] = "project-two"
    assert owner(field) != owner(other)
    state["template_sha256"] = "another-original"
    with pytest.raises(ValueError, match="originale"):
        synchronize(state)


def test_modified_persisted_plan_requires_revalidation():
    _, _, _, state = scenario()
    state["document_plan"]["sections"][0]["condition"] = "Changed"
    with pytest.raises(ValueError, match="originale"):
        synchronize(state)
