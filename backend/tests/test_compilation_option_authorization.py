"""V14 option-clause authorization: a branch selection is never a declaration authorization."""

import json
from copy import deepcopy
from io import BytesIO

import pytest
from docx import Document
from test_compilation_catalog import catalog, native_checkbox
from test_compilation_global_sessions import BASE, api, provider, resolve
from test_compilation_sessions import source
from test_docx_formtext import form_text

from app import compilation_global_dialogue as dialogue
from app import compilation_global_sessions as adapter
from app import compilation_sessions as sessions
from app.compilation_chat import ClarificationReply
from app.compilation_option_regions import option_regions
from app.compilation_procedural_consent import blocked_fields, requirements

__all__ = ["api", "provider"]

CLAUSE = ("☐ per le società DI ESSERE iscritto alla CCIAA per un’attività coerente a quella "
          "oggetto della classe merceologica per la quale si richiede l’iscrizione, con i "
          "seguenti dati:")
ASSENT = "Sì, autorizzo espressamente questa dichiarazione"


def cciaa_doc(doc=None):
    doc = doc or Document()
    doc.add_paragraph(CLAUSE)
    table = doc.add_table(rows=1, cols=4)
    table.rows[0].cells[0].text = "Provincia di iscrizione:"
    table.rows[0].cells[2].text = "numero di iscrizione:"
    return doc


def by_context(cat, text):
    """Text candidates whose own paragraph contains text (inline slots only)."""
    paragraphs = {p["id"]: p["text"] for b in cat["blocks"] for p in b["paragraphs"]}
    return [cid for cid, c in cat["candidates"].items()
            if c["type"] == "text" and text in paragraphs.get(c["context_id"], "")]


def test_regions_follow_structure_not_wording():
    doc = cciaa_doc()
    doc.add_paragraph("Telefono dell'operatore ____________")
    doc.add_paragraph("☐ di applicare il contratto ____________;")
    doc.add_paragraph("ed indica gli operatori che compongono il raggruppamento:")
    shared = doc.add_table(rows=1, cols=2)
    shared.rows[0].cells[0].text = "Denominazione"
    cell = doc.add_table(rows=2, cols=1)
    cell.rows[0].cells[0].text = "☐ titolare  ☐ altro ____________"
    cell.rows[1].cells[0].text = "dell’impresa ____________"
    doc.add_paragraph("Il sottoscritto ____________ ☐ in qualità di ____________")
    cat = catalog(doc)
    regions = {r["paragraph_id"]: r for r in option_regions(cat)}
    governed = {cid for r in regions.values() for cid in r["field_ids"]}
    cciaa = [cid for cid, c in cat["candidates"].items()
             if c["location"].get("table_id") == "table:0"]
    assert len(cciaa) == 2 and set(cciaa) <= governed
    assert by_context(cat, "applicare il contratto")[0] in governed
    # Shared content after an intervening paragraph, other table rows and text
    # before the marker are not attributed to the option.
    assert not any(c["location"].get("table_id") == "table:1" for cid, c in
                   cat["candidates"].items() if cid in governed)
    assert by_context(cat, "dell’impresa")[0] not in governed
    assert by_context(cat, "Telefono")[0] not in governed
    signer, capacity = sorted(by_context(cat, "Il sottoscritto"),
                              key=lambda cid: cat["candidates"][cid]["range"][0])
    assert signer not in governed and capacity in governed
    assert by_context(cat, "☐ altro")[0] in governed
    assert all(r["in_table_cell"] for r in regions.values() if "altro" in r["text"])


def test_native_option_controls_are_the_decision_not_governed_values():
    doc = Document()
    choice = doc.add_paragraph()
    native_checkbox(choice)  # Cesano layout: option control plus a legacy text field.
    choice.add_run(" Della Società di Ingegneria (3) ")
    form_text(choice)
    pair = doc.add_paragraph()
    native_checkbox(pair)
    pair.add_run(" SI ")
    native_checkbox(pair)
    pair.add_run(" NO")
    cat = catalog(doc)
    regions = option_regions(cat)
    boxes = {cid for cid, c in cat["candidates"].items() if c["type"] == "checkbox"}
    assert len(regions) == 1 and regions[0]["field_ids"] == by_context(cat, "Società")
    assert not boxes & set(regions[0]["field_ids"])
    assert set(regions[0]["own_controls"]) < boxes


def session_data(cat, values, *, flag=True, answers=None):
    fields = [{"candidate_id": cid, "validation": "GROUNDED_PROPOSAL",
               "value": values[cid], "conditions": []} if cid in values else
              {"candidate_id": cid, "validation": "REVIEW_REQUIRED", "value": None,
               "conditions": []} for cid in cat["candidates"]]
    data = {"audit": {"version": 12, "fields": fields, "conditions": {}},
            "catalog": cat, "writable_ids": sorted(cat["candidates"]),
            "artifact_path": "alpha/_global_compilations/s/a",
            "dialogue": {"answers": answers or {}, "user_values": {}}}
    if flag:
        data["option_authorization_revision"] = 1
    return data


def test_requirements_are_new_session_only_and_commitments_stay_identical():
    doc = cciaa_doc()
    doc.add_paragraph("di accettare ogni comunicazione alla PEC ____________ indicata")
    cat = catalog(doc)
    cells = [cid for cid, c in cat["candidates"].items() if c["location"].get("table_id")]
    pec = by_context(cat, "PEC")[0]
    values = {cells[0]: "Bari", cells[1]: "01234560981", pec: "pec@demo.it"}
    old, new = (requirements(session_data(cat, values, flag=flag)) for flag in (False, True))
    assert old == new[:1]  # Existing commitment keys, questions and scopes unchanged.
    option = new[1]
    assert option["authorization"]["kind"] == "option_clause"
    assert "attività coerente" in option["question"] and "01234560981" in option["question"]
    assert blocked_fields(session_data(cat, values, flag=False)) == {pec}
    assert blocked_fields(session_data(cat, values)) == {pec, *cells}
    # A FALSE group effect is not a value to write: nothing to authorize.
    assert blocked_fields(session_data(cat, {cells[0]: False})) == set()


def test_selection_receipts_never_unlock_but_the_exact_authorization_does():
    cat = catalog(cciaa_doc())
    cells = [cid for cid, c in cat["candidates"].items() if c["location"].get("table_id")]
    values = {cells[0]: "Bari", cells[1]: "01234560981"}
    data = session_data(cat, values)
    data["audit"]["conditions"]["c_societa"] = {
        "validation": "GROUNDED_PROPOSAL", "result": "unknown",
        "form": {"paragraph_id": "paragraph:0", "quote": "per le società"}}
    for field in data["audit"]["fields"]:
        if field["candidate_id"] in cells:
            field["conditions"] = ["c_societa"]
    branch = {"action": "CONDITION_TRUE", "applies": True, "provenance": "USER",
              "user_quote": "Sì, Mapi è una società"}
    data["dialogue"]["answers"]["condition:c_societa"] = branch
    assert blocked_fields(data) == set(cells)
    requirement = requirements(data)[0]
    receipt = {"action": "CONDITION_TRUE", "applies": True, "provenance": "USER",
               "user_quote": ASSENT, "question": requirement["question"],
               "authorization": requirement["authorization"]}
    for wrong in [{**receipt, "action": "SKIP", "applies": None},
                  {**receipt, "user_quote": "Non so, forse"},
                  {**receipt, "question": "Confermi che è una società?"}]:
        data["dialogue"]["answers"][requirement["key"]] = wrong
        assert blocked_fields(data) == set(cells)
    data["dialogue"]["answers"][requirement["key"]] = receipt
    assert blocked_fields(data) == set()
    changed = deepcopy(data)
    next(f for f in changed["audit"]["fields"] if f["candidate_id"] == cells[0])["value"] = "Lecce"
    assert blocked_fields(changed) == {cells[0], cells[1]}  # A new value needs a new assent.


@pytest.mark.anyio
@pytest.mark.parametrize("branch_condition", [False, True])
@pytest.mark.parametrize("action", ["CONDITION_TRUE", "SKIP", "OLD_SESSION"])
async def test_real_session_branch_choice_does_not_write_the_attested_cells(
    api, provider, monkeypatch, branch_condition, action,
):
    stream = BytesIO()
    cciaa_doc().save(stream)
    uploaded = await api.post("/api/projects/alpha/forms", files={
        "file": ("modulo.docx", stream.getvalue())})
    await source(api, "Mapi Demo S.r.l.; Camera di Commercio competente: Bari. Numero di "
                      "iscrizione al Registro delle Imprese: 01234560981.", name="visura.md")

    async def proposal(client, settings, body, timeout, **kwargs):
        inp = json.loads(body["messages"][1]["content"])
        cells = sorted(cid for cid, c in inp["form"]["candidates"].items() if c["type"] == "text")
        labels = {p["id"]: p["text"] for b in inp["form"]["blocks"] for p in b["paragraphs"]}
        label = {cell: next(p for p in labels if labels[p].startswith(text))
                 for cell, text in zip(cells, ["Provincia", "numero"], strict=True)}
        evidence = next(s for s in inp["sources"] if "01234560981" in s["content"])
        facts = {name: {"subject": "Mapi Demo S.r.l.", "entity": "company",
                        "aspect": "identity", "property": name, "value": value,
                        "evidence": [{"chunk_id": evidence["chunk_id"], "quote": quote}]}
                 for name, value, quote in [
                     ("provincia", "Bari", "Camera di Commercio competente: Bari"),
                     ("numero", "01234560981", "Registro delle Imprese: 01234560981")]}
        conditions = {"c_societa": {
            "form": {"paragraph_id": "paragraph:0", "quote": "per le società"},
            "aspect": "other", "operator": "semantic", "fact_ids": [],
            "reason": "Ramo riservato alle società",
            "scope": [{"kind": "table", "id": "table:0"}]}} if branch_condition else {}
        return json.dumps({"version": 12, "facts": facts, "conditions": conditions,
            "fields": {cell: {"label": name, "kind": "data",
                "disposition": "conditional" if branch_condition else "value",
                "conditions": list(conditions), "property": name, "fact_id": name,
                "form": {"paragraph_id": label[cell], "quote": labels[label[cell]].strip()},
                **({"reason": "Ramo da confermare"} if branch_condition else {})}
                for cell, name in zip(cells, ["provincia", "numero"], strict=True)},
            "bindings": {}, "choice_groups": {}, "coverage": [], "clarifications": {}}), \
            "simulated", {}

    monkeypatch.setattr(adapter, "request_model_content", proposal)
    session = (await api.post(BASE, json={"form_id": uploaded.json()["id"], "engine": "global",
                                         "start_in_chat": True})).json()
    if action == "OLD_SESSION":
        from app.db import connection

        with connection() as db:
            db.execute("BEGIN IMMEDIATE")
            row, state = adapter._row_state(db, "alpha", session["id"], session["version"])
            state["global_compilation"].pop("option_authorization_revision")
            session = sessions.save_state(db, row, state, "test_previous_revision")
    session = await resolve(api, session)
    slots = session["chat"]["question"]["slots"]
    authorization = [s for s in slots if s["key"].startswith("authorization:")]
    branch = [s for s in slots if s["key"] == "condition:c_societa"]
    cells = sorted(f["candidate_id"] for f in session["global_compilation"]["fields"])
    if action == "OLD_SESSION":
        assert not authorization  # Previous sessions keep their V13 behaviour.
    else:
        assert len(authorization) == 1
        assert "attività coerente" in authorization[0]["question"]
        assert "01234560981" in authorization[0]["question"]
    replies, message = [], []
    if branch_condition:
        assert branch
        replies.append(ClarificationReply(slot=branch[0]["slot"], action="CONDITION_TRUE",
                                          user_quote="Sì, Mapi è una società"))
        message.append("Sì, Mapi è una società")
    if replies:
        session, errors = dialogue.apply_replies("alpha", session["id"], session["version"],
                                                 replies, " ".join(message))
        assert not errors
        session = await resolve(api, session)
    ids = dialogue.export_selection(session)[0]
    if action == "OLD_SESSION":
        assert ids == cells
        return
    assert not set(cells) & set(ids)  # Branch confirmed or not, the attestation is missing.
    bypass = await api.post(f"{BASE}/{session['id']}/global-draft", json={
        "version": session["version"], "candidate_ids": cells, "allow_unresolved": True,
        "confirmed_condition_ids": ["c_societa"] if branch_condition else [],
        "user_quote": "Genera la bozza"})
    assert bypass.status_code in {409, 422}
    slot = next(s for s in session["chat"]["question"]["slots"]
                if s["key"].startswith("authorization:"))
    quote = ASSENT if action == "CONDITION_TRUE" else "Rinvia questa dichiarazione"
    session, errors = dialogue.apply_replies("alpha", session["id"], session["version"], [
        ClarificationReply(slot=slot["slot"], action=action, user_quote=quote)], quote)
    assert not errors
    session = await resolve(api, session)
    ids, conditions, _ = dialogue.export_selection(session)
    assert set(ids) == (set(cells) if action == "CONDITION_TRUE" else set())
    if action == "CONDITION_TRUE":
        output = await api.post(f"{BASE}/{session['id']}/global-draft", json={
            "version": session["version"], "candidate_ids": ids, "allow_unresolved": True,
            "confirmed_condition_ids": conditions, "user_quote": "Genera la bozza parziale"})
        assert output.status_code == 200, output.text
        downloads = output.json()["last_generation"]["downloads"]
        written = Document(BytesIO((await api.get(downloads["docx"])).content))
        assert {"Bari", "01234560981"} <= {c.text for c in written.tables[0].rows[0].cells}
