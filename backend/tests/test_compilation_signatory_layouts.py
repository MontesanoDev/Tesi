"""Same selected person, different FORM layouts; no provider or company inference."""

import copy
import json
from io import BytesIO
from pathlib import Path

import pytest
from docx import Document
from test_compilation_selected_person import NAME, PROFILE

from app.compilation_global import build_global_input, validate_global_proposal
from app.compilation_global_sessions import initial_state
from app.compilation_source_reuse import reuse_selected_person
from app.docx_templates import inspect_docx


def prepare(doc):
    stream = BytesIO()
    doc.save(stream)
    layout = inspect_docx(stream.getvalue(), require_candidates=False)
    sources = [{"chunk_id": -7, "file_id": -3, "chunk_index": 0, "source_name": "persone.md",
        "scope": "global", "role": "source", "category": "company", "project_id": None,
        "content": PROFILE}]
    raw = {"version": 12, "fields": {}, "facts": {}, "conditions": {}, "bindings": {},
           "choice_groups": {}, "clarifications": {}, "coverage": []}
    data = initial_state(layout)
    data.update(catalog=build_global_input(layout, "alpha", sources)["form"],
        audit=validate_global_proposal(json.dumps(raw), layout, "alpha", sources),
        user_context=[{"kind": "recorded_request", "provenance": "USER", "turn_id": 18,
            "user_quote": "Per questa domanda il firmatario scelto è Ada Neri."}],
        dialogue={"answers": {}, "user_values": {}})
    return data, layout, sources


def proposals(case):
    data, layout, sources = case
    return {f["candidate_id"]: f for f in reuse_selected_person(
        data, layout, "alpha", sources)["fields"] if f.get("source_reuse")}


def inline():
    doc = Document()
    doc.add_paragraph("Il/la sottoscritto/a ________ nato/a a ________, il ________, "
                      "C.F. ________, "
                      "residente a ________ (________) via/piazza ________ n. ________")
    return doc


def test_eight_personal_fields_recovered_from_partial_inline_grammar():
    data = prepare(inline())
    p = proposals(data)
    assert {cid: f["proposed_value"] for cid, f in p.items()} == {
        "p0.s0": NAME, "p0.s1": "Torino (TO)", "p0.s2": "12 gennaio 1985",
        "p0.s3": "NREADA85A52L219R", "p0.s4": "Milano", "p0.s5": "MI",
        "p0.s6": "Roma", "p0.s7": "24"}
    assert all(f["source_reuse"]["identity_evidence"] for f in p.values())
    assert all(f["entity"] == "person" and f["source_evidence"] for f in p.values())


def table_document(*, separate=False):
    doc = Document()
    rows = [("Il/La sottoscritto/a", ""), ("Data e luogo di nascita", ""),
            ("Codice fiscale", ""), ("Residenza", "")]
    for i, (label, _) in enumerate(rows):
        if i == 0 or separate:
            table = doc.add_table(rows=0, cols=2)
        row = table.add_row()
        row.cells[0].text = label
    return doc


@pytest.mark.parametrize("separate", [False, True])
def test_signatory_properties_can_span_label_value_rows_and_adjacent_tables(separate):
    p = proposals(prepare(table_document(separate=separate)))
    assert len(p) == 2
    assert {f["proposed_value"] for f in p.values()} == {NAME, "NREADA85A52L219R"}
    # Compound cells are recognised, but no concatenated or half value is cited
    # as if it were literally present in a single SOURCE assertion.


def test_properties_on_continuation_paragraphs_belong_only_to_the_anchored_signatory():
    doc = Document()
    doc.add_paragraph("Il sottoscritto ________")
    doc.add_paragraph("nato a ________ il ________")
    doc.add_paragraph("C.F. ________")
    doc.add_paragraph("Direttore tecnico")
    doc.add_paragraph("Codice fiscale ________")
    p = proposals(prepare(doc))
    assert set(p) == {"p0.s0", "p1.s0", "p1.s1", "p2.s0"}


@pytest.mark.parametrize("quote", ["Il rappresentante legale è Ada Neri.",
    "Forse il firmatario scelto è Ada Neri.", "Per un'altra gara il firmatario scelto è Ada Neri.",
    "Il firmatario scelto è Ada Neri oppure Bea Mori.",
    "Non confermo. Il firmatario scelto è Ada Neri?", "Il firmatario scelto è Ada."])
def test_unrelated_uncertain_or_partial_user_context_is_not_a_selection(quote):
    data = prepare(inline())
    data[0]["user_context"][0]["user_quote"] = quote
    assert not proposals(data)


@pytest.mark.parametrize("change", ["assistant", "foreign_kind", "no_turn", "conflict",
                                    "homonym", "general_source", "company_only", "missing_cf"])
def test_identity_and_provenance_remain_strict(change):
    data = prepare(inline())
    ctx = data[0]["user_context"]
    if change == "assistant":
        ctx[0]["provenance"] = "MODEL"
    elif change == "foreign_kind":
        ctx[0]["kind"] = "unrelated_history"
    elif change == "no_turn":
        ctx[0].pop("turn_id")
    elif change == "conflict":
        ctx.append({**ctx[0], "turn_id": 19,
                    "user_quote": "Per questa domanda il firmatario scelto è Bea Mori."})
    elif change == "homonym":
        data[2].append({**data[2][0], "chunk_id": -8, "file_id": -4,
                        "content": PROFILE.replace("85A52", "86A52")})
    elif change == "general_source":
        data[2][0]["category"] = "general"
    elif change == "company_only":
        data[2][0]["content"] = "Società Ada Neri S.r.l.; CF/PIVA IT01234567890"
    else:
        data[2][0]["content"] = PROFILE.replace("- Codice fiscale personale: NREADA85A52L219R.", "")
    assert not proposals(data)


def test_missing_personal_data_are_not_replaced_with_company_or_other_person_properties():
    data = prepare(inline())
    data[2][0]["content"] = data[2][0]["content"].replace("- Data di nascita: 12 gennaio 1985.", "")
    data[2].append({**data[2][0], "chunk_id": -8, "file_id": -4, "content":
        "Società: Alba; data costituzione: 12 gennaio 1985\n" + PROFILE.replace(
            "Ada Neri", "Bea Mori").replace("Nome: Ada", "Nome: Bea").replace(
            "Cognome: Neri", "Cognome: Mori").replace("85A52", "86A52")})
    assert "p0.s2" not in proposals(data)


def test_operator_transition_does_not_transfer_company_cf_or_address_to_signatory():
    doc = Document()
    doc.add_paragraph("Il sottoscritto ________ nato a ________ il ________ "
                      "della società ________ con codice fiscale ________ con sede in ________")
    assert set(proposals(prepare(doc))) == {"p0.s0", "p0.s1", "p0.s2"}


def test_no_hidden_applicability_or_state_mutation():
    data = prepare(inline())
    data[0]["audit"]["fields"][0].update(conditions=["unknown"])
    data[0]["audit"]["conditions"]["unknown"] = {"result": "unknown"}
    before = copy.deepcopy(data[0])
    p = proposals(data)
    assert p and all(f["condition_checks"]["unknown"]["result"] == "unknown" for f in p.values())
    assert data[0] == before


@pytest.mark.parametrize("boundary", ["Direttore tecnico", "Società", "Un'altra persona",
                                     "Altri dati", "Se procuratore:"])
def test_owner_does_not_cross_an_unproved_heading_or_other_subject(boundary):
    doc = Document()
    doc.add_paragraph("Il sottoscritto ________")
    doc.add_paragraph(boundary)
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Codice fiscale"
    p = proposals(prepare(doc))
    assert set(p) == {"p0.s0"}


def test_global_choice_does_not_select_one_of_two_different_signatory_blocks():
    doc = inline()
    doc.add_paragraph("Altra dichiarazione")
    doc.add_paragraph("Il sottoscritto ________ nato a ________ il ________")
    assert not proposals(prepare(doc))


@pytest.mark.parametrize("separate", [False, True])
def test_conditional_signatory_block_does_not_become_applicable_from_a_name(separate):
    doc = Document()
    if separate:
        doc.add_paragraph("Se procuratore:")
    doc.add_paragraph(("" if separate else "Se procuratore, ") + "il sottoscritto ________")
    doc.add_paragraph("Codice fiscale ________")
    assert not proposals(prepare(doc))


def test_explicit_skip_of_signatory_overrides_initial_selection():
    data = prepare(inline())
    data[0]["dialogue"]["answers"]["field:p0.s0"] = {
        "action": "SKIP", "provenance": "USER", "user_quote": "Lascia vuoto il firmatario"}
    assert not proposals(data)


@pytest.mark.parametrize("mutation", ["label_editable", "nested", "vertical_merge"])
def test_complex_table_is_a_boundary(mutation):
    doc = table_document()
    data, layout, sources = prepare(doc)
    table = data["catalog"]["blocks"][0]["tables"][0]
    if mutation == "label_editable":
        row = table["rows"][1]
        data["catalog"]["candidates"]["extra"] = {"context_id": row[0]["paragraph_ids"][0]}
    elif mutation == "nested":
        table["rows"][1][0]["table_ids"] = ["nested"]
    else:
        table["rows"][1][0]["vertical_merge"] = "restart"
    assert set(proposals((data, layout, sources))) == {"t0.r0.c1"}


@pytest.mark.parametrize("slug,expected", [("trapani", 8), ("bella", 4),
    ("catanzaro", 2), ("minervino-v12-61", 9)])
def test_four_actual_forms_use_the_same_identity_and_source_reuse_mechanism(slug, expected):
    fixtures = Path(__file__).parent / "fixtures"
    captured = json.loads((fixtures / "signatory-layouts-v1/cases.json").read_text())[slug]
    form_slug = "minervino" if slug == "minervino-v12-61" else slug
    layout = inspect_docx((fixtures / f"formtext-v1/{form_slug}/original.docx").read_bytes())
    raw = validate_global_proposal(json.dumps(captured["proposal"]), layout,
                                   captured["project_id"], captured["sources"])
    data = {**captured["state"], "audit": raw, "user_context": captured["user_context"]}
    before = copy.deepcopy(data)
    result = reuse_selected_person(data, layout, captured["project_id"], captured["sources"])
    recovered = [f for f in result["fields"] if f.get("source_reuse")]
    assert len(recovered) == expected
    assert all(f["source_evidence"] and f["source_reuse"]["identity_evidence"]
               and f["entity"] == "person" for f in recovered)
    assert data == before and result["resolved"] == result["writes"] == 0
    if slug == "trapani":
        assert {f["candidate_id"]: f["proposed_value"] for f in recovered} == {
            "p12.s0": "Leonardo Montemurro", "p12.s1": "Lecce (LE)",
            "p12.s2": "14 novembre 1982", "p12.s3": "MNTLRD82S14E506X",
            "p12.s4": "Bari", "p12.s5": "BA", "p12.s6": "Amedeo", "p12.s7": "24"}
    elif slug == "bella":
        assert all(f["candidate_id"] in {"p21.s0", "p21.s1", "p21.s2", "p21.s3"}
                   for f in recovered)  # Company CF/PIVA never become personal ones.
    elif slug == "catanzaro":
        assert {f["candidate_id"] for f in recovered} == {"t0.r0.c1", "t0.r2.c1"}
    data["user_context"] = []
    data["dialogue"] = {"answers": {}, "user_values": {}}
    without = reuse_selected_person(data, layout, captured["project_id"], captured["sources"])
    assert not any(f.get("source_reuse") for f in without["fields"])


@pytest.mark.parametrize("retraction", ["Non confermo questa scelta.",
    "Non ho ancora scelto chi firma.", "Il firmatario è ancora da decidere.",
    "Annullo la scelta del firmatario."])
def test_initial_selection_and_retraction_are_not_silently_resolved(retraction):
    data = prepare(inline())
    data[0]["user_context"][0]["user_quote"] += " " + retraction
    assert not proposals(data)


def test_incomplete_legacy_catalog_does_not_break_review_or_create_targets():
    from app.compilation_signatory_form import signatory_groups

    catalog = {"candidates": {"p0.s0": {"id": "p0.s0"}}, "blocks": [
        {"id": "paragraph:0", "paragraphs": [{"id": "paragraph:0", "text": "Dato: ______"}]}]}
    assert signatory_groups(catalog) == []


@pytest.mark.parametrize("entity,subject,value,proof", [
    ("company", "Alba S.r.l.", "01234567890", "CF della società Alba S.r.l.: 01234567890"),
    ("person", "Bea Mori", "MROBEA86A52F205L", "Firmatario: Bea Mori; CF: MROBEA86A52F205L"),
])
def test_preexisting_grounded_fact_cannot_override_selected_signatory_identity(
    entity, subject, value, proof,
):
    doc = Document()
    doc.add_paragraph("Il sottoscritto ________")
    doc.add_paragraph("Codice fiscale ________")
    data, layout, sources = prepare(doc)
    sources.append({**sources[0], "chunk_id": -8, "file_id": -4, "content": proof})
    field = next(f for f in data["audit"]["fields"] if f["candidate_id"] == "p1.s0")
    # Reproduce an independently grounded proposal; changing the selected person
    # must add a constraint even when the old fact/proof itself remains valid.
    data["audit"]["facts"]["other"] = {"entity": entity, "subject": subject,
        "property": "cf", "value": value, "validation": "GROUNDED_PROPOSAL"}
    field.update(fact_id="other", validation="GROUNDED_PROPOSAL", value=value,
        disposition="value", kind="data", form_evidence={"paragraph_id": "paragraph:1",
        "quote": "Codice fiscale"}, proposal={"fact_id": "other", "disposition": "value"},
        source_evidence=[{"chunk_id": -8, "quote": proof}])
    result = reuse_selected_person(data, layout, "alpha", sources)
    rejected = next(f for f in result["fields"] if f["candidate_id"] == "p1.s0")
    assert rejected["validation"] == "REJECTED"
    assert rejected.get("value") is None and not rejected.get("proposal_grounding")
    assert result["facts"]["other"]["validation"] == "GROUNDED_PROPOSAL"


def test_rejected_existing_subjects_update_counts_even_when_nothing_is_reused():
    doc = Document()
    doc.add_paragraph("Il sottoscritto ________")
    data, layout, sources = prepare(doc)
    data["audit"]["facts"]["other"] = {"entity": "person", "subject": "Bea Mori",
        "value": "Bea Mori", "validation": "GROUNDED_PROPOSAL"}
    data["audit"]["fields"][0].update(validation="GROUNDED_PROPOSAL", value="Bea Mori",
        fact_id="other", proposal={"fact_id": "other", "disposition": "value"})
    result = reuse_selected_person(data, layout, "alpha", sources)
    assert result["counts"] == {"REJECTED": 1}
    assert not result.get("source_reuse_count")
