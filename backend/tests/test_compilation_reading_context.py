"""Physical reading contexts and strict local anchors, without an AI provider."""

import json
from io import BytesIO
from zipfile import ZipFile

import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from test_compilation_catalog import encoded, native_checkbox

from app.compilation_catalog import MAX_PRECEDING_PARAGRAPHS, build_compilation_catalog
from app.compilation_global import build_global_input, validate_global_proposal
from app.docx_templates import inspect_docx


def data_for(doc):
    layout = inspect_docx(encoded(doc), require_candidates=False)
    return layout, build_global_input(layout, "alpha", [])


def mark_header(row, value=None):
    marker = OxmlElement("w:tblHeader")
    if value is not None:
        marker.set(qn("w:val"), value)
    row._tr.get_or_add_trPr().append(marker)


def test_first_data_row_is_not_a_header_or_local_anchor_of_later_rows():
    doc = Document()
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Persona firmataria"
    table.cell(1, 0).text = "Organizzazione rappresentata"
    _, data = data_for(doc)
    form = data["form"]
    assert form["version"] == 5
    assert form["candidates"]["t0.r1.c1"]["anchor_paragraph_ids"] == ["paragraph:2"]
    location = form["candidates"]["t0.r1.c1"]["location"]
    assert location == {"block_id": "paragraph:0", "table_id": "table:0", "row_index": 1}
    assert "reading_contexts" not in form


@pytest.mark.parametrize("marker,expected", [(None, True), ("true", True), ("1", True),
                                           ("on", True), ("false", False), ("0", False),
                                           ("off", False)])
def test_only_explicit_enabled_headers_supply_header_context_and_anchors(marker, expected):
    doc = Document()
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Registro"
    table.cell(0, 1).text = "Numero"
    table.cell(1, 0).text = "Registro categoria Alfa"
    mark_header(table.rows[0], marker)
    _, data = data_for(doc)
    field = data["form"]["candidates"]["t0.r1.c1"]
    assert ("paragraph:1" in field["anchor_paragraph_ids"]) is expected
    table_data = data["form"]["blocks"][0]["tables"][0]
    assert table_data.get("repeat_header_rows", []) == ([0] if expected else [])
    assert field["location"]["row_index"] == 1


@pytest.mark.parametrize("header", [False, True])
def test_header_grounding_is_explicit_and_cannot_reuse_an_unrelated_first_row_label(header):
    doc = Document()
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Valore categoria Alfa"
    table.cell(0, 1).text = "Identificativo"
    table.cell(1, 0).text = "Categoria Beta"
    if header:
        mark_header(table.rows[0])
    layout, _ = data_for(doc)
    source = {"chunk_id": -1, "file_id": -1, "chunk_index": 0, "source_name": "profilo",
              "role": "source", "scope": "global", "category": "company",
              "project_id": None, "content": "Identificativo: AZ-98765"}
    payload = {"fields": {"t0.r1.c1": {
        "label": "Identificativo", "entity": "company", "kind": "data",
        "disposition": "value", "value": "AZ-98765",
        "form": {"paragraph_id": "paragraph:1", "quote": "Identificativo"},
        "evidence": [{"chunk_id": -1, "quote": source["content"]}],
        "reason": "Solo grounding locale, nessuna accettazione semantica",
    }}}
    audit = validate_global_proposal(json.dumps(payload), layout, "alpha", [source])
    field = audit["fields"][0]
    assert field["validation"] == ("GROUNDED_PROPOSAL" if header else "REJECTED")
    assert field["semantic_review_required"] and field["written_value"] is None
    assert audit["writes"] == audit["resolved"] == 0


def test_preceding_context_is_ordered_bounded_and_never_a_local_label_anchor():
    doc = Document()
    for index in range(10):
        doc.add_paragraph(f"Istruzione precedente {index}")
    doc.add_paragraph()
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Valore locale"
    _, data = data_for(doc)
    form = data["form"]
    preceding = form["blocks"][-1]["preceding_body_paragraph_ids"]
    assert preceding == [f"paragraph:{i}" for i in range(10 - MAX_PRECEDING_PARAGRAPHS, 10)]
    assert not set(preceding) & set(form["candidates"]["t0.r0.c1"]["anchor_paragraph_ids"])
    # No paragraph text is duplicated in the new shared contexts.
    assert json.dumps(form, ensure_ascii=False).count("Istruzione precedente 9") == 1


def test_visual_bold_titles_are_observations_without_invented_sections_or_conditions():
    doc = Document()
    doc.add_paragraph().add_run("Titolo visuale categoria Gamma").bold = True
    doc.add_paragraph("Identificativo: __________")
    doc.add_paragraph().add_run("Titolo visuale categoria Delta").bold = True
    doc.add_paragraph("Identificativo: __________")
    _, data = data_for(doc)
    form = data["form"]
    assert form["blocks"][0]["paragraphs"][0]["all_text_direct_bold"] is True
    assert form["blocks"][2]["paragraphs"][0]["all_text_direct_bold"] is True
    assert form["candidates"]["p1.s0"]["section_ids"] == []
    assert form["candidates"]["p3.s0"]["section_ids"] == []
    assert form["candidates"]["p1.s0"]["location"] == {"block_id": "paragraph:1"}
    assert form["candidates"]["p3.s0"]["location"] == {"block_id": "paragraph:3"}
    assert "applicability" not in json.dumps(form)


@pytest.mark.parametrize("kind", ["mixed", "false", "empty", "inherited"])
def test_bold_signal_does_not_guess_from_partial_emphasis_or_style(kind):
    doc = Document()
    p = doc.add_paragraph()
    if kind == "mixed":
        p.add_run("Parte evidenziata").bold = True
        p.add_run(" e parte normale")
    elif kind == "false":
        p.add_run("Normale").bold = False
    elif kind == "inherited":
        p.style = doc.styles["Heading 1"]
        p.add_run("Titolo ereditato")
    else:
        p.add_run().bold = True
    result = build_compilation_catalog(inspect_docx(encoded(doc), require_candidates=False))
    assert "all_text_direct_bold" not in result["blocks"][0]["paragraphs"][0]


def test_nested_table_has_its_own_row_and_parent_path_without_outer_label_contamination():
    doc = Document()
    outer = doc.add_table(rows=1, cols=2)
    outer.cell(0, 0).text = "Soggetto esterno"
    inner = outer.cell(0, 1).add_table(rows=1, cols=2)
    inner.cell(0, 0).text = "Soggetto interno"
    layout, data = data_for(doc)
    form = data["form"]
    tables = form["blocks"][0]["tables"]
    assert tables[1]["parent_cell_id"] == "t0.r0.c1"
    field = form["candidates"]["t1.r0.c1"]
    assert field["location"] == {"block_id": "paragraph:0", "table_id": "table:1",
                                 "row_index": 0}
    assert "paragraph:0" not in field["anchor_paragraph_ids"]
    assert set(form["candidates"]) == layout.candidate_ids


def test_checkbox_and_text_positions_share_physical_row_but_not_writable_identity():
    doc = Document()
    table = doc.add_table(rows=1, cols=3)
    native_checkbox(table.cell(0, 0).paragraphs[0])
    table.cell(0, 1).text = "Opzione con valore associato"
    _, data = data_for(doc)
    form = data["form"]
    assert form["candidates"]["p0.cb0"]["location"] == (
        form["candidates"]["t0.r0.c2"]["location"])
    assert "reading_contexts" not in form
    assert form["candidates"]["p0.cb0"]["type"] == "checkbox"
    assert form["checkbox_write_supported"] is True


def test_merged_cells_keep_physical_rows_and_do_not_create_duplicate_targets():
    doc = Document()
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).merge(table.cell(0, 1)).text = "Avvertenza comune"
    table.cell(1, 0).text = "Dato richiesto"
    layout, data = data_for(doc)
    assert set(data["form"]["candidates"]) == layout.candidate_ids == {"t0.r1.c1"}
    assert data["form"]["candidates"]["t0.r1.c1"]["anchor_paragraph_ids"] == ["paragraph:1"]


def test_header_cannot_cross_table_boundaries():
    doc = Document()
    first = doc.add_table(rows=1, cols=2)
    first.cell(0, 0).text = "Intestazione tabella Alfa"
    first.cell(0, 1).text = "Valore Alfa"
    mark_header(first.rows[0])
    second = doc.add_table(rows=1, cols=2)
    second.cell(0, 0).text = "Dato Beta"
    _, data = data_for(doc)
    assert data["form"]["candidates"]["t1.r0.c1"]["anchor_paragraph_ids"] == ["paragraph:2"]


def test_word_namespace_prefix_is_not_assumed_for_visual_or_header_metadata():
    doc = Document()
    doc.add_paragraph().add_run("Titolo visuale").bold = True
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Categoria"
    table.cell(0, 1).text = "Identificativo"
    table.cell(1, 0).text = "Categoria Alfa"
    mark_header(table.rows[0])
    stream = BytesIO()
    with ZipFile(BytesIO(encoded(doc))) as original, ZipFile(stream, "w") as output:
        for name in original.namelist():
            data = original.read(name)
            if name == "word/document.xml":
                data = (data.replace(b"xmlns:w=", b"xmlns:alias=")
                        .replace(b"<w:", b"<alias:").replace(b"</w:", b"</alias:")
                        .replace(b" w:", b" alias:"))
            output.writestr(name, data)
    layout = inspect_docx(stream.getvalue())
    form = build_global_input(layout, "alpha", [])["form"]
    assert form["blocks"][0]["paragraphs"][0]["all_text_direct_bold"] is True
    assert form["blocks"][1]["tables"][0]["repeat_header_rows"] == [0]
    assert "paragraph:2" in form["candidates"]["t0.r1.c1"]["anchor_paragraph_ids"]
