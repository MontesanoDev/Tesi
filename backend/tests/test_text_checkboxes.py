"""Text checkbox regression: real marker, isolated profile, no inference."""
import json
from io import BytesIO
from zipfile import ZipFile

import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from test_compilation_catalog import encoded, native_checkbox

from app.compilation_catalog import build_compilation_catalog
from app.compilation_global_sessions import initial_state
from app.docx_templates import DocumentInputError, fill_docx, inspect_docx


def form_bytes(text='[] di avvalersi dei requisiti della seguente impresa ausiliaria:'):
    doc = Document()
    doc.add_paragraph(text)
    doc.add_paragraph('Denominazione: ________')
    out = BytesIO()
    doc.save(out)
    return out.getvalue()


def experimental_layout(monkeypatch, original=None):
    monkeypatch.setenv('MAPI_COMPILATION_PROFILE', 'checkbox-testuali-2026-10-26')
    layout = inspect_docx(original or form_bytes(), require_candidates=False)
    initial_state(layout)
    return layout


def test_avvalimento_marker_has_a_separate_destination(monkeypatch):
    layout = experimental_layout(monkeypatch)
    catalog = build_compilation_catalog(layout)
    assert catalog['candidates']['p0.txcb0']['control_type'] == 'text_ascii'
    assert catalog['blocks'][0]['paragraphs'][0]['text_with_choices'].startswith('[[p0.txcb0]][]')


def test_writer_marks_only_the_authorized_text_marker(monkeypatch):
    layout = experimental_layout(monkeypatch)
    result = fill_docx(layout, {}, checkbox_values={'p0.txcb0': True})
    assert Document(BytesIO(result)).paragraphs[1].text.startswith('[X] di avvalersi')
    assert Document(BytesIO(layout.original)).paragraphs[0].text.startswith('[] di avvalersi')


@pytest.mark.parametrize('marker,selected,expected', [
    ('[]', True, '[X]'), ('[ ]', True, '[X]'), ('[x]', False, '[ ]'),
    ('[X]', False, '[ ]'), ('[]', False, '[]'), ('[ ]', False, '[ ]'),
])
@pytest.mark.parametrize('table', [False, True])
def test_split_runs_cells_and_other_package_parts_stay_intact(
    monkeypatch, marker, selected, expected, table,
):
    doc = Document()
    p = doc.add_table(rows=1, cols=1).cell(0, 0).paragraphs[0] if table else doc.add_paragraph()
    p.add_run(marker[0]).bold = True
    p.add_run(marker[1:])
    p.add_run(' Scelta esplicita: ________')
    original = encoded(doc)
    layout = experimental_layout(monkeypatch, original)
    result = fill_docx(layout, {'p0.s0': 'Dato'}, checkbox_values={'p0.txcb0': selected})
    output = Document(BytesIO(result))
    p = output.tables[0].cell(0, 0).paragraphs[0] if table else output.paragraphs[1]
    assert p.text == expected + ' Scelta esplicita: Dato'
    assert p.runs[0].bold is True
    assert layout.original == original
    with ZipFile(BytesIO(original)) as a, ZipFile(BytesIO(result)) as b:
        assert all(a.read(n) == b.read(n) for n in a.namelist() if n != 'word/document.xml')
    old = build_compilation_catalog(inspect_docx(original, require_candidates=False))
    assert 'p0.txcb0' not in old['candidates']


@pytest.mark.parametrize('text', ['Testo [] ordinario', '[]', '[] prima [] seconda',
                                 '☐ simbolo', '□ simbolo', '[nota] testo'])
def test_ambiguous_or_nonleading_text_is_not_a_checkbox(monkeypatch, text):
    layout = experimental_layout(monkeypatch, form_bytes(text))
    assert 'p0.txcb0' not in build_compilation_catalog(layout)['candidates']
    with pytest.raises(DocumentInputError):
        fill_docx(layout, {}, checkbox_values={'p0.txcb0': True})


def test_native_control_ids_and_symbol_elements_remain_distinct(monkeypatch):
    doc = Document()
    p = doc.add_paragraph()
    native_checkbox(p)
    p.add_run(' Nativa')
    p = doc.add_paragraph('[] Testuale')
    p = doc.add_paragraph('[] Simbolo')
    sym = OxmlElement('w:sym')
    sym.set(qn('w:font'), 'Wingdings')
    sym.set(qn('w:char'), 'F0A8')
    p.add_run()._r.append(sym)
    layout = experimental_layout(monkeypatch, encoded(doc))
    catalog = build_compilation_catalog(layout)
    assert set(catalog['candidates']) == {'p0.cb0', 'p1.txcb0'}
    result = fill_docx(layout, {}, checkbox_values={'p0.cb0': True, 'p1.txcb0': True})
    assert Document(BytesIO(result)).paragraphs[3].text == '[] Simbolo'


def test_session_revision_controls_reopening_not_current_environment(monkeypatch):
    from app.compilation_sessions import read_original

    layout = experimental_layout(monkeypatch)
    state = {'engine': 'global', 'template_sha256': layout.sha256, 'global_compilation': {}}
    row = {'original': layout.original, 'state_json': json.dumps(state)}
    assert read_original(row).text_checkbox_revision == 0
    state['global_compilation']['text_checkbox_revision'] = 1
    row['state_json'] = json.dumps(state)
    monkeypatch.setenv('MAPI_COMPILATION_PROFILE', 'storico')
    assert 'p0.txcb0' in build_compilation_catalog(read_original(row))['candidates']


def test_new_controls_do_not_remove_the_existing_option_scope(monkeypatch):
    from test_compilation_form_option_scope import form

    from app.compilation_form_option_scope import regions

    original = form()
    before = regions(build_compilation_catalog(inspect_docx(original)), 2)
    after = regions(build_compilation_catalog(experimental_layout(monkeypatch, original)), 2)
    assert before and after == before


def test_omitted_form_condition_blocks_even_a_checkbox_only_region(monkeypatch):
    from app.compilation_form_option_scope import blocked_fields, regions

    doc = Document()
    doc.add_paragraph().add_run('(da compilare nel caso di partecipazione in RT)').bold = True
    doc.add_paragraph('[] di avvalersi dei requisiti della seguente impresa ausiliaria:')
    layout = experimental_layout(monkeypatch, encoded(doc))
    catalog = build_compilation_catalog(layout)
    assert regions(catalog, 2)[0]['candidate_ids'] == ['p1.txcb0']
    data = {'generation_revision': 3, 'form_option_scope_revision': 2,
            'catalog': catalog, 'audit': {'fields': [], 'conditions': {}}}
    assert blocked_fields(data) == {'p1.txcb0'}


def test_enabling_a_marker_preserves_existing_data_context(monkeypatch):
    original = form_bytes()
    before = build_compilation_catalog(inspect_docx(original))
    after = build_compilation_catalog(experimental_layout(monkeypatch, original))
    assert after['candidates']['p1.s0'] == before['candidates']['p1.s0']


def test_source_proposal_and_partial_clause_cannot_select_text_checkbox(monkeypatch):
    import copy

    from app.compilation_global import validate_global_proposal

    layout = experimental_layout(monkeypatch)
    paragraph = build_compilation_catalog(layout)['blocks'][0]['paragraphs'][0]
    label = paragraph['text'][2:].strip()
    form = {'paragraph_id': 'paragraph:0', 'quote': paragraph['text_with_choices']}
    source = {'chunk_id': -1, 'file_id': -1, 'chunk_index': 0, 'source_name': 'fonte.txt',
              'role': 'source', 'scope': 'global', 'category': 'company', 'project_id': None,
              'content': 'Impresa: Vega\nScelta: ' + label}
    group = {'label': 'Scelta', 'subject': 'Vega', 'property': 'scelta', 'selection': 'multiple',
             'min_selected': 0, 'conditions': [],
             'form': {'paragraph_id': 'paragraph:0', 'quote': paragraph['text']},
             'options': {'p0.txcb0': {'label': label, 'form': form}}}
    raw = {'version': 12, 'uncertain_paragraph_ids': [], 'conditions': {}, 'bindings': {},
           'choice_groups': {'g': group}, 'coverage': [], 'facts': {'f': {
               'subject': 'Vega', 'entity': 'company', 'aspect': 'legal_type',
               'property': 'scelta', 'value': label,
               'evidence': [{'chunk_id': -1, 'quote': 'Scelta: ' + label}]}},
           'fields': {'p0.txcb0': {'label': label, 'kind': 'choice', 'disposition': 'value',
                'form': form, 'conditions': [], 'fact_id': 'f', 'property': 'scelta',
                'choice': True, 'choice_group_id': 'g'}}}
    result = validate_global_proposal(json.dumps(raw), layout, 'alpha', [source])
    field = next(f for f in result['fields'] if f['candidate_id'] == 'p0.txcb0')
    assert result['choice_groups']['g']['validation'] == 'GROUNDED_GROUP'
    assert field['validation'] == 'REJECTED'
    assert any('scelta USER' in e for e in field['validation_errors'])
    partial = copy.deepcopy(raw)
    partial['choice_groups']['g']['options']['p0.txcb0']['form']['quote'] = (
        '[[p0.txcb0]][] di avvalersi')
    result = validate_global_proposal(json.dumps(partial), layout, 'alpha', [source])
    assert result['choice_groups']['g']['validation'] == 'REJECTED'


def test_text_checkboxes_do_not_change_the_prompt(monkeypatch):
    from app.compilation_global import TASK, build_global_input, global_request
    from app.config import AISettings

    settings = AISettings(api_key='fake', model='fake', base_url='https://example.test',
                          thinking=False)
    original = form_bytes()
    normal = build_global_input(inspect_docx(original), 'alpha', [])
    assert global_request(settings, normal)['messages'][0]['content'] == TASK
    experimental = build_global_input(experimental_layout(monkeypatch, original), 'alpha', [])
    assert global_request(settings, experimental)['messages'][0]['content'] == TASK


@pytest.mark.parametrize('separator,accepted', [('OPPURE', True), ('oppure', True),
                                               ('Dati di un altro soggetto', False)])
def test_explicit_alternatives_form_one_complete_group(monkeypatch, separator, accepted):
    from app.compilation_global import validate_global_proposal

    doc = Document()
    doc.add_paragraph('[ ] di possedere in proprio i requisiti richiesti;')
    doc.add_paragraph(separator)
    doc.add_paragraph('[] di avvalersi dei requisiti della seguente impresa ausiliaria:')
    layout = experimental_layout(monkeypatch, encoded(doc))
    catalog = build_compilation_catalog(layout)
    options = {}
    for n in (0, 2):
        p = catalog['blocks'][n]['paragraphs'][0]
        cid = f'p{n}.txcb0'
        options[cid] = {'label': p['text'].split(']', 1)[1].strip(),
                        'form': {'paragraph_id': p['id'], 'quote': p['text_with_choices']}}
    payload = {'version': 12, 'conditions': {}, 'facts': {}, 'fields': {}, 'coverage': [],
               'choice_groups': {'g': {'label': 'Requisiti', 'subject': 'Concorrente',
                    'property': 'requisiti', 'selection': 'single', 'min_selected': 0,
                    'form': {'paragraph_id': 'paragraph:0',
                             'quote': catalog['blocks'][0]['paragraphs'][0]['text']},
                    'conditions': [], 'options': options}}}
    result = validate_global_proposal(json.dumps(payload), layout, 'alpha', [])
    assert (result['choice_groups']['g']['validation'] == 'GROUNDED_GROUP') is accepted


def test_inserted_data_cannot_create_a_checkbox_destination(monkeypatch):
    layout = experimental_layout(monkeypatch, form_bytes('________'))
    with pytest.raises(DocumentInputError):
        fill_docx(layout, {'p0.s0': '[] nuova opzione'}, checkbox_values={'p0.txcb0': True})


def test_original_marker_id_survives_a_preceding_cell_paragraph_removal(monkeypatch):
    doc = Document()
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = 'Dato'
    table.cell(0, 1).add_paragraph('')
    doc.add_paragraph('[] Clausola da selezionare')
    layout = experimental_layout(monkeypatch, encoded(doc))
    assert 'p3.txcb0' in build_compilation_catalog(layout)['candidates']
    result = fill_docx(layout, {'t0.r0.c1': 'Valore'}, checkbox_values={'p3.txcb0': True})
    assert Document(BytesIO(result)).paragraphs[-1].text == '[X] Clausola da selezionare'
