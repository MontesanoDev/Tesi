"""Sparse software regressions; these do not measure real model semantics."""

import copy
import json
from io import BytesIO

import pytest
from docx import Document

from app.compilation_global import build_global_input, validate_global_proposal
from app.docx_templates import inspect_docx


@pytest.fixture
def case():
    doc = Document()
    doc.add_heading("Informazioni dell'organizzazione", level=1)
    table = doc.add_table(rows=3, cols=2)
    labels = ["Identificativo organizzazione", "Recapito principale",
              "Recapito secondario (solo se diverso dal principale)"]
    for index, label in enumerate(labels):
        table.cell(index, 0).text = label
    stream = BytesIO()
    doc.save(stream)
    layout = inspect_docx(stream.getvalue())
    source = {"chunk_id": -7, "file_id": -3, "chunk_index": 0,
              "source_name": "profilo.txt", "role": "source", "scope": "global",
              "category": "company", "project_id": None,
              "content": "Identificativo: AZ-98765; principale: Via Luna 4; "
                         "secondario: Via Luna 4"}
    facts, fields = {}, {}
    for index, (prop, value) in enumerate(zip(labels, ["AZ-98765", "Via Luna 4", "Via Luna 4"],
                                             strict=True)):
        fact_id = f"fact-{index}"
        facts[fact_id] = {
            "subject": "Organizzazione", "entity": "company", "aspect": "location",
            "property": prop, "value": value,
            "evidence": [{"chunk_id": -7, "quote": source["content"]}]}
        fields[f"t0.r{index}.c1"] = {
            "label": prop, "kind": "data", "disposition": "value", "property": prop,
            "fact_id": fact_id, "conditions": [],
            "form": {"paragraph_id": f"paragraph:{index * 2 + 1}", "quote": prop}}
    condition = {
        "form": {"paragraph_id": "paragraph:5", "quote": "solo se diverso dal principale"},
        "aspect": "location", "operator": "different", "fact_ids": ["fact-1", "fact-2"],
        "reason": "Confronto recapiti", "scope": [{"kind": "candidate", "id": "t0.r2.c1"}]}
    return layout, [source], {"version": 5, "facts": facts, "fields": fields,
                              "conditions": {"different": condition}}


def audit(case):
    return validate_global_proposal(json.dumps(case[2]), case[0], "alpha", case[1])


def by_id(result):
    return {field["candidate_id"]: field for field in result["fields"]}


def test_sparse_input_has_no_inventory_receipt_or_partition(case):
    data = build_global_input(case[0], "alpha", case[1])
    assert "form_review_paragraph_ids" not in data
    assert set(data["schema_output"]["properties"]) == {
        "version", "facts", "conditions", "fields", "uncertain_paragraph_ids"}
    assert len(data["form"]["candidates"]) == 3
    assert data["sources"][0]["role"] == "source"


def test_omitted_positions_stay_unassessed_in_canonical_backend_inventory(case):
    case[2]["fields"] = {"t0.r0.c1": case[2]["fields"]["t0.r0.c1"]}
    result = audit(case)
    assert result["candidate_count"] == 3
    assert result["unassigned_count"] == 2
    assert result["counts"] == {"GROUNDED_PROPOSAL": 1, "REVIEW_REQUIRED": 2}
    for field in result["fields"][1:]:
        assert field["proposal"] is None
        assert "value" not in field and "label" not in field
        assert field["written_value"] is None


def test_empty_proposal_is_not_a_completed_document(case):
    case[2].update(fields={}, conditions={}, facts={})
    result = audit(case)
    assert result["counts"] == {"REVIEW_REQUIRED": 3}
    assert result["unassigned_count"] == result["candidate_count"] == 3
    assert result["resolved"] == result["writes"] == 0


@pytest.mark.parametrize("unknown", ["invented", "paragraph:0", "table:0", "p999.s0"])
def test_invented_destinations_are_quarantined_without_discarding_valid_assignments(case, unknown):
    case[2]["fields"][unknown] = copy.deepcopy(case[2]["fields"]["t0.r0.c1"])
    result = audit(case)
    assert len(result["fields"]) == 3
    assert result["rejected_assignments"][unknown]["validation"] == "REJECTED"
    assert result["fields"][0]["validation"] == "GROUNDED_PROPOSAL"
    assert result["candidate_count"] == 3


@pytest.mark.parametrize("change", [{"unexpected": "x"}, {"fact_id": "absent"},
    {"property": "Altra proprietà"}, {"conditions": ["missing-condition"]},
    {"form": {"paragraph_id": "paragraph:3", "quote": "Recapito principale"}},
    {"disposition": "missing"}])
def test_invalid_assignment_is_local_and_does_not_erase_independent_sibling(case, change):
    case[2]["fields"]["t0.r0.c1"].update(change)
    result = by_id(audit(case))
    assert result["t0.r0.c1"]["validation"] == "REJECTED"
    assert result["t0.r1.c1"]["validation"] == "GROUNDED_PROPOSAL"


@pytest.mark.parametrize("change", [{"value": "Inventato"}, {"evidence": []},
    {"evidence": [{"chunk_id": -99, "quote": "AZ-98765"}]}, {"value": "non disponibili"},
    {"extra": True}])
def test_bad_fact_cannot_verify_any_dependent_field(case, change):
    case[2]["facts"]["fact-0"].update(change)
    result = audit(case)
    assert result["facts"]["fact-0"]["validation"] == "REJECTED"
    assert result["fields"][0]["validation"] == "REJECTED"
    assert result["fields"][1]["validation"] == "GROUNDED_PROPOSAL"


def test_orphan_fact_is_reported_and_not_used_as_an_assignment(case):
    case[2]["facts"]["orphan"] = copy.deepcopy(case[2]["facts"]["fact-0"])
    result = audit(case)
    assert result["unused_fact_ids"] == ["orphan"]
    assert "orphan" not in result["facts"]
    assert result["fields"][0]["validation"] == "GROUNDED_PROPOSAL"


@pytest.mark.parametrize("scope,affected", [
    ({"kind": "candidate", "id": "t0.r2.c1"}, ["t0.r2.c1"]),
    ({"kind": "paragraph", "id": "paragraph:5"}, ["t0.r2.c1"]),
    ({"kind": "table", "id": "table:0"}, ["t0.r0.c1", "t0.r1.c1", "t0.r2.c1"]),
    ({"kind": "section", "id": "paragraph:0"}, ["t0.r0.c1", "t0.r1.c1", "t0.r2.c1"]),
])
def test_scope_is_derived_from_real_typed_structural_location(case, scope, affected):
    case[2]["conditions"]["different"]["scope"] = [scope]
    result = audit(case)
    assert result["conditions"]["different"]["applies_to"] == affected
    assert result["conditions"]["different"]["scope"] == [scope]
    assert result["fields"][2]["validation"] == "REJECTED"
    # A nonlocal scope cannot convert a false predicate into local NA/value.
    if len(affected) > 1:
        assert result["fields"][0]["validation"] == "REJECTED"


@pytest.mark.parametrize("scope", [
    {"kind": "candidate", "id": "paragraph:5"},
    {"kind": "paragraph", "id": "t0.r2.c1"},
    {"kind": "section", "id": "paragraph:5"},
    {"kind": "table", "id": "table:99"},
    {"kind": "unknown", "id": "t0.r2.c1"},
])
def test_bad_condition_scope_is_local_but_cannot_disappear_from_its_origin(case, scope):
    case[2]["conditions"]["different"]["scope"] = [scope]
    result = audit(case)
    assert result["conditions"]["different"]["validation"] == "REJECTED"
    assert result["fields"][0]["validation"] == "GROUNDED_PROPOSAL"
    assert result["fields"][2]["validation"] == "REJECTED"


@pytest.mark.parametrize("change", [{"operator": "invented"}, {"fact_ids": None},
    {"scope": None}, {"form": {"paragraph_id": "paragraph:5", "quote": "Inventata"}},
    {"form": {"paragraph_id": "paragraph:99", "quote": "Inventata"}}])
def test_malformed_condition_does_not_discard_siblings_or_unguard_known_targets(case, change):
    case[2]["conditions"]["different"].update(change)
    result = audit(case)
    assert result["conditions"]["different"]["validation"] == "REJECTED"
    assert result["fields"][0]["validation"] == "GROUNDED_PROPOSAL"
    assert result["fields"][2]["validation"] == "REJECTED"


def test_false_local_condition_allows_only_non_applicability_without_value(case):
    case[2]["fields"]["t0.r2.c1"].update(disposition="not_applicable", fact_id=None)
    result = audit(case)
    field = result["fields"][2]
    assert field["validation"] == "GROUNDED_PROPOSAL"
    assert field["value"] is None
    assert field["source_evidence"][0]["role"] == "source"
    assert field["written_value"] is None


def test_uncertain_local_context_blocks_only_its_assignments(case):
    case[2]["uncertain_paragraph_ids"] = ["paragraph:3"]
    result = audit(case)
    assert result["fields"][0]["validation"] == "GROUNDED_PROPOSAL"
    assert result["fields"][1]["validation"] == "REJECTED"
    assert "incerto" in result["fields"][1]["validation_errors"][0]


def test_local_clause_cannot_be_omitted_by_the_assignment_or_scope(case):
    case[2]["conditions"]["different"]["scope"] = []
    result = audit(case)
    assert result["fields"][2]["validation"] == "REJECTED"
    assert "falsa" in result["fields"][2]["validation_errors"][0]


@pytest.mark.parametrize("kind", ["semantic", "different"])
def test_unknown_condition_never_verifies_a_value(case, kind):
    condition = case[2]["conditions"]["different"]
    condition["operator"] = kind
    if kind == "semantic":
        condition["fact_ids"] = []
    else:
        case[1][0]["content"] += "; altra sede: Via Sole 5"
        case[2]["facts"]["fact-2"].update(value="Via Sole 5", evidence=[{
            "chunk_id": -7, "quote": "altra sede: Via Sole 5"}])
    result = audit(case)
    assert result["conditions"]["different"]["result"] == "unknown"
    assert result["fields"][2]["validation"] == "REJECTED"


@pytest.mark.parametrize("origin", ["form", "user", "other_project", "general"])
def test_source_provenance_and_project_isolation_remain_mandatory(case, origin):
    source = case[1][0]
    if origin == "general":
        source["category"] = "general"
        result = audit(case)
        assert result["facts"]["fact-0"]["validation"] == "REJECTED"
        assert result["fields"][0]["validation"] == "REJECTED"
    else:
        if origin == "other_project":
            source.update(chunk_id=7, file_id=3, project_id="beta", scope="project:beta",
                          category=None)
        else:
            source["role"] = origin
        with pytest.raises(ValueError):
            audit(case)


def test_shared_source_is_preserved_but_does_not_establish_semantic_acceptance(case):
    result = audit(case)
    for field in result["fields"][:2]:
        proof = field["source_evidence"][0]
        assert (proof["file_id"], proof["chunk_id"], proof["scope"], proof["category"]) == (
            -3, -7, "global", "company")
        assert len(proof["content_sha256"]) == 64
        assert field["semantic_review_required"] and field["written_value"] is None
    assert result["resolved"] == result["writes"] == 0
    assert result["condition_inventory"]["semantic_completeness_verified"] is False


@pytest.mark.parametrize("mutation", ["truncated", "duplicate", "envelope", "budget"])
def test_unusable_global_json_is_still_rejected(case, mutation):
    if mutation == "envelope":
        case[2]["deferred"] = []
    elif mutation == "budget":
        case[2]["fields"] = {str(i): {} for i in range(129)}
    raw = json.dumps(case[2])
    if mutation == "truncated":
        raw = raw[:-1]
    elif mutation == "duplicate":
        raw = raw.replace('"version": 5', '"version": 5, "version": 5')
    with pytest.raises(ValueError):
        validate_global_proposal(raw, case[0], "alpha", case[1])
