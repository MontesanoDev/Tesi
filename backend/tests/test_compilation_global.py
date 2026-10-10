"""Software checks for the isolated proposal, not semantic model benchmarks."""

import json
import sqlite3
from io import BytesIO

import pytest
from docx import Document

from app.compilation_global import (
    MAX_SOURCE_CHARACTERS,
    build_global_input,
    global_request,
    validate_global_proposal,
)
from app.config import AISettings
from app.docx_templates import inspect_docx


@pytest.fixture
def layout():
    doc = Document()
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Identificativo organizzazione"
    table.cell(1, 0).text = "Recapito organizzazione"
    stream = BytesIO()
    doc.save(stream)
    return inspect_docx(stream.getvalue())


@pytest.fixture
def source():
    return {"chunk_id": -7, "file_id": -3, "chunk_index": 0,
            "source_name": "profilo.txt",
            "content": "Identificativo: AZ-98765; Recapito: Via Luna 4",
            "role": "source", "scope": "global", "category": "company", "project_id": None}


def response(layout, source):
    data = build_global_input(layout, "alpha", [source])
    records = {}
    for index, candidate in enumerate(data["form"]["candidates"]):
        records[candidate] = {
            "label": "Identificativo organizzazione" if index == 0 else "Recapito organizzazione",
            "entity": "company", "kind": "data", "disposition": "value",
            "value": "AZ-98765" if index == 0 else "Via Luna 4",
            "form": {"paragraph_id": f"paragraph:{index * 2}",
                     "quote": "Identificativo organizzazione" if index == 0
                     else "Recapito organizzazione"},
            "evidence": [{"chunk_id": source["chunk_id"], "quote": source["content"]}],
            "reason": "Proposta da verificare semanticamente",
        }
    return {"fields": records}


def check(payload, layout, source):
    return validate_global_proposal(json.dumps(payload), layout, "alpha", [source])


def test_complete_input_has_form_sources_separate_with_provenance(layout, source):
    data = build_global_input(layout, "alpha", [source])
    assert data["form"]["role"] == "form"
    assert data["sources"][0]["role"] == "source"
    assert data["sources"][0]["chunk_id"] == -7
    assert data["sources"][0]["category"] == "company"
    assert data["source_coverage"]["partial"] is False
    assert data["user_values"] == []
    body = global_request(AISettings("secret", "model", "https://example.invalid"), data)
    assert "secret" not in json.dumps(body)
    assert len(body["messages"]) == 2
    assert body["thinking"] == {"type": "disabled"}


@pytest.mark.parametrize("change", [
    {"role": "form"}, {"role": "user"}, {"scope": "global", "category": "generated"},
    {"chunk_id": 1, "file_id": 1, "project_id": "beta", "scope": "project:beta"},
    {"chunk_id": True}, {"chunk_index": None},
])
def test_input_rejects_wrong_role_scope_identity(layout, source, change):
    source.update(change)
    with pytest.raises(ValueError):
        build_global_input(layout, "alpha", [source])


def test_current_project_source_is_allowed_and_reference_preserved(layout, source):
    source.update(chunk_id=7, file_id=3, project_id="alpha", scope="project:alpha", category=None)
    result = check(response(layout, source), layout, source)
    assert result["counts"] == {"GROUNDED_PROPOSAL": 2}
    proof = result["fields"][0]["source_evidence"][0]
    assert (proof["project_id"], proof["chunk_id"], proof["file_id"]) == ("alpha", 7, 3)
    assert len(proof["content_sha256"]) == 64


def test_one_source_can_ground_two_different_proposals_but_never_promote_them(layout, source):
    result = check(response(layout, source), layout, source)
    assert result["counts"] == {"GROUNDED_PROPOSAL": 2}
    assert result["resolved"] == result["writes"] == 0
    assert all(f["semantic_review_required"] and f["written_value"] is None
               for f in result["fields"])
    assert [f["value"] for f in result["fields"]] == ["AZ-98765", "Via Luna 4"]


@pytest.mark.parametrize("mutation", ["missing", "unknown", "duplicate"])
def test_coverage_ids_are_checked_by_backend_not_model_counter(layout, source, mutation):
    payload = response(layout, source)
    if mutation == "missing":
        payload["fields"].pop("t0.r1.c1")
    elif mutation == "unknown":
        payload["fields"]["invented"] = payload["fields"].pop("t0.r1.c1")
    raw = json.dumps(payload)
    if mutation == "duplicate":
        raw = '{"fields":' + json.dumps(payload["fields"]) + ',"fields":{}}'
    with pytest.raises(ValueError):
        validate_global_proposal(raw, layout, "alpha", [source])


@pytest.mark.parametrize("change,expected", [
    ({"value": "AZ-9876"}, "numero"),
    ({"value": "Inventato"}, "letteralmente"),
    ({"evidence": []}, "senza SOURCE"),
    ({"evidence": [{"chunk_id": -99, "quote": "Identificativo: AZ-98765"}]}, "fuori"),
    ({"evidence": [{"chunk_id": -7, "quote": "Identificativo: INESISTENTE"}]}, "non presente"),
    ({"kind": "signature"}, "protetta"),
    ({"disposition": "missing"}, "inatteso"),
    ({"value": True}, "testuale"),
    ({"field_id": "inventato"}, "Extra inputs"),
])
def test_invalid_local_proposal_is_quarantined_without_losing_valid_sibling(
        layout, source, change, expected):
    payload = response(layout, source)
    payload["fields"]["t0.r0.c1"].update(change)
    result = check(payload, layout, source)
    assert result["counts"] == {"REJECTED": 1, "GROUNDED_PROPOSAL": 1}
    assert expected in " ".join(result["fields"][0]["validation_errors"])


def test_form_quote_from_unrelated_row_is_not_a_local_requirement(layout, source):
    payload = response(layout, source)
    payload["fields"]["t0.r0.c1"]["form"] = {
        "paragraph_id": "paragraph:2", "quote": "Recapito organizzazione",
    }
    result = check(payload, layout, source)
    assert result["fields"][0]["validation"] == "REJECTED"
    assert "contesto strutturale" in result["fields"][0]["validation_errors"][0]


def test_false_semantic_association_is_never_resolved_even_if_quote_is_literal(layout, source):
    payload = response(layout, source)
    payload["fields"]["t0.r0.c1"]["value"] = "Via Luna 4"
    result = check(payload, layout, source)
    # Lexical grounding is intentionally insufficient: no semantic acceptance
    # is forged for a wrong-property proposal, no session/export can consume it.
    assert result["fields"][0]["semantic_review_required"]
    assert result["fields"][0]["written_value"] is None
    assert result["resolved"] == result["writes"] == 0


def test_missing_and_non_applicable_are_distinct_and_need_no_factual_value(layout, source):
    payload = response(layout, source)
    first, second = payload["fields"].values()
    first.update(disposition="missing", value=None, evidence=[])
    second.update(disposition="not_applicable", value=None)
    result = check(payload, layout, source)
    assert result["dispositions"] == {"missing": 1, "not_applicable": 1}
    assert result["counts"] == {"GROUNDED_PROPOSAL": 2}
    second["evidence"] = []
    assert check(payload, layout, source)["fields"][1]["validation"] == "REJECTED"


def test_corpus_budget_fails_without_silent_partial_coverage(layout, source):
    source["content"] = "a" * (MAX_SOURCE_CHARACTERS + 1)
    with pytest.raises(ValueError, match="oltre budget"):
        build_global_input(layout, "alpha", [source])


def test_general_kb_cannot_ground_company_identity(layout, source):
    source["category"] = "general"
    result = check(response(layout, source), layout, source)
    assert result["counts"] == {"REJECTED": 2}
    assert "General KB" in result["fields"][0]["validation_errors"][0]


@pytest.mark.parametrize("value", [True, False, "false"])
def test_checkbox_is_typed_and_never_materialized_in_experiment(source, value):
    from test_compilation_catalog import encoded, native_checkbox

    doc = Document()
    p = doc.add_paragraph("Opzione documentata ")
    native_checkbox(p)
    layout = inspect_docx(encoded(doc), require_candidates=False)
    payload = {"fields": {"p0.cb0": {
        "label": "Opzione documentata", "entity": "company", "kind": "choice",
        "disposition": "value", "value": value,
        "form": {"paragraph_id": "paragraph:0", "quote": "Opzione documentata"},
        "evidence": [{"chunk_id": -7, "quote": source["content"]}],
        "reason": "Decisione semantica da valutare",
    }}}
    result = check(payload, layout, source)
    assert result["counts"] == {"REJECTED" if isinstance(value, str)
                                else "GROUNDED_PROPOSAL": 1}
    assert result["writes"] == result["resolved"] == 0


def test_real_sql_corpus_reread_excludes_forms_artifacts_and_other_projects():
    from scripts.benchmark_compilation_global import load_document_corpus

    with sqlite3.connect(":memory:") as db:
        db.row_factory = sqlite3.Row
        db.executescript("""
            CREATE TABLE projects(id TEXT);
            CREATE TABLE project_files(id INTEGER,project_id TEXT,name TEXT,
                                       kind TEXT,metadata TEXT);
            CREATE TABLE document_chunks(id INTEGER,file_id INTEGER,project_id TEXT,
                                         chunk_index INTEGER,content TEXT);
            CREATE TABLE global_documents(id INTEGER,name TEXT,category TEXT,metadata TEXT);
            CREATE TABLE global_document_chunks(id INTEGER,document_id INTEGER,
                                                chunk_index INTEGER,content TEXT);
            CREATE TABLE project_artifact_links(file_id INTEGER,project_id TEXT,
                                               artifact_id INTEGER);
            CREATE TABLE knowledge_artifacts(id INTEGER,project_id TEXT,kind TEXT);
            INSERT INTO projects VALUES ('alpha'),('beta');
            INSERT INTO project_files VALUES
                (1,'alpha','fonte.txt','source',''),(2,'alpha','modulo.txt','form',''),
                (3,'beta','segreto.txt','source',''),(4,'alpha','fatti.txt','artifact','');
            INSERT INTO document_chunks VALUES
                (1,1,'alpha',0,'Fonte alpha'),(2,2,'alpha',0,'Solo FORM'),
                (3,3,'beta',0,'Fonte beta'),(4,4,'alpha',0,'USER da non confondere con SOURCE');
            INSERT INTO global_documents VALUES
                (1,'azienda.txt','company',''),(2,'manuale.txt','general',''),
                (3,'modello.txt','template','');
            INSERT INTO global_document_chunks VALUES
                (1,1,0,'Fonte aziendale'),(2,2,0,'Fonte generale'),(3,3,0,'Solo template');
        """)
        sources = load_document_corpus(db, "alpha")
        assert {s["chunk_id"] for s in sources} == {1, -1, -2}
        assert all(s["role"] == "source" for s in sources)
        assert {s["category"] for s in sources} == {None, "company", "general"}
        assert load_document_corpus(db, "missing") == []
