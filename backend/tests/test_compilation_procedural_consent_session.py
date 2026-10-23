"""Real session/review/export persistence; provider transport alone is simulated."""

import json
from io import BytesIO

import pytest
from docx import Document
from test_compilation_global_sessions import BASE, api, provider, resolve
from test_compilation_sessions import source

from app import compilation_global_dialogue as dialogue
from app import compilation_global_sessions as adapter
from app.compilation_chat import ClarificationReply

__all__ = ["api", "provider"]


@pytest.fixture
def anyio_backend():
    return "asyncio"


@pytest.mark.anyio
@pytest.mark.parametrize("action", ["CONDITION_TRUE", "CONDITION_FALSE", "UNKNOWN", "SKIP"])
async def test_source_and_user_assent_remain_separate_through_real_session_export(
    api, provider, monkeypatch, action,
):
    doc = Document()
    doc.add_paragraph("PEC aziendale ________")
    doc.add_paragraph("di accettare ogni comunicazione alla PEC ________ assumendone il rischio")
    stream = BytesIO()
    doc.save(stream)
    uploaded = await api.post("/api/projects/alpha/forms", files={
        "file": ("modulo.docx", stream.getvalue())})
    await source(api, "Società Aurora; PEC aziendale: ufficio@aurora.demo.", name="recapiti.md")
    calls = []

    async def proposal(client, settings, body, timeout, **kwargs):
        calls.append(body)
        inp = json.loads(body["messages"][1]["content"])
        evidence = next(s for s in inp["sources"] if "ufficio@aurora.demo" in s["content"])
        return json.dumps({"version": 12, "facts": {"pec": {
            "subject": "Aurora", "entity": "company", "aspect": "identity", "property": "pec",
            "value": "ufficio@aurora.demo", "evidence": [{"chunk_id": evidence["chunk_id"],
                                                        "quote": evidence["content"]}]}},
            "fields": {cid: {"label": "PEC", "kind": "data", "disposition": "value",
                "conditions": [], "property": "pec", "fact_id": "pec",
                "form": {"paragraph_id": f"paragraph:{i}", "quote":
                         f"PEC aziendale [[{cid}]]" if i == 0 else f"PEC [[{cid}]]"}}
                for i, cid in enumerate(["p0.s0", "p1.s0"])},
            "conditions": {}, "bindings": {}, "choice_groups": {}, "coverage": [],
            "clarifications": {}}), "simulated", {}

    monkeypatch.setattr(adapter, "request_model_content", proposal)
    session = (await api.post(BASE, json={"form_id": uploaded.json()["id"], "engine": "global",
                                         "start_in_chat": True})).json()
    session = await resolve(api, session)
    question = session["chat"]["question"]["slots"][0]
    assert question["key"].startswith("authorization:")
    assert "assumendone il rischio" in question["question"]
    assert question["field_ids"] == ["p1.s0"]
    from app.compilation_procedural_consent import receipt_scope
    from app.db import connection

    with connection() as db:
        _, persisted = adapter._row_state(db, "alpha", session["id"], session["version"])
    from app.compilation_procedural_consent import requirements

    current = requirements(persisted["global_compilation"])[0]
    assert current == {k: question[k] for k in current}
    receipt_scope(persisted["global_compilation"], question)
    bypass = {"version": session["version"], "candidate_ids": ["p1.s0"],
              "allow_unresolved": True, "user_quote": "Genera la bozza"}
    denied = await api.post(f"{BASE}/{session['id']}/global-draft", json=bypass)
    assert denied.status_code == 409  # Also enforced below the chat's export selection.
    quote = {"CONDITION_TRUE": "Sì, autorizzo espressamente questa dichiarazione",
             "CONDITION_FALSE": "No, non autorizzo", "UNKNOWN": "Non lo so",
             "SKIP": "Rinvia questa dichiarazione"}[action]
    session, errors = dialogue.apply_replies("alpha", session["id"], session["version"], [
        ClarificationReply(slot=question["slot"], action=action, user_quote=quote)], quote)
    assert not errors
    session = await resolve(api, session)
    ids, conditions, users = dialogue.export_selection(session)
    assert ids == (["p0.s0", "p1.s0"] if action == "CONDITION_TRUE" else ["p0.s0"])
    assert not conditions and not users
    output = await api.post(f"{BASE}/{session['id']}/global-draft", json={
        "version": session["version"], "candidate_ids": ids, "allow_unresolved": True,
        "user_quote": "Genera la bozza parziale"})
    assert output.status_code == 200, output.text
    downloads = output.json()["last_generation"]["downloads"]
    report = (await api.get(downloads["report"])).json()
    assert report["written_field_count"] == len(ids)
    assert all(f["provenance"] == "SOURCE" for f in report["fields"])
    if action == "CONDITION_TRUE":
        receipt = report["conversation_review"]["answers"][question["key"]]
        assert receipt["authorization"]["values"] == {"p1.s0": "ufficio@aurora.demo"}
    written = Document(BytesIO((await api.get(downloads["docx"])).content))
    declaration = next(p.text for p in written.paragraphs if "di accettare" in p.text)
    assert ("ufficio@aurora.demo" in declaration) == (action == "CONDITION_TRUE")
    assert len(calls) == 1
