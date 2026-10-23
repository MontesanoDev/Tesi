"""Real API persistence/export; temporary stores and simulated provider only."""

import json
from copy import deepcopy
from io import BytesIO

import pytest
from docx import Document
from test_compilation_global_sessions import BASE, api, provider, resolve
from test_compilation_sessions import source
from test_compilation_table_form import layout_for, proposed

from app import compilation_global_sessions as adapter

__all__ = ["api", "provider"]


@pytest.fixture
def anyio_backend():
    return "asyncio"


@pytest.mark.anyio
@pytest.mark.parametrize("kind", ["previous", "header"])
@pytest.mark.parametrize("result", ["true", "false", "unknown"])
async def test_structural_form_proof_persists_and_fresh_export_keeps_applicability(
    api, provider, monkeypatch, kind, result,
):
    label = "Sede operativa (se diversa dalla sede legale)"
    rows = [[label], [""]] if kind == "previous" else [[label, "Recapito"], ["", ""]]
    layout = layout_for(rows)
    cid = "t0.r1.c0"
    content = "Sede operativa di Aurora: Via Uno 1."
    if result == "true":
        content += "\nSede operativa di Aurora: diversa dalla sede legale."
    elif result == "false":
        content += "\nSede legale di Aurora: Via Uno 1."
    raw, _ = proposed(layout, cid, content=content, value="Via Uno 1")
    upload = await api.post("/api/projects/alpha/forms", files={
        "file": ("modulo.docx", layout.original)})
    await source(api, content, name="sedi.txt")

    async def proposal(client, settings, body, timeout, **kwargs):
        inp = json.loads(body["messages"][1]["content"])
        proof = next(s for s in inp["sources"] if "Sede operativa di Aurora:" in s["content"])
        value = deepcopy(raw)
        value["facts"]["f"]["evidence"][0]["chunk_id"] = proof["chunk_id"]
        assert "_table_form_links" not in inp["form"]
        return json.dumps(value), "simulated", {}

    monkeypatch.setattr(adapter, "request_model_content", proposal)
    session = (await api.post(BASE, json={"form_id": upload.json()["id"], "engine": "global",
                                         "start_in_chat": True})).json()
    session = await resolve(api, session)
    item = session["global_compilation"]["fields"][0]
    assert item["validation"] == "GROUNDED_PROPOSAL"
    assert item["form_applicability"]["result"] == result
    assert item["selectable"] is (result == "true")
    response = await api.post(f"{BASE}/{session['id']}/global-draft", json={
        "version": session["version"], "candidate_ids": [cid], "allow_unresolved": True,
        "user_quote": "Genera la bozza parziale"})
    if result == "true":
        assert response.status_code == 200, response.text
        downloads = response.json()["last_generation"]["downloads"]
        report = (await api.get(downloads["report"])).json()
        assert report["written_field_count"] == 1
        written = Document(BytesIO((await api.get(downloads["docx"])).content))
        assert written.tables[0].cell(1, 0).text == "Via Uno 1"
    else:
        assert response.status_code == 409, response.text
        assert "Applicabilità FORM" in response.json()["detail"]
        current = (await api.get(f"{BASE}/{session['id']}")).json()
        assert current["last_generation"] is None
