"""Literal role traceability regressions, not independent semantic acceptance."""

import copy
import json
import unicodedata
from io import BytesIO

import pytest
from docx import Document

from app.compilation_global import build_global_input, validate_global_proposal
from app.docx_templates import inspect_docx


@pytest.fixture
def case():
    doc = Document()
    doc.add_heading("Dati generali", 1)
    doc.add_paragraph("Responsabile della verifica")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Nome e cognome"
    table.cell(1, 0).text = "Numero registrazione"
    stream = BytesIO()
    doc.save(stream)
    layout = inspect_docx(stream.getvalue())
    source = {"chunk_id": -7, "file_id": -3, "chunk_index": 0,
              "source_name": "incarico.txt", "role": "source", "scope": "global",
              "category": "company", "project_id": None,
              "content": "Responsabile della verifica: Ada Neri. Numero registrazione: 6234."}
    facts, fields = {}, {}
    for index, (prop, value) in enumerate([
        ("nome del responsabile della verifica", "Ada Neri"),
        ("numero registrazione del responsabile della verifica", "6234"),
    ]):
        ref = f"fact-{index}"
        facts[ref] = {"subject": "Ada Neri", "entity": "person", "aspect": "professional",
                      "property": prop, "value": value,
                      "evidence": [{"chunk_id": -7, "quote": source["content"]}]}
        fields[f"t0.r{index}.c1"] = {
            "label": table.cell(index, 0).text, "kind": "data", "disposition": "value",
            "property": prop, "fact_id": ref, "binding_id": "role", "conditions": [],
            "form": {"paragraph_id": f"paragraph:{2 + 2 * index}",
                     "quote": table.cell(index, 0).text}}
    binding = {"form": {"paragraph_id": "paragraph:1", "quote": "Responsabile della verifica"},
               "required_role": "Responsabile della verifica",
               "source_role": "Responsabile della verifica",
               "evidence": [{"chunk_id": -7, "quote": source["content"]}]}
    return layout, [source], {"version": 7, "conditions": {}, "facts": facts,
                              "fields": fields, "bindings": {"role": binding}}


def audit(case):
    return validate_global_proposal(json.dumps(case[2]), case[0], "alpha", case[1])


def test_explicit_role_is_shared_by_distinct_person_properties_without_becoming_resolved(case):
    result = audit(case)
    assert result["binding_checks"] == {"literal_match": 2}
    assert result["counts"] == {"GROUNDED_PROPOSAL": 2}
    for field in result["fields"]:
        assert field["subject_binding"]["result"] == "literal_match"
        assert field["semantic_review_required"] and field["written_value"] is None
        proof = field["subject_binding"]["source_evidence"][0]
        assert (proof["role"], proof["scope"], proof["category"]) == ("source", "global", "company")
        assert (proof["file_id"], proof["chunk_id"], proof["chunk_index"]) == (-3, -7, 0)
        assert proof["content_sha256"]
    assert [field["value"] for field in result["fields"]] == ["Ada Neri", "6234"]
    assert result["resolved"] == result["writes"] == 0


def test_different_documented_role_cannot_be_renamed_as_the_required_role(case):
    case[1][0]["content"] = "Revisore esterno: Ada Neri. Numero registrazione: 6234."
    for fact in case[2]["facts"].values():
        fact["evidence"][0]["quote"] = case[1][0]["content"]
    binding = case[2]["bindings"]["role"]
    binding.update(source_role="Revisore esterno", evidence=case[2]["facts"]["fact-0"]["evidence"])
    result = audit(case)
    for field in result["fields"]:
        assert field["validation"] == "REVIEW_REQUIRED"
        assert field["value"] is None and field["fact_id"] is None
        assert "differenti" in field["validation_errors"][0]
    assert result["facts"]["fact-0"]["validation"] == "GROUNDED_PROPOSAL"


@pytest.mark.parametrize("reference", [None, "invented"])
def test_person_value_without_real_binding_stays_under_review(case, reference):
    first = case[2]["fields"]["t0.r0.c1"]
    first["binding_id"] = reference
    result = audit(case)
    assert result["fields"][0]["validation"] == "REVIEW_REQUIRED"
    assert result["fields"][0]["value"] is None
    assert result["fields"][1]["validation"] == "GROUNDED_PROPOSAL"


@pytest.mark.parametrize("change", [
    {"required_role": "Ruolo inventato"}, {"source_role": "Ruolo inventato"},
    {"form": {"paragraph_id": "paragraph:999", "quote": "Responsabile della verifica"}},
    {"evidence": [{"chunk_id": -99, "quote": "Responsabile della verifica: Ada Neri"}]},
    {"evidence": [{"chunk_id": -7, "quote": "Estratto inventato"}]},
    {"evidence": []}, {"unexpected": True},
])
def test_invalid_binding_is_quarantined_and_cannot_supply_a_person_role(case, change):
    case[2]["bindings"]["role"].update(change)
    result = audit(case)
    assert result["bindings"]["role"]["validation"] == "REJECTED"
    assert result["counts"] == {"REVIEW_REQUIRED": 2}
    assert all(field["value"] is None for field in result["fields"])


@pytest.mark.parametrize("content", [
    "Ada Neri non è Responsabile della verifica: Ada Neri. Numero registrazione: 6234.",
    "Non risulta Responsabile della verifica: Ada Neri. Numero registrazione: 6234.",
    "Responsabile della verifica non attestato. Ada Neri. Numero registrazione: 6234.",
    "La responsabile della verifica è Ada Neri. Numero registrazione: 6234.",
    "Responsabile della verifica: Bea Mori. Ada Neri: numero registrazione 6234.",
])
def test_negative_narrative_or_other_person_is_not_an_explicit_attribute_pair(case, content):
    case[1][0]["content"] = content
    for fact in case[2]["facts"].values():
        fact["evidence"][0]["quote"] = content
    # Even trimming the preceding negation from the citation cannot pass.
    quote = ("Responsabile della verifica: Ada Neri" if
             "Responsabile della verifica: Ada Neri" in content else content)
    case[2]["bindings"]["role"]["evidence"][0]["quote"] = quote
    result = audit(case)
    assert result["binding_checks"] == {"review_required": 2}
    assert all(field["value"] is None for field in result["fields"])


def test_role_and_value_in_distinct_chunks_do_not_invent_a_cross_chunk_subject_join(case):
    role_source = copy.deepcopy(case[1][0])
    role_source.update(chunk_id=-8, chunk_index=1, content="Responsabile della verifica: Ada Neri.")
    case[1].append(role_source)
    case[1][0]["content"] = "Ada Neri. Numero registrazione: 6234."
    for fact in case[2]["facts"].values():
        fact["evidence"][0]["quote"] = case[1][0]["content"]
    case[2]["bindings"]["role"]["evidence"] = [{"chunk_id": -8, "quote": role_source["content"]}]
    result = audit(case)
    assert result["counts"] == {"REVIEW_REQUIRED": 2}
    assert "stesso chunk" in result["fields"][1]["validation_errors"][0]


def test_unique_same_file_rebinding_keeps_actual_role_proof_identity(case):
    other = copy.deepcopy(case[1][0])
    other.update(chunk_id=-8, chunk_index=1, content="Testo estraneo.")
    case[1].append(other)
    case[2]["bindings"]["role"]["evidence"][0]["chunk_id"] = -8
    result = audit(case)
    proof = result["fields"][0]["subject_binding"]["source_evidence"][0]
    assert proof["chunk_id"] == -7 and proof["cited_chunk_id"] == -8
    assert proof["citation_rebound"]
    assert result["fields"][0]["subject_binding"]["result"] == "literal_match"


def test_role_from_another_branch_is_not_accepted_for_local_candidate(case):
    case[2]["bindings"]["role"]["form"] = {
        "paragraph_id": "paragraph:4", "quote": "Numero registrazione"}
    case[2]["bindings"]["role"].update(required_role="Numero registrazione",
                                       source_role="Numero registrazione")
    result = audit(case)
    assert result["fields"][0]["subject_binding"]["result"] == "review_required"
    assert "esterno al contesto" in result["fields"][0]["validation_errors"][0]


def test_only_nearest_preceding_body_paragraph_is_allowed(case):
    doc = Document()
    doc.add_paragraph("Responsabile della verifica")
    doc.add_paragraph("Sezione diversa")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Nome e cognome"
    table.cell(1, 0).text = "Numero registrazione"
    stream = BytesIO()
    doc.save(stream)
    case = (inspect_docx(stream.getvalue()), case[1], case[2])
    case[2]["bindings"]["role"]["form"]["paragraph_id"] = "paragraph:0"
    result = audit(case)
    assert result["binding_checks"] == {"review_required": 2}
    assert "esterno al contesto" in result["fields"][0]["validation_errors"][0]


def test_current_outline_is_valid_role_context_but_not_a_field_label_anchor(case):
    # The heading is a valid role origin; it does not become an allowed field
    # label elsewhere in the table, which retains its separate row anchor.
    case[2]["bindings"]["role"].update(
        form={"paragraph_id": "paragraph:0", "quote": "Dati generali"},
        required_role="Dati generali", source_role="Dati generali")
    content = "Dati generali: Ada Neri. Numero registrazione: 6234."
    case[1][0]["content"] = content
    for fact in case[2]["facts"].values():
        fact["evidence"][0]["quote"] = content
    case[2]["bindings"]["role"]["evidence"][0]["quote"] = content
    result = audit(case)
    assert result["fields"][0]["subject_binding"]["result"] == "literal_match"
    case[2]["fields"]["t0.r0.c1"]["form"] = case[2]["bindings"]["role"]["form"]
    result = audit(case)
    assert result["fields"][0]["validation"] == "REJECTED"


def test_same_person_may_have_two_explicit_roles_without_role_inference(case):
    case[1][0]["content"] += "\nResponsabile del controllo: Ada Neri."
    for fact in case[2]["facts"].values():
        fact["evidence"][0]["quote"] = case[1][0]["content"]
    # Additional role does not contaminate the original, still literally attested.
    result = audit(case)
    assert result["binding_checks"] == {"literal_match": 2}


def test_role_normalization_only_changes_case_whitespace_and_canonical_unicode(case):
    content = "RESPONSABILE   DELLA VERIFICA:\nAda Neri. Numero registrazione: 6234."
    case[1][0]["content"] = content
    for fact in case[2]["facts"].values():
        fact["evidence"][0]["quote"] = content
    binding = case[2]["bindings"]["role"]
    binding["source_role"] = unicodedata.normalize("NFD", "RESPONSABILE DELLA VERIFICA")
    binding["evidence"][0]["quote"] = content
    # A newline inside the attribute must remain normal whitespace, not remove
    # the identity association. Explicit paragraph attribute syntax is enough.
    assert audit(case)["binding_checks"] == {"literal_match": 2}


def test_valid_binding_never_promotes_unknown_applicability(case):
    case[2]["conditions"]["unknown"] = {
        "form": {"paragraph_id": "paragraph:2", "quote": "Nome e cognome"},
        "aspect": "professional", "operator": "semantic", "fact_ids": [],
        "reason": "Applicabilità non stabilita", "scope": [{"kind": "table", "id": "table:0"}]}
    result = audit(case)
    assert result["counts"] == {"REJECTED": 2}
    assert result["binding_checks"] == {"literal_match": 2}
    assert result["resolved"] == result["writes"] == 0


@pytest.mark.parametrize("role", ["form", "user"])
def test_form_or_user_cannot_attest_person_roles(case, role):
    case[1][0]["role"] = role
    with pytest.raises(ValueError, match="Ruolo SOURCE"):
        audit(case)


def test_general_kb_and_foreign_projects_cannot_attest_roles(case):
    case[1][0]["category"] = "general"
    result = audit(case)
    assert result["counts"] == {"REJECTED": 2}
    assert result["bindings"]["role"]["validation"] == "REJECTED"
    case[1][0].update(chunk_id=7, file_id=3, scope="project:beta", project_id="beta")
    with pytest.raises(ValueError, match="fuori progetto"):
        audit(case)


def test_company_fact_does_not_need_a_person_role_binding(case):
    for fact in case[2]["facts"].values():
        fact.update(entity="company", subject="Alba")
    for field in case[2]["fields"].values():
        del field["binding_id"]
    case[2]["bindings"] = {}
    result = audit(case)
    assert result["counts"] == {"GROUNDED_PROPOSAL": 2}
    assert result["binding_checks"] == {}
    # This gate cannot independently certify a model-supplied entity type.
    assert all(field["semantic_review_required"] for field in result["fields"])


def test_invalid_person_binding_does_not_discard_independent_company_field(case):
    case[2]["bindings"]["role"]["source_role"] = "Ruolo inventato"
    case[2]["facts"]["fact-1"].update(entity="company", subject="Alba")
    del case[2]["fields"]["t0.r1.c1"]["binding_id"]
    result = audit(case)
    assert result["fields"][0]["validation"] == "REVIEW_REQUIRED"
    assert result["fields"][1]["validation"] == "GROUNDED_PROPOSAL"


def test_person_alternatives_stay_nonwritable_without_inventing_a_role(case):
    other = copy.deepcopy(case[2]["facts"]["fact-0"])
    other.update(value="Bea Mori")
    case[1][0]["content"] += " Responsabile della verifica: Bea Mori."
    other["evidence"][0]["quote"] = case[1][0]["content"]
    case[2]["facts"]["other"] = other
    field = case[2]["fields"]["t0.r0.c1"]
    field.pop("fact_id")
    field.pop("binding_id")
    field.update(disposition="review", review_kind="ambiguous", reason="Due nominativi",
                 alternative_fact_ids=["fact-0", "other"])
    reviewed = audit(case)["fields"][0]
    assert reviewed["validation"] == "REVIEW_REQUIRED"
    assert {a["value"] for a in reviewed["alternatives"]} == {"Ada Neri", "Bea Mori"}
    assert reviewed["value"] is reviewed["written_value"] is None


def test_schema_and_envelope_are_bounded_and_do_not_change_v6_replay(case):
    schema = build_global_input(case[0], "alpha", case[1])["schema_output"]
    assert schema["properties"]["version"]["const"] == 11
    assert schema["properties"]["bindings"]["maxProperties"] == 128
    assert schema["$defs"]["SubjectBinding"]["properties"]["evidence"]["maxItems"] == 2
    case[2]["bindings"] = {f"role-{i}": {} for i in range(129)}
    with pytest.raises(ValueError, match="bindings"):
        audit(case)
    del case[2]["bindings"]
    case[2]["version"] = 6
    for field in case[2]["fields"].values():
        del field["binding_id"]
    result = audit(case)
    assert result["counts"] == {"GROUNDED_PROPOSAL": 2}
    assert "bindings" not in result


@pytest.mark.parametrize("valid_binding", [True, False])
def test_conditional_person_keeps_the_full_role_gate_and_unknown_condition(case, valid_binding):
    case[2]["version"] = 8
    case[2]["conditions"]["unknown"] = {
        "form": {"paragraph_id": "paragraph:2", "quote": "Nome e cognome"},
        "aspect": "professional", "operator": "semantic", "fact_ids": [],
        "reason": "Applicabilità da confermare", "scope": [{"kind": "table", "id": "table:0"}]}
    for record in case[2]["fields"].values():
        record.update(disposition="conditional", conditions=["unknown"],
                      reason="Ramo da confermare")
    if not valid_binding:
        case[2]["bindings"]["role"]["source_role"] = "Altro ruolo"
    result = audit(case)
    assert result["counts"] == {"REVIEW_REQUIRED": 2}
    for field in result["fields"]:
        assert field["value"] is field["written_value"] is None
        assert ("proposed_value" in field) is valid_binding
        assert field["condition_checks"]["unknown"]["result"] == "unknown"
        assert field["subject_binding"]["result"] == (
            "literal_match" if valid_binding else "review_required")
    assert result["resolved"] == result["writes"] == 0
