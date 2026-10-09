import json
from io import BytesIO
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

import httpx
import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from app import document_compilation as compilation
from app import document_compilation_routes as routes
from app.artifacts import seed_markdown_artifacts
from app.db import connection, get_storage_path, init_database
from app.docx_templates import (
    DRAFT_NOTICE,
    DocumentInputError,
    DocxTooLargeError,
    fill_docx,
    inspect_docx,
    text_of,
)
from app.generation import GenerationError
from app.main import app
from app.seed import seed_database

CASE = Path(__file__).resolve().parents[2] / "demo-documents/bandi/catanzaro-dl-cse"


def template_bytes(labels=("Ragione sociale", "Codice fiscale", "Firma")):
    document = Document()
    document.add_paragraph("Domanda di partecipazione")
    document.sections[0].header.paragraphs[0].text = "Intestazione originale"
    table = document.add_table(rows=len(labels), cols=2)
    for index, label in enumerate(labels):
        table.cell(index, 0).text = label
    target = BytesIO()
    document.save(target)
    return target.getvalue()


def replace_part(data, name, content):
    target = BytesIO()
    with ZipFile(BytesIO(data)) as source, ZipFile(target, "w", ZIP_DEFLATED) as result:
        for entry in source.infolist():
            result.writestr(entry, content if entry.filename == name else source.read(entry))
    return target.getvalue()


def sources(text="Denominazione: Mapi Ingegneria S.r.l."):
    return compilation.CompilationSources(
        [
            {
                "id": "company:1",
                "document_id": 1,
                "source_name": "azienda.txt",
                "chunk_index": 0,
                "content": text,
                "scope": "company",
                "source_kind": "source",
            }
        ],
        1,
        len(text),
    )


def proposal(**overrides):
    return {
        "cell_id": "t0.r0.c1",
        "label": "Ragione sociale",
        "entity": "company",
        "kind": "data",
        "status": "proposed",
        "value": "Mapi Ingegneria S.r.l.",
        "evidence": [{"source_id": "company:1", "quote": "Denominazione: Mapi Ingegneria S.r.l."}],
        "reason": "Fonte aziendale",
        **overrides,
    }


def result(*fields):
    return json.dumps({"fields": list(fields), "warnings": []})


def test_writer_changes_only_authorized_cells_and_main_part():
    original = template_bytes()
    layout = inspect_docx(original)
    output = fill_docx(layout, {"t0.r0.c1": "Societa D'Amico & Figli <S.r.l.>\nBari"})
    doc = Document(BytesIO(output))
    assert doc.tables[0].cell(0, 1).text == "Societa D'Amico & Figli <S.r.l.>\nBari"
    assert doc.tables[0].cell(1, 1).text == ""
    assert doc.tables[0].cell(2, 1).text == ""
    assert doc.tables[0].cell(0, 0).text == "Ragione sociale"
    assert doc.paragraphs[0].text == DRAFT_NOTICE
    assert original == layout.original
    with ZipFile(BytesIO(original)) as before, ZipFile(BytesIO(output)) as after:
        assert before.namelist() == after.namelist()
        for name in before.namelist():
            if name != "word/document.xml":
                assert before.read(name) == after.read(name)
    with pytest.raises(DocumentInputError, match="non autorizzato"):
        fill_docx(layout, {"t0.r0.c0": "Sovrascrittura etichetta"})
    with pytest.raises(DocumentInputError, match="bozza gia generata"):
        inspect_docx(output)


def test_catanzaro_real_cells_and_package_are_preserved():
    data = (CASE / "modello/domanda-partecipazione.docx").read_bytes()
    mapping = json.loads((CASE / "mappa-campi.json").read_text())
    layout = inspect_docx(data)
    assert len(layout.catalog) == 50
    values = {}
    for field in mapping["fields"]:
        target = field["target"]
        cell_id = f"t{target['table']}.r{target['row']}.c{target['column']}"
        assert cell_id in layout.candidate_ids
        if field["state"] == "proposed_demo":
            values[cell_id] = field["suggested_value"]
    output = fill_docx(layout, values)
    with ZipFile(BytesIO(data)) as before, ZipFile(BytesIO(output)) as after:
        for name in before.namelist():
            if name != "word/document.xml":
                assert before.read(name) == after.read(name)
    doc = Document(BytesIO(output))
    assert len(list(doc.element.body.iter(qn("w:tbl")))) == 50
    assert len(list(doc.element.body.iter(qn("w:checkBox")))) == 28
    # Tests use the fixture, but the production discovery/prompt never imports it.
    assert "Mapi Ingegneria S.r.l." in text_of(doc.element)
    assert "Elisa Romano" in text_of(doc.element)


def test_controls_and_vertical_merge_continuations_are_not_writable():
    document = Document(BytesIO(template_bytes()))
    cell = document.tables[0].cell(0, 1)
    field = OxmlElement("w:fldChar")
    field.set(qn("w:fldCharType"), "begin")
    cell.paragraphs[0].add_run()._r.append(field)
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    cell.paragraphs[0].add_run()._r.append(end)
    document.tables[0].cell(1, 1).merge(document.tables[0].cell(2, 1))
    output = BytesIO()
    document.save(output)
    layout = inspect_docx(output.getvalue())
    assert "t0.r0.c1" not in layout.candidate_ids
    assert "t0.r1.c1" in layout.candidate_ids
    assert "t0.r2.c1" not in layout.candidate_ids


@pytest.mark.parametrize("data", [b"", b"%PDF-1.4", b"not a docx"])
def test_invalid_docx_is_rejected(data):
    with pytest.raises(DocumentInputError):
        inspect_docx(data)


def test_active_content_external_resources_and_dtd_are_rejected():
    data = template_bytes()
    target = BytesIO()
    with ZipFile(BytesIO(data)) as original, ZipFile(target, "w") as archive:
        for entry in original.infolist():
            archive.writestr(entry, original.read(entry))
        archive.writestr("word/vbaProject.bin", b"macro")
    with pytest.raises(DocumentInputError, match="macro"):
        inspect_docx(target.getvalue())
    malicious = b'<!DOCTYPE x [<!ENTITY a SYSTEM "file:///etc/passwd">]><x>&a;</x>'
    with pytest.raises(DocumentInputError, match="DTD"):
        inspect_docx(replace_part(data, "word/document.xml", malicious))
    relationships = (
        b'<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
        b'<Relationship Id="r1" Type="attachedTemplate" '
        b'TargetMode="External" Target="http://example.invalid/remote.dotm"/></Relationships>'
    )
    with pytest.raises(DocumentInputError, match="esterne"):
        inspect_docx(replace_part(data, "word/_rels/document.xml.rels", relationships))


def test_archive_limits_and_no_empty_fields(monkeypatch):
    from app import docx_templates

    data = template_bytes()
    monkeypatch.setattr(docx_templates, "MAX_UNPACKED_BYTES", 100)
    with pytest.raises(DocxTooLargeError):
        inspect_docx(data)
    monkeypatch.undo()
    document = Document()
    document.add_paragraph("Modello soltanto narrativo")
    output = BytesIO()
    document.save(output)
    with pytest.raises(DocumentInputError, match="Nessun campo"):
        inspect_docx(output.getvalue())


def test_proposals_have_checked_sources_and_missing_fields_remain_empty():
    layout = inspect_docx(template_bytes())
    missing = proposal(
        cell_id="t0.r1.c1",
        label="Codice fiscale",
        status="missing",
        value=None,
        evidence=[],
        reason="Non presente",
    )
    parsed = compilation.validate_proposals(result(proposal(), missing), layout, sources())
    assert parsed["fields"][0]["written_value"] == "Mapi Ingegneria S.r.l."
    assert parsed["fields"][0]["evidence"][0]["fragment"] == 1
    assert parsed["fields"][0]["evidence"][0]["page"] is None
    assert parsed["fields"][1]["written_value"] is None
    assert parsed["unclassified_cells"] == ["t0.r2.c1"]


@pytest.mark.parametrize(
    "field",
    [
        proposal(cell_id="t0.r0.c0"),
        proposal(cell_id="t999.r0.c1"),
        proposal(status="missing"),
        proposal(value={"text": "Mapi"}),
        proposal(kind="unsupported"),
    ],
)
def test_invalid_proposals_fail_closed(field):
    with pytest.raises(GenerationError):
        compilation.validate_proposals(result(field), inspect_docx(template_bytes()), sources())


@pytest.mark.parametrize("payload", ["[]", "null", "{}", "not json", '{"fields":"oops"}'])
def test_invalid_model_json_is_rejected(payload):
    with pytest.raises(GenerationError):
        compilation.validate_proposals(payload, inspect_docx(template_bytes()), sources())


def test_duplicate_cells_are_rejected():
    with pytest.raises(GenerationError, match="duplicato"):
        compilation.validate_proposals(
            result(proposal(), proposal()), inspect_docx(template_bytes()), sources()
        )


@pytest.mark.parametrize(
    "field",
    [
        proposal(value="Dato inventato"),
        proposal(evidence=[]),
        proposal(kind="declaration"),
        proposal(kind="choice"),
        proposal(cell_id="t0.r2.c1", label="Nome", kind="data"),
    ],
)
def test_unsupported_values_declarations_and_disguised_signatures_are_not_written(field):
    parsed = compilation.validate_proposals(
        result(field), inspect_docx(template_bytes()), sources()
    )
    assert parsed["fields"][0]["written_value"] is None
    assert parsed["fields"][0]["status"] == "needs_review"
    assert parsed["fields"][0]["validation_notes"]


def test_general_knowledge_is_not_company_evidence():
    context = sources()
    context.selected[0]["scope"] = "general"
    parsed = compilation.validate_proposals(
        result(proposal()), inspect_docx(template_bytes()), context
    )
    assert parsed["fields"][0]["written_value"] is None


def test_person_data_label_is_not_treated_as_a_signature():
    document = Document()
    document.add_paragraph("Nome del firmatario: ____")
    data = BytesIO()
    document.save(data)
    parsed = compilation.validate_proposals(
        result(proposal(cell_id="p0.s0", label="Nome del firmatario")),
        inspect_docx(data.getvalue()),
        sources(),
    )
    assert parsed["fields"][0]["written_value"] is not None


@pytest.fixture
def anyio_backend():
    return "asyncio"


@pytest.fixture
async def client(tmp_path, monkeypatch):
    monkeypatch.setenv("MAPI_DB_PATH", str(tmp_path / "test.db"))
    monkeypatch.setenv("MAPI_STORAGE_PATH", str(tmp_path / "uploads"))
    monkeypatch.setenv("MAPI_KNOWLEDGE_PATH", str(tmp_path / "knowledge"))
    init_database()
    seed_database()
    seed_markdown_artifacts()
    async with httpx.AsyncClient(
        transport=httpx.ASGITransport(app=app), base_url="http://test"
    ) as client:
        yield client


def test_persistence_rolls_back_files_if_write_fails(tmp_path, monkeypatch):
    monkeypatch.setenv("MAPI_DB_PATH", str(tmp_path / "test.db"))
    monkeypatch.setenv("MAPI_STORAGE_PATH", str(tmp_path / "uploads"))
    init_database()
    seed_database()
    original_write = Path.write_bytes

    def write(path, data):
        if path.name == "bozza.docx":
            raise OSError("Disk full")
        return original_write(path, data)

    monkeypatch.setattr(Path, "write_bytes", write)
    with pytest.raises(OSError):
        routes._persist(
            "fondo-riqualificazione-2027",
            "modulo.docx",
            b"original",
            b"draft",
            {"created_at": "now"},
        )
    assert not list(get_storage_path().rglob("template.docx"))
    with connection() as db:
        assert db.execute("SELECT COUNT(*) FROM document_compilations").fetchone()[0] == 0


@pytest.mark.anyio
async def test_previously_saved_reports_and_originals_remain_downloadable(client):
    # Historical metadata is returned as saved, without reviving its execution path.
    original = template_bytes()
    report = {
        "schema_version": 3, "created_at": "2026-01-01T00:00:00+00:00",
        "execution": {"batch_size": 32, "requests": 9, "completed_batches": 9},
        "fields": [], "warnings": [],
    }
    saved = routes._persist(
        "fondo-riqualificazione-2027", "modulo.docx", original, original, report,
    )
    detail = await client.get(
        f"/api/projects/fondo-riqualificazione-2027/document-compilations/{saved['id']}",
    )
    assert detail.json()["report"] == report
    assert (await client.get(saved["downloads"]["report"])).json() == report
    for kind in ("docx", "template"):
        assert (await client.get(saved["downloads"][kind])).content == original
