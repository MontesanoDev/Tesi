"""Generic choice groups: actual grounding/consistency, no semantic provider claims."""

import copy
import json
from io import BytesIO

import pytest
from docx import Document
from test_compilation_catalog import encoded, native_checkbox

from app.compilation_choices import materialize_choices
from app.compilation_global import build_global_input, validate_global_proposal
from app.compilation_global_sessions import initial_state
from app.docx_templates import fill_docx, inspect_docx


@pytest.fixture
def choice_case():
    doc = Document()
    doc.add_heading("Domanda", 1)
    doc.add_paragraph("Tipologia dell'operatore: selezionare una sola opzione.")
    for label, legacy in [("Organizzazione tecnica", True), ("Professionista autonomo", False)]:
        p = doc.add_paragraph()
        native_checkbox(p, legacy=legacy, checked=not legacy)
        p.add_run(" " + label)
    doc.add_paragraph("Denominazione: ________")
    layout = inspect_docx(encoded(doc))
    source = {"chunk_id": -7, "file_id": -3, "chunk_index": 0, "source_name": "profilo.txt",
              "role": "source", "scope": "global", "project_id": None, "category": "company",
              "content": "Denominazione: Vega S.r.l.\nTipologia: Organizzazione tecnica"}
    group = {"label": "Tipologia dell'operatore", "subject": "Vega S.r.l.",
             "property": "tipologia", "selection": "single", "min_selected": 1,
             "conditions": [], "form": {"paragraph_id": "paragraph:1",
             "quote": "Tipologia dell'operatore: selezionare una sola opzione."}, "options": {}}
    fields = {}
    for index, label in enumerate(["Organizzazione tecnica", "Professionista autonomo"]):
        cid = f"p{index + 2}.cb0"
        form = {"paragraph_id": f"paragraph:{index + 2}", "quote": f"[[{cid}]]" + (
            " " if index == 0 else "☑ ") + label}
        group["options"][cid] = {"label": label, "form": form}
        fields[cid] = {"label": label, "kind": "choice", "disposition": "value",
                       "form": form, "conditions": [], "fact_id": "tipo", "property": "tipologia",
                       "choice": index == 0, "choice_group_id": "g"}
    payload = {"version": 11, "uncertain_paragraph_ids": [], "conditions": {}, "bindings": {},
               "choice_groups": {"g": group}, "fields": fields, "facts": {"tipo": {
                   "subject": "Vega S.r.l.", "entity": "company", "aspect": "legal_type",
                   "property": "tipologia", "value": "Organizzazione tecnica", "evidence": [
                       {"chunk_id": -7, "quote": "Tipologia: Organizzazione tecnica"}]}}}
    return layout, [source], payload


def audit(case):
    return validate_global_proposal(json.dumps(case[2]), case[0], "alpha", case[1],
                                    expected_version=11)


def data_for(case):
    data = initial_state(case[0])
    data.update(catalog=build_global_input(case[0], "alpha", case[1])["form"], audit=audit(case))
    data["writable_ids"] += [cid for cid, c in data["catalog"]["candidates"].items()
                             if c["type"] == "checkbox"]
    data["dialogue"] = {}
    return data


def test_source_choice_grounding_and_boolean_provenance(choice_case):
    result = audit(choice_case)
    assert result["choice_groups"]["g"]["validation"] == "GROUNDED_GROUP"
    proposals = [f for f in result["fields"] if f.get("kind") == "choice"]
    assert [f["value"] for f in proposals] == [True, False]
    assert all(f["validation"] == "GROUNDED_PROPOSAL" for f in proposals)
    assert all(f["source_evidence"][0]["chunk_id"] == -7 for f in proposals)
    assert result["resolved"] == result["writes"] == 0


def test_field_reuses_group_marker_without_repeating_it(choice_case):
    for cid, field in choice_case[2]["fields"].items():
        field["form"]["quote"] = choice_case[2]["choice_groups"]["g"]["options"][cid]["label"]
    result = audit(choice_case)
    assert sum(f["validation"] == "GROUNDED_PROPOSAL" for f in result["fields"]) == 2


@pytest.mark.parametrize("bad_form", [
    {"paragraph_id": "paragraph:3", "quote": "Professionista autonomo"},
    {"paragraph_id": "paragraph:2", "quote": "Professionista autonomo"},
    {"paragraph_id": "paragraph:2", "quote": "[[p3.cb0]] Organizzazione tecnica"},
    {"paragraph_id": "paragraph:2", "quote": "tecnica"},
])
def test_reused_marker_never_repairs_a_contradictory_field_quote(choice_case, bad_form):
    choice_case[2]["fields"]["p2.cb0"]["form"] = bad_form
    f = next(f for f in audit(choice_case)["fields"] if f["candidate_id"] == "p2.cb0")
    assert f["validation"] != "GROUNDED_PROPOSAL"


def test_same_option_for_another_source_subject_is_not_accepted(choice_case):
    choice_case[1][0]["content"] = "Denominazione: Altra società\nTipologia: Organizzazione tecnica"
    result = audit(choice_case)
    assert not any(f["validation"] == "GROUNDED_PROPOSAL" for f in result["fields"])


def test_case_and_space_normalization_do_not_drop_valid_source_choices(choice_case):
    choice_case[1][0]["content"] = ("Denominazione: Vega S.r.l.\n"
                                   "Tipologia: organizzazione   tecnica.")
    fact = choice_case[2]["facts"]["tipo"]
    fact.update(value="organizzazione tecnica", evidence=[{
        "chunk_id": -7, "quote": "Tipologia: organizzazione tecnica."}])
    result = audit(choice_case)
    assert sum(f["validation"] == "GROUNDED_PROPOSAL" for f in result["fields"]) == 2
    assert all(f["source_evidence"][0]["support_quote"] == "Tipologia: organizzazione   tecnica."
               for f in result["fields"] if f.get("kind") == "choice")


def test_positive_choice_among_several_attributes_on_one_source_line(choice_case):
    choice_case[1][0]["content"] = ("Denominazione: Vega S.r.l. Tipologia: organizzazione tecnica. "
                                   "Codice: 2468. Nota: esempio.")
    fact = choice_case[2]["facts"]["tipo"]
    fact.update(value="organizzazione tecnica", evidence=[{
        "chunk_id": -7, "quote": "Tipologia: organizzazione tecnica."}])
    result = audit(choice_case)
    assert sum(f["validation"] == "GROUNDED_PROPOSAL" for f in result["fields"]) == 2
    assert all(f["source_evidence"][0]["support_quote"] == "Tipologia: organizzazione tecnica."
               for f in result["fields"] if f.get("kind") == "choice")


def test_inline_negative_attribute_still_rejects_cropped_choice(choice_case):
    choice_case[1][0]["content"] = ("Denominazione: Vega S.r.l. "
                                   "Tipologia: non Organizzazione tecnica. Codice: 2468.")
    choice_case[2]["facts"]["tipo"]["evidence"][0]["quote"] = "Organizzazione tecnica"
    assert not any(f["validation"] == "GROUNDED_PROPOSAL" for f in audit(choice_case)["fields"])


@pytest.mark.parametrize("body", ["non Organizzazione tecnica", "non è Organizzazione tecnica",
                                "non possiede Organizzazione tecnica",
                                "Organizzazione tecnica negata"])
def test_cropped_noun_in_negative_source_cannot_authorize_positive_choice(choice_case, body):
    choice_case[1][0]["content"] = "Denominazione: Vega S.r.l.\nTipologia: " + body
    choice_case[2]["facts"]["tipo"]["evidence"][0]["quote"] = "Organizzazione tecnica"
    result = audit(choice_case)
    assert not any(f["validation"] == "GROUNDED_PROPOSAL" for f in result["fields"])
    assert any("positiva" in error for f in result["fields"] for error in f["validation_errors"])


@pytest.mark.parametrize("change", ["wrong_subject", "wrong_property", "wrong_polarity",
                                   "no_group", "general_source", "no_source"])
def test_unsupported_checkbox_never_becomes_a_value(choice_case, change):
    payload = choice_case[2]
    if change == "wrong_subject":
        payload["choice_groups"]["g"]["subject"] = "Altra azienda"
    elif change == "wrong_property":
        payload["choice_groups"]["g"]["property"] = "altro dato"
    elif change == "wrong_polarity":
        payload["fields"]["p2.cb0"]["choice"] = False
    elif change == "no_group":
        payload["choice_groups"] = {}
    elif change == "general_source":
        choice_case[1][0]["category"] = "general"
    else:
        payload["facts"]["tipo"]["evidence"][0]["chunk_id"] = -99
    assert next(f for f in audit(choice_case)["fields"] if f["candidate_id"] == "p2.cb0")[
        "validation"] != "GROUNDED_PROPOSAL"


@pytest.mark.parametrize("change", ["unknown", "bad_quote", "wrong_marker", "unknown_id",
                                   "overlap", "omitted_middle"])
def test_unreliable_group_is_quarantined(choice_case, change):
    group = choice_case[2]["choice_groups"]["g"]
    if change == "unknown":
        group["selection"] = "unknown"
    elif change == "bad_quote":
        group["form"]["quote"] = "istruzione inventata"
    elif change == "wrong_marker":
        group["options"]["p2.cb0"]["form"]["quote"] = "[[p3.cb0]] Organizzazione tecnica"
    elif change == "unknown_id":
        group["options"]["p99.cb0"] = group["options"].pop("p2.cb0")
    elif change == "overlap":
        choice_case[2]["choice_groups"]["secondo"] = copy.deepcopy(group)
    else:
        group["options"].pop("p2.cb0")
        # Known group member missing: the proposed field cannot attach to it.
    result = audit(choice_case)
    field = next(f for f in result["fields"] if f["candidate_id"] == "p2.cb0")
    assert field["validation"] != "GROUNDED_PROPOSAL"


def test_export_clears_initial_checked_peer_and_always_starts_from_original(choice_case):
    data = data_for(choice_case)
    values, effects = materialize_choices(data, ["p2.cb0"], [])
    assert values == {"p2.cb0": True, "p3.cb0": False}
    assert effects == [{"candidate_id": "p3.cb0", "value": False,
                        "effect": "exclusive_alternative_cleared", "derived_from_candidate_id":
                        "p2.cb0", "choice_group_id": "g", "provenance": "SOURCE"}]
    generated = fill_docx(choice_case[0], {}, checkbox_values=values)
    doc = Document(BytesIO(generated))
    from test_docx_checkboxes import states

    assert states(generated) == {"p3.cb0": True, "p4.cb0": False}
    assert doc.paragraphs[-1].text == "Denominazione: ________"
    assert states(choice_case[0].original) == {"p2.cb0": False, "p3.cb0": True}


def test_export_subset_cannot_bypass_exclusive_group(choice_case):
    data = data_for(choice_case)
    next(f for f in data["audit"]["fields"] if f["candidate_id"] == "p3.cb0")["value"] = True
    with pytest.raises(ValueError, match="cardinalità"):
        materialize_choices(data, ["p2.cb0", "p3.cb0"], [])


def test_absent_source_keeps_options_as_form_context_not_facts(choice_case):
    choice_case[2]["fields"] = {}
    choice_case[2]["facts"] = {}
    result = audit(choice_case)
    assert result["choice_groups"]["g"]["validation"] == "GROUNDED_GROUP"
    assert not any(f["validation"] == "GROUNDED_PROPOSAL" for f in result["fields"])
    assert all(not f["source_evidence"] for f in result["fields"])


def test_multiple_choices_require_independent_positive_sources(choice_case):
    doc = Document(BytesIO(choice_case[0].original))
    doc.paragraphs[1].text = "Attività: selezionare una o più opzioni."
    doc.paragraphs[2].runs[-1].text = " Analisi"
    doc.paragraphs[3].runs[-1].text = " Verifica"
    choice_case = (inspect_docx(encoded(doc)), choice_case[1], choice_case[2])
    group = choice_case[2]["choice_groups"]["g"]
    group.update(selection="multiple", property="attività")
    group["form"]["quote"] = "Attività: selezionare una o più opzioni."
    for cid, label, symbol in [("p2.cb0", "Analisi", ""), ("p3.cb0", "Verifica", "☑")]:
        group["options"][cid].update(label=label, form={
            "paragraph_id": f"paragraph:{cid[1]}", "quote": f"[[{cid}]]{symbol} {label}"})
        choice_case[2]["fields"][cid].update(label=label, property="attività",
                                            form=group["options"][cid]["form"])
    choice_case[2]["facts"]["tipo"].update(property="attività", value="Analisi", evidence=[{
        "chunk_id": -7, "quote": "Attività: Analisi"}])
    other = copy.deepcopy(choice_case[2]["facts"]["tipo"])
    other.update(value="Verifica", evidence=[{
        "chunk_id": -7, "quote": "Attività aggiuntiva: Verifica"}])
    choice_case[1][0]["content"] = ("Denominazione: Vega S.r.l.\nAttività: Analisi\n"
                                   "Attività aggiuntiva: Verifica")
    choice_case[2]["facts"]["altro"] = other
    choice_case[2]["fields"]["p3.cb0"].update(choice=True, fact_id="altro")
    result = audit(choice_case)
    assert sum(f.get("value") is True for f in result["fields"]) == 2
    values, _ = materialize_choices(data_for(choice_case), ["p2.cb0", "p3.cb0"], [])
    assert values == {"p2.cb0": True, "p3.cb0": True}


def test_model_cannot_turn_explicit_exclusivity_into_multiple(choice_case):
    choice_case[2]["choice_groups"]["g"]["selection"] = "multiple"
    result = audit(choice_case)
    assert result["choice_groups"]["g"]["validation"] == "REJECTED"
    assert not any(f["validation"] == "GROUNDED_PROPOSAL" for f in result["fields"])


def test_two_exclusive_positive_sources_are_quarantined_together(choice_case):
    other = copy.deepcopy(choice_case[2]["facts"]["tipo"])
    other.update(value="Professionista autonomo", evidence=[{
        "chunk_id": -7, "quote": "Seconda tipologia: Professionista autonomo"}])
    choice_case[1][0]["content"] += "\nSeconda tipologia: Professionista autonomo"
    choice_case[2]["facts"]["altro"] = other
    choice_case[2]["fields"]["p3.cb0"].update(choice=True, fact_id="altro")
    result = audit(choice_case)
    assert not any(f["validation"] == "GROUNDED_PROPOSAL" for f in result["fields"])
    assert result["choice_groups"]["g"]["consistency"] == "conflicting"


def test_user_selection_is_scoped_and_has_no_source_provenance(choice_case):
    from app.compilation_chat import ClarificationReply
    from app.compilation_choices import user_choice
    from app.compilation_global_dialogue import question_pool

    choice_case[2]["fields"], choice_case[2]["facts"] = {}, {}
    data = data_for(choice_case)
    slot = question_pool(data)[0]
    assert slot["kind"] == "choice" and slot["alternatives"] == [
        "Organizzazione tecnica", "Professionista autonomo"]
    reply = ClarificationReply(action="VALUE", slot=1, value="siamo un professionista autonomo",
                               user_quote="siamo un professionista autonomo", selected_options=[2])
    values = user_choice(data, slot, reply)
    assert values["p3.cb0"]["value"] is True and values["p2.cb0"]["value"] is False
    assert all(v["provenance"] == "USER" and not v.get("source_evidence") for v in values.values())
    for invalid in [[1, 2], [0], [99], [1, 1]]:
        with pytest.raises(ValueError):
            user_choice(data, slot, reply.model_copy(update={"selected_options": invalid}))


def test_non_applicable_or_unknown_group_does_not_ask_for_a_choice(choice_case):
    from app.compilation_global_dialogue import question_pool

    choice_case[2]["fields"], choice_case[2]["facts"] = {}, {}
    choice_case[2]["conditions"] = {"c": {
        "form": {"paragraph_id": "paragraph:1", "quote": "Tipologia dell'operatore"},
        "aspect": "participation", "operator": "semantic", "fact_ids": [],
        "reason": "Ramo da confermare", "scope": [{"kind": "candidate", "id": "p2.cb0"},
                                                  {"kind": "candidate", "id": "p3.cb0"}]}}
    choice_case[2]["choice_groups"]["g"]["conditions"] = ["c"]
    data = data_for(choice_case)
    assert [s["kind"] for s in question_pool(data)] == ["applicability"]
    data["dialogue"]["answers"] = {"condition:c": {"applies": False}}
    assert question_pool(data) == []


def test_form_or_foreign_source_cannot_support_choices(choice_case):
    for mutation in [{"role": "form"}, {"scope": "project:beta", "project_id": "beta",
                                       "chunk_id": 7, "file_id": 3}]:
        bad = (choice_case[0], copy.deepcopy(choice_case[1]), copy.deepcopy(choice_case[2]))
        bad[1][0].update(mutation)
        with pytest.raises(ValueError):
            audit(bad)


