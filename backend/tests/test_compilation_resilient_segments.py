"""Rev 2 resilient FORM sections, simulated provider (zero LLM calls).

Ownership by dependency units, structural validation of every segment before the merge,
bounded retries, a request-digest cache for resumption and NON ACCERTATA destinations for
segments that were not analyzed. Valid JSON alone never authorizes a write.
"""

import asyncio
import inspect
import json
import re
from io import BytesIO

import pytest
from docx import Document
from test_compilation_global_sessions import BASE, api, provider
from test_compilation_p3 import fact
from test_compilation_sessions import source

from app import compilation_global_dialogue as dialogue
from app import compilation_global_sessions as adapter
from app import compilation_segments as segmentation
from app import compilation_sessions as sessions
from app.compilation_chat import ClarificationReply
from app.compilation_segments import (
    applicability_states,
    dependency_units,
    resilient_merge,
    resilient_partition,
    unassessed_fields,
    uncertain_fields,
)
from app.db import connection, get_storage_path
from app.docx_templates import inspect_docx
from app.generation import GenerationTruncatedError
from app.repository import save_conversation_turn

__all__ = ["api", "provider"]
CLAUSE = "☐ per le società DI ESSERE iscritto alla CCIAA con i seguenti dati:"
PREDICATE = "Sede operativa (se diversa dalla sede legale): ____________"
HEADING = "SEZIONE A - Da compilare solo per le società"


@pytest.fixture(autouse=True)
def candidate_runtime(monkeypatch):
    """The resilient candidate is opt-in: dialogue P3, derivation and generation rev 2."""
    monkeypatch.setenv("MAPI_DIALOGUE_REVISION", "1")
    monkeypatch.setenv("MAPI_DERIVATION_REVISION", "1")
    monkeypatch.setenv("MAPI_GENERATION_REVISION", "2")


def form_docx(count, *, clause=None, predicate=None, heading=False, headings=()):
    doc = Document()
    if heading:
        doc.add_paragraph().add_run(HEADING).bold = True
    for i in range(count):
        if i in headings:
            doc.add_paragraph().add_run(f"SEZIONE {i}").bold = True
        doc.add_paragraph(f"Dato {i}: ____________")
        if clause == i:
            doc.add_paragraph(CLAUSE)
            table = doc.add_table(rows=1, cols=4)
            table.rows[0].cells[0].text = "Provincia di iscrizione:"
            table.rows[0].cells[2].text = "numero di iscrizione:"
        if predicate == i:
            doc.add_paragraph(PREDICATE)
    stream = BytesIO()
    doc.save(stream)
    return stream.getvalue()


def model_form(data):
    """The FORM catalog as sent to the model (with anchor paragraphs)."""
    from app.compilation_global import build_global_input

    evidence = {"chunk_id": -1, "file_id": -21, "chunk_index": 0, "source_name": "dati.md",
                "role": "source", "scope": "global", "category": "company", "project_id": None,
                "content": "Codice: ABC123"}
    return build_global_input(inspect_docx(data), "alpha", [evidence])["form"]


def segment_of(body):
    """(index, owned candidate ids) from the backend instruction; (0, None) if monolithic."""
    if len(body["messages"]) != 3:
        return 0, None
    text = body["messages"][-1]["content"]
    index = int(re.match(r"SEGMENTO (\d+) DI", text)[1])
    return index, text.rsplit("[", 1)[1].rstrip("]").split(",")


def paragraph_of(inp, cid):
    key = inp["form"]["candidates"][cid]["context_id"]
    return next(p for b in inp["form"]["blocks"] for p in b["paragraphs"] if p["id"] == key)


def slot(inp, cid, fact_id="f", conditions=()):
    paragraph = paragraph_of(inp, cid)
    quote = next(part for part in re.split(r"(?<=\]\])", paragraph["text_with_fields"])
                 if f"[[{cid}]]" in part).strip()
    record = {"label": "Codice", "kind": "data", "disposition": "value",
              "conditions": list(conditions), "form": {"paragraph_id": paragraph["id"],
                                                       "quote": quote},
              "fact_id": fact_id, "property": "codice"}
    if conditions:
        record.update(disposition="conditional", reason="Applicabilità da confermare")
    return record


def envelope(inp, fields, **extra):
    evidence = next(s for s in inp["sources"] if "ABC123" in s["content"])
    return {"version": 12, "uncertain_paragraph_ids": [], "conditions": {}, "bindings": {},
            "choice_groups": {}, "clarifications": {}, "coverage": [],
            "facts": {"f": fact("ABC123", evidence["chunk_id"])}, "fields": fields, **extra}


def dato_ids(inp, ids):
    return [c for c in ids if paragraph_of(inp, c)["text"].startswith("Dato")]


def ordinary(inp, ids):
    """A valid answer: the first and last 'Dato' destinations of the segment."""
    chosen = dato_ids(inp, ids)
    return json.dumps(envelope(inp, {c: slot(inp, c) for c in {chosen[0], chosen[-1]}}))


class Model:
    """Scripted provider: plan[index] is a list of outcomes, one per attempt."""

    def __init__(self, plan=None):
        self.plan, self.calls = plan or {}, []

    async def __call__(self, client, settings, body, timeout, **kwargs):
        index, ids = segment_of(body)
        inp = json.loads(body["messages"][1]["content"])
        attempt = sum(1 for i, _ in self.calls if i == index)
        self.calls.append((index, attempt))
        outcomes = self.plan.get(index, [])
        outcome = outcomes[min(attempt, len(outcomes) - 1)] if outcomes else "ok"
        if isinstance(outcome, BaseException):
            raise outcome
        if callable(outcome):
            outcome = await outcome(inp, ids) if inspect.iscoroutinefunction(outcome) else \
                outcome(inp, ids)
        raw = ordinary(inp, ids) if outcome == "ok" else outcome
        return raw, "simulated", {"prompt_tokens": 100, "completion_tokens": 10,
                                  "total_tokens": 110, "prompt_cache_hit_tokens": 40}


async def started(api, monkeypatch, data, model, *, content="Codice: ABC123", start=None):
    upload = await api.post("/api/projects/alpha/forms", files={"file": ("modulo.docx", data)})
    await source(api, content, name="dati.md")
    monkeypatch.setattr(adapter, "request_model_content", model)
    session = (await api.post(BASE, json={"form_id": upload.json()["id"], "engine": "global",
                                         "start_in_chat": True})).json()
    if start:
        save_conversation_turn(session["conversation_id"], {
            "question": start, "answer": "Avvio della compilazione",
            "generation_status": "direct",
            "compilation": {"session_id": session["id"], "action": "start"}})
    return session


async def analyze(api, session):
    return await api.post(f"{BASE}/{session['id']}/resolve",
                          json={"version": session["version"], "automatic": True})


def state_of(session):
    with connection() as db:
        return json.loads(sessions._row(db, "alpha", session["id"])["state_json"])


def segments_of(session):
    return state_of(session)["global_compilation"]["audit"]["generation_segments"]


async def draft(api, session, ids, users=(), conditions=()):
    return await api.post(f"{BASE}/{session['id']}/global-draft", json={
        "version": session["version"], "candidate_ids": ids, "user_candidate_ids": list(users),
        "confirmed_condition_ids": list(conditions), "allow_unresolved": True,
        "user_quote": "Genera la bozza parziale"})


# T3/T4: per-segment validation before the merge, isolated failures, bounded retries.


def test_partition_owns_dependency_units_and_shows_the_local_context():
    catalog = model_form(form_docx(150, clause=70, heading=True))
    segments = resilient_partition(catalog)
    owned = [c for s in segments for c in s["candidate_ids"]]
    assert sorted(owned) == sorted(catalog["candidates"]) and len(owned) == len(set(owned))
    owner = {c: s["index"] for s in segments for c in s["candidate_ids"]}
    for unit in dependency_units(catalog):
        assert len({owner[c] for c in unit}) == 1  # A dependency unit is never split.
    region = [c for c, x in catalog["candidates"].items() if x["location"].get("table_id")]
    clause = next(b["id"] for b in catalog["blocks"] for p in b["paragraphs"]
                  if p["text"] == CLAUSE)
    holder = next(s for s in segments if region[0] in s["candidate_ids"])
    assert clause in holder["context_block_ids"] and {owner[c] for c in region} == {
        holder["index"]}
    heading = catalog["blocks"][0]["id"]
    assert all(heading in s["context_block_ids"] for s in segments)  # Nearest heading.
    order = [b["id"] for b in catalog["blocks"]]
    assert [b for s in segments for b in s["block_ids"]] == [
        b for b in order if any(b in s["block_ids"] for s in segments)]  # Document order.
    assert len(resilient_partition(model_form(form_docx(20)))) == 1


def test_a_tiny_section_joins_the_previous_pack_and_shared_parts_stay_in_the_prefix():
    from app.compilation_global import build_global_input
    from app.compilation_segments import segment_input

    catalog = model_form(form_docx(103, headings=(100,)))
    assert [len(s["candidate_ids"]) for s in resilient_partition(catalog)] == [60, 43]
    catalog = model_form(form_docx(150, headings=(100,)))
    assert [len(s["candidate_ids"]) for s in resilient_partition(catalog)] == [60, 40, 50]
    evidence = {"chunk_id": -1, "file_id": -21, "chunk_index": 0, "source_name": "dati.md",
                "role": "source", "scope": "global", "category": "company", "project_id": None,
                "content": "Codice: ABC123"}
    data = build_global_input(inspect_docx(form_docx(150)), "alpha", [evidence])
    segments = resilient_partition(data["form"])
    first, second = (json.dumps(segment_input(data, s), ensure_ascii=False) for s in segments[:2])
    shared = next(i for i, (a, b) in enumerate(zip(first, second, strict=False)) if a != b)
    plan = list(segment_input(data, segments[0])["reading_plan"])
    assert plan[-2:] == ["units", "source_role_targets"] or "units" not in plan
    # Everything before this segment's own units (SOURCE, schema, plan lines) is shared.
    assert first.index('"meaning"') < first.index('"units"') < shared


@pytest.mark.anyio
@pytest.mark.parametrize("failing", [1, 2, 3])
async def test_a_failed_segment_leaves_the_independent_segments_valid(
    api, provider, monkeypatch, failing,
):
    model = Model({failing: [GenerationTruncatedError(49152, 49152)]})
    session = await started(api, monkeypatch, form_docx(150), model)
    response = await analyze(api, session)
    assert response.status_code == 200, response.text
    session = response.json()
    report = segments_of(session)
    assert [s["status"] for s in report["segments"]] == [
        "failed" if s["index"] == failing else "analyzed" for s in report["segments"]]
    assert len(model.calls) == len(report["segments"]) + 1 and report["retries"] == 1
    failed = next(s for s in report["segments"] if s["index"] == failing)
    assert failed["attempts"] == 2 and "GenerationTruncatedError" in failed["error"]
    unanalyzed = set(report["unanalyzed_candidate_ids"])
    data = state_of(session)["global_compilation"]
    fields = {f["candidate_id"]: f for f in data["audit"]["fields"]}
    assert unanalyzed and all(fields[c]["coverage"]["state"] == "UNASSESSED"
                              and fields[c]["coverage"]["origin"] == "BACKEND"
                              for c in unanalyzed)
    ids, _, _ = dialogue.export_selection(session)
    assert ids and not set(ids) & unanalyzed
    assert len(ids) == 2 * (len(report["segments"]) - 1)  # Every independent proposal kept.
    exported = await draft(api, session, ids)
    assert exported.status_code == 200, exported.text


@pytest.mark.anyio
@pytest.mark.parametrize("broken", ["incomplete", "structure", "budget", "duplicate_key"])
async def test_broken_output_fails_only_its_segment_and_is_never_repaired(
    api, provider, monkeypatch, broken,
):
    def output(inp, ids):
        valid = ordinary(inp, ids)
        if broken == "incomplete":
            return valid[: len(valid) // 2]
        if broken == "structure":
            return json.dumps({**json.loads(valid), "fields": [], "facts": "nessuno"})
        if broken == "budget":
            return json.dumps(envelope(inp, {}, clarifications={
                c: {"label": "Dato", "reason": "Da chiarire", "form": {
                    "paragraph_id": paragraph_of(inp, c)["id"], "quote": "Dato"}}
                for c in ids[:17]}))
        return valid[:-1] + ',"facts":{}}'

    model = Model({2: [output]})
    session = await started(api, monkeypatch, form_docx(150), model)
    session = (await analyze(api, session)).json()
    report = segments_of(session)
    assert [s["status"] for s in report["segments"]] == ["analyzed", "failed", "analyzed"]
    assert report["segments"][1]["attempts"] == 2  # One bounded retry, then recorded failure.
    folder = get_storage_path().resolve() / state_of(session)["global_compilation"][
        "artifact_path"] / "segments" / "2"
    assert sorted(p.name for p in folder.glob("response-*.txt")) == [
        "response-1.txt", "response-2.txt"]  # Raw outputs kept as received, never repaired.
    ids, _, _ = dialogue.export_selection(session)
    assert len(ids) == 4 and not set(ids) & set(report["unanalyzed_candidate_ids"])


@pytest.mark.anyio
async def test_literal_control_characters_in_strings_are_read_without_other_changes(
    api, provider, monkeypatch,
):
    model = Model({2: [lambda inp, ids: ordinary(inp, ids).replace(
        '"label":"Codice"', '"label":"Codice\tsocietà"').replace(
        '"label": "Codice"', '"label": "Codice\tsocietà"')]})
    session = await started(api, monkeypatch, form_docx(150), model)
    session = (await analyze(api, session)).json()
    report = segments_of(session)
    assert [s["status"] for s in report["segments"]] == ["analyzed"] * 3
    assert report["retries"] == 0 and len(model.calls) == 3


@pytest.mark.anyio
async def test_a_successful_retry_is_recorded_and_the_analysis_retry_budget_is_bounded(
    api, provider, monkeypatch,
):
    transient = GenerationTruncatedError(49152, 49152)
    model = Model({1: [transient, "ok"], 2: [transient], 3: [transient]})
    session = await started(api, monkeypatch, form_docx(150), model)
    session = (await analyze(api, session)).json()
    report = segments_of(session)
    assert report["retries"] == 2  # At most two per analysis, one per segment.
    assert [(s["status"], s["attempts"]) for s in report["segments"]] == [
        ("analyzed", 2), ("failed", 2), ("failed", 1)]
    assert len(model.calls) == 5


@pytest.mark.anyio
async def test_no_analyzed_segment_fails_the_analysis_without_values(api, provider, monkeypatch):
    model = Model({i: [GenerationTruncatedError(49152, 49152)] for i in (1, 2, 3)})
    session = await started(api, monkeypatch, form_docx(150), model)
    response = await analyze(api, session)
    assert response.status_code >= 400
    stored = state_of(session)
    assert stored["status"] == "FAILED" and stored["global_compilation"]["audit"] is None


# T5: valid JSON with a correct SOURCE value that omits a FORM condition.


def condition(inp, paragraph_id, quote, scope=None):
    return {"form": {"paragraph_id": paragraph_id, "quote": quote},
            "aspect": "participation", "operator": "semantic", "fact_ids": [],
            "reason": "Applicabilità da confermare",
            "scope": scope or [{"kind": "paragraph", "id": paragraph_id}]}


@pytest.mark.anyio
@pytest.mark.parametrize("definer", ["same_segment", "other_segment"])
async def test_a_recognizable_omitted_condition_stays_conditional_and_is_not_exported(
    api, provider, monkeypatch, definer,
):
    """The field omits a condition that some segment anchored to its local FORM clause."""
    target = {}

    def owner(inp, ids):
        cid = dato_ids(inp, ids)[0]
        target["cid"] = cid
        extra = {}
        if definer == "same_segment":
            paragraph = paragraph_of(inp, cid)
            extra["conditions"] = {"c": condition(inp, paragraph["id"], paragraph["text"][:7])}
        return json.dumps(envelope(inp, {cid: slot(inp, cid)}, **extra))

    def other(inp, ids):
        payload = json.loads(ordinary(inp, ids))
        payload["conditions"] = {"c": condition(inp, "paragraph:0", "Dato 0", scope=[
            {"kind": "candidate", "id": target["cid"]}])}
        return json.dumps(payload)

    model = Model({2: [owner], 3: [other] if definer == "other_segment" else ["ok"]})
    session = await started(api, monkeypatch, form_docx(150), model)
    session = (await analyze(api, session)).json()
    data = state_of(session)["global_compilation"]
    field = next(f for f in data["audit"]["fields"] if f["candidate_id"] == target["cid"])
    assert field["proposal"]["conditions"] == []  # Omitted by the model's field record...
    # ...attached by the unchanged V12 validator, which then refuses the unconditional value.
    assert field["validation"] == "REJECTED"
    assert any("Condizione sconosciuta" in e for e in field["validation_errors"])
    assert target["cid"] not in applicability_states(data)  # Not a write: nothing to assess.
    ids, _, _ = dialogue.export_selection(session)
    assert target["cid"] not in ids and ids  # The other segments stay usable.


@pytest.mark.anyio
@pytest.mark.parametrize("kind", ["option_region", "form_predicate"])
async def test_backend_owned_applicability_vetoes_a_value_without_any_model_condition(
    api, provider, monkeypatch, kind,
):
    evidence_text = "Codice: ABC123\nProvincia di iscrizione: Torino\nSede operativa: Milano"
    target = {}

    def owner(inp, ids):
        evidence = next(s for s in inp["sources"] if "ABC123" in s["content"])
        labels = {p["id"]: p["text"] for b in inp["form"]["blocks"] for p in b["paragraphs"]}
        if kind == "option_region":
            cid = next(c for c in ids if inp["form"]["candidates"][c]["location"].get(
                "table_id"))
            label = next(p for p in labels if labels[p].startswith("Provincia"))
            form = {"paragraph_id": label, "quote": labels[label].strip()}
            value, quote, name = "Torino", "Provincia di iscrizione: Torino", "provincia"
        else:
            cid = next(c for c in ids if paragraph_of(inp, c)["text"].startswith("Sede"))
            paragraph = paragraph_of(inp, cid)
            form = {"paragraph_id": paragraph["id"], "quote": paragraph["text_with_fields"]}
            value, quote, name = "Milano", "Sede operativa: Milano", "sede operativa"
        target["cid"] = cid
        facts = {"v": {"subject": "Aurora S.r.l.", "entity": "company", "aspect": "identity",
                       "property": name, "value": value,
                       "evidence": [{"chunk_id": evidence["chunk_id"], "quote": quote}]}}
        return json.dumps({**envelope(inp, {cid: {
            "label": name, "kind": "data", "disposition": "value", "conditions": [],
            "fact_id": "v", "property": name, "form": form}}), "facts": facts})

    layout = form_docx(150, clause=100 if kind == "option_region" else None,
                       predicate=100 if kind == "form_predicate" else None)
    owner_index = next(s["index"] for s in resilient_partition(model_form(layout))
                       if any(c.startswith("t") or c == "p101.s0" for c in s["candidate_ids"]))
    model = Model({owner_index: [owner]})
    session = await started(api, monkeypatch, layout, model, content=evidence_text)
    session = (await analyze(api, session)).json()
    data = state_of(session)["global_compilation"]
    field = next(f for f in data["audit"]["fields"] if f["candidate_id"] == target["cid"])
    assert "GROUNDED_PROPOSAL" in {field.get("validation"), field.get("proposal_grounding")}
    state = applicability_states(data)[target["cid"]]
    reason = "autorizzazione_V14" if kind == "option_region" else "predicato_FORM"
    assert state["state"] == "CONDIZIONATA" and reason in state["reasons"]
    ids, _, _ = dialogue.export_selection(session)
    assert target["cid"] not in ids


@pytest.mark.anyio
async def test_a_value_under_a_clause_the_model_declared_uncertain_is_not_written(
    api, provider, monkeypatch,
):
    target = {}

    def uncertain(inp, ids):
        cid = dato_ids(inp, ids)[0]
        target["cid"] = cid
        return json.dumps(envelope(inp, {cid: slot(inp, cid)}, uncertain_paragraph_ids=[
            paragraph_of(inp, cid)["id"]]))

    model = Model({2: [uncertain]})
    session = await started(api, monkeypatch, form_docx(150), model)
    session = (await analyze(api, session)).json()
    data = state_of(session)["global_compilation"]
    field = next(f for f in data["audit"]["fields"] if f["candidate_id"] == target["cid"])
    # The unchanged V12 validator already refuses the model's own proposal there...
    assert field["validation"] == "REJECTED"
    assert "Inventario locale incerto" in field["validation_errors"][0]
    # ...and rev 2 keeps any later SOURCE write (e.g. a reuse) there CONDIZIONATA as well.
    assert target["cid"] in uncertain_fields(data)
    reused = {**data, "audit": {**data["audit"], "fields": [
        {**f, "proposal_grounding": "GROUNDED_PROPOSAL", "source_reuse": {"method": "test"}}
        if f["candidate_id"] == target["cid"] else f for f in data["audit"]["fields"]]}}
    assert applicability_states(reused)[target["cid"]] == {
        "state": "CONDIZIONATA", "reasons": ["paragrafo_locale_incerto"]}
    ids, _, _ = dialogue.export_selection(session)
    assert target["cid"] not in ids and ids
    forced = await draft(api, session, [*ids, target["cid"]])
    assert forced.status_code == 409


@pytest.mark.anyio
async def test_unrecognizable_omitted_condition_is_a_documented_residual_risk(
    api, provider, monkeypatch,
):
    """Counterexample: the heading restricts the section, the model omits it, nothing else
    knows the clause. The value is VERIFICATA only in the sense of "no recognized pending
    dependency"; this test pins the documented limit instead of hiding it."""
    model = Model()
    session = await started(api, monkeypatch, form_docx(150, heading=True), model)
    session = (await analyze(api, session)).json()
    data = state_of(session)["global_compilation"]
    states = applicability_states(data)
    assert states and {s["state"] for s in states.values()} == {"VERIFICATA"}
    ids, _, _ = dialogue.export_selection(session)
    assert ids  # Exported: the restriction in the heading was never recognized.


# T6/T7: shared conditions and conflicts.


@pytest.mark.anyio
async def test_equivalent_conditions_from_two_segments_become_one_answer(
    api, provider, monkeypatch,
):
    def scoped(inp, ids):
        chosen = dato_ids(inp, ids)[:2]
        heading = inp["form"]["blocks"][0]["paragraphs"][0]
        assert heading["text"] == HEADING  # The nearest heading is part of every context.
        return json.dumps(envelope(inp, {c: slot(inp, c, conditions=["societa"])
                                         for c in chosen}, conditions={"societa": condition(
            inp, heading["id"], "solo per le società", scope=[
                {"kind": "candidate", "id": c} for c in chosen])}))

    model = Model({1: [scoped], 2: [scoped], 3: [scoped]})
    session = await started(api, monkeypatch, form_docx(150, heading=True), model)
    session = (await analyze(api, session)).json()
    data = state_of(session)["global_compilation"]
    report = data["audit"]["generation_segments"]
    assert len(report["unified_conditions"]) == 1
    assert list(data["audit"]["conditions"]) == [report["unified_conditions"][0]["kept"]]
    key = report["unified_conditions"][0]["kept"]
    conditional = [f["candidate_id"] for f in data["audit"]["fields"] if key in f.get(
        "conditions", [])]
    assert len(conditional) == 6
    receipt = {"applies": True, "action": "CONDITION_TRUE", "provenance": "USER",
               "user_quote": "Sì, siamo una società"}
    after = {**data, "dialogue": {**(data.get("dialogue") or {}), "answers": {
        **((data.get("dialogue") or {}).get("answers") or {}), "condition:" + key: receipt}}}
    states = applicability_states(after)
    assert all(states[c]["state"] == "VERIFICATA" for c in conditional)  # One answer, 3 scopes.


def test_different_conditions_on_one_clause_and_conflicting_values_stay_visible():
    catalog = model_form(form_docx(150, heading=True))
    segments = resilient_partition(catalog)
    first, second = segments[0]["candidate_ids"][0], segments[1]["candidate_ids"][0]

    def record(cid, value, aspect):
        paragraph = catalog["candidates"][cid]["context_id"]
        quote = f"Dato {int(paragraph.split(':')[1]) - 1}: [[{cid}]]"
        return {"version": 12, "facts": {"f": fact(value)}, "conditions": {"c": {
            "form": {"paragraph_id": "paragraph:0", "quote": "solo per le società"},
            "aspect": aspect, "operator": "semantic", "fact_ids": [], "reason": "?"}},
            "fields": {cid: {"label": "Codice", "kind": "data", "disposition": "value",
                             "conditions": ["c"], "fact_id": "f", "property": "codice",
                             "form": {"paragraph_id": paragraph, "quote": quote}}}}

    merged, report = resilient_merge(segments, {1: record(first, "ABC123", "participation"),
                                                2: record(second, "XYZ789", "legal_form")})
    assert report["unified_conditions"] == []
    assert report["shared_conditions_not_unified"] == ["paragraph:0"]
    assert len(merged["conditions"]) == 2
    assert len(report["conflicts"]) == 1  # Same subject and property, different values.
    assert {merged["fields"][c]["disposition"] for c in (first, second)} == {"review"}


# T8: SOURCE reuse after the owner segment failed.

SIGNER = "Ada Neri"


@pytest.mark.anyio
async def test_reuse_into_an_unanalyzed_section_is_not_exported_but_a_user_value_is(
    api, provider, monkeypatch,
):
    from test_compilation_selected_person import case

    _, layout, sources = case()
    original = segmentation.resilient_partition
    monkeypatch.setattr(segmentation, "resilient_partition",
                        lambda catalog: original(catalog, target=3))

    def operator(inp, ids):
        chunk = next(s["chunk_id"] for s in inp["sources"]
                     if "- **Società:** Aurora Servizi S.r.l." in s["content"])
        return json.dumps({"version": 12, "facts": {"company": {
            "subject": "Aurora Servizi S.r.l.", "entity": "company", "aspect": "identity",
            "property": "denominazione", "value": "Aurora Servizi S.r.l.",
            "evidence": [{"chunk_id": chunk, "quote": "- **Società:** Aurora Servizi S.r.l."}]}},
            "fields": {"p2.s1": {"label": "Operatore economico", "kind": "data",
                                 "disposition": "value", "conditions": [],
                                 "property": "denominazione", "fact_id": "company",
                                 "form": {"paragraph_id": "paragraph:2",
                                          "quote": "operatore economico [[p2.s1]]"}}},
            "bindings": {}, "conditions": {}, "choice_groups": {}, "clarifications": {},
            "coverage": []})

    broken = GenerationTruncatedError(49152, 49152)
    model = Model({1: [broken], 2: [broken], 3: [operator]})
    upload = await api.post("/api/projects/alpha/forms", files={
        "file": ("modulo.docx", layout.original)})
    await source(api, sources[0]["content"], name="profili.md")
    monkeypatch.setattr(adapter, "request_model_content", model)
    session = (await api.post(BASE, json={"form_id": upload.json()["id"], "engine": "global",
                                         "start_in_chat": True})).json()
    save_conversation_turn(session["conversation_id"], {
        "question": f"Compila la domanda. Il firmatario scelto per questa domanda è {SIGNER}.",
        "answer": "Avvio della compilazione", "generation_status": "direct",
        "compilation": {"session_id": session["id"], "action": "start"}})
    session = (await analyze(api, session)).json()
    data = state_of(session)["global_compilation"]
    reuse = {f["candidate_id"] for f in data["audit"]["fields"] if f.get("source_reuse")}
    reuse &= unassessed_fields(data)  # Reuse into the p0/p1 sections whose owners failed.
    assert reuse
    states = applicability_states(data)
    assert {states[c]["state"] for c in reuse} == {"NON_ACCERTATA"}
    ids, _, _ = dialogue.export_selection(session)
    assert "p2.s1" in ids and not set(ids) & reuse
    forced = await draft(api, session, [*ids, sorted(reuse)[0]])
    assert forced.status_code == 409 and "non accertata" in forced.json()["detail"]
    with connection() as db:  # The same slot asked again later: a USER value is independent.
        db.execute("BEGIN IMMEDIATE")
        row, state = adapter._row_state(db, "alpha", session["id"], session["version"])
        state["global_compilation"]["dialogue"].update(phase="READY", error=None, question={
            "kind": "clarifications", "field_ids": [], "message": "?", "slots": [{
                "slot": 1, "key": "field:p0.s0", "kind": "value", "field_ids": ["p0.s0"],
                "label": "Il sottoscritto", "condition": "", "context": "Il sottoscritto",
                "alternatives": [], "impact": 1, "question": "Chi firma?"}]})
        session = sessions.save_state(db, row, state, "test_reask")
    session, errors = dialogue.apply_replies("alpha", session["id"], session["version"], [
        ClarificationReply(slot=1, action="VALUE", value=SIGNER, user_quote=SIGNER)], SIGNER)
    assert not errors
    session = (await analyze(api, session)).json()
    ids, _, users = dialogue.export_selection(session)
    assert "p0.s0" in users and not set(ids) & unassessed_fields(
        state_of(session)["global_compilation"])
    exported = await draft(api, session, ids, users)
    assert exported.status_code == 200, exported.text


# T9: interruption, resumption from the digest cache, invalidation by a different request.


class Interrupted(Exception):
    pass


@pytest.mark.anyio
async def test_resumption_reuses_validated_segments_and_never_merges_other_requests(
    api, provider, monkeypatch,
):
    async def late_crash(inp, ids):
        await asyncio.sleep(0.05)
        raise Interrupted("processo interrotto")

    model = Model({3: [late_crash]})
    session = await started(api, monkeypatch, form_docx(150), model)
    with pytest.raises(Interrupted):
        await analyze(api, session)
    stored = state_of(session)
    assert stored["status"] == "FAILED" and stored["global_compilation"]["audit"] is None
    folder = get_storage_path().resolve() / "alpha" / "_global_compilations" / session["id"]
    assert len(list((folder / "segment-cache").glob("*.json"))) == 2  # Validated work kept.
    current = (await api.get(f"{BASE}/{session['id']}")).json()
    resumed = adapter.control_session("alpha", session["id"], current["version"], "resume")
    model.plan, model.calls = {}, []
    session = (await analyze(api, resumed)).json()
    report = segments_of(session)
    assert [s["source"] for s in report["segments"]] == ["cache", "cache", "called"]
    assert model.calls == [(3, 0)]  # Only the missing segment is requested again.
    ids, _, _ = dialogue.export_selection(session)
    assert len(ids) == 6
    exported = await draft(api, session, ids)
    assert exported.status_code == 200, exported.text  # Same final gates after a resume.

    # A corrupted cache entry is ignored, never trusted.
    entry = sorted((folder / "segment-cache").glob("*.json"))[0]
    entry.write_text('{"raw": "{\\"version\\":12,')
    digest = entry.stem
    with connection() as db:
        db.execute("BEGIN IMMEDIATE")
        row, state = adapter._row_state(db, "alpha", session["id"],
                                        state_of_version(session))
        state["global_compilation"].update(phase="PENDING")
        state.update(status="CREATED")
        session = sessions.save_state(db, row, state, "test_reanalysis")
    model.calls = []
    session = (await analyze(api, session)).json()
    report = segments_of(session)
    corrupted = next(s for s in report["segments"] if s["digest"] == digest)
    assert corrupted["source"] == "called" and len(model.calls) == 1

    # A different SOURCE changes every request: nothing from the cache is merged.
    await source(api, "Codice: ABC123\nNumero: 4455", name="altri.md")
    with connection() as db:
        db.execute("BEGIN IMMEDIATE")
        row, state = adapter._row_state(db, "alpha", session["id"],
                                        state_of_version(session))
        state["global_compilation"].update(phase="PENDING")
        state.update(status="CREATED")
        session = sessions.save_state(db, row, state, "test_reanalysis")
    model.calls = []
    session = (await analyze(api, session)).json()
    assert [s["source"] for s in segments_of(session)["segments"]] == ["called"] * 3
    assert len(model.calls) == 3


def state_of_version(session):
    with connection() as db:
        return sessions._row(db, "alpha", session["id"])["version"]


# T10: opt-in only; earlier revisions keep their behaviour.


def test_resilient_generation_is_opt_in_and_separate_from_p3(monkeypatch):
    from app.compilation_global_sessions import initial_state

    layout = inspect_docx(form_docx(3))
    monkeypatch.delenv("MAPI_GENERATION_REVISION")
    assert "generation_revision" not in initial_state(layout)
    monkeypatch.setenv("MAPI_GENERATION_REVISION", "1")
    assert initial_state(layout)["generation_revision"] == 1
    monkeypatch.setenv("MAPI_GENERATION_REVISION", "2")
    assert initial_state(layout)["generation_revision"] == 2
    assert unassessed_fields({"generation_revision": 1, "audit": {"generation_segments": {
        "unanalyzed_candidate_ids": ["p0.s0"]}}}) == set()
