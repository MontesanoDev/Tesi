"""Local USER questions and source-role coverage, no real-model quality claims."""

import json
from io import BytesIO

import pytest
from docx import Document
from test_compilation_catalog import encoded
from test_compilation_global_sessions import BASE, proposal, resolve, start
from test_compilation_global_sessions import api as global_api
from test_compilation_global_sessions import provider as global_provider
from test_compilation_sessions import source

from app import compilation_global_dialogue as dialogue
from app import compilation_global_sessions as adapter
from app.compilation_chat import ClarificationReply
from app.compilation_global import build_global_input, validate_global_proposal
from app.docx_templates import inspect_docx

api, provider = global_api, global_provider


@pytest.fixture
def anyio_backend():
    return "asyncio"


@pytest.fixture
def professional():
    doc = Document()
    doc.add_paragraph("Il Responsabile delle verifiche indica:")
    for label in ["Nome e cognome", "Registro", "Data iscrizione"]:
        doc.add_paragraph(label + ": _________")
    layout = inspect_docx(encoded(doc))
    sources = [{"chunk_id": -7, "file_id": -3, "chunk_index": 0, "source_name": "profilo.md",
        "role": "source", "scope": "global", "project_id": None, "category": "company",
        "content": "### Dott. Berta Neri - Responsabile delle verifiche\n"
                   "Nome: Berta Neri.\nRegistro: R-236.\nData iscrizione: non disponibile."}]
    data = build_global_input(layout, "alpha", sources)
    payload = {"version": 11, "conditions": {}, "choice_groups": {}, "clarifications": {},
               "facts": {}, "fields": {}, "bindings": {}}
    return layout, sources, data, payload


def snapshot(case):
    layout, sources, model, payload = case
    data = adapter.initial_state(layout)
    data.update(audit=validate_global_proposal(json.dumps(payload), layout, "alpha", sources),
                catalog=model["form"], reading_plan=model["reading_plan"])
    return data


def request_field(paragraph=3, label="Data iscrizione"):
    return {"label": label, "form": {"paragraph_id": f"paragraph:{paragraph}",
            "quote": f"{label}: [[p{paragraph}.s0]]"}, "conditions": [],
            "reason": "Dato pertinente non trovato nel record esaminato"}


def add_professional_values(case):
    payload = case[3]
    payload["bindings"]["b"] = {"required_role": "Responsabile delle verifiche",
        "source_role": "Responsabile delle verifiche", "form": {
            "paragraph_id": "paragraph:0", "quote": "Il Responsabile delle verifiche indica:"},
        "evidence": [{"chunk_id": -7,
                      "quote": "### Dott. Berta Neri - Responsabile delle verifiche"}]}
    for index, (label, value) in enumerate([("Nome e cognome", "Berta Neri"),
                                           ("Registro", "R-236")], 1):
        fid, cid = f"f{index}", f"p{index}.s0"
        payload["facts"][fid] = {"subject": "Berta Neri", "entity": "person",
            "aspect": "professional", "property": label, "value": value,
            "evidence": [{"chunk_id": -7, "quote": "Nome: Berta Neri." if index == 1
                          else "Registro: R-236."}]}
        payload["fields"][cid] = {"label": label, "kind": "data", "disposition": "value",
            "form": {"paragraph_id": f"paragraph:{index}",
                     "quote": f"{label}: [[{cid}]]"}, "conditions": [], "fact_id": fid,
            "property": label, "binding_id": "b"}


def test_role_reading_targets_are_references_not_factual_values(professional):
    plan = professional[2]["reading_plan"]
    target = next(t for t in plan["source_role_targets"]
                  if t["form_paragraph_id"] == "paragraph:0")
    assert target["candidate_ids"] == ["p1.s0", "p2.s0", "p3.s0"]
    assert target["source_records"][0]["chunk_id"] == -7
    assert "value" not in target and "subject" not in target
    assert len(plan["units"]) == 1
    assert all(r["quote"] in professional[1][0]["content"]
               for r in plan["source_property_lines"])
    assert {r["quote"] for r in plan["source_property_lines"]} == {
        "Nome: Berta Neri.", "Registro: R-236.", "Data iscrizione: non disponibile."}
    professional[1][0]["category"] = "general"
    data = build_global_input(professional[0], "alpha", professional[1])
    assert data["reading_plan"]["source_role_targets"] == []
    assert data["reading_plan"]["source_property_lines"] == []


def test_literal_property_reading_budget_never_silently_truncates():
    from app.compilation_reading_plan import MAX_PROPERTY_LINES, property_lines

    source = {"chunk_id": -1, "category": "company",
              "content": "\n".join(f"Dato {i}: valore {i}" for i in range(300))}
    records, omitted = property_lines([source])
    assert len(records) == MAX_PROPERTY_LINES and omitted == 300 - MAX_PROPERTY_LINES
    assert all(r["quote"] in source["content"] and "value" not in r for r in records)


def test_missing_source_role_proposals_are_flagged_not_asked_from_user(professional):
    professional[3]["clarifications"]["p1.s0"] = request_field(1, "Nome e cognome")
    data = snapshot(professional)
    assert dialogue.question_pool(data) == []
    context = dialogue.model_context(data, [])
    assert context["unassessed_source_roles"] == [
        {"role": "Responsabile delle verifiche", "positions": 3}]
    assert data["audit"]["resolved"] == 0


def test_source_role_properties_and_missing_local_field_stay_separate(professional):
    add_professional_values(professional)
    professional[3]["clarifications"]["p3.s0"] = request_field()
    data = snapshot(professional)
    fields = {f["candidate_id"]: f for f in data["audit"]["fields"]}
    assert fields["p1.s0"]["validation"] == fields["p2.s0"]["validation"] == "GROUNDED_PROPOSAL"
    assert fields["p3.s0"]["source_evidence"] == []
    assert fields["p3.s0"]["value"] is None
    pool = dialogue.question_pool(data)
    assert len(pool) == 1 and pool[0]["kind"] == "value"
    assert pool[0]["field_ids"] == ["p3.s0"]
    assert not dialogue.model_context(data, pool)["unassessed_source_roles"]


@pytest.mark.parametrize("mutation", ["unknown", "quote", "source_value", "duplicate"])
def test_invalid_user_request_is_quarantined(professional, mutation):
    record = request_field()
    key = "p3.s0"
    if mutation == "unknown":
        key = "p999.s0"
    elif mutation == "quote":
        record["form"]["quote"] = "Una richiesta non presente"
    elif mutation == "source_value":
        record["value"] = "Inventato"
    else:
        add_professional_values(professional)
        key, record = "p1.s0", request_field(1, "Nome e cognome")
    professional[3]["clarifications"][key] = record
    data = snapshot(professional)
    assert data["audit"]["clarifications"][key]["validation"] == "REJECTED"


async def local_request_session(api, monkeypatch, *, conditional=False):
    doc = Document()
    doc.add_paragraph("Dati aziendali")
    table = doc.add_table(rows=3, cols=2)
    for row, label in zip(table.rows, ["Denominazione sociale", "Sede legale",
                                     "Carica del sottoscrittore"], strict=True):
        row.cells[0].text = label
    uploaded = await api.post("/api/projects/alpha/forms", files={
        "file": ("modulo.docx", encoded(doc))})
    assert uploaded.status_code == 201
    await source(api, "Denominazione sociale: Aurora Servizi S.r.l.\nSede legale: Via del Sole 9")

    async def request(client, settings, body, timeout, **kwargs):
        data = json.loads(body["messages"][1]["content"])
        payload = proposal(data, conditional=conditional)
        if conditional:
            payload["conditions"]["c"]["scope"] = [
                {"kind": "candidate", "id": cid} for cid in payload["fields"]]
        pid = next(p["id"] for b in data["form"]["blocks"] for p in b["paragraphs"]
                   if p["text"] == "Carica del sottoscrittore")
        payload["clarifications"] = {"t0.r2.c1": {"label": "Carica del sottoscrittore",
            "form": {"paragraph_id": pid, "quote": "Carica del sottoscrittore"},
            "conditions": [], "reason": "Scelta del soggetto che sottoscrive questa pratica"}}
        return json.dumps(payload), "simulated", {}

    monkeypatch.setattr(adapter, "request_model_content", request)
    s = (await api.post(BASE, json={"form_id": uploaded.json()["id"],
                                  "engine": "global", "start_in_chat": True})).json()
    return await resolve(api, s)


@pytest.mark.anyio
@pytest.mark.parametrize("action", ["SKIP", "UNKNOWN", "VALUE"])
async def test_local_user_reply_never_defers_independent_company_values(
    api, provider, monkeypatch, action,
):
    s = await local_request_session(api, monkeypatch)
    assert s["chat"]["question"]["slots"][0]["field_ids"] == ["t0.r2.c1"]
    args = {"value": "Delegato"} if action == "VALUE" else {}
    s, errors = dialogue.apply_replies("alpha", s["id"], s["version"], [ClarificationReply(
        slot=1, action=action, user_quote="Delegato" if args else "lascia vuoto", **args)],
        "Delegato" if args else "lascia vuoto")
    assert not errors
    s = await resolve(api, s)
    ids, conditions, users = dialogue.export_selection(s)
    assert ids == ["t0.r0.c1", "t0.r1.c1"] and not conditions
    assert users == (["t0.r2.c1"] if action == "VALUE" else [])
    generated = await api.post(f"{BASE}/{s['id']}/global-draft", json={"version": s["version"],
        "candidate_ids": ids, "user_candidate_ids": users, "allow_unresolved": True})
    assert generated.status_code == 200, generated.text
    downloads = generated.json()["last_generation"]["downloads"]
    report = (await api.get(downloads["report"])).json()
    assert report["written_field_count"] == (3 if args else 2)
    assert [f["provenance"] for f in report["fields"]] == (
        ["SOURCE", "SOURCE", "USER"] if args else ["SOURCE", "SOURCE"])
    word = Document(BytesIO((await api.get(downloads["docx"])).content))
    assert word.tables[0].cell(2, 1).text == ("Delegato" if args else "")
    assert (await api.get(f"{BASE}/{s['id']}")).json()["status"] == "GENERATED"


@pytest.mark.anyio
@pytest.mark.parametrize("question", [None, float("nan"), float("inf"), "test\ud800"])
async def test_malformed_chat_payload_is_a_controlled_client_error(api, provider, question):
    response = await api.post("/api/projects/alpha/answer",
        content=json.dumps({"question": question}).encode(),
        headers={"Content-Type": "application/json"})
    assert response.status_code == 422, response.text


@pytest.mark.anyio
async def test_random_chat_text_does_not_execute_commands_or_crash(api, provider, monkeypatch):
    from app import intents

    async def route(*args, **kwargs):
        return json.dumps({"intent": "reply", "answer": "Puoi precisare la richiesta?",
                           "target": "source", "queries": []}), "simulated", {}

    monkeypatch.setattr(intents, "request_model_content", route)
    messages = ["salta", "non lo so", "🧪🧪", "' OR 1=1 --", "<script>alert(1)</script>",
                "{} [] null", "basta… anzi no", "\x00\x01input", "società\u200bdi ingegneria"]
    for message in messages:
        response = await api.post("/api/projects/alpha/answer", json={"question": message})
        assert response.status_code == 200, (message, response.text)
        assert response.json().get("compilation") is None


@pytest.mark.anyio
@pytest.mark.parametrize("raw", ["[" * 10_000, '"test\\ud800"',
                                 '{"version":NaN}', '{"version":Infinity}', '[]'],
                         ids=["deep", "unicode", "nan", "infinite", "array"])
async def test_malformed_model_response_fails_without_http_500(api, provider, monkeypatch, raw):
    async def invalid(*args, **kwargs):
        return raw, "simulated", {}

    monkeypatch.setattr(adapter, "request_model_content", invalid)
    s = await start(api)
    response = await api.post(f"{BASE}/{s['id']}/resolve", json={"version": s["version"]})
    assert response.status_code == 502, response.text
    state = (await api.get(f"{BASE}/{s['id']}")).json()
    assert state["status"] == "FAILED" and state["last_generation"] is None


@pytest.mark.anyio
async def test_invalid_unicode_router_output_does_not_corrupt_retry_request(
    api, provider, monkeypatch,
):
    from app import intents

    calls = []

    async def route(client, settings, body, timeout, **kwargs):
        json.dumps(body, ensure_ascii=False).encode("utf-8")  # Actual HTTP JSON encoding boundary.
        calls.append(body)
        return json.dumps({"intent": "reply", "answer": "\ud800" if len(calls) == 1 else
                           "Puoi precisare la richiesta?", "target": "source", "queries": []},
                          ensure_ascii=False), "simulated", {}

    monkeypatch.setattr(intents, "request_model_content", route)
    response = await api.post("/api/projects/alpha/answer", json={"question": "🧪test"})
    assert response.status_code == 200 and len(calls) == 2
    assert response.json()["answer"] == "Puoi precisare la richiesta?"


@pytest.mark.anyio
async def test_condition_confirmation_and_local_skip_keep_all_source_values(
    api, provider, monkeypatch,
):
    from app import intents

    s = await local_request_session(api, monkeypatch, conditional=True)
    slots = s["chat"]["question"]["slots"]
    assert [q["kind"] for q in slots] == ["applicability", "value"]
    assert "2 dati collegati" in slots[0]["question"]
    calls = []

    async def route(client, settings, body, timeout, **kwargs):
        calls.append(body)
        if kwargs["response_schema"]["title"] == "ReplyScopeChecks":
            data = json.loads(body["messages"][1]["content"])
            assert len(data["proposals"]) == 1
            return '{"checks":[{"slot":1,"result":"SUPPORTED"}]}', "simulated", {}
        return (json.dumps({"intent": "answer_fields", "answer": "", "queries": [],
            "target": "form", "replies": [
                {"slot": 1, "action": "CONDITION_TRUE", "user_quote": "confermo il primo punto"},
                {"slot": 2, "action": "SKIP", "user_quote": "la carica lascia vuota"}]}),
                "simulated", {})

    monkeypatch.setattr(intents, "request_model_content", route)
    response = await api.post("/api/projects/alpha/answer", json={
        "question": "confermo il primo punto, la carica lascia vuota",
        "conversation_id": s["conversation_id"], "compilation_session_id": s["id"],
        "compilation_version": s["version"]})
    assert response.status_code == 200, response.text
    s = await resolve(api, (await api.get(f"{BASE}/{s['id']}")).json())
    assert dialogue.export_selection(s)[0] == ["t0.r0.c1", "t0.r1.c1"]
    assert s["global_compilation"]["dialogue"]["answers"]["field:t0.r2.c1"][
        "affected_fields"] == ["t0.r2.c1"]
    assert len(calls) == 2


@pytest.mark.anyio
async def test_narrow_skip_cannot_silently_defer_a_condition_block(api, provider, monkeypatch):
    from app import intents

    s = await local_request_session(api, monkeypatch, conditional=True)

    async def route(client, settings, body, timeout, **kwargs):
        if kwargs["response_schema"]["title"] == "ReplyScopeChecks":
            data = json.loads(body["messages"][1]["content"])
            assert data["proposals"][0]["action"] == "SKIP"
            assert data["slots"][0]["affected_field_count"] == 2
            return '{"checks":[{"slot":1,"result":"AMBIGUOUS"}]}', "simulated", {}
        return (json.dumps({"intent": "answer_fields", "answer": "", "queries": [],
            "target": "form", "replies": [{"slot": 1, "action": "SKIP",
                                            "user_quote": "solo la carica lascia vuota"}]}),
                "simulated", {})

    monkeypatch.setattr(intents, "request_model_content", route)
    response = await api.post("/api/projects/alpha/answer", json={
        "question": "solo la carica lascia vuota", "conversation_id": s["conversation_id"],
        "compilation_session_id": s["id"], "compilation_version": s["version"]})
    assert response.status_code == 200, response.text
    assert "più dati collegati" in response.json()["answer"]
    assert (await api.get(f"{BASE}/{s['id']}")).json() == s
