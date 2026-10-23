"""E1 structural fidelity and negative controls; no model quality assertions."""

import copy
import json
from io import BytesIO
from pathlib import Path

import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from lxml import etree

from app.compilation_form_view import (
    compact_form,
    compact_input,
    indexed_form,
    indexed_input,
    validate_form_view,
    validate_indexed_form,
)
from app.compilation_global import build_global_input
from app.docx_templates import inspect_docx

DEMO = Path(__file__).resolve().parents[2] / "demo-documents"


def data_for(doc):
    stream = BytesIO()
    doc.save(stream)
    return build_global_input(inspect_docx(stream.getvalue(), require_candidates=False), "p", [])


@pytest.fixture
def sample():
    doc = Document()
    doc.add_heading("Società Alfa", level=1)
    doc.add_paragraph("Nome: __________ CF: __________")
    doc.add_heading("Società Beta", level=1)
    doc.add_paragraph("Nome: __________")
    doc.add_paragraph("Firma: __________")
    return data_for(doc)


@pytest.mark.parametrize("relative,expected", [
    ("benchmark-minervino/v1/input/modulo.docx", 267),
    ("benchmark-compilazione/v1/modulo-benchmark.docx", 55),
    ("bandi/catanzaro-dl-cse/modello/domanda-partecipazione.docx", None),
])
@pytest.mark.parametrize("project,validator", [(compact_input, validate_form_view),
                                               (indexed_input, validate_indexed_form)])
def test_real_originals_keep_inventory_text_and_topology_without_gold(
        relative, expected, project, validator):
    original = (DEMO / relative).read_bytes()
    data = build_global_input(inspect_docx(original, require_candidates=False), "p", [])
    before = copy.deepcopy(data)
    projected = project(data)
    view, full = projected["form"], data["form"]
    checked = validator(full, view)
    assert checked["candidate_count"] == len(full["candidates"])
    if expected is not None:
        assert checked["candidate_count"] == expected
    assert data == before
    assert view["blocks"] == full["blocks"]
    assert {k: v for k, v in projected.items() if k != "form"} == {
        k: v for k, v in data.items() if k != "form"}
    assert len(json.dumps(view)) < len(json.dumps(full))
    assert "range" not in json.dumps(view) and "slot_context" not in json.dumps(view)


def test_grouped_contexts_keep_headings_continuations_types_and_signature(sample):
    form = sample["form"]
    view = compact_form(form)
    assert view["contexts"]["paragraph:1"]["fields"] == ["p1.s0", "p1.s1"]
    assert view["contexts"]["paragraph:1"]["defaults"]["section_ids"] == ["paragraph:0"]
    assert view["contexts"]["paragraph:3"]["defaults"]["section_ids"] == ["paragraph:2"]
    assert view["contexts"]["paragraph:4"]["field_options"]["p4.s0"]["write_protected"]
    # A detached caller cannot change the canonical text or its nested lists.
    view["blocks"][0]["paragraphs"][0]["text"] = "changed"
    assert form["blocks"][0]["paragraphs"][0]["text"] == "Società Alfa"


@pytest.mark.parametrize("legacy", [True, False])
@pytest.mark.parametrize("checked", [True, False])
def test_native_controls_keep_state_without_inventing_slots_ignored_by_parser(legacy, checked):
    doc = Document()
    p = doc.add_paragraph("Referente: __________ ")
    if legacy:
        begin = OxmlElement("w:fldChar")
        begin.set(qn("w:fldCharType"), "begin")
        control = etree.SubElement(etree.SubElement(begin, qn("w:ffData")), qn("w:checkBox"))
        etree.SubElement(control, qn("w:checked")).set(qn("w:val"), str(int(checked)))
        p.add_run()._r.append(begin)
        end = OxmlElement("w:fldChar")
        end.set(qn("w:fldCharType"), "end")
        p.add_run()._r.append(end)
    else:
        w14 = "http://schemas.microsoft.com/office/word/2010/wordml"
        sdt = OxmlElement("w:sdt")
        control = etree.SubElement(etree.SubElement(sdt, qn("w:sdtPr")), f"{{{w14}}}checkbox")
        etree.SubElement(control, f"{{{w14}}}checked").set(f"{{{w14}}}val", str(int(checked)))
        run = etree.SubElement(etree.SubElement(sdt, qn("w:sdtContent")), qn("w:r"))
        etree.SubElement(run, qn("w:t")).text = "☐"
        p._p.append(sdt)
    data = data_for(doc)
    view = compact_form(data["form"])
    opts = view["contexts"]["paragraph:0"]["field_options"]["p0.cb0"]
    assert opts["initial_checked"] is checked
    assert opts["type"] == "checkbox"
    assert opts["control_type"] == ("legacy" if legacy else "content_control")
    # Current parser skips inline slots in paragraphs with native controls.
    # E1 must preserve its inventory, not silently create unsupported targets.
    assert "p0.s0" not in data["form"]["candidates"]
    assert view["contexts"]["paragraph:0"]["fields"] == ["p0.cb0"]
    direct = indexed_form(data["form"])["candidates"]["p0.cb0"]
    assert direct["context_id"] == "paragraph:0"
    assert direct["type"] == "checkbox" and direct["initial_checked"] is checked
    assert direct["control_type"] == opts["control_type"]


def test_merged_nested_tables_repeat_headers_and_readonly_parts_are_exact():
    doc = Document()
    doc.add_heading("Prestazioni", level=1)
    table = doc.add_table(rows=3, cols=2)
    table.cell(0, 0).merge(table.cell(0, 1)).text = "Titolo comune"
    table.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
    table.cell(1, 0).merge(table.cell(2, 0)).text = "Etichetta verticale"
    nested = table.cell(1, 1).add_table(rows=1, cols=2)
    nested.cell(0, 0).text = "Dato interno"
    doc.sections[0].header.paragraphs[0].text = "Istruzione in sola lettura"
    full = data_for(doc)["form"]
    view = compact_form(full)
    assert view["blocks"] == full["blocks"]
    outer, inner = view["blocks"][1]["tables"]
    assert outer["repeat_header_rows"] == [0]
    assert outer["rows"][0][0]["grid_span"] == "2"
    assert outer["rows"][1][0]["vertical_merge"] == "restart"
    assert inner["parent_cell_id"] == outer["rows"][1][1]["id"]
    assert view["read_only_parts"] == full["read_only_parts"]
    direct = indexed_form(full)
    assert direct["blocks"] == full["blocks"]
    assert direct["read_only_parts"] == full["read_only_parts"]


def test_different_field_references_in_one_context_remain_explicit_overrides(sample):
    form = sample["form"]
    # Exercise legitimate unequal per-field context without promoting an ID heuristic.
    form["candidates"]["p1.s1"]["anchor_paragraph_ids"] = ["paragraph:0", "paragraph:1"]
    form["candidates"]["p1.s1"]["value_type"] = "tax_code"
    view = compact_form(form)
    assert "anchor_paragraph_ids" not in view["contexts"]["paragraph:1"]["defaults"]
    opts = view["contexts"]["paragraph:1"]["field_options"]
    assert opts["p1.s1"]["anchor_paragraph_ids"] == ["paragraph:0", "paragraph:1"]
    assert opts["p1.s1"]["value_type"] == "tax_code"


@pytest.mark.parametrize("mutation", ["omit", "duplicate", "role", "signature", "type",
                                      "text", "marker", "unknown", "reference"])
def test_independent_checker_rejects_lost_or_misassociated_information(sample, mutation):
    full = sample["form"]
    view = compact_form(full)
    group = view["contexts"]["paragraph:1"]
    if mutation == "omit":
        group["fields"].remove("p1.s1")
    elif mutation == "duplicate":
        group["fields"].append("p1.s0")
    elif mutation == "role":
        group["defaults"]["section_ids"] = ["paragraph:2"]
    elif mutation == "signature":
        view["contexts"]["paragraph:4"]["field_options"]["p4.s0"].pop("write_protected")
    elif mutation == "type":
        group.setdefault("field_options", {}).setdefault("p1.s0", {})["type"] = "checkbox"
    elif mutation in {"text", "marker"}:
        p = view["blocks"][1]["paragraphs"][0]
        key = "text" if mutation == "text" else "text_with_fields"
        p[key] = p[key].replace("Nome", "Qualifica")
    elif mutation == "unknown":
        group["fields"].append("p999.s0")
    else:
        group["defaults"]["anchor_paragraph_ids"] = ["paragraph:999"]
    with pytest.raises(ValueError):
        validate_form_view(full, view)


def test_checker_rejects_changed_table_column_association():
    doc = Document()
    table = doc.add_table(rows=2, cols=3)
    for i, text in enumerate(["Servizio", "Fascia A", "Fascia B"]):
        table.cell(0, i).text = text
    table.cell(1, 0).text = "Strutture"
    full = data_for(doc)["form"]
    view = compact_form(full)
    row = view["blocks"][0]["tables"][0]["rows"][1]
    row[1], row[2] = row[2], row[1]
    with pytest.raises(ValueError, match="topology"):
        validate_form_view(full, view)


def test_future_candidate_metadata_cannot_be_silently_dropped(sample):
    sample["form"]["candidates"]["p1.s0"]["new_semantic_constraint"] = "required"
    with pytest.raises(ValueError, match="Unsupported candidate properties"):
        compact_form(sample["form"])


def test_cannot_apply_projection_twice(sample):
    with pytest.raises(ValueError, match="canonical"):
        compact_form(compact_form(sample["form"]))


def test_direct_index_is_explicit_and_shared_context_preserves_scope(sample):
    full = sample["form"]
    view = indexed_form(full)
    assert view["candidates"].keys() == full["candidates"].keys()
    assert view["candidates"]["p1.s0"] == {"type": "text", "context_id": "paragraph:1"}
    assert view["contexts"]["paragraph:1"]["section_ids"] == ["paragraph:0"]
    assert view["contexts"]["paragraph:3"]["section_ids"] == ["paragraph:2"]
    assert view["candidates"]["p4.s0"]["write_protected"] is True
    assert all("fields" not in context for context in view["contexts"].values())
    view["contexts"]["paragraph:1"]["section_ids"].append("paragraph:2")
    assert full["candidates"]["p1.s0"]["section_ids"] == ["paragraph:0"]


def test_direct_index_keeps_unequal_references_as_field_overrides(sample):
    full = sample["form"]
    full["candidates"]["p1.s1"]["anchor_paragraph_ids"] = ["paragraph:0", "paragraph:1"]
    full["candidates"]["p1.s1"]["value_type"] = "tax_code"
    view = indexed_form(full)
    assert "anchor_paragraph_ids" not in view["contexts"]["paragraph:1"]
    assert view["candidates"]["p1.s0"]["anchor_paragraph_ids"] == ["paragraph:1"]
    assert view["candidates"]["p1.s1"]["anchor_paragraph_ids"] == ["paragraph:0", "paragraph:1"]
    assert view["candidates"]["p1.s1"]["value_type"] == "tax_code"


@pytest.mark.parametrize("mutation", ["omit", "extra_field", "missing_type", "wrong_context",
                                      "missing_context", "extra_context", "shadowed_context",
                                      "role", "signature", "range", "unknown", "marker"])
def test_direct_index_checker_rejects_changed_information(sample, mutation):
    full = sample["form"]
    view = indexed_form(full)
    record = view["candidates"]["p1.s0"]
    if mutation == "omit":
        view["candidates"].pop("p1.s0")
    elif mutation == "extra_field":
        view["candidates"]["p999.s0"] = record.copy()
    elif mutation == "missing_type":
        record.pop("type")
    elif mutation == "wrong_context":
        record["context_id"] = "paragraph:3"
    elif mutation == "missing_context":
        view["contexts"].pop("paragraph:1")
    elif mutation == "extra_context":
        view["contexts"]["paragraph:999"] = {}
    elif mutation == "shadowed_context":
        view["contexts"]["paragraph:1"]["context_id"] = "paragraph:3"
    elif mutation == "role":
        view["contexts"]["paragraph:1"]["section_ids"] = ["paragraph:2"]
    elif mutation == "signature":
        view["candidates"]["p4.s0"].pop("write_protected")
    elif mutation == "range":
        record["range"] = [0, 10]
    elif mutation == "unknown":
        record["new_constraint"] = "silently ignored?"
    else:
        view["blocks"][1]["paragraphs"][0]["text_with_fields"] = "[[p1.s1]]"
    with pytest.raises(ValueError):
        validate_indexed_form(full, view)


def test_direct_projection_rejects_unknown_properties_and_double_projection(sample):
    full = sample["form"]
    with pytest.raises(ValueError, match="canonical"):
        indexed_form(indexed_form(full))
    full["candidates"]["p1.s0"]["future_constraint"] = True
    with pytest.raises(ValueError, match="Unsupported"):
        indexed_form(full)
