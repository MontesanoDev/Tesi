"""Only casing/whitespace equivalence, after exact proposed fact identity."""

import copy
import json
from io import BytesIO

import pytest
from docx import Document

from app.compilation_global import validate_global_proposal
from app.docx_templates import inspect_docx


def result(first="Ingegnere", second="ingegnere", change=None, *, explicit_review=False):
    document = Document()
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Titolo dichiarato"
    document.add_paragraph("Contesto distinto")
    stream = BytesIO()
    document.save(stream)
    layout = inspect_docx(stream.getvalue())
    source = {"chunk_id": -7, "file_id": -3, "chunk_index": 0, "source_name": "profilo.txt",
              "role": "source", "scope": "global", "project_id": None, "category": "company",
              "content": f"Alba: titolo {first}.\nAlba: titolo {second}."}
    fact = {"subject": "Alba", "entity": "company", "aspect": "identity",
            "property": "titolo", "value": first,
            "evidence": [{"chunk_id": -7, "quote": source["content"]}]}
    other = {**copy.deepcopy(fact), "value": second, **(change or {})}
    payload = {"version": 12, "conditions": {}, "bindings": {}, "choice_groups": {},
               "clarifications": {}, "coverage": [], "facts": {"a": fact, "b": other},
               "fields": {"t0.r0.c1": {"label": "Titolo dichiarato", "kind": "data",
                   "disposition": "value", "fact_id": "a", "property": "titolo",
                   "conditions": [], "form": {"paragraph_id": "paragraph:0",
                                              "quote": "Titolo dichiarato"}}}}
    # Keep both facts referenced without selecting the second one for the field.
    payload["conditions"]["context"] = {"form": {"paragraph_id": "paragraph:2",
        "quote": "Contesto distinto"}, "aspect": "identity", "operator": "semantic",
        "fact_ids": ["b"], "reason": "Contesto separato", "scope": []}
    if explicit_review:
        field = payload["fields"]["t0.r0.c1"]
        field.pop("fact_id")
        field.update(disposition="review", alternative_fact_ids=["a", "b"],
                     review_kind="ambiguous", reason="Scelta esplicita del modello")
    return validate_global_proposal(json.dumps(payload), layout, "alpha", [source])


@pytest.mark.parametrize("first,second", [
    ("Ingegnere", "ingegnere"), ("INGEGNERE", "Ingegnere"),
    ("Ingegnere civile", "ingegnere   civile"),
    ("Ingegnere civile", "ingegnere\tcivile"),
])
def test_equivalent_values_keep_original_and_source_proofs(first, second):
    audit = result(first, second)
    field = audit["fields"][0]
    assert field["validation"] == "GROUNDED_PROPOSAL"
    assert field["value"] == first and field["fact_id"] == "a"
    assert field["source_evidence"] == audit["facts"]["a"]["source_evidence"]
    assert not field.get("alternatives") and field["written_value"] is None
    assert audit["facts"]["b"]["value"] == second


@pytest.mark.parametrize("first,second", [
    ("Ingegnere", "Architetto"), ("Ing.", "Ingegnere"), ("A-1", "A1"),
    ("Albo 123", "Albo 124"), ("Élite", "Elite"), ("Straße", "Strasse"),
])
def test_distinct_spellings_still_require_review(first, second):
    field = result(first, second)["fields"][0]
    assert field["validation"] == "REVIEW_REQUIRED" and field["value"] is None
    assert {a["value"] for a in field["alternatives"]} == {first, second}


@pytest.mark.parametrize("change", [{"subject": "Altra azienda"}, {"subject": "alba"},
    {"entity": "other"}, {"property": "altra proprietà"}, {"aspect": "location"}])
def test_equivalent_text_never_merges_other_subjects_or_properties(change):
    audit = result(change=change)
    assert audit["fields"][0]["fact_id"] == "a"
    assert audit["fields"][0]["source_evidence"] == audit["facts"]["a"]["source_evidence"]


def test_invalid_evidence_is_not_rescued_by_typographic_equivalence():
    audit = result(change={"evidence": [{"chunk_id": -7, "quote": "prova inventata"}]})
    assert audit["facts"]["b"]["validation"] == "REJECTED"
    assert audit["fields"][0]["fact_id"] == "a"


def test_explicit_redundant_review_cannot_become_a_fake_choice_or_automatic_value():
    audit = result(explicit_review=True)
    field = audit["fields"][0]
    assert field["validation"] == "REJECTED" and not field.get("alternatives")
    assert field["value"] is field["written_value"] is None
    # No fact was selected by that review record; never promote it to value.
    assert [f["value"] for f in audit["facts"].values()] == ["Ingegnere", "ingegnere"]
    assert all(f["source_evidence"] for f in audit["facts"].values())
