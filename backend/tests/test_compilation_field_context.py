"""Physical field compatibility and bounded FORM lists; no model or writer."""

import json
from io import BytesIO

import pytest
from docx import Document

from app.compilation_global import build_global_input, validate_global_proposal
from app.compilation_sessions import validate_value
from app.docx_templates import inspect_docx

ROLE = "revisore"
INTRO = "componenti con qualifica di revisore:"
ADDRESS = "Via delle Querce 24, 00100 Roma (RM)"


def layout_for(doc):
    stream = BytesIO()
    doc.save(stream)
    return inspect_docx(stream.getvalue(), require_candidates=False)


def source(content, chunk=-7):
    return {"chunk_id": chunk, "file_id": -3, "chunk_index": -chunk - 7,
            "source_name": "profili.txt", "role": "source", "scope": "global",
            "category": "company", "project_id": None, "content": content}


def checked(layout, cid, value, label="dato aziendale"):
    return validate_value(layout, {"id": cid, "label": label, "entity": "company", "kind": "data"},
                          value, source("Dato attestato: " + value), "Dato attestato: " + value)


@pytest.mark.parametrize("kind", ["explicit", "split_line", "split_paragraph", "table"])
def test_address_is_rejected_for_physical_municipality_even_with_misleading_model_label(kind):
    doc = Document()
    if kind == "table":
        table = doc.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "Comune sede legale"
        cid = "t0.r0.c1"
    elif kind == "explicit":
        doc.add_paragraph("Comune: ________")
        cid = "p0.s0"
    elif kind == "split_line":
        doc.add_paragraph("Sede legale in ________ prov. ________ via/piazza ________ n. ________")
        cid = "p0.s0"
    else:
        doc.add_paragraph("P. IVA ________ sede legale in ________")
        doc.add_paragraph("prov. ________ via/piazza ________ n. ________")
        cid = "p0.s1"
    result = checked(layout_for(doc), cid, ADDRESS, "indirizzo completo")
    assert result["written_value"] is None
    assert "incompatible_locality_value" in result["validation_codes"]
    assert result["value"] == ADDRESS  # Rejection, never an extracted/replaced city.


@pytest.mark.parametrize("city", ["Roma", "L'Aquila", "Viareggio", "Piazza Armerina", "Bari (BA)"])
def test_locality_spelling_is_preserved_and_street_like_city_names_are_not_rejected(city):
    doc = Document()
    doc.add_paragraph("Comune: ________")
    assert checked(layout_for(doc), "p0.s0", city)["written_value"] == city


@pytest.mark.parametrize("label", ["Indirizzo sede legale", "Via/piazza", "Sede legale in",
                                  "Comune e indirizzo completo"])
def test_full_address_remains_valid_when_form_requests_address_or_has_no_split_components(label):
    doc = Document()
    doc.add_paragraph(label + " ________")
    assert checked(layout_for(doc), "p0.s0", ADDRESS, "Comune")["written_value"] == ADDRESS


def test_another_persons_address_field_does_not_retype_previous_location():
    doc = Document()
    doc.add_paragraph("Sede legale in ________")
    doc.add_paragraph("Recapito del responsabile: via ________ n. ________")
    assert checked(layout_for(doc), "p0.s0", ADDRESS)["written_value"] == ADDRESS


def list_case(*, barrier=None, people=3):
    doc = Document()
    doc.add_paragraph("Denominazione ________; " + INTRO)
    fields, facts, bindings, groups, sources = {}, {}, {}, [], []
    for member in range(people):
        name, number = f"Persona Cognome{member}", str(6234 + member)
        quote = f"Revisore: {name}; Qualifica: Perito; Numero iscrizione: {number}."
        chunk = -7 - member
        sources.append(source(quote, chunk))
        bid = f"b{member}"
        bindings[bid] = {"form": {"paragraph_id": "paragraph:0", "quote": INTRO},
                         "required_role": ROLE, "source_role": ROLE,
                         "evidence": [{"chunk_id": chunk, "quote": quote}]}
        group = []
        for offset, (label, value) in enumerate([
            ("Nome e cognome", name), ("Qualifica", "Perito"), ("Numero iscrizione", number),
        ]):
            if member == 0 and offset == 1 and barrier:
                if barrier == "blank":
                    doc.add_paragraph()
                elif barrier == "prose":
                    doc.add_paragraph("Un altro gruppo di dati")
                elif barrier == "heading":
                    doc.add_heading("Altro gruppo", 1)
                elif barrier == "table":
                    doc.add_table(rows=1, cols=1).cell(0, 0).text = "Altro gruppo"
                elif barrier == "control":
                    from test_compilation_catalog import native_checkbox

                    native_checkbox(doc.add_paragraph("Altra opzione"))
                elif barrier == "unmarked_role":
                    doc.add_paragraph("Altro incarico ________")
                else:
                    doc.add_paragraph("Nuovo incarico: ________")
            prefix = f"{member + 1}) " if offset == 0 else ""
            paragraph = doc.add_paragraph(prefix + label + " ________")
            # This fixture uses only body paragraphs except explicit barriers.
            index = list(doc.element.body.iter(paragraph._p.tag)).index(paragraph._p)
            cid, fid = f"p{index}.s0", f"f{member}_{offset}"
            facts[fid] = {"subject": name, "entity": "person", "aspect": "professional",
                          "property": label, "value": value,
                          "evidence": [{"chunk_id": chunk, "quote": quote}]}
            fields[cid] = {"label": label, "kind": "data", "disposition": "conditional",
                           "form": {"paragraph_id": f"paragraph:{index}",
                                    "quote": f"{label} [[{cid}]]"}, "conditions": ["c"],
                           "reason": "Applicabilità da confermare", "property": label,
                           "fact_id": fid, "binding_id": bid}
            group.append(cid)
        groups.append(group)
    condition = {"form": {"paragraph_id": "paragraph:0", "quote": INTRO},
                 "aspect": "participation", "operator": "semantic", "fact_ids": [],
                 "scope": [{"kind": "candidate", "id": cid} for cid in fields],
                 "reason": "Partecipazione non confermata"}
    proposal = {"version": 12, "uncertain_paragraph_ids": [], "facts": facts,
                "fields": fields, "bindings": bindings, "conditions": {"c": condition},
                "coverage": [], "choice_groups": {}, "clarifications": {}}
    return layout_for(doc), sources, proposal, groups


def audit(data):
    return validate_global_proposal(json.dumps(data[2]), data[0], "alpha", data[1])


def test_literal_role_scopes_bounded_numbered_members_without_changing_form():
    data = list_case()
    original_form = build_global_input(data[0], "alpha", data[1])["form"]
    result = audit(data)
    assert build_global_input(data[0], "alpha", data[1])["form"] == original_form
    for field in result["fields"]:
        if field["candidate_id"] not in data[2]["fields"]:
            continue
        assert field["subject_binding"]["result"] == "literal_match"
        assert field["proposed_value"] == data[2]["facts"][field["proposal"]["fact_id"]]["value"]
        assert field["condition_checks"]["c"]["result"] == "unknown"
        assert field["value"] is field["written_value"] is None
    last = next(f for f in result["fields"] if f["candidate_id"] == data[3][-1][-1])
    assert last["subject_binding"]["form_context_attribution"]["format"] == "numbered_form_list"
    assert result["resolved"] == result["writes"] == 0


@pytest.mark.parametrize("barrier", ["blank", "prose", "heading", "table", "control", "new_role",
                                     "unmarked_role"])
def test_numbered_role_scope_stops_at_physical_or_explicit_role_boundary(barrier):
    data = list_case(barrier=barrier)
    field = next(f for f in audit(data)["fields"] if f["candidate_id"] == data[3][0][-1])
    assert field["subject_binding"]["result"] == "review_required"
    assert "proposal_grounding" not in field


def test_role_before_last_placeholder_cannot_be_borrowed_for_a_following_list():
    data = list_case()
    paragraph = data[0].document.paragraphs[0]
    paragraph.text = "Revisore ________; altri dati:"
    for binding in data[2]["bindings"].values():
        binding["form"]["quote"] = "Revisore"
    data = layout_for(data[0].document), data[1], data[2], data[3]
    field = next(f for f in audit(data)["fields"] if f["candidate_id"] == data[3][0][-1])
    assert field["subject_binding"]["result"] == "review_required"


def test_one_numbered_member_cannot_mix_subjects_from_different_source_records():
    data = list_case()
    cid = data[3][1][1]
    data[2]["facts"][data[2]["fields"][cid]["fact_id"]]["subject"] = "Altra Persona"
    field = next(f for f in audit(data)["fields"] if f["candidate_id"] == data[3][1][-1])
    assert field["subject_binding"]["result"] == "review_required"


def test_long_numbered_list_does_not_create_an_unbounded_role_section():
    data = list_case(people=5)
    field = next(f for f in audit(data)["fields"] if f["candidate_id"] == data[3][-1][-1])
    assert field["subject_binding"]["result"] == "review_required"


@pytest.mark.parametrize("number", ["1)", "4)"])
def test_restarted_or_nonconsecutive_numbering_does_not_extend_original_role(number):
    data = list_case()
    paragraph = data[0].document.paragraphs[4]
    paragraph.text = paragraph.text.replace("2)", number)
    data = layout_for(data[0].document), data[1], data[2], data[3]
    field = next(f for f in audit(data)["fields"] if f["candidate_id"] == data[3][1][-1])
    assert field["subject_binding"]["result"] == "review_required"


def test_insufficient_role_quote_is_not_repaired_by_numbering():
    data = list_case()
    for binding in data[2]["bindings"].values():
        binding["form"]["quote"] = "componenti"
    field = next(f for f in audit(data)["fields"] if f["candidate_id"] == data[3][-1][-1])
    assert field["subject_binding"]["result"] == "review_required"


def test_explicitly_false_local_condition_remains_rejected_with_valid_numbered_role():
    data = list_case()
    data[1].append(source("Codice A12.", -50))
    for fid, prop in [("a", "codice uno"), ("b", "codice due")]:
        data[2]["facts"][fid] = {"subject": "Alba", "entity": "company", "aspect": "identity",
                                  "property": prop, "value": "A12",
                                  "evidence": [{"chunk_id": -50, "quote": "Codice A12."}]}
    cid = data[3][0][-1]
    data[2]["conditions"]["false"] = {
        "form": {"paragraph_id": data[2]["fields"][cid]["form"]["paragraph_id"],
                 "quote": "Numero iscrizione"},
        "operator": "different", "aspect": "identity", "fact_ids": ["a", "b"],
        "reason": "Due valori uguali, condizione di differenza falsa",
        "scope": [{"kind": "candidate", "id": cid}]}
    data[2]["fields"][cid].update(disposition="value", conditions=["false"])
    field = next(f for f in audit(data)["fields"] if f["candidate_id"] == cid)
    assert field["subject_binding"]["result"] == "literal_match"
    assert field["validation"] == "REJECTED"
    assert any("Condizione falsa" in e for e in field["validation_errors"])
    assert field.get("value") is field["written_value"] is None
    assert "proposal_grounding" not in field


def test_form_list_extension_is_not_implicitly_applied_to_v10():
    data = list_case()
    data[2]["version"] = 10
    for key in ["coverage", "choice_groups", "clarifications"]:
        data[2].pop(key)
    field = next(f for f in audit(data)["fields"] if f["candidate_id"] == data[3][-1][-1])
    assert field["subject_binding"]["result"] == "review_required"
