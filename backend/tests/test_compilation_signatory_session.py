"""Persisted start selection and literal SOURCE reuse; real services, fake transport."""

import json
from io import BytesIO

import pytest
from docx import Document
from test_compilation_global_sessions import BASE, api, provider, resolve
from test_compilation_selected_person import NAME, PROFILE
from test_compilation_sessions import source
from test_compilation_signatory_layouts import inline, table_document

from app import compilation_global_sessions as adapter
from app.compilation_global_dialogue import export_selection
from app.repository import save_conversation_turn

__all__ = ["api", "provider"]


@pytest.fixture
def anyio_backend():
    return "asyncio"


@pytest.mark.anyio
@pytest.mark.parametrize("shape,count", [("inline", 8), ("table", 2), ("tables", 2)])
async def test_recorded_user_selection_reaches_source_reuse_review_and_ordinary_export(
    api, provider, monkeypatch, shape, count,
):
    document = inline() if shape == "inline" else table_document(separate=shape == "tables")
    stream = BytesIO()
    document.save(stream)
    form = (await api.post("/api/projects/alpha/forms", files={
        "file": ("domanda.docx", stream.getvalue())})).json()
    await source(api, PROFILE, name="anagrafica.md")
    calls = []

    async def propose(client, settings, body, timeout, **kwargs):
        calls.append(body)
        return json.dumps({"version": 12, "facts": {}, "fields": {}, "bindings": {},
            "conditions": {}, "clarifications": {}, "choice_groups": {}, "coverage": []}), \
            "simulated", {}

    monkeypatch.setattr(adapter, "request_model_content", propose)
    s = (await api.post(BASE, json={"form_id": form["id"], "engine": "global",
                                   "start_in_chat": True})).json()
    save_conversation_turn(s["conversation_id"], {"question":
        "Compila la bozza. Per questa domanda il firmatario scelto è Ada Neri.",
        "answer": "Il modello non fornisce dati personali", "generation_status": "direct",
        "compilation": {"session_id": s["id"], "action": "start"}})
    s = await resolve(api, s)
    ids, conditions, users = export_selection(s)
    assert len(ids) == count and not conditions and not users
    assert len(calls) == 1
    report = (await api.get(f"{BASE}/{s['id']}/global-report")).json()
    assert report["source_reuse_count"] == count
    generated = await api.post(f"{BASE}/{s['id']}/global-draft", json={
        "version": s["version"], "candidate_ids": ids, "allow_unresolved": True,
        "user_quote": "Genera la bozza parziale"})
    assert generated.status_code == 200, generated.text
    downloads = generated.json()["last_generation"]["downloads"]
    output = Document(BytesIO((await api.get(downloads["docx"])).content))
    text = "\n".join([*(p.text for p in output.paragraphs),
        *(c.text for t in output.tables for r in t.rows for c in r.cells)])
    assert NAME in text and "NREADA85A52L219R" in text
    report = (await api.get(downloads["report"])).json()
    assert report["written_field_count"] == count
    assert all(f["provenance"] == "SOURCE" for f in report["fields"])
    assert len(calls) == 1  # No global inference repeated after reading the start decision.
