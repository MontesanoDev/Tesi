from io import BytesIO
from pathlib import Path

import httpx
import pytest
from docx import Document
from pypdf import PdfWriter

from app import project_forms
from app.db import connection, get_storage_path, init_database
from app.fact_extraction import load_project_source_chunks
from app.main import app
from app.repository import create_project, delete_project, search_project_evidence
from app.schemas import ProjectCreate


@pytest.fixture
def anyio_backend():
    return "asyncio"


@pytest.fixture
async def client(tmp_path, monkeypatch):
    monkeypatch.setenv("MAPI_DB_PATH", str(tmp_path / "test.db"))
    monkeypatch.setenv("MAPI_STORAGE_PATH", str(tmp_path / "uploads"))
    monkeypatch.setenv("MAPI_KNOWLEDGE_PATH", str(tmp_path / "knowledge"))
    init_database()
    for title in ("Primo", "Secondo"):
        create_project(ProjectCreate(title=title, description="Progetto di prova"))
    async with httpx.AsyncClient(
        transport=httpx.ASGITransport(app=app), base_url="http://test",
    ) as result:
        yield result


def document_bytes(suffix):
    if suffix in ("txt", "md"):
        return "Domanda di partecipazione\r\nRagione sociale: ___\r\nCittà: ___".encode()
    output = BytesIO()
    if suffix == "docx":
        doc = Document()
        doc.add_paragraph("Un modulo senza posizioni compilabili e comunque archiviabile.")
        doc.save(output)
    else:
        pdf = PdfWriter()
        pdf.add_blank_page(width=200, height=200)
        pdf.write(output)
    return output.getvalue()


@pytest.mark.anyio
@pytest.mark.parametrize("suffix", ["docx", "txt"])
async def test_archive_indexes_text_without_compiling_or_altering_original(client, suffix):
    original = document_bytes(suffix)
    response = await client.post(
        "/api/projects/primo/forms", files={"file": (f"domanda.{suffix}", original)},
    )
    assert response.status_code == 201, response.text
    form = response.json()
    assert form["kind"] == "form"
    assert form["status"] == "Indicizzato"
    assert form["chunk_count"] > 0
    assert "storage_path" not in form
    assert (await client.get("/api/projects/primo/forms")).json() == [form]
    project = (await client.get("/api/projects/primo")).json()
    assert project["files"] == [form]
    assert project["source_count"] == 0
    assert project["model_count"] == 1
    downloaded = await client.get(f"/api/projects/primo/forms/{form['id']}/download")
    assert downloaded.status_code == 200
    assert downloaded.content == original
    assert "attachment" in downloaded.headers["content-disposition"]
    assert not load_project_source_chunks("primo")
    with connection() as db:
        for table in ("document_chunks", "document_chunks_fts"):
            assert db.execute(f"SELECT COUNT(*) FROM {table}").fetchone()[0] == form["chunk_count"]
        chunks = db.execute(
            "SELECT c.project_id, c.file_id, f.kind FROM document_chunks c "
            "JOIN project_files f ON f.id=c.file_id",
        ).fetchall()
        assert all(tuple(row) == ("primo", form["id"], "form") for row in chunks)
        assert db.execute("SELECT COUNT(*) FROM document_compilations").fetchone()[0] == 0


@pytest.mark.anyio
async def test_duplicate_names_are_independent_and_removal_preserves_other_forms(client):
    rows = []
    for content in (b"Primo originale", b"Secondo originale"):
        rows.append((await client.post(
            "/api/projects/primo/forms", files={"file": ("domanda.txt", content)},
        )).json())
    assert rows[0]["id"] != rows[1]["id"]
    assert (await client.delete(f"/api/projects/primo/forms/{rows[0]['id']}")).status_code == 204
    assert (await client.get("/api/projects/primo/forms")).json() == [rows[1]]
    assert (await client.get(
        f"/api/projects/primo/forms/{rows[1]['id']}/download",
    )).content == b"Secondo originale"
    assert (await client.get("/api/projects/primo")).json()["model_count"] == 1
    assert len(list(get_storage_path().rglob("*.txt"))) == 1
    with connection() as db:
        for table in ("document_chunks", "document_chunks_fts"):
            assert db.execute(
                f"SELECT COUNT(*) FROM {table} WHERE file_id=?", (rows[0]["id"],),
            ).fetchone()[0] == 0
            assert db.execute(
                f"SELECT COUNT(*) FROM {table} WHERE file_id=?", (rows[1]["id"],),
            ).fetchone()[0] > 0


@pytest.mark.anyio
async def test_project_and_source_endpoints_cannot_access_or_modify_other_roles(client):
    form = (await client.post(
        "/api/projects/primo/forms", files={"file": ("domanda.txt", b"Nome: ___")},
    )).json()
    source = (await client.post(
        "/api/projects/primo/files", files={"file": ("fonte.txt", b"Dato aziendale")},
    )).json()
    for url in (
        f"/api/projects/secondo/forms/{form['id']}",
        f"/api/projects/primo/forms/{source['id']}",
    ):
        assert (await client.get(f"{url}/download")).status_code == 404
        assert (await client.delete(url)).status_code == 404
    assert (await client.delete(f"/api/projects/primo/files/{form['id']}")).status_code == 404
    url = f"/api/projects/primo/files/{form['id']}/content"
    assert (await client.get(url)).status_code == 404
    assert (await client.put(url, json={"content": "Alterato"})).status_code == 404
    assert (await client.get(f"/api/projects/primo/forms/{form['id']}/download")).content == (
        b"Nome: ___"
    )
    assert len(load_project_source_chunks("primo")) == 1


@pytest.mark.anyio
@pytest.mark.parametrize("name, content, code", [
    ("modulo.exe", b"invalid", 415),
    ("modulo.pdf", document_bytes("pdf"), 415),
    ("modulo.md", document_bytes("md"), 415),
    ("vuoto.txt", b"", 422),
    ("spazi.txt", b"  \n", 422),
    ("binario.txt", b"\xff\x00", 422),
    ("falso.docx", b"invalid", 422),
    ("x" * 181 + ".txt", b"Nome: ___", 422),
])
async def test_invalid_uploads_leave_no_records_or_files(client, name, content, code):
    response = await client.post("/api/projects/primo/forms", files={"file": (name, content)})
    assert response.status_code == code, response.text
    assert (await client.get("/api/projects/primo/forms")).json() == []
    assert not list(get_storage_path().rglob("*.*"))


@pytest.mark.anyio
async def test_size_limit_and_missing_project(client, monkeypatch):
    monkeypatch.setattr(project_forms, "MAX_FILE_SIZE", 10)
    response = await client.post(
        "/api/projects/primo/forms", files={"file": ("modulo.txt", b"x" * 11)},
    )
    assert response.status_code == 413
    assert (await client.get("/api/projects/assente/forms")).status_code == 404
    assert (await client.post(
        "/api/projects/assente/forms", files={"file": ("modulo.txt", b"ok")},
    )).status_code == 404
    assert not list(get_storage_path().rglob("*.*"))


@pytest.mark.anyio
async def test_failed_save_rolls_back_file_and_database(client, monkeypatch):
    def fail(*args):
        raise OSError("disk failure")
    monkeypatch.setattr(project_forms, "touch_project", fail)
    result = await client.post(
        "/api/projects/primo/forms", files={"file": ("modulo.txt", b"Nome: ___")},
    )
    assert result.status_code == 500
    assert (await client.get("/api/projects/primo/forms")).json() == []
    assert not list(get_storage_path().rglob("*.txt"))


@pytest.mark.anyio
async def test_project_deleted_during_validation_cannot_leave_orphan_file(client, monkeypatch):
    original = project_forms._validate_file

    def delete_during_validation(data, suffix):
        original(data, suffix)
        delete_project("primo")

    monkeypatch.setattr(project_forms, "_validate_file", delete_during_validation)
    response = await client.post(
        "/api/projects/primo/forms", files={"file": ("modulo.txt", b"Nome: ___")},
    )
    assert response.status_code == 404
    assert not list(get_storage_path().rglob("*.txt"))


@pytest.mark.anyio
async def test_documentary_docx_keeps_tables_and_paragraphs_in_order(client):
    doc = Document()
    doc.add_paragraph("Domanda di partecipazione")
    doc.add_paragraph("Dati del sottoscrittore")
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Forma di partecipazione"
    table.cell(0, 1).text = "Raggruppamento temporaneo"
    doc.add_paragraph("Allegati richiesti: documento di identita")
    out = BytesIO()
    doc.save(out)
    form = (await client.post(
        "/api/projects/primo/forms", files={"file": ("domanda.docx", out.getvalue())},
    )).json()
    with connection() as db:
        text = db.execute(
            "SELECT content FROM document_chunks WHERE file_id=? ORDER BY chunk_index",
            (form["id"],),
        ).fetchone()[0]
    terms = ["sottoscrittore", "Forma di partecipazione", "Raggruppamento", "Allegati"]
    assert [text.index(term) for term in terms] == sorted(text.index(term) for term in terms)
    assert text.count("Raggruppamento") == 1


@pytest.mark.anyio
async def test_explicit_reindex_restores_legacy_form_without_changing_original(client, monkeypatch):
    from scripts.reindex_project_forms import main

    original = document_bytes("txt")
    form = (await client.post(
        "/api/projects/primo/forms", files={"file": ("domanda.txt", original)},
    )).json()
    with connection() as db:
        db.execute("DELETE FROM document_chunks WHERE file_id=?", (form["id"],))
        db.execute(
            "UPDATE project_files SET chunk_count=0, status='Caricato' WHERE id=?", (form["id"],),
        )
    assert search_project_evidence("primo", "Domanda", target="form") == []
    monkeypatch.setattr("sys.argv", [
        "reindex_project_forms", "--project", "primo", "--form-id", str(form["id"]),
    ])
    main()
    found = search_project_evidence("primo", "Domanda", target="form", form_id=form["id"])
    assert found and all(item["role"] == "form" for item in found)
    assert search_project_evidence("secondo", "Domanda", target="form") == []
    assert search_project_evidence("primo", "Domanda") == []
    downloaded = await client.get(f"/api/projects/primo/forms/{form['id']}/download")
    assert downloaded.content == original
    # Repeating the explicit operation replaces chunks, rather than duplicating them.
    main()
    assert (await client.get("/api/projects/primo/forms")).json()[0]["chunk_count"] == len(found)


@pytest.mark.anyio
async def test_failed_reindex_preserves_previous_chunks(client, monkeypatch):
    form = (await client.post(
        "/api/projects/primo/forms", files={"file": ("domanda.txt", document_bytes("txt"))},
    )).json()
    before = search_project_evidence("primo", "Domanda", target="form")
    def fail(*args):
        raise OSError("Errore dopo scrittura dei frammenti")
    monkeypatch.setattr(project_forms, "touch_project", fail)
    with pytest.raises(OSError):
        project_forms.reindex_form("primo", form["id"])
    assert search_project_evidence("primo", "Domanda", target="form") == before


@pytest.mark.anyio
async def test_reindex_command_rejects_wrong_project_or_backend_before_changes(client, monkeypatch):
    from scripts.reindex_project_forms import main

    form = (await client.post(
        "/api/projects/primo/forms", files={"file": ("domanda.txt", document_bytes("txt"))},
    )).json()
    before = search_project_evidence("primo", "Domanda", target="form")
    for options in (
        ["--project", "secondo", "--form-id", str(form["id"])],
        ["--project", "primo", "--sync-vectors"],
    ):
        monkeypatch.setattr("sys.argv", ["reindex_project_forms", *options])
        with pytest.raises(SystemExit) as error:
            main()
        assert error.value.code == 2
        assert search_project_evidence("primo", "Domanda", target="form") == before


@pytest.mark.anyio
async def test_failed_removal_is_retriable_and_project_deletion_cleans_forms(client, monkeypatch):
    form = (await client.post(
        "/api/projects/primo/forms", files={"file": ("modulo.txt", b"Originale")},
    )).json()
    with monkeypatch.context() as patch:
        def fail(*args, **kwargs):
            raise OSError("File occupato")
        patch.setattr(Path, "unlink", fail)
        assert (await client.delete(f"/api/projects/primo/forms/{form['id']}")).status_code == 500
    assert (await client.get("/api/projects/primo/forms")).json() == [form]
    assert (await client.get(f"/api/projects/primo/forms/{form['id']}/download")).content == (
        b"Originale"
    )
    assert (await client.delete("/api/projects/primo")).status_code == 204
    assert not (get_storage_path() / "primo").exists()
    with connection() as db:
        assert db.execute("SELECT COUNT(*) FROM project_files").fetchone()[0] == 0
