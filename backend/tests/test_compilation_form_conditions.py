"""Real FORM positions, distant parents and SOURCE/USER write constraints; no AI."""

from copy import deepcopy
from io import BytesIO

import pytest
from docx import Document
from test_compilation_condition_reuse import branch_field, state
from test_compilation_semantics import classified, source, source_review
from test_compilation_sessions import BASE, create, resolve, simulate
from test_compilation_sessions import api as api
from test_compilation_sessions import source as upload_source

from app import compilation_session_resolution as resolution
from app.compilation_clarifications import pending_slots, synchronize
from app.compilation_conditions import owner
from app.compilation_form_conditions import (
    blocked_previous_writes,
    required_applicability,
    synchronize_form_dependencies,
    write_blockers,
)
from app.compilation_semantics import review_sources
from app.compilation_session_models import CandidateMatches
from app.compilation_sessions import candidate_snapshots
from app.docx_templates import inspect_docx


@pytest.fixture
def anyio_backend():
    return "asyncio"


def form(*, identity=True, distant=True):
    doc = Document()
    if identity:
        doc.add_paragraph("Identificazione dell’operatore economico in relazione "
                          "alla propria ragione sociale")
    conditions = [
        "PROFESSIONISTA SINGOLO",
        "professionisti associati (studio associato) nel caso di rappresentante "
        "munito di idonei poteri",
        "professionisti associati (studio associato) in assenza di rappresentante "
        "munito di idonei poteri",
        "SOCIETÀ DI PROFESSIONISTI",
        "SOCIETÀ DI INGEGNERIA",
        "altro soggetto abilitato in forza del diritto nazionale",
        "CONSORZIO STABILE",
        "CONSORZIATA ESECUTRICE di Consorzio Stabile",
    ]
    for i, condition in enumerate(conditions):
        doc.add_table(rows=1, cols=1).cell(0, 0).text = (
            f"7.{chr(97+i)}) da compilare in caso di {condition}")
        if distant:
            for j in range(16):
                doc.add_table(rows=1, cols=1).cell(0, 0).text = f"Istruzione {j}"
        table = doc.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "Denominazione sociale"
    doc.add_paragraph(">>>>> PARTE SUCCESSIVA <<<<<")
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Nome del rappresentante"
    stream = BytesIO()
    doc.save(stream)
    layout = inspect_docx(stream.getvalue())
    fields = candidate_snapshots(layout, "alpha", 4, "form.docx")
    return layout, fields


def prepared(**kwargs):
    layout, fields = form(**kwargs)
    for field in fields:
        field.update(status="AMBIGUOUS", kind="data", entity="company", condition="",
                     requirement={"name": field["label_hint"]}, search={"status": "searched"})
    return layout, fields


def prove(field, *, provenance="SOURCE", applies=True):
    condition = field["form_dependency"]["condition"]
    field.update(condition=condition, applicability={
        "condition": condition, "applies": applies, "provenance": provenance,
        "evidence": [{"role": "source", "chunk_id": -8,
                      "quote": f"Tipologia di attività: {condition}."}],
        **({"user_quote": "risposta verificata"} if provenance == "USER" else {}),
    })


def core(field):
    return {k: deepcopy(field[k]) for k in (
        "status", "value", "provenance", "source_evidence", "condition", "applicability",
    ) if k in field}


@pytest.mark.parametrize("provenance", ["SOURCE", "USER"])
def test_distant_table_parents_exclude_identity_alternatives_not_participation(provenance):
    _, fields = prepared()
    engineering = fields[4]
    prove(engineering, provenance=provenance)
    s = state(*fields)
    synchronize(s)
    excluded = fields[:4] + fields[5:7]
    assert all(f["status"] == "NOT_APPLICABLE" for f in excluded)
    assert all(required_applicability(f) is False for f in excluded)
    assert required_applicability(engineering) is True
    assert required_applicability(fields[7]) is None  # Engineering may be a member.
    assert "form_dependency" not in fields[8]  # Stop at the real major boundary.
    assert all(f["form_dependency_decision"]["origin_field_id"] == engineering["id"]
               for f in excluded)
    assert all(f["form_dependency_decision"]["basis"]["provenance"] == provenance
               for f in excluded)
    assert not any("STUDIO" in slot["condition"].upper() or "SINGOLO" in slot["condition"]
                   for slot in pending_slots(s))
    before = deepcopy(s)
    synchronize(s)
    assert before == s


def test_no_exclusive_type_inference_without_form_identity_group():
    _, fields = prepared(identity=False)
    prove(fields[4])
    synchronize(state(*fields))
    assert fields[0]["status"] == fields[1]["status"] == "AMBIGUOUS"
    assert required_applicability(fields[7]) is None


def test_preserve_and_report_preexisting_wrong_branch_values_without_ready():
    _, fields = prepared(distant=False)
    prove(fields[4])
    for field in fields:
        field.update(status="RESOLVED", value="Aurora", provenance="SOURCE",
                     source_evidence=[{"role": "source", "quote": "Aurora"}])
    before = [core(f) for f in fields]
    s = state(*fields)
    synchronize(s)
    assert [core(f) for f in fields] == before
    assert len(blocked_previous_writes(s)) == 7  # Six excluded + one UNKNOWN member.
    assert s["status"] != "READY"
    assert not write_blockers(fields[4])
    assert "esclusa" in fields[1]["write_blockers"][0]
    assert "non accertata" in fields[7]["write_blockers"][0]


def test_contradictory_type_proofs_block_writes_even_with_direct_positive_proof():
    _, fields = prepared(distant=False)
    prove(fields[4])
    prove(fields[0], provenance="USER")
    synchronize(state(*fields))
    assert all(required_applicability(f) is None for f in fields[:7])
    assert all(f["status"] == "AMBIGUOUS" for f in fields[:7])


def test_same_type_source_user_contradiction_cannot_be_bypassed_by_own_positive_proof():
    _, fields = prepared(distant=False)
    prove(fields[4])
    denied = deepcopy(fields[4])
    denied["id"] = "separate-proof"
    prove(denied, provenance="USER", applies=False)
    before = deepcopy([fields[4]["applicability"], denied["applicability"]])
    synchronize(state(*fields, denied))
    assert all(required_applicability(f) is None for f in fields[:7] + [denied])
    assert [fields[4]["applicability"], denied["applicability"]] == before


@pytest.mark.parametrize("evidence", [[], [{"role": "form", "quote": "ingegneria"}]])
def test_missing_or_form_source_proof_does_not_select_identity_branch(evidence):
    _, fields = prepared(distant=False)
    prove(fields[4])
    fields[4]["applicability"]["evidence"] = evidence
    synchronize(state(*fields))
    assert all(required_applicability(f) is None for f in fields[:8])
    assert all(f["status"] == "AMBIGUOUS" for f in fields[:8])


@pytest.mark.parametrize("condition", [
    "soggetto residente in Italia",
    "SOCIETÀ DI INGEGNERIA in presenza di soggetto residente in Italia",
])
def test_local_person_predicates_keep_representative_and_director_distinct(condition):
    _, fields = prepared(distant=False)
    director = fields[4]
    representative = deepcopy(director)
    director.update(entity="person", condition=condition,
                    requirement={"person_role": "direttore tecnico"})
    representative.update(id="another", entity="person", condition=director["condition"],
                          requirement={"person_role": "legale rappresentante"})
    assert owner(director) != owner(representative)
    prove(fields[1])
    # Merely sharing a type branch never transfers a person's local predicate.
    assert owner(director)[0] == "local"


def test_user_can_establish_consortium_member_independently_of_engineering_type():
    _, fields = prepared(distant=False)
    prove(fields[4])
    member = fields[7]
    condition = member["form_dependency"]["condition"]
    member["form_dependency_user"] = {"condition": condition, "applies": True,
                                      "provenance": "USER", "user_quote": "si applica"}
    synchronize(state(*fields))
    assert required_applicability(fields[4]) is True
    assert required_applicability(member) is True
    assert fields[6]["status"] == "NOT_APPLICABLE"  # Own type is not the stable consortium.


@pytest.mark.parametrize("quote,opposite_excluded", [
    ("non siamo uno studio associato", True), ("no", False), ("non so", False),
])
def test_literal_user_parent_denial_reaches_both_qualifiers_but_short_no_does_not(
    quote, opposite_excluded,
):
    _, fields = prepared(distant=False)
    prove(fields[1], provenance="USER", applies=False)
    fields[1]["applicability"]["user_quote"] = quote
    before = deepcopy(fields[1]["applicability"])
    synchronize(state(*fields))
    assert (required_applicability(fields[2]) is False) == opposite_excluded
    assert fields[1]["applicability"] == before
    assert required_applicability(fields[7]) is None


def test_local_true_condition_cannot_bypass_unknown_parent():
    _, fields = prepared(distant=False)
    field = fields[7]
    field.update(condition="soggetto residente in Italia", applicability={
        "condition": "soggetto residente in Italia", "applies": True,
        "provenance": "USER", "user_quote": "si applica"})
    synchronize(state(*fields))
    assert required_applicability(field) is None


@pytest.mark.anyio
@pytest.mark.parametrize("parent", ["excluded", "unknown", "verified"])
async def test_accepted_semantic_source_value_still_requires_verified_parent(parent):
    layout, fields = form(distant=False)
    target = fields[1]
    await classified([target], ["Denominazione sociale"],
                     subject_anchor="Denominazione sociale", subject_relation="organization")
    assert target["semantic_validation"]["accepted"]
    if parent == "excluded":
        prove(fields[4])
    elif parent == "verified":
        prove(target)
        # The source test requires the interpretation digest to remain unchanged.
        target["form_dependency_user"] = {**target.pop("applicability"),
                                          "provenance": "USER", "user_quote": "si applica"}
        target["condition"] = ""
    synchronize_form_dependencies(state(*fields))
    evidence = source("Denominazione sociale: Aurora.")
    matches = CandidateMatches(fields=[{
        "candidate_id": target["id"], "supports": [{"source_id": 1,
            "quote": evidence["content"], "value": "Aurora"}], "reason": "Valore nella SOURCE",
    }])
    coverage = {target["id"]: [1]}
    await review_sources([target], matches, [evidence], coverage, source_review)
    resolution.apply_matches(layout, [target], matches, [evidence], coverage)
    if parent == "verified":
        assert target["status"] == "RESOLVED" and target["value"] == "Aurora"
    else:
        assert target["value"] is None and target["status"] != "RESOLVED"
        assert any("Sezione/condizione" in e or "Applicabilità" in e
                   for e in target["validation_errors"])


def test_blank_unlabelled_row_is_not_a_second_company_property():
    field = branch_field("generic", "")
    field.update(location={"kind": "table_cell"}, label_hint="",
                 structural={"row": [{"text": ""}, {"text": ""}]})
    assert "senza etichette" in write_blockers(field)[0]
    field["label_hint"] = "Sede legale"  # A real column header can establish the property.
    assert write_blockers(field) == []


@pytest.mark.anyio
async def test_partial_export_reports_blocked_old_write_and_checks_xml(api, monkeypatch):
    import json

    from test_compilation_conversation import mutate_state

    layout, fields = form(distant=False)
    session = await create(api, layout.original)
    engineering_id, studio_id = fields[4]["id"], fields[1]["id"]
    condition = fields[4]["form_dependency"]["condition"]

    def verified_user(saved):
        field = next(f for f in saved["fields"] if f["id"] == engineering_id)
        field["form_dependency_user"] = {
            "condition": condition, "applies": True, "provenance": "USER",
            "user_quote": "Siamo una società di ingegneria",
        }

    mutate_state(session, verified_user)
    await upload_source(api, "Denominazione sociale: Aurora.")
    simulate(monkeypatch, {engineering_id: ["Aurora"]})
    session = await resolve(api, session, [engineering_id])
    engineering = next(f for f in session["fields"] if f["id"] == engineering_id)
    assert engineering["status"] == "RESOLVED", engineering["validation_errors"]

    def historical_wrong_write(saved):
        field = next(f for f in saved["fields"] if f["id"] == studio_id)
        field.update(status="RESOLVED", value=engineering["value"], provenance="SOURCE",
                     source_evidence=deepcopy(engineering["source_evidence"]))

    mutate_state(session, historical_wrong_write)
    session = (await api.get(f"{BASE}/{session['id']}")).json()
    assert session["summary"]["resolved"] == 2
    assert session["summary"]["writable"] == session["summary"]["blocked_writes"] == 1
    guarded = await api.post(f"{BASE}/{session['id']}/finalize",
                             json={"version": session["version"]})
    assert guarded.status_code == 409
    response = await api.post(f"{BASE}/{session['id']}/finalize", json={
        "version": session["version"], "allow_unresolved": True,
    })
    assert response.status_code == 200, response.text
    after = response.json()
    assert after["summary"]["resolved"] == 2
    assert next(f for f in after["fields"] if f["id"] == studio_id)["value"] == "Aurora"
    downloads = after["last_generation"]["downloads"]
    report = json.loads((await api.get(downloads["report"])).content)
    assert report["written_field_count"] == 1
    assert report["blocked_previous_writes"][0]["field_id"] == studio_id
    assert report["blocked_previous_writes"][0]["source_evidence"] == engineering["source_evidence"]
    assert report["status"] == "needs_review" and not report["ready_for_submission"]
    doc = Document(BytesIO((await api.get(downloads["docx"])).content))
    assert doc.tables[fields[4]["location"]["table"]].cell(0, 1).text == "Aurora"
    assert doc.tables[fields[1]["location"]["table"]].cell(0, 1).text == ""


@pytest.mark.parametrize("provenance", ["SOURCE", "USER"])
def test_verified_condition_is_bound_before_form_review_even_when_model_omits_it(provenance):
    from app.compilation_session_models import CandidateMeanings

    _, fields = prepared(distant=False)
    field = fields[4]
    prove(field, provenance=provenance)
    before = deepcopy(field["applicability"])
    meanings = CandidateMeanings(fields=[{
        "candidate_id": field["id"], "classification": "data", "entity": "company",
        "requirement": {"name": "Denominazione sociale", "form_quote": field["context"],
                        "form_citation_id": 1},
        "form_quote": field["context"], "condition": "", "reason": "Classificatore incompleto",
    }])
    resolution.bind_meanings([field], meanings)
    assert meanings.fields[0].condition == field["condition"]
    resolution.apply_meanings([field], meanings)
    assert field["applicability"] == before
    assert field["condition"] == before["condition"]


@pytest.mark.parametrize("key,value", [("project_id", "beta"), ("file_id", 44),
                                        ("template_sha256", "other-original")])
def test_dependency_proofs_do_not_cross_project_file_or_template(key, value):
    _, fields = prepared(distant=False)
    prove(fields[4])
    fields[1]["form_evidence"][key] = value
    synchronize(state(*fields))
    assert required_applicability(fields[1]) is None
    assert fields[1]["status"] == "AMBIGUOUS"
