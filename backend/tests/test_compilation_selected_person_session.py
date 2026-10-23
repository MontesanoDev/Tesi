"""Actual CompilationSession persistence/export; only provider transports simulated."""

import json
from io import BytesIO

import pytest
from docx import Document
from test_compilation_global_sessions import BASE, api, provider, resolve
from test_compilation_selected_person import NAME, case
from test_compilation_sessions import source

from app import compilation_global_dialogue as dialogue
from app import compilation_global_sessions as adapter
from app.compilation_chat import ClarificationReply
from app.db import get_storage_path

__all__ = ["api", "provider"]


@pytest.fixture
def anyio_backend():
    return "asyncio"


@pytest.mark.anyio
@pytest.mark.parametrize("with_clarification", [True, False])
async def test_signatory_reply_source_reuse_review_export_and_fresh_source_guard(api, provider,
    monkeypatch, with_clarification,
):
    _, layout, sources = case()
    uploaded = await api.post("/api/projects/alpha/forms", files={
        "file": ("modulo.docx", layout.original)})
    await source(api, sources[0]["content"], name="profili.md")
    calls = []

    async def proposal(client, settings, body, timeout, **kwargs):
        calls.append(body)
        model_input = json.loads(body["messages"][1]["content"])
        chunk = next(s["chunk_id"] for s in model_input["sources"]
                     if "- **Società:** Aurora Servizi S.r.l." in s["content"])
        return json.dumps({"version": 12, "facts": {"company": {
            "subject": "Aurora Servizi S.r.l.", "entity": "company", "aspect": "identity",
            "property": "denominazione", "value": "Aurora Servizi S.r.l.",
            "evidence": [{"chunk_id": chunk, "quote": "- **Società:** Aurora Servizi S.r.l."}]}},
            "fields": {"p2.s1": {"label": "Operatore economico", "kind": "data",
                "disposition": "value", "conditions": [], "property": "denominazione",
                "fact_id": "company", "form": {"paragraph_id": "paragraph:2",
                    "quote": "operatore economico [[p2.s1]]"}}}, "bindings": {},
            "conditions": {}, "choice_groups": {}, "coverage": [], "clarifications": {
                "p0.s0": {"label": "Il sottoscritto", "reason": "Scelta della pratica",
                          "form": {"paragraph_id": "paragraph:0",
                                   "quote": "Il sottoscritto [[p0.s0]]"}}}
                if with_clarification else {}}), "simulated", {}

    monkeypatch.setattr(adapter, "request_model_content", proposal)
    s = (await api.post(BASE, json={"form_id": uploaded.json()["id"], "engine": "global",
                                  "start_in_chat": True})).json()
    s = await resolve(api, s)
    assert len(s["chat"]["question"]["slots"]) == 1
    assert s["chat"]["question"]["slots"][0]["field_ids"] == ["p0.s0"]
    raw_path = next(p for p in get_storage_path().rglob("response.txt")
                    if "dialogue" not in p.parts)
    before = raw_path.read_bytes()
    s, errors = dialogue.apply_replies("alpha", s["id"], s["version"], [
        ClarificationReply(slot=1, action="VALUE", value=NAME, user_quote=NAME)], NAME)
    assert not errors and s["global_compilation"]["proposal_count"] == 9
    effective_report = (await api.get(f"{BASE}/{s['id']}/global-report")).json()
    assert effective_report["source_reuse_count"] == 8
    original_audit = json.loads(raw_path.with_name("audit.json").read_text())
    assert original_audit.get("source_reuse_count") is None
    assert len(s["global_compilation"]["dialogue"]["user_values"]) == 1
    s = await resolve(api, s)
    ids, conditions, users = dialogue.export_selection(s)
    assert len(ids) == 9 and users == ["p0.s0"] and not conditions
    request = {"version": s["version"], "candidate_ids": ids, "user_candidate_ids": users,
               "confirmed_condition_ids": conditions, "allow_unresolved": True,
               "user_quote": "Genera la bozza parziale"}
    generated = await api.post(f"{BASE}/{s['id']}/global-draft", json=request)
    assert generated.status_code == 200, generated.text
    run = generated.json()["last_generation"]
    report = (await api.get(run["downloads"]["report"])).json()
    assert report["written_field_count"] == 10 and report["resolved"] == 0
    assert len([f for f in report["fields"] if f["provenance"] == "SOURCE"]) == 9
    assert len([f for f in report["fields"] if f["provenance"] == "USER"]) == 1
    assert report["checkbox_values"] == {} and not report["ready_for_submission"]
    draft = Document(BytesIO((await api.get(run["downloads"]["docx"])).content))
    assert "Torino" in "\n".join(p.text for p in draft.paragraphs)
    assert "NREADA85A52L219R" in "\n".join(p.text for p in draft.paragraphs)
    assert len(calls) == 1 and raw_path.read_bytes() == before
    # Changed evidence cannot reuse the stale review, even if the name is unchanged.
    from app.db import connection

    with connection() as db:
        db.execute("UPDATE global_document_chunks SET content="
                   "replace(content, '12 gennaio', '13 gennaio')")
    request["version"] = generated.json()["version"]
    rejected = await api.post(f"{BASE}/{s['id']}/global-draft", json=request)
    assert rejected.status_code == 409
