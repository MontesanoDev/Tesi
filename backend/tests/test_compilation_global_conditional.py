"""Conditional proposals retain data without approving applicability or writes.

These are deterministic software regressions, not measures of real AI quality.
"""

import json
from io import BytesIO

import pytest
from docx import Document

from app.compilation_global import build_global_input, validate_global_proposal
from app.docx_templates import inspect_docx


@pytest.fixture
def case():
    doc = Document()
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Codice organizzazione (se incaricata da terzi)"
    table.cell(1, 0).text = "Recapito principale"
    stream = BytesIO()
    doc.save(stream)
    layout = inspect_docx(stream.getvalue())
    source = {"chunk_id": -7, "file_id": -3, "chunk_index": 0,
              "source_name": "profilo.txt", "role": "source", "scope": "global",
              "category": "company", "project_id": None,
              "content": "Alba: codice AZ-98765; recapito Via Luna 4."}
    facts = {key: {"subject": "Alba", "entity": "company", "aspect": "identity",
                   "property": prop, "value": value,
                   "evidence": [{"chunk_id": -7, "quote": source["content"]}]}
             for key, prop, value in [("code", "codice", "AZ-98765"),
                                      ("address", "recapito", "Via Luna 4")]}
    fields = {
        "t0.r0.c1": {"label": "Codice organizzazione", "kind": "data",
                     "disposition": "conditional", "property": "codice", "fact_id": "code",
                     "conditions": ["branch"], "reason": "Incarico non documentato",
                     "form": {"paragraph_id": "paragraph:0", "quote": "Codice organizzazione"}},
        "t0.r1.c1": {"label": "Recapito principale", "kind": "data",
                     "disposition": "value", "property": "recapito", "fact_id": "address",
                     "conditions": [],
                     "form": {"paragraph_id": "paragraph:2", "quote": "Recapito principale"}},
    }
    condition = {"form": {"paragraph_id": "paragraph:0", "quote": "se incaricata da terzi"},
                 "aspect": "representation", "operator": "semantic", "fact_ids": [],
                 "reason": "Incarico non documentato",
                 "scope": [{"kind": "candidate", "id": "t0.r0.c1"}]}
    return layout, [source], {"version": 8, "facts": facts, "fields": fields,
                              "conditions": {"branch": condition}, "bindings": {}}


def audit(case):
    return validate_global_proposal(json.dumps(case[2]), case[0], "alpha", case[1])


def test_grounded_source_survives_unknown_applicability_but_never_selects_a_value(case):
    result = audit(case)
    field = result["fields"][0]
    assert field["validation"] == "REVIEW_REQUIRED"
    assert field["disposition"] == "conditional"
    assert field["proposal_grounding"] == "GROUNDED_PROPOSAL"
    assert field["proposed_value"] == "AZ-98765"
    assert field["proposed_fact_id"] == "code"
    assert field["value"] is field["fact_id"] is field["written_value"] is None
    assert field["condition_checks"]["branch"]["result"] == "unknown"
    assert field["semantic_review_required"]
    assert result["grounded_conditional_proposals"] == 1
    proof = field["source_evidence"][0]
    assert (proof["role"], proof["scope"], proof["category"]) == ("source", "global", "company")
    assert (proof["file_id"], proof["chunk_id"], proof["chunk_index"]) == (-3, -7, 0)
    assert proof["content_sha256"]
    assert result["fields"][1]["validation"] == "GROUNDED_PROPOSAL"
    assert result["resolved"] == result["writes"] == 0


def test_value_with_unknown_applicability_still_rejected_not_relabelled_by_backend(case):
    case[2]["fields"]["t0.r0.c1"]["disposition"] = "value"
    field = audit(case)["fields"][0]
    assert field["validation"] == "REJECTED"
    assert "sconosciuta" in field["validation_errors"][0]
    assert "proposed_value" not in field


@pytest.mark.parametrize("change", [
    {"conditions": []}, {"conditions": ["unknown-id"]}, {"conditions": ["branch", "branch"]},
    {"fact_id": "unknown-id"}, {"property": "diversa"}, {"reason": ""},
    {"form": {"paragraph_id": "paragraph:2", "quote": "Recapito principale"}},
    {"form": {"paragraph_id": "paragraph:0", "quote": "Inventato"}},
    {"kind": "signature"}, {"choice": True},
])
def test_conditional_record_cannot_bypass_local_schema_source_or_position_gates(case, change):
    case[2]["fields"]["t0.r0.c1"].update(change)
    result = audit(case)
    assert result["fields"][0]["validation"] == "REJECTED"
    assert "proposed_value" not in result["fields"][0]
    assert result["fields"][1]["validation"] == "GROUNDED_PROPOSAL"


@pytest.mark.parametrize("change", [
    {"value": "AZ-00000"}, {"value": "non disponibile"}, {"evidence": []},
    {"evidence": [{"chunk_id": -99, "quote": "AZ-98765"}]},
])
def test_conditional_source_must_actually_ground_a_real_value(case, change):
    case[2]["facts"]["code"].update(change)
    field = audit(case)["fields"][0]
    assert field["validation"] == "REJECTED"
    assert "proposed_value" not in field


@pytest.mark.parametrize("origin", ["form", "user", "foreign", "general"])
def test_conditional_proposal_does_not_relax_source_provenance_or_isolation(case, origin):
    source = case[1][0]
    if origin == "general":
        source["category"] = "general"
        assert audit(case)["fields"][0]["validation"] == "REJECTED"
    else:
        if origin == "foreign":
            source.update(chunk_id=7, file_id=3, project_id="beta", scope="project:beta",
                          category=None)
        else:
            source["role"] = origin
        with pytest.raises(ValueError):
            audit(case)


@pytest.mark.parametrize("operator", ["same", "different"])
def test_known_true_or_false_condition_cannot_be_disguised_as_unknown(case, operator):
    case[2]["facts"]["second-code"] = {**case[2]["facts"]["code"], "property": "altro codice"}
    condition = case[2]["conditions"]["branch"]
    condition.update(operator=operator, aspect="identity", fact_ids=["code", "second-code"])
    field = audit(case)["fields"][0]
    assert field["validation"] == "REJECTED"
    assert "proposed_value" not in field


def test_uncertain_or_invalid_condition_is_not_a_usable_conditional_proposal(case):
    case[2]["conditions"]["branch"]["form"]["quote"] = "Inventata"
    assert audit(case)["fields"][0]["validation"] == "REJECTED"
    case[2]["conditions"]["branch"]["form"]["quote"] = "se incaricata da terzi"
    case[2]["uncertain_paragraph_ids"] = ["paragraph:0"]
    assert audit(case)["fields"][0]["validation"] == "REJECTED"


def test_known_alternatives_cannot_be_hidden_by_a_conditional_choice(case):
    other = {**case[2]["facts"]["code"], "value": "AZ-12345"}
    case[1][0]["content"] += " Altro codice: AZ-12345."
    other["evidence"] = [{"chunk_id": -7, "quote": "Altro codice: AZ-12345."}]
    case[2]["facts"]["other"] = other
    case[2]["conditions"]["branch"].update(aspect="identity", fact_ids=["other"])
    field = audit(case)["fields"][0]
    assert field["validation"] == "REVIEW_REQUIRED" and field["disposition"] == "review"
    assert {a["value"] for a in field["alternatives"]} == {"AZ-98765", "AZ-12345"}
    assert "proposed_value" not in field


def test_empty_is_reported_explicitly_without_forcing_values_or_conditions(case):
    case[2].update(fields={}, facts={}, conditions={})
    result = audit(case)
    assert result["assessment"] == "empty_proposal"
    assert result["grounded_conditional_proposals"] == 0
    assert result["unassigned_count"] == 2
    assert result["resolved"] == result["writes"] == 0


@pytest.mark.parametrize("support", ["absent", "invalid_value", "wrong_dimension"])
def test_bad_optional_source_support_does_not_erase_real_unknown_form_condition(case, support):
    condition = case[2]["conditions"]["branch"]
    condition["fact_ids"] = ["support"]
    if support != "absent":
        case[2]["facts"]["support"] = {**case[2]["facts"]["code"]}
        if support == "invalid_value":
            case[2]["facts"]["support"]["value"] = "Inventato"
        else:
            case[2]["facts"]["support"]["aspect"] = "legal_type"
    result = audit(case)
    field, checked = result["fields"][0], result["conditions"]["branch"]
    assert checked["result"] == "unknown"
    assert checked["grounding_basis"] == "form_only_unknown"
    assert checked["source_support_validation"] == "REJECTED"
    assert checked["source_support_validation_errors"]
    assert field["proposed_value"] == "AZ-98765"
    assert field["proposed_fact_id"] == "code"
    assert field["value"] is field["written_value"] is None
    assert result["writes"] == result["resolved"] == 0


def test_unknown_condition_support_does_not_repair_a_bad_field_source(case):
    case[2]["facts"]["code"]["value"] = "Inventato"
    case[2]["conditions"]["branch"]["fact_ids"] = ["code"]
    field = audit(case)["fields"][0]
    assert field["validation"] == "REJECTED"
    assert "proposed_value" not in field


@pytest.mark.parametrize("operator", ["same", "different"])
def test_invalid_source_for_a_real_comparison_still_rejects_condition(case, operator):
    case[2]["conditions"]["branch"].update(operator=operator, fact_ids=["absent", "code"])
    result = audit(case)
    assert result["conditions"]["branch"]["validation"] == "REJECTED"
    assert result["fields"][0]["validation"] == "REJECTED"
    assert "proposed_value" not in result["fields"][0]


def test_form_condition_itself_must_be_grounded_despite_missing_source_support(case):
    case[2]["conditions"]["branch"].update(
        form={"paragraph_id": "paragraph:0", "quote": "Clausola inventata"},
        fact_ids=["absent"],
    )
    result = audit(case)
    assert result["conditions"]["branch"]["validation"] == "REJECTED"
    assert "proposed_value" not in result["fields"][0]


def test_unknown_value_is_still_rejected_when_optional_condition_support_is_bad(case):
    case[2]["fields"]["t0.r0.c1"]["disposition"] = "value"
    case[2]["conditions"]["branch"]["fact_ids"] = ["absent"]
    field = audit(case)["fields"][0]
    assert field["validation"] == "REJECTED"
    assert "proposed_value" not in field


def test_schema_bounded_and_historical_v7_does_not_accept_new_disposition(case):
    schema = build_global_input(case[0], "alpha", case[1])["schema_output"]
    assert schema["properties"]["version"]["const"] == 9
    assert schema["properties"]["fields"]["maxProperties"] == 400
    assert schema["$defs"]["ConditionalValue"]["properties"]["conditions"]["minItems"] == 1
    case[2]["version"] = 7
    result = audit(case)
    assert result["fields"][0]["validation"] == "REJECTED"
    assert result["fields"][1]["validation"] == "GROUNDED_PROPOSAL"
    assert "assessment" not in result
