"""Review survives insertions; changed facts/evidence require review again."""
import copy
import json
from io import BytesIO

import pytest
from docx import Document
from test_compilation_global_sessions import BASE, api, provider, resolve
from test_compilation_sessions import source

from app import compilation_global_dialogue as dialogue
from app import compilation_global_sessions as adapter
from app.config import AISettings
from app.db import get_storage_path
from app.generation import GenerationError

__all__ = ["api", "provider"]


@pytest.fixture
def anyio_backend():
    return "asyncio"


@pytest.fixture
def narrator(monkeypatch):
    calls = []
    monkeypatch.setattr(dialogue, "get_ai_settings", lambda: AISettings(
        api_key="test-secret", model="simulated", base_url="https://provider.invalid"))

    async def request(client, settings, body, timeout, **kwargs):
        context = json.loads(body["messages"][1]["content"])
        calls.append(context)
        return json.dumps({**({"message": "Ho conservato le proposte verificabili."}
                             if context["conversation_stage"] == "initial" else {}),
                           "questions": [], "held_proposal_slots": []}), "simulated", {}

    monkeypatch.setattr(dialogue, "request_model_content", request)
    return calls


def data(count=2, start=0):
    fields, facts, candidates, blocks = [], {}, {}, []
    for index in range(start, start + count):
        key, fact, paragraph = f"p{index}.s0", f"f{index}", f"paragraph:{index}"
        proof = {"chunk_id": 1, "file_id": 1, "source_name": "azienda.txt",
                 "role": "source", "category": "company", "scope": "project",
                 "project_id": "alpha", "quote": f"Dato documentato {index}: valore {index}",
                 "content_sha256": "original"}
        facts[fact] = {"subject": "Aurora", "entity": "company", "property": "dato",
                       "value": f"valore {index}", "source_evidence": [copy.deepcopy(proof)]}
        fields.append({"candidate_id": key, "fact_id": fact, "label": f"Dato {index}",
            "kind": "data", "disposition": "value", "validation": "GROUNDED_PROPOSAL",
            "value": f"valore {index}", "conditions": [], "validation_errors": [],
            "form_evidence": {"paragraph_id": paragraph, "quote": f"Dato {index}: ______"},
            "proposal": {"fact_id": fact}, "source_evidence": [proof]})
        candidates[key] = {"id": key, "anchor_paragraph_ids": [paragraph]}
        blocks.append({"id": paragraph, "paragraphs": [
            {"id": paragraph, "text": f"Dato {index}: ______"}]})
    return {"audit": {"version": 12, "fields": fields, "facts": facts,
                      "conditions": {}, "bindings": {}},
            "catalog": {"template_sha256": "same-original", "candidates": candidates,
                        "blocks": blocks}, "writable_ids": list(candidates), "dialogue": {}}


def selection(current):
    return dialogue.export_selection({"global_compilation": {
        **current, "fields": adapter.review_fields(current)}})[0]


@pytest.mark.anyio
async def test_new_properties_do_not_erase_forty_previous_reviews(narrator, tmp_path):
    current = data(40)
    current["dialogue"] = await dialogue.narrate(current, tmp_path)
    reviewed = current["dialogue"]["reviewed_candidate_ids"][:]
    extra = data(8, 100)
    current["audit"]["fields"] = extra["audit"]["fields"] + current["audit"]["fields"]
    current["audit"]["facts"].update(extra["audit"]["facts"])
    current["catalog"]["candidates"].update(extra["catalog"]["candidates"])
    current["catalog"]["blocks"] += extra["catalog"]["blocks"]
    current["writable_ids"] += extra["writable_ids"]
    current["dialogue"].update(user_turns=1, latest_turn={"accepted": [{"action": "VALUE"}]})
    current["dialogue"] = await dialogue.narrate(current, tmp_path)
    assert set(reviewed) <= set(current["dialogue"]["reviewed_candidate_ids"])
    assert len(current["dialogue"]["reviewed_candidate_ids"]) == 48
    assert len(narrator[-1]["document_values"]) == 8 and len(selection(current)) == 48


@pytest.mark.anyio
async def test_windows_cover_remaining_proposals_before_export(narrator, tmp_path):
    current = data(45)
    current["dialogue"] = await dialogue.narrate(current, tmp_path)
    assert len(current["dialogue"]["reviewed_candidate_ids"]) == 40
    assert current["dialogue"]["phase"] == "PENDING"
    assert current["dialogue"]["question"] is None and not selection(current)
    current["dialogue"] = await dialogue.narrate(current, tmp_path)
    assert current["dialogue"]["phase"] == "READY" and len(selection(current)) == 45
    assert [len(c["document_values"]) for c in narrator] == [40, 5]
    assert set(v["value"] for v in narrator[0]["document_values"]).isdisjoint(
        v["value"] for v in narrator[1]["document_values"])


@pytest.mark.anyio
@pytest.mark.parametrize("changed", ["value", "subject", "source", "role", "form", "condition"])
async def test_changed_proposal_cannot_reuse_review(narrator, tmp_path, changed):
    current = data()
    current["audit"]["conditions"]["c"] = {"result": "true", "reason": "Condizione iniziale"}
    current["audit"]["fields"][0]["conditions"] = ["c"]
    current["dialogue"] = await dialogue.narrate(current, tmp_path)
    field = current["audit"]["fields"][0]
    if changed == "value":
        field["value"] = "altro valore"
    elif changed == "subject":
        current["audit"]["facts"]["f0"]["subject"] = "Altra società"
    elif changed == "source":
        field["source_evidence"][0]["content_sha256"] = "updated-source"
    elif changed == "role":
        field["subject_binding"] = {"binding_id": "b", "result": "review_required"}
    elif changed == "form":
        field["form_evidence"]["quote"] = "Diversa richiesta: ______"
    else:
        current["audit"]["conditions"]["c"]["reason"] = "Condizione differente"
    assert "p0.s0" not in selection(current)
    current["dialogue"] = await dialogue.narrate(current, tmp_path)
    assert len(narrator[-1]["document_values"]) == 1
    assert "p1.s0" in current["dialogue"]["reviewed_candidate_ids"]


@pytest.mark.anyio
async def test_unrelated_answer_keeps_source_review(narrator, tmp_path):
    current = data()
    current["dialogue"] = await dialogue.narrate(current, tmp_path)
    current["dialogue"]["answers"] = {"condition:other": {
        "provenance": "USER", "action": "CONDITION_TRUE", "applies": True}}
    current["dialogue"] = await dialogue.narrate(current, tmp_path)
    assert not narrator[-1]["document_values"] and len(selection(current)) == 2


@pytest.mark.anyio
async def test_review_never_confirms_unknown_or_releases_holds(narrator, tmp_path, monkeypatch):
    current = data()
    current["audit"]["conditions"]["c"] = {"result": "unknown"}
    current["audit"]["fields"][0].update(conditions=["c"], condition_checks={
        "c": {"result": "unknown"}})

    async def hold(*args, **kwargs):
        return json.dumps({"message": "Una destinazione richiede ancora verifica.",
                           "questions": [], "held_proposal_slots": [2]}), "simulated", {}
    monkeypatch.setattr(dialogue, "request_model_content", hold)
    current["dialogue"] = await dialogue.narrate(current, tmp_path)
    assert current["dialogue"]["held_candidate_ids"] == ["p1.s0"] and not selection(current)
    current["audit"]["fields"][1]["value"] = "valore aggiornato"

    async def no_hold(*args, **kwargs):
        return json.dumps({"message": "Le proposte restano da confermare.",
                           "questions": [], "held_proposal_slots": []}), "simulated", {}
    monkeypatch.setattr(dialogue, "request_model_content", no_hold)
    current["dialogue"] = await dialogue.narrate(current, tmp_path)
    assert current["dialogue"]["held_candidate_ids"] == ["p1.s0"] and not selection(current)
    current["dialogue"]["answers"] = {"condition:c": {
        "provenance": "USER", "action": "CONDITION_TRUE", "applies": True}}
    assert selection(current) == ["p0.s0"]
    current["dialogue"]["answers"]["condition:c"]["applies"] = False
    assert not selection(current)


@pytest.mark.anyio
async def test_failed_next_window_keeps_receipts_but_blocks_export(narrator, tmp_path, monkeypatch):
    current = data(45)
    current["dialogue"] = await dialogue.narrate(current, tmp_path)
    before = copy.deepcopy(current["dialogue"])

    async def fail(*args, **kwargs):
        raise GenerationError("Simulated failure")
    monkeypatch.setattr(dialogue, "request_model_content", fail)
    current["dialogue"] = await dialogue.narrate(current, tmp_path)
    assert current["dialogue"]["reviewed_candidate_ids"] == before["reviewed_candidate_ids"]
    assert current["dialogue"]["error"] == "GenerationError" and not selection(current)


@pytest.mark.anyio
async def test_forty_five_fields_reach_persisted_docx(api, provider, monkeypatch):
    doc = Document()
    table = doc.add_table(rows=45, cols=2)
    for row in table.rows:
        row.cells[0].text, row.cells[1].text = "Denominazione sociale", "________"
    stream = BytesIO()
    doc.save(stream)
    uploaded = await api.post("/api/projects/alpha/forms", files={
        "file": ("modulo.docx", stream.getvalue())})
    await source(api, "Denominazione sociale: Aurora Servizi S.r.l.")
    calls = []

    async def proposal(client, settings, body, timeout, **kwargs):
        model_input = json.loads(body["messages"][1]["content"])
        calls.append(model_input)
        chunk = next(s["chunk_id"] for s in model_input["sources"] if "Aurora" in s["content"])
        paragraphs = {p["id"]: p["text"] for b in model_input["form"]["blocks"]
                      for p in b["paragraphs"]}
        fields = {}
        for index in range(45):
            key = f"t0.r{index}.c1"
            anchor = next(p for p in model_input["form"]["candidates"][key]["anchor_paragraph_ids"]
                          if paragraphs[p] == "Denominazione sociale")
            fields[key] = {"label": "Denominazione sociale", "kind": "data", "disposition": "value",
                "form": {"paragraph_id": anchor, "quote": "Denominazione sociale"},
                "conditions": [], "property": "denominazione", "fact_id": "company"}
        return json.dumps({"version": 12, "facts": {"company": {
            "subject": "Aurora Servizi S.r.l.", "entity": "company", "aspect": "identity",
            "property": "denominazione", "value": "Aurora Servizi S.r.l.",
            "evidence": [{"chunk_id": chunk,
                          "quote": "Denominazione sociale: Aurora Servizi S.r.l."}]}},
            "fields": fields, "conditions": {}, "bindings": {}, "choice_groups": {},
            "coverage": [], "clarifications": {}}), "simulated", {}
    monkeypatch.setattr(adapter, "request_model_content", proposal)
    current = (await api.post(BASE, json={"form_id": uploaded.json()["id"], "engine": "global",
                                         "start_in_chat": True})).json()
    current = await resolve(api, current)
    assert current["chat"]["auto_continue"] and current["chat"]["question"] is None
    current = (await api.get(f"{BASE}/{current['id']}")).json()
    current = await resolve(api, current)
    assert not current["chat"]["auto_continue"]
    ids, conditions, users = dialogue.export_selection(current)
    assert len(ids) == 45 and not conditions and not users
    generated = await api.post(f"{BASE}/{current['id']}/global-draft", json={
        "version": current["version"], "candidate_ids": ids, "allow_unresolved": True})
    assert generated.status_code == 200, generated.text
    exports = generated.json()["last_generation"]["downloads"]
    report = (await api.get(exports["report"])).json()
    written = Document(BytesIO((await api.get(exports["docx"])).content))
    assert report["written_field_count"] == 45 and report["resolved"] == 0
    assert all(row.cells[1].text == "Aurora Servizi S.r.l." for row in written.tables[0].rows)
    assert len(calls) == 1 and len(list(get_storage_path().rglob('dialogue/*/request.json'))) == 2
    # A subject change is invisible in the compact public field list. The full
    # audit must still invalidate its receipts before chat can select an export.
    from app import compilation_sessions as sessions
    from app.db import connection

    with connection() as db:
        row = sessions._row(db, "alpha", current["id"])
        state = json.loads(row["state_json"])
        state["global_compilation"]["audit"]["facts"]["company"]["subject"] = "Altro operatore"
        sessions.save_state(db, row, state, "test_changed_subject")
    refreshed = (await api.get(f"{BASE}/{current['id']}")).json()
    assert not refreshed["global_compilation"]["dialogue"]["reviewed_candidate_ids"]
    assert not dialogue.export_selection(refreshed)[0]


@pytest.mark.anyio
async def test_explicit_refresh_of_old_v12_reviews_them_without_implicit_upgrade(
    narrator, tmp_path,
):
    current = data()
    current["dialogue"] = {"phase": "READY", "reviewed_candidate_ids": ["p0.s0"],
                           "held_candidate_ids": []}
    assert selection(current) == ["p0.s0"]
    assert "reviewed_candidate_fingerprints" not in current["dialogue"]
    current["dialogue"] = await dialogue.narrate(current, tmp_path)
    assert len(narrator[-1]["document_values"]) == 2 and len(selection(current)) == 2
    assert current["dialogue"]["reviewed_candidate_fingerprints"]


@pytest.mark.anyio
async def test_prior_contract_is_not_upgraded(narrator, tmp_path):
    current = data()
    current["audit"]["version"] = 11
    current["dialogue"] = await dialogue.narrate(current, tmp_path)
    assert "reviewed_candidate_fingerprints" not in current["dialogue"]
    assert "review_pending" not in current["dialogue"]
    assert len(selection(current)) == 2


@pytest.mark.anyio
async def test_hold_index_cannot_refer_to_a_previous_window(narrator, tmp_path, monkeypatch):
    current = data(41)
    current["dialogue"] = await dialogue.narrate(current, tmp_path)

    async def bad_hold(*args, **kwargs):
        return json.dumps({"questions": [], "held_proposal_slots": [2]}), "simulated", {}

    monkeypatch.setattr(dialogue, "request_model_content", bad_hold)
    current["dialogue"] = await dialogue.narrate(current, tmp_path)
    assert current["dialogue"]["error"] == "ValueError"
    assert len(current["dialogue"]["reviewed_candidate_ids"]) == 40
    assert not current["dialogue"]["held_candidate_ids"] and not selection(current)
