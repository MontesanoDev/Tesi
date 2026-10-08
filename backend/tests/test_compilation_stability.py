"""LLM boundary failures preserve independent evidence, USER decisions and bounded work."""

import json
from copy import deepcopy
from io import BytesIO

import pytest
from docx import Document
from fastapi import HTTPException
from test_compilation_clarifications import answer, checkpoint, reply
from test_compilation_controls import current
from test_compilation_conversation import chat as chat
from test_compilation_conversation import mutate_state, say, start, step
from test_compilation_semantics import (
    approve_form,
    classified,
    layout_fields,
    meaning,
)
from test_compilation_semantics import (
    source as evidence,
)
from test_compilation_semantics import (
    source_review as approve_source,
)
from test_compilation_sessions import docx, simulate, source

from app import compilation_session_resolution as resolution
from app import compilation_sessions as sessions
from app import main
from app.compilation_chat import MAX_AUTO_STEPS, ChatControl
from app.compilation_semantics import (
    FormReview,
    HistoricalServices,
    SourceReview,
    review_meanings,
    review_sources,
)
from app.compilation_session_models import CandidateMatches, CandidateMeanings
from app.docx_templates import inspect_docx
from app.generation import GenerationError


@pytest.fixture
def anyio_backend():
    return "asyncio"


@pytest.mark.anyio
@pytest.mark.parametrize("bad", ["unknown", "other_batch", "duplicate", "missing", "invalid"])
async def test_form_review_isolates_bad_verdicts(bad):
    _, all_fields = layout_fields()
    fields = all_fields[:2]
    meanings = CandidateMeanings(fields=[meaning(f, name) for f, name in
                                         zip(fields, ["Denominazione", "Indirizzo"], strict=True)])

    async def assess(task, data, schema):
        raw = (await approve_form(task, data, schema)).model_dump()
        if bad == "missing":
            raw["fields"] = raw["fields"][1:]
        elif bad == "invalid":
            raw["fields"][0]["accepted"] = "invalid"
        else:
            extra = deepcopy(raw["fields"][0])
            extra["candidate_id"] = (all_fields[2]["id"] if bad == "other_batch" else
                                     "foreign" if bad == "unknown" else fields[0]["id"])
            raw["fields"].append(extra)
        return resolution.parse_structured(json.dumps(raw), schema)

    await review_meanings(fields, meanings, assess)
    resolution.apply_meanings(fields, meanings)
    assert fields[1]["requirement"] is not None
    assert bool(fields[0]["requirement"]) == (bad in {"unknown", "other_batch"})
    assert all_fields[2]["requirement"] is None


@pytest.mark.anyio
@pytest.mark.parametrize("bad", ["unknown", "other_batch", "duplicate", "missing", "invalid",
                                 "kind", "index", "kind_object"])
async def test_source_review_isolates_bad_verdicts_without_wrong_assignments(bad):
    layout, all_fields = layout_fields()
    fields = all_fields[:2]
    await classified(fields, ["Denominazione", "Indirizzo"])
    src = evidence("Denominazione: Aurora. Indirizzo: Via Roma 12. cooperativa")
    matches = CandidateMatches(fields=[{
        "candidate_id": f["id"], "reason": "Fonte", "supports": [
            {"source_id": 1, "quote": src["content"], "value": value},
        ],
    } for f, value in zip(fields, ["Aurora", "Via Roma 12"], strict=True)])
    untouched = deepcopy(all_fields[2])

    async def assess(task, data, schema):
        raw = (await approve_source(task, data, schema)).model_dump()
        if bad == "missing":
            raw["supports"] = raw["supports"][1:]
        elif bad == "invalid":
            raw["supports"][0]["accepted"] = "invalid"
        elif bad in {"kind", "index"}:
            extra = deepcopy(raw["supports"][0])
            extra[bad] = "invalid"
            raw["supports"].append(extra)
        elif bad == "kind_object":
            extra = deepcopy(raw["supports"][0])
            extra["kind"] = {"invalid": True}
            raw["supports"].append(extra)
        else:
            extra = deepcopy(raw["supports"][0])
            extra["candidate_id"] = (all_fields[2]["id"] if bad == "other_batch" else
                                     "foreign" if bad == "unknown" else fields[0]["id"])
            raw["supports"].append(extra)
        return resolution.parse_structured(json.dumps(raw), schema)

    coverage = {f["id"]: [1] for f in fields}
    await review_sources(fields, matches, [src], coverage, assess)
    resolution.apply_matches(layout, fields, matches, [src], coverage)
    assert fields[1]["status"] == "RESOLVED" and fields[1]["value"] == "Via Roma 12"
    assert (fields[0]["status"] == "RESOLVED") == (bad in {"unknown", "other_batch"})
    assert all_fields[2] == untouched


@pytest.mark.anyio
async def test_verdict_for_filtered_historical_offer_cannot_restore_it_or_crash_mapping():
    layout, fields = layout_fields("Servizi già eseguiti da cooperativa ______ Indirizzo ______")
    meanings = CandidateMeanings(fields=[
        meaning(fields[0], "Servizio", context_role="PAST_SERVICE",
                context_quote="Servizi già eseguiti"),
        meaning(fields[1], "Indirizzo"),
    ])
    await review_meanings(fields, meanings, approve_form)
    resolution.apply_meanings(fields, meanings)
    src = evidence("Denominazione: Aurora. Indirizzo: Via Roma 12. cooperativa")
    matches = CandidateMatches(fields=[{
        "candidate_id": f["id"], "reason": "Fonte", "supports": [
            {"source_id": 1, "quote": src["content"], "value": value},
        ],
    } for f, value in zip(fields, ["Aurora", "Via Roma 12"], strict=True)])

    async def assess(task, data, schema):
        if schema is HistoricalServices:
            return HistoricalServices(facts=[{
                "source_id": 99, "performer_quote": "Aurora", "service_quote": "servizio",
                "performance_quote": "Aurora ha eseguito un servizio.",
            }])
        raw = (await approve_source(task, data, schema)).model_dump()
        raw["supports"].append({"candidate_id": fields[0]["id"], "kind": "value",
                                "index": 0, "accepted": True, "reason": "Non richiesto"})
        return SourceReview.model_validate(raw)

    coverage = {f["id"]: [1] for f in fields}
    await review_sources(fields, matches, [src], coverage, assess)
    resolution.apply_matches(layout, fields, matches, [src], coverage)
    assert fields[0]["status"] == "MISSING" and fields[0]["value"] is None
    assert fields[1]["status"] == "RESOLVED" and fields[1]["value"] == "Via Roma 12"


@pytest.mark.anyio
@pytest.mark.parametrize("failure", ["provider", "timeout"])
async def test_known_model_outage_pauses_without_burning_remaining_fields(
    chat, monkeypatch, failure,
):
    state, _ = await start(chat)
    await source(chat[0], "Denominazione: Aurora. Partita IVA: IT12345678901.")
    simulate(monkeypatch)

    async def unavailable(*args, **kwargs):
        raise GenerationError("Dettaglio tecnico provider")

    async def broken(*args):
        if failure == "timeout":
            raise TimeoutError("Dettaglio tecnico timeout")
        return await resolution.request_structured("Test provider", {}, CandidateMatches)

    monkeypatch.setattr(resolution, "request_model_content", unavailable)
    monkeypatch.setattr(resolution, "match_candidates", broken)
    state = await step(chat[0], state)
    assert state["status"] != "FAILED" and state["chat"]["paused"]
    assert not state["chat"]["auto_continue"] and state["chat"]["question"] is None
    assert "Ho conservato" in state["chat"]["notice"]
    assert state["last_error"] is None
    assert set(state["chat_workflow"]["source_attempts"].values()) == {1}
    assert set(state["chat_workflow"]["analysis_attempts"].values()) == {1}
    before = deepcopy(state["fields"])
    state = sessions.chat_control("alpha", state["id"], state["version"],
                                  ChatControl(kind="resume", user_quote="riprendi"))
    assert state["fields"] == before and state["chat"]["auto_continue"]


@pytest.mark.anyio
async def test_internal_coverage_bug_is_not_recovered_as_model_output(chat, monkeypatch):
    state, _ = await start(chat)
    simulate(monkeypatch)

    async def corrupt(*args):
        return [], {"internal-wrong-id": []}, []

    monkeypatch.setattr(resolution, "retrieve_sources", corrupt)
    with pytest.raises(ValueError, match="Mapping interno"):
        await step(chat[0], state)
    state = await current(chat[0], state)
    assert state["status"] == "FAILED"
    assert not state["chat_workflow"].get("output_rejections")


@pytest.mark.anyio
async def test_consecutive_model_errors_two_resumes_and_partial_export_preserve_source_user(
    chat, monkeypatch,
):
    state, _ = await start(chat, content=docx(("Denominazione sociale", "CF professionista",
                                             "Referente", "Sede legale", "Altro dato"),
                                            context="Dati aziendali; professionista singolo"))
    await source(chat[0], "Denominazione sociale: Aurora Progetti S.r.l.\n"
                 "Sede legale: Via Roma 12, Bari.")
    simulate(monkeypatch, {"t0.r0.c1": ["Aurora Progetti S.r.l."],
                           "t0.r3.c1": ["Via Roma 12, Bari"]})
    classify, match = resolution.classify_candidates, resolution.match_candidates

    async def first(fields):
        result = await classify(fields)
        result.fields = result.fields[:-1]
        result.fields[1].condition = "professionista singolo"
        return result

    async def bad_source(*args):
        result = await match(*args)
        result.fields[-1].supports[0].quote = "Estratto inventato"
        return result

    monkeypatch.setattr(resolution, "classify_candidates", first)
    monkeypatch.setattr(resolution, "match_candidates", bad_source)
    state = await step(chat[0], state)
    mutate_state(state, lambda s: s["chat_workflow"].update(steps=MAX_AUTO_STEPS))
    state = await current(chat[0], state)
    assert state["chat"]["question"]["kind"] == "applicability"
    response = await say(chat, state, "No. Aurora è una società, non un professionista singolo.",
                         {"action": "CONDITION_FALSE"})
    assert response.status_code == 200
    state = await current(chat[0], state)
    response = await chat[0].patch(f"/api/projects/alpha/compilation-sessions/{state['id']}/fields",
        json={"version": state["version"], "fields": [
            {"field_id": "t0.r2.c1", "action": "set", "value": "Anna Bianchi"},
        ]})
    assert response.status_code == 200, response.text
    state = response.json()
    preserved = deepcopy(state["fields"][:3])

    async def rejected(*args):
        return resolution.parse_structured("{", CandidateMeanings)

    monkeypatch.setattr(resolution, "match_candidates", rejected)
    state = await step(chat[0], state)
    assert state["fields"][:3] == preserved
    monkeypatch.setattr(resolution, "classify_candidates", rejected)
    state = await step(chat[0], state)
    assert state["fields"][:3] == preserved
    assert state["chat_workflow"]["source_attempts"]["t0.r3.c1"] == 2
    assert state["chat_workflow"]["analysis_attempts"]["t0.r4.c1"] == 2
    assert not state["chat"]["auto_continue"]
    prior_version = state["version"]
    for _ in range(2):
        state = sessions.chat_control("alpha", state["id"], state["version"],
                                      ChatControl(kind="resume", user_quote="riprendi"))
        assert state["fields"][:3] == preserved and not state["chat"]["auto_continue"]
        assert not set(f["id"] for f in preserved).intersection(
            (state["chat"]["question"] or {}).get("field_ids", []))
    with pytest.raises(HTTPException) as stale:
        sessions.chat_control("alpha", state["id"], prior_version,
                              ChatControl(kind="resume", user_quote="riprendi"))
    assert stale.value.status_code == 409
    assert state["chat"]["verified"] == state["summary"]["resolved"] == 1
    assert state["chat"]["user_provided"] == state["summary"]["user_provided"] == 1
    assert state["fields"][1]["status"] == "NOT_APPLICABLE"
    assert state["fields"][1]["provenance"] == "USER"
    response = await chat[0].post(
        f"/api/projects/alpha/compilation-sessions/{state['id']}/finalize",
        json={"version": state["version"], "allow_unresolved": True},
    )
    assert response.status_code == 200, response.text
    exported = response.json()
    draft = await chat[0].get(exported["last_generation"]["downloads"]["docx"])
    document = Document(BytesIO(draft.content))
    assert document.tables[0].cell(0, 1).text == "Aurora Progetti S.r.l."
    assert document.tables[0].cell(1, 1).text == ""
    assert document.tables[0].cell(2, 1).text == "Anna Bianchi"
    assert document.tables[0].cell(3, 1).text == document.tables[0].cell(4, 1).text == ""
    report = (await chat[0].get(exported["last_generation"]["downloads"]["report"])).json()
    assert not report["ready_for_submission"]
    assert exported["fields"][:3] == preserved


@pytest.mark.anyio
@pytest.mark.parametrize("bad", ["unknown", "duplicate", "invalid", "candidate_id"])
async def test_partial_user_group_keeps_only_unambiguous_active_slots(chat, monkeypatch, bad):
    state = await checkpoint(chat, monkeypatch)
    raw = [reply(1, value="055123456", quote="055123456")]
    other = reply(2, value="Anna Bianchi", quote="Anna Bianchi")
    if bad == "unknown":
        other["slot"] = 99
    elif bad == "invalid":
        other["action"] = "INVALID"
    elif bad == "candidate_id":
        other["candidate_id"] = "t0.r0.c1"
    raw.append(other)
    if bad == "duplicate":
        raw.append(deepcopy(other))
    response = await say(chat, state, "1. 055123456; 2. Anna Bianchi", answer(*raw))
    assert response.status_code == 200, response.text
    state = await current(chat[0], state)
    assert state["fields"][0]["value"] == "055123456"
    assert state["fields"][0]["provenance"] == "USER"
    assert state["fields"][1]["value"] is None
    assert state["fields"][0]["id"] not in state["chat"]["question"]["field_ids"]


@pytest.mark.anyio
async def test_chat_provider_failure_is_friendly_and_standalone_pause_resume_still_work(
    chat, monkeypatch,
):
    state = await checkpoint(chat, monkeypatch)
    before = deepcopy(state["fields"])

    async def unavailable(*args, **kwargs):
        raise GenerationError("Traceback: errore tecnico riservato")

    monkeypatch.setattr(main, "plan_chat_turn", unavailable)
    response = await say(chat, state, "Anna Bianchi", answer(reply(2, value="Anna Bianchi")))
    assert response.status_code == 200
    assert "Ho conservato" in response.json()["answer"]
    assert "Traceback" not in json.dumps(response.json())
    assert (await current(chat[0], state))["version"] == state["version"]
    for command, outcome in [("basta", "paused"), ("riprendi", "resumed")]:
        response = await say(chat, state, command, answer())
        assert response.status_code == 200, response.text
        assert response.json()["compilation"]["action"] == outcome
        state = await current(chat[0], state)
        assert state["fields"] == before


def parent_form():
    document = Document()
    document.add_paragraph("Da compilare se la candidatura è presentata da un raggruppamento")
    document.add_paragraph("Sono elencati i componenti")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "Denominazione componente", "Codice fiscale"
    output = BytesIO()
    document.save(output)
    layout = inspect_docx(output.getvalue())
    fields = sessions.candidate_snapshots(layout, "alpha", 1, "modulo.docx")
    field = fields[0]
    item = meaning(field, "Denominazione componente", subject_anchor="componenti",
                   section_id="paragraph:1", context_role="DECLARATION",
                   context_quote="Sono elencati i componenti", condition_kind="subject_type")
    item["condition"] = "raggruppamento"
    return field, item


@pytest.mark.anyio
@pytest.mark.parametrize("correct_kind", [True, False])
async def test_unique_parent_condition_needs_independent_review_and_correct_subject(correct_kind):
    field, item = parent_form()
    meanings = CandidateMeanings(fields=[item])

    async def assess(task, data, schema):
        assert data["proposals"][0]["interpretation"]["semantic"]["section_id"] == "paragraph:0"
        result = (await approve_form(task, data, schema)).model_dump()
        if correct_kind:
            result["fields"][0]["section_condition"] = {
                "condition": "raggruppamento", "section_id": "paragraph:0",
                "condition_kind": "participation",
            }
        return FormReview.model_validate(result)

    await review_meanings([field], meanings, assess)
    resolution.apply_meanings([field], meanings)
    assert bool(field["requirement"]) == correct_kind
    assert field["value"] is None


@pytest.mark.anyio
async def test_invented_parent_condition_still_fails_before_review():
    field, item = parent_form()
    item["condition"] = "società inventata"
    meanings = CandidateMeanings(fields=[item])

    async def forbidden(*args):
        pytest.fail("Una condizione inventata non deve arrivare alla revisione")

    await review_meanings([field], meanings, forbidden)
    resolution.apply_meanings([field], meanings)
    assert field["requirement"] is None and field["status"] == "PENDING"


def test_reinterpretation_cannot_forget_a_persisted_user_condition():
    field, item = parent_form()
    field["condition"] = "raggruppamento"
    field["applicability"] = {"condition": "raggruppamento", "applies": True,
                              "provenance": "USER", "user_quote": "Sì"}
    proof = deepcopy(field["applicability"])
    item["condition"] = "componenti"
    resolution.apply_meanings([field], CandidateMeanings(fields=[item]))
    assert field["applicability"] == proof and field["condition"] == "raggruppamento"
    assert field["requirement"] is None
    assert "USER" in field["validation_errors"][0]


@pytest.mark.parametrize("provenance,status", [("USER", "USER_PROVIDED"), ("SOURCE", "RESOLVED")])
@pytest.mark.parametrize("exclusive", [False, True])
def test_dependency_propagation_never_discards_an_independent_persisted_value(
    provenance, status, exclusive,
):
    from app.compilation_clarifications import synchronize

    field, _ = parent_form()
    field.update(status=status, value="Aurora", provenance=provenance,
                 condition="raggruppamento", entity="company")
    field["semantic"] = {"section_id": "paragraph:1", "exclusive_group_id": "choice"}
    field["semantic_validation"] = {"accepted": True}
    other = deepcopy(field)
    other.update(id="other", value=None, status="MISSING", provenance=None)
    if exclusive:
        other["semantic"]["section_id"] = "paragraph:0"
        other["condition"] = "candidatura"
    other["applicability"] = {"condition": other["condition"], "applies": exclusive,
                              "provenance": "SOURCE", "evidence": [{"quote": "Fatto"}]}
    state = {"status": "WAITING_FOR_USER", "fields": [field, other],
             "chat_workflow": {"clarification_mode": "grouped", "steps": MAX_AUTO_STEPS}}
    before = deepcopy(field)
    synchronize(state)
    assert field == before


def test_omitted_user_condition_is_present_for_independent_form_review():
    field, item = parent_form()
    field["condition"] = "raggruppamento"
    field["applicability"] = {"condition": "raggruppamento", "applies": True,
                              "provenance": "USER", "user_quote": "Sì"}
    item["condition"] = ""
    meanings = resolution.bind_meanings([field], CandidateMeanings(fields=[item]))
    assert meanings.fields[0].condition == "raggruppamento"


@pytest.mark.anyio
async def test_ambiguous_parent_condition_location_is_not_guessed():
    field, item = parent_form()
    field["structural"]["form_sections"].append(
        {"id": "second-parent", "text": "Da compilare per un raggruppamento"},
    )

    async def forbidden(*args):
        pytest.fail("Due sezioni possibili richiedono una localizzazione esplicita")

    meanings = CandidateMeanings(fields=[item])
    await review_meanings([field], meanings, forbidden)
    resolution.apply_meanings([field], meanings)
    assert field["requirement"] is None


def test_deeply_nested_json_is_a_bounded_output_failure():
    with pytest.raises(resolution.CompilationOutputError):
        resolution.parse_structured("[" * 2000 + "]" * 2000, CandidateMeanings)
