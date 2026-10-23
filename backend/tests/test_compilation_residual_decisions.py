"""FORM-anchored decision questions; receipts never silently write their targets."""

import copy
import json

import pytest
from docx import Document
from test_compilation_catalog import encoded, native_checkbox
from test_compilation_global_sessions import BASE, api, provider, resolve

from app import compilation_global_dialogue as dialogue
from app import compilation_global_sessions as adapter
from app import compilation_sessions as sessions
from app.compilation_chat import ClarificationReply
from app.compilation_global import validate_global_proposal
from app.db import connection
from app.docx_templates import inspect_docx
from app.intents import ChatDecision

__all__ = ["api", "provider"]


@pytest.fixture
def anyio_backend():
    return "asyncio"


def document():
    doc = Document()
    doc.add_heading("Domanda", 1)
    doc.add_paragraph("Scegliere i settori interessati – max. 2 tipologie:")
    for title in ["Servizi di progettazione", "Coordinamento della sicurezza"]:
        doc.add_paragraph(title)
        table = doc.add_table(rows=3, cols=3)
        for cell, text in zip(table.rows[0].cells, ["Categoria di opere", "Fascia A", "Fascia B"]):
            cell.text = text
        table.cell(1, 0).text, table.cell(2, 0).text = "Strutture", "Impianti"
    doc.add_paragraph("DICHIARA")
    p = doc.add_paragraph()
    native_checkbox(p, legacy=False)
    p.add_run(" che i fatti, stati e qualità corrispondono a verità:")
    doc.add_paragraph("che il curriculum allegato è autentico;")
    doc.add_paragraph("di acconsentire al trattamento dei dati personali.")
    p = doc.add_paragraph()
    native_checkbox(p, legacy=False)
    p.add_run(" che in caso di conferimento dell'incarico i professionisti che eseguiranno "
              "la prestazione sono:")
    doc.add_paragraph("________________________")
    doc.add_paragraph("________________________")
    doc.add_paragraph("e che il professionista incaricato dell’integrazione delle prestazioni "
                      "specialistiche è:")
    doc.add_paragraph("________________________")
    doc.add_paragraph("________________________")
    doc.add_paragraph("Luogo e data, ________________________")
    doc.add_paragraph("Firma digitale")
    doc.add_paragraph("________________________")
    return encoded(doc)


def empty_proposal():
    return {"version": 12, "facts": {}, "fields": {}, "conditions": {}, "bindings": {},
            "choice_groups": {}, "clarifications": {}, "coverage": []}


def data():
    layout = inspect_docx(document())
    current = adapter.initial_state(layout)
    audit = validate_global_proposal(json.dumps(empty_proposal()), layout, "alpha", [])
    current.update(audit=audit,
                   dialogue={"answers": {}, "user_values": {}})
    return current


def decisions(current):
    return [q for q in dialogue.question_pool(current) if q["kind"] == "decision"]


def test_questions_group_actual_decisions_and_preserve_full_form_context():
    current = data()
    pool = decisions(current)
    assert len(pool) == 5
    assert len({q["key"] for q in pool}) == 5
    prefs = next(q for q in pool if q["decision_type"] == "preferences")
    assert len(prefs["field_ids"]) == 8 and "max. 2" in prefs["context"]
    assert "Fascia A" in prefs["context"] and "Strutture" in prefs["context"]
    attestation = next(q for q in pool if q["decision_type"] == "attestation")
    assert "curriculum allegato è autentico" in attestation["context"]
    assert "acconsentire al trattamento" in attestation["context"]
    assert [len(q["field_ids"]) for q in pool if q["decision_type"] == "designation"] == [3, 2]
    assert not any("Firma digitale" in q["context"] for q in pool)
    assert current["audit"]["choice_groups"] == {} and not current["dialogue"]["user_values"]


@pytest.mark.parametrize("action", ["UNKNOWN", "SKIP", "REFUSE", "VALUE"])
def test_answered_decision_not_repeated_and_resolved_targets_not_asked(action):
    current = data()
    pool = decisions(current)
    question = pool[0]
    current["dialogue"]["answers"][question["key"]] = {
        "action": action, "provenance": "USER", "user_quote": "una risposta", "value": "risposta"}
    assert question["key"] not in {q["key"] for q in decisions(current)}
    other = next(q for q in pool if q["decision_type"] == "location_date")
    current["dialogue"]["user_values"][other["field_ids"][0]] = {"value": "Roma, 20/10/2026"}
    assert other["key"] not in {q["key"] for q in decisions(current)}


def test_source_proposals_existing_questions_and_unknown_branches_not_duplicated():
    current = data()
    target = next(q for q in decisions(current) if q["decision_type"] == "designation")
    for field in current["audit"]["fields"]:
        if field["candidate_id"] in target["field_ids"]:
            field.update(proposal={"fact_id": "documented"}, validation="GROUNDED_PROPOSAL")
    assert target["key"] not in {q["key"] for q in decisions(current)}
    current = data()
    target = decisions(current)[0]
    current["audit"]["conditions"]["c"] = {"result": "unknown"}
    for field in current["audit"]["fields"]:
        if field["candidate_id"] in target["field_ids"]:
            field["conditions"] = ["c"]
    assert target["key"] not in {q["key"] for q in decisions(current)}
    current["audit"]["version"] = 11
    assert not decisions(current)


def test_coverage_claim_alone_cannot_create_a_question_or_native_choice_group():
    current = data()
    current["audit"]["coverage_groups"] = [{"outcome": "manual_input", "accepted_ids": ["fake"],
        "form_evidence": {"paragraph_id": "fake", "quote": "conferma tutto"}}]
    baseline = decisions(current)
    current["catalog"]["blocks"] = []
    assert not decisions(current) and len(baseline) == 5


def test_preference_receipt_identity_is_stable_after_one_target_is_resolved():
    current = data()
    target = next(q for q in decisions(current) if q["decision_type"] == "preferences")
    current["dialogue"]["answers"][target["key"]] = {
        "action": "VALUE", "provenance": "USER", "user_quote": "decisione già registrata"}
    current["dialogue"]["user_values"][target["field_ids"][0]] = {"value": "X"}
    assert not any(q["decision_type"] == "preferences" for q in decisions(current))


def test_false_condition_stays_false_even_with_a_contradicting_user_receipt():
    current = data()
    target = next(q for q in decisions(current) if q["decision_type"] == "designation")
    current["audit"]["conditions"]["c"] = {"result": "false"}
    current["dialogue"]["answers"]["condition:c"] = {"applies": True}
    for field in current["audit"]["fields"]:
        if field["candidate_id"] in target["field_ids"]:
            field["conditions"] = ["c"]
    assert target["key"] not in {q["key"] for q in decisions(current)}


def test_decision_fallback_and_partial_export_offer_when_narrator_omits_questions():
    pool = dialogue.question_pool(data())
    rendered = dialogue.render_dialogue(dialogue.DialoguePlan(
        message="Ho conservato i dati disponibili.", questions=[]), pool, exportable_count=0)
    assert len(rendered["question"]["slots"]) == 3
    assert all(s["kind"] == "decision" for s in rendered["question"]["slots"])
    assert "bozza parziale" in rendered["message"] and "non sai" in rendered["message"]
    assert rendered["empty_export_offered"] is True
    assert "senza aggiungere valori" in rendered["message"]


async def start_decisions(api, monkeypatch):
    uploaded = await api.post("/api/projects/alpha/forms", files={
        "file": ("domanda.docx", document())})
    assert uploaded.status_code == 201

    async def global_response(*args, **kwargs):
        return json.dumps(empty_proposal()), "simulated", {}

    monkeypatch.setattr(adapter, "request_model_content", global_response)
    started = (await api.post(BASE, json={"form_id": uploaded.json()["id"], "engine": "global",
                                        "start_in_chat": True})).json()
    return await resolve(api, started)


@pytest.mark.anyio
@pytest.mark.parametrize("action,value", [("UNKNOWN", None), ("SKIP", None), ("REFUSE", None),
    ("VALUE", "Scelgo Strutture, fascia A")])
async def test_decision_receipts_do_not_select_boxes_or_write_fields(
    api, provider, monkeypatch, action, value,
):
    current = await start_decisions(api, monkeypatch)
    slot = current["chat"]["question"]["slots"][0]
    assert slot["kind"] == "decision"
    message = value or "Non so rispondere a questa decisione"
    reply = ClarificationReply(slot=slot["slot"], action=action, value=value, user_quote=message)
    current, errors = dialogue.apply_replies("alpha", current["id"], current["version"],
                                             [reply], message)
    assert not errors
    current = await resolve(api, current)
    receipt = current["global_compilation"]["dialogue"]["answers"][slot["key"]]
    assert receipt["provenance"] == "USER" and receipt["write_authorized"] is False
    assert not current["global_compilation"]["dialogue"]["user_values"]
    assert not dialogue.export_selection(current)[0]
    assert slot["key"] not in {q["key"] for q in current["chat"]["question"]["slots"]}
    assert len(provider) == 0


@pytest.mark.anyio
async def test_finish_can_export_unchanged_original_with_decisions_still_open(
    api, provider, monkeypatch,
):
    current = await start_decisions(api, monkeypatch)
    assert current["chat"]["question"]["slots"]
    before = copy.deepcopy(current["global_compilation"]["dialogue"].get("answers", {}))
    decision = ChatDecision(action="compilation_generate", answer="", queries=[], target="form")
    _, command = await adapter.handle_chat_action(current, decision, current["version"],
                                                 "Non so rispondere, genera una bozza parziale")
    assert command["action"] == "generated"
    current = sessions.get_session("alpha", current["id"])
    report = (await api.get(current["last_generation"]["downloads"]["report"])).json()
    assert report["written_field_count"] == 0 and report["checkbox_values"] == {}
    assert report["conversation_review"]["answers"] == before
    with connection() as db:
        stored = json.loads(sessions._row(db, "alpha", current["id"])["state_json"])
    assert not stored["global_compilation"]["dialogue"].get("user_values")


@pytest.mark.anyio
@pytest.mark.parametrize("change", [
    {"action": "CONDITION_TRUE", "value": None},
    {"value": "dato inventato"}, {"selected_options": [1]},
    {"normalized_value": "01/01/2026"},
])
async def test_decision_cannot_act_as_applicability_or_choice_or_invent_a_value(
    api, provider, monkeypatch, change,
):
    current = await start_decisions(api, monkeypatch)
    slot = current["chat"]["question"]["slots"][0]
    message = "Scelgo Strutture, fascia A"
    reply = ClarificationReply(**{"slot": slot["slot"], "action": "VALUE", "value": message,
                                  "user_quote": message, **change})
    current, errors = dialogue.apply_replies("alpha", current["id"], current["version"],
                                             [reply], message)
    assert errors == [slot["slot"]]
    assert not current["global_compilation"]["dialogue"]["answers"]
    assert not current["global_compilation"]["dialogue"]["user_values"]
