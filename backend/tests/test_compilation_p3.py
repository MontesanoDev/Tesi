"""P3: generation by FORM segments and decision state of the dialogue (simulated provider)."""

import json
from io import BytesIO

import pytest
from docx import Document
from test_compilation_global_sessions import BASE, api, proposal, provider, resolve, start
from test_compilation_sessions import source

from app import compilation_decision_registry as registry
from app import compilation_global_dialogue as dialogue
from app import compilation_global_sessions as adapter
from app import compilation_sessions as sessions
from app.compilation_catalog import build_compilation_catalog
from app.compilation_chat import ClarificationReply, hedged, value_alternatives
from app.compilation_global import global_request, validate_global_proposal
from app.compilation_segments import attach_status, merge, parse, partition
from app.db import connection
from app.docx_templates import inspect_docx
from app.generation import GenerationTruncatedError

__all__ = ["api", "provider"]


@pytest.fixture(autouse=True)
def p3_runtime(monkeypatch):
    """P3 is opt-in: these tests run the private experiment runtime (both revisions)."""
    monkeypatch.setenv("MAPI_DIALOGUE_REVISION", "1")
    monkeypatch.setenv("MAPI_GENERATION_REVISION", "1")


def slots_docx(count, *, clause=None):
    doc = Document()
    for i in range(count):
        doc.add_paragraph(f"Dato {i}: ____________")
        if clause is not None and i == clause:
            doc.add_paragraph("☐ per le società DI ESSERE iscritto alla CCIAA con i seguenti dati:")
            table = doc.add_table(rows=1, cols=4)
            table.rows[0].cells[0].text = "Provincia di iscrizione:"
            table.rows[0].cells[2].text = "numero di iscrizione:"
    stream = BytesIO()
    doc.save(stream)
    return stream.getvalue()


def catalog_of(data):
    return build_compilation_catalog(inspect_docx(data))


def test_partition_covers_every_candidate_once_in_document_order():
    catalog = catalog_of(slots_docx(130))
    segments = partition(catalog)
    ids = [cid for s in segments for cid in s["candidate_ids"]]
    assert sorted(ids) == sorted(catalog["candidates"]) and len(ids) == len(set(ids))
    assert 2 <= len(segments) <= 6 and all(len(s["candidate_ids"]) <= 60 for s in segments)
    order = [b["id"] for b in catalog["blocks"]]
    blocks = [b for s in segments for b in s["block_ids"]]
    assert blocks == order  # Contiguous, never split, nothing lost.
    assert len(partition(catalog_of(slots_docx(20)))) == 1


def test_option_clause_stays_with_its_table_across_a_boundary():
    catalog = catalog_of(slots_docx(70, clause=59))
    segments = partition(catalog)
    table_cells = [cid for cid, c in catalog["candidates"].items() if c["location"].get("table_id")]
    owners = {s["index"] for s in segments for cid in table_cells if cid in s["candidate_ids"]}
    clause = next(b["id"] for b in catalog["blocks"] for p in b["paragraphs"]
                  if "CCIAA" in p["text"])
    assert len(owners) == 1
    assert clause in next(s for s in segments if s["index"] in owners)["block_ids"]


def test_parse_accepts_literal_tabs_but_never_repairs():
    assert parse('{"version":12,"note":"5.a)\tdati"}')["note"] == "5.a)\tdati"
    for broken in ['{"version":12,"fields":{', '{"version":12}}', '{"version":11}']:
        with pytest.raises(ValueError):
            parse(broken)


def fact(value, chunk=-1, quote=None):
    return {"subject": "Aurora S.r.l.", "entity": "company", "aspect": "identity",
            "property": "codice", "value": value,
            "evidence": [{"chunk_id": chunk, "quote": quote or f"Codice: {value}"}]}


def assignment(cid, fact_id, paragraph):
    return {"label": "Codice", "kind": "data", "disposition": "value", "conditions": [],
            "form": {"paragraph_id": paragraph, "quote": f"Dato {paragraph.split(':')[1]}: "
                     f"[[{cid}]]"}, "fact_id": fact_id, "property": "codice"}


def two_segment_case(second_value):
    data = slots_docx(70)
    layout, catalog = inspect_docx(data), catalog_of(data)
    segments = partition(catalog)
    first, last = segments[0]["candidate_ids"][0], segments[1]["candidate_ids"][-1]
    paragraph = {cid: catalog["candidates"][cid]["context_id"] for cid in (first, last)}
    content = "Codice: ABC123\nCodice: " + second_value
    payloads = {
        1: {"version": 12, "facts": {"f": fact("ABC123")}, "fields": {
            first: assignment(first, "f", paragraph[first]),
            last: assignment(last, "f", paragraph[last])}},  # Outside segment 1: dropped.
        2: {"version": 12, "facts": {"f": fact(second_value)}, "fields": {
            last: assignment(last, "f", paragraph[last])}},
    }
    source = {"chunk_id": -1, "file_id": -21, "chunk_index": 0, "source_name": "dati.md",
              "role": "source", "scope": "global", "category": "company", "project_id": None,
              "content": content}
    return layout, segments, payloads, [source], first, last


def test_merge_renames_deduplicates_drops_and_validates_like_one_proposal():
    layout, segments, payloads, sources, first, last = two_segment_case("ABC123")
    merged, report = merge(segments, payloads)
    assert report["dropped_out_of_segment"] == [{"segment": 1, "candidate_id": last}]
    assert list(merged["facts"]) == ["s1_f"] and report["deduplicated"]["facts"] == 1
    assert merged["fields"][first]["fact_id"] == merged["fields"][last]["fact_id"] == "s1_f"
    audit = validate_global_proposal(json.dumps(merged), layout, "alpha", sources)
    fields = {f["candidate_id"]: f for f in audit["fields"]}
    assert fields[first]["validation"] == fields[last]["validation"] == "GROUNDED_PROPOSAL"


def test_cross_segment_conflict_is_a_review_never_a_hidden_choice():
    layout, segments, payloads, sources, first, last = two_segment_case("XYZ789")
    merged, report = merge(segments, payloads)
    assert len(report["conflicts"]) == 1
    for cid in (first, last):
        assert merged["fields"][cid]["disposition"] == "review"
        assert sorted(merged["fields"][cid]["alternative_fact_ids"]) == ["s1_f", "s2_f"]
    audit = validate_global_proposal(json.dumps(merged), layout, "alpha", sources)
    for field in audit["fields"]:
        assert field["validation"] != "GROUNDED_PROPOSAL"
        assert field.get("proposal_grounding") != "GROUNDED_PROPOSAL"


def test_single_segment_keeps_identifiers_and_the_monolithic_request():
    catalog = catalog_of(slots_docx(5))
    segments = partition(catalog)
    payload = {"version": 12, "facts": {"f": fact("ABC")}, "fields": {}, "conditions": {}}
    merged, _ = merge(segments, {1: payload})
    assert list(merged["facts"]) == ["f"]
    settings = type("S", (), {"model": "simulated"})()
    assert len(global_request(settings, {"form": catalog})["messages"]) == 2


def test_failed_segment_is_visible_backend_coverage_not_a_model_explanation():
    layout, segments, payloads, sources, first, last = two_segment_case("ABC123")
    merged, report = merge(segments, {1: payloads[1]})
    report["segments"] = [{"index": 1, "status": "analyzed", "candidate_count": 60,
                           "candidate_ids": segments[0]["candidate_ids"], "error": None},
                          {"index": 2, "status": "failed", "candidate_count": 10,
                           "candidate_ids": segments[1]["candidate_ids"],
                           "error": "GenerationTruncatedError: limite"}]
    audit = attach_status(validate_global_proposal(json.dumps(merged), layout, "alpha",
                                                   sources), report)
    fields = {f["candidate_id"]: f for f in audit["fields"]}
    assert fields[last]["coverage"]["origin"] == "BACKEND"
    assert fields[last]["coverage"]["state"] == "UNASSESSED"
    assert "non analizzata" in fields[last]["coverage"]["reason"]
    assert audit["generation_segments"]["unanalyzed_candidate_ids"] == segments[1]["candidate_ids"]
    assert fields[first]["validation"] == "GROUNDED_PROPOSAL"


@pytest.mark.anyio
@pytest.mark.parametrize("failure", ["second", "all", "monolithic_runtime"])
async def test_segmented_session_keeps_successful_segments_and_never_retries(
    api, provider, monkeypatch, failure,
):
    if failure == "monolithic_runtime":
        monkeypatch.delenv("MAPI_GENERATION_REVISION")
    upload = await api.post("/api/projects/alpha/forms", files={
        "file": ("modulo.docx", slots_docx(70))})
    await source(api, "Codice: ABC123\nNumero: 4455", name="dati.md")
    calls = []

    async def model(client, settings, body, timeout, **kwargs):
        inp = json.loads(body["messages"][1]["content"])
        segment = body["messages"][-1]["content"] if len(body["messages"]) == 3 else ""
        calls.append(segment[:12])
        if failure == "all" or (failure == "second" and segment.startswith("SEGMENTO 2")):
            raise GenerationTruncatedError(49152, 49152)
        evidence = next(s for s in inp["sources"] if "ABC123" in s["content"])
        wanted = segment.rsplit("[", 1)[1].rstrip("]").split(",") if segment else list(
            inp["form"]["candidates"])
        fields = {}
        for cid in wanted[:1] + wanted[-1:]:
            paragraph = inp["form"]["candidates"][cid]["context_id"]
            fields[cid] = assignment(cid, "f", paragraph)
        return json.dumps({"version": 12, "uncertain_paragraph_ids": [], "conditions": {},
                           "bindings": {}, "choice_groups": {}, "clarifications": {},
                           "coverage": [], "facts": {"f": fact("ABC123", evidence["chunk_id"])},
                           "fields": fields}), "simulated", {"total_tokens": 10}

    monkeypatch.setattr(adapter, "request_model_content", model)
    session = (await api.post(BASE, json={"form_id": upload.json()["id"], "engine": "global",
                                         "start_in_chat": True})).json()
    response = await api.post(f"{BASE}/{session['id']}/resolve",
                              json={"version": session["version"], "automatic": True})
    if failure == "all":
        assert response.status_code >= 400 and calls == ["SEGMENTO 1 D", "SEGMENTO 2 D"]
        stored = (await api.get(f"{BASE}/{session['id']}")).json()
        assert stored["last_generation"] is None
        return
    session = response.json()
    audit = session["global_compilation"]
    if failure == "monolithic_runtime":
        assert calls == [""]  # One call, no segment instruction.
        return
    assert calls == ["SEGMENTO 1 D", "SEGMENTO 2 D"]  # No retry of the failed segment.
    with connection() as db:
        data = json.loads(sessions._row(db, "alpha", session["id"])["state_json"])
    segments = data["global_compilation"]["audit"]["generation_segments"]
    assert [s["status"] for s in segments["segments"]] == ["analyzed", "failed"]
    ids, _, _ = dialogue.export_selection(session)
    assert ids and set(ids).isdisjoint(segments["unanalyzed_candidate_ids"])
    draft = await api.post(f"{BASE}/{session['id']}/global-draft", json={
        "version": session["version"], "candidate_ids": ids, "allow_unresolved": True,
        "confirmed_condition_ids": [], "user_quote": "Genera la bozza"})
    assert draft.status_code == 200, draft.text
    assert audit["proposal_count"] >= 1


def test_negated_coordination_is_not_an_alternative_but_a_doubt_is():
    assert not value_alternatives("Confermo, senza raggruppamenti o consorzi", "Confermo")
    assert value_alternatives("Bari o Lecce", "Bari")
    assert value_alternatives("forse Bari", "Bari")
    assert hedged("Sì o forse no") and not hedged("senza raggruppamenti o consorzi")


def previous(session, *flags):
    with connection() as db:
        db.execute("BEGIN IMMEDIATE")
        row, state = adapter._row_state(db, "alpha", session["id"], session["version"])
        for flag in flags:
            state["global_compilation"].pop(flag, None)
        return sessions.save_state(db, row, state, "test_previous_revision")


async def conditional(api, monkeypatch, *, legacy=False):
    async def request(client, settings, body, timeout, **kwargs):
        data = json.loads(body["messages"][1]["content"])
        return json.dumps(proposal(data, conditional=True)), "simulated", {}

    monkeypatch.setattr(adapter, "request_model_content", request)
    session = await start(api)
    if legacy:
        session = previous(session, "dialogue_revision")
    return await resolve(api, session)


@pytest.mark.anyio
@pytest.mark.parametrize("legacy", [True, False])
async def test_guard_checked_condition_accepts_natural_negated_coordination(
    api, provider, monkeypatch, legacy,
):
    s = await conditional(api, monkeypatch, legacy=legacy)
    message = "Confermo la partecipazione singola, senza raggruppamenti o consorzi."
    s, errors = dialogue.apply_replies("alpha", s["id"], s["version"], [
        ClarificationReply(slot=1, action="CONDITION_TRUE", user_quote=message)], message)
    answers = s["global_compilation"]["dialogue"]["answers"]
    if legacy:
        assert errors == [1] and not answers
    else:
        assert not errors and answers["condition:c"]["applies"] is True
        assert answers["condition:c"]["ambiguity_check"] == "SEMANTIC_GUARD"
        assert not s["global_compilation"]["dialogue"]["user_values"]


@pytest.mark.anyio
@pytest.mark.parametrize("legacy", [True, False])
async def test_hold_on_a_condition_is_released_once_that_condition_is_answered(
    api, provider, monkeypatch, legacy,
):
    stage = {"n": 0}

    async def narrator(client, settings, body, timeout, **kwargs):
        context = json.loads(body["messages"][1]["content"])
        stage["n"] += 1
        if stage["n"] > 1:
            return json.dumps({"questions": []}), "simulated", {}
        condition = next(q["slot"] for q in context["question_pool"] if q["kind"] ==
                         "applicability")
        questions = [{"slot": condition, "confirmation": "la sezione si applica"}]
        hold = {"blocker": "applicability", "reason": "Dipende dalla sezione",
                "evidence_quote": "Sede legale", "condition_slot": condition}
        if legacy:
            return json.dumps({"message": "Ho trovato i dati della società.",
                               "questions": questions, "held_proposal_slots": [2],
                               "hold_reviews": [{"proposal_slot": 2, **hold}]}), "simulated", {}
        return json.dumps({"message": "Ho trovato i dati della società.",
                           "questions": questions,
                           "hold_groups": [{"proposal_slots": [2], **hold}]}), "simulated", {}

    s = await conditional(api, monkeypatch, legacy=legacy)
    monkeypatch.setattr(dialogue, "request_model_content", narrator)
    s = previous(s)
    with connection() as db:
        db.execute("BEGIN IMMEDIATE")
        row, state = adapter._row_state(db, "alpha", s["id"], s["version"])
        state["global_compilation"]["dialogue"] = None
        s = sessions.save_state(db, row, state, "test_new_review")
    s = await resolve(api, s)
    assert s["global_compilation"]["dialogue"]["held_candidate_ids"] == ["t0.r1.c1"]
    message = "Sì, la sezione si applica."
    s, errors = dialogue.apply_replies("alpha", s["id"], s["version"], [
        ClarificationReply(slot=1, action="CONDITION_TRUE", user_quote=message)], message)
    assert not errors
    s = await resolve(api, s)
    held = s["global_compilation"]["dialogue"]["held_candidate_ids"]
    ids = dialogue.export_selection(s)[0]
    if legacy:
        assert held == ["t0.r1.c1"] and ids == ["t0.r0.c1"]  # Previous sticky behaviour.
    else:
        assert held == [] and sorted(ids) == ["t0.r0.c1", "t0.r1.c1"]
        transition = s["global_compilation"]["dialogue"]["hold_transition"]
        assert transition["released_blocker_answered"] == ["t0.r1.c1"]


@pytest.mark.anyio
async def test_backend_asks_open_conditions_when_the_narrator_asks_nothing(
    api, provider, monkeypatch,
):
    async def silent(client, settings, body, timeout, **kwargs):
        return json.dumps({"message": "Ho trovato i dati della società."}), "simulated", {}

    s = await conditional(api, monkeypatch)
    monkeypatch.setattr(dialogue, "request_model_content", silent)
    with connection() as db:
        db.execute("BEGIN IMMEDIATE")
        row, state = adapter._row_state(db, "alpha", s["id"], s["version"])
        state["global_compilation"]["dialogue"] = None
        s = sessions.save_state(db, row, state, "test_new_review")
    s = await resolve(api, s)
    slots = s["chat"]["question"]["slots"]
    assert [q["key"] for q in slots] == ["condition:c"]
    assert slots[0]["question"].startswith("Confermi che si applica al vostro caso")


def test_grouped_holds_report_repeated_and_missing_proposals_apart():
    plan = dialogue.GroupedDialoguePlan(message="Riepilogo dei dati trovati.", hold_groups=[
        {"proposal_slots": [1, 2], "blocker": "form", "reason": "Dubbio",
         "evidence_quote": "Sede legale"},
        {"proposal_slots": [2, 5], "blocker": "source", "reason": "Altro dubbio",
         "evidence_quote": "Sede legale"}])
    expanded, irregular = dialogue.expanded_plan(plan, dialogue.DialoguePlan, 2)
    assert expanded.held_proposal_slots == [1, 2]
    assert [r.blocker for r in expanded.hold_reviews] == ["form", "form"]
    assert irregular["repeated_references"][2][0]["reason"] == "Altro dubbio"
    assert [r["proposal_slot"] for r in irregular["invalid_references"]] == [5]


def test_registry_never_answers_authorizations_values_or_choices():
    pool = [{"slot": 1, "kind": "applicability", "key": "authorization:x", "authorization": {},
             "label": "Autorizzazione", "condition": "☐ dichiaro", "field_ids": ["a"]},
            {"slot": 2, "kind": "applicability", "key": "condition:c", "label": "Ramo",
             "condition": "Dati aziendali", "field_ids": ["b"]},
            {"slot": 3, "kind": "value", "key": "field:z", "label": "Data",
             "condition": "", "field_ids": ["z"]},
            {"slot": 4, "kind": "choice", "key": "choice:g", "label": "Scelta",
             "condition": "", "field_ids": ["y"]}]
    assert [q["key"] for q in registry.open_slots(pool)] == ["condition:c"]


REQUEST = ("Compila la domanda per Aurora S.r.l. come società singola, senza raggruppamenti o "
           "consorzi. Non compilare dichiarazioni non confermate.")


@pytest.mark.anyio
@pytest.mark.parametrize("case", ["accepted", "not_literal", "guard_unrelated",
                                  "provider_failure"])
async def test_initial_request_registers_only_literal_guarded_branch_choices(
    api, provider, monkeypatch, case,
):
    from app import compilation_user_context

    monkeypatch.setattr(compilation_user_context, "registered_user_context", lambda *a: [
        {"kind": "recorded_request", "provenance": "USER", "turn_id": 1, "user_quote": REQUEST}])
    quote = "come società singola, senza raggruppamenti o consorzi"

    async def model(client, settings, body, timeout, **kwargs):
        system = body["messages"][0]["content"]
        if system.startswith("Leggi il messaggio iniziale"):
            if case == "provider_failure":
                raise GenerationTruncatedError(768, 768)
            slots = json.loads(body["messages"][1]["content"])["slots"]
            return json.dumps({"answers": [{"slot": slots[0]["slot"], "action": "CONDITION_TRUE",
                "user_quote": quote if case != "not_literal" else "partecipo da sola"}]}), \
                "simulated", {}
        result = "UNRELATED" if case == "guard_unrelated" else "SUPPORTED"
        return json.dumps({"checks": [{"slot": 1, "result": result}]}), "simulated", {}

    monkeypatch.setattr(registry, "request_model_content", model)
    s = await conditional(api, monkeypatch)
    answers = s["global_compilation"]["dialogue"].get("answers", {})
    if case != "accepted":
        assert not answers
        assert [q["key"] for q in s["chat"]["question"]["slots"]] == ["condition:c"]
        return
    receipt = answers["condition:c"]
    assert (receipt["origin"], receipt["provenance"], receipt["applies"]) == (
        "initial_request", "USER", True)
    assert receipt["user_quote"] in REQUEST
    message = s["global_compilation"]["dialogue"]["message"]
    assert "Dalla tua richiesta iniziale ho registrato" in message
    assert not s["global_compilation"]["dialogue"].get("user_values")
    ids, confirmed, _ = dialogue.export_selection(s)
    assert confirmed == ["c"] and sorted(ids) == ["t0.r0.c1", "t0.r1.c1"]
    draft = await api.post(f"{BASE}/{s['id']}/global-draft", json={
        "version": s["version"], "candidate_ids": ids, "allow_unresolved": True,
        "confirmed_condition_ids": confirmed, "user_quote": "Genera la bozza"})
    assert draft.status_code == 200, draft.text
