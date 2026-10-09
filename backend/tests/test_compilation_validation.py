import json
from io import BytesIO

import pytest
from docx import Document

from app import document_compilation as compilation
from app.docx_templates import inspect_docx
from app.generation import GenerationError

COMPANY = "Impresa Esempio S.r.l."
LEGAL_FORM = "Societa a responsabilita limitata"


def layout_for(kind="table"):
    doc = Document()
    doc.add_paragraph("Modulo dimostrativo indipendente dal bando Catanzaro")
    if kind == "paragraph":
        for label, marker in (("Denominazione", "nome"), ("Forma", "forma"), ("Firma", "firma")):
            doc.add_paragraph(f"{label}: {{{{{marker}}}}}")
    else:
        if kind == "shifted":
            extra = doc.add_table(rows=1, cols=1)
            extra.cell(0, 0).text = "Tabella descrittiva non compilabile"
        table = doc.add_table(rows=3, cols=2)
        for row, label in zip(table.rows, ("Denominazione", "Forma", "Firma"), strict=True):
            row.cells[0].text = label
    output = BytesIO()
    doc.save(output)
    return inspect_docx(output.getvalue())


def context():
    content = f"Denominazione: {COMPANY}. Forma giuridica: {LEGAL_FORM}."
    return compilation.CompilationSources([{
        "id": "company:77", "document_id": 77, "source_name": "impresa.txt",
        "chunk_index": 0, "content": content, "scope": "company", "source_kind": "source",
    }], 1, len(content))


def proposal(target, *, label="Denominazione", value=COMPANY, **changes):
    return {
        "cell_id": target, "label": label, "entity": "company", "kind": "data",
        "status": "proposed", "value": value,
        "evidence": [{"source_id": "company:77", "quote": context().selected[0]["content"]}],
        "reason": "Dato esplicito nella fonte", **changes,
    }


def response(*fields):
    return json.dumps({"fields": fields, "warnings": []})


def targets(layout):
    return (*layout.cells, *layout.slots)


@pytest.mark.parametrize("status", ["missing", "needs_review", "not_applicable"])
def test_omitted_empty_evidence_is_allowed_only_for_unwritten_fields(status):
    layout = layout_for()
    raw = proposal(targets(layout)[0], status=status, value=None)
    raw.pop("evidence")
    field = compilation.validate_proposals(response(raw), layout, context())["fields"][0]
    assert field["status"] == status
    assert field["written_value"] is None
    assert field["evidence"] == []


@pytest.mark.parametrize("status", ["missing", "needs_review", "not_applicable"])
@pytest.mark.parametrize("label", ["", "   "])
def test_unwritten_unlabelled_cell_uses_its_coordinate_without_inventing_meaning(status, label):
    layout = layout_for()
    target = targets(layout)[0]
    raw = proposal(target, label=label, status=status, value=None, evidence=[])
    field = compilation.validate_proposals(response(raw), layout, context())["fields"][0]
    assert field["label"] == target
    assert field["written_value"] is None
    assert raw["label"] == label


@pytest.mark.parametrize("case", ["proposed", "value", "null", "missing", "unknown_id"])
def test_unwritten_label_fallback_does_not_relax_write_or_coordinate_validation(case):
    layout = layout_for()
    raw = proposal(targets(layout)[0], label="", status="needs_review", value=None)
    if case == "proposed":
        raw.update(status="proposed", value=COMPANY)
    elif case == "value":
        raw["value"] = COMPANY
    elif case == "null":
        raw["label"] = None
    elif case == "missing":
        raw.pop("label")
    else:
        raw["cell_id"] = "t999.r0.c0"
    with pytest.raises(GenerationError):
        compilation.validate_proposals(response(raw), layout, context())


@pytest.mark.parametrize("case", [
    "proposed", "value", "null", "string", "object", "missing_status", "missing_value",
])
def test_empty_evidence_default_does_not_relax_malformed_or_writable_proposals(case):
    layout = layout_for()
    raw = proposal(targets(layout)[0], status="needs_review", value=None)
    raw.pop("evidence")
    if case == "proposed":
        raw.update(status="proposed", value=COMPANY)
    elif case == "value":
        raw["value"] = COMPANY
    elif case in ("null", "string", "object"):
        raw["evidence"] = {"null": None, "string": "", "object": {}}[case]
    else:
        raw.pop(case.removeprefix("missing_"))
    with pytest.raises(GenerationError):
        compilation.validate_proposals(response(raw), layout, context())


def test_user_input_is_citable_without_becoming_a_document_or_affecting_coverage():
    layout, sources = layout_for(), context()
    before = sources.coverage()
    entered = "Sottoscrittore: Giulia Bianchi. Partecipazione singola."
    catalog = compilation.source_catalog(sources, entered)
    user = next(s for s in catalog if s["id"] == "user:instructions")
    assert user["origin"] == "user"
    assert user["scope"] == "user"
    assert user["source_kind"] == "user_instructions"
    assert user["document_id"] is None
    assert user["chunk_index"] is None
    assert sources.coverage() == before
    assert "origin" not in sources.selected[0]
    raw = proposal(targets(layout)[0], entity="person", value="Giulia Bianchi", evidence=[{
        "source_id": "user:instructions", "quote": "Sottoscrittore: Giulia Bianchi.",
    }])
    field = compilation.validate_proposals(
        response(raw), layout, sources, instructions=entered,
    )["fields"][0]
    assert field["written_value"] == "Giulia Bianchi"
    assert field["evidence"][0]["origin"] == "user"
    assert field["evidence"][0]["fragment"] is None
    assert field["evidence"][0]["document_id"] is None
    absent = compilation.validate_proposals(response(raw), layout, sources)["fields"][0]
    assert absent["written_value"] is None
    assert "unknown_source" in absent["validation_codes"]


@pytest.mark.parametrize("kind,origin", [
    ("source", "document"), ("project_facts", "user"), ("call_facts", "extracted"),
])
def test_persisted_source_origin_is_explicit(kind, origin):
    sources = context()
    sources.selected[0]["source_kind"] = kind
    assert compilation.source_catalog(sources)[0]["origin"] == origin


@pytest.mark.parametrize("bad_evidence", [
    [{"source_id": "project:999", "quote": COMPANY}],
    [{"source_id": "company:77", "quote": "Sottoscrittore: Giulia Bianchi"}],
])
def test_false_citation_blocks_only_its_field_even_with_another_valid_citation(bad_evidence):
    layout = layout_for()
    first, second, _ = targets(layout)
    good = proposal(first)
    bad = proposal(second, evidence=[*good["evidence"], *bad_evidence])
    report = compilation.validate_proposals(response(good, bad), layout, context())
    assert report["fields"][0]["written_value"] == COMPANY
    blocked = report["fields"][1]
    assert blocked["status"] == "needs_review"
    assert blocked["written_value"] is None
    assert blocked["rejected_evidence"][0]["quote"] == bad_evidence[0]["quote"]
    assert len(blocked["evidence"]) == 1


@pytest.mark.parametrize("failure", ["unknown", "unwritable", "duplicate", "bad_schema"])
def test_all_proposals_undergo_structural_validation(failure):
    layout = layout_for()
    first, second, _ = targets(layout)
    extras = [proposal(second)]
    if failure == "unknown":
        extras = [proposal("t999.r0.c1")]
    elif failure == "unwritable":
        extras = [proposal("t0.r0.c0")]
    elif failure == "duplicate":
        extras.append(proposal(second))
    else:
        extras = [proposal(second, value={"unexpected": "object"})]
    with pytest.raises(GenerationError):
        compilation.validate_proposals(
            response(proposal(first), *extras), layout, context(),
        )


def test_complete_document_keeps_evidence_and_signature_checks():
    layout = layout_for()
    first, second, signature = targets(layout)
    report = compilation.validate_proposals(
        response(proposal(first, evidence=[]), proposal(second), proposal(signature)),
        layout, context(),
    )
    assert len(report["fields"]) == 3
    assert report["fields"][0]["written_value"] is None
    assert report["fields"][1]["written_value"] == COMPANY
    assert report["fields"][2]["written_value"] is None
    assert "value_not_in_quote" in report["fields"][0]["validation_codes"]
    assert "protected_field" in report["fields"][2]["validation_codes"]


@pytest.mark.parametrize("kind", ["table", "shifted", "paragraph"])
@pytest.mark.parametrize("error", ["bad_value", "bad_quote", "unknown_source"])
def test_blocked_proposal_stays_unwritten(kind, error):
    layout = layout_for(kind)
    first, second, _ = targets(layout)
    bad = proposal(second, label="Forma", value=LEGAL_FORM)
    if error == "bad_value":
        bad["value"] = "Forma giuridica inventata"
    elif error == "bad_quote":
        bad["evidence"][0]["quote"] = "Citazione inventata"
    else:
        bad["evidence"][0]["source_id"] = "company:inesistente"
    report = compilation.validate_proposals(
        response(proposal(first), bad), layout, context(),
    )
    assert report["fields"][0]["written_value"] == COMPANY
    assert report["fields"][1]["written_value"] is None
    assert report["fields"][1]["status"] == "needs_review"
    assert report["fields"][1]["validation_codes"]


@pytest.mark.parametrize("kind", ["declaration", "choice", "signature", "disguised_signature"])
def test_user_sources_cannot_enable_protected_fields(kind):
    layout = layout_for()
    target = targets(layout)[2 if kind == "disguised_signature" else 0]
    report = compilation.validate_proposals(
        response(proposal(
            target, kind="data" if kind == "disguised_signature" else kind,
            value="Giulia Bianchi", entity="person", evidence=[{
                "source_id": "user:instructions", "quote": "Giulia Bianchi",
            }],
        )),
        layout, context(), instructions="Giulia Bianchi",
    )
    assert report["fields"][0]["written_value"] is None
    assert report["fields"][0]["validation_codes"] == ["protected_field"]


@pytest.mark.parametrize("original,altered", [
    ("2026-09-17", "2026-09-18"), ("01234567890", "01234567891"), ("1250000", "1520000"),
])
def test_dates_identifiers_and_amounts_are_not_fuzzy_matched(original, altered):
    layout, sources = layout_for(), context()
    sources.selected[0]["content"] = original
    field = compilation.validate_proposals(response(proposal(
        targets(layout)[0], value=altered,
        evidence=[{"source_id": "company:77", "quote": original}],
    )), layout, sources)["fields"][0]
    assert field["written_value"] is None
    assert "value_not_in_quote" in field["validation_codes"]
