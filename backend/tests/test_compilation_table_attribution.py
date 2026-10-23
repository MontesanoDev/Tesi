"""Table subject attribution: controlled proposals, no new model inference."""

import json
from copy import deepcopy
from io import BytesIO
from pathlib import Path

import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from app.compilation_attribution import check_source_target
from app.compilation_catalog import build_compilation_catalog
from app.compilation_global import validate_global_proposal
from app.docx_templates import inspect_docx

FIXTURE = Path(__file__).parent / "fixtures"


@pytest.mark.parametrize("cid,fid", [
    ("t0.r0.c1", "f_company_name"), ("t0.r2.c1", "f_cf"),
    ("t26.r0.c1", "f_company_name"), ("t26.r1.c1", "f_forma"),
])
def test_real_form_and_sources_reject_injected_company_fact_in_personal_table_field(cid, fid):
    case = json.loads((FIXTURE / "source-attribution-v1/catanzaro-real.json").read_text())
    layout = inspect_docx((FIXTURE / "formtext-v1/catanzaro/original.docx").read_bytes())
    catalog = build_compilation_catalog(layout)
    tables = {t["id"]: t for b in catalog["blocks"] for t in b.get("tables", [])}
    paragraphs = {p["id"]: p["text"] for b in catalog["blocks"] for p in b["paragraphs"]}
    c = catalog["candidates"][cid]
    row = tables[c["location"]["table_id"]]["rows"][c["location"]["row_index"]]
    pid = row[0]["paragraph_ids"][0]
    proposal = deepcopy(case["proposal"])
    proposal["fields"][cid] = {
        "label": paragraphs[pid], "kind": "data", "disposition": "conditional",
        "form": {"paragraph_id": pid, "quote": paragraphs[pid]},
        "conditions": ["c_partecipante"] + (["c_soc_ing"] if cid.startswith("t26.") else []),
        "property": proposal["facts"][fid]["property"], "fact_id": fid,
        "reason": "Applicabilità distinta dal soggetto della proprietà",
    }
    result = validate_global_proposal(json.dumps(proposal), layout,
                                      case["project_id"], case["sources"])
    field = next(f for f in result["fields"] if f["candidate_id"] == cid)
    assert result["facts"][fid]["validation"] == "GROUNDED_PROPOSAL"
    assert field["validation"] == "REJECTED"
    assert not field.get("value") and "proposal_grounding" not in field
    check = field["source_target_check"]
    assert check["expected_entity"] == "person"
    assert "source_subject_type_mismatch" in check["validation_codes"]
    assert check["form_evidence"]["format"] == "table_label_value"
    assert field["source_evidence"]


def table_case(labels=None, *, entity="company", target_row=1, binding=False):
    doc = Document()
    table = doc.add_table(rows=0, cols=2)
    for label in labels or ["Firmatario", "PEC", "Società", "PEC"]:
        row = table.add_row()
        row.cells[0].text = label
    # Existing role rules require a real repeated header, not an inferred one.
    table.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
    stream = BytesIO()
    doc.save(stream)
    layout = inspect_docx(stream.getvalue(), require_candidates=False)
    catalog = build_compilation_catalog(layout)
    cid = f"t0.r{target_row}.c1"
    paragraphs = {p["id"]: p["text"] for b in catalog["blocks"] for p in b["paragraphs"]}
    meta = next(t for b in catalog["blocks"] for t in b.get("tables", []))
    pid = meta["rows"][target_row][0]["paragraph_ids"][0]
    role_pid = meta["rows"][0][0]["paragraph_ids"][0]
    content = ("Firmatario: Ada Neri; PEC personale: persona@alba.demo." if entity == "person"
               else "Impresa Alba; PEC aziendale: ufficio@alba.demo.")
    value = "persona@alba.demo" if entity == "person" else "ufficio@alba.demo"
    source = {"chunk_id": -7, "file_id": -3, "chunk_index": 0, "source_name": "profili.txt",
              "role": "source", "scope": "global", "project_id": None,
              "category": "company", "content": content}
    proposal = {"version": 12, "uncertain_paragraph_ids": [], "coverage": [],
                "choice_groups": {}, "clarifications": {}, "conditions": {}, "bindings": {},
                "facts": {"f": {"entity": entity, "subject": "Ada Neri" if entity == "person"
                                 else "Alba", "aspect": "identity", "property": "pec",
                                 "value": value, "evidence": [{"chunk_id": -7, "quote": content}]}},
                "fields": {cid: {"label": "Etichetta MODEL generica", "kind": "data",
                                 "disposition": "value", "conditions": [], "property": "pec",
                                 "fact_id": "f", "form": {"paragraph_id": pid,
                                                          "quote": paragraphs[pid]}}}}
    if binding:
        proposal["fields"][cid]["binding_id"] = "b"
        proposal["bindings"]["b"] = {
            "form": {"paragraph_id": role_pid, "quote": "Firmatario"},
            "required_role": "Firmatario", "source_role": "Firmatario",
            "evidence": [{"chunk_id": -7, "quote": content}]}
    return layout, [source], proposal, cid


def audit(data):
    return validate_global_proposal(json.dumps(data[2]), data[0], "alpha", data[1])


@pytest.mark.parametrize("row,entity,binding,accepted", [
    (1, "company", False, False), (1, "person", True, True),
    (3, "company", False, True), (3, "person", True, False),
])
def test_mixed_person_company_table_does_not_transfer_properties(row, entity, binding, accepted):
    data = table_case(entity=entity, target_row=row, binding=binding)
    field = next(f for f in audit(data)["fields"] if f["candidate_id"] == data[3])
    assert (field["validation"] == "GROUNDED_PROPOSAL") is accepted
    if accepted:
        assert field["value"] == data[2]["facts"]["f"]["value"]
        assert "source_target_check" not in field  # Existing review fingerprint unchanged.
        if entity == "person":
            # The table veto must not turn an unmarked row into a valid role citation.
            tr = data[0].document.tables[0].rows[0]._tr
            tr.get_or_add_trPr().remove(tr.find(f"{qn('w:trPr')}/{qn('w:tblHeader')}"))
            stream = BytesIO()
            data[0].document.save(stream)
            layout = inspect_docx(stream.getvalue(), require_candidates=False)
            other = next(f for f in audit((layout, *data[1:]))["fields"]
                         if f["candidate_id"] == data[3])
            assert other["subject_binding"]["result"] == "review_required"
            assert not other.get("value") and "proposal_grounding" not in other
    else:
        assert "source_subject_type_mismatch" in field["source_target_check"]["validation_codes"]


@pytest.mark.parametrize("change", ["missing_binding", "different_person", "different_role",
                                     "invented_quote"])
def test_personal_table_compatibility_cannot_relax_existing_role_and_source_checks(change):
    data = table_case(entity="person", binding=True)
    if change == "missing_binding":
        data[2]["fields"][data[3]].pop("binding_id")
    elif change == "different_person":
        data[2]["facts"]["f"]["subject"] = "Bea Mori"
    elif change == "different_role":
        data[2]["bindings"]["b"]["source_role"] = "Direttore tecnico"
    else:
        data[2]["facts"]["f"]["evidence"][0]["quote"] = "Inventata: persona@alba.demo"
    field = next(f for f in audit(data)["fields"] if f["candidate_id"] == data[3])
    assert not field.get("value") and "proposal_grounding" not in field


@pytest.mark.parametrize("labels", [
    ["Firmatario", "", "PEC"],
    ["Firmatario", "Altro soggetto", "PEC"],
    ["Firmatario", "Dichiarazioni e attestazioni", "PEC"],
])
def test_subject_is_not_carried_across_empty_or_unrecognised_rows(labels):
    data = table_case(labels, target_row=2)
    result = audit(data)
    check = check_source_target(data[3], build_compilation_catalog(data[0]), result["facts"]["f"])
    assert check["result"] == "not_determined"


def test_subject_is_not_borrowed_from_another_table():
    data = table_case(["PEC"], target_row=0)
    document = data[0].document
    other = document.add_table(rows=1, cols=2)
    other.cell(0, 0).text = "Firmatario"
    stream = BytesIO()
    document.save(stream)
    data = inspect_docx(stream.getvalue(), require_candidates=False), *data[1:]
    result = audit(data)
    check = check_source_target(data[3], build_compilation_catalog(data[0]), result["facts"]["f"])
    assert check["result"] == "not_determined"


def test_location_qualification_in_table_is_not_replaced_by_source_residence():
    data = table_case(["Firmatario", "Comune studio professionale"], entity="person", binding=True)
    content = "Firmatario: Ada Neri; Comune di residenza: Torino."
    data[1][0]["content"] = content
    data[2]["facts"]["f"].update(property="comune", value="Torino",
                                 evidence=[{"chunk_id": -7, "quote": content}])
    data[2]["fields"][data[3]]["property"] = "comune"
    data[2]["bindings"]["b"]["evidence"] = [{"chunk_id": -7, "quote": content}]
    field = next(f for f in audit(data)["fields"] if f["candidate_id"] == data[3])
    assert field["validation"] == "REJECTED"
    assert "source_property_scope_mismatch" in field["source_target_check"]["validation_codes"]


def test_vertical_merge_stops_subject_inheritance():
    data = table_case(["Firmatario", "PEC"])
    table = data[0].document.tables[0]
    merge = OxmlElement("w:vMerge")
    merge.set(qn("w:val"), "restart")
    table.cell(0, 0)._tc.get_or_add_tcPr().append(merge)
    stream = BytesIO()
    data[0].document.save(stream)
    layout = inspect_docx(stream.getvalue(), require_candidates=False)
    check = check_source_target(data[3], build_compilation_catalog(layout),
                                {"entity": "company", "subject": "Alba"})
    assert check["result"] == "not_determined"


def test_bare_economic_operator_is_a_boundary_not_proof_of_company_nature():
    data = table_case(["Firmatario", "PEC", "Operatore economico", "PEC"], target_row=3)
    check = check_source_target(data[3], build_compilation_catalog(data[0]),
                                {"entity": "company", "subject": "Alba"})
    assert check["result"] == "not_determined"


def test_nested_table_is_not_an_inherited_subject_label():
    data = table_case(["Firmatario", "PEC", "PEC"], target_row=2)
    data[0].document.tables[0].cell(1, 0).add_table(rows=1, cols=1).cell(0, 0).text = "Società"
    stream = BytesIO()
    data[0].document.save(stream)
    layout = inspect_docx(stream.getvalue(), require_candidates=False)
    check = check_source_target(data[3], build_compilation_catalog(layout),
                                {"entity": "company", "subject": "Alba"})
    assert check["result"] == "not_determined"


def test_editable_field_in_label_cell_is_not_a_subject_anchor():
    data = table_case(["Firmatario ________", "PEC"])
    check = check_source_target(data[3], build_compilation_catalog(data[0]),
                                {"entity": "company", "subject": "Alba"})
    assert check["result"] == "not_determined"


def test_multiple_record_columns_are_not_inferred_as_a_vertical_label_value_table():
    doc = Document()
    table = doc.add_table(rows=2, cols=4)
    for cell, label in zip(table.rows[0].cells, ["Nome e cognome", "Nascita", "Residenza", "CF"]):
        cell.text = label
    stream = BytesIO()
    doc.save(stream)
    layout = inspect_docx(stream.getvalue(), require_candidates=False)
    check = check_source_target("t0.r1.c3", build_compilation_catalog(layout),
                                {"entity": "company", "subject": "Alba"})
    assert check["result"] == "not_determined"


def test_valid_personal_table_proof_does_not_make_branch_applicable():
    data = table_case(entity="person", binding=True)
    pid = data[2]["fields"][data[3]]["form"]["paragraph_id"]
    data[2]["conditions"]["c"] = {
        "form": {"paragraph_id": pid, "quote": "PEC"}, "aspect": "participation",
        "operator": "semantic", "fact_ids": [],
        "scope": [{"kind": "candidate", "id": data[3]}], "reason": "Decisione da acquisire"}
    data[2]["fields"][data[3]].update(disposition="conditional", conditions=["c"],
                                      reason="Ramo non confermato")
    field = next(f for f in audit(data)["fields"] if f["candidate_id"] == data[3])
    assert field["proposal_grounding"] == "GROUNDED_PROPOSAL"
    assert field["condition_checks"]["c"]["result"] == "unknown"
    assert field["value"] is field["written_value"] is None
