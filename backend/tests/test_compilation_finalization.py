"""Real SQLite lock/rollback regressions with isolated storage and simulated AI."""

import sqlite3
from contextlib import closing, contextmanager
from io import BytesIO

import pytest
from docx import Document
from test_compilation_sessions import BASE, create, resolve, simulate, source
from test_compilation_sessions import anyio_backend as anyio_backend
from test_compilation_sessions import api as api

from app import compilation_sessions as sessions
from app import document_compilation_routes as routes
from app.db import connection, get_db_path, get_storage_path


def exclusive_writer(monkeypatch, change_source=None):
    """Trigger the same cache-spill lock as a large report, with a tiny cache."""
    original_persist = routes._persist
    probes = []

    @contextmanager
    def small_cache():
        with connection() as db:
            assert db.execute("PRAGMA journal_mode").fetchone()[0] == "delete"
            db.execute("PRAGMA cache_size = 1")
            db.execute("PRAGMA cache_spill = ON")
            yield db

    def persist(*args, on_saved):
        def save_under_lock(db, generated):
            # The callback reads the session before rechecking its SOURCE. This
            # additional page read can evict dirty report pages from a tiny cache.
            db.execute("SELECT state_json FROM compilation_sessions").fetchall()
            # Prove another connection really cannot read before the callback.
            with closing(sqlite3.connect(get_db_path(), timeout=0)) as reader:
                with pytest.raises(sqlite3.OperationalError, match="database is locked"):
                    reader.execute("SELECT id FROM projects").fetchall()
            probes.append(True)
            if change_source:
                change_source(db)
            return on_saved(db, generated)

        return original_persist(*args, on_saved=save_under_lock)

    monkeypatch.setattr(routes, "connection", small_cache)
    monkeypatch.setattr(sessions, "_persist", persist)
    return probes


async def ready_session(api, monkeypatch, source_scope):
    state = await create(api)
    await source(
        api, "Denominazione sociale: Aurora Progetti S.r.l.",
        project="alpha" if source_scope == "project" else None,
    )
    simulate(monkeypatch, {"t0.r0.c1": ["Aurora Progetti S.r.l."]})
    state = await resolve(api, state)
    response = await api.patch(
        f"{BASE}/{state['id']}/fields",
        json={"version": state["version"], "fields": [
            {"field_id": "t0.r1.c1", "value": "Via delle Rose 12, Firenze"},
        ]},
    )
    assert response.status_code == 200, response.text
    state = response.json()
    assert state["status"] == "READY"
    return state


@pytest.mark.anyio
@pytest.mark.parametrize("source_scope", ["global", "project"])
async def test_finalize_with_source_survives_exclusive_report_write(api, monkeypatch, source_scope):
    state = await ready_session(api, monkeypatch, source_scope)
    probes = exclusive_writer(monkeypatch)
    response = await api.post(
        f"{BASE}/{state['id']}/finalize", json={"version": state["version"]},
    )
    assert response.status_code == 200, response.text
    generated = response.json()
    assert generated["status"] == "GENERATED" and probes == [True]
    assert generated["version"] == state["version"] + 1
    assert generated["fields"] == state["fields"]
    downloads = generated["last_generation"]["downloads"]
    draft = await api.get(downloads["docx"])
    document = Document(BytesIO(draft.content))
    assert document.tables[0].cell(0, 1).text == "Aurora Progetti S.r.l."
    assert document.tables[0].cell(1, 1).text == "Via delle Rose 12, Firenze"
    report = (await api.get(downloads["report"])).json()
    assert report["fields"][0]["evidence"][0]["origin"] == "document"
    assert report["session_fields"][0]["source_evidence"][0]["role"] == "source"
    assert report["fields"][1]["evidence"][0]["origin"] == "user"
    with connection() as db:
        assert db.execute("SELECT COUNT(*) FROM document_compilations").fetchone()[0] == 1


@pytest.mark.anyio
@pytest.mark.parametrize("source_scope", ["global", "project"])
async def test_source_changed_in_save_transaction_rejects_and_rolls_back(
    api, monkeypatch, source_scope,
):
    state = await ready_session(api, monkeypatch, source_scope)
    evidence = state["fields"][0]["source_evidence"][0]
    table = "global_document_chunks" if source_scope == "global" else "document_chunks"
    chunk_id = abs(evidence["chunk_id"])

    def change_source(db):
        db.execute(f"UPDATE {table} SET content='Fonte cambiata' WHERE id=?", (chunk_id,))

    probes = exclusive_writer(monkeypatch, change_source)
    response = await api.post(
        f"{BASE}/{state['id']}/finalize", json={"version": state["version"]},
    )
    assert response.status_code == 409, response.text
    assert "SOURCE" in response.json()["detail"] and probes == [True]
    current = (await api.get(f"{BASE}/{state['id']}")).json()
    assert current == state
    assert current["last_generation"] is None
    with connection() as db:
        assert db.execute("SELECT COUNT(*) FROM document_compilations").fetchone()[0] == 0
        assert db.execute(
            f"SELECT content FROM {table} WHERE id=?", (chunk_id,),
        ).fetchone()[0] == evidence["content"]
    assert not list(get_storage_path().rglob("bozza.docx"))
    assert not list(get_storage_path().rglob("report.json"))
