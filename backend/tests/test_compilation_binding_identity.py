"""SOURCE binding recoveries: direct proofs and bounded identity, no provider."""

import json
from io import BytesIO

import pytest
from docx import Document

from app.compilation_global import validate_global_proposal
from app.docx_templates import inspect_docx

TITLE = "## Soci della società"
ORDER = "Ordine degli Ingegneri di Torino"
HEADER = ("| Persona | Qualifica professionale | Albo | Sezione | Provincia | "
          "Iscrizione dal | N. iscrizione | Codice fiscale |")
ROW = f"| Ada Neri | Ingegnere | {ORDER} | A | Torino (TO) | 09/02/2017 | 6234 | NREADA85A41L219R |"
TABLE = f"{TITLE}\n{HEADER}\n| --- | --- | --- | --- | --- | --- | --- | --- |\n{ROW}"
DIRECTOR = (f"### Ing. Ada Neri - Direttore tecnico\n"
            f"- Ordine professionale: {ORDER}.\n- Numero iscrizione: 6234.")


def source(content, chunk):
    return {"chunk_id": chunk, "file_id": chunk - 20, "chunk_index": 0,
            "source_name": "anagrafica.md", "role": "source", "scope": "global",
            "category": "company", "project_id": None, "content": content}


def case(value="Ingegnere", prop="qualifica_professionale", quote=None):
    quote = quote or f"- Dato: {value}."
    content = f"{DIRECTOR}\n{quote}"
    doc = Document()
    doc.add_paragraph("Socio")
    doc.add_paragraph("Dato: ________")
    stream = BytesIO()
    doc.save(stream)
    layout = inspect_docx(stream.getvalue(), require_candidates=False)
    proposal = {
        "version": 12, "uncertain_paragraph_ids": [], "choice_groups": {},
        "clarifications": {}, "coverage": [],
        "conditions": {"c": {"form": {"paragraph_id": "paragraph:0", "quote": "Socio"},
                             "aspect": "participation", "operator": "semantic", "fact_ids": [],
                             "scope": [{"kind": "candidate", "id": "p1.s0"}],
                             "reason": "Scelta della pratica non registrata"}},
        "bindings": {"b": {
            "form": {"paragraph_id": "paragraph:0", "quote": "Socio"},
            "required_role": "Socio", "source_role": "Socio",
            "evidence": [{"chunk_id": -7, "quote": TITLE}, {"chunk_id": -7, "quote": ROW}]}},
        "facts": {"f": {"subject": "Ada Neri", "entity": "person", "aspect": "professional",
                          "property": prop, "value": value,
                          "evidence": [{"chunk_id": -9, "quote": quote}]}},
        "fields": {"p1.s0": {
            "label": "Dato", "kind": "data", "disposition": "conditional", "conditions": ["c"],
            "form": {"paragraph_id": "paragraph:1", "quote": "Dato: [[p1.s0]]"},
            "fact_id": "f", "property": prop, "binding_id": "b", "reason": "Ramo sconosciuto"}},
    }
    return layout, [source(TABLE, -7), source(content, -9)], proposal


def audit(data):
    return validate_global_proposal(json.dumps(data[2]), data[0], "alpha", data[1])


def assert_grounded(data, method):
    result = audit(data)
    field = result["fields"][0]
    assert field["proposal_grounding"] == "GROUNDED_PROPOSAL"
    assert field["proposed_value"] == data[2]["facts"]["f"]["value"]
    check = field["subject_binding"]
    assert check["result"] == "literal_match"
    assert check["recovery"]["method"] == method
    assert check["recovery"]["original_validation_errors"]
    assert field["condition_checks"]["c"]["result"] == "unknown"
    assert field["validation"] == "REVIEW_REQUIRED"
    assert field["value"] is field["written_value"] is None
    assert result["writes"] == result["resolved"] == 0
    assert result["facts"]["f"]["value"] == data[2]["facts"]["f"]["value"]
    for proof in check["source_evidence"]:
        actual = next(s for s in data[1] if s["chunk_id"] == proof["chunk_id"])
        assert proof["quote"] in actual["content"]
        assert proof["content_sha256"]
    return result


def assert_blocked(data):
    result = audit(data)
    field = result["fields"][0]
    assert "proposal_grounding" not in field
    assert field["subject_binding"]["result"] == "review_required"
    assert field["value"] is field["written_value"] is None
    assert result["writes"] == result["resolved"] == 0
    return result


@pytest.mark.parametrize("value,prop", [
    ("Ingegnere", "qualifica_professionale"), (ORDER, "ordine_professionale"),
    ("A", "sezione"), ("6234", "numero_iscrizione"),
])
def test_binding_row_exact_is_independent_complete_value_proof(value, prop):
    result = assert_grounded(case(value, prop), "binding_source_record")
    check = result["fields"][0]["subject_binding"]
    assert check["value_support"]["comparison"] == "exact"
    assert check["identity_support"]["method"] == "direct_role_record"
    assert {p["chunk_id"] for p in check["source_evidence"]} == {-7, -9}
    assert HEADER in {p["quote"] for p in check["source_evidence"]}


@pytest.mark.parametrize("change", ["other_person", "second_bad_citation", "wrong_column",
                                     "ambiguous_row", "invented_binding_quote", "wrong_subject",
                                     "conflicting_identifier"])
def test_binding_row_exact_cannot_hide_wrong_person_or_unsupported_attribution(change):
    data = case()
    if change == "other_person":
        data[1][-1]["content"] = data[1][-1]["content"].replace("Ada Neri", "Bea Mori")
    elif change == "second_bad_citation":
        data[1].append(source("### Bea Mori - Altro ruolo\n- Dato: Ingegnere.", -11))
        data[2]["facts"]["f"]["evidence"].append({"chunk_id": -11, "quote": "- Dato: Ingegnere."})
    elif change == "wrong_column":
        data[2]["facts"]["f"].update(property="provincia", value="A")
        data[2]["fields"]["p1.s0"]["property"] = "provincia"
        data[1][-1]["content"] += "\n- Dato: A."
        data[2]["facts"]["f"]["evidence"] = [{"chunk_id": -9, "quote": "- Dato: A."}]
    elif change == "ambiguous_row":
        data[1][0]["content"] += "\n" + ROW.replace("6234", "7777")
    elif change == "invented_binding_quote":
        data[2]["bindings"]["b"]["evidence"][0]["quote"] = "## Titolo inventato"
    elif change == "conflicting_identifier":
        data[1][-1]["content"] += "\n- Codice fiscale: NREADA86A41L219R."
    else:
        data[2]["facts"]["f"]["subject"] = "Bea Mori"
    assert_blocked(data)


@pytest.mark.parametrize("value,prop", [
    ("Ing. Ada Neri", "nome"), ("Arch. Ada Neri", "nome"),
    ("9 febbraio 2017", "iscrizione_dal"), ("2017-02-09", "iscrizione_dal"),
])
def test_binding_row_equivalent_value_never_merges_other_attributes(value, prop):
    data = case(value, prop)
    if value.startswith("Arch."):
        data[1][0]["content"] = TABLE.replace("Ingegnere", "Architetto")
        data[2]["bindings"]["b"]["evidence"][-1]["quote"] = ROW.replace("Ingegnere", "Architetto")
        data[1][-1]["content"] = data[1][-1]["content"].replace("Ing. Ada", "Arch. Ada")
    result = assert_grounded(data, "binding_source_record")
    support = result["fields"][0]["subject_binding"]
    assert support["value_support"]["comparison"] in {"professional_title", "calendar_date"}
    assert support["identity_support"]["method"] == "direct_role_record"
    assert result["facts"]["f"]["value"] == value


@pytest.mark.parametrize("value,prop", [
    ("A. Neri", "nome"), ("Ing. Ada Nero", "nome"), ("Ada", "nome"),
    ("31 febbraio 2017", "iscrizione_dal"), ("09/02/17", "iscrizione_dal"),
    ("2 settembre 2017", "iscrizione_dal"), ("9 febbraio 2017", "numero_iscrizione"),
    ("Ing. Ingegnere", "qualifica_professionale"),
])
def test_equivalence_is_not_fuzzy_identity_date_guess_or_qualification(value, prop):
    assert_blocked(case(value, prop))


def identity_case(value="NREADA85A41L219R", prop="codice_fiscale"):
    data = case(value, prop, ROW)
    data[1][0]["content"] = DIRECTOR
    data[1][-1]["content"] = TABLE
    binding = data[2]["bindings"]["b"]
    binding.update(required_role="Direttore tecnico", source_role="Direttore tecnico",
                   evidence=[{"chunk_id": -7, "quote": DIRECTOR.splitlines()[0]}])
    # Keep FORM literal and applicability unresolved.
    document = data[0].document
    document.paragraphs[0].text = "Direttore tecnico"
    stream = BytesIO()
    document.save(stream)
    layout = inspect_docx(stream.getvalue(), require_candidates=False)
    binding["form"]["quote"] = "Direttore tecnico"
    data[2]["conditions"]["c"]["form"]["quote"] = "Direttore tecnico"
    return layout, data[1], data[2]


@pytest.mark.parametrize("value,prop", [
    ("NREADA85A41L219R", "codice_fiscale"), ("Torino (TO)", "provincia"),
])
def test_identity_requires_name_order_and_registration_without_transferring_role(value, prop):
    result = assert_grounded(identity_case(value, prop), "verified_person_records")
    support = result["fields"][0]["subject_binding"]
    assert support["identity_support"]["method"] == "professional_registration"
    assert support["identity_support"]["key"]["order"] == ORDER.casefold()
    assert support["identity_support"]["key"]["number"] == "6234"
    assert support["role_support"]["role"] == "Direttore tecnico"
    assert support["value_support"]["property"] == prop
    assert {p["chunk_id"] for p in support["source_evidence"]} == {-7, -9}
    assert DIRECTOR in {p["quote"] for p in support["source_evidence"]}


def test_identity_can_use_explicit_shared_personal_identifier():
    data = identity_case()
    data[1][0]["content"] = ("### Ing. Ada Neri - Direttore tecnico\n"
                             "- Codice fiscale: NREADA85A41L219R.")
    result = assert_grounded(data, "verified_person_records")
    support = result["fields"][0]["subject_binding"]
    assert support["identity_support"]["method"] == "personal_identifier"


def test_missing_attributes_and_explicit_scope_limit_do_not_negate_the_asserted_role():
    data = identity_case()
    data[1][0]["content"] += (
        "\n- Data di nascita e residenza: non disponibili.\n\n"
        "Il ruolo aziendale non equivale a una designazione per uno specifico incarico.")
    result = assert_grounded(data, "verified_person_records")
    assert result["fields"][0]["condition_checks"]["c"]["result"] == "unknown"


@pytest.mark.parametrize("change", [
    "name_only", "number_only", "order_only", "different_name", "similar_name",
    "different_order", "different_number", "conflicting_cf", "duplicate_homonym",
    "uncited_conflict", "general_identity", "negative_identity", "wrong_person_quote",
    "wrong_role", "power", "qualification", "invalid_cf_quote", "wrong_province_column",
    "same_key_different_person", "scope_note_with_negated_role", "province_without_registration",
    "conflicting_key_in_one_record",
    "conflicting_table_aliases", "role_key_belongs_to_other_person",
    "uncited_named_conflict", "named_key_different_person",
])
def test_identity_join_rejects_homonyms_conflicts_unproved_roles_and_other_people(change):
    data = identity_case()
    if change == "name_only":
        data[1][0]["content"] = DIRECTOR.splitlines()[0]
    elif change == "number_only":
        data[1][0]["content"] = DIRECTOR.replace(f"- Ordine professionale: {ORDER}.\n", "")
    elif change == "order_only":
        data[1][0]["content"] = DIRECTOR.replace("- Numero iscrizione: 6234.", "")
    elif change in {"different_name", "similar_name"}:
        replacement = "Bea Mori" if change == "different_name" else "Ada Nero"
        data[1][-1]["content"] = TABLE.replace("Ada Neri", replacement)
        data[2]["facts"]["f"]["evidence"][0]["quote"] = ROW.replace("Ada Neri", replacement)
    elif change == "different_order":
        data[1][0]["content"] = DIRECTOR.replace("Torino", "Genova")
    elif change == "different_number":
        data[1][0]["content"] = DIRECTOR.replace("6234", "7777")
    elif change == "conflicting_cf":
        data[1][0]["content"] += "\n- Codice fiscale: NREADA86A41L219R."
    elif change == "duplicate_homonym":
        data[1][-1]["content"] += "\n" + ROW.replace("6234", "7777")
    elif change == "uncited_conflict":
        data[1].append(source(DIRECTOR + "\n- Codice fiscale: NREADA86A41L219R.", -11))
    elif change == "general_identity":
        data[1][0]["category"] = "general"
    elif change == "negative_identity":
        data[1][0]["content"] += "\n- Incarico revocato."
    elif change == "wrong_person_quote":
        extra = source(TABLE.replace("Ada Neri", "Bea Mori"), -11)
        data[1].append(extra)
        data[2]["facts"]["f"]["evidence"].append(
            {"chunk_id": -11, "quote": ROW.replace("Ada Neri", "Bea Mori")})
    elif change == "wrong_role":
        data[2]["bindings"]["b"]["required_role"] = "Firmatario"
    elif change in {"power", "qualification"}:
        prop, value = ("potere_firma", "NREADA85A41L219R") if change == "power" else (
            "qualifica_professionale", "Ingegnere")
        data[2]["facts"]["f"].update(property=prop, value=value)
        data[2]["fields"]["p1.s0"]["property"] = prop
    elif change == "wrong_province_column":
        data = identity_case("A", "provincia")
    elif change == "same_key_different_person":
        data[1].append(source(DIRECTOR.replace("Ada Neri", "Bea Mori"), -11))
    elif change == "scope_note_with_negated_role":
        data[1][0]["content"] += (
            "\nIl ruolo aziendale non equivale a una designazione; Ada non è direttore tecnico.")
    elif change == "province_without_registration":
        data = identity_case("Torino (TO)", "provincia")
        data[1][0]["content"] = DIRECTOR.splitlines()[0] + "\n- Codice fiscale: NREADA85A41L219R."
    elif change == "conflicting_key_in_one_record":
        data[1][0]["content"] += "\n- Numero iscrizione: 8888."
    elif change == "conflicting_table_aliases":
        row = ROW.replace(f"| {ORDER} |", f"| {ORDER} | Ordine degli Ingegneri di Genova |")
        data[1][-1]["content"] = TABLE.replace(HEADER, HEADER.replace(
            "| Albo |", "| Albo | Ordine professionale |" )).replace(
            "| --- | --- | --- | --- | --- | --- | --- | --- |",
            "| --- | --- | --- | --- | --- | --- | --- | --- | --- |").replace(ROW, row)
        data[2]["facts"]["f"]["evidence"][0]["quote"] = row
    elif change == "role_key_belongs_to_other_person":
        data[1][0]["content"] = DIRECTOR.replace(
            f"- Ordine professionale: {ORDER}.",
            f"- **Bea Mori:** altra persona.\n- Ordine professionale: {ORDER}.")
    elif change in {"uncited_named_conflict", "named_key_different_person"}:
        name = "Ada Neri" if change == "uncited_named_conflict" else "Bea Mori"
        number = "8888" if change == "uncited_named_conflict" else "6234"
        data[1].append(source(
            f"- **{name}:** direttore tecnico; {ORDER}, iscrizione con numero {number}.", -11))
    else:
        data[2]["facts"]["f"]["evidence"][0]["quote"] = "| Ada Neri | Ingegnere |"
        # The quote does not contain the proposed CF even though the full row does.
        data[1][-1]["content"] += "\n| Ada Neri | Ingegnere |"
    assert_blocked(data)


def test_long_structural_column_cannot_crash_recovery_or_drop_its_proof():
    data = case()
    data[1][0]["content"] = TABLE.replace("Codice fiscale", "Etichetta " + "x" * 900)
    assert_blocked(data)


def test_unrepresentable_role_record_cannot_discard_identity_evidence():
    data = identity_case()
    data[1][0]["content"] += "\n- Annotazione: " + "x" * 900
    assert_blocked(data)


@pytest.mark.parametrize("factory,method", [
    (case, "binding_source_record"), (identity_case, "verified_person_records"),
])
def test_authorized_project_source_keeps_its_scope_and_can_prove_personal_facts(factory, method):
    data = factory()
    for item in data[1]:
        item.update(chunk_id=-item["chunk_id"], file_id=-item["file_id"],
                    scope="project:alpha", project_id="alpha", category=None)
    for section in [data[2]["facts"], data[2]["bindings"]]:
        for item in section.values():
            for proof in item["evidence"]:
                proof["chunk_id"] = -proof["chunk_id"]
    result = assert_grounded(data, method)
    assert all(p["scope"] == "project:alpha" and p["project_id"] == "alpha"
               for p in result["fields"][0]["subject_binding"]["source_evidence"])
    data[1][0].update(scope="project:beta", project_id="beta")
    with pytest.raises(ValueError, match="fuori progetto"):
        audit(data)
