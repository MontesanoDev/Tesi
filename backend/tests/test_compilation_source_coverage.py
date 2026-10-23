"""Registered selections and structural proofs; no provider or fixture IDs."""

import copy
import json
from io import BytesIO

import pytest
from docx import Document
from test_compilation_binding_identity import DIRECTOR, ORDER, ROW, TABLE, identity_case
from test_compilation_selected_person import NAME, PROFILE, case, reused

from app.compilation_global import build_global_input, validate_global_proposal
from app.docx_templates import inspect_docx


def chosen():
    data, layout, sources = case()
    data["dialogue"] = {"answers": {}, "user_values": {}}
    form = {"paragraph_id": "paragraph:0", "quote": "Il sottoscritto"}
    data["audit"]["conditions"]["selection"] = {
        "validation": "GROUNDED_PROPOSAL", "result": "unknown", "form_evidence": form,
        "validation_errors": []}
    data["dialogue"]["answers"]["condition:selection"] = {
        "provenance": "USER", "action": "CONDITION_TRUE", "applies": True,
        "question": "Confermi la partecipazione nella qualità indicata?",
        "user_quote": "Confermo la partecipazione; "
                      f"per questa pratica il firmatario scelto è {NAME}."}
    return data, layout, sources


def test_registered_choice_reuses_nine_source_properties_without_new_user_value():
    from app.compilation_source_reuse import source_reuse_requested

    data, layout, sources = chosen()
    before = copy.deepcopy(data)
    assert source_reuse_requested(data)
    audit = reused(data, layout, sources)
    fields = [f for f in audit["fields"] if f.get("source_reuse")]
    assert len(fields) == 9 and data == before
    name = next(f for f in fields if f["candidate_id"] == "p0.s0")
    assert name["proposed_value"] == NAME and name["source_evidence"]
    assert name["source_reuse"]["selection_receipt"]["action"] == "CONDITION_TRUE"
    assert all(f["written_value"] is None and f["semantic_review_required"] for f in fields)
    assert data["dialogue"]["user_values"] == {}


@pytest.mark.parametrize("quote", [
    "Confermo soltanto la tipologia societaria.",
    f"La fonte dice che il firmatario scelto è {NAME}.",
    f"Non confermo che il firmatario scelto è {NAME}.",
    f"Forse il firmatario scelto è {NAME}.",
    f"Il firmatario scelto è {NAME} oppure Bea Mori.",
    f'Nel documento: "il firmatario scelto è {NAME}".',
    "Per questa pratica il firmatario scelto è Ada.",
])
def test_company_role_narration_negation_or_ambiguous_name_is_not_a_selection(quote):
    data, layout, sources = chosen()
    data["dialogue"]["answers"]["condition:selection"]["user_quote"] = quote
    assert not any(f.get("source_reuse") for f in reused(data, layout, sources)["fields"])


@pytest.mark.parametrize("change", ["unknown", "false", "unregistered", "bad_form",
                                    "homonym", "conflict", "other_person"])
def test_selection_keeps_identity_receipt_and_source_constraints(change):
    data, layout, sources = chosen()
    receipt = data["dialogue"]["answers"]["condition:selection"]
    if change == "unknown":
        receipt.update(action="UNKNOWN", applies=None)
    elif change == "false":
        receipt.update(action="CONDITION_FALSE", applies=False)
    elif change == "unregistered":
        data["dialogue"]["answers"] = {}
        data["user_context"] = [{"user_quote": receipt["user_quote"]}]
    elif change == "bad_form":
        data["audit"]["conditions"]["selection"]["validation_errors"] = ["FORM non valido"]
    elif change == "homonym":
        sources.append({**sources[0], "chunk_id": -8, "file_id": -4,
                        "content": PROFILE.replace("85A52", "86A52")})
    elif change == "conflict":
        data["audit"]["conditions"]["other"] = data["audit"]["conditions"]["selection"]
        data["dialogue"]["answers"]["condition:other"] = {
            **receipt, "user_quote": "Per questa pratica il firmatario scelto è Bea Mori."}
    else:
        sources[0]["content"] = sources[0]["content"].replace(NAME, "Bea Mori")
    assert not any(f.get("source_reuse") for f in reused(data, layout, sources)["fields"])


def test_binding_rejected_proposal_can_be_superseded_only_by_independent_matching_source():
    data, layout, sources = chosen()
    target = data["audit"]["fields"][0]
    data["audit"]["facts"]["biography"] = {
        "validation": "GROUNDED_PROPOSAL", "subject": NAME, "entity": "person",
        "property": "nome", "value": NAME}
    target.update(proposal={"fact_id": "biography", "disposition": "conditional"},
                  validation="REVIEW_REQUIRED",
                  validation_errors=["Binding di ruolo rifiutato"])
    audit = reused(data, layout, sources)
    recovered = next(f for f in audit["fields"] if f["candidate_id"] == target["candidate_id"])
    assert recovered["source_reuse"]["original_validation_errors"] == target["validation_errors"]
    assert recovered["proposal"] == target["proposal"]
    assert audit["unassigned_count"] == data["audit"]["unassigned_count"] - 8
    assert recovered["proposal_grounding"] == "GROUNDED_PROPOSAL"
    for mutation in [{"value": "Bea Mori"}, {"subject": "Bea Mori"},
                     {"validation": "REJECTED"}]:
        bad = copy.deepcopy(data)
        bad["audit"]["facts"]["biography"].update(mutation)
        field = reused(bad, layout, sources)["fields"][0]
        assert not field.get("source_reuse")


def test_explicit_model_alternatives_or_user_override_remain_protected():
    data, layout, sources = chosen()
    target = data["audit"]["fields"][0]
    target.update(proposal={"disposition": "review"}, alternatives=[{"value": "Bea Mori"}])
    assert not reused(data, layout, sources)["fields"][0].get("source_reuse")
    data["dialogue"]["user_values"]["p1.s3"] = {"value": "dato USER"}
    assert not next(f for f in reused(data, layout, sources)["fields"]
                    if f["candidate_id"] == "p1.s3").get("source_reuse")


def test_narrator_can_review_source_reuse_while_retaining_failed_original_binding():
    from app.compilation_global_dialogue import model_context, question_pool

    data, layout, sources = chosen()
    data["audit"]["fields"][0]["subject_binding"] = {
        "result": "review_required", "form_evidence": None}
    data["audit"] = reused(data, layout, sources)
    context = model_context(data, question_pool(data))
    assert len(context["document_values"]) == 10
    assert any(role["subject"] == NAME for role in context["role_contexts"])
    assert data["audit"]["fields"][0]["subject_binding"]["form_evidence"] is None


def test_a_later_direct_signatory_deferral_is_not_replaced_by_an_older_condition_choice():
    data, layout, sources = chosen()
    data["dialogue"]["answers"]["field:p0.s0"] = {
        "action": "SKIP", "provenance": "USER", "user_quote": "Il firmatario lo decidiamo dopo"}
    assert not any(f.get("source_reuse") for f in reused(data, layout, sources)["fields"])


def form_case(*, split=False):
    doc = Document()
    if split:
        doc.add_paragraph("Operatore economico ________ con sede in")
        doc.add_paragraph("________ (________), tel. ________")
        owner, target = "p0.s0", "p1.s0"
        form = {"paragraph_id": "paragraph:1", "quote": "[[p1.s0]] ("}
    else:
        doc.add_paragraph("Codice fiscale: ________")
        owner, target = None, "p0.s0"
        form = {"paragraph_id": "paragraph:0", "quote": "Codice fiscale: [[p0.s0]] iscritto"}
    stream = BytesIO()
    doc.save(stream)
    layout = inspect_docx(stream.getvalue())
    src = {"chunk_id": -7, "file_id": -3, "chunk_index": 0, "source_name": "azienda.md",
           "role": "source", "scope": "global", "project_id": None, "category": "company",
           "content": "Società: Alba S.r.l.; Comune: Torino; Codice: 01234567890."}
    payload = {"version": 12, "fields": {target: {"label": "Comune" if split else "Codice",
        "kind": "data", "disposition": "value", "conditions": [], "fact_id": "value",
        "property": "comune" if split else "codice", "form": form}},
        "facts": {"value": {"subject": "Alba S.r.l.", "entity": "company", "aspect": "identity",
            "property": "comune" if split else "codice", "value": "Torino" if split
            else "01234567890", "evidence": [{"chunk_id": -7, "quote": src["content"]}]}},
        "conditions": {}, "bindings": {}, "clarifications": {}, "choice_groups": {}, "coverage": []}
    if owner:
        payload["facts"]["owner"] = {**payload["facts"]["value"], "property": "nome",
                                      "value": "Alba S.r.l."}
        payload["fields"][owner] = {**payload["fields"][target], "label": "Operatore economico",
            "fact_id": "owner", "property": "nome", "form": {"paragraph_id": "paragraph:0",
                "quote": "Operatore economico [[p0.s0]]"}}
    return layout, [src], payload, target


def form_audit(data):
    return validate_global_proposal(json.dumps(data[2]), data[0], "alpha", data[1])


@pytest.mark.parametrize("split", [False, True])
def test_backend_recovers_only_actual_same_slot_form_proofs(split):
    data = form_case(split=split)
    original = build_global_input(data[0], "alpha", data[1])["form"]
    field = next(f for f in form_audit(data)["fields"] if f["candidate_id"] == data[3])
    assert field["validation"] == "GROUNDED_PROPOSAL"
    assert field["destination_check"]["recovery"]["original_quote"] == data[2]["fields"][
        data[3]]["form"]["quote"]
    assert build_global_input(data[0], "alpha", data[1])["form"] == original
    assert field["written_value"] is None and field["semantic_review_required"]


@pytest.mark.parametrize("change", ["invented_prefix", "wrong_marker", "wrong_paragraph",
                                    "bad_source", "old_contract"])
def test_form_recovery_cannot_fix_false_label_pointer_source_or_historical_contract(change):
    data = form_case()
    proposal = data[2]["fields"][data[3]]
    if change == "invented_prefix":
        proposal["form"]["quote"] = "Nascita [[p0.s0]] iscritto"
    elif change == "wrong_marker":
        proposal["form"]["quote"] = "Codice fiscale: [[p0.s99]] iscritto"
    elif change == "wrong_paragraph":
        proposal["form"]["paragraph_id"] = "paragraph:99"
    elif change == "old_contract":
        data[2]["version"] = 11
        data[2].pop("coverage")
    else:
        data[2]["facts"]["value"]["evidence"][0]["quote"] = "Citazione inventata"
    assert form_audit(data)["fields"][0]["validation"] == "REJECTED"


@pytest.mark.parametrize("change", ["owner", "address", "interposed", "heading", "current_heading",
                                    "not_operator_name"])
def test_adjacent_location_proof_cannot_cross_owner_boundary_or_accept_full_address(change):
    data = form_case(split=True)
    if change == "owner":
        data[2]["facts"]["value"]["subject"] = "Altra azienda"
    elif change == "not_operator_name":
        data[2]["facts"]["owner"]["value"] = "Torino"
    elif change == "address":
        value = "Via Roma 24, Torino (TO)"
        data[1][0]["content"] += " Indirizzo: " + value
        data[2]["facts"]["value"]["value"] = value
        data[2]["facts"]["value"]["evidence"][0]["quote"] = data[1][0]["content"]
    else:
        first = data[0].document.paragraphs[0]
        if change in {"heading", "current_heading"}:
            data[0].document.paragraphs[0 if change == "heading" else 1].style = "Heading 1"
        else:
            first.text = first.text.replace("con sede in", "recapito del referente")
        stream = BytesIO()
        data[0].document.save(stream)
        data = inspect_docx(stream.getvalue()), *data[1:]
    field = next(f for f in form_audit(data)["fields"] if f["candidate_id"] == data[3])
    assert field["validation"] != "GROUNDED_PROPOSAL"


@pytest.mark.parametrize("value,prop", [
    ("Torino (TO)", "provincia_albo"), ("09/02/2017", "data_iscrizione_albo"),
    ("6234", "numero_iscrizione_albo"),
])
def test_professional_attributes_keep_exact_registration_link_and_unknown(value, prop):
    data = identity_case(value, prop)
    field = form_audit(data)["fields"][0]
    assert field["proposal_grounding"] == "GROUNDED_PROPOSAL"
    check = field["subject_binding"]
    assert check["identity_support"]["key"] == {"order": ORDER.casefold(), "number": "6234"}
    assert field["condition_checks"]["c"]["result"] == "unknown"
    assert field["value"] is field["written_value"] is None


def test_adjacent_operator_must_pass_later_role_checks_even_when_catalog_order_differs():
    data = form_case(split=True)
    doc = Document()
    for _ in range(99):
        doc.add_paragraph()
    for paragraph in data[0].document.paragraphs:
        doc.add_paragraph(paragraph.text)
    stream = BytesIO()
    doc.save(stream)
    layout = inspect_docx(stream.getvalue())
    payload = copy.deepcopy(data[2])
    payload["fields"] = {"p99.s0": {**payload["fields"]["p0.s0"], "binding_id": "invalid",
        "form": {"paragraph_id": "paragraph:99", "quote": "Operatore economico [[p99.s0]]"}},
        "p100.s0": {**payload["fields"]["p1.s0"], "form": {
            "paragraph_id": "paragraph:100", "quote": "[[p100.s0]] ("}}}
    payload["bindings"] = {"invalid": {"required_role": "Richiedente", "source_role": "Richiedente",
        "form": {"paragraph_id": "paragraph:99", "quote": "con sede in"},
        "evidence": [{"chunk_id": -7, "quote": data[1][0]["content"]}]}}
    result = form_audit((layout, data[1], payload))
    city = next(f for f in result["fields"] if f["candidate_id"] == "p100.s0")
    assert city["validation"] == "REJECTED" and not city.get("proposal_grounding")


@pytest.mark.parametrize("change", ["name_only", "other_registration", "other_person",
    "conflict", "wrong_column", "power", "qualification"])
def test_registration_attribute_recovery_never_transfers_unproved_roles_or_other_people(change):
    data = identity_case("09/02/2017", "data_iscrizione_albo")
    if change == "name_only":
        data[1][0]["content"] = DIRECTOR.splitlines()[0]
    elif change == "other_registration":
        data[1][0]["content"] = DIRECTOR.replace("6234", "9999")
    elif change == "other_person":
        data[1][-1]["content"] = TABLE.replace("Ada Neri", "Ada Nero")
        data[2]["facts"]["f"]["evidence"][0]["quote"] = ROW.replace("Ada Neri", "Ada Nero")
    elif change == "conflict":
        data[1].append({**data[1][0], "chunk_id": -11,
                       "content": DIRECTOR.replace("6234", "9999")})
    else:
        prop = {"wrong_column": "numero_iscrizione_albo", "power": "poteri_firma",
                "qualification": "qualifica_professionale"}[change]
        data[2]["facts"]["f"]["property"] = prop
        data[2]["fields"]["p1.s0"]["property"] = prop
    assert not form_audit(data)["fields"][0].get("proposal_grounding")
