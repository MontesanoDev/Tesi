"""V12 persistence, export gates and report: isolated API and simulated provider."""

import json
from io import BytesIO

import pytest
from docx import Document
from test_compilation_global_sessions import BASE, proposal, resolve, start
from test_compilation_global_sessions import api as global_api
from test_compilation_global_sessions import provider as global_provider

from app import compilation_global_sessions as adapter

api, provider = global_api, global_provider


@pytest.fixture
def anyio_backend():
    return "asyncio"


@pytest.mark.anyio
async def test_session_retains_reason_but_it_cannot_authorize_an_export(api, provider, monkeypatch):
    calls = []

    async def request(client, settings, body, timeout, **kwargs):
        data = json.loads(body["messages"][1]["content"])
        calls.append(data)
        p = proposal(data)
        del p["fields"]["t0.r1.c1"]
        p["coverage"] = [{"candidate_ids": ["t0.r1.c1"], "outcome": "manual_input",
                          "reason": "Richiesta di conferma della sede da utilizzare",
                          "form": {"paragraph_id": "paragraph:0", "quote": "Dati aziendali"}}]
        return json.dumps(p), "simulated", {}

    monkeypatch.setattr(adapter, "request_model_content", request)
    s = await resolve(api, await start(api))
    assert len(calls) == 1
    summary = s["global_compilation"]["coverage_summary"]
    assert summary["proposed"] == summary["explained_unfilled"] == 1
    assert summary["unexplained"] == 0
    assert not s["chat"]["auto_continue"] and not s["last_generation"]
    assert s == (await api.get(f"{BASE}/{s['id']}")).json()
    report = (await api.get(f"{BASE}/{s['id']}/global-report")).json()
    assert report["fields"][1]["coverage"]["reason"].startswith("Richiesta di conferma")
    assert report["fields"][1]["source_evidence"] == []
    request = {"version": s["version"], "candidate_ids": ["t0.r1.c1"],
               "allow_unresolved": True}
    response = await api.post(f"{BASE}/{s['id']}/global-draft", json=request)
    assert response.status_code == 422
    response = await api.post(f"{BASE}/{s['id']}/global-draft", json={
        **request, "candidate_ids": ["t0.r0.c1"]})
    assert response.status_code == 200, response.text
    s = response.json()
    generated = (await api.get(s["last_generation"]["downloads"]["report"])).json()
    assert generated["global_audit"]["coverage_summary"] == summary
    assert generated["fields"][0]["provenance"] == "SOURCE"
    doc = Document(BytesIO((await api.get(s["last_generation"]["downloads"]["docx"])).content))
    assert doc.tables[0].cell(0, 1).text == "Aurora Servizi S.r.l."
    assert doc.tables[0].cell(1, 1).text == ""


@pytest.mark.anyio
async def test_unexplained_positions_have_a_backend_notice_even_if_narrator_omits_it(
    api, provider, monkeypatch,
):
    async def request(client, settings, body, timeout, **kwargs):
        data = json.loads(body["messages"][1]["content"])
        p = proposal(data)
        del p["fields"]["t0.r1.c1"]
        p["coverage"] = []
        return json.dumps(p), "simulated", {}

    monkeypatch.setattr(adapter, "request_model_content", request)
    s = await resolve(api, await start(api))
    assert s["global_compilation"]["coverage_summary"]["unexplained"] == 1
    assert "manca ancora un esito esplicito" in s["global_compilation"]["dialogue"]["message"]
    assert s["global_compilation"]["proposal_count"] == 1
