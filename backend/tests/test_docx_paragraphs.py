import random
from io import BytesIO
from pathlib import Path
from zipfile import ZipFile

import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from lxml import etree

from app import docx_templates
from app.docx_templates import (
    DocumentInputError,
    fill_docx,
    inspect_docx,
    is_signature_target,
)


def save(document):
    data = BytesIO()
    document.save(data)
    return data.getvalue()


def test_paragraph_only_form_keeps_labels_styles_and_other_parts():
    document = Document()
    document.sections[0].header.paragraphs[0].text = "Intestazione: ______"
    document.sections[0].footer.paragraphs[0].text = "Documento originale"
    paragraph = document.add_paragraph()
    paragraph.add_run("Il sottoscritto ").bold = True
    paragraph.add_run("__").italic = True
    paragraph.add_run("___").underline = True
    paragraph.add_run(", nato a ")
    paragraph.add_run("___").bold = True
    paragraph.add_run(", partecipa alla procedura 1234/2026.")
    document.add_paragraph("Societa: {{ragione_sociale}}. Indirizzo: [DA COMPILARE].")
    document.add_paragraph("Il richiedente deve verificare la dichiarazione prestampata.")
    original = save(document)
    layout = inspect_docx(original)
    assert layout.candidate_ids == {"p0.s0", "p0.s1", "p1.s0", "p1.s1"}
    assert layout.catalog == []
    output = fill_docx(
        layout, {"p0.s0": "Luca Ferri", "p0.s1": "Bari", "p1.s0": "Mapi & Figli <S.r.l.>"}
    )
    result = Document(BytesIO(output))
    assert (
        result.paragraphs[1].text
        == "Il sottoscritto Luca Ferri, nato a Bari, partecipa alla procedura 1234/2026."
    )
    assert result.paragraphs[1].runs[0].bold is True
    assert result.paragraphs[1].runs[1].italic is True
    assert result.paragraphs[1].runs[2].underline is True
    assert result.paragraphs[1].runs[4].bold is True
    assert result.paragraphs[2].text == "Societa: Mapi & Figli <S.r.l.>. Indirizzo: [DA COMPILARE]."
    assert result.paragraphs[3].text == document.paragraphs[2].text
    with ZipFile(BytesIO(original)) as before, ZipFile(BytesIO(output)) as after:
        for name in before.namelist():
            if name != "word/document.xml":
                assert before.read(name) == after.read(name)
    assert layout.original == original


@pytest.mark.parametrize(
    "marker",
    [
        "___",
        "_ _ _",
        "....",
        ". . . .",
        "\u2026\u2026",
        "\u2026...\u2026..\u2026",
        ".... \u2026\u2026 . . . .",
        "\u2026\u00a0.\u00a0\u2026",
        "{{nome}}",
        "[INSERIRE NOME]",
        "[indicare la sede]",
    ],
)
def test_supported_markers_are_replaced_literally(marker):
    document = Document()
    document.add_paragraph(f"Nome: {marker}; testo originale.")
    result = Document(BytesIO(fill_docx(inspect_docx(save(document)), {"p0.s0": "Luca"})))
    assert result.paragraphs[1].text == "Nome: Luca; testo originale."


@pytest.mark.parametrize("marker", ["(__)", "(_ _)", "( __ )", "(\u00a0_\u00a0_\u00a0)"])
def test_short_parenthesized_fields_survive_split_runs_without_changing_surrounding_text(marker):
    document = Document()
    paragraph = document.add_paragraph()
    for character in f"Comune: ____ {marker}; riferimento AB__12; nota __testo__.":
        paragraph.add_run(character)
    layout = inspect_docx(save(document))
    assert layout.candidate_ids == {"p0.s0", "p0.s1"}
    result = Document(BytesIO(fill_docx(layout, {"p0.s0": "Torino", "p0.s1": "TO"})))
    expected_marker = marker.replace(layout.slots["p0.s1"].placeholder, "TO")
    assert result.paragraphs[1].text == (
        f"Comune: Torino {expected_marker}; riferimento AB__12; nota __testo__."
    )


@pytest.mark.parametrize(
    "text",
    [
        "Proseguire... poi leggere.",
        "Norma [12] e articolo 3",
        "Nome:       ",
        "Testo libero senza campi",
        "Riferimento AB__12, __nome__, segno (_) e __ isolato",
    ],
)
def test_prose_punctuation_and_unmarked_spaces_are_not_fields(text):
    document = Document()
    document.add_paragraph(text)
    with pytest.raises(DocumentInputError, match="Nessun campo supportato"):
        inspect_docx(save(document))


@pytest.mark.parametrize("seed", range(12))
def test_replacements_survive_random_run_boundaries_and_multiple_slots(seed):
    text = "Nome {{nome}}; indirizzo _____; CF ____; luogo ....; note {{note}}."
    rng = random.Random(seed)
    document = Document()
    paragraph = document.add_paragraph()
    start = 0
    while start < len(text):
        size = rng.randint(1, 9)
        paragraph.add_run(text[start : start + size]).bold = bool(start % 2)
        start += size
    values = {
        "p0.s0": "Anna D'Amico",
        "p0.s1": "Via Roma 7\n70100 Bari",
        "p0.s2": "ABC123",
        "p0.s3": "Bari",
        "p0.s4": "A\tB & <C>",
    }
    output = fill_docx(inspect_docx(save(document)), dict(reversed(list(values.items()))))
    result = Document(BytesIO(output))
    assert (
        result.paragraphs[1].text
        == "Nome Anna D'Amico; indirizzo Via Roma 7\n70100 Bari; CF ABC123; "
        "luogo Bari; note A\tB & <C>."
    )


def test_mixed_cells_inline_cell_fields_and_body_fields_are_distinct():
    document = Document()
    document.add_paragraph("Nome: ____")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Societa"
    table.cell(1, 0).text = "Dati"
    table.cell(1, 1).text = "Sede: ____; CF: ____"
    height = OxmlElement("w:trHeight")
    height.set(qn("w:val"), "100")
    height.set(qn("w:hRule"), "exact")
    table.rows[1]._tr.get_or_add_trPr().append(height)
    layout = inspect_docx(save(document))
    assert set(layout.cells) == {"t0.r0.c1"}
    assert set(layout.slots) == {"p0.s0", "p4.s0", "p4.s1"}
    result = Document(
        BytesIO(fill_docx(layout, {"p0.s0": "Luca", "p4.s0": "Bari", "t0.r0.c1": "Mapi"}))
    )
    assert result.tables[0].cell(0, 0).text == "Societa"
    assert result.tables[0].cell(0, 1).text == "Mapi"
    assert result.tables[0].cell(1, 1).text == "Sede: Bari; CF: ____"
    assert result.tables[0].rows[1]._tr.trPr.find(qn("w:trHeight")).get(qn("w:hRule")) == "atLeast"


def test_breaks_and_tabs_do_not_merge_short_unrelated_marks():
    document = Document()
    paragraph = document.add_paragraph("Nome: __")
    paragraph.add_run().add_break()
    paragraph.add_run("__ ")
    paragraph.add_run().add_tab()
    paragraph.add_run("_ Data: ____/____/________")
    layout = inspect_docx(save(document))
    assert len(layout.slots) == 3
    result = Document(BytesIO(fill_docx(layout, {"p0.s0": "15", "p0.s1": "09", "p0.s2": "2000"})))
    assert result.paragraphs[1].text == "Nome: __\n__ \t_ Data: 15/09/2000"


@pytest.mark.parametrize(
    "paragraphs,target",
    [
        (["Firma: ____"], "p0.s0"),
        (["Nome ____; Firma ____"], "p0.s1"),
        (["Firma del rappresentante", "______"], "p1.s0"),
        (["______", "Firma del rappresentante"], "p0.s0"),
        (["________ (firma)"], "p0.s0"),
        (["{{firma}}"], "p0.s0"),
        (["{{firma_digitale}}"], "p0.s0"),
    ],
)
def test_signature_slots_are_never_written_even_when_addressed_directly(paragraphs, target):
    document = Document()
    for text in paragraphs:
        document.add_paragraph(text)
    layout = inspect_docx(save(document))
    assert is_signature_target(layout, target)
    with pytest.raises(DocumentInputError, match="firme"):
        fill_docx(layout, {target: "Luca Ferri"})


def test_signature_label_in_table_applies_to_inline_slots():
    document = Document()
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Firma"
    table.cell(0, 1).text = "____ / ____"
    layout = inspect_docx(save(document))
    assert all(is_signature_target(layout, target) for target in layout.slots)
    with pytest.raises(DocumentInputError, match="firme"):
        fill_docx(layout, {"p1.s0": "Luca Ferri"})


def test_control_contents_and_spanning_fields_stay_untouched_and_are_reported():
    document = Document()
    document.add_paragraph("Nome: ____")
    opening = document.add_paragraph()
    start = OxmlElement("w:fldChar")
    start.set(qn("w:fldCharType"), "begin")
    opening.add_run()._r.append(start)
    protected = document.add_paragraph("Risultato del campo: ____")
    closing = document.add_paragraph()
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    closing.add_run()._r.append(end)
    document.add_paragraph("Sede: ____")
    controlled = document.add_paragraph("Nato a: ____")
    control = OxmlElement("w:sdt")
    content = OxmlElement("w:sdtContent")
    controlled._p.addprevious(control)
    control.append(content)
    content.append(controlled._p)
    before = etree.tostring(control)
    layout = inspect_docx(save(document))
    assert layout.candidate_ids == {"p0.s0", "p4.s0"}
    assert {item["paragraph"] for item in layout.unsupported_locations} == {2, 3, 4, 6}
    output = Document(BytesIO(fill_docx(layout, {"p0.s0": "Luca", "p4.s0": "Bari"})))
    assert output.paragraphs[3].text == protected.text
    assert etree.tostring(output.element.body.find(qn("w:sdt"))) == before


def test_unterminated_field_does_not_expose_following_placeholders():
    document = Document()
    start = OxmlElement("w:fldChar")
    start.set(qn("w:fldCharType"), "begin")
    document.add_paragraph().add_run()._r.append(start)
    document.add_paragraph("Nome ____")
    with pytest.raises(DocumentInputError, match="Nessun campo supportato"):
        inspect_docx(save(document))


def test_total_candidate_limit_includes_paragraph_slots(monkeypatch):
    document = Document()
    document.add_paragraph("Nome ____; Sede ____; CF ____")
    monkeypatch.setattr(docx_templates, "MAX_CANDIDATES", 2)
    with pytest.raises(DocumentInputError, match="troppi campi"):
        inspect_docx(save(document))


def test_writer_rejects_unknown_slots_and_preserves_all_unwritten_markers():
    document = Document()
    document.add_paragraph("Nome ____; Cognome ____")
    layout = inspect_docx(save(document))
    for target in ("p0.s2", "p9.s0", "p0", "t0.r0.c1"):
        with pytest.raises(DocumentInputError, match="non autorizzato"):
            fill_docx(layout, {target: "Luca"})
    result = Document(BytesIO(fill_docx(layout, {"p0.s1": "Ferri"})))
    assert result.paragraphs[1].text == "Nome ____; Cognome Ferri"


def test_unknown_inline_containers_do_not_join_unrelated_text_runs():
    document = Document()
    document.add_paragraph("Nome: ____")
    paragraph = document.add_paragraph("__")
    custom = OxmlElement("w:customXml")
    run = OxmlElement("w:r")
    text = OxmlElement("w:t")
    text.text = "contenuto originale"
    run.append(text)
    custom.append(run)
    paragraph._p.append(custom)
    paragraph.add_run("___")
    original = etree.tostring(paragraph._p)
    layout = inspect_docx(save(document))
    assert layout.candidate_ids == {"p0.s0"}
    assert layout.unsupported_locations[0]["paragraph"] == 2
    output = Document(BytesIO(fill_docx(layout, {"p0.s0": "Luca"})))
    assert etree.tostring(output.paragraphs[2]._p) == original


def test_unknown_container_in_an_empty_cell_is_not_a_whole_cell_target():
    document = Document()
    document.add_paragraph("Nome: ____")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Controllo"
    table.cell(0, 1).paragraphs[0]._p.append(OxmlElement("w:customXml"))
    layout = inspect_docx(save(document))
    assert layout.candidate_ids == {"p0.s0"}
    assert layout.unsupported_locations[0]["paragraph"] == 3


@pytest.mark.parametrize("seed", range(12))
def test_mixed_dotted_area_is_one_field_across_runs_without_eating_abbreviations(seed):
    text = "prov. ....\u2026.. via: ....\u2026 n. . . . .; tel. .... \u2026\u2026"
    doc = Document()
    paragraph = doc.add_paragraph()
    rng = random.Random(seed)
    position = 0
    while position < len(text):
        end = position + rng.randint(1, 8)
        paragraph.add_run(text[position:end]).italic = bool(position % 2)
        position = end
    layout = inspect_docx(save(doc))
    assert len(layout.slots) == 4
    values = dict(zip(layout.slots, ("BA", "Via Roma", "7", "080 123456"), strict=True))
    output = Document(BytesIO(fill_docx(layout, values)))
    assert output.paragraphs[1].text == "prov. BA via: Via Roma n. 7; tel. 080 123456"


def test_dots_do_not_merge_across_labels_newlines_tabs_or_date_separators():
    doc = Document()
    doc.add_paragraph("Nome .... Cognome \u2026\u2026\n....\t\u2026\u2026 Data ..../..../........")
    layout = inspect_docx(save(doc))
    assert len(layout.slots) == 7
    assert all(
        "\t" not in s.placeholder and "\n" not in s.placeholder for s in layout.slots.values()
    )


def test_native_form_regression_merges_contacts_and_preserves_abbreviation_dots():
    path = Path(__file__).resolve().parents[2] / (
        "demo-documents/bandi/minervino-elenco-sia/originali/domanda-iscrizione.docx"
    )
    layout = inspect_docx(path.read_bytes())
    assert len(layout.cells) == 67
    assert len(layout.slots) == 187
    paragraph = next(p for p in layout.paragraph_catalog if p["paragraph"] == 42)
    assert len(paragraph["fields"]) == 6
    assert paragraph["text_with_fields"].startswith("prov. [[p42.s0]] via/piazza ")
    assert "n. [[p42.s2]]" in paragraph["text_with_fields"]
    assert layout.field_types["p42.s4"] == layout.field_types["p42.s5"] == "email"
    assert layout.field_types["p16.s0"] == "email"  # Label ends the preceding paragraph.
    assert len([s for s in layout.slots.values() if s.paragraph_index == 40]) == 1
    assert layout.slots["p15.s1"].placeholder == "__"
    assert layout.slots["p12.s2"].placeholder == "__"
    assert layout.slots["p13.s1"].placeholder == "__"
    values = {
        "p15.s0": "Bari", "p15.s1": "BA",
        "p42.s4": "segreteria@azienda.demo", "p42.s5": "azienda@pec.demo",
    }
    output = Document(BytesIO(fill_docx(layout, values)))
    assert output.paragraphs[16].text.startswith("Bari (BA), tel. ")
    assert "e-mail segreteria@azienda.demo pec azienda@pec.demo e composta" in (
        output.paragraphs[43].text
    )
