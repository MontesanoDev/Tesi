"""V13 S1: independent structural/factual proofs, adversarial records and gates."""

import copy
import json
from io import BytesIO
from pathlib import Path

import pytest
from docx import Document
from test_compilation_binding_identity import DIRECTOR, ORDER, ROW, TABLE, source
from test_compilation_selected_person import NAME, PROFILE, case
from test_compilation_source_coverage import form_case

from app.compilation_global import build_global_input, validate_global_proposal
from app.compilation_global_dialogue import export_selection
from app.compilation_global_sessions import review_fields
from app.compilation_source_reuse import reuse_selected_person
from app.docx_templates import inspect_docx

FIXTURES = Path(__file__).parent / 'fixtures/v13-s1'


def grounded(field):
    return field.get('validation') == 'GROUNDED_PROPOSAL' or field.get(
        'proposal_grounding') == 'GROUNDED_PROPOSAL'


def run(data, layout, sources):
    original = copy.deepcopy(data)
    audit = reuse_selected_person({**data, 'backend_revision': 13}, layout, 'alpha', sources)
    assert data == original
    assert audit['conditions'] == data['audit']['conditions']
    assert all(f['written_value'] is None for f in audit['fields'])
    return audit


def choice(text):
    data, layout, sources = case()
    data['dialogue'] = {'answers': {}, 'user_values': {}}
    data['user_context'] = [{'kind': 'recorded_request', 'provenance': 'USER',
                             'turn_id': 1, 'user_quote': text}]
    return data, layout, sources


@pytest.mark.parametrize('text', [
    f'Il firmatario scelto per questa domanda è {NAME}.',
    f'Il firmatario scelto per la domanda è {NAME}.',
    f'Il firmatario scelto per questa pratica è {NAME}.',
    f'Il firmatario scelto per la pratica è {NAME}.',
    f'Il firmatario scelto per questa procedura è {NAME}.',
    f'Il firmatario scelto per la procedura è {NAME}.',
    f'Il firmatario designato per questa domanda è {NAME}.',
    f'Il firmatario designato per questa pratica sarà {NAME}.',
    f'Il firmatario scelto per questa domanda: {NAME}.',
    f'Per questa pratica il firmatario scelto è {NAME}.',
    f'Compila una bozza. Il firmatario scelto per questa domanda è {NAME}, dati dalle SOURCE.',
    f'Confermo; il firmatario scelto per questa domanda è {NAME}.',
])
def test_explicit_scoped_selection_recovers_source_not_user_facts(text):
    data, layout, sources = choice(text)
    result = run(data, layout, sources)
    reused = [f for f in result['fields'] if f.get('source_reuse')]
    assert len(reused) == 9
    assert all(f['source_evidence'] and f['source_reuse']['identity_evidence'] for f in reused)
    assert data['dialogue']['user_values'] == {}
    assert all(f['source_reuse']['selection_receipt']['provenance'] == 'USER' for f in reused)


@pytest.mark.parametrize('mutation', [
    'negated', 'hypothetical', 'question', 'alternative', 'quotation', 'source_narration',
    'role_only', 'surname_only', 'other_person', 'MODEL_context', 'assistant_context',
    'no_turn', 'string_turn', 'wrong_kind', 'SKIP', 'UNKNOWN', 'REFUSE', 'two_names',
])
def test_selection_does_not_manufacture_identity_or_authorization(mutation):
    sentence = f'Il firmatario scelto per questa domanda è {NAME}.'
    texts = {'negated': 'Non confermo: ' + sentence, 'hypothetical': 'Forse ' + sentence,
             'question': sentence[:-1] + '?', 'alternative': sentence[:-1] + ' oppure Bea Mori.',
             'quotation': 'Nel documento: "' + sentence + '"',
             'source_narration': 'La fonte dice che ' + sentence,
             'role_only': f'{NAME} è rappresentante legale.',
             'surname_only': sentence.replace(NAME, 'Neri'),
             'other_person': sentence.replace(NAME, 'Bea Mori')}
    data, layout, sources = choice(texts.get(mutation, sentence))
    context = data['user_context'][0]
    if mutation in {'MODEL_context', 'assistant_context'}:
        context['provenance'] = mutation.split('_')[0]
    elif mutation == 'no_turn':
        context.pop('turn_id')
    elif mutation == 'string_turn':
        context['turn_id'] = '1'
    elif mutation == 'wrong_kind':
        context['kind'] = 'source_fact'
    elif mutation in {'SKIP', 'UNKNOWN', 'REFUSE'}:
        data['dialogue']['answers']['field:p0.s0'] = {'action': mutation, 'provenance': 'USER'}
    elif mutation == 'two_names':
        data['user_context'].append({**context, 'turn_id': 2,
            'user_quote': sentence.replace(NAME, 'Bea Mori')})
    assert not any(f.get('source_reuse') for f in run(data, layout, sources)['fields'])


def inline_case(label='Il sottoscritto', *, table=False, twice=False):
    doc = Document()
    if table:
        cells = doc.add_table(rows=1, cols=2).rows[0].cells
        cells[0].text = label + ' ________'
        if twice:
            cells[1].text = label + ' ________'
    else:
        doc.add_paragraph(label + ' ________')
        if twice:
            doc.add_paragraph(label + ' ________')
    stream = BytesIO()
    doc.save(stream)
    layout = inspect_docx(stream.getvalue())
    sources = [source(PROFILE, -7)]
    payload = {'version': 12, 'facts': {}, 'fields': {}, 'bindings': {}, 'conditions': {},
               'choice_groups': {}, 'clarifications': {}, 'coverage': []}
    data = {'audit': validate_global_proposal(json.dumps(payload), layout, 'alpha', sources),
            'catalog': build_global_input(layout, 'alpha', sources)['form'],
            'writable_ids': list(layout.slots) + list(layout.cells),
            'user_context': [{'kind': 'recorded_request', 'provenance': 'USER', 'turn_id': 1,
                             'user_quote': f'Il firmatario scelto per questa domanda è {NAME}.'}],
            'dialogue': {'answers': {}, 'user_values': {}}}
    return data, layout, sources


@pytest.mark.parametrize('label,table', [
    ('Il sottoscritto', False), ('La sottoscritta', False), ('Firmatario', False),
    ('Sottoscrittore', False), ('Il sottoscritto', True), ('La sottoscritta', True),
    ('Firmatario', True), ('Sottoscrittore', True),
])
def test_missing_model_name_uses_its_own_paragraph_not_empty_companion(label, table):
    data, layout, sources = inline_case(label, table=table)
    result = run(data, layout, sources)
    recovered = [f for f in result['fields'] if f.get('source_reuse')]
    assert len(recovered) == 1 and recovered[0]['proposed_value'] == NAME
    assert recovered[0]['candidate_id'] in layout.slots
    assert recovered[0]['form_evidence']['paragraph_id'] == data['catalog']['candidates'][
        recovered[0]['candidate_id']]['context_id']


@pytest.mark.parametrize('label,table,twice', [
    ('Socio', False, False), ('Direttore tecnico', False, False),
    ('Firma del sottoscritto', False, False), ('Il sottoscritto se autorizzato', False, False),
    ('Società sottoscritta', True, False), ('Nome e cognome del socio', True, False),
    ('Il sottoscritto', True, True), ('Il sottoscritto', False, True),
])
def test_name_selection_cannot_choose_another_role_signature_or_ambiguous_block(
    label, table, twice,
):
    data, layout, sources = inline_case(label, table=table, twice=twice)
    assert not any(f.get('source_reuse') for f in run(data, layout, sources)['fields'])


ROLE_PROPERTIES = [
    ('nome_completo', 'Nome e cognome', NAME),
    ('nome', 'Nome e cognome', NAME),
    ('nome_e_cognome', 'Nome e cognome', NAME),
    ('qualifica_professionale', 'Qualifica professionale', 'Ingegnere'),
    ('ordine_professionale', 'Iscritto all’Ordine professionale', ORDER),
    ('sezione_ordine', 'Albo professionale, sezione', 'A'),
    ('sezione_albo', 'Albo professionale, sezione', 'A'),
    ('sezione', 'Albo professionale, sezione', 'A'),
    ('provincia_ordine', 'Ordine professionale della provincia di', 'Torino (TO)'),
    ('provincia_albo', 'Ordine professionale della provincia di', 'Torino (TO)'),
    ('numero_iscrizione_ordine', 'Albo professionale con il n.', '6234'),
    ('numero_iscrizione_albo', 'Albo professionale con il n.', '6234'),
    ('numero_iscrizione', 'Albo professionale con il n.', '6234'),
    ('codice_fiscale', 'Codice fiscale', 'NREADA85A41L219R'),
    ('data_iscrizione_albo', 'Iscritto all’Ordine professionale dal', '09/02/2017'),
    ('iscrizione_dal', 'Iscritto all’Ordine professionale dal', '09/02/2017'),
]


@pytest.mark.parametrize('where', ['body_heading', 'cell_heading', 'body_inline'])
def test_inline_table_name_never_bypasses_an_unconfirmed_conditional_heading(where):
    doc = Document()
    if where == 'body_heading':
        doc.add_paragraph('Se procuratore:')
    if where == 'body_inline':
        doc.add_paragraph('Se procuratore: il sottoscritto ________')
    else:
        cell = doc.add_table(rows=1, cols=2).cell(0, 0)
        if where == 'cell_heading':
            cell.text = 'Se procuratore:'
            cell.add_paragraph('Il sottoscritto ________')
        else:
            cell.text = 'Il sottoscritto ________'
    stream = BytesIO()
    doc.save(stream)
    layout = inspect_docx(stream.getvalue())
    sources = [source(PROFILE, -7)]
    payload = {'version': 12, 'facts': {}, 'fields': {}, 'bindings': {}, 'conditions': {},
               'choice_groups': {}, 'coverage': [], 'clarifications': {}}
    audit = validate_global_proposal(json.dumps(payload), layout, 'alpha', sources)
    data = {'audit': audit, 'catalog': build_global_input(layout, 'alpha', sources)['form'],
            'writable_ids': list(layout.slots), 'dialogue': {}, 'user_context': [{
                'kind': 'recorded_request', 'provenance': 'USER', 'turn_id': 1,
                'user_quote': f'Il firmatario scelto per questa domanda è {NAME}.'}]}
    assert not any(f.get('source_reuse') for f in run(data, layout, sources)['fields'])


def role_case(prop='qualifica_professionale', label='Qualifica professionale', value='Ingegnere'):
    doc = Document()
    doc.add_paragraph('Direttore tecnico')
    doc.add_paragraph(label + ': ________')
    stream = BytesIO()
    doc.save(stream)
    layout = inspect_docx(stream.getvalue())
    sources = [source(TABLE, -7), source(DIRECTOR, -9)]
    payload = {'version': 12, 'uncertain_paragraph_ids': [], 'choice_groups': {},
        'coverage': [], 'clarifications': {}, 'conditions': {},
        'bindings': {'b': {'form': {'paragraph_id': 'paragraph:0', 'quote': 'Direttore tecnico'},
            'required_role': 'Direttore tecnico', 'source_role': 'Direttore tecnico',
            'evidence': [{'chunk_id': -9, 'quote': DIRECTOR.splitlines()[0]}]}},
        'facts': {'f': {'entity': 'person', 'subject': NAME, 'aspect': 'professional',
            'property': prop, 'value': value, 'evidence': [{'chunk_id': -7, 'quote': ROW}]}},
        'fields': {'p1.s0': {'label': label, 'kind': 'data', 'disposition': 'value',
            'conditions': [], 'property': prop, 'fact_id': 'f', 'binding_id': 'b',
            'form': {'paragraph_id': 'paragraph:1', 'quote': label + ': [[p1.s0]]'}}}}
    return layout, sources, payload


def role_audit(case_data):
    layout, sources, payload = case_data
    audit = validate_global_proposal(json.dumps(payload), layout, 'alpha', sources)
    data = {'audit': audit, 'catalog': build_global_input(layout, 'alpha', sources)['form'],
            'writable_ids': list(layout.slots), 'dialogue': {}}
    return run(data, layout, sources)


@pytest.mark.parametrize('prop,label,value', ROLE_PROPERTIES)
def test_typed_property_from_same_verified_registration_preserves_all_provenance(
    prop, label, value,
):
    result = role_audit(role_case(prop, label, value))
    field = next(f for f in result['fields'] if f['candidate_id'] == 'p1.s0')
    assert grounded(field) and field.get('proposed_value', field.get('value')) == value
    assert field['subject_binding']['result'] == 'literal_match'
    assert field['subject_binding']['identity_support']['key'] == {
        'order': ORDER.casefold(), 'number': '6234'}
    assert field['source_evidence'][0]['quote'] == ROW
    assert result['facts']['f']['value'] == value


@pytest.mark.parametrize('mutation', [
    'other_person', 'homonym_CF', 'other_number', 'other_order', 'role_missing_number',
    'role_missing_order', 'role_revoked', 'role_negated', 'role_future', 'role_other',
    'role_quote_fabricated', 'FORM_quote_fabricated', 'FORM_wrong_paragraph', 'FORM_wrong_role',
    'SOURCE_quote_fabricated', 'SOURCE_general_role', 'SOURCE_general_value', 'SOURCE_form',
    'entity_company', 'wrong_qualification_column', 'duplicate_role', 'duplicate_person_row',
    'second_person_same_CF', 'second_person_same_registration', 'conflicting_extra_record',
    'missing_value', 'unrelated_property', 'missing_binding', 'reversed_role', 'multiple_CF_values',
])
def test_cross_record_recovery_refuses_wrong_people_records_scopes_and_properties(mutation):
    data = role_case()
    _, sources, payload = data
    fact, binding = payload['facts']['f'], payload['bindings']['b']
    if mutation == 'other_person':
        fact['subject'] = 'Bea Mori'
    elif mutation == 'homonym_CF':
        sources[-1]['content'] += '\n- Codice fiscale: NREADA86A41L219R.'
    elif mutation == 'other_number':
        sources[-1]['content'] = DIRECTOR.replace('6234', '9123')
    elif mutation == 'other_order':
        sources[-1]['content'] = DIRECTOR.replace(ORDER, 'Ordine degli Architetti di Torino')
    elif mutation in {'role_missing_number', 'role_missing_order'}:
        needle = 'Numero iscrizione:' if mutation.endswith('number') else 'Ordine professionale:'
        sources[-1]['content'] = '\n'.join(x for x in DIRECTOR.splitlines() if needle not in x)
    elif mutation in {'role_revoked', 'role_negated', 'role_future'}:
        word = {'role_revoked': 'revocato', 'role_negated': 'non',
                'role_future': 'futuro'}[mutation]
        sources[-1]['content'] = DIRECTOR.replace('Direttore tecnico', word + ' Direttore tecnico')
        binding['evidence'][0]['quote'] = sources[-1]['content'].splitlines()[0]
    elif mutation in {'role_other', 'reversed_role'}:
        binding['source_role' if mutation == 'role_other' else 'required_role'] = 'Socio'
    elif mutation == 'role_quote_fabricated':
        binding['evidence'][0]['quote'] = '### Ada Neri - falso ruolo'
    elif mutation == 'FORM_quote_fabricated':
        binding['form']['quote'] = 'Firmatario'
    elif mutation == 'FORM_wrong_paragraph':
        binding['form']['paragraph_id'] = 'paragraph:1'
    elif mutation == 'FORM_wrong_role':
        binding['required_role'] = binding['source_role'] = 'Socio'
    elif mutation == 'SOURCE_quote_fabricated':
        fact['evidence'][0]['quote'] = 'Citazione inventata'
    elif mutation in {'SOURCE_general_role', 'SOURCE_general_value'}:
        sources[-1 if mutation.endswith('role') else 0]['category'] = 'general'
    elif mutation == 'SOURCE_form':
        sources[0]['role'] = 'form'
    elif mutation == 'entity_company':
        fact['entity'] = 'company'
    elif mutation == 'wrong_qualification_column':
        fact['property'] = 'sezione_ordine'
    elif mutation == 'duplicate_role':
        sources[-1]['content'] += '\n' + DIRECTOR
    elif mutation == 'duplicate_person_row':
        sources[0]['content'] += '\n' + ROW.replace('6234', '9123')
    elif mutation in {'second_person_same_CF', 'second_person_same_registration'}:
        other = ROW.replace(NAME, 'Bea Mori')
        if mutation.endswith('CF'):
            other = other.replace('6234', '9123')
        else:
            other = other.replace('NREADA85A41L219R', 'MROBEA86A41L219R')
        sources[0]['content'] += '\n' + other
    elif mutation == 'conflicting_extra_record':
        sources.append(source(DIRECTOR.replace('6234', '9123'), -11))
    elif mutation == 'missing_value':
        fact['value'] = 'Dato non presente'
    elif mutation == 'unrelated_property':
        fact['property'] = 'fatturato'
    elif mutation == 'missing_binding':
        payload['fields']['p1.s0'].pop('binding_id')
    elif mutation == 'multiple_CF_values':
        sources[-1]['content'] += ('\n- Codice fiscale: NREADA85A41L219R.'
                                  '\n- Codice fiscale: NREADA86A41L219R.')
    if mutation == 'SOURCE_form':
        with pytest.raises(ValueError, match='Ruolo SOURCE non valido'):
            role_audit(data)
        return
    result = role_audit(data)
    field = next(f for f in result['fields'] if f['candidate_id'] == 'p1.s0')
    assert not grounded(field)


@pytest.mark.parametrize('condition,receipt', [
    ('unknown', None), ('false', None), ('unknown', 'UNKNOWN'),
    ('unknown', 'SKIP'), ('unknown', 'REFUSE'), ('unknown', 'CONDITION_FALSE'),
])
def test_source_recovery_cannot_export_false_unknown_or_deferred_conditions(condition, receipt):
    data, layout, sources = choice(f'Il firmatario scelto per questa domanda è {NAME}.')
    data['audit']['conditions']['c'] = {'result': condition, 'validation_errors': [],
                                      'validation': 'GROUNDED_PROPOSAL'}
    for f in data['audit']['fields']:
        f['conditions'] = ['c']
        f['condition_checks'] = {'c': {'result': condition}}
        if condition == 'false':
            f.update(validation='REVIEW_REQUIRED', value=None)
            f.pop('proposal_grounding', None)
    if receipt:
        data['dialogue']['answers']['condition:c'] = {'action': receipt, 'provenance': 'USER',
                                                    'applies': False if receipt.endswith('FALSE')
                                                    else None, 'user_quote': 'Rinvio'}
    result = run(data, layout, sources)
    data.update(audit=result, fields=review_fields({**data, 'audit': result}))
    data['dialogue'].update(phase='READY', reviewed_candidate_ids=data['writable_ids'])
    assert export_selection({'global_compilation': data})[0] == []


@pytest.mark.parametrize('mutation', ['valid', 'unknown_owner', 'owner_false', 'wrong_company',
    'owner_not_name', 'missing_owner_fact', 'unproved_owner', 'SOURCE_missing',
    'wrong_quote', 'wrong_marker', 'boundary', 'wrong_subject_type'])
def test_conditional_operator_continuity_requires_literal_same_company_and_destination(mutation):
    layout, sources, payload, cid = form_case(split=True)
    owner_id = 'p0.s0'
    audit = validate_global_proposal(json.dumps(payload), layout, 'alpha', sources)
    fields = {f['candidate_id']: f for f in audit['fields']}
    owner = fields[owner_id]
    if mutation in {'valid', 'unknown_owner', 'owner_false'}:
        fid = owner['fact_id']
        owner.update(validation='REVIEW_REQUIRED', proposal_grounding='GROUNDED_PROPOSAL',
                     proposed_fact_id=fid, proposed_value=owner['value'], fact_id=None, value=None)
        owner['condition_checks'] = {'c': {'result': 'false' if mutation == 'owner_false'
                                          else 'unknown'}}
    elif mutation == 'wrong_company':
        audit['facts'][owner['fact_id']]['subject'] = 'Altra società'
    elif mutation == 'owner_not_name':
        audit['facts'][owner['fact_id']]['value'] = 'Torino'
    elif mutation == 'missing_owner_fact':
        audit['facts'].pop(owner['fact_id'])
    elif mutation == 'unproved_owner':
        owner.update(validation='REJECTED', proposal_grounding=None)
    elif mutation == 'SOURCE_missing':
        sources[0]['content'] = 'Nessun dato documentato'
    elif mutation == 'wrong_quote':
        fields[cid]['validation_errors'] = ['Estratto non presente nel FORM']
    elif mutation == 'wrong_marker':
        fields[cid]['destination_check'] = {'recovery': {'method': 'unproved'}}
    elif mutation == 'boundary':
        fields[cid]['validation_errors'].append('Cambio soggetto')
    elif mutation == 'wrong_subject_type':
        audit['facts'][payload['fields'][cid]['fact_id']]['entity'] = 'person'
    # Reproduce the final V12 gate independently of the MODEL condition syntax.
    if mutation in {'valid', 'unknown_owner', 'owner_false'}:
        fields[cid].update(validation='REJECTED', validation_errors=[
            'Continuità FORM senza operatore verificato'], proposal_grounding=None)
    data = {'audit': audit, 'catalog': build_global_input(layout, 'alpha', sources)['form'],
            'writable_ids': list(layout.slots), 'dialogue': {}}
    result = run(data, layout, sources)
    field = next(f for f in result['fields'] if f['candidate_id'] == cid)
    if mutation in {'valid', 'unknown_owner'}:
        assert grounded(field)
    elif mutation == 'owner_false':
        assert not grounded(field)
    elif grounded(fields[cid]):
        assert field == fields[cid]  # An already-valid V12 field is never rewritten.
    else:
        assert not grounded(field)


def frozen_case(slug):
    folder = FIXTURES / slug
    ctx = json.loads((folder / 'session-context.json').read_text())
    sources = json.loads((FIXTURES / 'SOURCE.json').read_text())
    layout = inspect_docx((folder / 'FORM-originale.docx').read_bytes())
    audit = validate_global_proposal((folder / 'risposta-globale.txt').read_text(),
                                    layout, ctx['project_id'], sources, expected_version=12)
    data = {**ctx, 'catalog': build_global_input(layout, ctx['project_id'], sources)['form'],
            'audit': audit}
    return data, layout, sources


@pytest.mark.parametrize('cid', ['p12.s0', 'p12.s1', 'p12.s2', 'p12.s3', 'p13.s0', 'p13.s1',
    'p13.s2', 'p13.s3', 'p14.s0', 'p15.s0', 'p51.s0', 'p52.s0', 'p53.s0', 'p53.s1',
    'p53.s2', 'p53.s4', 'craa:p5.s0'])
def test_each_original_factual_omission_is_independently_validable_with_frozen_MODEL_USER(cid):
    slug, cid = cid.split(':') if ':' in cid else ('minervino', cid)
    data, layout, sources = frozen_case(slug)
    old = reuse_selected_person(data, layout, data['project_id'], sources)
    assert not grounded(next(f for f in old['fields'] if f['candidate_id'] == cid))
    result = reuse_selected_person({**data, 'backend_revision': 13}, layout,
                                   data['project_id'], sources)
    field = next(f for f in result['fields'] if f['candidate_id'] == cid)
    assert grounded(field) and field['source_evidence'] and field['written_value'] is None
    assert result['conditions'] == old['conditions']
    assert result['facts'] == old['facts'] and result['bindings'] == old['bindings']


@pytest.mark.parametrize('number', [3, 4, 5, 6, 10, 11, 12])
def test_all_seven_original_Cesano_company_to_person_rejections_survive(number):
    data, layout, sources = frozen_case('cesano')
    result = reuse_selected_person({**data, 'backend_revision': 13}, layout,
                                   data['project_id'], sources)
    field = next(f for f in result['fields'] if f['candidate_id'] == 'p12.ft' + str(number))
    assert not grounded(field) and field['validation'] == 'REJECTED'
    assert 'source_subject_type_mismatch' in field['source_target_check']['validation_codes']


@pytest.mark.parametrize('role', ['Socio', 'Associato', 'Amministratore',
    'Amministratore unico', 'Legale rappresentante', 'Rappresentante legale',
    'Procuratore', 'Vice direttore tecnico', 'Firmatario', 'Professionista designato'])
def test_new_subject_in_same_paragraph_cannot_inherit_an_earlier_director_role(role):
    result = role_audit(role_case('nome_completo', role + ', Nome e cognome', NAME))
    field = next(f for f in result['fields'] if f['candidate_id'] == 'p1.s0')
    assert not grounded(field)


@pytest.mark.parametrize('topology', ['same_document_chunks', 'same_chunk_sections',
                                     'named_role_record', 'authorized_project_sources'])
def test_same_identity_keys_support_distinct_source_record_topologies(topology):
    data = role_case()
    _, sources, payload = data
    if topology == 'same_document_chunks':
        sources[-1]['file_id'] = sources[0]['file_id']
        sources[-1]['chunk_index'] = 1
    elif topology == 'same_chunk_sections':
        sources[0]['content'] += '\n' + DIRECTOR
        sources.pop()
        payload['bindings']['b']['evidence'][0]['chunk_id'] = -7
    elif topology == 'named_role_record':
        text = (f'- **{NAME}:** Direttore tecnico di Aurora Servizi S.r.l.; '
                f'ingegnere, {ORDER}, sezione A, numero 6234.')
        sources[-1]['content'] = text
        payload['bindings']['b']['evidence'][0]['quote'] = text
    else:
        for s in sources:
            s.update(scope='project:alpha', project_id='alpha', category=None,
                     chunk_id=-s['chunk_id'], file_id=-s['file_id'])
        payload['facts']['f']['evidence'][0]['chunk_id'] = 7
        payload['bindings']['b']['evidence'][0]['chunk_id'] = 9
    field = next(f for f in role_audit(data)['fields'] if f['candidate_id'] == 'p1.s0')
    assert grounded(field) and field['subject_binding']['identity_support']


@pytest.mark.parametrize('slug', ['minervino', 'cesano', 'craa'])
def test_historical_state_without_new_revision_remains_exact_and_recovery_is_idempotent(slug):
    data, layout, sources = frozen_case(slug)
    first = reuse_selected_person(data, layout, data['project_id'], sources)
    assert 'factual_recovery' not in first
    assert reuse_selected_person({**data, 'audit': first}, layout,
                                 data['project_id'], sources) == first
    new = reuse_selected_person({**data, 'backend_revision': 13}, layout,
                               data['project_id'], sources)
    assert reuse_selected_person({**data, 'audit': new, 'backend_revision': 13}, layout,
                                 data['project_id'], sources) == new
