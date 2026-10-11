"""Build the authored compilation fixture; no provider, database or application writes.

Expected semantics are authored below, before model evaluation. XML references
only bind those expectations to physical positions; they do not infer answers.
"""

from __future__ import annotations

import argparse
import hashlib
import json
from datetime import UTC, datetime
from io import BytesIO
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile, ZipInfo

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from lxml import etree

ROOT = Path(__file__).resolve().parents[2]
FIXTURE = ROOT / "demo-documents/benchmark-compilazione/v1"
PROJECT = "benchmark-compilazione-v1"
W14 = "http://schemas.microsoft.com/office/word/2010/wordml"


def encoded(value):
    return (json.dumps(value, ensure_ascii=False, indent=2) + "\n").encode()


def stable_docx(document):
    document.core_properties.created = datetime(2026, 1, 1, tzinfo=UTC)
    document.core_properties.modified = datetime(2026, 1, 1, tzinfo=UTC)
    original, result = BytesIO(), BytesIO()
    document.save(original)
    with ZipFile(original) as source, ZipFile(result, "w", ZIP_DEFLATED) as target:
        for name in sorted(source.namelist()):
            info = ZipInfo(name, date_time=(2026, 1, 1, 0, 0, 0))
            info.compress_type = ZIP_DEFLATED
            target.writestr(info, source.read(name))
    return result.getvalue()


def build_fixture():
    doc = Document()
    doc.core_properties.title = "Modulo labirinto — benchmark compilazione V1"
    doc.core_properties.author = "Mapi — fixture sintetica"
    doc.add_heading("Modulo labirinto", 0)
    doc.add_paragraph("Documento sintetico dedicato ai test. Non utilizzare per pratiche reali.")
    doc.add_paragraph(
        "Compilare solo i rami pertinenti. I dati di persone, operatore economico e "
        "raggruppamento hanno soggetti distinti. Una casella vuota non è una risposta negativa."
    )
    doc.sections[0].header.paragraphs[0].text = (
        "LABORATORIO MAPI — le dichiarazioni prestampate non attestano fatti aziendali."
    )
    doc.sections[0].footer.paragraphs[0].text = "Fixture V1 — bozza da revisionare"
    records, sections = [], {}

    def section(key, title, *, level=None, parent=None, condition=None):
        p = doc.add_heading(title, level) if level else doc.add_paragraph()
        if not level:
            p.add_run(title).bold = True
        sections[key] = {"title": title, "parent": parent, "condition": condition,
                         "heading": p._p, "word_heading": bool(level)}

    def remember(key, label, target, branch, *, state="RESOLVED", value=None,
                 sources=(), subject="operatore", reason="", tags=(), checkbox=False):
        records.append({"key": key, "label": label, "target": target, "section": branch,
                        "expected_state": state, "expected_value": value,
                        "source_chunk_ids": list(sources), "subject": subject,
                        "reason": reason, "tags": list(tags), "checkbox": checkbox})

    def table(branch, specs, *, parent=None):
        t = (parent or doc).add_table(rows=len(specs), cols=2)
        t.style = "Table Grid"
        for row, (key, label, options) in zip(t.rows, specs, strict=True):
            row.cells[0].text = label
            remember(key, label, row.cells[1]._tc, branch, **options)
        return t

    def checkbox(key, label, branch, *, legacy=False, **options):
        p = doc.add_paragraph()
        if legacy:
            begin = OxmlElement("w:fldChar")
            begin.set(qn("w:fldCharType"), "begin")
            data = etree.SubElement(begin, qn("w:ffData"))
            control = etree.SubElement(data, qn("w:checkBox"))
            etree.SubElement(control, qn("w:checked")).set(qn("w:val"), "0")
            p.add_run()._r.append(begin)
            end = OxmlElement("w:fldChar")
            end.set(qn("w:fldCharType"), "end")
            p.add_run()._r.append(end)
        else:
            sdt = OxmlElement("w:sdt")
            props = etree.SubElement(sdt, qn("w:sdtPr"))
            control = etree.SubElement(props, f"{{{W14}}}checkbox")
            etree.SubElement(control, f"{{{W14}}}checked").set(f"{{{W14}}}val", "0")
            content = etree.SubElement(sdt, qn("w:sdtContent"))
            etree.SubElement(etree.SubElement(content, qn("w:r")), qn("w:t")).text = "☐"
            p._p.append(sdt)
        p.add_run(" " + label)
        remember(key, label, p._p, branch, checkbox=True, **options)

    section("operatore", "1. Dati dell'operatore economico", level=1)
    legal_address = "Via delle Querce 18, 60121 Ancona (AN)"
    company = [
        ("denominazione", "Denominazione sociale", "Orione Tecnica S.r.l."),
        ("forma", "Forma giuridica", "società a responsabilità limitata"),
        ("cf", "Codice fiscale operatore", "01234567897"),
        ("piva", "Partita IVA", "01234567897"),
        ("sede", "Sede legale", legal_address),
        ("telefono", "Telefono", "+39 071 000 1234"),
        ("pec", "PEC", "orione@pec.example.org"),
    ]
    table("operatore", [(k, label, {"value": value, "sources": [-1]})
                        for k, label, value in company])
    table("operatore", [
        ("operativa", "Sede operativa (compilare solo se diversa dalla sede legale)", {
            "state": "NOT_APPLICABLE", "sources": [-1, -3],
            "reason": "Le fonti attestano lo stesso indirizzo per entrambe le sedi.",
            "tags": ["local_condition", "same_address"]}),
        ("corrispondenza", "Recapito per la corrispondenza", {
            "state": "AMBIGUOUS", "sources": [-3],
            "reason": "Due recapiti ammessi senza priorità; nessuna scelta autorizzata."}),
    ])

    section("firma", "2. Il sottoscritto — dati della persona che firma questa domanda", level=1)
    doc.add_paragraph(
        "La carica societaria non identifica da sola il firmatario della presente pratica. "
        "Per gli estremi della procura verificare prima se il sottoscrittore "
        "agisce come procuratore."
    )
    table("firma", [
        ("firmatario", "Il sottoscritto — Nome e cognome", {
            "state": "MISSING", "subject": "firmatario pratica",
            "reason": "La fonte nomina l'amministratore, non designa il firmatario.",
            "tags": ["same_label", "administrator_is_not_signer"]}),
        ("carica", "Ruolo/carica nella presente domanda", {
            "state": "MISSING", "subject": "firmatario pratica"}),
        ("procura", "Estremi procura (solo se il sottoscrittore agisce come procuratore)", {
            "state": "NEEDS_APPLICABILITY", "subject": "firmatario pratica",
            "reason": "Non è noto se il firmatario agisca come procuratore.",
            "tags": ["unknown_condition"]}),
        ("notaio", "Notaio della procura (stessa condizione della voce precedente)", {
            "state": "NEEDS_APPLICABILITY", "subject": "firmatario pratica",
            "tags": ["inherited_condition"]}),
    ])
    for key, label in [("firma_legale", "Il firmatario agisce come legale rappresentante"),
                       ("firma_procura", "Il firmatario agisce come procuratore")]:
        checkbox(key, label, "firma", state="MISSING", subject="firmatario pratica")

    section("tipo", "3. Tipologia dell'operatore richiedente", level=1)
    doc.add_paragraph(
        "Selezionare esattamente una delle cinque tipologie seguenti. In questo modulo "
        "sono alternative per il medesimo operatore; non descrivono soci né altri partecipanti."
    )
    for i, label in enumerate(["Società di ingegneria", "Professionista singolo",
                               "Studio associato", "Società di professionisti",
                               "Consorzio stabile"]):
        checkbox(f"tipo_{i}", label, "tipo", legacy=i % 2 == 0,
                 value=i == 0, sources=[-1], tags=["exclusive_choice"])

    section("rami", "4. Identificazione — compilare esclusivamente il ramo pertinente", level=1)
    section("singolo", "4.a — Da compilare solo in caso di professionista singolo",
            parent="rami", condition="professionista singolo")
    table("singolo", [(key, label, {"state": "NOT_APPLICABLE", "sources": [-1],
                                   "subject": "professionista singolo",
                                   "tags": ["visual_heading", "excluded_branch", "same_label"]})
                      for key, label in [("singolo_nome", "Nome e cognome"),
                                         ("singolo_ordine", "Ordine professionale"),
                                         ("singolo_albo", "Numero iscrizione albo")]])
    section("associato", "4.b — Da compilare solo in caso di studio associato",
            parent="rami", condition="studio associato")
    doc.add_paragraph("Il sottoscritto in questo ramo è il rappresentante dello studio associato.")
    table("associato", [(key, label, {"state": "NOT_APPLICABLE", "sources": [-1],
                                     "subject": "studio associato",
                                     "tags": ["visual_heading", "excluded_branch", "same_label"]})
                        for key, label in [("studio_denom", "Denominazione sociale"),
                                           ("studio_nome", "Il sottoscritto — Nome e cognome"),
                                           ("studio_albo", "Numero iscrizione albo")]])
    section("ingegneria", "4.c — Da compilare solo in caso di società di ingegneria",
            parent="rami", condition="società di ingegneria")
    doc.add_paragraph(
        "I dati societari seguenti riguardano l'operatore. I dati professionali successivi "
        "riguardano il direttore tecnico, che non è automaticamente il sottoscrittore."
    )
    table("ingegneria", [
        ("ing_denom", "Denominazione sociale", {"value": company[0][2], "sources": [-1]}),
        ("ing_forma", "Forma giuridica", {"value": company[1][2], "sources": [-1]}),
    ])
    section("direttore", "Direttore tecnico della società di ingegneria", parent="ingegneria")
    table("direttore", [(key, label, {"value": value, "sources": [-2],
                                     "subject": "direttore tecnico", "tags": ["same_label"]})
                        for key, label, value in [
                            ("dt_nome", "Nome e cognome", "Giulia Serra"),
                            ("dt_qualifica", "Qualifica professionale", "Ingegnere"),
                            ("dt_ordine", "Ordine professionale",
                             "Ordine degli Ingegneri di Ancona"),
                            ("dt_albo", "Numero iscrizione albo", "5739"),
                        ]] + [
        ("dt_data", "Data di abilitazione", {"state": "MISSING", "subject": "direttore tecnico"}),
        ("organigramma", "Riferimento all'organigramma allegato", {"state": "MISSING"}),
    ])
    section("secondo_dt", "Solo se è stato nominato un secondo direttore tecnico",
            parent="ingegneria", condition="nomina di un secondo direttore tecnico")
    outer = doc.add_table(rows=1, cols=1)
    outer.cell(0, 0).text = "Anagrafica del secondo direttore tecnico (tabella annidata)"
    table("secondo_dt", [(key, label, {"state": "NEEDS_APPLICABILITY",
                                       "subject": "secondo direttore tecnico",
                                       "tags": ["nested_table", "unknown_condition"]})
                         for key, label in [("secondo_nome", "Nome e cognome"),
                                            ("secondo_albo", "Numero iscrizione albo")]],
          parent=outer.cell(0, 0))
    section("consorzio", "4.d — Da compilare solo in caso di consorzio stabile",
            parent="rami", condition="consorzio stabile")
    table("consorzio", [(key, label, {"state": "NOT_APPLICABLE", "sources": [-1],
                                     "subject": "consorzio", "tags": ["excluded_branch"]})
                        for key, label in [("consorzio_denom", "Denominazione sociale"),
                                           ("consorzio_nome", "Il sottoscritto — Nome e cognome")]])

    section("partecipazione", "5. Modalità di partecipazione alla presente procedura", level=1)
    doc.add_paragraph(
        "Selezionare una sola modalità. La forma giuridica dell'operatore non determina "
        "se questa domanda sia presentata singolarmente o in raggruppamento."
    )
    for key, label in [("singola", "Partecipazione singola"), ("mandataria", "Mandataria di RTI"),
                       ("mandante", "Mandante di RTI")]:
        checkbox(key, label, "partecipazione", state="MISSING", subject="partecipazione pratica",
                 tags=["legal_type_is_not_participation"])
    table("partecipazione", [
        ("rti_denom", "Denominazione raggruppamento (solo se si partecipa in RTI)", {
            "state": "NEEDS_APPLICABILITY", "subject": "raggruppamento"}),
        ("partner", "Denominazione sociale dell'altro componente (solo in RTI)", {
            "state": "NEEDS_APPLICABILITY", "subject": "altro componente",
            "tags": ["same_label", "different_subject"]}),
    ])

    section("requisiti", "6. Requisiti, registrazioni e allegati", level=1)
    doc.add_paragraph(
        "Il dichiarante attesta di possedere la certificazione ISO 9001. Questa è una "
        "dichiarazione da verificare, non un'attestazione già acquisita."
    )
    table("requisiti", [
        ("polizza", "Numero della polizza professionale vigente", {
            "state": "CONFLICTING", "sources": [-4, -5],
            "reason": "Due attestazioni di pari data indicano numeri incompatibili."}),
        ("iso_numero", "Numero certificazione ISO 9001", {"state": "MISSING"}),
        ("iso_scadenza", "Scadenza certificazione ISO 9001", {"state": "MISSING"}),
    ])
    checkbox("iso_possesso", "Dichiaro il possesso della ISO 9001", "requisiti", state="MISSING",
             tags=["form_is_not_fact"])
    table("requisiti", [(key, label, {"value": value, "sources": [-1]})
                        for key, label, value in [
                            ("cciaa", "Iscrizione CCIAA", "Camera di Commercio delle Marche"),
                            ("rea", "Numero iscrizione REA", "AN-204060"),
                            ("data_cciaa", "Data iscrizione CCIAA", "14/03/2019"),
                        ]])
    table("requisiti", [("procedura", "Codice della presente procedura", {
        "value": "LAB-ORIONE-2026", "sources": [101], "subject": "procedura",
        "tags": ["project_source"]})])

    section("finale", "7. Protocollo, annotazioni e firma", level=1)
    merged = doc.add_table(rows=3, cols=2)
    merged.style = "Table Grid"
    merged.cell(0, 0).merge(merged.cell(0, 1)).text = "Riferimenti della presente domanda"
    for ri, key, label in [(1, "protocollo", "Numero protocollo della domanda"),
                           (2, "data_domanda", "Data di sottoscrizione della domanda")]:
        merged.cell(ri, 0).text = label
        remember(key, label, merged.cell(ri, 1)._tc, "finale", state="MISSING",
                 subject="pratica", tags=["merged_table", "do_not_use_unrelated_number"])
    p = doc.add_paragraph("Firma del sottoscrittore: __________")
    remember("firma_manuale", "Firma del sottoscrittore", p._p, "finale", state="MANUAL",
             subject="firmatario pratica", tags=["signature_protected"])
    table("finale", [("spazio", "Spazio riservato all'ufficio — non compilare", {
        "state": "DO_NOT_FILL", "subject": "ufficio", "tags": ["candidate_is_not_required"]})])

    paragraphs = list(doc.element.body.iter(qn("w:p")))
    paragraph_ids = {p: f"paragraph:{i}" for i, p in enumerate(paragraphs)}
    cells = {}
    for ti, t in enumerate(doc.element.body.iter(qn("w:tbl"))):
        for ri, row in enumerate(t.findall(qn("w:tr"))):
            for ci, cell in enumerate(row.findall(qn("w:tc"))):
                cells[cell] = f"t{ti}.r{ri}.c{ci}"
    for entry in records:
        target = entry.pop("target")
        if entry.pop("checkbox"):
            entry["candidate_id"] = "p" + paragraph_ids[target].split(":")[1] + ".cb0"
            entry["type"] = "checkbox"
        elif target.tag == qn("w:p"):
            entry["candidate_id"] = "p" + paragraph_ids[target].split(":")[1] + ".s0"
            entry["type"] = "text"
        else:
            entry["candidate_id"] = cells[target]
            entry["type"] = "text"
    for entry in sections.values():
        entry["paragraph_id"] = paragraph_ids[entry.pop("heading")]

    source_specs = [
        (-1, -1, "01-profilo-aziendale.txt", "company",
         "Dati interamente fittizi. Denominazione: Orione Tecnica S.r.l.\n"
         "Forma giuridica: società a responsabilità limitata. Tipologia operatore: società di "
         "ingegneria. Codice fiscale operatore: 01234567897. Partita IVA: 01234567897.\n"
         f"Sede legale: {legal_address}. Telefono: +39 071 000 1234. PEC: orione@pec.example.org.\n"
         "Iscrizione CCIAA: Camera di Commercio delle Marche. Numero REA: AN-204060. "
         "Data iscrizione CCIAA: 14/03/2019. Capitale sociale: EUR 50000.\n"
         "Amministratore unico: Paolo Conti. Questo profilo non designa il sottoscrittore "
         "della presente domanda né la modalità di partecipazione alla procedura."),
        (-2, -2, "02-direttore-tecnico.txt", "company",
         "Orione Tecnica S.r.l. — direttore tecnico: Giulia Serra. Qualifica professionale: "
         "Ingegnere. Ordine professionale: Ordine degli Ingegneri di Ancona, n. 5739. "
         "Telefono ufficio: +39 071 000 5678. Identificativo documento: DT-8821. "
         "Il numero 8821 identifica il documento, non l'iscrizione all'albo. "
         "Questa scheda non riporta la data di abilitazione e non contiene un organigramma."),
        (-3, -3, "03-recapiti.txt", "company",
         f"Orione Tecnica S.r.l. — sede operativa: {legal_address}. "
         "La sede operativa coincide con la sede legale. Per la corrispondenza sono "
         "utilizzabili due recapiti, senza preferenza indicata: Via del Porto 2, Ancona; "
         "Via dei Tigli 7, Jesi. Nessuno dei due è indicato come nuova sede legale."),
        (-4, -4, "04-polizza-a.txt", "company",
         "Attestazione del 01/09/2026: Orione Tecnica S.r.l., polizza professionale vigente "
         "RC-OR-2026-A. Nessuna priorità su altre attestazioni della stessa data."),
        (-5, -5, "05-polizza-b.txt", "company",
         "Attestazione del 01/09/2026: Orione Tecnica S.r.l., polizza professionale vigente "
         "RC-OR-2026-B. Nessuna priorità su altre attestazioni della stessa data."),
        (101, 101, "06-procedura.txt", None,
         "La presente procedura ha codice LAB-ORIONE-2026. Importo stimato: EUR 250000. "
         "Il codice della procedura non è il protocollo assegnato alla domanda del partecipante."),
        (-201, -201, "07-conoscenza-generale.txt", "general",
         "Esempio didattico riferito a un soggetto diverso: Vega Restauri S.r.l., "
         "partita IVA 09999999999, referente Marco Riva, certificato ISO 9001 GEN-7788. "
         "Non sono dati dell'operatore che presenta questa domanda."),
    ]
    sources = [{"chunk_id": ci, "file_id": fi, "chunk_index": 0, "source_name": name,
                "role": "source", "scope": "global" if ci < 0 else f"project:{PROJECT}",
                "project_id": None if ci < 0 else PROJECT, "category": cat, "content": text}
               for ci, fi, name, cat, text in source_specs]
    excluded = [{"chunk_id": 901, "file_id": 901, "chunk_index": 0,
                 "source_name": "altro-progetto.txt", "role": "source",
                 "scope": "project:benchmark-estraneo", "project_id": "benchmark-estraneo",
                 "category": None,
                 "content": "Orione Tecnica S.r.l. — data abilitazione Giulia Serra: 12/06/2014."}]
    original = stable_docx(doc)
    gold = {"version": 1, "project_id": PROJECT, "form": "modulo-benchmark.docx",
            "original_sha256": hashlib.sha256(original).hexdigest(),
            "annotation_origin": "Authored fixture semantics, independent of model responses",
            "sections": sections, "fields": records,
            "notes": ["Expected states are evaluation targets, not current resolver states.",
                      "A source-supported value in a conditional branch still "
                      "requires applicability.",
                      "Model output and local grounding are scored separately; neither is a DOCX."]}
    files = {"modulo-benchmark.docx": original, "sources.json": encoded(sources),
             "excluded-sources.json": encoded(excluded), "gold.json": encoded(gold)}
    files.update({"sources/" + s["source_name"]: (s["content"] + "\n").encode() for s in sources})
    files["manifest.json"] = encoded({
        "fixture_version": 1, "project_id": PROJECT, "form": "modulo-benchmark.docx",
        "sources": "sources.json", "gold": "gold.json",
        "files": {name: hashlib.sha256(data).hexdigest() for name, data in sorted(files.items())},
    })
    return files


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output", type=Path, default=FIXTURE)
    parser.add_argument("--check", action="store_true", help="Verify without writing any files")
    args = parser.parse_args()
    files = build_fixture()
    if args.check:
        different = [name for name, data in files.items()
                     if not (args.output / name).is_file()
                     or (args.output / name).read_bytes() != data]
        if different:
            raise SystemExit("Fixture diversa: " + ", ".join(different))
    else:
        for name, data in files.items():
            path = args.output / name
            path.parent.mkdir(parents=True, exist_ok=True)
            path.write_bytes(data)
    gold = json.loads(files["gold.json"])
    print(json.dumps({"fixture": str(args.output), "fields": len(gold["fields"]),
                      "original_sha256": gold["original_sha256"], "check_only": args.check}))


if __name__ == "__main__":
    main()
