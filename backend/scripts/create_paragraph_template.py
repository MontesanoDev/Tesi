"""Create a synthetic DOCX form for testing paragraph and table placeholders."""

import argparse
from pathlib import Path

from docx import Document
from docx.shared import Cm, Pt

ROOT = Path(__file__).resolve().parents[2]


def create_template(path: Path) -> None:
    document = Document()
    section = document.sections[0]
    section.page_height, section.page_width = Cm(29.7), Cm(21)
    section.top_margin = section.bottom_margin = Cm(2)
    document.styles["Normal"].font.name = "Calibri"
    document.styles["Normal"].font.size = Pt(11)
    document.add_heading("Anagrafica del concorrente", 0)
    document.add_paragraph("MODELLO DIMOSTRATIVO - NON VALIDO PER CANDIDATURE REALI")
    paragraph = document.add_paragraph("Il sottoscritto ")
    paragraph.add_run("___").italic = True
    paragraph.add_run("___").underline = True
    paragraph.add_run(", nato a ______ il ___/___/______, in qualita di ______.")
    document.add_paragraph("Societa: {{ragione_sociale}}. Sede legale: [INSERIRE SEDE LEGALE].")
    document.add_paragraph(
        "Codice fiscale dell'impresa: {{codice_fiscale}}. Partita IVA: {{partita_iva}}."
    )
    document.add_heading("Recapiti", 1)
    table = document.add_table(rows=2, cols=2)
    table.style = "Table Grid"
    table.cell(0, 0).text = "Email"
    table.cell(1, 0).text = "PEC"
    document.add_paragraph("Luogo: ........; data: ........")
    document.add_paragraph("Firma")
    document.add_paragraph("________________________")
    section.footer.paragraphs[0].text = "Mapi - esempio tecnico con dati da verificare"
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("xb") as output:
        document.save(output)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--output",
        type=Path,
        default=ROOT / "demo-documents/modelli/modulo-paragrafi.docx",
    )
    args = parser.parse_args()
    if args.output.exists():
        parser.error("Il file esiste gia: scegli un percorso nuovo")
    create_template(args.output)
    print(args.output.resolve())


if __name__ == "__main__":
    main()
