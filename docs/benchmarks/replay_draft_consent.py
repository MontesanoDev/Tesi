"""Read-only live checkpoint -> private snapshot -> real HTTP consent/export."""

import argparse
import asyncio
import hashlib
import json
import os
import shutil
import sqlite3
import tempfile
import time
from io import BytesIO
from pathlib import Path

import httpx
from docx import Document

from app import intents
from app.db import get_db_path, get_storage_path
from app.main import app
from scripts.benchmark_compilation_dialogue import real_settings


def digest(blob):
    return hashlib.sha256(blob).hexdigest()


def checkpoint(path, session_id):
    with sqlite3.connect(path.as_uri() + "?mode=ro", uri=True) as db:
        r = db.execute(
            "select version,state_json,original from compilation_sessions where id=?", (session_id,)
        ).fetchone()
        return {
            "version": r[0],
            "state_sha256": digest(r[1].encode()),
            "original_sha256": digest(r[2]),
        }


async def main(args):
    session_id = args.session_id
    live = get_db_path().resolve()
    storage = get_storage_path().resolve()
    before = checkpoint(live, session_id)
    with sqlite3.connect(live.as_uri() + "?mode=ro", uri=True) as db:
        project = db.execute(
            "select project_id from compilation_sessions where id=?", (session_id,)
        ).fetchone()[0]
    settings = real_settings(project)
    if before["version"] != args.expected_version:
        raise ValueError("La sessione è cambiata: congelare un nuovo checkpoint esplicito")
    application = Path(__file__).resolve().parents[2] / "backend/app"
    frozen_code = {p.name: digest(p.read_bytes()) for p in application.glob("*.py")}
    args.output.mkdir(parents=True, exist_ok=False)

    def save(name, data):
        p = args.output / name
        p.parent.mkdir(parents=True, exist_ok=True)
        p.write_text(
            json.dumps(data, ensure_ascii=False, indent=2, default=str).replace(
                settings.api_key, "[REDACTED]"
            )
            + "\n"
        )

    save(
        "protocol.json",
        {
            "session_id": session_id,
            "expected_version": args.expected_version,
            "message": args.message,
            "transport": "HTTP ASGI",
            "real_provider": settings.provider,
            "model": settings.model,
            "new_analysis": False,
            "live_database_writes": False,
            "max_http_user_requests": 1,
            "max_planner_attempts": 2,
            "application_sha256": frozen_code,
        },
    )
    calls = []
    original = intents.request_model_content

    async def capture(client, config, body, timeout, **kwargs):
        idx = len(calls) + 1
        item = {"request": body, "schema": kwargs.get("response_schema")}
        calls.append(item)
        start = time.perf_counter()
        try:
            raw, model, usage = await original(client, config, body, timeout, **kwargs)
            item.update(raw=raw, model=model, usage=usage)
            return raw, model, usage
        except Exception as exc:
            item.update(error=type(exc).__name__, message=str(exc))
            raise
        finally:
            item["seconds"] = time.perf_counter() - start
            save(f"call-{idx}.json", item)

    intents.request_model_content = capture
    env = {
        k: os.environ.get(k) for k in ("MAPI_DB_PATH", "MAPI_STORAGE_PATH", "MAPI_KNOWLEDGE_PATH")
    }
    try:
        with tempfile.TemporaryDirectory(prefix="mapi-draft-consent-") as tmp:
            private = Path(tmp)
            dbpath = private / "private.sqlite"
            uploads = private / "uploads"
            with (
                sqlite3.connect(live.as_uri() + "?mode=ro", uri=True) as src,
                sqlite3.connect(dbpath) as dest,
            ):
                src.backup(dest)
                dest.row_factory = sqlite3.Row
                row = dest.execute(
                    "select * from compilation_sessions where id=?", (session_id,)
                ).fetchone()
                assert row["version"] == args.expected_version, "Checkpoint changed during copy"
                state = json.loads(row["state_json"])
                from app.compilation_global_sessions import session_payload

                initial = session_payload(row, state)
                save("initial-session.json", initial)
            shutil.copyfile(live.with_suffix(".ai-key"), dbpath.with_suffix(".ai-key"))
            os.chmod(dbpath.with_suffix(".ai-key"), 0o600)
            for relative in (
                state["global_compilation"]["artifact_path"],
                state["original_storage_path"],
            ):
                source = storage / relative
                target = uploads / relative
                target.parent.mkdir(parents=True, exist_ok=True)
                if source.is_dir():
                    shutil.copytree(source, target)
                else:
                    shutil.copyfile(source, target)
            os.environ.update(
                MAPI_DB_PATH=str(dbpath),
                MAPI_STORAGE_PATH=str(uploads),
                MAPI_KNOWLEDGE_PATH=str(private / "knowledge"),
            )
            async with httpx.AsyncClient(
                transport=httpx.ASGITransport(app=app), base_url="http://private", timeout=400
            ) as client:
                payload = {
                    "question": args.message,
                    "document_id": row["form_id"],
                    "conversation_id": row["conversation_id"],
                    "compilation_session_id": session_id,
                    "compilation_version": row["version"],
                }
                save("request.json", payload)
                response = await client.post(
                    f"/api/projects/{row['project_id']}/answer", json=payload
                )
                save("response.json", {"status": response.status_code, "body": response.json()})
                result = (
                    await client.get(
                        f"/api/projects/{row['project_id']}/compilation-sessions/{session_id}"
                    )
                ).json()
                save("final-session.json", result)
                checks = {
                    "http_status": response.status_code,
                    "generation_status": response.json().get("generation_status"),
                    "compilation": response.json().get("compilation"),
                    "status": result["status"],
                    "initial_version": row["version"],
                    "final_version": result["version"],
                    "original_immutable": None,
                    "materialized_values": [],
                    "calls": len(calls),
                }
                if result.get("last_generation"):
                    downloads = {}
                    for kind, url in result["last_generation"]["downloads"].items():
                        r = await client.get(url)
                        assert r.status_code == 200
                        downloads[kind] = r.content
                        (
                            args.output
                            / ("generated-report.json" if kind == "report" else kind + ".docx")
                        ).write_bytes(r.content)
                    report = json.loads(downloads["report"])
                    out = Document(BytesIO(downloads["docx"]))
                    checks["original_immutable"] = downloads["template"] == bytes(row["original"])
                    for field in report["fields"]:
                        pos = field["location"]
                        from docx.oxml.ns import qn
                        from docx.text.paragraph import Paragraph

                        from app.docx_templates import DRAFT_NOTICE

                        paragraphs = [
                            Paragraph(p, out._body).text for p in out.element.body.iter(qn("w:p"))
                        ]
                        assert paragraphs[0] == DRAFT_NOTICE
                        paragraphs = paragraphs[1:]
                        text = (
                            paragraphs[pos["paragraph"] - 1]
                            if pos.get("kind") == "paragraph"
                            else ""
                        )
                        checks["materialized_values"].append(
                            {
                                "candidate_id": field["cell_id"],
                                "label": field["label"],
                                "value": field["written_value"],
                                "provenance": field["provenance"],
                                "source_evidence": field.get("source_evidence"),
                                "present_in_destination": str(field["written_value"]) in text,
                                "destination_text": text,
                            }
                        )
                    checks.update(
                        written_count=report["written_field_count"],
                        all_materialized=all(
                            x["present_in_destination"] for x in checks["materialized_values"]
                        ),
                        source_values_unchanged=result["global_compilation"]["fields"]
                        == initial["global_compilation"]["fields"],
                    )
                if result.get("last_generation"):
                    checks.update(verify_word(downloads, report))
                artifact = uploads / state["global_compilation"]["artifact_path"]
                if (artifact / "routing").is_dir():
                    for trace in (artifact / "routing").glob("*/trace.json"):
                        save(
                            "routing/" + trace.parent.name + "/trace.json",
                            json.loads(trace.read_text()),
                        )
                save("checks.json", checks)
                print(
                    json.dumps(
                        {k: v for k, v in checks.items() if k != "materialized_values"},
                        ensure_ascii=False,
                    ),
                    flush=True,
                )
    finally:
        for k, v in env.items():
            if v is None:
                os.environ.pop(k, None)
            else:
                os.environ[k] = v
        after = checkpoint(live, session_id)
        save(
            "integrity.json",
            {
                "before": before,
                "after": after,
                "live_session_unchanged": before == after,
                "private_directory_removed": not private.exists(),
                "application_unchanged": all(
                    digest((application / name).read_bytes()) == sha
                    for name, sha in frozen_code.items()
                ),
            },
        )


def verify_word(downloads, report):
    """Reconstruct expected text using original slot offsets, not the writer."""
    from collections import defaultdict

    from docx.oxml.ns import qn
    from docx.text.paragraph import Paragraph

    from app.docx_templates import DRAFT_NOTICE, inspect_docx

    original = downloads["template"]
    layout = inspect_docx(original, require_candidates=False)
    expected = [
        Paragraph(p, layout.document._body).text
        for p in layout.document.element.body.iter(qn("w:p"))
    ]
    updates = defaultdict(list)
    unchecked = []
    for field in report["fields"]:
        cid = field["cell_id"]
        if cid not in layout.slots:
            unchecked.append(cid)
            continue
        slot = layout.slots[cid]
        updates[slot.paragraph_index].append((slot, field["written_value"]))
    for index, replacements in updates.items():
        for slot, value in sorted(replacements, key=lambda item: item[0].start, reverse=True):
            assert expected[index][slot.start : slot.end] == slot.placeholder
            expected[index] = expected[index][: slot.start] + value + expected[index][slot.end :]
    doc = Document(BytesIO(downloads["docx"]))
    actual = [Paragraph(p, doc._body).text for p in doc.element.body.iter(qn("w:p"))]
    assert actual[0] == DRAFT_NOTICE
    actual = actual[1:]
    mismatches = [i for i, (a, b) in enumerate(zip(expected, actual, strict=True)) if a != b]
    return {
        "all_text_positions_correct": not mismatches and not unchecked,
        "paragraph_mismatches": mismatches,
        "unchecked_candidates": unchecked,
        "unresolved_placeholders_unchanged": not mismatches,
    }


if __name__ == "__main__":
    p = argparse.ArgumentParser()
    p.add_argument("--output", type=Path, required=True)
    p.add_argument("--message", default="si")
    p.add_argument("--session-id", required=True)
    p.add_argument("--expected-version", type=int, required=True)
    asyncio.run(main(p.parse_args()))
